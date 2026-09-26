"""Astra-4 counterexamples, evaluated from the proposed JSON formulas."""
import contextlib
import copy
import importlib.util
import io
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import unittest
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tests/fixtures/hort_finance'))
sys.path.insert(0, str(ROOT / 'skills/knuaf-doc/scripts'))
from formula_eval import Evaluator, FormulaError, referenced_cells
import gg_hort_workbook as h
from synthetic import build, entries

LANES = ROOT / '.hw-work/fin/lanes'
PARAMS = Path(os.environ.get('HFIN_PARAMS', ROOT / 'skills/knuaf-doc/references/hort-env-systems/finance-params.json'))
PUBLIC = Path(os.environ.get('HFIN_PUBLIC_DIR', ROOT / 'skills/knuaf-doc/references/hort-env-systems'))
WORKBOOK = Path(os.environ['HFIN_H01']) if os.environ.get('HFIN_H01') else None
SHA = 'e3c9defe376fa74413641408900f6bb656017a389fb9745d9dac221a10951c1c'


def load(path):
    return json.loads(path.read_text(encoding='utf-8'))


def actions(params):
    return load(PUBLIC/'finance-transform.json')['entries']


def matches_ref(formula, cell):
    return bool(re.search(r'(?<![A-Z0-9])\$?' + re.escape(re.match(r'[A-Z]+',cell).group()) +
                          r'\$?' + re.escape(re.search(r'\d+',cell).group()) + r'(?!\d)', formula, re.I))


class FormulaContract(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.params = load(PARAMS)
        cls.entries = actions(cls.params)
        cls.formula_entries = [a for a in cls.entries if a.get('after_formula')]

    def evaluate(self, sheet, cell, inputs):
        return Evaluator(self.entries, inputs).cell(sheet, cell)

    def test_T_REQ_1_all_area_unit_slots(self):
        sheet = '12 .경비계획'
        keys = [(k,v) for k,v in self.params['parameters'].items()
                if k.endswith('.coverage_ratio') and v.get('method_tag') == 'area_unit']
        self.assertEqual(len(keys), 25)
        formula_map = {(a['sheet'],a['cell']):a['after_formula'] for a in self.formula_entries}
        checked = 0
        for key, param in keys:
            sh, cell = param['canonical_cell'].split('!')
            row = re.search(r'\d+', cell).group()
            area_cell = 'L'+row
            targets = [c for (s,c),f in formula_map.items()
                       if s == sh and matches_ref(f, area_cell) and matches_ref(f,'M'+row)]
            self.assertTrue(targets, key)
            values = {(sh,'I2'):2027, (sh,'I3'):1000, (sh,'I4'):1000,
                      (sh,'I5'):1000, (sh,'K'+row):'부지',
                      (sh,cell):None, (sh,'M'+row):5,
                      (sh,'N'+row):2027, (sh,'O'+row):0,
                      (sh,'P'+row):None}
            self.assertEqual(self.evaluate(sh,area_cell,values),1000,key)
            for target in targets:
                got = self.evaluate(sh,target,values)
                self.assertAlmostEqual(got,1000*5/1000,places=8,msg=key+' -> '+target)
                checked += 1
        self.assertGreaterEqual(checked,25)

    def test_T_REQ_2_flat_material_quote_without_basis_area(self):
        sheet = '9. 영농자재구매계획'
        keys = [(k,v) for k,v in self.params['parameters'].items() if k.endswith('.basis_m2')]
        self.assertEqual(len(keys),30)
        for key, param in keys:
            sh, cell = param['canonical_cell'].split('!')
            row = int(re.search(r'\d+',cell).group())
            target = next((a['cell'] for a in self.formula_entries if a['sheet']==sh
                           and a['cell'].startswith('E') and matches_ref(a['after_formula'],cell)),None)
            self.assertIsNotNone(target,key)
            values = {(sh,'U'+str(row)):'synthetic item', (sh,'AB'+str(row)):'quote',
                      (sh,'V'+str(row)):'부지', (sh,'W'+str(row)):None,
                      (sh,'X'+str(row)):7, (sh,'Y'+str(row)):1,
                      ('12 .경비계획','I3'):1000}
            self.assertAlmostEqual(self.evaluate(sh,target,values),7,msg=key)

    def _enumerate_optional_key_dependency(self):
        params = self.params['parameters']
        self.assertEqual(len(params),1464)
        by_coord={tuple(v['canonical_cell'].split('!',1)):(k,v) for k,v in params.items()
                  if '!' in (v.get('canonical_cell') or '')}
        self.assertEqual(len(by_coord),1433)
        pairs=0; optional_keys=set(); unused_pairs=0; failures=[]
        for action in self.formula_entries:
            refs=referenced_cells(action['after_formula'],action['sheet'])
            for coord in refs:
                if coord not in by_coord: continue
                key,param=by_coord[coord]
                req=param.get('required_when','').lower()
                if req.startswith('always') or req.startswith('each active plan year'):
                    continue
                if req.startswith('corresponding plan year active') and param.get('explicit_none_calculation')!=0:
                    continue
                if req.startswith('corresponding plan year and month row active') and (
                    param.get('explicit_none_calculation')!=0 and
                    not key.endswith('.price_krw_per_unit')):
                    # Production quantity is required for each active month;
                    # the price key is nullable when there is no sale.
                    continue
                pairs+=1
                optional_keys.add(key)
                values={}
                for sh,cell in refs:
                    other=by_coord.get((sh,cell),(None,{}))[1]
                    unit=other.get('unit_code')
                    if unit in ('YEAR','DATE'): value=2027
                    elif unit=='MONTH': value=1
                    elif cell.startswith(('K','V')) and sh in ('12 .경비계획','9. 영농자재구매계획'):
                        value='부지'
                    elif cell.startswith('U') and sh=='9. 영농자재구매계획': value='synthetic item'
                    elif cell.startswith('AB') and sh=='9. 영농자재구매계획': value='quote'
                    elif unit in ('TEXT','EVIDENCE_REF'): value='synthetic'
                    else: value=1
                    values[(sh,cell)]=value
                # Drive the exact allowed absence branch. A zero-derived fact
                # enters the calculation layer as 0, never as an empty cell.
                if key.startswith('loan.'):
                    for sh,cell in refs:
                        if sh=='5. 원리금상환계획' and cell.startswith('C') and cell[1:].isdigit() and int(cell[1:])>=40:
                            values[(sh,cell)]=0
                if key.startswith('cost.material.'):
                    row=re.search(r'\d+',coord[1]).group()
                    for sh,cell in refs:
                        if sh==coord[0] and cell in ('U'+row,'AB'+row):
                            values[(sh,cell)]='' if cell.startswith('U') else '해당 없음'
                if key.startswith('cost.overhead.'):
                    for sh,cell in refs:
                        if sh=='12 .경비계획' and cell.startswith('P'): values[(sh,cell)]=0
                    if '.direct_quote_krw1000' in key:
                        row=re.search(r'\d+',coord[1]).group()
                        values[(coord[0],'Q'+row)]=''
                        values[(coord[0],'R'+row)]=''
                if key.startswith('cost.labor.') and key.endswith('.hours'):
                    row=re.search(r'\d+',coord[1]).group()
                    for col in ('AD','AE','AF','AG','AH','AI'):
                        values[(coord[0],col+row)]=''
                if key.startswith(('packaging.','transport.')) and key.endswith('_per_sales_unit'):
                    row=re.search(r'\d+',coord[1]).group()
                    price_row=str(int(row)-1)
                    values[(coord[0],coord[1][0]+price_row)]=0
                if key.startswith('sales.') and key.endswith('.price_krw_per_unit'):
                    col=re.match(r'[A-Z]+',coord[1]).group()
                    row=int(re.search(r'\d+',coord[1]).group())
                    values[(coord[0],f'{col}{row-1}')]=0
                    for sh,cell in refs:
                        if sh==coord[0] and cell.startswith(col) and cell!=coord[1]:
                            values[(sh,cell)]=0
                if key.startswith('sales.') and ('stock_qty' in key):
                    for sh,cell in refs:
                        if sh==' 7. 판매계획': values[(sh,cell)]=0
                values[coord]=0 if param.get('explicit_none_calculation')==0 else None
                try:
                    Evaluator([action],values).cell(action['sheet'],action['cell'])
                except FormulaError as exc:
                    failures.append(('allowed branch',key,action['sheet'],action['cell'],str(exc)))
                parts=key.split('.')
                if parts[0] in ('invest','loan') and len(parts)>2:
                    group='.'.join(parts[:2])
                elif parts[:2] in (['cost','overhead'],['cost','labor']) and len(parts)>4:
                    group='.'.join(parts[:4])
                elif parts[:2]==['asset','opening_equipment'] and len(parts)>3:
                    group='.'.join(parts[:3])
                else: group=None
                if group:
                    unused_pairs+=1
                    unused=values.copy()
                    for othercoord,(otherkey,otherparam) in by_coord.items():
                        if otherkey.startswith(group+'.'):
                            unused[othercoord]=(0 if otherparam.get('explicit_none_calculation')==0 else None)
                    try:
                        Evaluator([action],unused).cell(action['sheet'],action['cell'])
                    except FormulaError as exc:
                        failures.append(('unused group',key,action['sheet'],action['cell'],str(exc)))
        return pairs, len(optional_keys), unused_pairs, failures

    def test_T_REQ_3a_optional_branch_enumeration(self):
        pairs, keys, unused, _ = self._enumerate_optional_key_dependency()
        self.assertEqual((pairs, keys, unused), (3082, 1098, 2213))

    def test_T_REQ_3b_optional_branch_formula_evaluation(self):
        _, _, _, failures = self._enumerate_optional_key_dependency()
        with self.assertRaises(FormulaError):
            self.evaluate('12 .경비계획','C16',{
                ('12 .경비계획','P14'):None,('12 .경비계획','Q14'):3,
                ('12 .경비계획','R14'):'synthetic evidence'})
        with self.assertRaises(FormulaError):
            self.evaluate('11. 인건비계획','G9',{
                ('11. 인건비계획','AD10'):None,
                ('11. 인건비계획','AI10'):'synthetic evidence'})
        if failures:
            from collections import Counter
            kinds=Counter((f[0],f[1].split('.')[0],f[2],re.match(r'[A-Z]+',f[3]).group()) for f in failures)
            self.fail('optional branch formula errors: '+repr(failures[:10])+
                      ' (total '+str(len(failures))+', groups '+repr(kinds.most_common(15))+')')

    def test_T_REQ_4_explicit_none_wording_and_state(self):
        for key,param in self.params['parameters'].items():
            wording = ' '.join(str(param.get(f) or '') for f in
                               ('specification','required_when','output_conversion')).lower()
            if 'explicit_none' in wording:
                self.assertIn('explicit_none',param.get('allowed_answer_states',[]),key)
                if 'explicit_none=0' in wording:
                    self.assertEqual(param.get('explicit_none_calculation'),0,key)

    def test_T_ZERO_1_derived_zero_visible_blank(self):
        keys = [(k,v) for k,v in self.params['parameters'].items()
                if v.get('explicit_none_calculation') == 0 and v.get('explicit_none_output') == 'blank']
        self.assertEqual(len(keys),140)
        public=load(PUBLIC/'finance-transform.json')
        mapped={(e['sheet'],e['cell']):e for e in public['entries']}
        import openpyxl
        wb=(openpyxl.load_workbook(WORKBOOK,read_only=True,data_only=False)
            if WORKBOOK is not None and WORKBOOK.is_file() else None)
        try:
            for key,param in keys:
                self.assertEqual(param.get('explicit_none_display_strategy'),
                                 'zero_hidden_number_format',key)
                self.assertEqual(param.get('explicit_none_number_format_policy'),
                                 'preserve_source_nonzero_sections',key)
                coord=tuple(param['canonical_cell'].split('!',1))
                number_format=mapped[coord].get('number_format','')
                parts=number_format.split(';')
                self.assertGreaterEqual(len(parts),3,key)
                self.assertEqual(parts[2],'',key)
                if wb is not None:
                    source_format=wb[coord[0]][coord[1]].number_format
                    if source_format not in (None,'General'):
                        original=source_format.split(';')
                        self.assertEqual(parts[0],original[0],key)
                        if len(original)>1:self.assertEqual(parts[1],original[1],key)
                    else:
                        self.assertEqual(number_format,param['explicit_none_number_format_fallback'],key)
        finally:
            if wb is not None: wb.close()
        for target,price in (('E41','M41'),('E46','N41'),('E17','O41'),
                             ('F17','P41'),('G17','Q41'),('E39','M39'),
                             ('E44','N39'),('E16','O39'),('F16','P39'),('G16','Q39')):
            values = {('15. 손익계획',price):0}
            self.assertEqual(self.evaluate('15. 손익계획',target,values),0)

    def test_T_TRN_1_formula_based_transport_and_mutant(self):
        sheet = '15. 손익계획'
        sales = [630,650,700,720,750]
        expected = [18.9,19.5,21.0,21.6,22.5]
        values = {(sheet,'C39'):sales[0],(sheet,'C44'):sales[1],
                  (sheet,'M39'):1200,(sheet,'M40'):1/3}
        for col in 'MNOPQ':
            values[(sheet,col+'41')]=300
            values[(sheet,col+'42')]=0.1
        for cell,qty in zip(('O40','O49','O58'),sales[2:]):
            values[(' 7. 판매계획',cell)]=qty
        targets = ['E41','E46','E17','F17','G17']
        for cell,want in zip(targets,expected):
            self.assertAlmostEqual(self.evaluate(sheet,cell,values),want,places=8)
        self.assertAlmostEqual(self.evaluate(sheet,'E39',values),252,places=8)
        mutant=copy.deepcopy(self.entries)
        next(a for a in mutant if a['sheet']==sheet and a['cell']=='C41')['after_formula']='=C39/10'
        self.assertNotAlmostEqual(Evaluator(mutant,values).cell(sheet,'E41'),expected[0],places=8)



class PublicContract(unittest.TestCase):
    def setUp(self):
        lane_tmp = ROOT / '.hw-work/fin/lanes/I1-TRANSFORMER/tmp'
        lane_tmp.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=lane_tmp)
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'public-synthetic.xlsx'
        self.roster = h._json(h.REF / 'profile.json')['workbook_reference']['sheets']
        build(self.path, self.roster)
        self.sha = h.file_digest(self.path)
        self.docs = {
            h.FILES[0]: {'schema': 'knuaf-hort-finance-transform/v2', 'source_id': 'H01',
                         'source_sha256': self.sha, 'entry_count': len(entries(self.roster[0])),
                         'entries': entries(self.roster[0])},
            h.FILES[1]: {'schema': 'knuaf-hort-finance-params/v2', 'source_id': 'H01',
                         'source_sha256': self.sha, 'param_count': 1,
                         'parameters': {'synthetic': {'canonical_cell': self.roster[0]+'!D1',
                                                     'data_type': 'decimal', 'unit_code': 'KRW_1000',
                                                     'allowed_range': {'min': None, 'max': None},
                                                     'required_when': 'synthetic required',
                                                     'required_if': {'all': []},
                                                     'allowed_answer_states': ['provided'],
                                                     'output_conversion': 'identity'}},
                         'rules': {key: {} for key in h._AUDIT_RULE}},
            h.FILES[2]: {'schema': 'knuaf-hort-finance-verify-classes/v2', 'source_id': 'H01',
                         'source_sha256': self.sha, 'classes': ['numeric_required'],
                         'exhaustiveness': 'synthetic', 'explicit_none_zero_keys': [],
                         'explicit_none_blank_keys': [], 'rule_ids': [],
                         'registry': [{'sheet': e['sheet'], 'cell': e['cell'],
                                       'action': e['action'], 'result_kind': 'numeric_required',
                                       'unit': 'KRW_1000', 'printed': True,
                                       'core_reachable': True,
                                       'expected_state': 'formula'}
                                      for e in entries(self.roster[0])],
                         'teacher_example_coords': [],
                         'registry_count': len(entries(self.roster[0]))},
        }

    def audit(self, docs=None):
        return h.audit_public_documents(docs or self.docs, self.path, self.sha)

    def assert_code_path(self, docs, code, path):
        errors = self.audit(docs)
        self.assertIn({'code': code, 'path': path}, errors)

    def test_T_PUB_1_original_numeric_in_new_formula(self):
        self.assertEqual(self.audit(), [])
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[0]]['entries'][0]['after_formula'] = '=99'
        self.assert_code_path(docs, 'source_formula_number', h.FILES[0]+'.entries.0.after_formula')

    def test_T_PUB_2_original_numeric_in_after_value(self):
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[0]]['entries'][1]['after_value'] = 99
        self.assert_code_path(docs, 'source_numeric_value', h.FILES[0]+'.entries.1.after_value')

    def test_T_PUB_3_original_text_in_string_fields(self):
        for name, mutate, path in (
            (h.FILES[1], lambda d: d[h.FILES[1]]['parameters']['synthetic'].update(required_when='fixture'),
             h.FILES[1]+'.parameters.synthetic.required_when'),
            (h.FILES[1], lambda d: d[h.FILES[1]]['parameters']['synthetic'].update(output_conversion='fixture'),
             h.FILES[1]+'.parameters.synthetic.output_conversion'),
            (h.FILES[0], lambda d: d[h.FILES[0]]['entries'][1].update(after_value='fixture'),
             h.FILES[0]+'.entries.1.after_value'),
            (h.FILES[0], lambda d: d[h.FILES[0]]['entries'][0].update(reason='fixture'),
             h.FILES[0]+'.entries.0.reason'),
        ):
            docs = copy.deepcopy(self.docs)
            mutate(docs)
            self.assert_code_path(docs, 'source_string_value', path)
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[0]]['entries'][0]['sheet'] = 'fixture'
        self.assert_code_path(docs, 'entry_sheet', h.FILES[0]+'.entries[0].sheet')
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[1]]['parameters']['fixture'] = docs[h.FILES[1]]['parameters'].pop('synthetic')
        self.assert_code_path(docs, 'source_string_key', h.FILES[1]+'.parameters.fixture')

    def test_T_PUB_4_nested_unknown_and_forbidden_controls(self):
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[1]]['rules']['year_axis']['unapproved_field'] = 1
        self.assert_code_path(docs, 'rule_unknown_field', h.FILES[1]+'.rules.year_axis')
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[1]]['rules']['explicit_none_zero']['requires'] = [{'unapproved_field': 1}]
        self.assert_code_path(docs, 'rule_nested_value', h.FILES[1]+'.rules.explicit_none_zero')
        for field in ('before_formula', 'before_digest'):
            docs = copy.deepcopy(self.docs)
            docs[h.FILES[0]]['entries'][0][field] = 'x'
            self.assert_code_path(docs, 'entry_unknown_field', h.FILES[0]+'.entries[0].'+field)
            self.assert_code_path(docs, 'forbidden_key', h.FILES[0]+'.entries.0.'+field)

    def test_T_PUB_5_before_state_delete_and_source_sha(self):
        docs = copy.deepcopy(self.docs)
        del docs[h.FILES[0]]['entries'][0]['before_state']
        self.assert_code_path(docs, 'before_state_missing_or_invalid', h.FILES[0]+'.entries[0].before_state')
        for name in h.FILES:
            docs = copy.deepcopy(self.docs)
            del docs[name]['source_sha256']
            self.assert_code_path(docs, 'source_sha_missing_or_mismatch', name+'.source_sha256')
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[0]]['entries'][0]['before_state'] = 'absent'
        self.assert_code_path(docs, 'before_state_source_mismatch', h.FILES[0]+'.entries[0].before_state')

    def test_T_PUB_6_label_key_and_original_number_token(self):
        self.assertTrue(h._ALLOWED_CONSTANTS)
        self.assertTrue(all(isinstance(v, str) and v.strip() for v in h._ALLOWED_CONSTANTS.values()))
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[0]]['entries'][1]['label_key'] = 'copy_retained_label'
        self.assert_code_path(docs, 'label_key_contract', h.FILES[0]+'.entries[1].label_key')
        docs = copy.deepcopy(self.docs)
        docs[h.FILES[0]]['entries'][0]['after_formula'] = '=(2+3)+99'
        self.assert_code_path(docs, 'source_formula_number', h.FILES[0]+'.entries.0.after_formula')

    # The audit against the real H01 is a local-only step (FD-A): run
    # `gg_hort_workbook.py audit-public --source <H01>` where the workbook
    # exists.  CI has no H01, so the suite tests the auditor on synthetic docs.


if __name__ == "__main__":
    unittest.main()
