"""통합: WP-3 ``review_plan`` → ``check``의 ``finance_plan_review`` 행.

HANDOFF-WP3-to-WP1의 처분을 잠근다.

- 통합문서 경로·해시가 맞는 ``.xlsx`` 원자료가 있으면 검사가 실제로 돈다.
- ``blocked``는 통과가 아니라 미수행이며 제출 후보를 막는다.
- ``pending``은 경고라 막지 않는다.
- 원자료가 없거나 해시가 낡으면 이 검사는 행을 만들지 않는다(기존
  ``source_hash``/``source_missing`` 검사가 그 상태를 소유한다).
"""

import sys
from pathlib import Path

from tests._harness import ContractCase, bind_major, runtime, write_text

SCRIPTS = Path(__file__).resolve().parent.parent / "skills" / "knuaf-doc" / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

CHECK_ID = "finance_plan_review"

def _workbook(path, *, profit=(-10, -5, 0, 5, 10), sales=(100, 200, 300, 400, 500),
              cash=(0, 0, 0, 0, 0), years=(2027, 2028, 2029, 2030, 2031)):
    from openpyxl import Workbook

    wb = Workbook()
    assets = wb.active
    assets.title = "1. 기초재무상태조사"
    assets["B6"] = "(1) 토지"
    assets["B36"] = "(6) 현금 및 준현금"
    assets["E39"] = 1000

    plan = wb.create_sheet("12 손익계획")
    plan["B4"] = "과목"
    for offset, year in enumerate(years):
        plan.cell(4, 3 + offset, "%d년" % year)
    plan["B6"] = "I. 매출액"
    plan["B26"] = "X. 당기순이익"
    for offset in range(len(years)):
        plan.cell(6, 3 + offset, sales[offset])
        plan.cell(26, 3 + offset, profit[offset])

    flow = wb.create_sheet("14. 현금흐름계획")
    flow["B5"] = "구  분"
    for offset, year in enumerate(years):
        flow.cell(6, 4 + offset, "%d년" % year)
    flow["B24"] = "누적수지균형(D)\n(D=전년도D+금년도C)"
    for offset in range(len(years)):
        flow.cell(24, 4 + offset, cash[offset])

    wb.save(path)
    wb.close()

class FinancePlanReviewTests(ContractCase):
    def _project(self, *, workbook=True, mutate=None, stale=False,
                 profit=(-10, -5, 0, 5, 10)):
        core = runtime("gg_core")
        root = self.make_project()
        bind_major(root)
        if workbook:
            path = write_text(root, "finance/plan.xlsx", "")
            _workbook(path, profit=profit)
            if mutate:
                mutate(path)
            source = {"collection": "sources", "value": {
                "id": "plan", "path": "finance/plan.xlsx"}}
            core.apply(root, {"request_id": "integ-plan-source",
                              "ops": [source]}, core.load(root)["revision"])
            if stale:
                path.write_bytes(path.read_bytes() + b"")
                path.touch()
                write_text(root, "finance/plan.xlsx", "변조된 바이트")
        return core, root

    def _rows(self, report):
        return [r for r in report if r["check_id"] == CHECK_ID]

    def _codes(self, rows):
        return {r["evidence"][0].get("code") for r in rows if r.get("evidence")}

    def test_deficit_plan_reports_a_failing_row_that_blocks(self):
        core, root = self._project(profit=(-30, -20, -10, -5, 0))
        p = core.load(root)
        rows = self._rows(core.checks(root, p))
        self.assertTrue(rows, "finance_plan_review 행이 없음")
        self.assertTrue(all(r["target"] == "finance" for r in rows))
        deficit = [r for r in rows
                   if r["evidence"][0].get("code") == "PLAN_ALL_YEARS_DEFICIT"]
        self.assertEqual(1, len(deficit))
        self.assertEqual("fail", deficit[0]["status"])
        self.assertEqual("error", deficit[0]["severity"])
        self.assertTrue(core.blocks_skill_candidate(deficit[0]))

    def test_unreadable_profit_is_blocked_not_a_pass(self):
        core, root = self._project(
            mutate=lambda path: self._drop_profit_cache(path))
        p = core.load(root)
        rows = self._rows(core.checks(root, p))
        blocked = [r for r in rows if r["status"] == "blocked"]
        self.assertEqual(
            {"PLAN_NET_PROFIT_UNREADABLE"},
            {r["evidence"][0]["code"] for r in blocked})
        self.assertTrue(all(core.blocks_skill_candidate(r) for r in blocked))
        self.assertTrue(all(r.get("recovery") for r in blocked))
        # 미수행을 통과로 소비하지 않는다: 같은 검사에 pass 행이 없다.
        self.assertNotIn("pass", {r["status"] for r in rows})

    def test_undeclared_opening_assets_warns_without_blocking(self):
        core, root = self._project()
        p = core.load(root)
        rows = [r for r in self._rows(core.checks(root, p))
                if r["evidence"][0].get("code") == "OPENING_ASSETS_UNDECLARED"]
        self.assertEqual(1, len(rows))
        self.assertEqual("pending", rows[0]["status"])
        self.assertEqual("warning", rows[0]["severity"])
        self.assertFalse(core.blocks_skill_candidate(rows[0]))

    def test_no_current_workbook_source_makes_no_row(self):
        core, root = self._project(stale=True)
        p = core.load(root)
        self.assertEqual([], self._rows(core.checks(root, p)))
        self.assertIn("source_hash",
                      {r["check_id"] for r in core.checks(root, p)})

    def test_project_without_workbook_keeps_checks_unchanged(self):
        core, root = self._project(workbook=False)
        p = core.load(root)
        self.assertEqual([], self._rows(core.checks(root, p)))

    @staticmethod
    def _drop_profit_cache(path):
        import openpyxl

        workbook = openpyxl.load_workbook(path)
        sheet = workbook["12 손익계획"]
        for offset in range(5):
            sheet.cell(26, 3 + offset, "=B6*0.1")
        workbook.save(path)
        workbook.close()

class _MultiWorkbookCase(ContractCase):
    """INTEG-R1 반례용 바닥: 여러 통합문서·재무 게이트 기록."""

    def setUp(self):
        self.core = runtime("gg_core")
        self.root = self.make_project()
        bind_major(self.root)

    def apply(self, ops):
        revision = self.core.load(self.root)["revision"]
        return self.core.apply(self.root, {
            "request_id": "integ-r1:%d" % revision, "ops": ops}, revision)

    def add_source(self, sid, rel, **kwargs):
        path = write_text(self.root, rel, "")
        _workbook(path, **kwargs)
        self.apply([{"collection": "sources", "value": {"id": sid, "path": rel}}])
        return path

    def record_gate(self, state, profit=None, *, period="2030"):
        from tests._harness import claim_of, fact_op, source_op

        facts = [fact_op("gate-state", "finance.gate.state", state, "",
                         scope="project",
                         period=period if profit is not None else None,
                         verification="claim_supported",
                         source_id="calculation", claim_id="state")]
        if profit is not None:
            facts.append(fact_op(
                "gate-profit", "finance.gate.target_profit", str(profit),
                "원", value_type="decimal", period=period, scope="annual",
                kind="observation", verification="claim_supported",
                source_id="calculation", claim_id="profit"))
            facts[-1]["value"].update(
                finance_role="context", meaning_id="sales.net_profit",
                measure={"kind": "monetary_total", "currency": "KRW",
                         "unit": "원"})
        claims = {}
        for fact in facts:
            claims.update(claim_of(
                fact["value"], fact["value"]["source_refs"][0]["claim_id"]))
        write_text(self.root, "calculation.json", "{}")
        source = source_op("calculation.json", claims=claims)
        source["value"]["id"] = "calculation"
        self.apply([source, *facts])

    def accept_deficit(self, value="proceed_with_deficit"):
        from tests._harness import claim_of, fact_op, source_op

        field = "finance.gate.acceptance"
        profit = self.core.load(self.root)["facts"]["gate-profit"]["value"]
        write_text(self.root, "gate-interview.txt",
                   "질문: 연도별 손익 2029=-200, 2030=%s원. 목표 2030년.\n"
                   "근거 있는 규모·판로 조정 후 재검산 / 적자를 알고 진행 / 모름.\n"
                   "학생 답: %s\n" % (profit, value))
        fact = fact_op("gate-answer", field, value, "", scope="project",
                       verification="claim_supported",
                       source_id="gate-interview", claim_id="answer",
                       kind="reported_fact")
        fact["value"]["question_ref"] = {"id": field, "decision_revision": 1}
        source = source_op("gate-interview.txt",
                           claims=claim_of(fact["value"], "answer"))
        source["value"]["id"] = "gate-interview"
        self.apply([source, fact, {"collection": "questions", "value": {
            "id": field, "field_id": field, "attempts": 1,
            "decision_revision": 1,
            "source_ref": {"id": "gate-interview", "revision": 1,
                           "locator": "lines 1-2"}}}])

    def rows(self):
        return [r for r in self.core.checks(self.root, self.core.load(self.root))
                if r["check_id"] == CHECK_ID]

    def codes(self, rows=None):
        rows = self.rows() if rows is None else rows
        return [r["evidence"][0].get("code") for r in rows if r.get("evidence")]

class WorkbookSelectionTests(_MultiWorkbookCase):
    """INTEG-R1-01: 현재 계획을 개정 수·등록 순서로 추정하지 않는다."""

    def test_ancillary_workbook_does_not_block_the_current_plan(self):
        self.add_source("plan", "finance/plan.xlsx", profit=(1, 2, 3, 4, 5))
        quote = write_text(self.root, "quote.xlsx", "")
        from openpyxl import Workbook

        book = Workbook()
        book.active.title = "시설 견적"
        book.active["A1"] = "synthetic quote"
        book.save(quote)
        book.close()
        self.apply([{"collection": "sources", "value": {
            "id": "quote", "path": "quote.xlsx"}}])

        rows = self.rows()
        self.assertNotIn("PLAN_SHEET_MISSING", self.codes(rows))
        self.assertFalse([r for r in rows
                          if self.core.blocks_skill_candidate(r)])
        # 확정할 수 없다는 사실은 숨기지 않는다(차단 아님).
        ambiguous = [r for r in rows
                     if "FINANCE_PLAN_WORKBOOK_AMBIGUOUS" in self.codes([r])]
        self.assertEqual(1, len(ambiguous))
        self.assertEqual("pending", ambiguous[0]["status"])

    def test_old_revision_does_not_outrank_the_current_plan(self):
        self.add_source("old", "finance/old.xlsx", profit=(1, 2, 3, 4, 5))
        self.apply([{"collection": "sources", "value": {
            "id": "old", "path": "finance/old.xlsx"}}])
        self.add_source("current", "finance/current.xlsx",
                        profit=(-50, -40, -30, -20, -10))

        rows = self.rows()
        # 과거 계획의 결과로 현재 계획을 대신하지 않는다.
        self.assertNotIn("PLAN_ALL_YEARS_DEFICIT", self.codes(rows))
        self.assertIn("FINANCE_PLAN_WORKBOOK_AMBIGUOUS", self.codes(rows))

    def test_gate_target_selects_the_matching_plan(self):
        years = (2026, 2027, 2028, 2029, 2030)
        self.add_source("old", "finance/old.xlsx",
                        profit=(1, 2, 3, 4, 27050), years=years)
        self.add_source("current", "finance/current.xlsx",
                        profit=(-50, -40, -30, -20, -10), years=years)
        self.record_gate("deficit_pending", -10000)

        rows = self.rows()
        self.assertNotIn("FINANCE_PLAN_WORKBOOK_AMBIGUOUS", self.codes(rows))
        deficit = [r for r in rows
                   if r["evidence"][0].get("code") == "PLAN_ALL_YEARS_DEFICIT"]
        self.assertEqual(1, len(deficit))
        self.assertEqual("fail", deficit[0]["status"])
        self.assertEqual([-50.0, -40.0, -30.0, -20.0, -10.0],
                         deficit[0]["evidence"][0]["net_profit"])

    def test_near_miss_candidates_stay_ambiguous(self):
        """게이트 -10,000원 vs 후보 -11/-12천원은 어느 쪽도 일치가 아니다."""
        years = (2026, 2027, 2028, 2029, 2030)
        self.add_source("a", "finance/a.xlsx",
                        profit=(-50, -40, -30, -20, -11), years=years)
        self.add_source("b", "finance/b.xlsx",
                        profit=(-50, -40, -30, -20, -12), years=years)
        self.record_gate("deficit_pending", -10000)

        rows = self.rows()
        self.assertEqual(["FINANCE_PLAN_WORKBOOK_AMBIGUOUS"], self.codes(rows))
        self.assertEqual("pending", rows[0]["status"])
        # 어느 후보의 적자 판정도 현재 계획의 결과로 소비되지 않는다.
        self.assertNotIn("PLAN_ALL_YEARS_DEFICIT", self.codes(rows))

    def test_sub_won_conversion_difference_still_matches(self):
        """같은 값의 천원↔원 환산은 부동소수 오차만 있어도 일치다."""
        years = (2026, 2027, 2028, 2029, 2030)
        self.add_source("a", "finance/a.xlsx",
                        profit=(1, 2, 3, 4, 25753.87399530496), years=years)
        self.add_source("b", "finance/b.xlsx",
                        profit=(-50, -40, -30, -20, -10), years=years)
        self.record_gate("deficit_pending", 25753873.99530496)

        rows = self.rows()
        # a가 일치 후보로 선택되면 b의 적자는 판정되지 않고 모호 경고도 없다.
        self.assertNotIn("FINANCE_PLAN_WORKBOOK_AMBIGUOUS", self.codes(rows))
        self.assertNotIn("PLAN_ALL_YEARS_DEFICIT", self.codes(rows))
        self.assertEqual(["OPENING_ASSETS_UNDECLARED"], self.codes(rows))

class AcceptedDeficitTests(_MultiWorkbookCase):
    """INTEG-R1-03: 유효한 적자 수용은 같은 적자를 다시 차단하지 않는다."""

    def _accepted(self, profit=(-50, -40, -30, -20, -10), target=-10000):
        self.add_source("plan", "finance/plan.xlsx", profit=profit,
                        years=(2026, 2027, 2028, 2029, 2030))
        self.record_gate("deficit_accepted", target)
        self.accept_deficit()
        return self.rows()

    def test_accepted_deficit_is_a_warning_not_a_block(self):
        rows = self._accepted()
        self.assertEqual("deficit_accepted",
                         self.core.finance_gate(self.core.load(self.root))["state"])
        deficit = [r for r in rows
                   if r["evidence"][0].get("code") == "PLAN_ALL_YEARS_DEFICIT"]
        self.assertEqual(1, len(deficit))
        self.assertEqual("pending", deficit[0]["status"])
        self.assertEqual("warning", deficit[0]["severity"])
        self.assertFalse(self.core.blocks_skill_candidate(deficit[0]))
        self.assertTrue(deficit[0]["evidence"][0]["accepted"])

    def test_deficit_before_acceptance_still_blocks(self):
        self.add_source("plan", "finance/plan.xlsx",
                        profit=(-50, -40, -30, -20, -10),
                        years=(2026, 2027, 2028, 2029, 2030))
        self.record_gate("deficit_pending", -10000)
        rows = self.rows()
        deficit = [r for r in rows
                   if r["evidence"][0].get("code") == "PLAN_ALL_YEARS_DEFICIT"]
        self.assertEqual(1, len(deficit))
        self.assertEqual("fail", deficit[0]["status"])
        self.assertTrue(self.core.blocks_skill_candidate(deficit[0]))

    def test_acceptance_for_other_numbers_does_not_exempt(self):
        rows = self._accepted(profit=(-90, -80, -70, -60, -20),
                              target=-10000)
        deficit = [r for r in rows
                   if r["evidence"][0].get("code") == "PLAN_ALL_YEARS_DEFICIT"]
        self.assertEqual(1, len(deficit))
        self.assertEqual("fail", deficit[0]["status"])
        self.assertTrue(self.core.blocks_skill_candidate(deficit[0]))

    def test_acceptance_does_not_cover_a_near_miss_amount(self):
        """수용 -10,000원 vs 계획 -11천원(-11,000원)은 다른 적자다."""
        rows = self._accepted(profit=(-50, -40, -30, -20, -11),
                              target=-10000)
        self.assertEqual("deficit_accepted",
                         self.core.finance_gate(self.core.load(self.root))["state"])
        deficit = [r for r in rows
                   if r["evidence"][0].get("code") == "PLAN_ALL_YEARS_DEFICIT"]
        self.assertEqual(1, len(deficit))
        self.assertEqual("fail", deficit[0]["status"])
        self.assertEqual("error", deficit[0]["severity"])
        self.assertNotIn("accepted", deficit[0]["evidence"][0])
        self.assertTrue(self.core.blocks_skill_candidate(deficit[0]))

    def test_blocked_plan_stays_blocked_after_acceptance(self):
        self.add_source("plan", "finance/plan.xlsx",
                        profit=(-50, -40, -30, -20, -10),
                        years=(2026, 2027, 2028, 2029, 2030))
        path = self.root / "finance/plan.xlsx"
        from openpyxl import load_workbook

        book = load_workbook(path)
        for offset in range(5):
            book["12 손익계획"].cell(26, 3 + offset, "=B6*0.1")
        book.save(path)
        book.close()
        self.apply([{"collection": "sources", "value": {
            "id": "plan", "path": "finance/plan.xlsx"}}])
        self.record_gate("deficit_accepted", -10000)
        self.accept_deficit()
        rows = self.rows()
        blocked = [r for r in rows if r["status"] == "blocked"]
        self.assertEqual({"PLAN_NET_PROFIT_UNREADABLE"},
                         {r["evidence"][0]["code"] for r in blocked})
        self.assertTrue(all(self.core.blocks_skill_candidate(r)
                            for r in blocked))

class ProfileScopeTests(_MultiWorkbookCase):
    """INTEG-R1-02: 17시트 검사는 17시트 프로필에만 연결한다."""

    def test_horticulture_18_sheet_profile_is_not_reviewed(self):
        import json

        hort = runtime("gg_hort_workbook")
        roster = json.loads(
            (hort.REF / "profile.json").read_text(encoding="utf-8")
        )["workbook_reference"]["sheets"]
        root = self.make_project()
        bind_major(root, "hort_env_systems")
        path = write_text(root, "finance/hort.xlsx", "")
        from openpyxl import Workbook

        book = Workbook()
        book.remove(book.active)
        for name in roster:
            book.create_sheet(name)
        book.save(path)
        book.close()
        core = runtime("gg_core")
        core.apply(root, {"request_id": "integ-hort", "ops": [
            {"collection": "sources", "value": {
                "id": "hort", "path": "finance/hort.xlsx"}}]},
            core.load(root)["revision"])

        self.assertTrue(core.finance_gate(core.load(root))["supported"])
        self.assertEqual([], [r for r in core.checks(root, core.load(root))
                              if r["check_id"] == CHECK_ID])

