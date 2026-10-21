"""WP-3 regression tests: five-year trajectory and area denominators.

- all-deficit plans fail, loss-with-cash-growth asks for an explanation;
- opening assets are compared per category, not only by total;
- body/bed area denominators are checked against one declaration.
"""

from pathlib import Path
import sys
import tempfile
import unittest

TESTS = Path(__file__).resolve().parent
SCRIPTS = TESTS.parent / "skills" / "knuaf-doc" / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import gg_finance_tables as ft  # noqa: E402


def plan_workbook(path, *, years=(2027, 2028, 2029, 2030, 2031),
                  sales=(100, 200, 300, 400, 500),
                  profit=(-10, -5, 0, 5, 10),
                  cash=(0, 0, 0, 0, 0),
                  loan=(0, 0, 0, 0, 0),
                  facility_invest=(0, 0, 0, 0, 0),
                  machine_invest=(0, 0, 0, 0, 0),
                  depreciation=100, bed_note=False):
    from openpyxl import Workbook
    wb = Workbook()
    assets = wb.active
    assets.title = "1. 기초재무상태조사"
    assets["B6"] = "(1) 토지"
    assets["B9"] = "답"
    assets["D9"] = 1500
    assets["E9"] = 112283.094
    assets["B36"] = "(6) 현금 및 준현금"
    assets["C39"] = "예금"
    assets["E39"] = 87716.906
    assets["B41"] = "합계"
    assets["E41"] = 87716.906
    assets["B43"] = "(7) 재고생산자재"

    plan = wb.create_sheet("12 손익계획")
    plan["B4"] = "과목"
    for offset, year in enumerate(years):
        plan.cell(4, 3 + offset, "%d년" % year)
    plan["B6"] = "I. 매출액"
    plan["B26"] = "X. 당기순이익"
    for offset in range(len(years)):
        plan.cell(6, 3 + offset, sales[offset])
        plan.cell(26, 3 + offset, profit[offset])

    flow = wb.create_sheet("14. 현금흐름계획")
    flow["B5"] = "구  분"
    for offset, year in enumerate(years):
        flow.cell(6, 4 + offset, "%d년" % year)
    flow["C12"] = "융자금"
    flow["C15"] = "시설투자"
    flow["C16"] = "농기계투자"
    flow["B24"] = "누적수지균형(D)\n(D=전년도D+금년도C)"
    for offset in range(len(years)):
        flow.cell(12, 4 + offset, loan[offset])
        flow.cell(15, 4 + offset, facility_invest[offset])
        flow.cell(16, 4 + offset, machine_invest[offset])
        flow.cell(24, 4 + offset, cash[offset])

    dep = wb.create_sheet("10. 감가상각비계획 ")
    for index in range(len(years)):
        base = 4 + index * 10
        dep.cell(base + 2, 5, "계")
        dep.cell(base + 2, 9, "계")
        dep.cell(base + 8, 2, "감 가 상 각 비")
        dep.cell(base + 8, 3, depreciation / 2)
        dep.cell(base + 8, 5, depreciation / 2)
        dep.cell(base + 8, 9, depreciation / 2)

    invest = wb.create_sheet("3.투자계획")
    invest["C19"] = "비닐하우스 신축"
    invest["D19"] = "1,500㎡ 파이프비닐"

    production = wb.create_sheet("6. 생산계획")
    production["B8"] = "재배면적(㎡)"
    production["H8"] = 1500

    expense = wb.create_sheet("9 .경비계획")
    expense["B8"] = "경영면적(m2) > 1년차"
    expense["D8"] = 1500
    if bed_note:
        expense["B20"] = "* 배지: 물량 = (1,500㎡ ÷ 1.9m) × 0.15m = 118.4㎥"
    wb.save(path)
    wb.close()


def codes(result):
    return [issue["code"] for issue in result["issues"]]




class PlanReviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "plan.xlsx"

    def review(self, **kwargs):
        return ft.review_plan(self.path, **kwargs)

    def test_all_years_deficit_is_an_error(self):
        plan_workbook(self.path, profit=(-10, -20, -30, -40, -50))
        result = self.review()
        self.assertIn("PLAN_ALL_YEARS_DEFICIT", codes(result))

    def test_growth_jump_and_unrecovered_trajectory(self):
        plan_workbook(self.path, sales=(100, 200, 700, 800, 900),
                      profit=(-100, -50, -10, 20, 40))
        result = self.review()
        self.assertIn("PLAN_GROWTH_JUMP", codes(result))
        self.assertIn("PLAN_TRAJECTORY_NOT_RECOVERED", codes(result))
        self.assertNotIn("PLAN_ALL_YEARS_DEFICIT", codes(result))

    def test_loss_with_cash_growth_requires_explanation(self):
        plan_workbook(self.path, profit=(-10, -20, -30, -40, -50),
                      cash=(100000, 100000, 100000, 100000, 100000),
                      depreciation=500)
        result = self.review()
        issue = next(i for i in result["issues"]
                     if i["code"] == "CASH_PROFIT_DIVERGENCE")
        self.assertEqual(issue["evidence"]["depreciation"], 500)
        self.assertEqual(issue["evidence"]["cash_before"], 87716.906)

    def test_loss_with_falling_cash_is_not_flagged(self):
        plan_workbook(self.path, profit=(-10, 20, 30, 40, 50),
                      cash=(1000, 2000, 3000, 4000, 5000))
        result = self.review()
        self.assertNotIn("CASH_PROFIT_DIVERGENCE", codes(result))

    def test_opening_assets_compare_per_category(self):
        plan_workbook(self.path)
        result = self.review(declared_assets={"cash": 200000})
        mismatches = [i for i in result["issues"]
                      if i["code"] == "OPENING_ASSETS_MISMATCH"]
        self.assertEqual({i["evidence"]["role"] for i in mismatches},
                         {"land", "cash"})
        matched = self.review(declared_assets={
            "land": 112283.094, "cash": 87716.906})
        self.assertNotIn("OPENING_ASSETS_MISMATCH", codes(matched))

    def test_undeclared_opening_assets_is_a_warning(self):
        plan_workbook(self.path)
        result = self.review()
        issue = next(i for i in result["issues"]
                     if i["code"] == "OPENING_ASSETS_UNDECLARED")
        self.assertEqual(issue["severity"], "warning")
        self.assertEqual(issue["evidence"]["total"], 200000.0)

    def test_area_denominators_compare_body_and_bed_base(self):
        plan_workbook(self.path, bed_note=True)
        result = self.review(
            area_declarations={"facility_m2": 1500, "cultivation_m2": 1500,
                               "management_m2": 1500, "bed_m2": 789},
            body_text="재배면적은 1,000 ㎡로 잡았다.")
        self.assertIn("AREA_BODY_MISMATCH", codes(result))
        self.assertIn("AREA_BED_BASE_MISMATCH", codes(result))
        clean = self.review(
            area_declarations={"facility_m2": 1500, "cultivation_m2": 1500,
                               "management_m2": 1500, "bed_m2": 1500},
            body_text="재배면적은 1,500 ㎡로 잡았다.")
        self.assertNotIn("AREA_BODY_MISMATCH", codes(clean))
        self.assertNotIn("AREA_BED_BASE_MISMATCH", codes(clean))

    def test_area_declaration_mismatch(self):
        plan_workbook(self.path)
        result = self.review(area_declarations={"cultivation_m2": 1200})
        self.assertIn("AREA_DECLARATION_MISMATCH", codes(result))
