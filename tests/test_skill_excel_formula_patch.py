"""gg_excel_formula_patch: the only path that changes the school form's formulas.

Migrated from the GUI contract suite (cb72ae1).  Filling writes student
values; this rewrites the workbook's own arithmetic.  The governing rule is
that you may not change a formula unless you correctly state what is in the
cell right now — a patch written against a misread formula is refused rather
than applied to whatever happens to be there.
"""
import json
import tempfile
import unittest
from pathlib import Path

from tests._harness import ContractCase, runtime

SHEET = "1. 입력"


class FormulaPatchTests(ContractCase):
    def setUp(self):
        self.fp = runtime("gg_excel_formula_patch")
        import openpyxl as _openpyxl

        self.openpyxl = _openpyxl
        tmp = tempfile.TemporaryDirectory(prefix="p1-formula-")
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)
        wb = self.openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET
        ws["C1"] = 600
        ws["C2"] = 2
        ws["C3"] = "=C1*2"
        ws["C4"] = 1200
        ws["B5"] = "병합"
        ws.merge_cells("B5:D5")
        self.source = self.tmp / "학교 양식.xlsx"
        wb.save(self.source)

    def patch(self, cell="C3", drop=(), **over):
        base = {"sheet": SHEET, "cell": cell,
                "expected_formula": "=C1*2", "new_formula": "=C1*C2",
                "reason": "면적 단가가 상수로 고정돼 있어 단가 셀을 참조하도록 수정",
                "evidence_ref": {"source_id": "src-school", "revision": 1,
                                 "locator": "p.7"}}
        base.update(over)
        for key in (drop,) if isinstance(drop, str) else drop:
            base.pop(key, None)
        return base

    def value_patch(self, cell="C4", expected=1200, **over):
        return self.patch(cell=cell, drop="expected_formula",
                          expected_value=expected, **over)

    def write_map(self, patches, sha=None):
        m = self.tmp / "formula-map.json"
        m.write_text(json.dumps({
            "schema": self.fp.MAP_SCHEMA,
            "source": {"sha256": sha or self.fp.sha256(self.source)},
            "patches": patches}, ensure_ascii=False), encoding="utf-8")
        return m

    def run_patch(self, patches, sha=None, out_name="고친 사본.xlsx"):
        return self.fp._patch_copy(
            self.source, self.write_map(patches, sha), self.tmp / out_name)

    def run_cli(self, patches, out="고친 사본.xlsx", receipt="영수증.json"):
        # Policy B: the CLI publishes through the major output guard, so it
        # needs an explicit major bound in the canonical record.
        from tests._harness import bind_major
        project = self.tmp / "project"
        if not (project / "project.json").exists():
            runtime("gg_core").init(project)
            bind_major(project)
        return self.fp.cli([
            "--source", str(self.source),
            "--map", str(self.write_map(patches)),
            "--out", str(self.tmp / out),
            "--receipt", str(self.tmp / receipt),
            "--project", str(project),
            "--major", "specialty_crops"])

    # --- the happy path ---------------------------------------------------

    def test_a_reviewed_formula_is_replaced_and_the_source_is_untouched(self):
        before = self.source.read_bytes()
        receipt = self.run_patch([self.patch()])
        ws = self.openpyxl.load_workbook(
            self.tmp / "고친 사본.xlsx")[SHEET]
        self.assertEqual("=C1*C2", str(ws["C3"].value).replace(" ", ""))
        self.assertEqual(before, self.source.read_bytes())
        self.assertGreaterEqual(receipt["formulaCachesInvalidated"], 1)

    def test_the_receipt_records_both_the_old_and_the_new_formula(self):
        receipt = self.run_patch([self.patch()])
        blob = json.dumps(receipt["patched"][0], ensure_ascii=False)
        self.assertIn("C1*2", blob)
        self.assertIn("C1*C2", blob)
        self.assertIn("면적 단가", blob)

    # --- you must know what you are changing ------------------------------

    def test_a_formula_that_is_not_what_the_patch_claims_is_refused(self):
        with self.assertRaisesRegex(ValueError, "expected formula mismatch"):
            self.run_patch([self.patch(expected_formula="=C1*99")])

    def test_a_patch_against_a_non_formula_cell_is_refused(self):
        with self.assertRaisesRegex(ValueError, "not a formula cell"):
            self.run_patch([self.patch(cell="C1")])

    def test_a_source_changed_since_review_is_refused(self):
        with self.assertRaisesRegex(RuntimeError, "source version changed"):
            self.run_patch([self.patch()], sha="0" * 64)

    def test_a_merged_non_anchor_cannot_be_patched(self):
        with self.assertRaisesRegex(ValueError, "missing cell|merged non-anchor"):
            self.run_patch([self.patch(cell="C5")])

    def test_two_patches_to_one_cell_are_refused(self):
        with self.assertRaisesRegex(ValueError, "duplicate patch target"):
            self.run_patch([self.patch(), self.patch(new_formula="=C1+C2")])

    def test_a_sheet_that_does_not_exist_is_refused(self):
        with self.assertRaisesRegex(ValueError, "missing sheet"):
            self.run_patch([self.patch(sheet="없는 시트")])

    # --- restoring a lost calculation from a constant ---------------------

    def test_a_numeric_constant_may_become_a_formula_when_stated_exactly(self):
        receipt = self.run_patch([self.value_patch()])
        ws = self.openpyxl.load_workbook(
            self.tmp / "고친 사본.xlsx")[SHEET]
        self.assertEqual("=C1*C2", str(ws["C4"].value).replace(" ", ""))
        self.assertIn("value_to_formula", json.dumps(receipt, ensure_ascii=False))

    def test_a_constant_that_is_not_the_stated_number_is_refused(self):
        with self.assertRaisesRegex(ValueError, "expected value mismatch"):
            self.run_patch([self.value_patch(expected=999)])

    def test_expected_value_may_not_target_a_formula_cell(self):
        with self.assertRaisesRegex(
                ValueError, "expected_value target is a formula cell"):
            self.run_patch([self.value_patch(cell="C3")])

    def test_a_patch_must_state_exactly_one_precondition(self):
        with self.assertRaisesRegex(
                ValueError, "exactly one of expected_formula or expected_value"):
            self.run_patch([self.patch(drop="expected_formula")])
        with self.assertRaisesRegex(
                ValueError, "exactly one of expected_formula or expected_value"):
            self.run_patch([self.patch(expected_value=1200)])

    def test_a_non_finite_expected_value_is_refused(self):
        for bad in (float("inf"), float("nan"), True, "1200", None):
            with self.subTest(bad=bad):
                with self.assertRaises(ValueError):
                    self.run_patch([self.value_patch(expected=bad)])

    # --- formulas that must never be written ------------------------------

    def test_external_and_dde_formulas_are_refused(self):
        for hostile in (
            '=DDE("cmd","/c calc","x")',
            '=WEBSERVICE("http://example.invalid/x")',
            '=HYPERLINK("http://example.invalid","click")',
            "=[다른파일.xlsx]Sheet1!A1",
            "=#REF!*2",
            "=C1|C2",
        ):
            with self.subTest(hostile=hostile):
                with self.assertRaisesRegex(
                        ValueError, "unsafe external/DDE formula"):
                    self.run_patch([self.patch(new_formula=hostile)])

    def test_a_hostile_expected_formula_is_refused_too(self):
        with self.assertRaisesRegex(ValueError, "unsafe external/DDE formula"):
            self.run_patch([self.patch(
                expected_formula='=WEBSERVICE("http://x.invalid")')])

    def test_an_empty_or_non_string_formula_is_refused(self):
        for empty in ("", "   ", "=", None, 5):
            with self.subTest(empty=empty):
                with self.assertRaises(ValueError):
                    self.run_patch([self.patch(new_formula=empty)])

    # --- every patch needs a reason and evidence --------------------------

    def test_every_patch_needs_full_evidence(self):
        for missing in ("source_id", "revision", "locator"):
            with self.subTest(missing=missing):
                ev = {"source_id": "s", "revision": 1, "locator": "p.7"}
                ev.pop(missing)
                with self.assertRaisesRegex(ValueError, "evidence_ref"):
                    self.run_patch([self.patch(evidence_ref=ev)])

    def test_every_patch_needs_a_reason(self):
        for reason in ("", "   ", None):
            with self.subTest(reason=reason):
                with self.assertRaisesRegex(
                        ValueError, "reason must be non-empty"):
                    self.run_patch([self.patch(reason=reason)])

    def test_a_map_with_no_patches_is_refused(self):
        with self.assertRaisesRegex(ValueError, "non-empty patches"):
            self.run_patch([])

    # --- publication is all-or-nothing ------------------------------------

    def test_the_cli_publishes_the_copy_and_the_receipt_together(self):
        self.assertEqual(0, self.run_cli([self.patch()]))
        self.assertTrue((self.tmp / "고친 사본.xlsx").is_file())
        receipt = json.loads(
            (self.tmp / "영수증.json").read_text(encoding="utf-8"))
        self.assertTrue(receipt["patched"])

    def test_the_cli_refuses_to_overwrite_an_existing_output(self):
        self.assertEqual(0, self.run_cli([self.patch()]))
        self.assertEqual(2, self.run_cli([self.patch()]))

    def test_a_refused_patch_leaves_no_output_and_no_staging_file(self):
        self.assertEqual(2, self.run_cli(
            [self.patch(expected_formula="=WRONG()")]))
        self.assertFalse((self.tmp / "고친 사본.xlsx").exists())
        self.assertFalse((self.tmp / "영수증.json").exists())
        self.assertEqual([], [p.name for p in self.tmp.glob(".*tmp-*")])


if __name__ == "__main__":
    unittest.main()
