"""Finance asset selection, investigation safety, and OOXML copy boundaries."""
from copy import copy
import json
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from tests._harness import ContractCase, bind_major, runtime


class FinanceTemplateIntegrationTests(ContractCase):
    def setUp(self):
        self.tpl = runtime("gg_excel_template")
        self.mc = runtime("gg_major_contract")
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.folder = Path(tmp.name)
        self.project = self.make_project()
        bind_major(self.project)
        self.context = self.mc.output_context(self.project, "specialty_crops")

    def workbook(self):
        path = self.folder / "student.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Student"
        ws["A1"] = "합성표제"
        ws["B2"] = 123
        ws["C2"] = "=B2*2"
        ws["D2"] = "merged"
        ws.merge_cells("D2:E2")
        ws["B2"].alignment = Alignment(wrap_text=False, horizontal="right")
        ws["C2"].alignment = Alignment(wrap_text=True, horizontal="center")
        ws["F2"] = 12
        ws["F2"].alignment = Alignment(horizontal="left")
        wb.save(path)
        return path

    def write_map(self, source, entries):
        path = self.folder / "map.json"
        path.write_text(json.dumps({
            "schema": "gg-xlsx-template-map/v1", "source": {"sha256": self.tpl.sha256(source)},
            "layout": {"variant": "unknown", "mapOrigin": "reviewed_custom"},
            "entries": entries}), encoding="utf-8")
        return path

    def test_student_copy_preserves_values_and_refuses_overwrite(self):
        source = self.workbook()
        original = source.read_bytes()
        out = self.folder / "copy.xlsx"
        receipt = self.tpl.create_workbook(source, out, context=self.context, template_ref="X02")
        self.assertEqual(out.read_bytes(), original)
        self.assertEqual(receipt["source"]["kind"], "student")
        self.assertEqual(receipt["source"]["sha256"], receipt["output"]["sha256"])
        with self.assertRaises(FileExistsError):
            self.tpl.create_workbook(source, out, context=self.context)
        self.assertEqual(source.read_bytes(), original)
        self.assertEqual(out.read_bytes(), original)
        with self.assertRaises(FileNotFoundError):
            self.tpl.create_workbook(self.folder / "missing.xlsx", self.folder / "missing-out.xlsx",
                                     context=self.context, template_ref="X01")
        self.assertFalse((self.folder / "missing-out.xlsx").exists())

    def test_no_source_requires_major_and_explicit_x01_choice(self):
        for context, ref, reason in ((None, "X01", "output_context_required"),
                                     (self.context, None, "student_template_required"),
                                     (self.context, "X02", "student_template_required")):
            out = self.folder / "held.xlsx"
            with self.assertRaises(self.mc.OutputHeldError) as raised:
                self.tpl.create_workbook(None, out, context=context, template_ref=ref)
            self.assertEqual(raised.exception.reason, reason)
            self.assertFalse(out.exists())
        out = self.folder / "bundled.xlsx"
        receipt = self.tpl.create_workbook(None, out, context=self.context, template_ref="X01")
        self.assertEqual(receipt["source"]["kind"], "bundled")
        self.assertEqual(self.tpl.sha256(out), self.tpl.BUNDLED_TEMPLATE_SHA256)

    def test_other_majors_never_silently_fall_back_to_x01(self):
        for major in ("hort_env_systems", "fruit_trees"):
            project = self.make_project()
            bind_major(project, major)
            context = self.mc.output_context(project, major)
            out = self.folder / (major + ".xlsx")
            with self.assertRaises(self.mc.OutputHeldError):
                self.tpl.create_workbook(None, out, context=context, template_ref="X01")
            self.assertFalse(out.exists())
        inventory = {"survey": {"scope_label": "fruit", "scope_id": "/synthetic/fruit",
                                 "status": "complete", "evidence": {
                                     "verified": True, "basis": "user_attested"}}, "files": []}
        with self.assertRaises(self.mc.OutputHeldError) as raised:
            self.tpl.create_workbook(None, out, context=context, template_ref="X01",
                                     reference_inventory=inventory)
        self.assertEqual(raised.exception.reason, "unsupported_output")
        self.assertFalse(out.exists())

    def test_inventory_is_not_promoted_by_editing_status_alone(self):
        source = self.workbook()
        path = self.folder / "inventory.json"
        data = self.tpl.inspect_source(source, path, None)
        observed = {entry["range"] for entry in data["entries"]}
        self.assertIn("B2", observed)
        self.assertNotIn("C2", observed)
        self.assertNotIn("E2", observed)
        self.assertTrue(all(entry["action"] == "preserve" and not entry["explicit"]
                            for entry in data["entries"]))
        data["mapStatus"] = "ok"
        path.write_text(json.dumps(data), encoding="utf-8")
        out = self.folder / "unsafe.xlsx"
        with self.assertRaisesRegex(ValueError, "map_layout_invalid"):
            self.tpl.blank_copy(source, path, out, context=self.context)
        self.assertFalse(out.exists())

    def test_shared_formula_followers_are_not_literal_candidates(self):
        inventory = [{"sheet": "s", "cell": "A1", "value": 4,
                      "formula": None, "hasFormula": True},
                     {"sheet": "s", "cell": "A2", "value": 5,
                      "formula": None, "hasFormula": False}]
        self.assertEqual([entry["range"] for entry in self.tpl.inventory_entries(inventory, {})], ["A2"])

    def test_allow_missing_is_explicit_partial_and_keeps_formula(self):
        source = self.workbook()
        original = source.read_bytes()
        path = self.write_map(source, [
            {"sheet": "Student", "cell": "B2", "action": "clear"},
            {"sheet": "Student", "cell": "C2", "action": "clear"},
            {"sheet": "Student", "cell": "Z99", "action": "clear"},
            {"sheet": "Student", "cell": "Y98", "action": "preserve"}])
        out = self.folder / "partial.xlsx"
        with self.assertRaisesRegex(ValueError, "missing"):
            self.tpl.blank_copy(source, path, out, context=self.context)
        self.assertFalse(out.exists())
        receipt = self.tpl.blank_copy(source, path, out, allow_missing=True, context=self.context)
        self.assertEqual(receipt["missingCells"], [{"sheet": "Student", "cell": "Y98"},
                                                   {"sheet": "Student", "cell": "Z99"}])
        self.assertEqual(receipt["output"]["status"], "partial")
        self.assertFalse(receipt["publicationReady"])
        wb = load_workbook(out)
        self.addCleanup(wb.close)
        self.assertIsNone(wb["Student"]["B2"].value)
        self.assertEqual(wb["Student"]["C2"].value, "=B2*2")
        self.assertEqual(source.read_bytes(), original)
        path = self.write_map(source, [{"sheet": "Absent", "cell": "A1", "action": "clear"}])
        with self.assertRaisesRegex(ValueError, "missing sheet"):
            self.tpl.blank_copy(source, path, self.folder / "bad-sheet.xlsx", True, context=self.context)
        self.assertFalse((self.folder / "bad-sheet.xlsx").exists())

    def test_number_only_print_keeps_alignment_and_defined_names_order(self):
        source = self.workbook()
        # A legal externalReferences predecessor exposes the old sheets+1 insertion.
        ns = self.tpl.NS_MAIN
        with zipfile.ZipFile(source) as archive:
            parts = {info.filename: archive.read(info.filename) for info in archive.infolist()}
        book = ET.fromstring(parts["xl/workbook.xml"])
        names = book.find(f"{{{ns}}}definedNames")
        if names is not None:
            book.remove(names)
        sheets = book.find(f"{{{ns}}}sheets")
        book.insert(list(book).index(sheets) + 1, ET.Element(f"{{{ns}}}externalReferences"))
        parts["xl/workbook.xml"] = ET.tostring(book)
        with zipfile.ZipFile(source, "w") as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
        path = self.folder / "print-map.json"
        path.write_text(json.dumps({"schema": "gg-xlsx-print-map/v1",
            "source": {"sha256": self.tpl.sha256(source)}, "sheets": [{
                "sheet": "Student", "reason": "synthetic numeric display", "fit_width": 1,
                "fit_height": 0, "print_titles": {"rows": [1]},
                "number_formats": [{"cells": ["B2", "C2", "F2"], "format": "#,##0"}]}]}), encoding="utf-8")
        out = self.folder / "print.xlsx"
        runtime("gg_excel_print").apply(source, path, out, self.folder / "print-receipt.json",
                                         context=self.context)
        before, after = load_workbook(source), load_workbook(out)
        self.addCleanup(before.close)
        self.addCleanup(after.close)
        for cell in ("B2", "C2", "F2"):
            self.assertEqual(copy(before["Student"][cell].alignment), copy(after["Student"][cell].alignment))
            self.assertEqual(before["Student"][cell].value, after["Student"][cell].value)
            self.assertEqual(after["Student"][cell].number_format, "#,##0")
        self.assertEqual(after["Student"].print_title_rows, "$1:$1")
        with zipfile.ZipFile(out) as archive:
            children = [node.tag.rsplit("}", 1)[-1] for node in ET.fromstring(archive.read("xl/workbook.xml"))]
        self.assertLess(children.index("externalReferences"), children.index("definedNames"))
        self.assertLess(children.index("definedNames"), children.index("calcPr"))
