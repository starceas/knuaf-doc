"""W-A: school_17_sheet_v1 세 상승률 분리 + 공식 통계 레지스트리.

입력은 상승률 r(0.05=5%), 내부 배율은 1+r. 일반물가(r_g)→자재·수도광열비,
임금(r_w)→노무비, 판매가(r_s)→매출. 같은 가격에 률을 두 번 곱하지 않고,
판매가에 총지수를 자동 대입하지 않는다. 모든 거부는 ValueError.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

from tests._harness import (
    ContractCase, bind_major, runtime, write_text,
)

YEARS = [2027, 2028, 2029, 2030, 2031]  # writing_year=2026


def _spec(**over):
    """Synthetic school_17_sheet_v1 spec (no real student data)."""
    spec = dict(
        profile="school_17_sheet_v1", writing_year=2026,
        source_refs=["synthetic:합성-원답변"], unit="천원", quantity_unit="kg",
        production_evidence={k: "합성 근거" for k in
                             ("area", "seedlings", "yield", "growth",
                              "commodity")},
        crops=["합성작목"], area_m2="1000", production_kg="2000",
        purchase_price="3", direct_price="5", purchase_share="0.5",
        direct_share="0.5", land="0", facility="10000", equipment="2000",
        equity="6000", loan="6000", loan_rate="0.02", salvage="1000",
        materials="500", labor="800", packing="100", transport="100",
        household="2000", repair_facility_rate="0.01",
        repair_equipment_rate="0.02", utility_per_10a="50",
        price_assumptions={
            "general": {
                "source_id": "official.kosis.cpi.total",
                "base_year": 2020,
                "application_base_year": 2027,
            },
            "wage": {
                "source_id": "official.kosis.farm_purchase.labor",
                "base_year": 2020,
                "application_base_year": 2027,
            },
            "sales": {
                "status": "not_applied",
                "reason": "학생 작목 판매가 수열 확인 불가",
            },
        },
        grace=1, term=5, life=10)
    spec.update(over)
    return spec


def _student_source(sid="student.sample.series", **over):
    """공식 출처 계약을 충족하는 합성 학생 수열(실제 공표값이 아닌 테스트 선언)."""
    entry = dict(
        id=sid, agency="합성 공공기관(테스트)",
        agency_kind="public_corporation",
        statistic_name="합성 수매단가",
        item_name="합성작목 수매 단가", unit="원/kg", base_year=None,
        observations=[
            {"year": 2023, "value": 100.0},
            {"year": 2024, "value": 104.0},
            {"year": 2025, "value": 108.16},
        ],
        url="https://synthetic.example.go.kr/price/test",
        evidence_locator="합성 위치 — 실제 공표값이 아닌 테스트 선언",
        dimensions={
            "region": "전국", "crop": "합성작목", "channel": "농가수취"},
        allowed_roles=["sales"], unconfirmed_years=[],
        retrieved_at="2026-10-07", revision=1)
    entry.update(over)
    return entry


class RegistryIntegrityTests(ContractCase):
    """번들 레지스트리는 R1이 실제 조회한 공표값만 싣는다."""

    def setUp(self):
        self.pa = runtime("gg_price_assumptions")

    def test_registry_loads_r1_values_verbatim(self):
        reg = self.pa.load_registry()
        cpi = reg["official.kosis.cpi.total"]
        self.assertEqual([100.0, 102.5, 107.72, 111.59, 114.18, 116.61],
                         [o["value"] for o in cpi["observations"]])
        labor = reg["official.kosis.farm_purchase.labor"]
        self.assertEqual(137.9, labor["observations"][-1]["value"])
        self.assertEqual(2020, labor["base_year"])
        self.assertEqual("DT_1J62", labor["table_id"])
        kamis = reg["official.kamis.retail.rice_20kg.top"]
        self.assertIn(2020, kamis["unconfirmed_years"])
        self.assertEqual([59080, 51336, 53278, 53512, 58872],
                         [o["value"] for o in kamis["observations"]])
        potato = reg["official.rda.crop_income.spring_potato.output_price"]
        self.assertEqual([898, 940, 1193, 1300, 1190],
                         [o["value"] for o in potato["observations"]])
        self.assertIn(2025, potato["unconfirmed_years"])
        for sid, entry in reg.items():
            self.assertIn(entry["agency_kind"],
                          self.pa.OFFICIAL_AGENCY_KINDS, sid)

    def test_kamis_rice_is_reference_only(self):
        """V1-03: 번들 KAMIS 쌀 수열은 참고 전용 — 어떤 역할에도 대입 불가."""
        reg = self.pa.load_registry()
        kamis = reg["official.kamis.retail.rice_20kg.top"]
        self.assertEqual([], kamis["allowed_roles"])

    def test_registry_cagr_matches_r1(self):
        reg = self.pa.load_registry()
        labor = {o["year"]: o["value"] for o in
                 reg["official.kosis.farm_purchase.labor"]["observations"]}
        self.assertAlmostEqual(
            self.pa.cagr(labor[2020], labor[2025], 5), 0.06638212, places=6)
        kamis = {o["year"]: o["value"] for o in
                 reg["official.kamis.retail.rice_20kg.top"]["observations"]}
        self.assertLess(
            self.pa.cagr(kamis[2021], kamis[2025], 4), 0)  # 하락 수열

    def test_farm_sales_total_is_aggregate_no_sales_role(self):
        reg = self.pa.load_registry()
        agg = reg["official.kosis.farm_sales.total"]
        self.assertTrue(agg["aggregate"])
        self.assertNotIn("sales", agg["allowed_roles"])


class NormalizeValidationTests(ContractCase):
    def setUp(self):
        self.pa = runtime("gg_price_assumptions")
        self.sx = runtime("gg_school_excel")

    def norm(self, spec):
        return self.pa.normalize_price_assumptions(spec, YEARS)

    def test_missing_block_rejected(self):
        spec = _spec()
        del spec["price_assumptions"]
        with self.assertRaisesRegex(ValueError, "price_assumptions"):
            self.norm(spec)

    def test_missing_role_rejected(self):
        spec = _spec()
        del spec["price_assumptions"]["wage"]
        with self.assertRaisesRegex(ValueError, "wage"):
            self.norm(spec)

    def test_legacy_inflation_field_rejected_with_guidance(self):
        spec = _spec(inflation="1.02")
        with self.assertRaisesRegex(ValueError, "price_assumptions"):
            self.sx.validate(spec)

    def test_missing_source_id_rejected(self):
        spec = _spec()
        del spec["price_assumptions"]["general"]["source_id"]
        with self.assertRaisesRegex(ValueError, "출처 id"):
            self.norm(spec)

    def test_unknown_source_id_rejected(self):
        spec = _spec()
        spec["price_assumptions"]["general"]["source_id"] = "unknown.series"
        with self.assertRaisesRegex(ValueError, "알 수 없는 출처"):
            self.norm(spec)

    def test_missing_base_year_key_rejected(self):
        spec = _spec()
        del spec["price_assumptions"]["general"]["base_year"]
        with self.assertRaisesRegex(ValueError, "기준연도"):
            self.norm(spec)

    def test_base_year_mismatch_rejected(self):
        spec = _spec()
        spec["price_assumptions"]["general"]["base_year"] = 2015
        with self.assertRaisesRegex(ValueError, "기준연도 혼합"):
            self.norm(spec)

    def test_student_source_null_base_year_allowed(self):
        spec = _spec(
            price_sources=[_student_source()],
        )
        spec["price_assumptions"]["sales"] = {
            "source_id": "student.sample.series",
            "base_year": None,
            "application_base_year": 2027,
        }
        a = self.norm(spec)["sales"]
        self.assertEqual("applied", a["status"])
        self.assertIsNone(a["base_year"])
        self.assertAlmostEqual(
            a["rate"], self.pa.cagr(100.0, 108.16, 2), places=12)

    def test_nonpositive_observation_rejected(self):
        bad = _student_source(observations=[
            {"year": 2023, "value": 0},
            {"year": 2025, "value": 108.0}])
        spec = _spec(price_sources=[bad])
        spec["price_assumptions"]["sales"] = {
            "source_id": "student.sample.series",
            "base_year": None,
            "application_base_year": 2027,
        }
        with self.assertRaisesRegex(ValueError, "양수"):
            self.norm(spec)

    def test_zero_interval_rejected(self):
        spec = _spec()
        spec["price_assumptions"]["general"]["observation"] = {
            "start_year": 2022, "end_year": 2022}
        with self.assertRaisesRegex(ValueError, "연간격"):
            self.norm(spec)

    def test_unconfirmed_observation_year_rejected(self):
        spec = _spec()
        spec["price_assumptions"]["general"]["observation"] = {
            "start_year": 2020, "end_year": 2019}  # 역방향도 거부
        with self.assertRaisesRegex(ValueError, "연간격"):
            self.norm(spec)
        spec["price_assumptions"]["general"]["observation"] = {
            "start_year": 2020, "end_year": 2040}  # 없는 연도
        with self.assertRaisesRegex(ValueError, "확인 불가"):
            self.norm(spec)

    def test_mixed_source_id_collision_rejected(self):
        """번들 id를 학생 블록이 다른 기준연도로 재정의하면 혼합 비교 거부."""
        collision = _student_source(
            "official.kosis.cpi.total",
            agency="다른기관", base_year=2015, unit="지수(2015=100)")
        spec = _spec(price_sources=[collision])
        with self.assertRaisesRegex(ValueError, "혼합"):
            self.norm(spec)

    def test_sales_aggregate_index_rejected(self):
        """판매가에 전국 총지수 자동 대입 금지 — 명시 선택해도 거부."""
        spec = _spec()
        spec["price_assumptions"]["sales"] = {
            "source_id": "official.kosis.farm_sales.total",
            "base_year": 2020,
            "application_base_year": 2027,
        }
        with self.assertRaisesRegex(ValueError, "사용 불가"):
            self.norm(spec)

    def test_role_mismatch_rejected(self):
        """노무비 지수를 일반물가 역할에 대입 금지."""
        spec = _spec()
        spec["price_assumptions"]["general"] = {
            "source_id": "official.kosis.farm_purchase.labor",
            "base_year": 2020,
            "application_base_year": 2027,
        }
        with self.assertRaisesRegex(ValueError, "사용 불가"):
            self.norm(spec)

    def test_sales_not_applied_requires_reason(self):
        spec = _spec()
        spec["price_assumptions"]["sales"] = {"status": "not_applied"}
        with self.assertRaisesRegex(ValueError, "사유"):
            self.norm(spec)

    def test_not_applied_is_explicit_choice(self):
        a = self.norm(_spec())["sales"]
        self.assertEqual("not_applied", a["status"])
        self.assertEqual(1.0, self.pa.multiplier(a, 2029))
        self.assertIn("미적용", self.pa.footnote(a))

    def test_negative_rate_allowed(self):
        spec = _spec(
            price_sources=[_student_source(observations=[
                {"year": 2021, "value": 59080},
                {"year": 2025, "value": 58872}])])
        spec["price_assumptions"]["sales"] = {
            "source_id": "student.sample.series",
            "base_year": None,
            "application_base_year": 2027,
        }
        a = self.norm(spec)["sales"]
        self.assertLess(a["rate"], 0)

    def test_rate_floor_rejected(self):
        spec = _spec()
        spec["price_assumptions"]["general"]["rate"] = -1.0
        with self.assertRaisesRegex(ValueError, "배율"):
            self.norm(spec)

    def test_application_base_year_must_precede_plan(self):
        spec = _spec()
        spec["price_assumptions"]["general"]["application_base_year"] = 2028
        with self.assertRaisesRegex(ValueError, "기초연도"):
            self.norm(spec)

    def test_application_count_boundary_3_vs_2(self):
        """2022→2025는 3회 적용 — 기초연도가 1년 어긋나면 배율이 한 번 다름."""
        spec = _spec()
        a2 = self.norm(_spec(**{
            "price_assumptions": dict(spec["price_assumptions"], general={
                "source_id": "official.kosis.cpi.total", "base_year": 2020,
                "application_base_year": 2025})}))["general"]
        a3 = self.norm(_spec(**{
            "price_assumptions": dict(spec["price_assumptions"], general={
                "source_id": "official.kosis.cpi.total", "base_year": 2020,
                "application_base_year": 2024})}))["general"]
        m2 = self.pa.multiplier(a2, 2027)
        m3 = self.pa.multiplier(a3, 2027)
        self.assertAlmostEqual(m3 / m2, 1.0 + a2["rate"], places=12)
        self.assertEqual(2, self.pa.applications(a2, 2027))
        self.assertEqual(3, self.pa.applications(a3, 2027))


class SeparateRatesModelTests(ContractCase):
    """세 률이 서로 다를 때 자재·노무·매출이 각자 자기 률로 연도별 증가."""

    def setUp(self):
        self.pa = runtime("gg_price_assumptions")
        self.sx = runtime("gg_school_excel")
        self.verify = runtime("gg_school_verify")
        self.mc = runtime("gg_major_contract")

    def _divergent_spec(self):
        """상승률은 명시하지 않는다 — 수열의 확정 관측 CAGR이 적용률(V1-01)."""
        return _spec(
            price_sources=[_student_source()],
            price_assumptions={
                "general": {
                    "source_id": "official.kosis.cpi.total",
                    "base_year": 2020,
                    "application_base_year": 2027},
                "wage": {
                    "source_id": "official.kosis.farm_purchase.labor",
                    "base_year": 2020,
                    "application_base_year": 2027},
                "sales": {
                    "source_id": "student.sample.series",
                    "base_year": None,
                    "application_base_year": 2027},
            })

    def test_three_rates_diverge_per_year(self):
        rg = self.pa.cagr(100.0, 116.61, 5)   # CPI 2020→2025
        rw = self.pa.cagr(100.0, 137.9, 5)    # 농가구입가격지수 노무비 2020→2025
        rs = self.pa.cagr(100.0, 108.16, 2)   # 합성 수매단가 2023→2025
        exp = self.verify.calculate_school_expected(self._divergent_spec())
        yrs = exp["expected_years"]
        for i in range(5):
            self.assertAlmostEqual(
                float(yrs[i]["materials"]), 500 * (1 + rg)**i, places=6)
            self.assertAlmostEqual(
                float(yrs[i]["labor"]), 800 * (1 + rw)**i, places=6)
            base_rev = 2000 * 0.5 * 3 / 1000 + 2000 * 0.5 * 5 / 1000
            self.assertAlmostEqual(
                float(yrs[i]["revenue"]), base_rev * (1 + rs)**i, places=6)

    def test_no_double_application_on_labor(self):
        """노무비는 임금률만 — 같은 총액에 일반물가 재적용 없음."""
        rg = self.pa.cagr(100.0, 116.61, 5)
        rw = self.pa.cagr(100.0, 137.9, 5)
        exp = self.verify.calculate_school_expected(self._divergent_spec())
        self.assertAlmostEqual(
            float(exp["expected_years"][4]["labor"]),
            800 * (1 + rw)**4, places=6)
        # 일반물가까지 곱해졌다면 800*((1+rw)*(1+rg))^4와 달라야 한다
        self.assertNotAlmostEqual(
            float(exp["expected_years"][4]["labor"]),
            800 * ((1 + rw) * (1 + rg))**4, places=6)

    def test_workbook_formulas_use_separate_rates(self):
        spec = self._divergent_spec()
        norm = self.pa.normalize_price_assumptions(spec, YEARS)
        step_g = 1.0 + norm["general"]["rate"]
        step_w = 1.0 + norm["wage"]["rate"]
        ms_2028 = self.pa.multiplier(norm["sales"], 2028)
        root = self.make_project()
        bind_major(root, "specialty_crops")
        ctx = self.mc.output_context(root, "specialty_crops")
        out = Path(root) / "school.xlsx"
        self.sx.school_workbook(spec, out, context=ctx)
        wb = load_workbook(out)
        mat = wb["7. 영농자재소요계획"]
        labor = wb["8. 노무비계획"]
        sales = wb[" 5. 판매계획"]
        exp = wb["9 .경비계획"]
        self.assertEqual(f"=C4*{step_g}", mat["C5"].value)
        self.assertEqual(f"=C4*{step_w}", labor["C5"].value)
        self.assertEqual(f"=G22*$O$7/1000*{ms_2028}", sales["E39"].value)
        # 수도광열비는 일반물가와 같은 률(2년차 블록 C38 = 전년 C18 ×(1+r_g))
        self.assertEqual(f"=C18*{step_g}", exp["C38"].value)
        # 각주에 통계명·기준연도·관측기간·산식이 드러남
        self.assertIn("소비자물가지수", mat["B12"].value)
        self.assertIn("기준연도 2020", labor["B12"].value)
        self.assertIn("농가판매및구입가격조사", labor["B12"].value)
        self.assertIn("산식", sales["B42"].value)
        self.assertIn("출처:", exp["B31"].value)

    def test_sales_not_applied_note_is_visible(self):
        root = self.make_project()
        bind_major(root, "specialty_crops")
        ctx = self.mc.output_context(root, "specialty_crops")
        out = Path(root) / "school.xlsx"
        self.sx.school_workbook(_spec(), out, context=ctx)
        wb = load_workbook(out)
        sales = wb[" 5. 판매계획"]
        self.assertIn("판매가 상승률 미적용(명시 선택)", sales["B42"].value)
        # 미적용이면 5년 모두 같은 가격 수식(배율 접미사 없음)
        self.assertEqual("=E22*$O$7/1000", sales["D39"].value)
        self.assertEqual("=M22*$O$7/1000", sales["H39"].value)

    def test_manifest_records_assumptions(self):
        root = self.make_project()
        bind_major(root, "specialty_crops")
        ctx = self.mc.output_context(root, "specialty_crops")
        out = Path(root) / "school.xlsx"
        self.sx.school_workbook(_spec(), out, context=ctx)
        manifest = json.loads(
            out.with_suffix(".manifest.json").read_text(encoding="utf-8"))
        pa = manifest["price_assumptions"]
        self.assertEqual("applied", pa["general"]["status"])
        self.assertEqual("official.kosis.cpi.total", pa["general"]["source_id"])
        self.assertEqual("not_applied", pa["sales"]["status"])


class V1CounterexampleTests(ContractCase):
    """V1-01~04 독립검토 반례를 잠그는 회귀.

    재현 입력: work/V1-scratch/probe_prices.py / price-probes.json.
    """

    def setUp(self):
        self.pa = runtime("gg_price_assumptions")

    def norm(self, spec):
        return self.pa.normalize_price_assumptions(spec, YEARS)

    def _sales_applied(self, sid, base_year=None, **extra):
        a = {"source_id": sid, "base_year": base_year,
             "application_base_year": 2027}
        a.update(extra)
        return a

    # ---- V1-01: r은 관측 수열의 CAGR로만 ----

    def test_v1_01_declared_rate_off_cagr_rejected(self):
        """probe declared_90pct_CPI: 임의 r=0.9를 CPI에 선언해도 거부."""
        spec = _spec()
        spec["price_assumptions"]["general"]["rate"] = 0.9
        with self.assertRaisesRegex(ValueError, "불일치"):
            self.norm(spec)

    def test_v1_01_declared_rate_within_tolerance_is_crosscheck(self):
        """관측 CAGR과 허용오차 내 선언은 검산으로 받되 적용값은 CAGR."""
        ref = self.pa.cagr(100.0, 116.61, 5)
        spec = _spec()
        spec["price_assumptions"]["general"]["rate"] = 0.0312
        a = self.norm(spec)["general"]
        self.assertAlmostEqual(ref, a["rate"], places=12)
        self.assertEqual(0.0312, a["declared_rate"])
        self.assertEqual("series_cagr", a["rate_source"])

    def test_v1_01_declared_rate_string_within_tolerance(self):
        spec = _spec()
        spec["price_assumptions"]["general"]["rate"] = " 0.0312 "
        a = self.norm(spec)["general"]
        self.assertEqual(0.0312, a["declared_rate"])

    def test_v1_01_no_observations_rate_rejected(self):
        """probe no_observations_20pct: 관측 없는 수열에 r만 둘 수 없다."""
        spec = _spec(price_sources=[_student_source(observations=[])])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series", rate=0.2)
        with self.assertRaisesRegex(ValueError, "관측"):
            self.norm(spec)

    def test_v1_01_single_observation_rejected(self):
        spec = _spec(price_sources=[_student_source(observations=[
            {"year": 2025, "value": 108.16}])])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "미만"):
            self.norm(spec)

    def test_v1_01_window_endpoint_not_observed_rejected(self):
        """관측 구간 끝점이 확정 관측에 없으면 거부(보간·대체 금지)."""
        spec = _spec()
        spec["price_assumptions"]["general"]["observation"] = {
            "start_year": 2018, "end_year": 2025}  # 2018 미관측
        with self.assertRaisesRegex(ValueError, "확인 불가"):
            self.norm(spec)
        # 부분 구간 2020~2024는 끝점이 확정 관측이라 허용
        spec["price_assumptions"]["general"]["observation"] = {
            "start_year": 2020, "end_year": 2024}
        a = self.norm(spec)["general"]
        self.assertAlmostEqual(
            self.pa.cagr(100.0, 114.18, 4), a["rate"], places=12)

    # ---- V1-02: 공식 출처 계약 ----

    def test_v1_02_private_blog_rejected(self):
        """probe private_blog_source: 비공식 출처는 선언해도 거부."""
        blog = _student_source(
            agency="개인 블로그", agency_kind=None,
            url="https://example.com/blog/price")
        spec = _spec(price_sources=[blog])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "agency_kind|기관 종류"):
            self.norm(spec)

    def test_v1_02_official_kind_wrong_domain_rejected(self):
        """기관 종류를 꾸며도 url 도메인이 비공식이면 거부."""
        blog = _student_source(url="https://blog.example.com/p")
        spec = _spec(price_sources=[blog])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "url"):
            self.norm(spec)

    def test_v1_02_http_url_rejected(self):
        blog = _student_source(url="http://kosis.kr/insecure")
        spec = _spec(price_sources=[blog])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "url"):
            self.norm(spec)

    def test_v1_02_missing_locator_rejected(self):
        spec = _spec(price_sources=[
            _student_source(evidence_locator=None)])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "위치"):
            self.norm(spec)

    def test_v1_02_missing_dimensions_rejected(self):
        spec = _spec(price_sources=[_student_source(dimensions=None)])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "차원"):
            self.norm(spec)

    def test_v1_02_allowed_roles_string_rejected(self):
        """probe allowed_roles_string_substring: 문자열에 부분일치 우회 금지."""
        spec = _spec(price_sources=[
            _student_source(allowed_roles="nonsales")])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "allowed_roles"):
            self.norm(spec)

    def test_v1_02_missing_allowed_roles_rejected(self):
        entry = _student_source()
        del entry["allowed_roles"]
        spec = _spec(price_sources=[entry])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "allowed_roles"):
            self.norm(spec)

    def test_v1_02_unconfirmed_overlap_rejected(self):
        """probe unconfirmed_value_present: 관측·미확인 연도 교집합 금지."""
        spec = _spec(price_sources=[
            _student_source(unconfirmed_years=[2025])])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "겹침"):
            self.norm(spec)

    def test_v1_02_unconfirmed_endpoint_rejected(self):
        """봄감자 2025는 미확인 — 관측 구간 끝점으로 못 쓴다."""
        spec = _spec(crops=["봄감자"], region="전국")
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "official.rda.crop_income.spring_potato.output_price",
            observation={"start_year": 2020, "end_year": 2025})
        with self.assertRaisesRegex(ValueError, "확인 불가"):
            self.norm(spec)

    def test_v1_02_source_id_case_or_space_rejected(self):
        """id 정규화 우회 금지: 대소문자·공백 변형은 미등록 id."""
        for variant in ("OFFICIAL.KOSIS.CPI.TOTAL",
                        " official.kosis.cpi.total "):
            spec = _spec()
            spec["price_assumptions"]["general"]["source_id"] = variant
            with self.assertRaisesRegex(ValueError, "알 수 없는 출처",
                                        msg=variant):
                self.norm(spec)

    # ---- V1-03: 판매가 수열 차원·역할 ----

    def test_v1_03_kamis_rice_sales_rejected(self):
        """probe rice_retail_to_unrelated_crop: 쌀 소매가를 타 작목에 대입 금지."""
        spec = _spec(crops=["봄감자"], region="전국")
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "official.kamis.retail.rice_20kg.top")
        with self.assertRaisesRegex(ValueError, "사용 불가"):
            self.norm(spec)

    def test_v1_03_dimension_mismatch_rejected(self):
        """probe potato_wrong_region_cultivation: 지역·재배형태 불일치 거부."""
        spec = _spec(crops=["봄감자"], region="제주", cultivation="시설")
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "official.rda.crop_income.spring_potato.output_price")
        with self.assertRaisesRegex(ValueError, "불일치"):
            self.norm(spec)

    def test_v1_03_target_dimensions_on_assumption(self):
        """target_dimensions를 가정 블록에 선언해도 같은 규칙 적용."""
        spec = _spec(crops=["봄감자"])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "official.rda.crop_income.spring_potato.output_price",
            target_dimensions={"channel": "소매"})
        with self.assertRaisesRegex(ValueError, "불일치"):
            self.norm(spec)

    def test_v1_03_series_without_crop_rejected(self):
        """sales 수열은 작목 차원을 명시해야 계획과 대조 가능."""
        entry = _student_source(dimensions={"region": "전국"})
        spec = _spec(price_sources=[entry])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "crop|차원"):
            self.norm(spec)

    def test_v1_03_matching_dimensions_accepted(self):
        """봄감자 전국·노지·농가수취·원/kg 대상은 감자 수열과 일치."""
        spec = _spec(
            crops=["봄감자"], region="전국",
            price_dimensions={
                "cultivation": "년 1기작/10a",
                "unit": "원/kg",
                "channel": "농가수취(주산물 단가)"})
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "official.rda.crop_income.spring_potato.output_price")
        a = self.norm(spec)["sales"]
        self.assertEqual("applied", a["status"])
        self.assertAlmostEqual(
            self.pa.cagr(898.0, 1190.0, 4), a["rate"], places=12)

    def test_v1_03_crop_alias_normalization(self):
        """별칭·한글 차원 키도 정본 키로 대조된다."""
        entry = _student_source(dimensions={
            "작목": "합성작목", "거래단계": "농가수취", "단위": "원/kg"})
        spec = _spec(price_sources=[entry],
                     price_dimensions={"상품": "합성작목"})
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series",
            target_dimensions={"crop": ["합성작목"]})
        # 상품(product) 차원이 계획에만 있고 수열에 없음 → 거부
        with self.assertRaisesRegex(ValueError, "차원|명시"):
            self.norm(spec)

    # ---- V1-04: 유한성 ----

    def test_v1_04_nonfinite_observation_rejected(self):
        for bad in (float("inf"), float("-inf"), float("nan")):
            entry = _student_source(observations=[
                {"year": 2023, "value": 100.0},
                {"year": 2025, "value": bad}])
            spec = _spec(price_sources=[entry])
            spec["price_assumptions"]["sales"] = self._sales_applied(
                "student.sample.series")
            with self.assertRaisesRegex(ValueError, "유한", msg=bad):
                self.norm(spec)

    def test_v1_04_json_overflow_observation_rejected(self):
        """1e309 같은 JSON 오버플로 값은 inf로 파싱돼 거부된다."""
        entry = _student_source()
        entry["observations"][1]["value"] = 1e309
        entry = json.loads(json.dumps(entry))  # JSON 왕복 후에도 inf
        self.assertEqual(float("inf"),
                         entry["observations"][1]["value"])
        spec = _spec(price_sources=[entry])
        spec["price_assumptions"]["sales"] = self._sales_applied(
            "student.sample.series")
        with self.assertRaisesRegex(ValueError, "유한"):
            self.norm(spec)

    def test_v1_04_bool_or_string_observation_rejected(self):
        for bad in (True, "104"):
            entry = _student_source(observations=[
                {"year": 2023, "value": 100.0},
                {"year": 2025, "value": bad}])
            spec = _spec(price_sources=[entry])
            spec["price_assumptions"]["sales"] = self._sales_applied(
                "student.sample.series")
            with self.assertRaisesRegex(ValueError, "수치", msg=repr(bad)):
                self.norm(spec)

    def test_v1_04_nonfinite_declared_rate_rejected(self):
        for bad in (float("inf"), float("nan"), " INF ", "NaN", True):
            spec = _spec()
            spec["price_assumptions"]["general"]["rate"] = bad
            with self.assertRaisesRegex(ValueError, "유한|수치",
                                        msg=repr(bad)):
                self.norm(spec)

    def test_v1_04_multiplier_nonfinite_rejected(self):
        a = {"role": "general", "status": "applied", "rate": 2.0,
             "application_base_year": 1000}
        with self.assertRaisesRegex(ValueError, "유한하지 않음"):
            self.pa.multiplier(a, 3000)
