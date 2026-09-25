"""P9 layout-guard contracts for gg_excel_template (C6).

All workbooks are synthetic; no professor or school file is read.  The
anchors come from the shipped C1 extract, so the stand-ins carry only
the two distinguishing cells (계 total row, depreciation column J7).
"""
import json
import tempfile
from pathlib import Path
from unittest import mock

from tests._harness import ContractCase, bind_major, runtime
from tests._harness import write_text

INV = "3.투자계획"
DEP = "10. 감가상각비계획 "


def _workbook(path, *, variant):
    """Minimal two-sheet stand-in carrying the layout anchors."""
    from openpyxl import Workbook
    wb = Workbook()
    inv = wb.active
    inv.title = INV
    dep = wb.create_sheet(DEP)
    if variant == "x01":
        inv["C19"] = "입력"
        inv["C24"] = "계"
    elif variant == "x02":
        inv["C19"] = "입력"
        inv["C25"] = "계"
        dep["J7"] = "대식물"
    wb.save(path)


def _map(path, tpl, source, *, layout=None, status=None,
         entries=None):
    data = {"schema": "gg-xlsx-template-map/v1",
            "source": {"sha256": tpl.sha256(source)},
            "entries": entries or [
                {"sheet": INV, "cell": "C19", "action": "clear",
                 "semanticField": "합성", "reason": "합성",
                 "explicit": True}]}
    if layout is not None:
        data["layout"] = layout
    if status is not None:
        data["mapStatus"] = status
    path.write_text(json.dumps(data, ensure_ascii=False),
                    encoding="utf-8")
    return path


def _ctx(case):
    root = case.make_project()
    bind_major(root)
    return runtime("gg_major_contract").output_context(
        root, "specialty_crops")


class DetectLayoutTests(ContractCase):
    def test_detects_both_variants_and_unknown(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory() as tmp:
            x01 = Path(tmp) / "a.xlsx"
            x02 = Path(tmp) / "b.xlsx"
            other = Path(tmp) / "c.xlsx"
            _workbook(x01, variant="x01")
            _workbook(x02, variant="x02")
            _workbook(other, variant="other")
            self.assertEqual(tpl.detect_layout(x01), "x01")
            self.assertEqual(tpl.detect_layout(x02), "x02")
            self.assertEqual(tpl.detect_layout(other), "unknown")


class InspectLayoutTests(ContractCase):
    def test_x01_marks_map_ok(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x01.xlsx"
            _workbook(src, variant="x01")
            r = tpl.inspect_source(src, None, None)
            self.assertEqual(r["layout"],
                             {"variant": "x01",
                              "mapOrigin": "default_x01"})
            self.assertEqual(r["mapStatus"], "ok")
            self.assertIsNone(r["mapReason"])

    def test_x02_marks_manual_map_required(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x02.xlsx"
            _workbook(src, variant="x02")
            r = tpl.inspect_source(src, None, None)
            self.assertEqual(r["layout"]["variant"], "x02")
            self.assertEqual(r["mapStatus"], "manual_map_required")
            self.assertEqual(r["mapReason"], "layout_variant_unsupported")


class MalformedAnchorTests(ContractCase):
    """A2: an empty or malformed anchor list must never classify — a
    layout-unrecognised source refuses the legacy map, never re-uses
    the X01 coordinates."""

    def setUp(self):
        self.tpl = runtime("gg_excel_template")

    def _refused_with(self, anchors):
        tpl = self.tpl
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x02.xlsx"
            _workbook(src, variant="x02")
            with mock.patch.object(tpl, "_load_layout_anchors",
                                   return_value=anchors):
                self.assertEqual(tpl.detect_layout(src), "unknown")
                # a legacy map has no layout block -> default_x01,
                # refused because detection is unknown
                m = _map(Path(tmp) / "map.json", tpl, src)
                out = Path(tmp) / "out.xlsx"
                with self.assertRaisesRegex(
                        ValueError, "layout_variant_unsupported"):
                    tpl.blank_copy(src, m, out, context=_ctx(self))
                self.assertFalse(out.exists())

    def test_empty_anchor_list_fails_closed(self):
        shipped = self.tpl._load_layout_anchors()
        for variant in ("x01", "x02"):
            with self.subTest(variant=variant):
                self._refused_with(dict(shipped, **{variant: []}))

    def test_malformed_anchors_fail_closed(self):
        shipped = self.tpl._load_layout_anchors()
        malformed = [
            [{"sheet": DEP, "cell": "J7"}],                   # no expect
            [{"sheet": DEP, "cell": "BAD CELL",
              "expect": {"equals": "x"}}],                   # bad cell ref
            [{"sheet": DEP, "cell": "J7",
              "expect": {"bogus": True}}],                   # unknown predicate
            [{"sheet": DEP, "cell": "J7",
              "expect": {"equals": "x", "present": True}}],  # two predicates
            [{"sheet": 7, "cell": "J7", "expect": {"present": True}}],
        ]
        for anchors in malformed:
            with self.subTest(anchors=anchors):
                self._refused_with(dict(shipped, x02=anchors))


class FormulaAwareAnchorTests(ContractCase):
    """A3: a formula cell without a cached value counts as present,
    never empty — an X02 whose J7 formula cache is invalidated still
    detects as x02."""

    def test_uncached_j7_formula_detects_x02(self):
        tpl = runtime("gg_excel_template")
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x02f.xlsx"
            wb = Workbook()
            inv = wb.active
            inv.title = INV
            dep = wb.create_sheet(DEP)
            inv["C19"] = "입력"
            inv["C25"] = "계"
            dep["J7"] = "=SUM(1,2)"   # formula, no cached value
            wb.save(src)
            self.assertEqual(tpl.detect_layout(src), "x02")
            # a reviewed_custom map for the detected variant clears; the
            # blank copy must still detect x02 after cache invalidation
            m = _map(Path(tmp) / "map.json", tpl, src,
                     layout={"variant": "x02",
                             "mapOrigin": "reviewed_custom"})
            out = Path(tmp) / "out.xlsx"
            receipt = tpl.blank_copy(src, m, out, context=_ctx(self))
            self.assertTrue(out.exists())
            self.assertEqual([c["cell"] for c in receipt["cleared"]],
                             ["C19"])
            self.assertEqual(tpl.detect_layout(out), "x02")


class ClearLayoutGuardTests(ContractCase):
    def test_legacy_map_on_x02_refused(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x02.xlsx"
            _workbook(src, variant="x02")
            # a legacy map has no layout block -> treated as default_x01
            m = _map(Path(tmp) / "map.json", tpl, src)
            out = Path(tmp) / "out.xlsx"
            with self.assertRaisesRegex(
                    ValueError, "layout_variant_unsupported"):
                tpl.blank_copy(src, m, out, context=_ctx(self))
            self.assertFalse(out.exists())

    def test_status_edited_to_ok_still_refused(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x02.xlsx"
            _workbook(src, variant="x02")
            m = _map(Path(tmp) / "map.json", tpl, src, status="ok")
            out = Path(tmp) / "out.xlsx"
            with self.assertRaisesRegex(
                    ValueError, "layout_variant_unsupported"):
                tpl.blank_copy(src, m, out, context=_ctx(self))
            self.assertFalse(out.exists())

    def test_reviewed_custom_matching_variant_accepted(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x02.xlsx"
            _workbook(src, variant="x02")
            m = _map(Path(tmp) / "map.json", tpl, src,
                     layout={"variant": "x02",
                             "mapOrigin": "reviewed_custom"})
            out = Path(tmp) / "out.xlsx"
            receipt = tpl.blank_copy(src, m, out, context=_ctx(self))
            self.assertTrue(out.exists())
            self.assertEqual(receipt["cleared"][0]["cell"], "C19")

    def test_reviewed_custom_mismatch_refused(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x02.xlsx"
            _workbook(src, variant="x02")
            m = _map(Path(tmp) / "map.json", tpl, src,
                     layout={"variant": "x01",
                             "mapOrigin": "reviewed_custom"})
            out = Path(tmp) / "out.xlsx"
            with self.assertRaisesRegex(
                    ValueError, "map_layout_invalid"):
                tpl.blank_copy(src, m, out, context=_ctx(self))
            self.assertFalse(out.exists())

    def test_x01_anchored_legacy_map_unchanged(self):
        tpl = runtime("gg_excel_template")
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "x01.xlsx"
            _workbook(src, variant="x01")
            m = _map(Path(tmp) / "map.json", tpl, src)
            out = Path(tmp) / "out.xlsx"
            receipt = tpl.blank_copy(src, m, out, context=_ctx(self))
            wb = load_workbook(out)
            try:
                self.assertIsNone(wb[INV]["C19"].value)
            finally:
                wb.close()
            self.assertEqual([c["cell"] for c in receipt["cleared"]],
                             ["C19"])


if __name__ == "__main__":
    import unittest
    unittest.main()
