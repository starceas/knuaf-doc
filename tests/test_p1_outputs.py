"""P1-B behavior tests on generated synthetic outputs.

Every test exercises a produced artifact (DOCX/XLSX) or a fake-COM
boundary — never source-string checks, real Office applications, real
school templates, or user documents.
"""
import contextlib
import io
import json
import os
import subprocess
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

from tests._harness import ContractCase, bind_major, runtime


def _bound_context(case):
    """Policy B: a bound specialty project and the explicit major for a
    guarded output call (inputs stay synthetic)."""
    project = case.make_project()
    bind_major(project)
    return runtime("gg_major_contract").output_context(
        project, "specialty_crops")


class WindowsComBlockTests(ContractCase):
    """C6 (fixed in P1): the Windows COM path must refuse before any COM
    call. Fake pythoncom/win32com modules count every entry point — all
    counters must stay at zero. Baseline red: run_word/run_excel reached
    EnsureDispatch/Quit on main 29abec0."""

    @staticmethod
    def _fake_com():
        calls = {"CoInitialize": 0, "EnsureDispatch": 0,
                 "Dispatch": 0, "Quit": 0}

        def _app():
            return types.SimpleNamespace(
                Quit=lambda: calls.__setitem__("Quit", calls["Quit"] + 1))

        pythoncom = types.ModuleType("pythoncom")
        pythoncom.CoInitialize = lambda: calls.__setitem__(
            "CoInitialize", calls["CoInitialize"] + 1)
        pythoncom.CoUninitialize = lambda: None
        client = types.ModuleType("win32com.client")
        client.gencache = types.SimpleNamespace(
            EnsureDispatch=lambda name: (
                calls.__setitem__(
                    "EnsureDispatch", calls["EnsureDispatch"] + 1),
                _app())[1])
        client.Dispatch = lambda name: (
            calls.__setitem__("Dispatch", calls["Dispatch"] + 1),
            _app())[1]
        win32com = types.ModuleType("win32com")
        win32com.client = client
        return calls, {
            "pythoncom": pythoncom,
            "win32com": win32com,
            "win32com.client": client,
        }

    def _assert_blocked_helper(self, fn):
        calls, fakes = self._fake_com()
        err = io.StringIO()
        with mock.patch.dict(sys.modules, fakes), \
                contextlib.redirect_stderr(err):
            with self.assertRaises(SystemExit) as cm:
                fn("in.docx", "out.pdf")
        self.assertEqual(2, cm.exception.code)
        self.assertIn("차단", err.getvalue())
        self.assertEqual(
            {"CoInitialize": 0, "EnsureDispatch": 0,
             "Dispatch": 0, "Quit": 0},
            calls, "a COM entry point was reached despite the block")

    def test_helper_run_word_blocks_before_com(self):
        win = runtime("gg_office_win")
        self._assert_blocked_helper(win.run_word)

    def test_helper_run_excel_blocks_before_com(self):
        win = runtime("gg_office_win")
        self._assert_blocked_helper(win.run_excel)

    def test_wrapper_returns_block_without_spawning(self):
        office = runtime("gg_office")
        runtime("gg_office_win")
        with mock.patch.object(subprocess, "run") as run_mock:
            res = office.run_windows_com("word", ["a.docx", "b.pdf"])
        run_mock.assert_not_called()
        self.assertEqual(2, res.returncode)
        self.assertIn("차단", res.stderr)

    def test_engine_dispatch_on_win32_returns_block(self):
        office = runtime("gg_office")
        runtime("gg_office_win")
        with mock.patch.object(sys, "platform", "win32"), \
                mock.patch.object(subprocess, "run") as run_mock:
            res_w = office._run_word_engine(Path("a.docx"), Path("b.pdf"))
            res_x = office._run_excel_engine(Path("a.xlsx"), Path("b.pdf"))
        run_mock.assert_not_called()
        self.assertEqual(2, res_w.returncode)
        self.assertEqual(2, res_x.returncode)

    def test_macos_path_not_touched(self):
        """The C6 block must not change the macOS dispatch branch."""
        office = runtime("gg_office")
        with mock.patch.object(sys, "platform", "darwin"), \
                mock.patch.object(office, "ensure_word_foreground_session"), \
                mock.patch.object(office, "run_osascript",
                                  return_value=subprocess.CompletedProcess(
                                      ["osascript"], 0, "", "")) as osh:
            res = office._run_word_engine(Path("a.docx"), Path("b.pdf"))
        osh.assert_called_once()
        self.assertEqual(0, res.returncode)


WIDE_TABLE_MD = """앞 문단이다.

| A | B | C | D | E | F | G | H | I | J | K | L |
|---|---|---|---|---|---|---|---|---|---|---|---|
| 12345 | 2 | 3 | 4 | 5 | 6 | 7 | 8 | 9 | 10 | 11 | 12 |
| 12345 | 2 | 3 | 4 | 5 | 6 | 7 | 8 | 9 | 10 | 11 | 12 |

뒤 문단이다.
"""


class WideTableOrientationTests(ContractCase):
    """H2 (fixed in P1): a wide table must not flip the document section
    to landscape. Asserts on the generated DOCX artifact, not on source.
    Baseline red: fit_table ran with allow_landscape default and rotated
    the section on main 29abec0."""

    def test_wide_table_keeps_portrait_and_content(self):
        build = runtime("build_docx")
        reports = []
        doc = build.convert(WIDE_TABLE_MD, table_reports=reports,
                            context=_bound_context(self))
        self.assertTrue(doc.tables, "expected a rendered table")
        table = doc.tables[0]
        self.assertEqual(12, len(table.columns))
        self.assertEqual("12345", table.rows[1].cells[0].text)
        for sec in doc.sections:
            self.assertLess(
                sec.page_width, sec.page_height,
                "a section was flipped to landscape by a wide table")
        self.assertTrue(
            reports and all(not r.get("landscape") for r in reports),
            f"fit report shows a landscape attempt: {reports!r}")

    def test_narrow_table_still_renders_normally(self):
        build = runtime("build_docx")
        doc = build.convert(
            "문단.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n끝.\n",
            context=_bound_context(self))
        self.assertEqual(1, len(doc.tables))
        self.assertEqual("1", doc.tables[0].rows[1].cells[0].text)
        for sec in doc.sections:
            self.assertLess(sec.page_width, sec.page_height)


SCHOOL_MD = """겉표지
농업전문학사 학위논문
영농창업계획
합성 논문 제목
2026년 2월
한국농수산대학교
특용작물전공
홍길동
표제면
지도교수 김교수
농업전문학사 학위논문
영농창업계획
합성 논문 제목
이 논문을 농업전문학사 학위논문으로 제출함
2026년 1월
한국농수산대학교
특용작물전공
홍길동
제출서
인준서
위 내용을 심사하여 인준함
2026년 2월
위원장 김위원 (인)
위원 이위원 (인)
위원 박위원 (인)
위원 최위원 (인)
목차
요약
합성 요약 본문이다.
감사의 글
Ⅰ. 머리말
본문 첫 문단이다.
"""


class UnrenderedNodeTests(ContractCase):
    """H3 (fixed in P1): nodes between 목차 and 본문 that the renderer does
    not consume must fail with a diagnostic naming the lost content.
    Baseline red: a stray node in the gap was silently dropped on main
    29abec0."""

    def test_stray_node_between_toc_and_body_rejected(self):
        build = runtime("build_docx")
        stray = SCHOOL_MD.replace(
            "목차\n요약", "목차\n채택되지 않은 삽입 단락\n요약")
        with self.assertRaises(ValueError) as cm:
            build.convert(stray, context=_bound_context(self))
        msg = str(cm.exception)
        self.assertIn("렌더되지 않는", msg)
        self.assertIn("채택되지 않은 삽입 단락", msg)

    def test_recognized_gap_nodes_still_render(self):
        build = runtime("build_docx")
        doc = build.convert(SCHOOL_MD, context=_bound_context(self))
        texts = [p.text for p in doc.paragraphs]
        self.assertTrue(
            any("합성 요약 본문이다" in t for t in texts),
            "요약 section content missing from generated document")
        self.assertTrue(
            any("본문 첫 문단이다" in t for t in texts),
            "body content missing from generated document")


class MissingTemplateCellTests(ContractCase):
    """H5 (fixed in P1): a map referencing an absent template cell must
    fail before any partial output is published. Synthetic workbooks only —
    no real school template is used. Baseline red: blank_copy silently
    skipped missing cells on main 29abec0."""

    @staticmethod
    def _source(path):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "시트1"
        ws["A1"] = "지울값"
        ws["B2"] = "보존값"
        wb.save(path)

    @staticmethod
    def _map(folder, tpl, source, entries):
        m = Path(folder) / "map.json"
        m.write_text(json.dumps({
            "schema": "gg-xlsx-template-map/v1",
            "source": {"sha256": tpl.sha256(source)},
            "entries": entries,
        }, ensure_ascii=False), encoding="utf-8")
        return m

    def test_missing_cell_rejected_before_output(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory(prefix="knuaf-h5-") as tmp:
            src = Path(tmp) / "src.xlsx"
            self._source(src)
            m = self._map(tmp, tpl, src, [
                {"sheet": "시트1", "cell": "Z99", "action": "clear",
                 "semanticField": "합성", "reason": "없는 셀 참조",
                 "explicit": True},
            ])
            out = Path(tmp) / "out.xlsx"
            with self.assertRaisesRegex(ValueError, "Z99"):
                tpl.blank_copy(src, m, out, context=_bound_context(self))
            self.assertFalse(
                out.exists(), "partial output must not be published")

    def test_missing_sheet_still_rejected(self):
        tpl = runtime("gg_excel_template")
        with tempfile.TemporaryDirectory(prefix="knuaf-h5-") as tmp:
            src = Path(tmp) / "src.xlsx"
            self._source(src)
            m = self._map(tmp, tpl, src, [
                {"sheet": "없는시트", "cell": "A1", "action": "clear",
                 "semanticField": "합성", "reason": "없는 시트 참조",
                 "explicit": True},
            ])
            out = Path(tmp) / "out.xlsx"
            with self.assertRaisesRegex(ValueError, "없는시트"):
                tpl.blank_copy(src, m, out, context=_bound_context(self))
            self.assertFalse(out.exists())

    def test_existing_cells_cleared_and_preserved(self):
        tpl = runtime("gg_excel_template")
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(prefix="knuaf-h5-") as tmp:
            src = Path(tmp) / "src.xlsx"
            self._source(src)
            m = self._map(tmp, tpl, src, [
                {"sheet": "시트1", "cell": "A1", "action": "clear",
                 "semanticField": "합성", "reason": "지움 대상",
                 "explicit": True},
                {"sheet": "시트1", "cell": "B2", "action": "preserve",
                 "semanticField": "합성", "reason": "보존 대상",
                 "explicit": True},
            ])
            out = Path(tmp) / "out.xlsx"
            receipt = tpl.blank_copy(src, m, out, context=_bound_context(self))
            opened = load_workbook(out)
            ws = opened["시트1"]
            self.assertIsNone(ws["A1"].value)
            self.assertEqual("보존값", ws["B2"].value)
            opened.close()
            self.assertTrue(
                any(c["sheet"] == "시트1" and c["cell"] == "A1"
                    for c in receipt["cleared"]),
                f"cleared receipt missing 시트1!A1: {receipt['cleared']!r}")


class Utf8IoTests(ContractCase):
    """C2 (fixed in P1): text I/O pins UTF-8 instead of the locale default.
    The subprocess starts with UTF-8 filenames, then sets LC_CTYPE=C
    with -X utf8=0 to force ASCII text I/O independently of filenames.
    It runs the real inspect_source (UTF-8 write)
    and blank_copy (UTF-8 read) on Korean content. Baseline red: default
    encoding was used on main 29abec0."""

    def test_korean_map_roundtrip_under_ascii_default(self):
        tpl = runtime("gg_excel_template")
        scripts = str(Path(tpl.__file__).resolve().parent)
        snippet = r"""
import json, locale, sys
from pathlib import Path
sys.path.insert(0, sys.argv[1])
assert sys.flags.utf8_mode == 0
assert sys.getfilesystemencoding().lower().replace("-", "") == "utf8"
locale.setlocale(locale.LC_CTYPE, "C")
enc = locale.getpreferredencoding(False)
print("preferred:", enc)
assert any(t in enc.lower() for t in ("ascii", "ansi")), \
    "band not active: " + enc
root = Path(sys.argv[2])
probe = root / "implicit.txt"
try:
    probe.write_text("한글")
except UnicodeEncodeError:
    pass
else:
    raise AssertionError("implicit Korean text write must fail")
probe.write_text("한글", encoding="utf-8")
try:
    probe.read_text()
except UnicodeDecodeError:
    pass
else:
    raise AssertionError("implicit Korean text read must fail")
import gg_excel_template as tpl
from openpyxl import Workbook
src = root / "원본.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "시트1"
ws["A1"] = "한국어 값"
wb.save(src)
m = root / "map.json"
tpl.inspect_source(src, m, None)
custom = root / "custom.json"
custom.write_text(json.dumps({
    "schema": "gg-xlsx-template-map/v1",
    "source": {"sha256": tpl.sha256(src)},
    "entries": [{"sheet": "시트1", "cell": "A1", "action": "clear",
                 "semanticField": "한국어 필드", "reason": "비ASCII 사유",
                 "explicit": True}],
}, ensure_ascii=False), encoding="utf-8")
out = root / "out.xlsx"
import gg_major_contract as mc
tpl.blank_copy(src, custom, out,
               context=mc.output_context(sys.argv[3], "specialty_crops"))
print("PASS")
"""
        env = dict(os.environ, LC_ALL="C.UTF-8", PYTHONUTF8="0")
        env.pop("PYTHONIOENCODING", None)
        # Policy B: the bound project is prepared in this (UTF-8) process;
        # the ASCII-locale driver only names it as the output context.
        project = self.make_project()
        bind_major(project)
        with tempfile.TemporaryDirectory(prefix="knuaf-c2-") as tmp:
            driver = Path(tmp) / "driver.py"
            driver.write_text(snippet, encoding="utf-8")
            res = subprocess.run(
                [sys.executable, "-X", "utf8=0", str(driver), scripts, tmp,
                 str(project)],
                capture_output=True, text=True, env=env, cwd=tmp)
        self.assertEqual(0, res.returncode,
                         f"stdout={res.stdout}\nstderr={res.stderr}")
        self.assertIn("PASS", res.stdout)


if __name__ == "__main__":
    unittest.main()
