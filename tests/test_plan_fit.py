"""Cultivation-method equipment, grade and channel price fit checks.

Synthetic workbooks reproduce only the labels of the school 17-sheet layout
that ``gg_plan_fit`` reads. All names and numbers are synthetic.
"""
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook

from tests._harness import SCRIPTS, ContractCase, runtime
from tests.test_skill_school_verify import _source, base_spec

GRADE_ROWS = ("상품", "중품", "하품")


ASSET_BLOCKS = (  # (제목 행, 제목, 금액 열, 금액 머리글) — 학과 17시트 양식 배치
    (6, "(1) 토지", 5, "평가액(천원)"), (12, "(2) 건물및구축물", 5, "평가액(천원)"),
    (18, "(3) 대식물", 5, "평가액(천원)"), (24, "(4) 대동물", 8, "금액(천원)"),
    (30, "(5) 대농기계 및 차량운반구 등", 5, "구입가격(천원)"),
    (36, "(6) 현금 및 준현금", 5, "평가액(천원)"), (43, "(7) 재고생산자재", 6, "금액(천원)"),
    (49, "(8) 재고농산물", 6, "평가액(천원)"))
YEARS = (2027, 2028, 2029, 2030, 2031)
# Sheets the checker does not read; their B2 title and one key header identify
# the school template (X01). Same cells as the template.
X01_OTHER_SHEETS = (
    ("목록", "작 성 양 식", "E8", "1. 기초재무상태조사"),
    ("2. 중장기영농목표", "2. 중장기 영농목표", "B9", "구분"),
    ("4. 원리금상환계획", "4. 원리금 상환계획", "B6", "연도"),
    ("6. 생산계획", "6. 생산계획", "B6", "구분"),
    ("7. 영농자재소요계획", "7. 영농자재 소요계획", "B4", "적용기준"),
    ("9 .경비계획", "9. 경비계획", "B4", "항목"),
    ("11. 생산원가계획", "11. 생산원가계획", "B4", "과목"),
    ("13. 추정대차대조표", "13. 추정대차대조표", "B5", "과목"),
    ("14. 현금흐름계획", "14. 현금흐름계획", "B5", "구분"),
    ("15.추정소득분석", "15. 추정소득분석", "B4", "비목별"),
)


def plan_book(path, *, investments=(("비닐하우스", 50000), ("고설베드", 20000)),
              opening=(), depreciation=(("비닐하우스", 50000),),
              tasks=(("정식·베드관리", 5), ("수확 및 선별", 30)),
              prices=None, history=None, shares=None, notes=(),
              sga=(0, 0, 0, 0, 0), formula_cells=()):
    """Build a synthetic workbook with the school template's table counts.

    Only labels and positions the checker reads are reproduced: eight asset
    blocks, five investment year cells, five depreciation year tables, two
    labour tables and five plan years. ``opening`` fills the land block,
    ``investments``/``depreciation``/``tasks`` fill the first year/table."""
    prices = prices or {("purchase", "상품"): 3000, ("direct", "상품"): 6000}
    shares = shares or {("purchase", "상품"): 0.7, ("direct", "상품"): 0.3}
    history = history or {}
    wb = Workbook()
    base = wb.active
    base.title = "1. 기초재무상태조사"
    base["B4"] = "가. 자산"
    for title_row, title, col, head in ASSET_BLOCKS:
        # 양식 머리글: B 종류, C 명칭·주소, D·E·F·G 수량·금액·연도, H 비고(대동물은 H 금액).
        base.cell(title_row, 2, title)
        for c, text in zip(range(2, 9), ("종류", "명칭", "수량", "구분", "연수", "연도", "비고")):
            base.cell(title_row + 2, c, text)
        base.cell(title_row + 2, col, head)
    for i, (name, amount, note) in enumerate(opening):
        base.cell(9 + i, 2, name)
        base.cell(9 + i, 5, amount)
        base.cell(9 + i, 8, note)
    base["B55"] = "나. 부채"

    inv = wb.create_sheet("3.투자계획")
    inv["B17"], inv["C17"], inv["E17"], inv["I17"] = "연도별", "투자대상", "투자금액(천원)", "비   고"
    inv["C18"], inv["D18"], inv["E18"], inv["F18"], inv["G18"], inv["H18"] = (
        "투자항목", "규격", "계", "자부담", "보조", "융자")
    inv["B19"] = YEARS[0]
    for i, (name, amount) in enumerate(investments):
        inv.cell(19 + i, 3, name)
        inv.cell(19 + i, 5, amount)
    total = 19 + len(investments)
    inv.cell(total, 3, "계")
    inv.cell(total, 5, sum(a for _, a in investments if isinstance(a, (int, float))))
    for i, year in enumerate(YEARS[1:], start=1):
        inv.cell(total + i, 2, year)
        inv.cell(total + i, 5, 0)

    sales = wb.create_sheet(" 5. 판매계획")
    hist_cols = (5, 7, 9, 11, 13)
    for channel, row, label in (("purchase", 7, "평균가격(수매)"),
                                ("direct", 11, "평균가격(직거래)")):
        for c, year in zip(hist_cols, range(2021, 2026)):
            sales.cell(row - 1, c, year)
        sales.cell(row - 1, 15, "평균가격")
        sales.cell(row, 2, label)
        for offset, grade in enumerate(GRADE_ROWS):
            sales.cell(row + offset, 3, grade)
            price = prices.get((channel, grade))
            if price is not None:
                sales.cell(row + offset, 15, price)
            for c, value in zip(hist_cols, history.get((channel, grade), ())):
                sales.cell(row + offset, c, value)
    for i, text in enumerate(notes):
        sales.cell(14 + i, 2, text)
    for c, year in zip(range(4, 14, 2), YEARS):
        sales.cell(19, c, "%d년" % year)
        sales.cell(20, c, "판매비율")
        sales.cell(20, c + 1, "판매량")
    sales.cell(19, 14, "비고")
    for channel, row, label in (("purchase", 22, "수매시장"),
                                ("direct", 26, "직거래")):
        sales.cell(row, 2, label)
        for offset, grade in enumerate(GRADE_ROWS):
            sales.cell(row + offset, 3, grade)
            for c in range(4, 14, 2):
                sales.cell(row + offset, c, shares.get((channel, grade), 0))

    labor = wb.create_sheet("8. 노무비계획")
    for title_row, title, rows in ((22, "나. 노동비 산출근거(1년차)", tasks),
                                   (36, "다. 노동비 산출근거(2년차 이후)", ())):
        labor.cell(title_row, 2, title)
        labor.cell(title_row + 2, 2, "항목")
        labor.cell(title_row + 2, 9, "작업일수")
        for i, (name, days) in enumerate(rows):
            labor.cell(title_row + 4 + i, 2, name)
            labor.cell(title_row + 4 + i, 9, days)
        labor.cell(title_row + 4 + max(len(rows), 1), 2, "합계")

    dep = wb.create_sheet("10. 감가상각비계획 ")
    for block, year in enumerate(YEARS):
        top = 4 + 12 * block
        dep.cell(top, 2, "%s. %d년" % ("가나다라마"[block], year))
        dep.cell(top + 2, 2, "구       분")
        dep.cell(top + 2, 3, "영농시설")
        dep.cell(top + 4, 2, "취  득  가  격")
        if block == 0:
            for i, (name, cost) in enumerate(depreciation):
                dep.cell(top + 3, 3 + i, name)
                dep.cell(top + 4, 3 + i, cost)

    pl = wb.create_sheet("12 손익계획")
    pl["B4"] = "과목"
    for c, year in enumerate(YEARS, start=3):
        pl.cell(4, c, "%d년" % year)
    pl["B11"], pl["B12"], pl["B13"] = "IV. 판매비와 관리비", " 1. 포장비", " 2. 운반비"
    for c, value in enumerate(sga, start=3):
        pl.cell(13, c, value)
    pl["B15"] = "V. 영업이익"

    for name, title, cell, head in X01_OTHER_SHEETS:
        other = wb.create_sheet(name)
        other["B2"], other[cell] = title, head

    for sheet, cell, formula in formula_cells:
        wb[sheet][cell] = formula
    wb.save(path)
    return path


class _PlanFitCase(ContractCase):
    def setUp(self):
        self.pf = runtime("gg_plan_fit")
        tmp = tempfile.TemporaryDirectory(prefix="knuaf-plan-fit-")
        self.addCleanup(tmp.cleanup)
        self.dir = Path(tmp.name)
        self.count = 0

    GOOD_NOTES = (
        "* 수매 단가: 합성 기관 농가수취가 5개년 평균. 등급 구분 없는 자료라 상품 단일 등급, 중·하품 0은 계획 가정.",
        "* 직거래 단가: 합성 기관 소매가 5개년 평균. 택배비는 판매관리비 운반비로 계상.",
    )

    def review(self, method="protected_soilless", allow=None, edit=None, **book):
        book.setdefault("notes", self.GOOD_NOTES)
        book.setdefault("sga", (100, 200, 300, 300, 300))
        self.count += 1
        path = plan_book(self.dir / ("plan-%d.xlsx" % self.count), **book)
        if edit:
            wb = load_workbook(path)
            edit(wb)
            wb.save(path)
            wb.close()
        return self.pf.review_fit(path, method=method, allow=allow)

    @staticmethod
    def codes(result, severity=None):
        return [i["code"] for i in result["issues"]
                if severity is None or i["severity"] == severity]


class PlanFitTests(_PlanFitCase):
    def test_plan_matching_its_method_passes(self):
        result = self.review()
        self.assertEqual("pass", result["status"], result["issues"])
        self.assertEqual([], self.codes(result))

    def test_zero_amount_equipment_row_is_blocked_in_every_sheet(self):
        result = self.review(
            investments=(("비닐하우스", 50000), ("관리기", 0)),
            depreciation=(("비닐하우스", 50000), ("트랙터(미보유)", 0)))
        unowned = [i["evidence"]["cell"] for i in result["issues"]
                   if i["code"] == "UNOWNED_ROW"]
        self.assertEqual(["C20", "D7"], unowned)
        self.assertEqual("fail", result["status"])

    def test_tillage_machine_depends_on_cultivation_method(self):
        rows = (("비닐하우스", 50000), ("소형 관리기", 3000))
        soilless = self.review(investments=rows)
        self.assertIn("EQUIPMENT_METHOD_MISMATCH", self.codes(soilless, "error"))
        soil = self.review(method="시설토경", investments=rows,
                           tasks=(("경운", 4), ("수확", 30)))
        self.assertEqual("pass", soil["status"], soil["issues"])
        self.assertEqual([], self.codes(soil))

    def test_open_field_rejects_hydroponic_equipment(self):
        result = self.review(method="노지",
                             investments=(("양액공급기", 8000),),
                             depreciation=(("양액공급기", 8000),))
        self.assertIn("EQUIPMENT_METHOD_MISMATCH", self.codes(result, "error"))

    def test_student_confirmed_exception_downgrades_to_warning(self):
        result = self.review(
            investments=(("비닐하우스", 50000), ("소형 관리기", 3000)),
            allow={"소형 관리기": "하우스 바닥 정리에 사용(학생 확인)"})
        self.assertEqual(["EQUIPMENT_METHOD_EXCEPTION"], self.codes(result))
        self.assertEqual("pass", result["status"])

    def test_tillage_task_with_days_fails_and_empty_label_warns(self):
        worked = self.review(tasks=(("경운", 6), ("수확", 30)))
        self.assertIn("TASK_METHOD_MISMATCH", self.codes(worked, "error"))
        label = self.review(tasks=(("경운", 0), ("수확", 30)))
        self.assertEqual(["TASK_METHOD_LABEL"], self.codes(label))
        self.assertEqual("pass", label["status"])

    def test_opening_zero_row_is_a_warning_with_its_note(self):
        result = self.review(opening=(("전", 0, "첫해 매입 예정"),))
        issue = next(i for i in result["issues"] if i["code"] == "OPENING_ZERO_ROW")
        self.assertEqual("warning", issue["severity"])
        self.assertEqual("첫해매입예정", issue["evidence"]["note"])

    def test_unsold_grade_must_not_carry_a_copied_price(self):
        copied = self.review(prices={("purchase", "상품"): 3000, ("purchase", "중품"): 3000,
                                     ("direct", "상품"): 6000})
        cells = [i["evidence"]["cell"] for i in copied["issues"]
                 if i["code"] == "SALES_UNUSED_PRICE"]
        self.assertEqual(["O8"], cells)
        history_only = self.review(history={("direct", "하품"): (1, 2, 3, 4, 5)})
        self.assertIn("SALES_UNUSED_PRICE", self.codes(history_only))
        zero = self.review(prices={("purchase", "상품"): 3000, ("purchase", "중품"): 0,
                                   ("direct", "상품"): 6000})
        self.assertNotIn("SALES_UNUSED_PRICE", self.codes(zero))

    def test_sold_grades_need_distinct_descending_prices(self):
        shares = {("purchase", "상품"): 0.5, ("purchase", "중품"): 0.2,
                  ("direct", "상품"): 0.3}
        equal = self.review(shares=shares, prices={
            ("purchase", "상품"): 3000, ("purchase", "중품"): 3000.02,
            ("direct", "상품"): 6000})
        self.assertIn("GRADE_PRICE_EQUAL", self.codes(equal, "error"))
        reversed_ = self.review(shares=shares, prices={
            ("purchase", "상품"): 2000, ("purchase", "중품"): 2500,
            ("direct", "상품"): 6000})
        self.assertIn("GRADE_PRICE_ORDER", self.codes(reversed_, "error"))
        ordered = self.review(shares=shares, prices={
            ("purchase", "상품"): 3000, ("purchase", "중품"): 2000,
            ("direct", "상품"): 6000})
        self.assertEqual("pass", ordered["status"], ordered["issues"])

    def test_sold_grade_without_price_fails(self):
        result = self.review(prices={("purchase", "상품"): 3000})
        self.assertIn("SALES_PRICE_MISSING", self.codes(result, "error"))

    def test_two_channels_with_the_same_price_fail(self):
        result = self.review(prices={("purchase", "상품"): 3000, ("direct", "상품"): 3000})
        self.assertIn("CHANNEL_PRICE_EQUAL", self.codes(result, "error"))

    def test_single_channel_plan_is_not_compared_across_channels(self):
        result = self.review(shares={("purchase", "상품"): 1.0},
                             prices={("purchase", "상품"): 3000})
        self.assertEqual("pass", result["status"], result["issues"])

    def test_channel_cost_needs_selling_expense_or_a_note(self):
        bare = self.review(notes=bare_notes(), sga=(0, 0, 0, 0, 0))
        self.assertEqual(["CHANNEL_COST_UNREFLECTED"] * 2, self.codes(bare, "error"))
        expensed = self.review(notes=bare_notes(), sga=(0, 50, 50, 50, 50))
        self.assertNotIn("CHANNEL_COST_UNREFLECTED", self.codes(expensed))
        noted = self.review(sga=(0, 0, 0, 0, 0))
        self.assertNotIn("CHANNEL_COST_UNREFLECTED", self.codes(noted))

    def test_missing_basis_notes_warn_without_blocking(self):
        result = self.review(notes=())
        self.assertEqual("pass", result["status"])
        self.assertEqual({"GRADE_BASIS_NOTE_MISSING", "CHANNEL_SOURCE_NOTE_MISSING"},
                         set(self.codes(result, "warning")))

    def test_uncached_price_formula_blocks_instead_of_guessing(self):
        result = self.review(formula_cells=((" 5. 판매계획", "O8", "=O7"),))
        self.assertEqual("blocked", result["status"])
        self.assertIn("PLAN_FIT_UNCACHED", self.codes(result))
        self.assertNotIn("SALES_UNUSED_PRICE", self.codes(result))

    def test_without_method_equipment_is_not_judged(self):
        result = self.review(method=None,
                             investments=(("비닐하우스", 50000), ("관리기", 4000)))
        self.assertEqual(["CULTIVATION_METHOD_UNDECLARED"], self.codes(result))

    def test_method_aliases_resolve_and_unknown_is_refused(self):
        self.assertEqual("protected_soilless", self.pf.resolve_method("고설 수경"))
        self.assertEqual("open_field", self.pf.resolve_method("open_field"))
        self.assertEqual("mushroom", self.pf.resolve_method("균상재배"))
        with self.assertRaises(ValueError):
            self.pf.resolve_method("공중재배")

    def test_cli_exit_codes(self):
        good = plan_book(self.dir / "cli-good.xlsx", notes=self.GOOD_NOTES,
                         sga=(1, 1, 1, 1, 1))
        bad = plan_book(self.dir / "cli-bad.xlsx", notes=self.GOOD_NOTES,
                        sga=(1, 1, 1, 1, 1),
                        prices={("purchase", "상품"): 3000, ("direct", "상품"): 3000})
        script = str(SCRIPTS / "gg_plan_fit.py")

        def run(*args):
            proc = subprocess.run([sys.executable, script, *args],
                                  capture_output=True, text=True, encoding="utf-8")
            return proc.returncode, json.loads(proc.stdout)

        self.assertEqual(0, run(str(good), "--method", "고설수경")[0])
        code, out = run(str(bad), "--method", "고설수경")
        self.assertEqual((1, "fail"), (code, out["status"]))
        # 사용 오류(모르는 재배방식·이유 없는 --allow)는 확인 불가(2)와 다른 3이다.
        self.assertEqual(3, run(str(good), "--method", "공중재배")[0])
        self.assertEqual(3, run(str(good), "--allow", "이유없음")[0])


class UnreadableIsNotPassTests(_PlanFitCase):
    """Values or tables the check cannot read must block, never pass.

    Six cases from the independent review: a readable copied history fails,
    an uncached history formula and unreadable investment/labour tables
    block instead of disappearing as "no items"."""

    def test_baseline_passes(self):
        self.assertEqual("pass", self.review(method="고설수경")["status"])

    def test_unsold_grade_history_value_fails(self):
        result = self.review(method="고설수경", history={
            ("purchase", "상품"): (3000,), ("purchase", "중품"): (3000,)})
        self.assertEqual("fail", result["status"])
        self.assertIn("SALES_UNUSED_PRICE", self.codes(result))

    def test_uncached_history_formula_blocks(self):
        result = self.review(method="고설수경",
                             history={("purchase", "상품"): (3000,)},
                             formula_cells=((" 5. 판매계획", "E8", "=E7"),))
        self.assertEqual("blocked", result["status"])
        cells = [i["evidence"]["cell"] for i in result["issues"]
                 if i["code"] == "PLAN_FIT_UNCACHED"]
        self.assertEqual(["E8"], cells)

    def test_missing_investment_sheet_blocks(self):
        result = self.review(method="고설수경",
                             edit=lambda wb: wb.remove(wb["3.투자계획"]))
        self.assertEqual("blocked", result["status"])
        self.assertEqual(["3.투자계획"], self.unreadable_sheets(result))

    def test_unreadable_investment_header_blocks(self):
        def rename(wb):
            wb["3.투자계획"]["C18"] = "투자 항목명"
        result = self.review(method="고설수경", investments=(("관리기", 0),),
                             edit=rename)
        self.assertEqual("blocked", result["status"])
        self.assertEqual(["3.투자계획"], self.unreadable_sheets(result))

    def test_unreadable_labour_header_blocks(self):
        def rename(wb):
            wb["8. 노무비계획"]["I24"] = "작업일수(일)"
        result = self.review(method="고설수경", tasks=(("경운", 6),), edit=rename)
        self.assertEqual("blocked", result["status"])
        self.assertEqual(["8. 노무비계획"], self.unreadable_sheets(result))

    def test_empty_but_readable_tables_still_pass(self):
        result = self.review(method="고설수경", investments=(), depreciation=(),
                             tasks=())
        self.assertEqual("pass", result["status"], result["issues"])

    @staticmethod
    def unreadable_sheets(result):
        return [i["evidence"]["sheet"] for i in result["issues"]
                if i["code"] == "PLAN_FIT_TABLE_UNREADABLE"]


def _cells(sheet, **cells):
    def edit(wb):
        for ref, value in cells.items():
            wb[sheet][ref] = value
    return edit


class PartlyUnreadableIsNotPassTests(_PlanFitCase):
    """A readable row or block must not hide another one the check cannot read.

    The first four cases come from the second independent review; the rest
    are the same failure in the other readers."""

    def check(self, edit=None, **book):
        book.setdefault("sga", (1, 1, 1, 1, 1))
        return self.review(method="고설수경", edit=edit, **book)

    def assertBlocked(self, result, sheet=None):
        self.assertEqual("blocked", result["status"], result["issues"])
        if sheet is not None:
            self.assertIn(sheet, [i["evidence"].get("sheet") for i in result["issues"]
                                  if i["status"] == "blocked"])

    # --- review cases -----------------------------------------------------

    def test_asset_control_warns(self):
        result = self.check(opening=(("전", 0, ""),))
        self.assertEqual("pass", result["status"])
        self.assertEqual(["OPENING_ZERO_ROW"], self.codes(result))

    def test_asset_amount_is_never_taken_as_header(self):
        result = self.check(opening=(("전", 0, ""),),
                            edit=_cells("1. 기초재무상태조사", E8=None))
        self.assertBlocked(result, "1. 기초재무상태조사")
        self.assertNotIn("OPENING_ZERO_ROW", self.codes(result))

    @staticmethod
    def two_depreciation_blocks(broken):
        """First year table: 관리기 0; second: 비닐하우스 50000. ``broken`` makes
        only the first table's cost header unreadable."""
        def edit(wb):
            ws = wb["10. 감가상각비계획 "]
            ws["B8"] = "취득가격(천원)" if broken else "취득가격"
            ws["C7"], ws["C8"] = "관리기", 0
            ws["C19"], ws["C20"] = "비닐하우스", 50000
        return edit

    def test_depreciation_control_reads_both_blocks(self):
        result = self.check(edit=self.two_depreciation_blocks(False))
        self.assertEqual("fail", result["status"])
        self.assertEqual({"UNOWNED_ROW", "EQUIPMENT_METHOD_MISMATCH"}, set(self.codes(result)))

    def test_one_unreadable_depreciation_block_blocks(self):
        result = self.check(edit=self.two_depreciation_blocks(True))
        self.assertBlocked(result, "10. 감가상각비계획")

    def test_depreciation_section_without_cost_row_blocks(self):
        result = self.check(edit=_cells("10. 감가상각비계획 ", C19="관리기",
                                        B20="취득가(원)", C20=0))
        self.assertBlocked(result, "10. 감가상각비계획")

    # --- same failure in the other readers ---------------------------------

    def test_text_amount_blocks(self):
        self.assertBlocked(self.check(investments=(("관리기", "미정"),)), "3.투자계획")
        self.assertBlocked(self.check(opening=(("전", "미정", ""),)), "1. 기초재무상태조사")
        self.assertBlocked(self.check(tasks=(("경운", "미정"),)), "8. 노무비계획")

    def test_amount_without_item_name_blocks(self):
        # 첫해 블록 안(투자항목 2개 아래, '계' 위)에 이름 없는 금액 행을 끼운다.
        def nameless(wb):
            ws = wb["3.투자계획"]
            ws.insert_rows(21)
            ws["E21"] = 4000
        self.assertBlocked(self.check(edit=nameless), "3.투자계획")
        self.assertBlocked(self.check(edit=_cells("8. 노무비계획", B27=None)), "8. 노무비계획")

    def test_year_subtotal_row_is_not_a_nameless_item(self):
        # 연도 칸(B)만 있고 투자항목 없는 연도 행(금액 0)은 해당 없음이다.
        self.assertEqual("pass", self.check()["status"])
        self.assertEqual("pass", self.check(investments=())["status"])

    def test_cash_named_in_second_column_is_read(self):
        result = self.check(opening=((None, 0, ""),),
                            edit=_cells("1. 기초재무상태조사", C9="예금"))
        self.assertEqual(["OPENING_ZERO_ROW"], self.codes(result))

    def test_second_labour_table_without_days_header_blocks(self):
        result = self.check(edit=_cells("8. 노무비계획", B40="항목", I40="일수",
                                        B41="경운", I41=5, B42="합계"))
        self.assertBlocked(result, "8. 노무비계획")

    def test_labour_table_without_total_row_blocks(self):
        self.assertBlocked(self.check(edit=_cells("8. 노무비계획", B28=None)), "8. 노무비계획")

    def test_sales_value_without_year_or_share_header_blocks(self):
        self.assertBlocked(self.check(edit=_cells(" 5. 판매계획", E10=None, E12=3000)), "5. 판매계획")
        self.assertBlocked(self.check(edit=_cells(" 5. 판매계획", H20=None, H23=0.2)), "5. 판매계획")

    def test_sales_text_or_wrong_grade_blocks(self):
        self.assertBlocked(self.check(edit=_cells(" 5. 판매계획", D23="약간")), "5. 판매계획")
        self.assertBlocked(self.check(edit=_cells(" 5. 판매계획", C8="특품")), "5. 판매계획")
        self.assertBlocked(self.check(edit=_cells(" 5. 판매계획", C27="특품")), "5. 판매계획")

    def test_selling_expense_text_blocks(self):
        result = self.check(sga=(0, 0, 0, 0, 0), edit=_cells("12 손익계획", C13="별도"))
        self.assertBlocked(result, "12 손익계획")


class SharedTableCheckTests(_PlanFitCase):
    """Every table goes through one check: expected counts from the school
    layout, and the table's real column range (not the last recognised header).

    Each case removes or blanks one header at the table's last year/column and
    leaves a value under it; a control keeps the header."""

    def check(self, edit=None, **book):
        book.setdefault("sga", (1, 1, 1, 1, 1))
        return self.review(method="고설수경", edit=edit, **book)

    def assertBlockedOn(self, result, sheet):
        self.assertEqual("blocked", result["status"], result["issues"])
        self.assertIn(sheet, [i["evidence"].get("sheet") for i in result["issues"]
                              if i["status"] == "blocked"])

    # --- review case: last share year header missing -----------------------

    def test_last_share_year_control_reads_the_sale(self):
        result = self.check(edit=_cells(" 5. 판매계획", L23=0.2))
        self.assertEqual("fail", result["status"])
        self.assertIn("SALES_PRICE_MISSING", self.codes(result))

    def test_last_share_year_header_missing_blocks(self):
        result = self.check(edit=_cells(" 5. 판매계획", L20=None, L23=0.2))
        self.assertBlockedOn(result, "5. 판매계획")

    def test_last_volume_year_header_missing_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells(" 5. 판매계획", M20=None)), "5. 판매계획")

    # --- same boundary in the other tables ----------------------------------

    def test_last_history_year_header_missing_blocks(self):
        result = self.check(edit=_cells(" 5. 판매계획", M6=None, M8=3000))
        self.assertBlockedOn(result, "5. 판매계획")

    def test_value_after_average_price_column_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells(" 5. 판매계획", P8=3000)), "5. 판매계획")

    def test_last_selling_expense_year_header_missing_blocks(self):
        result = self.check(sga=(0, 0, 0, 0, 0), edit=_cells("12 손익계획", G4=None, G13=500))
        self.assertBlockedOn(result, "12 손익계획")

    def test_last_investment_year_missing_blocks(self):
        def drop_last_year(wb):
            ws = wb["3.투자계획"]
            last = max(r for r in range(19, ws.max_row + 1) if ws.cell(r, 2).value)
            ws.cell(last, 2).value = None
        self.assertBlockedOn(self.check(edit=drop_last_year), "3.투자계획")

    def test_value_right_of_investment_note_column_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells("3.투자계획", J19=4000)), "3.투자계획")

    def test_last_depreciation_year_table_missing_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells("10. 감가상각비계획 ", B54=None)),
                             "10. 감가상각비계획")

    def test_depreciation_cost_under_column_without_name_blocks(self):
        result = self.check(edit=_cells("10. 감가상각비계획 ", D8=7000))
        self.assertBlockedOn(result, "10. 감가상각비계획")

    def test_depreciation_value_right_of_last_header_blocks(self):
        # 머리글 '영농시설'을 지우면 첫해 표 머리글은 C7 장비 이름까지다. D8 값은 표 밖 열이다.
        result = self.check(edit=_cells("10. 감가상각비계획 ", C6=None, D8=7000))
        self.assertBlockedOn(result, "10. 감가상각비계획")

    def test_last_asset_block_missing_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells("1. 기초재무상태조사", B49=None)),
                             "1. 기초재무상태조사")

    def test_asset_value_right_of_block_header_blocks(self):
        result = self.check(opening=(("전", 100, ""),),
                            edit=_cells("1. 기초재무상태조사", I9=500))
        self.assertBlockedOn(result, "1. 기초재무상태조사")

    def test_second_labour_table_missing_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells("8. 노무비계획", B36=None, B38=None)),
                             "8. 노무비계획")

    def test_labour_value_under_unnamed_column_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells("8. 노무비계획", J26=3)), "8. 노무비계획")

    def test_table_to_the_right_of_a_gap_is_not_read(self):
        # The wage observation block (registered from the template at the
        # '인건비 평균 상승률' title) is a known area outside the labour table.
        result = self.check(edit=_cells("8. 노무비계획", W26="자료집",
                                        X24="인건비 평균 상승률", Y26=117156))
        self.assertEqual("pass", result["status"], result["issues"])

    def test_value_right_after_last_header_column_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells("8. 노무비계획", K24="자가 노동비", L26=3)),
                             "8. 노무비계획")

    # --- residual-value cases from the fourth review ------------------------

    def test_unnamed_depreciation_column_with_zero_money_rows_is_empty(self):
        result = self.check(edit=_cells("10. 감가상각비계획 ", D6="대농기구", D8=0,
                                        B9="감가상각비", D9=0))
        self.assertEqual("pass", result["status"], result["issues"])

    def test_unnamed_depreciation_column_with_residual_charge_blocks(self):
        result = self.check(edit=_cells("10. 감가상각비계획 ", D6="대농기구", D8=0,
                                        B9="감가상각비", D9=100))
        self.assertBlockedOn(result, "10. 감가상각비계획")

    def test_labour_value_one_empty_column_after_last_header_blocks(self):
        result = self.check(edit=_cells("8. 노무비계획", K24="자가 노동비", M26=5))
        self.assertBlockedOn(result, "8. 노무비계획")

    def test_value_anywhere_outside_recognised_areas_blocks(self):
        self.assertBlockedOn(self.check(edit=_cells("3.투자계획", Z40=1)), "3.투자계획")
        self.assertBlockedOn(self.check(edit=_cells("12 손익계획", H13=500)), "12 손익계획")
        self.assertBlockedOn(self.check(edit=_cells("1. 기초재무상태조사", J22="x")),
                             "1. 기초재무상태조사")

    def test_registered_template_leftover_cell_is_not_read(self):
        result = self.check(edit=_cells("1. 기초재무상태조사", J21=" w"))
        self.assertEqual("pass", result["status"], result["issues"])


def _four_year_plan(wb):
    """Drop the last plan year from every table the checker reads."""
    sales = wb[" 5. 판매계획"]
    for row in range(19, 30):
        sales.cell(row, 12).value = sales.cell(row, 13).value = None
    inv = wb["3.투자계획"]
    last = max(r for r in range(19, inv.max_row + 1) if inv.cell(r, 2).value)
    inv.cell(last, 2).value = inv.cell(last, 5).value = None
    dep = wb["10. 감가상각비계획 "]
    for ref in ("B52", "B54", "C54", "B56"):
        dep[ref] = None
    pl = wb["12 손익계획"]
    pl["G4"] = pl["G13"] = None


def _one_labour_table(wb):
    labor = wb["8. 노무비계획"]
    for ref in ("B36", "B38", "I38", "B41"):
        labor[ref] = None


class ScopeAndStateTests(_PlanFitCase):
    """Workbooks that are not the school template are not checked; template
    variants the checker does not support are 확인 불가 (blocked), never a plan
    error; only contradictions in values actually read fail."""

    def cli(self, path):
        proc = subprocess.run([sys.executable, str(SCRIPTS / "gg_plan_fit.py"), str(path),
                               "--method", "고설수경"],
                              capture_output=True, text=True, encoding="utf-8")
        return proc.returncode, json.loads(proc.stdout), proc.stderr.strip()

    def other_book(self, name, build):
        wb = Workbook()
        build(wb)
        path = self.dir / name
        wb.save(path)
        return path

    def assertNotApplicable(self, path):
        code, out, notice = self.cli(path)
        self.assertEqual((0, "not_applicable"), (code, out["status"]))
        self.assertEqual(["PLAN_FIT_NOT_APPLICABLE"], self.codes(out))
        self.assertEqual(1, len(notice.splitlines()), notice)

    def assertUnverified(self, result, sheets):
        self.assertEqual("blocked", result["status"], result["issues"])
        self.assertEqual([], self.codes(result, "error"))
        blocked = {i["evidence"].get("sheet") for i in result["issues"] if i["status"] == "blocked"}
        self.assertEqual(set(sheets), blocked)
        self.assertTrue(all(i["severity"] == "unverified" for i in result["issues"]
                            if i["status"] == "blocked"))

    def test_other_workbook_is_not_applicable(self):
        def build(wb):
            wb.active.title = "사업계획"
            wb.active["B2"] = "사업계획"
        self.assertNotApplicable(self.other_book("other.xlsx", build))

    def test_workbook_with_template_sheet_names_only_is_not_applicable(self):
        def build(wb):
            wb.active.title = "목록"
            for name, *_ in X01_OTHER_SHEETS[1:]:
                wb.create_sheet(name)
            for name in ("1. 기초재무상태조사", "3.투자계획", " 5. 판매계획",
                         "8. 노무비계획", "10. 감가상각비계획 ", "12 손익계획"):
                wb.create_sheet(name)
        self.assertNotApplicable(self.other_book("names-only.xlsx", build))

    def test_four_year_plan_is_unverified_not_failed(self):
        result = self.review(method="고설수경", edit=_four_year_plan)
        # 판매관리비(12 손익계획)는 판매표를 읽은 뒤에만 본다. 판매표가 확인 불가라 읽지 않는다.
        self.assertUnverified(result, {"3.투자계획", "5. 판매계획", "10. 감가상각비계획"})
        path = self.dir / ("plan-%d.xlsx" % self.count)
        code, out, notice = self.cli(path)
        self.assertEqual((2, "blocked"), (code, out["status"]))
        self.assertEqual(1, len(notice.splitlines()), notice)

    def test_one_labour_table_is_unverified_not_failed(self):
        self.assertUnverified(self.review(method="고설수경", edit=_one_labour_table),
                              {"8. 노무비계획"})

    def test_contradiction_in_read_values_fails_beside_unverified_table(self):
        result = self.review(method="고설수경", edit=_one_labour_table,
                             investments=(("관리기", 0),))
        self.assertEqual("fail", result["status"])
        self.assertIn("UNOWNED_ROW", self.codes(result, "error"))
        self.assertIn("PLAN_FIT_TABLE_UNREADABLE", self.codes(result, "unverified"))


def bare_notes():
    return ("* 등급: 상품 단일 등급 가정, 수매 도매 경락가와 직거래 소매가 평균.",)


class LegacyChannelPriceTests(ContractCase):
    """The legacy generator path refuses one channel's mean copied to both."""

    def setUp(self):
        self.sv = runtime("gg_school_verify")
        tmp = tempfile.TemporaryDirectory(prefix="knuaf-plan-fit-legacy-")
        self.addCleanup(tmp.cleanup)
        self.book = Path(tmp.name) / "any.xlsx"
        Workbook().save(self.book)

    def reasons(self, spec):
        return [i["reason"] for i in self.sv.verify_school_recalculated(spec, self.book)]

    def same_price_spec(self, **over):
        values = [6, 8, 10, 8, 8]
        return base_spec(purchase_price="8", direct_price="8", price_sources=[
            _source("student.purchase", "수매", values),
            _source("student.direct", "직거래", values)], **over)

    def test_equal_channel_prices_fail(self):
        self.assertTrue(any("수매·직거래 단가 동일" in r
                            for r in self.reasons(self.same_price_spec())))

    def test_distinct_channel_prices_pass_this_rule(self):
        self.assertFalse(any("단가 동일" in r for r in self.reasons(base_spec())))

    def test_single_active_channel_is_not_compared(self):
        spec = self.same_price_spec(purchase_share="1", direct_share="0")
        self.assertFalse(any("단가 동일" in r for r in self.reasons(spec)))
