"""Numeric diagnostics remain visible without blocking readable prose."""
import copy
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import Workbook, load_workbook
from tests._harness import SCRIPTS, runtime


class BodyPlanCheckTests(unittest.TestCase):
    """Keep review case names while migrating their verdicts to warnings."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.xlsx = self.root / 'plan.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = '판매계획'
        for row, values in enumerate(((0, 2400.4), (0.15, 2500), (0.3, 2600), (0.45, 2700), (0.6, 2800)), 2):
            sheet.cell(row, 2, values[0])
            sheet.cell(row, 3, values[1])
        book.save(self.xlsx)
        book.close()
        self.mapping = {'channels': [{'name': '직거래', 'aliases': ['직거래', '온라인'], 'years': {
            str(2040 + i): {'share': {'sheet': '판매계획', 'cell': f'B{i+2}'},
                           'price': {'sheet': '판매계획', 'cell': f'C{i+2}'}}
            for i in range(5)
        }}]}
        self.mod = runtime('gg_body_plan_check')

    def _write_package(self, name, parts):
        path = self.root / name
        with ZipFile(path, 'w') as package:
            for part, data in parts.items():
                package.writestr(part, data)
        return path

    def test_late_start_in_narrative_is_rejected_even_with_matching_ratio_elsewhere(self):
        result = self.mod.crosscheck('온라인·직거래 2042년 소량 시작\n2041년 직거래 15%', self.xlsx, self.mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual(result['issues'][0]['expected'], '2041')
        self.assertEqual(result['issues'][0]['metric'], 'start')
        self.assertEqual(result['issues'][0]['line'], 1)

    def test_metrics_before_year_cannot_hide_behind_a_matching_claim(self):
        for metric, claim in (
            ('start', '직거래 시작은 2042년이다.'),
            ('price', '직거래 단가 9,999원/kg은 2041년 계획이다.'),
            ('share', '직거래 비율 99%는 2041년 계획이다.'),
        ):
            with self.subTest(metric=metric):
                result = self.mod.crosscheck(claim + '\n2041년 직거래 15%', self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                self.assertEqual([r['status'] for r in result['comparisons']], ['match'])
                self.assertEqual([(r['line'], r['metric'], r['status'], r['quote'])
                                  for r in result['issues']], [(1, metric, 'not_checked', claim)])

    def test_unparsed_prefix_and_matching_suffix_on_same_line_do_not_pass(self):
        claim = '직거래 단가 9,999원/kg은 2041년 계획이며 판매비율은 15%다.'
        result = self.mod.crosscheck(claim, self.xlsx, self.mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual([(r['metric'], r['status']) for r in result['comparisons']], [('share', 'match')])
        self.assertEqual([(r['line'], r['metric']) for r in result['issues']], [(1, 'price')])

    def test_metrics_between_years_cannot_be_confirmed_against_the_earlier_plan(self):
        for metric, claim, unparsed in (
            ('price', '직거래 2041년 15%로 시작하며, 단가 2,500원/kg은 2042년 기준이다.', '2,500원/kg'),
            ('share', '직거래 2041년 시작했고, 비율 15%는 2042년 계획이다.', '15%'),
        ):
            with self.subTest(metric=metric):
                result = self.mod.crosscheck(claim, self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                issue = next(r for r in result['issues'] if r['metric'] == metric)
                self.assertEqual((issue['status'], issue['candidate_years'], issue['unparsed'], issue['quote']),
                                 ('not_checked', [2041, 2042], unparsed, claim))
                self.assertFalse(any(r['metric'] == metric for r in result['comparisons']))

    def test_year_before_metric_controls_still_report_the_later_year_mismatch(self):
        for metric, claim, actual, expected, cell in (
            ('price', '직거래 2041년 15%로 시작하며, 2042년 단가 2,500원/kg이다.', '2500', '2600', 'C4'),
            ('share', '직거래 2041년 시작했고, 2042년 비율 15%로 계획한다.', '15', '30', 'B4'),
        ):
            with self.subTest(metric=metric):
                result = self.mod.crosscheck(claim, self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                mismatch = next(r for r in result['issues'] if r['status'] == 'warning')
                self.assertEqual((mismatch['year'], mismatch['metric'], mismatch['actual'],
                                  mismatch['expected'], mismatch['locator']['cell']),
                                 (2042, metric, actual, expected, cell))

    def test_normal_claims_and_directly_year_bound_multiple_periods_pass(self):
        for claim in (
            '2041년 직거래 15%로 시작\n2042년 직거래 30%, 2,600원/kg',
            '직거래 2040년 0%, 2041년 비율 15%, 2042년 30%, 2043년 45%, 2044년 60%',
            '온라인 2041년 시작, 2042년 단가 2,600원/kg, 2043년 가격 2,700원/kg',
        ):
            with self.subTest(claim=claim):
                result = self.mod.crosscheck(claim, self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                self.assertEqual(result['issues'], [])
                self.assertTrue(all(r['actual'] == r['expected'] for r in result['comparisons']))

    def test_ambiguous_start_share_and_price_before_third_fourth_and_fifth_years(self):
        shares = {2040: 0, 2041: 15, 2042: 30, 2043: 45}
        for following in (2042, 2043, 2044):
            previous = following - 1
            prefix = '직거래 ' + ', '.join(f'{year}년 {shares[year]}%' for year in range(2040, following))
            for metric, text in (
                ('start', '시작'),
                ('share', f'비율 {shares[previous]}%'),
                ('price', f'단가 {2500 + (previous - 2041) * 100:,}원/kg'),
            ):
                for separator in (', ', '로 계획하며 '):
                    with self.subTest(following=following, metric=metric, separator=separator):
                        claim = f'{prefix}{separator}{text}은 {following}년 계획이다.'
                        result = self.mod.crosscheck(claim, self.xlsx, self.mapping)
                        self.assertEqual(result['status'], 'warning')
                        issue = next(r for r in result['issues']
                                     if r['metric'] == metric and r['unparsed'] in text)
                        self.assertEqual((issue['status'], issue['candidate_years'], issue['quote']),
                                         ('not_checked', [previous, following], claim))
                        self.assertFalse(any(r['metric'] == metric and r['year'] == previous
                                             for r in result['comparisons']))

    def test_explicit_later_years_retain_start_share_and_price_mismatches(self):
        shares = {2040: 0, 2041: 15, 2042: 30, 2043: 45}
        for year in (2042, 2043, 2044):
            prefix = '직거래 ' + ', '.join(f'{prior}년 {shares[prior]}%' for prior in range(2040, year))
            for metric, text, expected in (
                ('start', '시작', '2041'),
                ('share', f'비율 {shares[year - 1]}%', str(shares[year - 1] + 15)),
                ('price', f'단가 {2500 + (year - 2042) * 100:,}원/kg', str(2600 + (year - 2042) * 100)),
            ):
                with self.subTest(year=year, metric=metric):
                    result = self.mod.crosscheck(f'{prefix}, {year}년 {text}', self.xlsx, self.mapping)
                    self.assertEqual(result['status'], 'warning')
                    self.assertEqual([(r['year'], r['metric'], r['status'], r['expected'])
                                      for r in result['issues']], [(year, metric, 'warning', expected)])

    def test_even_an_adjacent_year_does_not_override_an_unseparated_following_year(self):
        for claim, metric in (
            ('직거래 2041년 단가 2,500원/kg은 2042년 기준이다.', 'price'),
            ('직거래 2041년 비율 15%는 2042년 계획이다.', 'share'),
            ('직거래 2041년 시작은 2042년이다.', 'start'),
        ):
            with self.subTest(metric=metric):
                result = self.mod.crosscheck(claim, self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                self.assertEqual(result['comparisons'], [])
                self.assertEqual([(r['metric'], r['candidate_years']) for r in result['issues']],
                                 [(metric, [2041, 2042])])

    def test_reviewed_shared_year_claims_compare_all_years(self):
        for claim, metrics in (
            ('직거래 2041년 및 2042년 비율은 모두 30%다.', {'share'}),
            ('직거래 2041년과 2042년 비율은 모두 30%다.', {'share'}),
            ('직거래 2041년–2042년 비율은 매년 30%다.', {'share'}),
            ('직거래 2041년 및 2042년 단가는 모두 2,600원/kg이다.', {'price'}),
            ('직거래 2041년과 2042년 단가는 모두 2,600원/kg이다.', {'price'}),
            ('직거래 (2041년 및 2042년) 비율은 모두 30%다.', {'share'}),
            ('직거래 2041년 및 2042년 단가는 2,600원/kg, 판매비율은 30%다.', {'price', 'share'}),
        ):
            with self.subTest(claim=claim):
                result = self.mod.crosscheck(claim, self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                self.assertEqual({(r['year'], r['metric'], r['status']) for r in result['comparisons']},
                                 {(year, metric, 'warning' if year == 2041 else 'match')
                                  for year in (2041, 2042) for metric in metrics})

    def test_group_grammar_covers_lists_ranges_and_unmarked_endpoints(self):
        for expression, years in (
            ('2041년와 2042년', [2041, 2042]),
            ('2041년, 2042년', [2041, 2042]),
            ('2040, 2041 및 2042년', [2040, 2041, 2042]),
            ('2041–2043년', [2041, 2042, 2043]),
            ('2041~2043', [2041, 2042, 2043]),
            ('2041년부터 2043년까지', [2041, 2042, 2043]),
            ('2041년에서 2043년까지', [2041, 2042, 2043]),
            ('2041년, 2042–2044년', [2041, 2042, 2043, 2044]),
        ):
            with self.subTest(expression=expression):
                result = self.mod.crosscheck(f'직거래 {expression} 비율은 모두 30%다.', self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                self.assertEqual([r['year'] for r in result['comparisons']], years)
        mapping = copy.deepcopy(self.mapping)
        for year in ('2041', '2042'):
            mapping['channels'][0]['years'][year] = mapping['channels'][0]['years']['2042'].copy()
        result = self.mod.crosscheck('직거래 2041년 및 2042년 단가 2,600원/kg, 비율 30%', self.xlsx, mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual({(r['year'], r['metric']) for r in result['comparisons']},
                         {(year, metric) for year in (2041, 2042) for metric in ('price', 'share')})
        self.assertEqual(result['issues'], [])

    def test_range_interior_is_compared_even_when_both_endpoints_match(self):
        for metric, text, column in (('share', '비율 30%', 'B'), ('price', '단가 2,600원/kg', 'C')):
            with self.subTest(metric=metric):
                mapping = copy.deepcopy(self.mapping)
                for year, row in (('2041', 4), ('2042', 3), ('2043', 4)):
                    mapping['channels'][0]['years'][year][metric]['cell'] = f'{column}{row}'
                result = self.mod.crosscheck(f'직거래 2041~2043년 {text}', self.xlsx, mapping)
                self.assertEqual(result['status'], 'warning')
                self.assertEqual([(r['year'], r['status']) for r in result['comparisons']],
                                 [(2041, 'match'), (2042, 'warning'), (2043, 'match')])
                self.assertEqual(result['issues'][0]['locator']['cell'], f'{column}3')

    def test_uncertain_groups_preserve_every_member_instead_of_comparing_a_subset(self):
        for expression, candidates in (
            ('2041년 또는 2042년', [2041, 2042]),
            ('2041년 대비 2042년', [2041, 2042]),
            ('2041/2042년', [2041, 2042]),
            ('2042년~2041년', [2041, 2042]),
            ('(2041년 및 2042년 중 2043년)', [2041, 2042, 2043]),
        ):
            with self.subTest(expression=expression):
                claim = f'직거래 {expression} 단가 2,600원/kg, 비율 30%'
                result = self.mod.crosscheck(claim + '\n2044년 직거래 60%', self.xlsx, self.mapping)
                self.assertEqual(result['status'], 'warning')
                self.assertEqual([(r['line'], r['year'], r['status']) for r in result['comparisons']],
                                 [(2, 2044, 'match')])
                self.assertEqual({(r['metric'], tuple(r['candidate_years']), r['status'], r['quote'])
                                  for r in result['issues']},
                                 {(metric, tuple(candidates), 'not_checked', claim) for metric in ('price', 'share')})

    def test_later_group_does_not_hide_behind_an_earlier_matching_period(self):
        result = self.mod.crosscheck('직거래 2040년 0%, 2041년 및 2042년 비율은 모두 30%', self.xlsx, self.mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual([(r['year'], r['status']) for r in result['comparisons']],
                         [(2040, 'match'), (2041, 'warning'), (2042, 'match')])

    def test_unmapped_range_year_blocks_even_when_all_known_years_match(self):
        mapping = copy.deepcopy(self.mapping)
        mapping['channels'][0]['years']['2043']['share']['cell'] = 'B6'
        result = self.mod.crosscheck('직거래 2043년~2045년 비율은 모두 60%', self.xlsx, mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual([(r['year'], r['status']) for r in result['comparisons']],
                         [(2043, 'match'), (2044, 'match'), (2045, 'not_checked')])
        self.assertIsNone(result['issues'][0]['expected'])

    def test_cli_summary_counts_comparisons_and_unchecked_claims_for_every_status(self):
        for claim, mapping, status, code, comparisons, unchecked in (
            ('직거래 2041년 비율 15%', self.mapping, 'warning', 0, 1, 0),
            ('직거래 2041년 대비 2042년 비율 30%', self.mapping, 'warning', 0, 0, 1),
            ('직거래 2041년 대비 2042년 비율 30%\n2044년 직거래 60%', self.mapping, 'warning', 0, 1, 1),
            ('직거래 2041년 및 2042년 비율 30%', self.mapping, 'warning', 0, 2, 0),
            ('직거래 2041년 비율 15%', {}, 'blocked', 2, 0, 0),
        ):
            with self.subTest(status=status, comparisons=comparisons):
                body, locator = self.root / 'body.md', self.root / 'map.json'
                body.write_text(claim, encoding='utf-8')
                locator.write_text(json.dumps(mapping, ensure_ascii=False), encoding='utf-8')
                result = subprocess.run([sys.executable, str(SCRIPTS / 'gg_body_plan_check.py'), '--body', str(body),
                                         '--xlsx', str(self.xlsx), '--map', str(locator)], capture_output=True, text=True)
                self.assertEqual(result.returncode, code, result.stderr)
                payload = json.loads(result.stdout)
                self.assertEqual(payload['status'], status)
                self.assertEqual((payload['summary']['comparison_count'], payload['summary']['not_checked_count']),
                                 (comparisons, unchecked))

    def test_start_ratio_and_rounded_price_match(self):
        result = self.mod.crosscheck('2041년 직거래 15%로 시작\n2042년 온라인 30%\n2040년 직거래 단가 2,400원/kg', self.xlsx, self.mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual([(r['metric'], r['actual'], r['expected']) for r in result['comparisons']],
                         [('start', '2041', '2041'), ('share', '15', '15'), ('share', '30', '30'), ('price', '2400', '2400')])

    def test_ratio_and_price_mismatch_name_their_cells(self):
        result = self.mod.crosscheck('2041년 직거래 20%, 단가 2,600원/kg', self.xlsx, self.mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual([r['locator']['cell'] for r in result['issues']], ['B3', 'C3'])

    def test_zero_claims_does_not_pass(self):
        result = self.mod.crosscheck('판로는 학생에게 확인한다.', self.xlsx, self.mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual(result['comparisons'], [])

    def test_missing_cache_is_blocked_not_zero(self):
        book = Workbook()
        book.active.title = '판매계획'
        book.active['B2'] = '=0.1'
        book.save(self.xlsx)
        book.close()
        with self.assertRaisesRegex(ValueError, 'cached value'):
            self.mod.crosscheck('2040년 직거래 10%', self.xlsx, self.mapping)

    def test_gapped_plan_period_is_rejected(self):
        mapping = copy.deepcopy(self.mapping)
        del mapping['channels'][0]['years']['2041']
        with self.assertRaisesRegex(ValueError, 'contiguous'):
            self.mod.crosscheck('2042년 직거래 시작', self.xlsx, mapping)

    def test_cli_reports_mismatch_and_nonzero_exit(self):
        body = self.root / 'body.md'
        body.write_text('온라인 2042년 소량 시작', encoding='utf-8')
        locator = self.root / 'map.json'
        locator.write_text(json.dumps(self.mapping, ensure_ascii=False), encoding='utf-8')
        result = subprocess.run([sys.executable, str(SCRIPTS / 'gg_body_plan_check.py'), '--body', str(body),
                                 '--xlsx', str(self.xlsx), '--map', str(locator)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = json.loads(result.stdout)
        self.assertEqual(payload['status'], 'warning')
        self.assertEqual([(r['metric'], r['actual'], r['expected']) for r in payload['issues']],
                         [('start', '2042', '2041')])

    def test_cli_with_matching_and_unparsed_claims_returns_not_checked(self):
        locator = self.root / 'map.json'
        locator.write_text(json.dumps(self.mapping, ensure_ascii=False), encoding='utf-8')
        for text in ('직거래 시작은 2042년이다.', '직거래 단가 9,999원/kg은 2041년 계획이다.'):
            with self.subTest(text=text):
                body = self.root / 'body.md'
                body.write_text(text + '\n2041년 직거래 15%', encoding='utf-8')
                result = subprocess.run([sys.executable, str(SCRIPTS / 'gg_body_plan_check.py'),
                                         '--body', str(body), '--xlsx', str(self.xlsx),
                                         '--map', str(locator)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                payload = json.loads(result.stdout)
                self.assertEqual(payload['status'], 'warning')
                self.assertEqual(payload['issues'][0]['quote'], text)

    def test_reviewed_respective_and_growth_claims_are_nonblocking_warnings(self):
        locator = self.root / 'map.json'
        locator.write_text(json.dumps(self.mapping, ensure_ascii=False), encoding='utf-8')
        for claim in (
            '직거래 2041년 및 2042년의 비율은 각각 15%와 30%다.',
            '직거래 2041년과 2042년의 단가는 각각 2,500원/kg과 2,600원/kg이다.',
            '직거래 2041년과 2042년은 각각 15%, 30%다.',
            '직거래 2041년부터 2042년까지 비율을 15%에서 30%로 늘린다.',
        ):
            with self.subTest(claim=claim):
                body = self.root / 'body.md'
                body.write_text(claim, encoding='utf-8')
                result = subprocess.run([sys.executable, str(SCRIPTS / 'gg_body_plan_check.py'),
                                         '--body', str(body), '--xlsx', str(self.xlsx),
                                         '--map', str(locator)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertEqual(json.loads(result.stdout)['status'], 'warning')

    def test_unreadable_inputs_remain_blocked_with_summary_and_nonzero_exit(self):
        body, locator = self.root / 'body.md', self.root / 'map.json'
        body.write_text('직거래 2041년 비율 15%', encoding='utf-8')
        locator.write_text(json.dumps(self.mapping, ensure_ascii=False), encoding='utf-8')
        corrupt = self.root / 'corrupt.xlsx'
        corrupt.write_bytes(b'not a zip')
        wrong_format = self.root / 'unsupported.xls'
        wrong_format.write_bytes(b'not an OOXML workbook')
        with ZipFile(self.xlsx) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        malformed = self._write_package('malformed.xlsx', parts | {'xl/workbook.xml': b'<broken'})
        package_cases = []
        for index, part in enumerate(('[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml',
                                      'xl/_rels/workbook.xml.rels', 'xl/worksheets/sheet1.xml')):
            path = self._write_package(f'missing-part-{index}.xlsx',
                                       {name: data for name, data in parts.items() if name != part})
            package_cases.append((f'missing {part}', body, path, locator))
        for index, part in enumerate(('[Content_Types].xml', '_rels/.rels',
                                      'xl/_rels/workbook.xml.rels', 'xl/worksheets/sheet1.xml')):
            path = self._write_package(f'broken-part-{index}.xlsx', parts | {part: b'<broken'})
            package_cases.append((f'broken {part}', body, path, locator))
        invalid_json = self.root / 'invalid.json'
        invalid_json.write_text('{', encoding='utf-8')
        invalid_refs = self.root / 'invalid-refs.json'
        mapping = copy.deepcopy(self.mapping)
        mapping['channels'][0]['years']['2041']['share']['sheet'] = '없는 시트'
        invalid_refs.write_text(json.dumps(mapping, ensure_ascii=False), encoding='utf-8')
        for label, body_path, xlsx, map_path in (
            ('corrupt archive', body, corrupt, locator),
            ('missing workbook', body, self.root / 'missing.xlsx', locator),
            ('malformed workbook XML', body, malformed, locator),
            ('unsupported workbook format', body, wrong_format, locator),
            ('invalid map JSON', body, self.xlsx, invalid_json),
            ('invalid map reference', body, self.xlsx, invalid_refs),
            ('missing map', body, self.xlsx, self.root / 'missing.json'),
            ('missing body', self.root / 'missing.md', self.xlsx, locator),
            *package_cases,
        ):
            with self.subTest(label=label):
                result = subprocess.run([sys.executable, str(SCRIPTS / 'gg_body_plan_check.py'),
                                         '--body', str(body_path), '--xlsx', str(xlsx),
                                         '--map', str(map_path)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 2, result.stderr)
                payload = json.loads(result.stdout)
                self.assertEqual((payload['status'], payload['message']), ('blocked', '확인 불가'))
                self.assertEqual((payload['summary']['comparison_count'], payload['summary']['not_checked_count']),
                                 (0, 0))

    def test_shared_string_references_require_a_readable_table_and_in_range_index(self):
        main = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        pkg = '{http://schemas.openxmlformats.org/package/2006/relationships}'
        ct = '{http://schemas.openxmlformats.org/package/2006/content-types}'
        with ZipFile(self.xlsx) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        strings = (f'<sst xmlns="{main[1:-1]}"><si><t>unused</t></si>'
                   '<si><t>0.15</t></si></sst>').encode()
        parts['xl/sharedStrings.xml'] = strings
        types = ET.fromstring(parts['[Content_Types].xml'])
        ET.SubElement(types, ct + 'Override', PartName='/xl/sharedStrings.xml',
                      ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
        parts['[Content_Types].xml'] = ET.tostring(types)
        relationships = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        ET.SubElement(relationships, pkg + 'Relationship', Id='rIdShared',
                      Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings',
                      Target='sharedStrings.xml')
        parts['xl/_rels/workbook.xml.rels'] = ET.tostring(relationships)
        sheet = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        cell = sheet.find(f'.//{main}c[@r="B3"]')
        cell.set('t', 's')
        cell.find(main + 'v').text = '1'
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
        valid = self._write_package('shared-valid.xlsx', parts)
        result = self.mod.crosscheck('직거래 2041년 비율15%', valid, self.mapping)
        self.assertEqual(result['status'], 'warning')
        self.assertEqual([(r['year'], r['metric'], r['actual'], r['expected'], r['status'])
                          for r in result['comparisons']], [(2041, 'share', '15', '15', 'match')])
        self.assertEqual(result['issues'], [])
        body, locator = self.root / 'body.md', self.root / 'map.json'
        body.write_text('직거래 2041년 비율15%', encoding='utf-8')
        locator.write_text(json.dumps(self.mapping, ensure_ascii=False), encoding='utf-8')
        for label, index, table in (
            ('missing', '1', None),
            ('broken', '1', b'<broken'),
            ('empty', '0', f'<sst xmlns="{main[1:-1]}"/>'.encode()),
            ('negative', '-1', strings),
            ('past-end', '2', strings),
            ('noninteger', '0.5', strings),
            ('missing-index', None, strings),
        ):
            with self.subTest(label=label):
                changed = parts.copy()
                if table is None:
                    del changed['xl/sharedStrings.xml']
                else:
                    changed['xl/sharedStrings.xml'] = table
                sheet = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
                cell = sheet.find(f'.//{main}c[@r="B3"]')
                value = cell.find(main + 'v')
                if index is None:
                    cell.remove(value)
                else:
                    value.text = index
                changed['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
                path = self._write_package(f'shared-{label}.xlsx', changed)
                result = subprocess.run([sys.executable, str(SCRIPTS / 'gg_body_plan_check.py'),
                                         '--body', str(body), '--xlsx', str(path),
                                         '--map', str(locator)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 2, result.stderr)
                payload = json.loads(result.stdout)
                self.assertEqual((payload['status'], payload['message']), ('blocked', '확인 불가'))
                self.assertEqual((payload['summary']['comparison_count'], payload['summary']['not_checked_count']),
                                 (0, 0))

    def test_every_referenced_sheet_is_checked_using_its_relationship_target(self):
        main = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        doc = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
        ct = '{http://schemas.openxmlformats.org/package/2006/content-types}'
        book = load_workbook(self.xlsx)
        book.create_sheet('별도')['A1'] = '참고'
        base = self.root / 'two-sheets.xlsx'
        book.save(base)
        book.close()
        with ZipFile(base) as source:
            parts = {name: source.read(name) for name in source.namelist()}
        workbook = ET.fromstring(parts['xl/workbook.xml'])
        sheet = workbook.find(f'{main}sheets/{main}sheet[@name="별도"]')
        relationships = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        relationship = next(rel for rel in relationships if rel.get('Id') == sheet.get(doc + 'id'))
        old = relationship.get('Target').lstrip('/')
        target = 'xl/worksheets/renamed.xml'
        relationship.set('Target', 'worksheets/renamed.xml')
        parts['xl/_rels/workbook.xml.rels'] = ET.tostring(relationships)
        parts[target] = parts.pop(old)
        types = ET.fromstring(parts['[Content_Types].xml'])
        types.find(f'{ct}Override[@PartName="/{old}"]').set('PartName', '/' + target)
        parts['[Content_Types].xml'] = ET.tostring(types)
        valid = self._write_package('renamed-valid.xlsx', parts)
        result = self.mod.crosscheck('직거래 2041년 비율15%', valid, self.mapping)
        self.assertEqual([(r['year'], r['actual'], r['expected'], r['status']) for r in result['comparisons']],
                         [(2041, '15', '15', 'match')])
        self.assertEqual(result['issues'], [])
        body, locator = self.root / 'body.md', self.root / 'map.json'
        body.write_text('직거래 2041년 비율15%', encoding='utf-8')
        locator.write_text(json.dumps(self.mapping, ensure_ascii=False), encoding='utf-8')
        shared = ET.fromstring(parts[target])
        cell = shared.find(f'.//{main}c[@r="A1"]')
        cell.clear()
        cell.attrib.update(r='A1', t='s')
        ET.SubElement(cell, main + 'v').text = '0'
        for label, data in (('missing', None), ('broken', b'<broken'), ('shared', ET.tostring(shared))):
            with self.subTest(label=label):
                changed = parts.copy()
                if data is None:
                    del changed[target]
                else:
                    changed[target] = data
                path = self._write_package(f'unmapped-{label}.xlsx', changed)
                result = subprocess.run([sys.executable, str(SCRIPTS / 'gg_body_plan_check.py'),
                                         '--body', str(body), '--xlsx', str(path),
                                         '--map', str(locator)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 2, result.stderr)
                payload = json.loads(result.stdout)
                self.assertEqual((payload['status'], payload['message']), ('blocked', '확인 불가'))
                self.assertEqual((payload['summary']['comparison_count'], payload['summary']['not_checked_count']),
                                 (0, 0))
