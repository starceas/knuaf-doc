"""Policy B common output guard — unit contract (G-B, owner decision
2026-09-25).

``gg_major_contract.authorize_output`` is the single common check every
output entry point calls.  These cases pin the A32-L1..L6 expectations at
the guard itself; the per-entry-point wiring is covered by
``test_major_guard_paper`` and ``test_major_guard_outputs``.
"""
import copy
from pathlib import Path
import tempfile

from tests._harness import ContractCase, bind_major, fact_op, runtime, source_op, write_text

SPECIALTY = "specialty_crops"
INSECTS = "industrial_insects"


class OutputGuardContractTests(ContractCase):
    def setUp(self):
        self.mc = runtime("gg_major_contract")
        self.core = runtime("gg_core")

    def _held(self, reason, output, context, **kw):
        with self.assertRaises(self.mc.OutputHeldError) as caught:
            self.mc.authorize_output(output, context, **kw)
        self.assertEqual(caught.exception.reason, reason, caught.exception.detail)
        self.assertIsInstance(caught.exception, self.mc.MajorContractError)
        self.assertIn("guidance", caught.exception.detail)
        return caught.exception

    def _ctx(self, root, major_id=None):
        return self.mc.output_context(root, major_id)

    def test_l1_unmarked_input_is_held_without_specialty_inference(self):
        """A32-L1: no ID and no binding — flat and specialty free-label
        specs are both held; nothing is inferred from labels or titles."""
        root = self.make_project()
        before = self._tree_bytes(root)
        for spec in (
            {"writing_year": 2026},
            {"writing_year": 2026, "major": "특용작물전공",
             "title": "인삼 재배 창업계획", "crops": ["인삼"]},
        ):
            with self.subTest(spec=spec):
                self._held("major_id_required", self.mc.OUTPUT_SCHOOL_PAPER,
                           self._ctx(root), spec=spec)
        self.assertEqual(self._tree_bytes(root), before)

    def test_l2_binding_alone_does_not_authorize(self):
        """A32-L2: a valid canonical binding without an explicit request ID
        is held; the ID is never silently filled in from the canonical."""
        root = self.make_project()
        bind_major(root)
        for output in self.mc.OUTPUT_REQUIREMENTS:
            with self.subTest(output=output):
                self._held("major_id_required", output, self._ctx(root),
                           spec={"writing_year": 2026})

    def test_l3_explicit_id_with_matching_binding_is_authorized(self):
        """A32-L3: explicit specialty ID (spec, profile, or context) plus the
        matching binding authorizes every output specialty declares."""
        root = self.make_project()
        revision = bind_major(root)
        named = (
            ({"major_id": SPECIALTY}, None),
            ({"school_profile": {"major_id": SPECIALTY}}, None),
            ({"major_id": SPECIALTY,
              "school_profile": {"major_id": SPECIALTY}}, SPECIALTY),
            (None, SPECIALTY),
        )
        for output in self.mc.OUTPUT_REQUIREMENTS:
            for spec, ctx_major in named:
                with self.subTest(output=output, spec=spec, ctx=ctx_major):
                    auth = self.mc.authorize_output(
                        output, self._ctx(root, ctx_major), spec=spec)
                    self.assertEqual(auth.to_dict(), {
                        "output": output,
                        "major_id": SPECIALTY,
                        "module_version": "1.0.0",
                        "contract_version": self.mc.CONTRACT_VERSION,
                        "project_revision": revision,
                        "binding_evidence": "facts:selected_major",
                    })

    def test_l4_unregistered_conflicting_or_mismatched_major_is_refused(self):
        """A32-L4: an unregistered fruit ID, a spec/profile/context conflict,
        and an ID that differs from the binding are all refused; nothing is
        converted to specialty."""
        root = self.make_project()
        bind_major(root)
        paper = self.mc.OUTPUT_SCHOOL_PAPER
        self._held("unknown_major", paper, self._ctx(root),
                   spec={"major_id": "fruit_trees"})
        self._held("unknown_major", paper, self._ctx(root, "fruit_trees"))
        self._held("major_id_conflict", paper, self._ctx(root), spec={
            "major_id": SPECIALTY, "school_profile": {"major_id": INSECTS}})
        self._held("major_id_conflict", paper, self._ctx(root, INSECTS),
                   spec={"major_id": SPECIALTY})
        self._held("major_binding_mismatch", paper, self._ctx(root, INSECTS))

        insect_root = self.make_project()
        bind_major(insect_root, INSECTS)
        self._held("major_binding_mismatch", paper, self._ctx(insect_root),
                   spec={"major_id": SPECIALTY})

    def test_unsupported_output_follows_module_declaration(self):
        """Capability is the bound module's own declaration — the guard is
        not a specialty-only rule.  industrial_insects declares no paper,
        no workbook, and no finance calculation."""
        root = self.make_project()
        bind_major(root, INSECTS)
        for output in self.mc.OUTPUT_REQUIREMENTS:
            with self.subTest(output=output):
                self._held("unsupported_output", output,
                           self._ctx(root, INSECTS))

    def test_l5_null_empty_or_padded_ids_are_invalid_not_missing(self):
        """A32-L5: a present-but-null/empty ID is a validation failure."""
        root = self.make_project()
        bind_major(root)
        paper = self.mc.OUTPUT_SCHOOL_PAPER
        for spec in ({"major_id": None}, {"major_id": ""},
                     {"major_id": " specialty_crops"},
                     {"major_id": 7},
                     {"school_profile": {"major_id": None}},
                     {"major_id": SPECIALTY,
                      "school_profile": {"major_id": ""}}):
            with self.subTest(spec=spec):
                self._held("major_id_invalid", paper, self._ctx(root),
                           spec=spec)
        self._held("major_id_invalid", paper, self._ctx(root, ""))
        # An explicit null in the context is invalid, not "not named" —
        # even when the spec itself names a valid major.
        self._held("major_id_invalid", paper,
                   self.mc.OutputContext(root, None),
                   spec={"major_id": SPECIALTY})
        self._held("major_id_invalid", paper,
                   {"project_root": root, "major_id": None},
                   spec={"major_id": SPECIALTY})

    def test_l5_invalid_bindings_are_held(self):
        """A32-L5: duplicate/conflicting provided facts, unreviewed
        verification, missing/broken source refs, and module version
        mismatch hold the output; an unprovided fact is no binding."""
        root = self.make_project()
        bind_major(root)
        good = self.core.load(root)
        fact = good["facts"]["selected_major"]

        def variant(mutate):
            project = copy.deepcopy(good)
            mutate(project, project["facts"]["selected_major"])
            return project

        def duplicate(project, f):
            project["facts"]["selected_major_2"] = dict(
                copy.deepcopy(f), id="selected_major_2")

        def conflicting(project, f):
            project["facts"]["selected_major_2"] = dict(
                copy.deepcopy(f), id="selected_major_2", value=INSECTS)

        invalid = {
            "duplicate": duplicate,
            "conflicting": conflicting,
            "unreviewed": lambda p, f: f.update(verification="source_located"),
            "no_source_refs": lambda p, f: f.update(source_refs=[]),
            "broken_source_ref": lambda p, f: f.update(
                source_refs=[dict(f["source_refs"][0], id="gone")]),
            "version_mismatch": lambda p, f: f.update(module_version="0.9.0"),
            "no_version": lambda p, f: f.pop("module_version"),
            "unregistered_binding": lambda p, f: f.update(value="fruit_trees"),
        }
        self.assertEqual(fact["answer_state"], "provided")
        for name, mutate in invalid.items():
            with self.subTest(case=name):
                error = self._held(
                    "major_binding_invalid", self.mc.OUTPUT_SCHOOL_PAPER,
                    self._ctx(root, SPECIALTY), project=variant(mutate))
                self.assertIn("binding_reason", error.detail)
        with self.subTest(case="not_provided"):
            self._held(
                "major_binding_required", self.mc.OUTPUT_SCHOOL_PAPER,
                self._ctx(root, SPECIALTY), project=variant(
                    lambda p, f: f.update(answer_state="unknown")))

    def test_l6_explicit_id_without_binding_is_held(self):
        """A32-L6: an explicit specialty ID with no canonical binding is
        held for every output."""
        root = self.make_project()
        before = self._tree_bytes(root)
        for output in self.mc.OUTPUT_REQUIREMENTS:
            with self.subTest(output=output):
                self._held("major_binding_required", output,
                           self._ctx(root, SPECIALTY),
                           spec={"major_id": SPECIALTY})
        self.assertEqual(self._tree_bytes(root), before)

    def test_context_is_required_and_carries_no_permission(self):
        """A46-T5 (guard level): a missing context holds; caller-written
        flags such as ``allowed: True`` are never read — the decision is
        recomputed from the canonical record on disk."""
        root = self.make_project()
        paper = self.mc.OUTPUT_SCHOOL_PAPER
        spec = {"major_id": SPECIALTY, "allowed": True}
        self._held("output_context_required", paper, None, spec=spec)
        self._held("output_context_required", paper, {"major_id": SPECIALTY})
        self._held("output_context_required", paper, str(root), spec=spec)
        self._held("major_binding_required", paper,
                   {"project_root": root, "major_id": SPECIALTY,
                    "allowed": True, "authorized": True}, spec=spec)
        bind_major(root)
        auth = self.mc.authorize_output(
            paper, {"project_root": str(root), "major_id": SPECIALTY})
        self.assertEqual(auth.major_id, SPECIALTY)

    def test_unreadable_project_and_unknown_output(self):
        empty = Path(tempfile.mkdtemp(prefix="knuaf-noproject-"))
        self.addCleanup(lambda: __import__("shutil").rmtree(empty, True))
        self._held("project_unreadable", self.mc.OUTPUT_SCHOOL_PAPER,
                   self._ctx(empty, SPECIALTY))
        self.assertEqual(list(empty.iterdir()), [])
        root = self.make_project()
        bind_major(root)
        self._held("unknown_output", "school_hwpx", self._ctx(root, SPECIALTY))

    def test_reconfirm_detects_revision_or_binding_change(self):
        """A staged output is made visible only if the canonical revision
        and binding are still the ones that were authorized."""
        root = self.make_project()
        bind_major(root)
        ctx = self._ctx(root, SPECIALTY)
        auth = self.mc.authorize_output(self.mc.OUTPUT_SCHOOL_WORKBOOK, ctx)
        self.assertEqual(self.mc.reconfirm_output(auth, ctx), auth)
        write_text(root, "note.txt", "메모\n")
        note = fact_op("note", "farm.note", "메모", "", source_id="note")
        src = source_op("note.txt")
        src["value"]["id"] = "note"
        self.core.apply(root, {"request_id": "note", "ops": [src, note]},
                        auth.project_revision)
        with self.assertRaises(self.mc.OutputHeldError) as caught:
            self.mc.reconfirm_output(auth, ctx)
        self.assertEqual(caught.exception.reason, "project_revision_changed")

    def test_refusal_never_writes_canonical_state(self):
        """Holding an output leaves every persisted byte unchanged — it does
        not register a major, approve, or bump the revision."""
        root = self.make_project()
        bind_major(root, INSECTS)
        before = self._tree_bytes(root)
        for output in self.mc.OUTPUT_REQUIREMENTS:
            for ctx_major, spec in ((None, None), (SPECIALTY, None),
                                    (INSECTS, None),
                                    (None, {"major_id": None})):
                with self.subTest(output=output, ctx=ctx_major, spec=spec):
                    with self.assertRaises(self.mc.OutputHeldError):
                        self.mc.authorize_output(
                            output, self._ctx(root, ctx_major), spec=spec)
        self.assertEqual(self._tree_bytes(root), before)


    def test_l4_known_major_labels_conflicting_with_explicit_id(self):
        """A32-L4: a cover/profile label naming another known major (fruit,
        insect) conflicts with an explicit specialty ID; labels never infer
        a major on their own."""
        root = self.make_project()
        bind_major(root)
        paper = self.mc.OUTPUT_SCHOOL_PAPER
        for spec in (
            {"major_id": SPECIALTY,
             "school_profile": {"department": "과수학과"}},
            {"major_id": SPECIALTY, "major": "과 수 전공"},
            {"major_id": SPECIALTY, "school_profile": {"major": "Fruit Trees"}},
            {"major_id": SPECIALTY, "department": "산업곤충학과"},
            {"major_id": SPECIALTY, "major": "특용작물·과수"},
        ):
            with self.subTest(spec=spec):
                self._held("major_marker_conflict", paper, self._ctx(root),
                           spec=spec)
        auth = self.mc.authorize_output(paper, self._ctx(root), spec={
            "major_id": SPECIALTY,
            "school_profile": {"department": "특용작물학과"}})
        self.assertEqual(auth.major_id, SPECIALTY)

    def test_registered_major_without_outputs_is_refused(self):
        """A registered peer module (synthetic fruit_trees) that declares no
        paper/workbook/finance output is refused by capability — the guard
        reads the module declaration, not a specialty-only rule."""
        mc = self.mc
        fruit = mc.declare_module(
            major_id="fruit_trees", module_version="0.0.1",
            capabilities={"question": "supported", "document": "supported",
                          "evidence": "unsupported",
                          "finance": "unsupported"},
            question_schema=(), document_plan=(),
            evidence_applicability={}, finance_capabilities=(),
            validation_rules=(),
            supported_outputs=("question_list", "document_plan"),
        )
        registry = mc.ModuleRegistry(
            mc.MODULES + (fruit,), pack_owner=mc.load_pack_owner())
        root = self.make_project()
        write_text(root, "major-answer.txt", "전공 선택: fruit_trees\n")
        fact = fact_op("selected_major", "common.major_id", "fruit_trees", "",
                       scope="project", verification="claim_supported",
                       source_id="major-answer")
        fact["value"]["module_version"] = "0.0.1"
        src = source_op("major-answer.txt")
        src["value"]["id"] = "major-answer"
        self.core.apply(root, {"request_id": "fruit", "ops": [src, fact]}, 0)
        before = self._tree_bytes(root)
        for output in mc.OUTPUT_REQUIREMENTS:
            with self.subTest(output=output):
                self._held("unsupported_output", output,
                           self._ctx(root, "fruit_trees"), registry=registry)
        # Without the synthetic registration the same ID is unknown.
        self._held("unknown_major", mc.OUTPUT_SCHOOL_PAPER,
                   self._ctx(root, "fruit_trees"))
        self.assertEqual(self._tree_bytes(root), before)

    def _files(self):
        from docx import Document
        from openpyxl import Workbook
        d = Path(tempfile.mkdtemp(prefix="knuaf-kinds-"))
        self.addCleanup(lambda: __import__("shutil").rmtree(d, True))
        Workbook().save(d / "재무.xlsx")
        Document().save(d / "논문.docx")
        (d / "본문.md").write_text("# 본문\n", encoding="utf-8")
        (d / "out.pdf").write_bytes(b"%PDF-1.7\n%%EOF\n")
        (d / "renamed.docx").write_bytes((d / "재무.xlsx").read_bytes())
        (d / "finance.dat").write_bytes((d / "재무.xlsx").read_bytes())
        (d / "fake.xlsx").write_text("not a workbook", encoding="utf-8")
        return d

    def test_multi_file_outputs_need_each_kind(self):
        """Adopt/export/Office publishing several files checks every file's
        kind from suffix AND bytes; a PDF takes its known generator's kind
        and needs both only when its origin is unknown."""
        mc = self.mc
        d = self._files()
        paper, book = mc.OUTPUT_SCHOOL_PAPER, mc.OUTPUT_SCHOOL_WORKBOOK
        self.assertEqual(mc.output_kinds_for_file(d / "재무.xlsx"), (book,))
        self.assertEqual(mc.output_kinds_for_file(d / "논문.docx"), (paper,))
        self.assertEqual(mc.output_kinds_for_file(d / "본문.md"), (paper,))
        self.assertEqual(mc.output_kinds_for_file(d / "out.pdf"),
                         (paper, book))
        self.assertEqual(
            mc.output_kinds_for_file(d / "out.pdf", pdf_origin="paper"),
            (paper,))
        self.assertEqual(
            mc.output_kinds_for_file(d / "out.pdf", pdf_origin="workbook"),
            (book,))
        for name in ("renamed.docx", "finance.dat", "fake.xlsx"):
            with self.subTest(name=name):
                with self.assertRaises(mc.OutputHeldError) as caught:
                    mc.output_kinds_for_file(d / name)
                self.assertEqual(caught.exception.reason, "unknown_output")
        with self.assertRaises(mc.OutputHeldError):
            mc.output_kinds_for_file(d / "out.pdf", pdf_origin="slides")

        root = self.make_project()
        bind_major(root)
        auths = mc.authorize_outputs(
            mc.output_kinds_for_file(d / "out.pdf") + (paper,),
            self._ctx(root, SPECIALTY))
        self.assertEqual([a.output for a in auths], [paper, book])
        self.assertEqual({a.project_revision for a in auths}, {1})
        with self.assertRaises(mc.OutputHeldError):
            mc.authorize_outputs((), self._ctx(root, SPECIALTY))

    def test_bundle_authorizations_share_one_revision(self):
        """``authorize_outputs`` reads the canonical once: a revision bump
        between kinds cannot mix revisions inside one bundle."""
        mc = self.mc
        root = self.make_project()
        bind_major(root)
        real_load = self.core.load
        calls = []
        bumping = []

        def load_then_bump(path):
            if bumping:  # apply's own internal reads are not guard reads
                return real_load(path)
            calls.append(path)
            project = real_load(path)
            if len(calls) == 1:
                write_text(root, "bump.txt", "메모\n")
                src = source_op("bump.txt")
                src["value"]["id"] = "bump"
                bumping.append(True)
                try:
                    self.core.apply(root, {"request_id": "bump",
                                           "ops": [src]},
                                    project["revision"])
                finally:
                    bumping.clear()
                self.assertEqual(real_load(root)["revision"], 2)
            return project

        mc.gg_core.load = load_then_bump
        self.addCleanup(setattr, mc.gg_core, "load", real_load)
        auths = mc.authorize_outputs(
            (mc.OUTPUT_SCHOOL_PAPER, mc.OUTPUT_SCHOOL_WORKBOOK),
            self._ctx(root, SPECIALTY))
        self.assertEqual(len(calls), 1)
        self.assertEqual({a.project_revision for a in auths}, {1})

    def test_label_matching_has_word_boundaries(self):
        """A label split across a department suffix ("…학과 수료") is not a
        fruit label; spaced letters ("과 수 전공") still are."""
        mc = self.mc
        for text in ("특용작물학과 수료", "특용작물학과수료", "특용작물전공 수석",
                     "grapefruit 가공"):
            with self.subTest(text=text):
                self.assertEqual(mc.marker_conflicts(
                    {"school_profile": {"department": text}}, SPECIALTY), [])
        for text in ("과 수 전공", "과 수전공", "과수학과", "원예과수학과",
                     "특용작물 과수", "특용작물·과수", "식용곤충학과", "Fruit",
                     "Fruit.", "Insects/특용"):
            with self.subTest(text=text):
                self.assertNotEqual(mc.marker_conflicts(
                    {"school_profile": {"department": text}}, SPECIALTY), [])

    def test_reconfirm_binding_removed_or_replaced(self):
        """A binding removed or replaced after authorization refuses with
        the promised ``project_revision_changed`` (cause kept)."""
        mc = self.mc
        for change, cause in (("unknown", "major_binding_required"),
                              (INSECTS, "major_binding_mismatch")):
            with self.subTest(change=change):
                root = self.make_project()
                bind_major(root)
                ctx = self._ctx(root, SPECIALTY)
                auth = mc.authorize_output(mc.OUTPUT_SCHOOL_PAPER, ctx)
                fact = dict(self.core.load(root)["facts"]["selected_major"])
                if change == "unknown":
                    fact["answer_state"] = "unknown"
                    fact["value"] = None
                else:
                    fact["value"] = INSECTS
                    fact["module_version"] = "0.1.0"
                try:
                    self.core.apply(root, {"request_id": "swap:" + change,
                                           "ops": [{"collection": "facts",
                                                    "value": fact}]},
                                    auth.project_revision)
                except Exception as error:  # pragma: no cover - diagnostic
                    self.fail("apply refused %s: %r" % (change, error))
                with self.assertRaises(mc.OutputHeldError) as caught:
                    mc.reconfirm_output(auth, ctx)
                self.assertEqual(caught.exception.reason,
                                 "project_revision_changed")
                self.assertEqual(caught.exception.detail["cause"], cause)

    def test_container_formats_are_confirmed_not_assumed(self):
        """Compound (OLE) and ZIP containers must name their real document
        kind: an XLS renamed .hwp, a broken ZIP named .hwpx, or a bare text
        file named .xlsx is refused; real HWP/HWPX/XLS shapes pass."""
        mc = self.mc
        d = Path(tempfile.mkdtemp(prefix="knuaf-ole-"))
        self.addCleanup(lambda: __import__("shutil").rmtree(d, True))
        magic = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
        pad = b"\x00" * 512
        xls = magic + pad + "Workbook".encode("utf-16-le") + pad
        hwp = (magic + pad + "FileHeader".encode("utf-16-le") + pad
               + b"HWP Document File" + pad)
        (d / "x.xls").write_bytes(xls)
        (d / "renamed.hwp").write_bytes(xls)
        (d / "doc.hwp").write_bytes(hwp)
        (d / "renamed.xls").write_bytes(hwp)
        (d / "broken.hwpx").write_bytes(b"PKbroken")
        import zipfile
        with zipfile.ZipFile(d / "doc.hwpx", "w") as z:
            z.writestr("mimetype", "application/hwp+zip")
        with zipfile.ZipFile(d / "nothing.hwpx", "w") as z:
            z.writestr("a.txt", "x")
        # A real workbook archive that also claims an HWPX mimetype.
        from openpyxl import Workbook
        Workbook().save(d / "book.xlsx")
        with zipfile.ZipFile(d / "book.xlsx") as src, \
                zipfile.ZipFile(d / "mixed.hwpx", "w") as z:
            z.writestr("mimetype", "application/hwp+zip")
            for item in src.infolist():
                z.writestr(item, src.read(item.filename))
        book, paper = mc.OUTPUT_SCHOOL_WORKBOOK, mc.OUTPUT_SCHOOL_PAPER
        self.assertEqual(mc.output_kinds_for_file(d / "x.xls"), (book,))
        self.assertEqual(mc.output_kinds_for_file(d / "doc.hwp"), (paper,))
        self.assertEqual(mc.output_kinds_for_file(d / "doc.hwpx"), (paper,))
        for name in ("renamed.hwp", "renamed.xls", "broken.hwpx",
                     "nothing.hwpx", "mixed.hwpx", "missing.pdf"):
            with self.subTest(name=name):
                with self.assertRaises(mc.OutputHeldError) as caught:
                    mc.output_kinds_for_file(d / name)
                self.assertEqual(caught.exception.reason, "unknown_output")

    def test_json_companion_inherits_main_kinds(self):
        """A receipt/manifest companion carries its main output's kinds; a
        deliverable companion is judged on its own bytes.  A planned PDF
        (before the engine runs) takes its engine's kind."""
        mc = self.mc
        d = self._files()
        (d / "재무.xlsx.manifest.json").write_text("{}", encoding="utf-8")
        main = mc.output_kinds_for_file(d / "재무.xlsx")
        self.assertEqual(mc.companion_kinds(d / "재무.xlsx.manifest.json", main),
                         main)
        self.assertEqual(mc.companion_kinds(d / "논문.docx", main),
                         (mc.OUTPUT_SCHOOL_PAPER,))
        with self.assertRaises(mc.OutputHeldError):
            mc.companion_kinds(d / "finance.dat", main)
        # Workbook bytes (or a JSON non-object) renamed .json do not
        # inherit the main output's kinds.
        (d / "finance.json").write_bytes((d / "재무.xlsx").read_bytes())
        (d / "list.json").write_text("[1, 2]", encoding="utf-8")
        for name in ("finance.json", "list.json", "absent.json"):
            with self.subTest(name=name):
                with self.assertRaises(mc.OutputHeldError) as caught:
                    mc.companion_kinds(d / name, (mc.OUTPUT_SCHOOL_PAPER,))
                self.assertEqual(caught.exception.reason, "unknown_output")
        self.assertEqual(mc.PDF_ORIGINS["paper"], (mc.OUTPUT_SCHOOL_PAPER,))
        self.assertEqual(mc.PDF_ORIGINS["workbook"],
                         (mc.OUTPUT_SCHOOL_WORKBOOK,))
