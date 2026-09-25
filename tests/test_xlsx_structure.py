"""Synthetic XLSX structure contracts on the deployed main runtime.

All workbooks are built in-test with openpyxl; no real templates or school
files are used.
"""
import json
import tempfile
from pathlib import Path
import unittest
import zipfile
import xml.etree.ElementTree as ET

from tests._harness import ContractCase, bind_major, runtime

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _workbook(path, *, hidden=False, print_area=False):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws["A1"] = "header"
    if hidden:
        wb.create_sheet("Hidden").sheet_state = "hidden"
    if print_area:
        ws.print_area = "A1:B9"
    wb.save(path)


def _strip_defined_names(source, dest):
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(dest, "w") as zout:
        for entry in zin.infolist():
            data = zin.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                tree = ET.fromstring(data)
                for names in tree.findall("{%s}definedNames" % NS):
                    tree.remove(names)
                data = ET.tostring(tree)
            zout.writestr(entry, data)


def _map_file(folder, printing, source):
    m = Path(folder) / "map.json"
    m.write_text(json.dumps({
        "schema": printing.SCHEMA,
        "source": {"sha256": printing.sha256(source)},
        "sheets": [{"sheet": "Visible", "reason": "test", "fit_width": 1,
                    "fit_height": 0, "print_titles": {"rows": [1]}}],
    }), encoding="utf-8")
    return m


class XlsxInspectTests(ContractCase):
    def test_corrupt_xlsx_returns_issue(self):
        """H7 (fixed in P1): a broken XLSX must become a structured issue,
        not a propagated BadZipFile. Baseline red: inspect_outputs raised
        BadZipFile on main 29abec0."""
        sheet = runtime("gg_school_excel")
        with tempfile.TemporaryDirectory(prefix="knuaf-xlsx-") as tmp:
            root = Path(tmp)
            (root / "broken.xlsx").write_bytes(b"NOT A ZIP FILE")
            p = {"outputs": {"broken": {"path": "broken.xlsx",
                                        "format": "xlsx"}}}
            issues = sheet.inspect_outputs(root, p)
        self.assertIsInstance(issues, list)
        self.assertTrue(
            any("zip" in str(r[1]).lower() for r in issues),
            f"expected a structured zip issue, got {issues!r}",
        )

    def test_visible_sheet_count(self):
        """H15-count (fixed in P1): hidden sheets must not count toward the
        worksheet/page lower bound. Baseline red: returned 2 for a
        visible+hidden workbook on main 29abec0."""
        office = runtime("gg_office")
        with tempfile.TemporaryDirectory(prefix="knuaf-xlsx-") as tmp:
            src = Path(tmp) / "source.xlsx"
            _workbook(src, hidden=True)
            self.assertEqual(1, office.workbook_worksheet_count(src))

    def test_all_hidden_workbook_rejected(self):
        """H15-count edge: a workbook with zero visible sheets raises a
        clear error instead of returning a 0 lower bound."""
        office = runtime("gg_office")
        with tempfile.TemporaryDirectory(prefix="knuaf-xlsx-") as tmp:
            src = Path(tmp) / "all_hidden.xlsx"
            _workbook(src)
            # openpyxl refuses to save a fully hidden workbook, so mark the
            # sheet hidden by editing xl/workbook.xml inside the zip.
            raw = src.read_bytes()
            with zipfile.ZipFile(src) as zin, \
                    zipfile.ZipFile(src.parent / "hidden.xlsx", "w") as zout:
                for entry in zin.infolist():
                    data = zin.read(entry.filename)
                    if entry.filename == "xl/workbook.xml":
                        data = data.replace(
                            b'state="visible"', b'state="hidden"')
                    zout.writestr(entry, data)
            hidden = src.parent / "hidden.xlsx"
            with self.assertRaisesRegex(ValueError, "visible"):
                office.workbook_worksheet_count(hidden)

    def test_valid_workbook_passes_without_school_judgment(self):
        """Positive: a non-school workbook is inspected without raising."""
        sheet = runtime("gg_school_excel")
        with tempfile.TemporaryDirectory(prefix="knuaf-xlsx-") as tmp:
            root = Path(tmp)
            _workbook(root / "plain.xlsx")
            p = {"outputs": {"o": {"path": "plain.xlsx", "format": "xlsx"}}}
            issues = sheet.inspect_outputs(root, p)
            self.assertIsInstance(issues, list)


class PrintTitleTests(ContractCase):
    def _ctx(self):
        """Policy B: a bound specialty project and explicit major for the
        print-layout copy (the workbooks stay synthetic)."""
        project = self.make_project()
        bind_major(project)
        return runtime("gg_major_contract").output_context(
            project, "specialty_crops")

    def test_titles_container_created_when_missing(self):
        """H6 (fixed in P1): print_titles on a workbook with no definedNames
        element must create the container, not a bare definedName.
        Baseline red: apply wrote a bare <definedName> sibling and
        openpyxl could not read print_title_rows on main 29abec0."""
        printing = runtime("gg_excel_print")
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(prefix="knuaf-xlsx-") as tmp:
            root = Path(tmp)
            src = root / "source.xlsx"
            _workbook(src)
            stripped = root / "stripped.xlsx"
            _strip_defined_names(src, stripped)
            mapping = _map_file(root, printing, stripped)
            out = root / "out.xlsx"
            printing.apply(stripped, mapping, out, root / "receipt.json",
                           context=self._ctx())
            with zipfile.ZipFile(out) as z:
                tree = ET.fromstring(z.read("xl/workbook.xml"))
            self.assertEqual(1, len(tree.findall(
                "{%s}definedNames/{%s}definedName" % (NS, NS))),
                "print title must sit inside a definedNames container")
            self.assertEqual(0, len(tree.findall(
                "{%s}definedName" % NS)),
                "bare definedName sibling must not be written")
            opened = load_workbook(out)
            self.assertEqual("$1:$1", opened["Visible"].print_title_rows)
            opened.close()

    def test_titles_apply_is_idempotent(self):
        """H6 edge: re-running apply on an already-fixed workbook keeps a
        single container and rewrites the same $1:$1 title."""
        printing = runtime("gg_excel_print")
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(prefix="knuaf-xlsx-") as tmp:
            root = Path(tmp)
            src = root / "source.xlsx"
            _workbook(src)
            stripped = root / "stripped.xlsx"
            _strip_defined_names(src, stripped)
            mapping = _map_file(root, printing, stripped)
            out1 = root / "out1.xlsx"
            printing.apply(stripped, mapping, out1, root / "receipt1.json",
                           context=self._ctx())
            mapping2 = _map_file(root, printing, out1)
            out2 = root / "out2.xlsx"
            printing.apply(out1, mapping2, out2, root / "receipt2.json",
                           context=self._ctx())
            with zipfile.ZipFile(out2) as z:
                tree = ET.fromstring(z.read("xl/workbook.xml"))
            self.assertEqual(1, len(tree.findall(
                "{%s}definedNames" % NS)))
            opened = load_workbook(out2)
            self.assertEqual("$1:$1", opened["Visible"].print_title_rows)
            opened.close()

    def test_titles_with_existing_defined_names(self):
        """Positive: when the container already exists the same apply
        produces a readable $1:$1 print title."""
        printing = runtime("gg_excel_print")
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(prefix="knuaf-xlsx-") as tmp:
            root = Path(tmp)
            src = root / "source.xlsx"
            _workbook(src, print_area=True)
            mapping = _map_file(root, printing, src)
            out = root / "out.xlsx"
            printing.apply(src, mapping, out, root / "receipt.json",
                           context=self._ctx())
            opened = load_workbook(out)
            self.assertEqual("$1:$1", opened["Visible"].print_title_rows)
            opened.close()
            with zipfile.ZipFile(out) as z:
                tree = ET.fromstring(z.read("xl/workbook.xml"))
            self.assertEqual(
                0, len(tree.findall("{%s}definedName" % NS)))


class MergeMarkerTests(ContractCase):
    def test_merge_markers_are_empty(self):
        """H1 (fixed in P1): '<' and '^' are merge continuation markers, not
        filled data cells. Baseline red: both returned True on main 29abec0."""
        doc = runtime("gg_document")
        self.assertFalse(doc._filled_cell("<"))
        self.assertFalse(doc._filled_cell("^"))
        # 유효 값은 여전히 채워진 칸으로 보존
        self.assertTrue(doc._filled_cell("내용"))
        self.assertTrue(doc._filled_cell("0"))


if __name__ == "__main__":
    unittest.main()
