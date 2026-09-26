"""Small, deliberately bounded evaluator for proposed Excel formulas in JSON.

It evaluates the actual formula text, including lazily selected IF branches. A
missing input is None; formula errors are represented by FormulaError.
"""
import re


class FormulaError(Exception):
    pass


TOKEN = re.compile(r'''\s*(?:
    (?P<string>"(?:[^"]|"")*")|
    (?P<ref>(?:(?:'[^']+'|[\w.가-힣]+)!)?\$?[A-Z]{1,3}\$?\d+)|
    (?P<number>\d+(?:\.\d+)?)|
    (?P<op><=|>=|<>|[+*/^(),:;=<>&-])|
    (?P<name>[A-Za-z_][A-Za-z_0-9.]*)
)''', re.X)


def tokenize(formula):
    body = formula[1:] if formula.startswith('=') else formula
    out = []
    pos = 0
    while pos < len(body):
        match = TOKEN.match(body, pos)
        if not match:
            raise FormulaError('unsupported syntax at offset %d' % pos)
        out.append((match.lastgroup, match.group(match.lastgroup)))
        pos = match.end()
    out.append(('end', ''))
    return out


class Parser:
    def __init__(self, formula):
        self.tokens = tokenize(formula)
        self.index = 0

    def peek(self):
        return self.tokens[self.index][1]

    def pop(self):
        item = self.tokens[self.index]
        self.index += 1
        return item

    def expect(self, value):
        if self.pop()[1] != value:
            raise FormulaError('expected %s' % value)

    def parse(self):
        tree = self.expr(0)
        if self.tokens[self.index][0] != 'end':
            raise FormulaError('trailing expression')
        return tree

    def expr(self, min_bp):
        kind, value = self.pop()
        if value in ('+', '-'):
            left = ('unary', value, self.expr(70))
        elif value == '(':
            left = self.expr(0)
            self.expect(')')
        elif kind == 'number':
            left = ('number', float(value))
        elif kind == 'string':
            left = ('string', value[1:-1].replace('""', '"'))
        elif kind == 'ref':
            left = ('ref', value)
        elif kind == 'name':
            if self.peek() == '(':
                self.pop()
                args = []
                if self.peek() != ')':
                    while True:
                        args.append(self.expr(0))
                        if self.peek() not in (',', ';'):
                            break
                        self.pop()
                self.expect(')')
                left = ('call', value.upper(), args)
            else:
                left = ('name', value.upper())
        else:
            raise FormulaError('invalid expression')
        precedence = {'=': 10, '<>': 10, '<': 10, '>': 10, '<=': 10,
                      '>=': 10, '&': 15, '+': 20, '-': 20, '*': 30, '/': 30,
                      '^': 40, ':': 50}
        while self.peek() in precedence:
            op = self.peek()
            bp = precedence[op]
            if bp < min_bp:
                break
            self.pop()
            right = self.expr(bp + (0 if op == '^' else 1))
            left = ('binary', op, left, right)
        return left


def normalize_ref(raw, sheet):
    if '!' in raw:
        sh, cell = raw.rsplit('!', 1)
        sheet = sh.strip("'").replace("''", "'")
    else:
        cell = raw
    return sheet, cell.replace('$', '').upper()


def referenced_cells(formula, sheet):
    """Return every direct coordinate, expanding rectangular range references."""
    root=Parser(formula).parse()
    found=set()
    def walk(node):
        if node[0]=='ref':
            found.add(normalize_ref(node[1],sheet))
        elif node[0]=='binary' and node[1]==':':
            a,b=node[2],node[3]
            if a[0]!='ref' or b[0]!='ref':
                raise FormulaError('unsupported range reference')
            s1,c1=normalize_ref(a[1],sheet)
            s2,c2=normalize_ref(b[1],s1)
            if s1!=s2: raise FormulaError('cross-sheet range')
            from openpyxl.utils.cell import range_boundaries, get_column_letter
            x1,y1,x2,y2=range_boundaries(c1+':'+c2)
            for y in range(y1,y2+1):
                for x in range(x1,x2+1):
                    found.add((s1,f'{get_column_letter(x)}{y}'))
        else:
            for child in node[1:]:
                if isinstance(child,tuple): walk(child)
                elif isinstance(child,list):
                    for n in child:
                        if isinstance(n,tuple): walk(n)
    walk(root)
    return found


def _number(value):
    if value is None or value == '':
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    raise FormulaError('non-numeric operand')


def _truth(value):
    return bool(value) if value is not None else False


def _flatten(values):
    for value in values:
        if isinstance(value, list):
            yield from _flatten(value)
        else:
            yield value


def _criterion(value, rule):
    if isinstance(rule,str):
        match=re.fullmatch(r'(<=|>=|<>|=|<|>)?(.*)',rule)
        op, target=match.groups()
        op=op or '='
        if target=='*': return value not in (None,'')
        try: target=float(target); value=_number(value)
        except (ValueError,FormulaError):
            value='' if value is None else str(value)
        return {'=':lambda:value==target,'<>':lambda:value!=target,
                '<':lambda:value<target,'>':lambda:value>target,
                '<=':lambda:value<=target,'>=':lambda:value>=target}[op]()
    return value==rule


class Evaluator:
    def __init__(self, actions, inputs=None, defaults=None):
        self.formulas = {(a['sheet'], a['cell'].upper()): a['after_formula']
                         for a in actions if a.get('after_formula')}
        self.inputs = inputs or {}
        self.defaults = defaults or {}
        self.stack = set()
        self.cache = {}

    def cell(self, sheet, cell):
        key = (sheet, cell.upper())
        if key in self.inputs:
            return self.inputs[key]
        if key in self.cache:
            return self.cache[key]
        if key in self.stack:
            raise FormulaError('circular reference')
        if key not in self.formulas:
            return self.defaults.get(key)
        self.stack.add(key)
        try:
            result = self.eval(Parser(self.formulas[key]).parse(), sheet)
            self.cache[key] = result
            return result
        finally:
            self.stack.remove(key)

    def eval(self, node, sheet):
        kind = node[0]
        if kind in ('number', 'string'):
            return node[1]
        if kind == 'ref':
            return self.cell(*normalize_ref(node[1], sheet))
        if kind == 'name':
            if node[1] == 'TRUE': return True
            if node[1] == 'FALSE': return False
            raise FormulaError('unsupported name')
        if kind == 'unary':
            n = _number(self.eval(node[2], sheet))
            return n if node[1] == '+' else -n
        if kind == 'binary':
            op = node[1]
            if op == ':':
                a, b = node[2], node[3]
                if a[0] != 'ref' or b[0] != 'ref':
                    raise FormulaError('unsupported range')
                s1, c1 = normalize_ref(a[1], sheet)
                s2, c2 = normalize_ref(b[1], s1)
                if s1 != s2:
                    raise FormulaError('cross-sheet range')
                def split(cell):
                    col, row = re.fullmatch(r'([A-Z]+)(\d+)', cell).groups()
                    n = 0
                    for ch in col: n = n * 26 + ord(ch) - 64
                    return n, int(row)
                x1, y1 = split(c1); x2, y2 = split(c2)
                out = []
                for y in range(min(y1,y2), max(y1,y2)+1):
                    for x in range(min(x1,x2), max(x1,x2)+1):
                        n = x; col = ''
                        while n:
                            n, rem = divmod(n-1,26); col = chr(65+rem)+col
                        out.append(self.cell(s1, f'{col}{y}'))
                return out
            a = self.eval(node[2], sheet)
            b = self.eval(node[3], sheet)
            if op in ('=', '<>', '<', '>', '<=', '>='):
                a = '' if a is None else a
                b = '' if b is None else b
                try:
                    return {'=': lambda:a==b, '<>':lambda:a!=b, '<':lambda:a<b,
                            '>':lambda:a>b, '<=':lambda:a<=b, '>=':lambda:a>=b}[op]()
                except TypeError:
                    return False
            if op == '&': return str(a if a is not None else '')+str(b if b is not None else '')
            a = _number(a); b = _number(b)
            if op == '+': return a+b
            if op == '-': return a-b
            if op == '*': return a*b
            if op == '/':
                if b == 0: raise FormulaError('division by zero')
                return a/b
            if op == '^': return a**b
        if kind == 'call':
            fn, args = node[1:]
            if fn == 'IF':
                if len(args) not in (2,3): raise FormulaError('IF arity')
                return self.eval(args[1 if _truth(self.eval(args[0], sheet)) else 2], sheet) if len(args)==3 or _truth(self.eval(args[0],sheet)) else False
            if fn == 'IFERROR':
                try: return self.eval(args[0], sheet)
                except FormulaError: return self.eval(args[1], sheet)
            if fn == 'ISNUMBER':
                try: val = self.eval(args[0], sheet)
                except FormulaError: return False
                return isinstance(val, (int,float)) and not isinstance(val,bool)
            if fn == 'NA': raise FormulaError('#N/A')
            vals = [self.eval(a, sheet) for a in args]
            if fn == 'AND': return all(_truth(v) for v in vals)
            if fn == 'OR': return any(_truth(v) for v in vals)
            if fn == 'NOT': return not _truth(vals[0])
            flat = list(_flatten(vals))
            nums = [float(v) for v in flat if isinstance(v,(int,float)) and not isinstance(v,bool)]
            if fn == 'SUM': return sum(nums)
            if fn == 'MAX': return max(nums) if nums else 0
            if fn == 'MIN': return min(nums) if nums else 0
            if fn == 'COUNT': return len(nums)
            if fn == 'AVERAGE': return sum(nums)/len(nums) if nums else 0
            if fn == 'LEN': return len(str(vals[0] or ''))
            if fn == 'RIGHT': return str(vals[0] or '')[-int(_number(vals[1])):]
            if fn == 'ROUND': return round(_number(vals[0]), int(_number(vals[1])))
            if fn == 'INT': return int(_number(vals[0])//1)
            if fn == 'TEXT': return str(vals[0] if vals[0] is not None else '')
            if fn in ('COUNTIF','SUMIF'):
                sample=vals[0] if isinstance(vals[0],list) else [vals[0]]
                matched=[i for i,v in enumerate(sample) if _criterion(v,vals[1])]
                if fn=='COUNTIF': return len(matched)
                sums=vals[2] if len(vals)>2 and isinstance(vals[2],list) else sample
                return sum(_number(sums[i]) for i in matched if i<len(sums))
            if fn in ('SUMIFS','COUNTIFS'):
                if fn=='SUMIFS':
                    sums=vals[0] if isinstance(vals[0],list) else [vals[0]]
                    criteria=vals[1:]
                else:
                    criteria=vals
                    sums=criteria[0] if isinstance(criteria[0],list) else [criteria[0]]
                count=0;total=0.0
                for i in range(len(sums)):
                    if all(_criterion((criteria[j] if isinstance(criteria[j],list) else [criteria[j]])[i],
                                      criteria[j+1]) for j in range(0,len(criteria),2)):
                        count+=1;total+=_number(sums[i])
                return total if fn=='SUMIFS' else count
            if fn == 'PMT':
                rate,nper,pv=map(_number,vals[:3])
                if nper<=0: raise FormulaError('invalid PMT term')
                return -pv/nper if rate==0 else -pv*rate*(1+rate)**nper/((1+rate)**nper-1)
            if fn in ('DATE','YEAR','MONTH','DAY'):
                import datetime
                epoch=datetime.datetime(1899,12,30)
                if fn=='DATE':
                    d=datetime.datetime(int(_number(vals[0])),int(_number(vals[1])),int(_number(vals[2])))
                    return (d-epoch).days
                d=epoch+datetime.timedelta(days=int(_number(vals[0])))
                return {'YEAR':d.year,'MONTH':d.month,'DAY':d.day}[fn]
            raise FormulaError('unsupported function '+fn)
        raise FormulaError('unsupported AST')
