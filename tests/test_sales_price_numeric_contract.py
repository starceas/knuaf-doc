"""DR-01..05: independent arithmetic, refusal paths, and actual workbook routes.

All prices in these tests are explicitly synthetic official-schema declarations.
"""
import copy
import json
from decimal import Decimal
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from tests._harness import ContractCase, bind_major, runtime
from tests.test_price_assumptions import _spec, _student_source
from tests.test_template_wage_link import fixture, assumptions, S5
from tests.test_hort_price_contract import _fact, _project


def source(sid="synthetic.purchase", values=(100,200,300,400,500), channel="수매", month=None):
    entry = _student_source(sid, observations=[{"year":y,"value":v} for y,v in zip(range(2021,2026),values)])
    entry["dimensions"]["channel"] = channel
    if month is not None:
        entry["dimensions"]["month"] = str(month)
    return entry


def sales_spec():
    spec = _spec(purchase_price="300",direct_price="500",
        price_dimensions={"unit":"원/kg"},
        price_sources=[source(),source("synthetic.direct",(300,400,500,600,700),"직거래")],
        price_assumptions={
            "general":{"source_id":"official.kosis.cpi.total","base_year":2020,"application_base_year":2027},
            "wage":{"source_id":"official.kosis.farm_purchase.labor","base_year":2020,"application_base_year":2027},
            "sales":{"application_base_year":2027,"anchors":{
                "purchase":{"source_id":"synthetic.purchase","base_year":None,"target_dimensions":{"channel":"수매","grade":"상품"}},
                "direct":{"source_id":"synthetic.direct","base_year":None,"target_dimensions":{"channel":"직거래","grade":"상품"}}}}})
    for entry in spec["price_sources"]:
        if entry["id"] in ("synthetic.purchase", "synthetic.direct"):
            entry["dimensions"]["grade"] = "상품"
    return spec


def horticulture_inputs():
    h = runtime("gg_hort_workbook")
    public = h.contracts()["finance-params.json"]["parameters"]
    keys = [k for k in public if k.startswith("price.observed.y") or
            k.startswith("sales.y") and k.endswith(".price_krw_per_unit")]
    keys += ["plan.start_year","sales.crop_item","sales.price_per_unit"]
    params = {k:copy.deepcopy(public[k]) for k in keys}
    # Exercise this complete pricing subgraph without unrelated farm inputs.
    for p in params.values(): p["required_if"]={"any":[]}
    facts = [_fact("plan.start_year",2027),_fact("sales.crop_item","합성작목"),_fact("sales.price_per_unit","kg")]
    sources = {"s1":{"revision":1,"kind":"stat","source_type":"official_stat","grade":1}}
    for month in range(1,13):
        sid=f"synthetic.month.{month}"
        sources[sid]={"revision":1,"kind":"stat","source_type":"official_stat","grade":1,
                      "price_series":source(sid,tuple(month*100+v for v in (10,20,30,40,50)),"농가수취",month)}
        facts.append(_fact(f"sales.y1.m{month}.price_krw_per_unit",month*100+30,
            refs=[{"id":sid,"revision":1}],price_assumption={"writing_year":2026,
                "application_base_year":2027,"source_id":sid,"base_year":None,
                "target_dimensions":{"crop":"합성작목","region":"전국","unit":"원/kg","channel":"농가수취"}}))
    return _project(facts,sources),params


class SalesNumericContract(ContractCase):
    def setUp(self):
        self.pa=runtime("gg_price_assumptions")

    def normalize(self,spec=None,registry=None,years=(2027,2028,2029,2030,2031)):
        return self.pa.normalize_price_role(spec or sales_spec(),"sales",list(years),registry=registry)

    def test_mean_and_cpi_are_independent(self):
        a=self.normalize(); r=(116.61/102.5)**.25-1
        self.assertEqual(300,a["anchors"]["purchase"]["base_price"])
        self.assertEqual(500,a["anchors"]["direct"]["base_price"])
        self.assertAlmostEqual(r,a["rate"],places=14)
        self.assertNotAlmostEqual((500/100)**.25-1,a["rate"])
        self.assertEqual(4,a["interval_years"])
        for i,y in enumerate(range(2027,2032)):
            self.assertAlmostEqual(300*(1+r)**i,300*self.pa.multiplier(a,y),places=10)
        self.assertEqual(1,self.pa.multiplier(a,2027))

    def test_start_year_moves_without_observation_shift(self):
        spec=sales_spec(); spec["price_assumptions"]["sales"]["application_base_year"]=2030
        a=self.normalize(spec,years=range(2030,2035)); b=self.normalize()
        self.assertEqual(a["price_observations"],b["price_observations"])
        self.assertEqual(a["rate"],b["rate"])
        self.assertEqual(1,self.pa.multiplier(a,2030))

    def test_writing_year_changes_required_window(self):
        spec=sales_spec(); spec["writing_year"]=2027
        with self.assertRaises(self.pa.SalesPriceHeld) as cm:self.normalize(spec)
        self.assertEqual([2026],cm.exception.detail["missing_years"])

    def test_missing_writing_year_never_derived_from_start(self):
        spec=sales_spec(); del spec["writing_year"]
        with self.assertRaisesRegex(ValueError,"writing_year"):self.normalize(spec)

    def test_application_base_is_exact_start(self):
        for wrong in (2026,2028,True,"2027"):
            spec=sales_spec();spec["price_assumptions"]["sales"]["application_base_year"]=wrong
            with self.subTest(wrong=wrong),self.assertRaisesRegex(ValueError,"첫 계획연도"):self.normalize(spec)

    def test_price_missing_middle_four_duplicate_future_and_forged(self):
        for mutation in ("middle","four","duplicate","future","forged"):
            spec=sales_spec(); entry=next(e for e in spec["price_sources"] if e["id"]=="synthetic.purchase")
            raw=spec["price_assumptions"]["sales"]["anchors"]["purchase"]
            if mutation=="middle":entry["observations"].pop(2)
            if mutation=="four":entry["observations"].pop()
            if mutation=="duplicate":entry["observations"][2]["year"]=2022
            if mutation=="future":entry["observations"][-1]["year"]=2026
            if mutation=="forged":raw.update(observation_years=list(range(2021,2026)),observed_values=[100,200,301,400,500])
            with self.subTest(mutation=mutation),self.assertRaises(ValueError):self.normalize(spec)

    def test_cpi_missing_middle_and_duplicate(self):
        for mutation in ("middle","duplicate"):
            registry=copy.deepcopy(self.pa.load_registry()); obs=registry[self.pa.CPI_SOURCE_ID]["observations"]
            if mutation=="middle":obs[:]=[o for o in obs if o["year"]!=2023]
            else:obs.append(copy.deepcopy(obs[2]))
            with self.subTest(mutation=mutation),self.assertRaises(ValueError):self.normalize(registry=registry)

    def test_wrong_cpi_metadata_and_substitute(self):
        for key,value in (("table_id","DT_1J60"),("frequency","monthly"),("item_name","식료품"),("dimensions",{"region":"제주","item":"총지수"}),("base_year",2021)):
            registry=copy.deepcopy(self.pa.load_registry());registry[self.pa.CPI_SOURCE_ID][key]=value
            with self.subTest(key=key),self.assertRaises(ValueError):self.normalize(registry=registry)
        spec=sales_spec();spec["price_assumptions"]["sales"]["cpi_source_id"]="official.kosis.farm_sales.total"
        with self.assertRaisesRegex(ValueError,"DT_1J22003"):self.normalize(spec)

    def test_not_applied_status_and_mode_cannot_bypass(self):
        for field in ("status","mode"):
            spec=sales_spec();spec["price_assumptions"]["sales"][field]="not_applied"
            with self.subTest(field=field),self.assertRaisesRegex(ValueError,"not_applied"):self.normalize(spec)

    def test_hold_preserves_research_and_selection_without_approving(self):
        for research in ({"status":"not_started","candidates":[]},
                         {"status":"failed","query":"official price","retrieved_at":"2026-10-13","result":"connection failed","candidates":[]},
                         {"status":"searched","candidates":[{"id":"candidate","years":[2021,2022],"unit":"원/kg","evidence":"official lookup","condition_difference":"missing years"}]}):
            spec=sales_spec(); raw=spec["price_assumptions"]["sales"];raw["research"]=research
            raw["selection"]={"status":"selected","candidate_id":"candidate","answer_ref":"synthetic-answer"}
            next(e for e in spec["price_sources"] if e["id"]=="synthetic.purchase")["observations"].pop()
            with self.subTest(research=research),self.assertRaises(self.pa.SalesPriceHeld) as cm:self.normalize(spec)
            self.assertEqual(research,cm.exception.detail["research"])
            self.assertEqual(raw["selection"],cm.exception.detail["selection"])
            self.assertEqual([2025],cm.exception.detail["missing_years"])

    def test_cost_role_needs_no_sales_placeholder(self):
        for role,sid in (("general","official.kosis.cpi.total"),("wage","official.kosis.farm_purchase.labor")):
            a=self.pa.normalize_price_role({"price_assumptions":{role:{"source_id":sid,"base_year":2020,"application_base_year":2027}}},role,[2027])
            self.assertEqual("series_cagr",a["rate_source"])
        self.assertEqual({2021:102.5,2025:116.61},self.pa.check_observations(self.pa.load_registry()[self.pa.CPI_SOURCE_ID],[2021,2025],[102.5,116.61]))

    def test_zero_and_negative_complete_cpi_are_valid(self):
        for r in (0,-.02):
            registry=copy.deepcopy(self.pa.load_registry())
            registry[self.pa.CPI_SOURCE_ID]["observations"]=[{"year":y,"value":100*(1+r)**i} for i,y in enumerate(range(2021,2026))]
            a=self.normalize(registry=registry)
            self.assertAlmostEqual(r,a["rate"],places=12)

    def test_school_channels_and_conflicting_direct_value(self):
        root=self.make_project();bind_major(root);ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        out=root/"sales.xlsx";spec=sales_spec()
        runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
        wb=load_workbook(out);sales=wb[" 5. 판매계획"]
        self.assertEqual("=AVERAGE(E7,G7,I7,K7,M7)",sales["O7"].value)
        self.assertEqual([100,200,300,400,500],[sales[c+"7"].value for c in "EGIKM"])
        self.assertEqual([300,400,500,600,700],[sales[c+"10"].value for c in "EGIKM"])
        self.assertEqual("=E22*$O$7/1000",sales["D39"].value)
        self.assertEqual("=E26*$O$10/1000",sales["D40"].value)
        expected=runtime("gg_school_verify").calculate_school_expected(spec)
        self.assertEqual(Decimal(800),expected["expected_years"][0]["revenue"])
        self.assertAlmostEqual(800*((116.61/102.5)**.25)**2,float(expected["expected_years"][2]["revenue"]),places=10)
        receipt=json.loads(out.with_suffix(".manifest.json").read_text())
        self.assertEqual(500,receipt["price_assumptions"]["sales"]["anchors"]["direct"]["base_price"])
        wb.close();spec["purchase_price"]="999"
        with self.assertRaisesRegex(ValueError,"불일치"):runtime("gg_school_excel").school_workbook(spec,root/"bad.xlsx",context=ctx)
        self.assertFalse((root/"bad.xlsx").exists())

    def test_school_valid_lower_grade_records_refused_before_output(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        school=runtime("gg_school_excel")
        for key in ("purchase","direct"):
            for grade in ("중품","하품"):
                spec=sales_spec()
                raw=spec["price_assumptions"]["sales"]["anchors"][key]
                raw["target_dimensions"]["grade"]=grade
                next(e for e in spec["price_sources"] if e["id"]==raw["source_id"])["dimensions"]["grade"]=grade
                # Valid source/own target agreement is insufficient for
                # the school workbook's fixed 상품 output row.
                self.assertEqual(grade,self.normalize(spec)["anchors"][key]["dimensions"]["grade"])
                out=root/(key+grade+".xlsx")
                with self.subTest(key=key,grade=grade),self.assertRaisesRegex(ValueError,"school.*grade"):
                    school.school_workbook(spec,out,context=ctx)
                self.assertFalse(out.exists())
                self.assertFalse(out.with_suffix(".manifest.json").exists())
                # Refusal also preserves an already-existing output pair.
                school.school_workbook(sales_spec(),out,context=ctx)
                previous={p:p.read_bytes() for p in (out,out.with_suffix(".manifest.json"))}
                before=copy.deepcopy(spec)
                with self.assertRaisesRegex(ValueError,"school.*grade"):
                    school.school_workbook(spec,out,context=ctx)
                self.assertEqual(before,spec)
                self.assertEqual(previous,{p:p.read_bytes() for p in previous})

    def test_school_unknown_source_grade_cannot_be_filled_as_product(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        for key in ("purchase","direct"):
            spec=sales_spec()
            raw=spec["price_assumptions"]["sales"]["anchors"][key]
            raw["target_dimensions"].pop("grade")
            next(e for e in spec["price_sources"] if e["id"]==raw["source_id"])["dimensions"].pop("grade")
            self.normalize(spec)  # shared contract still permits unknown grade
            before=copy.deepcopy(spec);out=root/(key+"-unknown.xlsx")
            with self.subTest(key=key),self.assertRaisesRegex(ValueError,"grade"):
                runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
            self.assertFalse(out.exists());self.assertFalse(out.with_suffix(".manifest.json").exists())
            self.assertEqual(before,spec)

    def test_school_known_product_source_binds_omitted_target_grade(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        spec=sales_spec()
        for raw in spec["price_assumptions"]["sales"]["anchors"].values():
            raw["target_dimensions"].pop("grade")
        before=copy.deepcopy(spec);out=root/"known-product.xlsx"
        runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
        receipt=json.loads(out.with_suffix(".manifest.json").read_text())
        for a in receipt["price_assumptions"]["sales"]["anchors"].values():
            self.assertEqual("상품",a["dimensions"]["grade"])
            self.assertEqual("상품",a["target_dimensions"]["grade"])
        self.assertEqual(before,spec)
        book=runtime("gg_workbook_audit").Workbook.load(out)
        self.assertEqual(300,book.evaluate(S5,"D39"))
        self.assertEqual(500,book.evaluate(S5,"D40"))
        self.assertAlmostEqual(800*((116.61/102.5)**.25)**2,
            book.evaluate(S5,"F39")+book.evaluate(S5,"F40"),places=10)

    def test_school_inactive_unknown_grade_remains_input_only(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        spec=sales_spec();spec.pop("direct_price")
        spec.update(purchase_share="1",direct_share="0")
        raw=spec["price_assumptions"]["sales"]["anchors"]["direct"]
        raw["target_dimensions"].pop("grade")
        next(e for e in spec["price_sources"] if e["id"]==raw["source_id"])["dimensions"].pop("grade")
        before=copy.deepcopy(spec);out=root/"inactive-unknown.xlsx"
        runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
        book=runtime("gg_workbook_audit").Workbook.load(out)
        self.assertEqual(600,book.evaluate(S5,"D39"))
        self.assertEqual(0,book.evaluate(S5,"O10"))
        receipt=json.loads(out.with_suffix(".manifest.json").read_text())
        self.assertEqual({"purchase"},set(receipt["price_assumptions"]["sales"]["anchors"]))
        self.assertEqual(before,spec)

    def test_school_row_dimensions_accept_aliases_and_singleton_lists(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        spec=sales_spec()
        for raw in spec["price_assumptions"]["sales"]["anchors"].values():
            e=next(e for e in spec["price_sources"] if e["id"]==raw["source_id"])
            for dims in (raw["target_dimensions"],e["dimensions"]):
                channel,grade=dims.pop("channel"),dims.pop("grade")
                # Source schema requires strings; target conditions also
                # accept equivalent singleton lists under check_dimensions.
                dims["거래단계"]=channel if dims is e["dimensions"] else [channel]
                dims["등급"]=grade if dims is e["dimensions"] else [grade]
                dims["단위"]="원/kg" if dims is e["dimensions"] else ["원/kg"]
        out=root/"aliases.xlsx"
        runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
        self.assertEqual(Decimal(800),runtime("gg_school_verify").calculate_school_expected(spec)["expected_years"][0]["revenue"])

    def test_school_valid_wrong_channel_or_unit_refused_without_conversion(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        for axis,value in (("channel","도매"),("unit","원/10kg")):
            spec=sales_spec()
            for raw in spec["price_assumptions"]["sales"]["anchors"].values():
                e=next(e for e in spec["price_sources"] if e["id"]==raw["source_id"])
                raw["target_dimensions"][axis]=value;e["dimensions"][axis]=value
                if axis=="unit":e["unit"]=value
            if axis=="unit":spec["price_dimensions"]["unit"]=value
            self.normalize(spec)
            out=root/(axis+"-mismatch.xlsx")
            with self.subTest(axis=axis),self.assertRaisesRegex(ValueError,"채널 불일치|원/kg"):
                runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
            self.assertFalse(out.exists());self.assertFalse(out.with_suffix(".manifest.json").exists())

    def test_fr01_swapped_source_and_target_rejected_before_output(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        spec=sales_spec();spec.pop("purchase_price");spec.pop("direct_price")
        spec.update(purchase_share="0.75",direct_share="0.25")
        anchors=spec["price_assumptions"]["sales"]["anchors"]
        anchors["purchase"],anchors["direct"]=anchors["direct"],anchors["purchase"]
        out=root/"swapped.xlsx"
        before=self._tree_bytes(root)
        with self.assertRaisesRegex(ValueError,"purchase.*수매.*채널 불일치"):
            runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
        self.assertFalse(out.exists());self.assertFalse(out.with_suffix(".manifest.json").exists())
        self.assertEqual(before,self._tree_bytes(root))
        with self.assertRaisesRegex(ValueError,"채널 불일치"):
            runtime("gg_school_verify").calculate_school_expected(spec)

    def test_fr01_asymmetric_correct_channels_revenue_700(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        spec=sales_spec();spec.pop("purchase_price");spec.pop("direct_price")
        spec.update(purchase_share="0.75",direct_share="0.25")
        out=root/"asymmetric.xlsx"
        runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
        book=runtime("gg_workbook_audit").Workbook.load(out)
        self.assertEqual(450,book.evaluate(" 5. 판매계획","D39"))
        self.assertEqual(250,book.evaluate(" 5. 판매계획","D40"))
        expected=runtime("gg_school_verify").calculate_school_expected(spec)
        self.assertEqual(Decimal(700),expected["expected_years"][0]["revenue"])
        manifest=json.loads(out.with_suffix(".manifest.json").read_text())
        for key,channel in (("purchase","수매"),("direct","직거래")):
            self.assertEqual(channel,manifest["price_assumptions"]["sales"]["anchors"][key]["dimensions"]["channel"])

    def test_fr02_inactive_direct_source_kept_in_input_revenue_600(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        spec=sales_spec();spec.pop("direct_price")
        spec.update(purchase_share="1",direct_share="0")
        before=copy.deepcopy(spec);out=root/"purchase-only.xlsx"
        runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
        expected=runtime("gg_school_verify").calculate_school_expected(spec)
        self.assertEqual(Decimal(600),expected["expected_years"][0]["revenue"])
        self.assertEqual(Decimal(0),expected["effective_direct_price"])
        book=runtime("gg_workbook_audit").Workbook.load(out)
        self.assertEqual(300,book.evaluate(" 5. 판매계획","O7"))
        self.assertEqual(0,book.evaluate(" 5. 판매계획","O10"))
        self.assertEqual(600,book.evaluate(" 5. 판매계획","D39"))
        self.assertEqual(0,book.evaluate(" 5. 판매계획","D40"))
        manifest=json.loads(out.with_suffix(".manifest.json").read_text())
        self.assertEqual({"purchase"},set(manifest["price_assumptions"]["sales"]["anchors"]))
        self.assertEqual(before,spec)

    def test_fr03_flat_single_channel_acyclic_xlsx_and_manifest(self):
        root=self.make_project();bind_major(root)
        ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        for key,amount in (("purchase",600),("direct",1000)):
            with self.subTest(key=key):
                spec=sales_spec();spec.pop("purchase_price");spec.pop("direct_price")
                spec.update(purchase_share="1" if key=="purchase" else "0",direct_share="1" if key=="direct" else "0")
                spec["price_assumptions"]["sales"]={**spec["price_assumptions"]["sales"]["anchors"][key],"application_base_year":2027}
                _,_,prices=runtime("gg_school_excel").validate(spec)
                self.assertIsNot(prices["sales"],prices["sales"]["anchors"][key])
                json.dumps(prices)
                out=root/(key+"-flat.xlsx")
                runtime("gg_school_excel").school_workbook(spec,out,context=ctx)
                self.assertTrue(out.exists());self.assertTrue(out.with_suffix(".manifest.json").is_file())
                w=load_workbook(out);self.assertEqual(17,len(w.sheetnames));w.close()
                manifest=json.loads(out.with_suffix(".manifest.json").read_text())
                self.assertEqual({key},set(manifest["price_assumptions"]["sales"]["anchors"]))
                expected=runtime("gg_school_verify").calculate_school_expected(spec)
                self.assertEqual(Decimal(amount),expected["expected_years"][0]["revenue"])
                book=runtime("gg_workbook_audit").Workbook.load(out)
                self.assertEqual(amount,book.evaluate(" 5. 판매계획","D39")+book.evaluate(" 5. 판매계획","D40"))

    def test_d8_distinct_actual_anchors(self):
        root=self.make_project();bind_major(root);ctx=runtime("gg_major_contract").output_context(root,"specialty_crops")
        src=root/"source.xlsx";fixture(src);cfg=assumptions()
        for index,(anchor,raw) in enumerate(cfg["sales"]["anchors"].items(),1):
            e=next(e for e in cfg["price_sources"] if e["id"]==raw["source_id"])
            e["observations"]=[{"year":y,"value":index*100+v} for y,v in zip(range(2021,2026),[10,20,30,40,50])]
        patch=runtime("gg_excel_formula_patch");plan=patch.materialize_d8(src,cfg)
        mp=root/"map.json";mp.write_text(json.dumps(plan));out=root/"result.xlsx"
        receipt=patch.patch_copy(src,mp,out,context=ctx);book=runtime("gg_workbook_audit").Workbook.load(out)
        for index,(row,anchor) in enumerate({"38":"O7","39":"O8","40":"O9","41":"O11","42":"O12","43":"O13"}.items(),1):
            b=index*100+30
            self.assertEqual(b,book.evaluate(S5,anchor))
            for year, column in enumerate("DFHJL"):
                self.assertAlmostEqual(b*.1*((116.61/102.5)**.25)**year,
                                       book.evaluate(S5,column+row),places=10)
        self.assertEqual([],patch.materialize_d8(out,cfg)["patches"])
        self.assertEqual(6,len(receipt["d8"]["rates"]["sales"]["anchors"]))

    def d8_inputs(self):
        root = self.make_project()
        src = root / "source.xlsx"
        fixture(src)
        return root, src, assumptions(), runtime("gg_excel_formula_patch")

    def test_d8_swapped_channel_records_rejected(self):
        _, src, cfg, patch = self.d8_inputs()
        for left, right in (("O7", "O11"), ("O8", "O12"), ("O9", "O13")):
            bad = copy.deepcopy(cfg)
            anchors = bad["sales"]["anchors"]
            anchors[left], anchors[right] = anchors[right], anchors[left]
            before = copy.deepcopy(bad)
            with self.subTest(pair=(left, right)), self.assertRaisesRegex(ValueError, "D8.*channel"):
                patch.materialize_d8(src, bad)
            self.assertEqual(before, bad)

    def test_d8_swapped_grade_records_rejected(self):
        _, src, cfg, patch = self.d8_inputs()
        for left, right in (("O7", "O8"), ("O8", "O9"), ("O11", "O12"), ("O12", "O13")):
            bad = copy.deepcopy(cfg)
            anchors = bad["sales"]["anchors"]
            anchors[left], anchors[right] = anchors[right], anchors[left]
            with self.subTest(pair=(left, right)), self.assertRaisesRegex(ValueError, "D8.*grade"):
                patch.materialize_d8(src, bad)

    def test_d8_missing_target_cannot_bypass_row_condition(self):
        _, src, cfg, patch = self.d8_inputs()
        # A valid O7 source normalizes against a matching common target;
        # putting it in O11 must still fail without an anchor-local target.
        target = cfg["sales"]["target_dimensions"]
        target.update(channel="수매", grade="상품")
        raw = cfg["sales"]["anchors"]["O7"]
        raw.pop("target_dimensions")
        cfg["sales"]["anchors"] = {"O11": raw}
        w = load_workbook(src)
        for row in (38, 39, 40, 42, 43):
            for col in "CEGIK": w[S5][col+str(row)] = 0
        w.save(src); w.close()
        with self.assertRaisesRegex(ValueError, "D8.*channel"):
            patch.materialize_d8(src, cfg)

    def test_d8_missing_grade_rejected(self):
        _, src, cfg, patch = self.d8_inputs()
        raw = cfg["sales"]["anchors"]["O7"]
        raw["target_dimensions"].pop("grade")
        source = next(e for e in cfg["price_sources"] if e["id"] == raw["source_id"])
        source["dimensions"].pop("grade")
        with self.assertRaisesRegex(ValueError, "D8.*grade"):
            patch.materialize_d8(src, cfg)

    def test_d8_row_dimensions_accept_aliases_and_singleton_lists(self):
        _, src, cfg, patch = self.d8_inputs()
        for raw in cfg["sales"]["anchors"].values():
            e = next(e for e in cfg["price_sources"] if e["id"] == raw["source_id"])
            raw["target_dimensions"]["거래단계"] = [raw["target_dimensions"].pop("channel")]
            raw["target_dimensions"]["등급"] = [raw["target_dimensions"].pop("grade")]
            e["dimensions"]["거래단계"] = e["dimensions"].pop("channel")
            e["dimensions"]["등급"] = e["dimensions"].pop("grade")
        self.assertEqual(6, len(patch.materialize_d8(src, cfg)["d8"]["rates"]["sales"]["anchors"]))
        # No new common-target input is required when each anchor already
        # declares the full condition under the existing shared contract.
        common = cfg["sales"].pop("target_dimensions")
        for raw in cfg["sales"]["anchors"].values(): raw["target_dimensions"].update(common)
        self.assertEqual(6, len(patch.materialize_d8(src, cfg)["d8"]["rates"]["sales"]["anchors"]))

    def test_d8_only_existing_kg_unit_supported(self):
        _, src, cfg, patch = self.d8_inputs()
        cfg["sales"]["target_dimensions"]["unit"] = "원/g"
        for raw in cfg["sales"]["anchors"].values():
            e = next(e for e in cfg["price_sources"] if e["id"] == raw["source_id"])
            e["unit"] = e["dimensions"]["unit"] = "원/g"
        with self.assertRaisesRegex(ValueError, "원/kg|unit"):
            patch.materialize_d8(src, cfg)

    def test_d8_inactive_missing_anchor_allowed_but_wrong_record_rejected(self):
        _, src, cfg, patch = self.d8_inputs()
        w = load_workbook(src)
        for col in "CEGIK": w[S5][col+"41"] = 0
        w.save(src); w.close()
        cfg["sales"]["anchors"].pop("O11")
        patch.materialize_d8(src, cfg)
        cfg["sales"]["anchors"]["O11"] = copy.deepcopy(cfg["sales"]["anchors"]["O7"])
        with self.assertRaisesRegex(ValueError, "D8.*channel"):
            patch.materialize_d8(src, cfg)

    def test_d8_active_missing_and_unknown_anchors_rejected(self):
        _, src, cfg, patch = self.d8_inputs()
        bad = copy.deepcopy(cfg); bad["sales"]["anchors"].pop("O7")
        with self.assertRaisesRegex(ValueError, "active sales anchor"):
            patch.materialize_d8(src, bad)
        bad = copy.deepcopy(cfg); bad["sales"]["anchors"]["O10"] = bad["sales"]["anchors"].pop("O7")
        with self.assertRaisesRegex(ValueError, "D8 sales anchors"):
            patch.materialize_d8(src, bad)

    def test_d8_common_plan_conflict_rejected_without_publication(self):
        root, src, cfg, patch = self.d8_inputs()
        cfg["sales"]["anchors"]["O7"]["target_dimensions"]["crop"] = "other"
        with self.assertRaises(ValueError): patch.materialize_d8(src, cfg)
        # CLI failure preserves a prior map and creates no new map.
        import contextlib, io
        config = root / "assumptions.json"; config.write_text(json.dumps(cfg))
        for exists in (False, True):
            out = root / ("held-existing.json" if exists else "held-new.json")
            if exists: out.write_text("prior map")
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                rc = patch.cli(["materialize-d8", "--source", str(src), "--assumptions", str(config), "--out", str(out)])
            self.assertNotEqual(0, rc)
            self.assertEqual("prior map" if exists else None, out.read_text() if out.exists() else None)

    def test_hort_all_sixty_prices_and_production_growth_preserved(self):
        h=runtime("gg_hort_workbook");project,params=horticulture_inputs()
        values,derivations,_=h._facts(project,params)
        self.assertEqual(60,len([d for d in derivations if d["rule_id"]=="SALES-FIVE-YEAR-MEAN-CPI"]))
        for month in range(1,13):
            for year in range(1,6):
                key=f"sales.y{year}.m{month}.price_krw_per_unit"
                self.assertAlmostEqual((month*100+30)*((116.61/102.5)**.25)**(year-1),float(values[key]),places=9)
                self.assertIn("AVERAGE" if year==1 else f"^{year-1}",h._sales_price_formula(key,values))
        public=h.contracts()["finance-params.json"]["parameters"]
        self.assertIsNone(public["sales.y2.growth_rate"].get("price_contract"))
        self.assertEqual(" 7. 판매계획!O84",public["sales.y2.growth_rate"]["canonical_cell"])

    def test_hort_partial_month_wrong_month_and_quote_refused(self):
        h=runtime("gg_hort_workbook")
        for mutation in ("partial","month","quote","override"):
            project,params=horticulture_inputs();e=project["sources"]["synthetic.month.1"]["price_series"]
            fact=project["facts"]["f-sales.y1.m1.price_krw_per_unit"]
            if mutation=="partial":e["observations"].pop(2)
            if mutation=="month":e["dimensions"]["month"]="2"
            if mutation=="quote":fact.pop("price_assumption")
            if mutation=="override":fact["value"]=999
            with self.subTest(mutation=mutation),self.assertRaises(h.Held):h._facts(project,params)

    def test_hort_json_partial_average_stays_blank(self):
        h=runtime("gg_hort_workbook");tr=h.contracts()["finance-transform.json"]
        row=next(e for e in tr["entries"] if e["sheet"]==" 7. 판매계획" and e["cell"]=="C13")
        self.assertIn("COUNT(C8:C12)<>5",row["after_formula"])
        evaluator=runtime("gg_hort_formula_eval").Evaluator([row],inputs={(" 7. 판매계획","C8"):100})
        self.assertEqual("",evaluator.cell(" 7. 판매계획","C13"))

    def test_actual_hort_ooxml_prices_reach_visible_revenues(self):
        import zipfile
        from tests.fixtures.hort_finance.synthetic import build
        h=runtime("gg_hort_workbook");project,params=horticulture_inputs()
        values,_,_=h._facts(project,params)
        docs=h.contracts();sheet=" 7. 판매계획"
        root=self.make_project();src=root/"synthetic.xlsx"
        roster=h._json(h.REF/"profile.json")["workbook_reference"]["sheets"]
        build(src,roster)
        original=load_workbook(src);self.assertEqual(roster,original.sheetnames);original.close()
        selected=[]
        for e in docs["finance-transform.json"]["entries"]:
            if e["sheet"]!=sheet:continue
            if e.get("param_key") in params or e["cell"] in [c+str(r) for c in "CDEFGHIJKLMN" for r in [13,24,25,33,34,42,43,51,52,60,61]]:
                selected.append({**e,"before_state":"absent"})
        # Supply real numeric quantities to the actual public revenue formulas.
        for c in "CDEFGHIJKLMN":
            for row in [22,31,40,49,58]:
                selected.append({"sheet":sheet,"cell":c+str(row),"action":"formula",
                                 "before_state":"absent","after_formula":"=10"})
        transform={"entries":selected}
        with mock.patch.object(h,"SOURCE_SHA",h.file_digest(src)):
            payload,counts=h._zip_transform(src,transform,values,params)
        out=root/"actual-output.xlsx";out.write_bytes(payload)
        generated=load_workbook(out);self.assertEqual(roster,generated.sheetnames)
        visible_sheet=generated[sheet]
        with zipfile.ZipFile(out) as z:
            actions=[];inputs={}
            from xml.etree import ElementTree as ET
            for name,path in h.workbook_sheets(z):
                if name!=sheet:continue
                for cell in ET.fromstring(z.read(path)).iter(h.local("c")):
                    f=cell.find(h.local("f"));v=cell.find(h.local("v"))
                    if f is not None:actions.append({"sheet":name,"cell":cell.get("r"),"after_formula":"="+f.text})
                    elif v is not None:inputs[(name,cell.get("r"))]=float(v.text)
        model=runtime("gg_hort_formula_eval").Evaluator(actions,inputs=inputs)
        for month,c in enumerate("CDEFGHIJKLMN",1):
            for year,(price_row,visible,revenue) in enumerate(zip([72,82,92,102,112],[24,33,42,51,60],[25,34,43,52,61]),1):
                expected=(month*100+30)*((116.61/102.5)**.25)**(year-1)
                self.assertAlmostEqual(expected,model.cell(sheet,c+str(price_row)),places=9)
                self.assertAlmostEqual(expected,model.cell(sheet,c+str(visible)),places=9)
                self.assertAlmostEqual(expected*.01,model.cell(sheet,c+str(revenue)),places=9)
                for ref in (c+str(price_row),c+str(visible),c+str(revenue)):
                    self.assertEqual(next(a["after_formula"] for a in actions if a["cell"]==ref),visible_sheet[ref].value)
        self.assertEqual(60,len([a for a in actions if a["cell"] in [c+str(r) for c in "CDEFGHIJKLMN" for r in [72,82,92,102,112]]]))
        self.assertGreater(counts["input_param"],60)
        generated.close()
