"""P3 C-lane: gg_core.checks() wiring of the S-lane fact_semantics check.

Structural connection only (design-r2 §5/§7): every fact is classified
by ``gg_fact_semantics.classify_fact`` and each returned issue surfaces
under ``check_id="fact_semantics"``.  The registry file IS deployed
(references/fact-semantics.json) and ``_semantic_context`` resolves it
relative to gg_core.py's own deployment location (scripts/ ->
knuaf-doc/references/) — never inside the user's project root
(Lane 28) — falling back to ``registry=None`` only when the deployed
file itself is missing, unreadable or corrupt.  A legacy fact with no
P3 metadata still classifies resolved and stays silent, while explicit
metadata conflicts surface real issues alongside the existing
check_ids (the new check is additive, not a replacement).  No
InputSnapshot files are deployed, so inputs/input_fingerprint stay
empty/None.
"""
import hashlib
import json
import unittest
from unittest import mock
from pathlib import Path

from tests._harness import (
    ContractCase,
    bind_major,
    claim_of,
    fact_op,
    runtime,
    section_claim,
    section_op,
    source_op,
    write_text,
)


def _seed(root):
    """One source+fact+section triple whose only fact is a plain text
    fact with no P3 metadata — the shape every pre-P3 project carries.
    unit="" makes the text fact unambiguously non-dimensional so the
    S-lane legacy path resolves it (unit=None is deliberately read as
    dimension-ambiguous -> SEMANTIC_UNRESOLVED)."""
    write_text(root, "answer.txt", "농장명: 행복농장\n")
    write_text(root, "intro.md", "농장명은 행복농장이다.\n")
    fact = fact_op("farm_name", "farm_name", "행복농장", "", claim_id="c1")
    ops = [
        source_op(claims=claim_of(fact["value"])),
        fact,
        section_op(
            "intro", "머리말", "intro.md",
            claims=[section_claim(fact["value"], "농장명은 행복농장이다")]),
    ]
    return runtime("gg_core").apply(
        root, {"request_id": "seed", "ops": ops}, 0)


def _seed_none_unit(root):
    """Same triple as _seed() but with unit=None — the shape legacy
    projects actually stored before P3.  The S-lane legacy branch only
    resolves silently when unit is a str ("" included); a missing unit
    is deliberately read as dimension-ambiguous -> SEMANTIC_UNRESOLVED."""
    write_text(root, "answer.txt", "농장명: 행복농장\n")
    write_text(root, "intro.md", "농장명은 행복농장이다.\n")
    fact = fact_op("farm_name", "farm_name", "행복농장", None, claim_id="c1")
    ops = [
        source_op(claims=claim_of(fact["value"])),
        fact,
        section_op(
            "intro", "머리말", "intro.md",
            claims=[section_claim(fact["value"], "농장명은 행복농장이다")]),
    ]
    return runtime("gg_core").apply(
        root, {"request_id": "seed", "ops": ops}, 0)


def _bad_measure_fact():
    """A provided monetary fact carrying P3 ``measure`` metadata whose
    currency is invalid per S-lane ``validate_metadata`` (KRW only)."""
    fact = fact_op(
        "sales", "매출액", "1000", "천원",
        value_type="decimal", period="2026년", scope="농장 전체",
        claim_id="c1")
    fact["value"]["measure"] = {
        "kind": "monetary_total", "currency": "USD"}
    return fact


def _seed_bad_measure(root):
    """A project whose sole fact carries the invalid measure.

    Lane 23 fixture repair: ``apply`` now write-rejects bad P3 metadata
    on NEW fact records (mode="new" admission), so the invalid measure
    can no longer be committed through apply.  It is injected into the
    stored canonical record directly, simulating exactly what it models:
    a fact that predates admission (or arrived through a path without
    it) — the read-time ``fact_semantics`` diagnosis each consumer test
    asserts on is unchanged."""
    core = runtime("gg_core")
    write_text(root, "answer.txt", "매출액: 1000천원\n")
    fact = _bad_measure_fact()
    bad_measure = fact["value"].pop("measure")
    p = core.apply(
        root,
        {"request_id": "seed",
         "ops": [source_op(claims=claim_of(fact["value"])), fact]},
        0,
    )
    p["facts"]["sales"]["measure"] = bad_measure
    (Path(root) / "project.json").write_text(
        json.dumps(p, ensure_ascii=False, indent=2), encoding="utf-8")
    return core.load(root)


class FactSemanticsWiringTests(ContractCase):
    def test_legacy_fact_produces_no_fact_semantics_entries(self):
        """A fact with no meaning_id/measure/finance_role must classify
        resolved with zero issues — checks() reports nothing for it."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed(root)
        report = core.checks(root, p)
        self.assertEqual(
            [],
            [r for r in report if r["check_id"] == "fact_semantics"],
        )

    def test_invalid_measure_surfaces_as_fact_semantics(self):
        """A fact whose measure carries an invalid currency surfaces a
        fact_semantics entry carrying the S-lane issue's status and
        reason."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_bad_measure(root)
        entries = [
            r for r in core.checks(root, p)
            if r["check_id"] == "fact_semantics" and r["target"] == "sales"
        ]
        self.assertTrue(entries)
        for r in entries:
            self.assertIn(r["status"], ("fail", "blocked"))
            self.assertTrue(r["reason"])
        # invalid currency is a hard S-lane failure, not a soft block
        self.assertTrue(any(r["status"] == "fail" for r in entries))

    def test_semantic_context_shape_and_revision(self):
        """_semantic_context returns exactly the documented Context keys
        and pins the project's actual current revision."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed(root)
        ctx = core._semantic_context(root, p)
        self.assertEqual(
            {"registry", "inputs", "evidence", "target_refs",
             "input_fingerprint", "input_revision"},
            set(ctx),
        )
        # Lane 28: the registry resolves deployment-relative, so it is
        # unconditionally the real installed bytes here.
        self.assertEqual(
            hashlib.sha256(
                SemanticContextRegistryTests._installed_bytes(core)
            ).hexdigest(),
            ctx["registry"]["sha256"])
        self.assertEqual([], ctx["inputs"])
        self.assertEqual(
            {"sources": {}, "reviews": {}, "approvals": {}},
            ctx["evidence"],
        )
        self.assertEqual([], ctx["target_refs"])
        self.assertIsNone(ctx["input_fingerprint"])
        self.assertEqual(p["revision"], ctx["input_revision"])

    def test_existing_checks_coexist_with_fact_semantics(self):
        """Additive, not a replacement: tampering with the source bytes
        still produces the pre-existing source_hash entry while the
        bad-measure fact produces its fact_semantics entry."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_bad_measure(root)
        write_text(root, "answer.txt", "tampered source bytes\n")
        ids = {
            (r["check_id"], r["target"]) for r in core.checks(root, p)
        }
        self.assertIn(("source_hash", "answer"), ids)
        self.assertIn(("fact_semantics", "sales"), ids)

    def test_unmetadata_legacy_fact_with_none_unit_blocks_submission(self):
        """unit=None is not unit="": the S-lane legacy branch only
        resolves silently when unit is a str ("" included), so a missing
        unit is deliberately read as dimension-ambiguous — the fact
        surfaces one blocked fact_semantics entry that blocks
        submission_candidate."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_none_unit(root)
        entries = [
            r for r in core.checks(root, p)
            if r["check_id"] == "fact_semantics"
            and r["target"] == "farm_name"
        ]
        self.assertEqual(1, len(entries), entries)
        entry = entries[0]
        self.assertEqual("blocked", entry["status"])
        self.assertEqual("error", entry["severity"])
        self.assertTrue(core.blocks_skill_candidate(entry))


def _specialty(root):
    """Policy B output context: explicit specialty major for a project
    bound with ``bind_major``."""
    return runtime("gg_major_contract").output_context(
        root, "specialty_crops")


def _finance_spec(**over):
    """Minimal calculable single_annual_cash_v1 spec — the same shape the
    P1 harness feeds gg_finance.workbook (test_p1_core.py)."""
    spec = dict(
        profile="single_annual_cash_v1", crops=["합성작목"],
        accounting_basis="cash_pre_tax_no_inventory",
        source_refs=["synthetic:행복농장-원답변"], unit="원", quantity_unit="kg",
        land="100", facility="300", equity="200", loan="200", loan_rate=".05",
        discount_rate=".05", salvage="0",
        investment_basis="school_farm_new_business", owner_labor_in_costs=False,
        start_year=2030, years=3, life=3, grace=0, term=3,
        repayment="equal_principal",
        periods=[dict(year=2030 + i, quantity="100", sold="90", loss="10",
                      price="5", variable_cost="1", fixed_cost="100",
                      household="50") for i in range(3)],
    )
    spec.update(over)
    return spec


class ExcelTemplatePresaveTests(ContractCase):
    """Lane 4: build_excel_finance_template.build() runs the C-lane
    economic validation before writing — hard "fail" issues refuse the
    build (no output file), soft "blocked" provenance issues never refuse
    a draft and ride along as economic_validation_issues."""

    def _write_spec(self, root, spec):
        write_text(root, "spec.json", json.dumps(spec, ensure_ascii=False))
        return "spec.json", "out.xlsx"

    def test_fail_conflict_refuses_build_no_file(self):
        """A real status="fail" conflict (a periods row where
        quantity≠sold+loss) refuses the build before workbook() runs —
        status="blocked" JSON and no output file."""
        mod = runtime("build_excel_finance_template")
        root = self.make_project()
        bind_major(root)
        spec = _finance_spec()
        spec["periods"][0]["loss"] = "9"  # 100 ≠ 90 + 9
        spec_path, out = self._write_spec(root, spec)
        result = mod.build(root, spec_path, out, major_id="specialty_crops")
        self.assertEqual("blocked", result["status"])
        self.assertTrue(result["issues"])
        self.assertTrue(
            any(i["status"] == "fail" for i in result["issues"]))
        self.assertFalse((Path(root) / out).exists())

    def test_clean_spec_builds(self):
        """A conflict-free spec builds normally; the issues key is always
        present on success (empty when nothing was found)."""
        mod = runtime("build_excel_finance_template")
        root = self.make_project()
        bind_major(root)
        spec_path, out = self._write_spec(root, _finance_spec())
        result = mod.build(root, spec_path, out, major_id="specialty_crops")
        self.assertEqual("calculated", result["status"])
        self.assertTrue((Path(root) / out).exists())
        self.assertEqual([], result["economic_validation_issues"])

    def test_blocked_only_issues_do_not_refuse_draft(self):
        """Provenance-only "blocked" issues (a declared source_ref with no
        context to resolve it against — always the case for this draft
        CLI) never refuse the draft build but surface on the result."""
        mod = runtime("build_excel_finance_template")
        root = self.make_project()
        bind_major(root)
        spec = _finance_spec()
        spec["input_inventory"] = {
            "site_area": {"status": "user_answer",
                          "source_ref": "src:unlinked", "value": "1"},
        }
        spec_path, out = self._write_spec(root, spec)
        result = mod.build(root, spec_path, out, major_id="specialty_crops")
        self.assertEqual("calculated", result["status"])
        self.assertTrue((Path(root) / out).exists())
        issues = result["economic_validation_issues"]
        self.assertTrue(issues)
        self.assertTrue(all(i["status"] == "blocked" for i in issues))

    def test_semantic_unit_outside_vocabulary_rejected(self):
        """Lane-4 gap fix: a semanticField map entry's unit is now checked
        against the economic unit vocabulary — previously '드럼통' would
        have been written into the workbook silently."""
        fill = runtime("gg_excel_fill")
        tpl = runtime("gg_excel_template")
        root = self.make_project()
        bind_major(root)
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "시트1"
        wb.active["A1"] = "기존값"
        src = Path(root) / "tpl.xlsx"
        wb.save(src)
        m = Path(root) / "map.json"
        m.write_text(json.dumps({
            "schema": "gg-xlsx-fill-map/v1",
            "template": {"sha256": tpl.sha256(src)},
            "entries": [{
                "sheet": "시트1", "cell": "A1", "role": "input",
                "period": "2030년", "source_note": "합성", "editable": True,
                "semanticField": "sales_amount", "unit": "드럼통",
            }],
        }, ensure_ascii=False), encoding="utf-8")
        v = Path(root) / "values.json"
        v.write_text(json.dumps(
            {"schema": "gg-xlsx-fill-values/v1", "values": []},
            ensure_ascii=False), encoding="utf-8")
        out = Path(root) / "filled.xlsx"
        with self.assertRaises(ValueError) as cm:
            fill.fill_copy(src, m, v, out, context=_specialty(root))
        self.assertIn("vocabulary", str(cm.exception))
        self.assertFalse(out.exists())

    def test_semantic_unit_in_vocabulary_still_fills(self):
        """Positive control: a recognized unit ('원') on a semanticField
        entry does not hit the new check — the fill completes and records
        the mapped cell as omitted (no value supplied)."""
        fill = runtime("gg_excel_fill")
        tpl = runtime("gg_excel_template")
        root = self.make_project()
        bind_major(root)
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "시트1"
        wb.active["A1"] = "기존값"
        src = Path(root) / "tpl.xlsx"
        wb.save(src)
        m = Path(root) / "map.json"
        m.write_text(json.dumps({
            "schema": "gg-xlsx-fill-map/v1",
            "template": {"sha256": tpl.sha256(src)},
            "entries": [{
                "sheet": "시트1", "cell": "A1", "role": "input",
                "period": "2030년", "source_note": "합성", "editable": True,
                "semanticField": "sales_amount", "unit": "원",
            }],
        }, ensure_ascii=False), encoding="utf-8")
        v = Path(root) / "values.json"
        v.write_text(json.dumps(
            {"schema": "gg-xlsx-fill-values/v1", "values": []},
            ensure_ascii=False), encoding="utf-8")
        out = Path(root) / "filled.xlsx"
        receipt = fill.fill_copy(src, m, v, out, context=_specialty(root))
        self.assertEqual("filled", receipt["output"]["status"])
        self.assertTrue(out.exists())


def _seed_econ(root):
    """A project whose sole fact is a provided monetary fact carrying P3
    economic metadata — the kind a calculation-kind review covers."""
    write_text(root, "answer.txt", "매출액: 1000천원\n")
    fact = fact_op("sales", "매출액", "1000", "천원",
                   value_type="decimal", period="2026년",
                   scope="농장 전체", claim_id="c1")
    fact["value"]["measure"] = {"kind": "monetary_total", "currency": "KRW"}
    fact["value"]["finance_role"] = "plan"
    return runtime("gg_core").apply(
        root,
        {"request_id": "seed",
         "ops": [source_op(claims=claim_of(fact["value"])), fact]},
        0)


def _semantics_report(sem, core, root, p, refs, *, registry=None):
    """An honest gg-finance-semantics-report/1 for ``refs`` built by
    re-interpreting the CURRENT project state — what a real economic
    review run persists at review.path.  ``registry`` defaults to the
    live context registry (the deployed file, resolved
    deployment-relative since Lane 28 — it is unconditionally present
    in this tree); pass an explicit dict to author a report declared
    against a different (stale) registry."""
    if registry is None:
        registry = core._semantic_context(root, p)["registry"]
    frs = [
        sem.classify_fact(
            p["facts"][r["id"]], registry=registry, bindings=[])
        for r in refs
        if r["collection"] == "facts" and r["id"] in p["facts"]]
    consumers = sem.resolve_consumers(p, inputs=[], registry=registry)
    report = sem.make_semantics_report(
        registry=registry, target_refs=refs,
        input_revision=p["revision"],
        input_fingerprint=core.fingerprint(root, p, refs),
        facts=frs, consumers=consumers, body_result=None)
    report["review"] = {
        "author_id": "author", "reviewer_id": "reviewer",
        "review_kinds": ["calculation"],
        "findings": [], "disposition": "resolved"}
    return report


def _drifted_context(core):
    """Context manager patching ``core._semantic_context`` so its
    registry carries a different version/sha256 — simulating a registry
    drift AFTER records were committed.  Lane 28: the deployed registry
    resolves unconditionally, so drift can no longer be staged by
    installing bytes late into a project root; patching the read is the
    equivalent mechanism (and leaves the real file untouched)."""
    real_ctx = core._semantic_context

    def ctx(root, p):
        out = real_ctx(root, p)
        if out.get("registry") is not None:
            out = dict(out)
            out["registry"] = {
                **out["registry"],
                "version": "drifted", "sha256": "0" * 64}
        return out

    return mock.patch.object(core, "_semantic_context", ctx)


def _bind_economic_review(core, root):
    """Full honest fixture: economic fact + persisted semantics report +
    schema-2 observation + applied bound calculation review.  Returns
    the post-apply project dict (canonical revision R+1; the review keeps
    the observed input_revision R)."""
    sem = runtime("gg_fact_semantics")
    p = _seed_econ(root)
    refs = core.sort_target_refs(core._export_refs(p))
    report = _semantics_report(sem, core, root, p, refs)
    write_text(root, "report.json", json.dumps(report, ensure_ascii=False))
    obs = {
        "author_session": "s-author", "reviewer_session": "s-reviewer",
        "author_id": "author", "reviewer_id": "reviewer",
        "review_kinds": ["calculation"],
        "target_refs": refs,
        "input_fingerprint": core.fingerprint(root, p, refs),
        "report_path": "report.json",
        "report_hash": core.digest(
            (Path(root) / "report.json").read_bytes()),
        "input_revision": p["revision"],
    }
    receipt = core.ingest_review_observation(root, obs, "observer-1")
    value = {
        "id": "rev-calc", "review_kind": "calculation",
        "author_id": "author", "reviewer_id": "reviewer",
        "target_refs": refs,
        "input_fingerprint": core.fingerprint(root, p, refs),
        "findings": [], "disposition": "resolved", "status": "pass",
        "path": "report.json", "coverage": {"econ": ["매출액 검토"]},
        "provenance_path": receipt["path"],
        "provenance_hash": receipt["hash"],
        "input_revision": p["revision"],
    }
    return core.apply(
        root,
        {"request_id": "bind", "ops": [{"collection": "reviews",
                                        "value": value}]},
        p["revision"])


class FinanceReviewFreshnessTests(ContractCase):
    """Lane 5/18: _finance_review_state via the additive
    finance_review_freshness read path — §7 freshness states for
    content/logic/calculation-kind reviews over economic-scoped target
    facts."""

    def test_content_kind_over_economic_fact_unresolved(self):
        """Lane 18 headline change: content/logic/calculation all route
        through the shared freshness function — a content-kind review
        over an economic-scoped fact with no bound report is now
        ("unresolved", [SEMANTIC_REPORT_REQUIRED]), not not_applicable
        (pre-Lane-18 the guard returned not_applicable for any
        non-calculation kind)."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        p["reviews"]["rev-content"] = {
            "id": "rev-content", "review_kind": "content",
            "author_id": "author", "reviewer_id": "reviewer",
            "target_refs": [{"collection": "facts", "id": "sales"}],
            "findings": [], "disposition": "resolved", "status": "pass"}
        state, issues = core.finance_review_freshness(
            root, p, "rev-content")
        self.assertEqual("unresolved", state)
        self.assertEqual(
            {"SEMANTIC_REPORT_REQUIRED"}, {i["code"] for i in issues})

    def test_non_economic_scope_stays_not_applicable(self):
        """Kind widening did not widen SCOPE: content- and
        calculation-kind reviews over a NON-economic fact (plain
        text, no finance_role/meaning_id/measure) still return
        ("not_applicable", [])."""
        core = runtime("gg_core")
        for kind in ("content", "calculation"):
            with self.subTest(kind=kind):
                root = self.make_project()
                p = _seed(root)  # text-only fact — no economic marker
                p["reviews"]["rev-x"] = {
                    "id": "rev-x", "review_kind": kind,
                    "author_id": "author", "reviewer_id": "reviewer",
                    "target_refs": [
                        {"collection": "facts", "id": "farm_name"}],
                    "findings": [], "disposition": "resolved",
                    "status": "pass"}
                state, issues = core.finance_review_freshness(
                    root, p, "rev-x")
                self.assertEqual("not_applicable", state)
                self.assertEqual([], issues)

    def test_docx_xlsx_render_kinds_stay_not_applicable(self):
        """The widening stops at exactly content/logic/calculation —
        artifact kinds over an economic-scoped fact still short-circuit
        to ("not_applicable", [])."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        for kind in ("docx", "xlsx", "render"):
            with self.subTest(kind=kind):
                p["reviews"]["rev-" + kind] = {
                    "id": "rev-" + kind, "review_kind": kind,
                    "author_id": "author", "reviewer_id": "reviewer",
                    "target_refs": [
                        {"collection": "facts", "id": "sales"}],
                    "findings": [], "disposition": "resolved",
                    "status": "pass"}
                state, issues = core.finance_review_freshness(
                    root, p, "rev-" + kind)
                self.assertEqual("not_applicable", state)
                self.assertEqual([], issues)

    def test_logic_kind_fresh_with_bound_report(self):
        """A logic-kind review over an economic fact, bound to a
        well-formed byte-matched report declaring review_kinds
        ["logic"], is ("fresh", []) — the same freshness AND-list a
        calculation review gets."""
        core = runtime("gg_core")
        root = self.make_project()
        p2, _ = _bind_kind_review(core, root, "logic", ["logic"])
        state, issues = core.finance_review_freshness(
            root, p2, "rev-logic")
        self.assertEqual("fresh", state)
        self.assertEqual([], issues)

    def test_logic_kind_stale_on_review_kinds_mismatch(self):
        """The report↔review kind-membership conjunct is enforced: the
        same binding shape but with the report declaring
        review_kinds ["calculation"] while the record is logic-kind ->
        ("stale", [REGISTRY_STALE 'review_kinds 불일치'])."""
        core = runtime("gg_core")
        root = self.make_project()
        p2, _ = _bind_kind_review(core, root, "logic", ["calculation"])
        state, issues = core.finance_review_freshness(
            root, p2, "rev-logic")
        self.assertEqual("stale", state)
        self.assertTrue(any(
            i["code"] == "REGISTRY_STALE" and "review_kinds" in i["reason"]
            for i in issues))

    def test_calculation_review_without_report_unresolved(self):
        """A calculation-kind review over an economic fact but with no
        bound report (no path/report_hash on the record) -> unresolved
        with a SEMANTIC_REPORT_REQUIRED issue."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        p["reviews"]["rev-calc"] = {
            "id": "rev-calc", "review_kind": "calculation",
            "author_id": "author", "reviewer_id": "reviewer",
            "target_refs": [{"collection": "facts", "id": "sales"}],
            "findings": [], "disposition": "resolved", "status": "pass"}
        state, issues = core.finance_review_freshness(
            root, p, "rev-calc")
        self.assertEqual("unresolved", state)
        self.assertEqual(
            {"SEMANTIC_REPORT_REQUIRED"},
            {i["code"] for i in issues})

    def test_byte_mismatched_report_invalid(self):
        """A calculation review whose bound report file's bytes no longer
        match report_hash -> invalid (the P2 v2 observation check also
        sees the same tamper)."""
        core = runtime("gg_core")
        root = self.make_project()
        p2 = _bind_economic_review(core, root)
        write_text(root, "report.json", '{"tampered": true}\n')
        state, issues = core.finance_review_freshness(
            root, p2, "rev-calc")
        self.assertEqual("invalid", state)
        self.assertTrue(issues)

    def test_revision_mismatch_is_stale(self):
        """Structurally valid, byte-matched report whose
        semantics.input_revision differs from the review's ->
        stale, citing the revision mismatch."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        report["semantics"]["input_revision"] = p["revision"] + 99
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        p["reviews"]["rev-calc"] = {
            "id": "rev-calc", "review_kind": "calculation",
            "author_id": "author", "reviewer_id": "reviewer",
            "target_refs": refs,
            "input_fingerprint": core.fingerprint(root, p, refs),
            "findings": [], "disposition": "resolved", "status": "pass",
            "path": "report.json",
            "report_hash": core.digest(
                (Path(root) / "report.json").read_bytes()),
            "coverage": {"econ": ["매출액 검토"]},
            "input_revision": p["revision"],
        }
        state, issues = core.finance_review_freshness(
            root, p, "rev-calc")
        self.assertEqual("stale", state)
        self.assertTrue(any(
            "input_revision" in (i.get("reason") or "") for i in issues))

    def test_consistent_review_is_fresh(self):
        """A fully consistent calculation review + report + matching v2
        observation -> ("fresh", [])."""
        core = runtime("gg_core")
        root = self.make_project()
        p2 = _bind_economic_review(core, root)
        state, issues = core.finance_review_freshness(
            root, p2, "rev-calc")
        self.assertEqual("fresh", state)
        self.assertEqual([], issues)

    def test_malformed_economic_metadata_still_economic_scoped(self):
        """Lane-5b regression: finance_role="" is present-but-malformed —
        classify_fact flags METADATA_INVALID and returns finance_role
        verbatim (""), so a truthiness check would wave the review
        through as not_applicable while the fact's own issues never
        surface.  With is-not-None semantics the review stays
        economic-scoped; no bound report -> unresolved."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        del p["facts"]["sales"]["measure"]
        p["facts"]["sales"]["finance_role"] = ""
        p["reviews"]["rev-calc"] = {
            "id": "rev-calc", "review_kind": "calculation",
            "author_id": "author", "reviewer_id": "reviewer",
            "target_refs": [{"collection": "facts", "id": "sales"}],
            "findings": [], "disposition": "resolved", "status": "pass"}
        state, issues = core.finance_review_freshness(
            root, p, "rev-calc")
        self.assertNotEqual("not_applicable", state)
        self.assertEqual("unresolved", state)
        self.assertEqual(
            {"SEMANTIC_REPORT_REQUIRED"},
            {i["code"] for i in issues})


class ChecksFinanceReviewWiringTests(ContractCase):
    """Lane 6: checks() surfaces each calculation-kind economic review's
    freshness under check_id="finance_review" — additive alongside the
    Lane-2 fact_semantics wiring, and silent when there is nothing to
    report."""

    @staticmethod
    def _unbound_calc_review():
        """A calculation-kind review over the 'sales' economic fact with
        no bound finance-semantics report (no path/report_hash)."""
        return {
            "id": "rev-calc", "review_kind": "calculation",
            "author_id": "author", "reviewer_id": "reviewer",
            "target_refs": [{"collection": "facts", "id": "sales"}],
            "findings": [], "disposition": "resolved", "status": "pass"}

    def test_unresolved_review_surfaces_blocked_entry(self):
        """Economic calculation review with no bound report -> checks()
        carries at least one status="blocked" finance_review entry (the
        unresolved / SEMANTIC_REPORT_REQUIRED case rides through with
        its own per-issue status, not a hardcoded fail)."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        p["reviews"]["rev-calc"] = self._unbound_calc_review()
        entries = [
            r for r in core.checks(root, p)
            if r["check_id"] == "finance_review"]
        self.assertTrue(entries)
        self.assertTrue(all(e["target"] == "rev-calc" for e in entries))
        self.assertTrue(any(e["status"] == "blocked" for e in entries))

    def test_fresh_review_adds_no_entries(self):
        """Fully consistent bound review + report + v2 observation ->
        checks() carries zero finance_review entries (additive contract:
        nothing added when nothing to report)."""
        core = runtime("gg_core")
        root = self.make_project()
        p2 = _bind_economic_review(core, root)
        self.assertEqual(
            [],
            [r for r in core.checks(root, p2)
             if r["check_id"] == "finance_review"])

    def test_unbound_content_kind_adds_finance_review_entry(self):
        """Lane 18: a content-kind review over an economic fact now
        routes through _finance_review_state — with no bound report it
        is unresolved -> checks() carries ONE blocked finance_review
        entry (pre-Lane-18 it was not_applicable and added zero)."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        p["reviews"]["rev-content"] = {
            "id": "rev-content", "review_kind": "content",
            "author_id": "author", "reviewer_id": "reviewer",
            "target_refs": [{"collection": "facts", "id": "sales"}]}
        entries = [r for r in core.checks(root, p)
                   if r["check_id"] == "finance_review"]
        self.assertEqual(1, len(entries))
        self.assertEqual("rev-content", entries[0]["target"])
        self.assertEqual("blocked", entries[0]["status"])
        self.assertIn("보고서", entries[0]["reason"])

    def test_fact_semantics_and_finance_review_coexist(self):
        """One checks() call carries BOTH check_ids together: the
        bad-currency fact emits its fact_semantics entry while the
        unbound calculation review emits its finance_review entry —
        true additivity, neither crowds out the other."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_bad_measure(root)
        p["reviews"]["rev-calc"] = self._unbound_calc_review()
        ids = {(r["check_id"], r["target"]) for r in core.checks(root, p)}
        self.assertIn(("fact_semantics", "sales"), ids)
        self.assertIn(("finance_review", "rev-calc"), ids)


def _calc_review_value(core, root, p, *, rid="rev-calc", path="report.json"):
    """A calculation-kind review value over every current ref (the
    seeded project contains exactly the economic 'sales' fact), with
    all P2 required fields satisfied."""
    refs = core.sort_target_refs(core._export_refs(p))
    return {
        "id": rid, "review_kind": "calculation",
        "author_id": "author", "reviewer_id": "reviewer",
        "target_refs": refs,
        "input_fingerprint": core.fingerprint(root, p, refs),
        "findings": [], "disposition": "resolved", "status": "pass",
        "path": path, "coverage": {"econ": ["매출액 검토"]}}


class ApplyEconomicReviewBindingTests(ContractCase):
    """Lane 7: _apply_locked's admission for a NEW calculation-kind
    review over an economic-scoped fact — a readable, hash-matched
    gg-finance-semantics-report/1 report must be bound at review.path
    before the canonical commit."""

    def test_new_economic_review_without_report_rejected(self):
        """'No bound report' rejected.  path absent -> the pre-existing
        required-fields gate rejects (invalid_change); path bound to a
        plain non-report file reaches the NEW schema check -> rejected
        (invalid_change, '의미 보고서 스키마 불일치')."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        value = _calc_review_value(core, root, p)
        del value["path"]
        with self.assertRaises(core.OperationError) as cm:
            core.apply(
                root,
                {"request_id": "r-missing",
                 "ops": [{"collection": "reviews", "value": value}]},
                p["revision"])
        self.assertEqual("invalid_change", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        write_text(root, "report.md", "검토 메모 — 의미 보고서 아님\n")
        value = _calc_review_value(core, root, p, path="report.md")
        with self.assertRaises(core.OperationError) as cm:
            core.apply(
                root,
                {"request_id": "r-notreport",
                 "ops": [{"collection": "reviews", "value": value}]},
                p["revision"])
        self.assertEqual("invalid_change", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertIn("의미 보고서 스키마 불일치", str(cm.exception))
        self.assertNotIn(
            "rev-calc", core.load(root).get("reviews", {}))

    def test_report_hash_mismatch_rejected(self):
        """Byte-mismatched declared hash -> ValueError.  Public apply
        derives review.report_hash from the live file bytes itself
        (post-check invariant), so the helper is exercised directly —
        the mismatch branch is the same code path apply runs."""
        core = runtime("gg_core")
        root = self.make_project()
        _seed_econ(root)
        write_text(root, "report.json", json.dumps(
            {"schema": "gg-finance-semantics-report/1",
             "semantics": {}}, ensure_ascii=False))
        with self.assertRaises(ValueError) as cm:
            core._validate_new_economic_review_binding(
                root,
                {"path": "report.json", "report_hash": "0" * 64},
                p=core.load(root), target_refs=[])
        self.assertIn("해시 불일치", str(cm.exception))

    def test_valid_report_apply_succeeds(self):
        """Positive control: bound gg-finance-semantics-report/1 whose
        bytes hash-match -> the new calculation review commits."""
        core = runtime("gg_core")
        root = self.make_project()
        p2 = _bind_economic_review(core, root)
        review = p2["reviews"]["rev-calc"]
        self.assertEqual("calculation", review["review_kind"])
        self.assertEqual("report.json", review["path"])

    def test_non_calculation_review_needs_no_report(self):
        """A NEW content-kind review over a NON-economic fact is not in
        scope for the admission check — applies with a plain non-report
        path.  (Lane 27 fixture repair: the original seeded the
        ECONOMIC fact and asserted the now-removed non-calculation
        bypass; the surviving contract is scope-based, not kind-based,
        so the seed is now the genuinely non-economic _seed_plain.)"""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_plain(root)
        write_text(root, "report.md", "content review notes\n")
        value = _calc_review_value(core, root, p, path="report.md")
        value["review_kind"] = "content"
        p2 = core.apply(
            root,
            {"request_id": "r-content",
             "ops": [{"collection": "reviews", "value": value}]},
            p["revision"])
        self.assertIn("rev-calc", p2["reviews"])

    def test_existing_review_update_not_readmitted(self):
        """prev is not None -> the new check is skipped.  Updating an
        already-registered calculation review keeps working even when
        the update's path does not point at a semantics report (the
        registration-time binding was already proven)."""
        core = runtime("gg_core")
        root = self.make_project()
        p2 = _bind_economic_review(core, root)
        write_text(root, "revised.md", "재검토 메모\n")
        value = _calc_review_value(core, root, p2, path="revised.md")
        value["findings"] = [{"severity": "info",
                              "status": "resolved", "note": "재검"}]
        p3 = core.apply(
            root,
            {"request_id": "r-update",
             "ops": [{"collection": "reviews", "value": value}]},
            p2["revision"])
        self.assertEqual("revised.md", p3["reviews"]["rev-calc"]["path"])


def _seed_plain(root):
    """A project whose sole fact carries NO P3 economic metadata — the
    non-economic counterpart of _seed_econ."""
    write_text(root, "answer.txt", "매출액: 1000천원\n")
    fact = fact_op("sales", "매출액", "1000", "천원",
                   value_type="decimal", period="2026년",
                   scope="농장 전체", claim_id="c1")
    return runtime("gg_core").apply(
        root,
        {"request_id": "seed",
         "ops": [source_op(claims=claim_of(fact["value"])), fact]},
        0)


def _observation(core, root, p, *, kinds=("calculation",),
                 report_path="report.json"):
    """A schema-2 observation dict bound to ``report_path`` over every
    current ref — the same shape _bind_economic_review ingests."""
    refs = core.sort_target_refs(core._export_refs(p))
    return {
        "author_session": "s-author", "reviewer_session": "s-reviewer",
        "author_id": "author", "reviewer_id": "reviewer",
        "review_kinds": list(kinds),
        "target_refs": refs,
        "input_fingerprint": core.fingerprint(root, p, refs),
        "report_path": report_path,
        "report_hash": core.digest(
            (Path(root) / report_path).read_bytes()),
        "input_revision": p["revision"],
    }


def _bind_kind_review(core, root, kind, report_kinds):
    """``_bind_economic_review`` generalized to any review_kind (Lane
    18): the report's review block declares ``report_kinds`` while the
    observation + review record use ``kind``.  Freshness holds only
    when the report's declared review_kinds set-equals ``[kind]`` —
    ``validate_report_binding`` compares them sorted."""
    sem = runtime("gg_fact_semantics")
    p = _seed_econ(root)
    refs = core.sort_target_refs(core._export_refs(p))
    report = _semantics_report(sem, core, root, p, refs)
    report["review"]["review_kinds"] = list(report_kinds)
    write_text(root, "report.json", json.dumps(report, ensure_ascii=False))
    obs = _observation(
        core, root, p, kinds=(kind,), report_path="report.json")
    receipt = core.ingest_review_observation(root, obs, "observer-1")
    rid = "rev-" + kind
    value = {
        "id": rid, "review_kind": kind,
        "author_id": "author", "reviewer_id": "reviewer",
        "target_refs": refs,
        "input_fingerprint": core.fingerprint(root, p, refs),
        "findings": [], "disposition": "resolved", "status": "pass",
        "path": "report.json", "coverage": {"econ": ["매출액 검토"]},
        "provenance_path": receipt["path"],
        "provenance_hash": receipt["hash"],
        "input_revision": p["revision"],
    }
    p2 = core.apply(
        root,
        {"request_id": "bind-" + kind,
         "ops": [{"collection": "reviews", "value": value}]},
        p["revision"])
    return p2, p2["reviews"][rid]


class IngestEconomicObservationTests(ContractCase):
    """Lane 8: ingest_review_observation admission — an observation
    declaring a calculation-kind review over an economic-scoped fact
    must bind a gg-finance-semantics-report/1 report; all other
    observations pass through on byte-identity alone (unchanged)."""

    def test_calculation_economic_bad_report_rejected(self):
        """Byte-matched but non-report JSON (missing semantics key /
        wrong schema) -> economic_report_schema_invalid, not_committed."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        write_text(root, "report.md", "검토 메모 — 의미 보고서 아님\n")
        obs = _observation(core, root, p, report_path="report.md")
        with self.assertRaises(core.OperationError) as cm:
            core.ingest_review_observation(root, obs, "observer-1")
        self.assertEqual(
            "economic_report_schema_invalid",
            cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])

    def test_calculation_economic_valid_report_accepted(self):
        """Positive control: byte-matched gg-finance-semantics-report/1
        -> the observation records and returns a receipt."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        receipt = core.ingest_review_observation(
            root, _observation(core, root, p), "observer-1")
        self.assertTrue(receipt["path"])
        self.assertTrue(receipt["hash"])

    def test_non_calculation_observation_not_gated(self):
        """review_kinds without 'calculation' over NON-economic target
        refs -> the economic report check does not run; a garbage bound
        report still records (byte-identity was already proven
        upstream).  (Lane 27 fixture repair: the original seeded the
        ECONOMIC fact and asserted the now-removed non-calculation
        bypass; the surviving exemption is scope-based — content/logic
        kinds over economic refs are now gated — so the seed is now the
        genuinely non-economic _seed_plain.)"""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_plain(root)
        write_text(root, "report.md", "content review notes\n")
        receipt = core.ingest_review_observation(
            root,
            _observation(core, root, p,
                         kinds=("content",), report_path="report.md"),
            "observer-1")
        self.assertTrue(receipt["path"])

    def test_calculation_noneconomic_targets_not_gated(self):
        """calculation kind but zero economic-scoped target facts ->
        not gated; a non-report-shaped bound file still records."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_plain(root)
        write_text(root, "report.md", "calc notes — not a report\n")
        receipt = core.ingest_review_observation(
            root,
            _observation(core, root, p, report_path="report.md"),
            "observer-1")
        self.assertTrue(receipt["path"])


def _xlsx_bytes(label):
    """A real (synthetic) workbook's bytes: policy B judges an adopted
    output by its content, so an ``.xlsx`` fixture must be a workbook."""
    import io
    from openpyxl import Workbook
    wb = Workbook()
    wb.active["A1"] = label
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


BYTES = _xlsx_bytes("BYTES")
BYTES2 = _xlsx_bytes("BYTES2")


def _write_xlsx(root, name, data):
    (Path(root) / name).write_bytes(data)


def _bind_specialty(root):
    """Explicit specialty binding (policy B) on a seeded project; returns
    the reloaded canonical record."""
    bind_major(root)
    return runtime("gg_core").load(root)


def _output_value_for(core, root, p, *, path="out.xlsx", data=BYTES,
                      checks=None):
    """An adopt_output output_value over every current ref, mirroring
    test_p2_core's _output_value shape; ``checks`` is attached only when
    given (it is not a required output_value key — only server keys are
    rejected by _check_adopt_output_value)."""
    refs = core.sort_target_refs(core._export_refs(p))
    ov = {
        "id": Path(path).stem,
        "path": path,
        "format": "xlsx",
        "file_hash": core.digest(data),
        "target_refs": refs,
        "input_fingerprint": core.fingerprint(root, p, refs),
    }
    if checks is not None:
        ov["checks"] = checks
    return ov


class AdoptEconomicOutputTests(ContractCase):
    """Lane 9/19: adopt_output admission — a NEW output over an
    economic-scoped fact must carry ≥1 checks entry resolving to a
    byte-matched, schema-valid gg-finance-semantics-report/1 report
    whose semantics.output_file_hash equals the adopted output's own
    digest (Lane 19 — a report bound to a different or no output file
    never qualifies); non-economic outputs keep checks as inert caller
    metadata."""

    def test_economic_output_without_report_checks_rejected(self):
        """No checks key, and a checks entry bound to a non-report file,
        both reject with economic_evidence_missing / not_committed."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        _write_xlsx(root, "out.xlsx", BYTES)
        ov = _output_value_for(core, root, p)
        with self.assertRaises(core.OperationError) as cm:
            core.adopt_output(root, ov, p["revision"], "adopt-none", major_id="specialty_crops")
        self.assertEqual(
            "economic_evidence_missing", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        write_text(root, "report.md", "검토 메모 — 의미 보고서 아님\n")
        ov = _output_value_for(
            core, root, p,
            checks=[{"evidence_path": "report.md",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.md").read_bytes())}])
        with self.assertRaises(core.OperationError) as cm:
            core.adopt_output(root, ov, p["revision"], "adopt-bad", major_id="specialty_crops")
        self.assertEqual(
            "economic_evidence_missing", cm.exception.result["reason"])
        self.assertNotIn("out", core.load(root).get("outputs", {}))

    def test_economic_output_with_valid_report_check_adopts(self):
        """Positive control: checks entry binding a byte-matched
        gg-finance-semantics-report/1 whose semantics.output_file_hash
        equals the adopted output's own digest -> the output is adopted.
        (Lane 19 fixture repair: make_semantics_report emits
        output_file_hash=None, which no longer qualifies — the report
        must be bound to THIS output's bytes.)"""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        _write_xlsx(root, "out.xlsx", BYTES)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        report["semantics"]["output_file_hash"] = core.digest(BYTES)
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        ov = _output_value_for(
            core, root, p,
            checks=[{"evidence_path": "report.json",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.json").read_bytes())}])
        result = core.adopt_output(root, ov, p["revision"], "adopt-ok", major_id="specialty_crops")
        self.assertEqual("adopted", result["status"])
        self.assertIn("out", core.load(root)["outputs"])

    def test_report_bound_to_other_output_rejected(self):
        """Lane 19 headline fix: a byte-matched, schema-valid report
        authored for a DIFFERENT output (output_file_hash set but not
        matching this output's bytes) no longer qualifies -> rejected
        with economic_evidence_missing, nothing committed."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        _write_xlsx(root, "out.xlsx", BYTES)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        report["semantics"]["output_file_hash"] = core.digest(
            b"OTHER-OUTPUT-BYTES")
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        ov = _output_value_for(
            core, root, p,
            checks=[{"evidence_path": "report.json",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.json").read_bytes())}])
        with self.assertRaises(core.OperationError) as cm:
            core.adopt_output(root, ov, p["revision"], "adopt-foreign", major_id="specialty_crops")
        self.assertEqual(
            "economic_evidence_missing", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertNotIn("out", core.load(root).get("outputs", {}))

    def test_null_output_file_hash_rejected(self):
        """A byte-matched, schema-valid report whose
        semantics.output_file_hash is null proves nothing about THIS
        output's bytes -> does not qualify -> economic_evidence_missing."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        _write_xlsx(root, "out.xlsx", BYTES)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        self.assertIsNone(report["semantics"]["output_file_hash"])
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        ov = _output_value_for(
            core, root, p,
            checks=[{"evidence_path": "report.json",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.json").read_bytes())}])
        with self.assertRaises(core.OperationError) as cm:
            core.adopt_output(root, ov, p["revision"], "adopt-null", major_id="specialty_crops")
        self.assertEqual(
            "economic_evidence_missing", cm.exception.result["reason"])
        self.assertNotIn("out", core.load(root).get("outputs", {}))

    def test_missing_source_still_output_source_missing(self):
        """L19-001 ordering regression: an economic output whose source
        file is absent must surface output_source_missing — the
        canonical source-read block runs BEFORE the economic-evidence
        loop again.  The checks entry is otherwise fully qualifying
        (byte-matched, schema-valid, output_file_hash bound to the
        intended bytes), so under the broken ordering this raised
        economic_evidence_missing instead."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        report["semantics"]["output_file_hash"] = core.digest(BYTES)
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        # meta["path"]="missing.xlsx" is deliberately never written.
        ov = _output_value_for(
            core, root, p, path="missing.xlsx",
            checks=[{"evidence_path": "report.json",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.json").read_bytes())}])
        with self.assertRaises(core.OperationError) as cm:
            core.adopt_output(root, ov, p["revision"], "adopt-missing", major_id="specialty_crops")
        self.assertEqual(
            "output_source_missing", cm.exception.result["reason"])
        self.assertNotIn(
            "missing", core.load(root).get("outputs", {}))

    def test_managed_source_path_still_invalid_output_metadata(self):
        """L19-001 ordering regression: an economic output whose path is
        .gg-artifacts/-managed must surface invalid_output_metadata
        before the economic-evidence loop — never
        economic_evidence_missing."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        ov = _output_value_for(
            core, root, p, path=".gg-artifacts/managed/out.xlsx",
            checks=[{"evidence_path": "report.json",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.json").read_bytes())}])
        with self.assertRaises(core.OperationError) as cm:
            core.adopt_output(root, ov, p["revision"], "adopt-managed", major_id="specialty_crops")
        self.assertEqual(
            "invalid_output_metadata", cm.exception.result["reason"])

    def test_noneconomic_output_ignores_checks(self):
        """No economic-scoped target fact -> the gate never runs; a
        garbage checks list flows through as inert metadata."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_plain(root)
        p = _bind_specialty(root)
        _write_xlsx(root, "out.xlsx", BYTES)
        ov = _output_value_for(
            core, root, p, checks=["garbage", {"unrelated": 1}])
        result = core.adopt_output(root, ov, p["revision"], "adopt-ne", major_id="specialty_crops")
        self.assertEqual("adopted", result["status"])


class ExportGateOrderingTests(ContractCase):
    """Lane 10 Part A: export_locked runs gate() fresh (line ~3460)
    BEFORE its receipt/existing-cache lookup (~line 3509), so a
    submission_candidate can never ride a cached bundle past a
    newly-invalid finance_review state; draft/review keep the blocked
    content per design ('blocked 내용을 보존해 출력 가능')."""

    def test_submission_blocked_on_invalid_finance_review(self):
        """Bound economic calculation review whose report bytes are then
        corrupted -> _finance_review_state = invalid -> finance_review
        entry (reason '보고서 파일 해시 불일치') -> export blocked BEFORE
        any cache path, with the finance-specific reason in the detail.
        A draft export on the SAME project still generates — kind-scoped
        gating, blocked content preserved."""
        core = runtime("gg_core")
        root = self.make_project()
        _bind_economic_review(core, root)
        bind_major(root)
        write_text(root, "report.json", '{"schema":"tampered"}')
        with self.assertRaises(core.OperationError) as cm:
            core.export(root, "submission_candidate", major_id="specialty_crops")
        self.assertEqual(
            "submission_candidate_blocked", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        # The finance_review check's own reason reaches the block detail
        # — proof gate() ran fresh on this call (the generic
        # review_calculation entry carries a different reason string).
        self.assertIn("보고서 파일 해시 불일치", str(cm.exception))
        self.assertEqual(
            "generated", core.export(root, "draft", major_id="specialty_crops")["status"])


class PaperNativeFieldValidationTests(ContractCase):
    """Lane 13 — DESIGN.md §7 pre-save native numeric adapter for paper:
    the docstring-declared str|int economic fields are rendered under
    fixed-unit labels (연간매출액/순이익 천원, 생산규모 ㎡), so malformed
    values must be refused before any publication — not an opaque-prose
    boundary (critic L12-A-001 correction of the Lane-12 reading)."""

    def _spec(self, root, **extra):
        data = {"author": "합성", "writing_year": 2026,
                "school_profile": {"mode": "school", "school": "합성대학교",
                                   "department": "특용작물학과"}}
        data.update(extra)
        write_text(root, "paper-spec.json",
                   json.dumps(data, ensure_ascii=False))
        return Path(root) / "paper-spec.json"

    def test_malformed_native_field_refused_before_save(self):
        core = runtime("gg_core")
        root = self.make_project()
        _seed_econ(root)
        bind_major(root)
        spec = self._spec(root, start_sales_thousand="99999달러")
        with self.assertRaises(core.OperationError) as cm:
            core.paper(root, "paper-spec.json", spec.read_bytes(),
                       "build/paper-out.md", major_id="specialty_crops")
        result = cm.exception.result
        self.assertEqual("paper", result["operation"])
        self.assertEqual("native_field_invalid", result["reason"])
        self.assertEqual("not_committed", result["commit_state"])
        self.assertIn("start_sales_thousand", str(cm.exception))
        # refusal precedes any save: no requested file, no managed bundle
        self.assertFalse((Path(root) / "build/paper-out.md").exists())
        rev = core.load(root)["revision"]
        self.assertFalse(
            (Path(root) / "build" / str(rev) / "paper").exists())

    def test_malformed_comma_grouping_refused(self):
        """Lane 13b — commas are never simply deleted: only proper
        thousands grouping (1-3 digits, then 3-digit groups) passes;
        "1,,2", "12,34", "1,23,456" are all refused pre-save."""
        for bad in ("1,,2", "12,34", "1,23,456"):
            with self.subTest(value=bad):
                core = runtime("gg_core")
                root = self.make_project()
                _seed_econ(root)
                bind_major(root)
                spec = self._spec(root, start_sales_thousand=bad)
                with self.assertRaises(core.OperationError) as cm:
                    core.paper(root, "paper-spec.json",
                               spec.read_bytes(), "build/paper-out.md", major_id="specialty_crops")
                result = cm.exception.result
                self.assertEqual("paper", result["operation"])
                self.assertEqual("native_field_invalid",
                                 result["reason"])
                self.assertEqual("not_committed",
                                 result["commit_state"])
                self.assertFalse(
                    (Path(root) / "build/paper-out.md").exists())
                rev = core.load(root)["revision"]
                self.assertFalse(
                    (Path(root) / "build" / str(rev) / "paper").exists())

    def test_second_field_also_checked(self):
        """The check is field-generic, not hardcoded to one key."""
        core = runtime("gg_core")
        root = self.make_project()
        _seed_econ(root)
        bind_major(root)
        spec = self._spec(root, start_profit_thousand="약 5000")
        with self.assertRaises(core.OperationError) as cm:
            core.paper(root, "paper-spec.json", spec.read_bytes(),
                       "build/paper-out.md", major_id="specialty_crops")
        self.assertEqual("native_field_invalid",
                         cm.exception.result["reason"])
        self.assertIn("start_profit_thousand", str(cm.exception))
        self.assertFalse((Path(root) / "build/paper-out.md").exists())

    def test_well_formed_native_values_generate(self):
        """int, digit string, and comma-thousands string all pass."""
        core = runtime("gg_core")
        root = self.make_project()
        _seed_econ(root)
        bind_major(root)
        spec = self._spec(
            root,
            start_sales_thousand=12000,
            target_sales_thousand="15000",
            start_profit_thousand="3,000")
        value = core.paper(
            root, "paper-spec.json", spec.read_bytes(),
            "build/paper-out.md", major_id="specialty_crops")
        self.assertEqual("generated", value["status"])
        out = (Path(root) / "build/paper-out.md").read_text(
            encoding="utf-8")
        self.assertIn("12000", out)
        self.assertIn("3,000", out)

    def test_absent_optional_fields_never_raise(self):
        """No target_* fields at all — zero spurious issues."""
        core = runtime("gg_core")
        root = self.make_project()
        _seed_econ(root)
        bind_major(root)
        spec = self._spec(root, start_sales_thousand=12000)
        value = core.paper(
            root, "paper-spec.json", spec.read_bytes(),
            "build/paper-out.md", major_id="specialty_crops")
        self.assertEqual("generated", value["status"])


class ReviewObservationEconomicCompositionTests(ContractCase):
    """Lane 14 — ACCEPTANCE.md test_all_public_paths contract: calling
    review_observation_state/valid DIRECTLY with p= must surface
    economic staleness ("state stale; valid false").  P2 diagnostics
    stay unmasked, and legacy 2-arg call sites keep byte-identical
    behavior (deliberate backward-compat boundary — the design row's
    composition is opt-in via p=, which gate() always supplies)."""

    @staticmethod
    def _stale_economic_fixture(core, root):
        """_bind_economic_review (P2-valid provenance + bound report),
        then an in-memory economic-fact mutation so the input
        fingerprint drifts -> _finance_review_state reports 'stale'
        while the P2 observation itself remains 'valid'."""
        p2 = _bind_economic_review(core, root)
        p2["facts"]["sales"]["finance_role"] = "actual"
        return p2, p2["reviews"]["rev-calc"]

    def test_stale_economic_review_composes(self):
        """The decisive acceptance case: direct call with p= returns
        ("stale", <reason>); review_observation_valid returns False."""
        core = runtime("gg_core")
        root = self.make_project()
        p2, review = self._stale_economic_fixture(core, root)
        state, reason = core.review_observation_state(
            root, review, p=p2)
        self.assertEqual("stale", state)
        self.assertTrue(reason)
        self.assertFalse(
            core.review_observation_valid(root, review, p=p2))

    def test_p2_invalid_never_masked(self):
        """Broken P2 provenance still returns the ORIGINAL P2 'invalid'
        + its own reason — economic composition never overrides P2
        diagnostics."""
        core = runtime("gg_core")
        root = self.make_project()
        p2 = _bind_economic_review(core, root)
        review = p2["reviews"]["rev-calc"]
        write_text(root, review["provenance_path"], "tampered\n")
        state, reason = core.review_observation_state(
            root, review, p=p2)
        self.assertEqual("invalid", state)
        self.assertIn("관측 기록 해시 불일치", reason)

    def test_non_economic_review_unaffected_by_p(self):
        """A content-kind review over NON-economic facts is economic
        not_applicable — the composed result stays the P2 'valid'.
        (Lane 18 fixture repair: the original seeded the ECONOMIC fact
        with a plain-text report bound — post-widening that fixture is
        economic-scoped and would be invalid, so the seed is now the
        genuinely non-economic _seed.)"""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed(root)
        write_text(root, "report.md", "content notes\n")
        obs = _observation(
            core, root, p, kinds=("content",), report_path="report.md")
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        value = _calc_review_value(core, root, p, path="report.md")
        value["review_kind"] = "content"
        value["provenance_path"] = receipt["path"]
        value["provenance_hash"] = receipt["hash"]
        value["input_revision"] = p["revision"]
        p2 = core.apply(
            root,
            {"request_id": "bind-content",
             "ops": [{"collection": "reviews", "value": value}]},
            p["revision"])
        review = p2["reviews"]["rev-calc"]
        self.assertEqual(
            ("valid", ""),
            core.review_observation_state(root, review, p=p2))
        self.assertTrue(
            core.review_observation_valid(root, review, p=p2))

    def test_two_arg_call_keeps_prior_behavior(self):
        """The SAME stale fixture called WITHOUT p= still returns
        ("valid", "") — pre-Lane-14 call sites are unchanged."""
        core = runtime("gg_core")
        root = self.make_project()
        p2, review = self._stale_economic_fixture(core, root)
        self.assertEqual(
            ("valid", ""),
            core.review_observation_state(root, review))
        self.assertTrue(core.review_observation_valid(root, review))

    @staticmethod
    def _born_stale_fixture(core, root):
        """P2-valid bound review whose report was AUTHORED claiming
        input_revision R+99.  All bindings stay consistent with those
        exact bytes/fields (observation hashes the mutated report;
        review.report_hash is derived from the same bytes at apply), so
        the only failing freshness conjunct is the
        obs/review/report input_revision triple — zero fact, target_refs
        or fingerprint-input drift."""
        sem = runtime("gg_fact_semantics")
        p = _seed_econ(root)
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(sem, core, root, p, refs)
        report["semantics"]["input_revision"] = p["revision"] + 99
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        obs = _observation(core, root, p, report_path="report.json")
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        value = _calc_review_value(core, root, p)
        value["provenance_path"] = receipt["path"]
        value["provenance_hash"] = receipt["hash"]
        value["input_revision"] = p["revision"]
        p2 = core.apply(
            root,
            {"request_id": "bind",
             "ops": [{"collection": "reviews", "value": value}]},
            p["revision"])
        return p2, p2["reviews"]["rev-calc"]

    def test_gate_excludes_stale_review(self):
        """Isolated trigger (L14-001 fix): the bound report claims
        input_revision R+99 while obs/review pin R — every OTHER gate()
        conjunct (status/disposition/findings/coverage/fingerprint) is
        provably still satisfied, so the exclusion comes specifically
        from the new 'stale' state, not a pre-existing conjunct."""
        core = runtime("gg_core")
        root = self.make_project()
        p2, review = self._born_stale_fixture(core, root)
        # isolation proof: the fingerprint conjunct still holds —
        # fact content and target_refs were never touched.
        self.assertEqual(
            review["input_fingerprint"],
            core.fingerprint(root, p2, review["target_refs"]))
        state, reason = core.review_observation_state(
            root, review, p=p2)
        self.assertEqual("stale", state)
        self.assertIn("input_revision", reason)
        entries = [
            e for e in core.gate(root, p2)
            if e["check_id"] == "review_calculation"]
        self.assertTrue(entries)
        self.assertTrue(all(e["status"] == "blocked" for e in entries))
        self.assertTrue(
            any("input_revision" in e["reason"] for e in entries))
        # negative control: identical binding with a consistent
        # input_revision satisfies review_calculation normally.
        root0 = self.make_project()
        p0 = _bind_economic_review(core, root0)
        self.assertEqual(
            [],
            [e for e in core.gate(root0, p0)
             if e["check_id"] == "review_calculation"
             and e["status"] == "blocked"])

    def test_gate_excludes_content_review_bound_to_non_report(self):
        """Lane 18 gate-level proof: a P2-VALID content-kind review over
        an economic fact bound to a plain non-report file used to satisfy
        review_content eligibility through the generic P2 path.  Now the
        shared freshness function runs for content kind too: report.md is
        not a gg-finance-semantics-report/1 -> "invalid" overrides the
        composed state -> the review drops out of review_content
        eligibility AND a finance_review entry appears in gate()'s
        output."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _seed_econ(root)
        write_text(root, "report.md", "content notes\n")
        # Lane 27 fixture repair: the widened write-time admission now
        # refuses this record shape (content kind + economic refs +
        # non-report path) at ingest/apply, so the committed record is
        # injected directly as a pre-Lane-27 legacy shape — the same
        # post-commit pattern gate()'s other pre-admission fixtures use.
        # The provenance points at a hand-written schema-1 receipt whose
        # fields satisfy _p2_review_observation_state's v1 checks.
        refs = core.sort_target_refs(core._export_refs(p))
        md_hash = core.digest((Path(root) / "report.md").read_bytes())
        obs = {
            "schema": "gg-review-observation/1",
            "observed_by": "observer-1",
            "author_session": "s-author",
            "reviewer_session": "s-reviewer",
            "author_id": "author", "reviewer_id": "reviewer",
            "review_kinds": ["content"],
            "target_refs": refs,
            "input_fingerprint": core.fingerprint(root, p, refs),
            "report_hash": md_hash,
        }
        obs_bytes = json.dumps(obs, ensure_ascii=False).encode("utf-8")
        obs_dir = Path(root) / ".gg-observations"
        obs_dir.mkdir(parents=True, exist_ok=True)
        (obs_dir / "obs-content.json").write_bytes(obs_bytes)
        p["reviews"]["rev-content"] = {
            "id": "rev-content", "review_kind": "content",
            "author_id": "author", "reviewer_id": "reviewer",
            "target_refs": refs,
            "input_fingerprint": core.fingerprint(root, p, refs),
            "findings": [], "disposition": "resolved", "status": "pass",
            "path": "report.md", "report_hash": md_hash,
            "coverage": {"econ": ["매출액 검토"]},
            "provenance_path": ".gg-observations/obs-content.json",
            "provenance_hash": core.digest(obs_bytes),
            "input_revision": p["revision"], "revision": 1,
        }
        p2 = p
        # P2 base is valid — the exclusion below is attributable solely
        # to the widened economic-freshness guard.
        review = p2["reviews"]["rev-content"]
        self.assertEqual(
            "valid", core._p2_review_observation_state(root, review)[0])
        entries = core.gate(root, p2)
        fin = [e for e in entries
               if e["check_id"] == "finance_review"
               and e["target"] == "rev-content"]
        self.assertTrue(fin)
        self.assertTrue(any("보고서" in e["reason"] for e in fin))
        content = [e for e in entries
                   if e["check_id"] == "review_content"]
        self.assertTrue(content)
        self.assertTrue(all(e["status"] == "blocked" for e in content))
        self.assertTrue(any("보고서" in e["reason"] for e in content))


class SemanticContextRegistryTests(ContractCase):
    """Lane 17/28 / design-r2 §7: ``_semantic_context`` resolves the
    deployed registry file ``skills/knuaf-doc/references/
    fact-semantics.json`` relative to ``gg_core.py``'s own location
    (scripts/ -> references/) — never inside the user's project root.
    A bare ``make_project`` fixture therefore gets the REAL deployed
    registry: its sha256, version, meanings and consumer_bindings flow
    through to ``classify_fact``/``resolve_consumers``/
    ``validate_report_binding``.  Only a genuinely missing, unreadable,
    corrupt or wrong-schema DEPLOYED file falls back to
    ``registry=None`` — simulated by patching ``__file__`` to a
    deployment tree without one, never by editing the real file."""

    REGISTRY_REL = Path(
        "skills/knuaf-doc/references/fact-semantics.json")

    @staticmethod
    def _installed_bytes(core):
        """The real deployed registry bytes, located via the imported
        module's own path (scripts/ -> knuaf-doc/references/)."""
        return (Path(core.__file__).resolve().parent.parent
                / "references" / "fact-semantics.json").read_bytes()

    def test_installed_registry_flows_through(self):
        """The deployed registry resolves with NO project-root copy —
        Lane 28: _semantic_context reads the file relative to
        gg_core.py's own location, so a bare project still gets the
        real registry's sha256/version/meanings."""
        core = runtime("gg_core")
        root = self.make_project()
        raw = self._installed_bytes(core)
        registry = core._semantic_context(root, core.load(root))[
            "registry"]
        self.assertIsNotNone(registry)
        # digest + version recomputed from the fixture bytes, never
        # hardcoded.
        self.assertEqual(
            hashlib.sha256(raw).hexdigest(), registry["sha256"])
        self.assertEqual(
            json.loads(raw).get("version"), registry["version"])
        self.assertIsInstance(registry["consumer_bindings"], list)

    def test_registry_resolves_from_deployment_not_project_root(self):
        """Lane 28 decisive regression: a project root containing a
        DECOY file at the old lookup location
        (root/skills/knuaf-doc/references/fact-semantics.json) still
        resolves the REAL deployed registry — proving the source moved
        off the project root entirely."""
        core = runtime("gg_core")
        root = self.make_project()
        decoy = Path(root) / self.REGISTRY_REL
        decoy.parent.mkdir(parents=True, exist_ok=True)
        decoy.write_bytes(
            b'{"schema": "gg-fact-semantics-registry/1",'
            b' "version": "decoy", "meanings": {},'
            b' "consumer_bindings": []}')
        registry = core._semantic_context(root, core.load(root))[
            "registry"]
        self.assertIsNotNone(registry)
        raw = self._installed_bytes(core)
        self.assertEqual(
            hashlib.sha256(raw).hexdigest(), registry["sha256"])
        self.assertNotEqual("decoy", registry["version"])
        self.assertEqual(
            json.loads(raw).get("version"), registry["version"])

    def test_meanings_passthrough_not_stub(self):
        core = runtime("gg_core")
        root = self.make_project()
        registry = core._semantic_context(root, core.load(root))[
            "registry"]
        self.assertTrue(registry["meanings"])
        self.assertIn("sales.revenue", registry["meanings"])

    def test_absent_registry_falls_back_to_none(self):
        """Deployed file genuinely absent -> registry=None.  Lane 28:
        simulated by resolving __file__ against a scripts dir with no
        sibling references/ — the real deployed file is never touched."""
        core = runtime("gg_core")
        root = self.make_project()
        fake = str(Path(root) / "pkg" / "scripts" / "gg_core.py")
        with mock.patch.object(core, "__file__", fake):
            self.assertIsNone(
                core._semantic_context(root, core.load(root))[
                    "registry"])

    def test_corrupt_and_wrong_schema_fall_back(self):
        """Corrupt/wrong-schema deployed file -> registry=None.  The
        variant bytes are written under a FAKE deployment tree reached
        by patching __file__ — the real file is never modified."""
        core = runtime("gg_core")
        for raw in (b"{not json",
                    b'{"schema": "other/9", "meanings": {}}'):
            tmp = Path(self.make_project())
            refs = tmp / "pkg" / "references"
            refs.mkdir(parents=True)
            (refs / "fact-semantics.json").write_bytes(raw)
            fake = str(tmp / "pkg" / "scripts" / "gg_core.py")
            root = self.make_project()
            with mock.patch.object(core, "__file__", fake):
                self.assertIsNone(
                    core._semantic_context(root, core.load(root))[
                        "registry"])


class ApplySectionClaimShapeTests(ContractCase):
    """design-r2 §7 apply row: a NEW section op that opts into the P3
    list-shaped ``claims`` gets its identifier structure validated
    BEFORE canonical commit — forbidden ``output_id``/``output_pointer``
    keys and plural-or-missing ``fact_id``/``projection_id`` are
    write-rejected.  Dict-shaped or absent ``claims`` stay legacy:
    preserved exactly as written and diagnosed only downstream in
    ``checks()``, never rejected at apply.  Field-value completeness
    (``unit``/``period``/``quote`` falsy) is likewise a checks()
    concern, so seeded fixtures carrying them keep committing.
    """

    def _apply_section(self, core, root, claims, *, ops_extra=()):
        write_text(root, "intro.md", "농장명은 행복농장이다.\n")
        value = {"id": "intro", "title": "머리말", "path": "intro.md",
                 "order": 1, "status": "drafting"}
        if claims is not None:
            value["claims"] = claims
        return core.apply(
            root,
            {"request_id": "claim-shape",
             "ops": [{"collection": "sections", "value": value},
                     *ops_extra]},
            0)

    def _assert_rejected(self, core, root, claims):
        before = (root / "project.json").read_bytes()
        with self.assertRaises(core.OperationError) as cm:
            self._apply_section(core, root, claims)
        self.assertEqual("invalid_change", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertIn("section claim 구조 오류", str(cm.exception))
        self.assertEqual(before, (root / "project.json").read_bytes())
        self.assertEqual(0, core.load(root)["revision"])

    @staticmethod
    def _valid_claim():
        fact = fact_op("sales", "매출액", "1000", "천원", period="2026년")
        return fact["value"], section_claim(fact["value"], "행복농장")

    def test_output_id_key_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        claim = dict(self._valid_claim()[1], output_id="out-1")
        self._assert_rejected(core, root, [claim])

    def test_output_pointer_key_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        claim = dict(self._valid_claim()[1], output_pointer="outputs/x")
        self._assert_rejected(core, root, [claim])

    def test_both_identifiers_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        claim = dict(self._valid_claim()[1], projection_id="proj-1")
        self._assert_rejected(core, root, [claim])

    def test_no_identifier_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        claim = {k: v for k, v in self._valid_claim()[1].items()
                 if k != "fact_id"}
        self._assert_rejected(core, root, [claim])

    def test_non_dict_entry_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        self._assert_rejected(core, root, ["행복농장"])

    def test_wellformed_claim_commits(self):
        core = runtime("gg_core")
        root = self.make_project()
        _, claim = self._valid_claim()
        p = self._apply_section(core, root, [claim])
        stored = p["sections"]["intro"]
        self.assertEqual(1, p["revision"])
        self.assertEqual([claim], stored["claims"])
        self.assertIn("draft_hash", stored)

    def test_field_completeness_defers_to_checks(self):
        """A list claim with a sound identifier but falsy unit/period
        commits — value completeness is diagnosed downstream in
        checks(), not write-rejected (the boundary that keeps seeded
        unit=None/period=None fixtures valid)."""
        core = runtime("gg_core")
        root = self.make_project()
        claim = dict(self._valid_claim()[1], unit=None, period=None)
        p = self._apply_section(core, root, [claim])
        self.assertEqual([claim], p["sections"]["intro"]["claims"])

    def test_fail_closed_filter_rejects_mixed_claim(self):
        """L16-001 fail-closed proof: the write-time filter defers ONLY
        the two named reason patterns (unsupported key / missing field)
        and rejects everything else.  This claim trips a DEFERRED reason
        (missing 'unit') and a NON-deferred one (forbidden
        'output_pointer') in the same entry — the apply still rejects,
        and the deferred reason is filtered out of the raised message.
        Since validate_claim_shape is frozen and cannot be made to emit
        an unfamiliar reason, this mixed-issue claim is the strongest
        in-repo proof that deferral applies only to the two exempted
        patterns while every other issue stays write-fatal."""
        core = runtime("gg_core")
        root = self.make_project()
        claim = dict(self._valid_claim()[1], output_pointer="outputs/x")
        del claim["unit"]
        before = (root / "project.json").read_bytes()
        with self.assertRaises(core.OperationError) as cm:
            self._apply_section(core, root, [claim])
        self.assertEqual("invalid_change", cm.exception.result["reason"])
        self.assertIn("금지된 claim key", str(cm.exception))
        self.assertNotIn("필수 필드", str(cm.exception))
        self.assertEqual(before, (root / "project.json").read_bytes())
        self.assertEqual(0, core.load(root)["revision"])

    def test_legacy_dict_and_absent_claims_preserved(self):
        """Dict-shaped (or absent) ``claims`` is the LEGACY shape: it is
        preserved byte-for-byte and diagnosed only in checks(), even
        when its ad-hoc content would fail validate_claim_shape were it
        mistakenly treated as a list entry (no identifier, no required
        fields)."""
        core = runtime("gg_core")
        root = self.make_project()
        legacy = {"c1": {"field_id": "farm_name", "value": "행복농장",
                         "unit": None, "period": None, "scope": "farm",
                         "answer_state": "provided",
                         "kind": "reported_fact"}}
        write_text(root, "second.md", "두 번째 절.\n")
        p = self._apply_section(
            core, root, legacy,
            ops_extra=[{"collection": "sections", "value": {
                "id": "second", "title": "둘째", "path": "second.md",
                "order": 2, "status": "drafting"}}])
        self.assertEqual(legacy, p["sections"]["intro"]["claims"])
        self.assertNotIn("claims", p["sections"]["second"])


def _plan_fact(fid, field_id, value, claim_id):
    """A provided monetary fact carrying the P3 metadata that makes
    classify_fact resolve it as finance_role='plan' — i.e. a REQUIRED
    body item per build_check_set's required_ids rule."""
    fact = fact_op(fid, field_id, value, "천원",
                   value_type="decimal", period="2026년",
                   scope="농장 전체", claim_id=claim_id)
    fact["value"]["measure"] = {"kind": "monetary_total",
                                "currency": "KRW"}
    fact["value"]["finance_role"] = "plan"
    return fact


def _plan_claim(fid, field_id, quote, value, **over):
    claim = {"fact_id": fid, "field_id": field_id, "quote": quote,
             "value": value, "unit": "천원", "period": "2026년",
             "scope": "농장 전체"}
    claim.update(over)
    return claim


def _seed_plan_sections(root, claims):
    """A complete-mode project (every section carries a list-shaped
    claims): two plan-role monetary facts plus one section whose claims
    list the caller controls."""
    write_text(root, "answer.txt", "매출액: 1000천원 / 이익: 200천원\n")
    write_text(
        root, "fin.md",
        "농장 전체 2026년 매출액은 1,000천원이고 이익은 200천원이다.\n")
    sales = _plan_fact("sales", "매출액", "1000", "c1")
    profit = _plan_fact("profit", "이익", "200", "c2")
    return runtime("gg_core").apply(
        root,
        {"request_id": "seed", "ops": [
            source_op(claims={**claim_of(sales["value"]),
                              **claim_of(profit["value"], "c2")}),
            sales, profit,
            section_op("fin", "재무", "fin.md", claims=claims)]},
        0)


class BodyCrosscheckSectionsWiringTests(ContractCase):
    """Lane 21 / architect finding C7-FINAL-1: checks() now calls the
    real §6 entry point ``gg_finance_body.crosscheck_sections`` —
    additively under check_id ``body_finance_crosscheck_v2`` and gated
    on the same complete_mode litmus crosscheck_sections applies
    internally (every section carrying list-shaped claims).  This
    delivers Lane 16's deferral promise (unsupported-key/missing-field
    CLAIM_INVALID issues finally surface at read time) and required-set
    completeness (CLAIM_PLAN_MISSING) that the legacy claimed-subset
    scan could never express."""

    def _v2(self, core, root, p):
        return [
            r for r in core.checks(root, p)
            if r["check_id"] == "body_finance_crosscheck_v2"]

    def test_unsupported_claim_key_surfaces_claim_invalid(self):
        """Deferred-issue delivery #1: a stored claim carrying an
        unsupported key commits at apply (deferred) and now surfaces a
        CLAIM_INVALID finding at checks() — previously promised, never
        delivered."""
        core = runtime("gg_core")
        root = self.make_project()
        claims = [
            _plan_claim("sales", "매출액", "매출액은 1,000천원", "1000",
                        bogus_key="x"),
            _plan_claim("profit", "이익", "이익은 200천원", "200"),
        ]
        p = _seed_plan_sections(root, claims)
        v2 = self._v2(core, root, p)
        # the malformed claim still satisfies its required item, so the
        # ONLY v2 finding is the deferred shape issue itself.
        self.assertEqual(1, len(v2))
        self.assertIn("지원하지 않는 claim key", v2[0]["reason"])
        self.assertEqual("fail", v2[0]["status"])

    def test_missing_required_field_surfaces_claim_invalid(self):
        """Deferred-issue delivery #2: a stored claim missing 'period'
        surfaces CLAIM_INVALID '필수 필드 누락' at checks() time."""
        core = runtime("gg_core")
        root = self.make_project()
        sales = _plan_claim("sales", "매출액", "매출액은 1,000천원", "1000")
        del sales["period"]
        claims = [sales,
                  _plan_claim("profit", "이익", "이익은 200천원", "200")]
        p = _seed_plan_sections(root, claims)
        v2 = self._v2(core, root, p)
        self.assertEqual(1, len(v2))
        self.assertIn("필수 필드 누락: period", v2[0]["reason"])

    def test_unclaimed_required_fact_is_claim_plan_missing(self):
        """The 'checked set shrinks to claimed subset' defect fixed:
        with two required plan facts but claims covering only 'sales',
        the legacy path stays silent on 'profit' (it only ever checks
        claimed facts) while the v2 path reports CLAIM_PLAN_MISSING."""
        core = runtime("gg_core")
        root = self.make_project()
        claims = [_plan_claim("sales", "매출액", "매출액은 1,000천원",
                              "1000")]
        p = _seed_plan_sections(root, claims)
        report = core.checks(root, p)
        ids = {(r["check_id"], r["target"]) for r in report}
        # legacy path verified silent on the omitted fact — the defect.
        self.assertNotIn(("body_finance_crosscheck", "profit"), ids)
        v2 = [r for r in report
              if r["check_id"] == "body_finance_crosscheck_v2"]
        self.assertEqual(1, len(v2))
        self.assertEqual("profit", v2[0]["target"])
        self.assertIn("필수 fact/projection", v2[0]["reason"])
        self.assertEqual("fail", v2[0]["status"])

    def test_noncomplete_sections_stay_legacy_only(self):
        """Backward-compat: a section WITHOUT list claims keeps the
        legacy path's exclusive ownership — its text scan still fires
        under 'body_finance_crosscheck' while the v2 path emits nothing
        (its internal fallback would re-run the same crosscheck_body
        scan and double-report under a second check_id)."""
        core = runtime("gg_core")
        root = self.make_project()
        write_text(root, "answer.txt", "매출액: 1000천원\n")
        write_text(root, "fin.md", "본문에 수치 없음.\n")
        sales = fact_op("sales", "매출액", "1000", "천원",
                        value_type="decimal", period="2026년",
                        scope="농장 전체", claim_id="c1")
        p = core.apply(
            root,
            {"request_id": "seed", "ops": [
                source_op(claims=claim_of(sales["value"])), sales,
                section_op("fin", "재무", "fin.md")]},
            0)
        report = core.checks(root, p)
        self.assertTrue([
            r for r in report
            if r["check_id"] == "body_finance_crosscheck"])
        self.assertEqual([], [
            r for r in report
            if r["check_id"] == "body_finance_crosscheck_v2"])


class FactSemanticsRegistryWiringTests(ContractCase):
    """Lane 22 / architect finding C7-FINAL-2: checks()'s per-fact
    ``classify_fact`` now receives the real installed registry through
    the shared ``_semantic_context`` call — no longer a hardcoded None.

    Honest scope note (verified by reading gg_fact_semantics):
    ``classify_fact`` internally calls ``validate_metadata(...,
    mode="legacy")``, whose only registry consumer — the
    meanings-membership check — is gated behind ``mode == "new"``, and
    no other branch reads ``registry`` at all.  So this wiring is a
    forward-looking consistency correction with NO currently-observable
    behavior difference: these tests prove the plumbing (registry
    reaches every classify call, one context build per checks() pass)
    and the no-diff invariant, not a false behavior change."""

    def test_installed_registry_reaches_every_classify_call(self):
        """Plumbing proof: with the real registry installed, EVERY
        classify_fact invocation inside checks() — the fact_semantics
        loop, the Lane-21 required-item adapter, and any review-scope
        classification — receives the registry dict (sha256-matched to
        the installed bytes), and _semantic_context is built exactly
        once per checks() pass."""
        core = runtime("gg_core")
        root = self.make_project()
        raw = SemanticContextRegistryTests._installed_bytes(core)
        p = _seed(root)
        expected_sha = hashlib.sha256(raw).hexdigest()
        sem = runtime("gg_fact_semantics")
        real_classify = sem.classify_fact
        seen_registries = []

        def spy(fact, *, registry, bindings):
            seen_registries.append(registry)
            return real_classify(fact, registry=registry,
                                 bindings=bindings)

        real_ctx = core._semantic_context
        ctx_calls = []

        def ctx_spy(r, proj):
            ctx_calls.append(r)
            return real_ctx(r, proj)

        sem.classify_fact = spy
        core._semantic_context = ctx_spy
        try:
            core.checks(root, p)
        finally:
            sem.classify_fact = real_classify
            core._semantic_context = real_ctx
        self.assertEqual(1, len(ctx_calls))
        self.assertTrue(seen_registries)
        self.assertTrue(all(
            isinstance(reg, dict)
            and reg.get("sha256") == expected_sha
            for reg in seen_registries))

    def test_registry_presence_does_not_change_fact_semantics(self):
        """No-diff invariant: the same issue-emitting fact (invalid
        measure currency) produces byte-identical fact_semantics entries
        with and without the installed registry — mode='legacy' has no
        reachable registry-dependent branch today, and a bogus
        meaning_id still escapes membership rejection (that check is
        mode='new'-only), so nothing about emitted issues may differ."""
        core = runtime("gg_core")
        root_no = self.make_project()
        p_no = _seed_bad_measure(root_no)
        root_yes = self.make_project()
        p_yes = _seed_bad_measure(root_yes)
        # sanity: root_yes resolves the deployed registry; root_no is
        # driven through a deployment tree WITHOUT one (Lane 28: the
        # file comes from gg_core.py's own location, so absence is
        # simulated by patching __file__ — never by deleting the real
        # file).
        self.assertIsNotNone(
            core._semantic_context(root_yes, p_yes)["registry"])
        fake = str(Path(root_no) / "pkg" / "scripts" / "gg_core.py")

        def fact_semantics(root, p):
            return [
                (r["check_id"], r["target"], r["status"], r["reason"])
                for r in core.checks(root, p)
                if r["check_id"] == "fact_semantics"]

        with mock.patch.object(core, "__file__", fake):
            self.assertIsNone(
                core._semantic_context(root_no, p_no)["registry"])
            fs_no = fact_semantics(root_no, p_no)
        self.assertTrue(fs_no)
        self.assertEqual(fs_no, fact_semantics(root_yes, p_yes))


class ApplyNewFactMetadataAdmissionTests(ContractCase):
    """Lane 23 / architect finding C7-FINAL-3 (design-r2 §2): a NEW fact
    record's explicit P3 metadata is admitted under
    ``validate_metadata(mode="new")`` BEFORE canonical commit — invalid
    finance_role, unknown measure keys, non-KRW currency and
    unregistered meaning_id all write-reject instead of only surfacing
    later in checks().  Updates (prev is not None) are not re-admitted,
    mirroring the reviews admission gate; pure-legacy facts pay zero
    friction."""

    def _apply_fact(self, core, root, fact_op_dict, *, request_id,
                    expected=0):
        write_text(root, "answer.txt", "매출액: 1000천원\n")
        return core.apply(
            root,
            {"request_id": request_id,
             "ops": [source_op(claims=claim_of(fact_op_dict["value"])),
                     fact_op_dict]},
            expected)

    def _assert_rejected(self, core, root, fact_op_dict):
        before = (root / "project.json").read_bytes()
        with self.assertRaises(core.OperationError) as cm:
            self._apply_fact(core, root, fact_op_dict,
                             request_id="bad-fact")
        self.assertEqual("invalid_change", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertIn("신규 사실 메타데이터 오류", str(cm.exception))
        self.assertEqual(before, (root / "project.json").read_bytes())
        self.assertEqual(0, core.load(root)["revision"])

    @staticmethod
    def _fact():
        return fact_op("sales", "매출액", "1000", "천원",
                       value_type="decimal", period="2026년",
                       scope="농장 전체", claim_id="c1")

    def test_bogus_finance_role_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        fact = self._fact()
        fact["value"]["finance_role"] = "bogus"
        self._assert_rejected(core, root, fact)

    def test_unknown_measure_key_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        fact = self._fact()
        fact["value"]["measure"] = {
            "kind": "monetary_total", "bogus_key": 1}
        self._assert_rejected(core, root, fact)

    def test_non_krw_currency_rejected(self):
        core = runtime("gg_core")
        root = self.make_project()
        fact = self._fact()
        fact["value"]["measure"] = {
            "kind": "monetary_total", "currency": "USD"}
        self._assert_rejected(core, root, fact)

    def test_unregistered_meaning_id_rejected_with_registry(self):
        """With the installed registry live, mode='new' membership
        applies: an unknown meaning_id write-rejects while a declared
        one commits.  Lane 28: the deployed registry resolves
        unconditionally — no project-root copy is needed."""
        core = runtime("gg_core")
        root = self.make_project()
        raw = SemanticContextRegistryTests._installed_bytes(core)
        known_meaning = next(iter(json.loads(raw)["meanings"]))

        bad = self._fact()
        bad["value"]["meaning_id"] = "bogus.unregistered"
        self._assert_rejected(core, root, bad)

        good = self._fact()
        good["value"]["meaning_id"] = known_meaning
        p = self._apply_fact(core, root, good, request_id="good-fact")
        self.assertEqual(1, p["revision"])
        self.assertEqual(
            known_meaning, p["facts"]["sales"]["meaning_id"])

    def test_wellformed_metadata_commits(self):
        core = runtime("gg_core")
        root = self.make_project()
        fact = self._fact()
        fact["value"]["meaning_id"] = "sales.revenue"
        fact["value"]["measure"] = {
            "kind": "monetary_total", "currency": "KRW", "unit": "천원"}
        fact["value"]["finance_role"] = "plan"
        p = self._apply_fact(core, root, fact, request_id="good-fact")
        self.assertEqual(1, p["revision"])
        stored = p["facts"]["sales"]
        self.assertEqual("plan", stored["finance_role"])
        self.assertEqual("sales.revenue", stored["meaning_id"])

    def test_legacy_fact_without_metadata_commits(self):
        """Zero-friction regression proof: a fact carrying NO P3
        metadata (meaning_id/measure/finance_role all absent) applies
        exactly as before — validate_metadata's early [] return."""
        core = runtime("gg_core")
        root = self.make_project()
        p = self._apply_fact(core, root, self._fact(),
                             request_id="legacy-fact")
        self.assertEqual(1, p["revision"])
        self.assertEqual("sales", p["facts"]["sales"]["id"])

    def test_update_not_readmitted(self):
        """prev is not None -> the admission check is skipped.  Updating
        an existing fact to carry metadata that would fail mode='new'
        still commits — the record's metadata was proven at creation,
        and re-validating every revision could reject unrelated edits
        (mirrors the reviews gate's update-not-readmitted rule)."""
        core = runtime("gg_core")
        root = self.make_project()
        p = self._apply_fact(core, root, self._fact(),
                             request_id="create-fact")
        updated = self._fact()
        updated["value"]["value"] = "2000"
        updated["value"]["finance_role"] = "bogus"
        p2 = self._apply_fact(core, root, updated,
                              request_id="update-fact",
                              expected=p["revision"])
        self.assertEqual(p["revision"] + 1, p2["revision"])
        self.assertEqual("bogus", p2["facts"]["sales"]["finance_role"])


def _output_formats_op():
    """The output_formats rule as a canonical op — fingerprint() digests
    p["rules"], so the rule must exist BEFORE any fingerprint-bearing
    record is committed (it cannot be patched onto p post-hoc)."""
    return {"collection": "rules", "value": {
        "id": "output_formats", "formats": ["xlsx"],
        "source_refs": [{"id": "answer", "locator": "l1",
                         "revision": 1}]}}


def _gate_econ_output_project(core, root, *, with_review):
    """An adopt-registered xlsx output over the economic 'sales' fact
    carrying every legacy gate() eligibility conjunct: live file hash,
    live fingerprint, coverage of all facts/sections, and all four xlsx
    check_ids each once with pass status and byte-matched evidence.
    With ``with_review`` a genuinely fresh bound calculation review
    over the SAME target_refs is committed BEFORE the adopt, so both
    sides' refs are identical ({answer,sales} — _export_refs grows to
    include the output only after adoption)."""
    sem = runtime("gg_fact_semantics")
    p = _seed_econ(root)
    p = _bind_specialty(root)
    p = core.apply(
        root, {"request_id": "fmt", "ops": [_output_formats_op()]},
        p["revision"])
    refs = core.sort_target_refs(core._export_refs(p))
    _write_xlsx(root, "out.xlsx", BYTES)
    report = _semantics_report(sem, core, root, p, refs)
    report["semantics"]["output_file_hash"] = core.digest(BYTES)
    write_text(root, "report.json",
               json.dumps(report, ensure_ascii=False))
    if with_review:
        obs = _observation(core, root, p)
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        value = _calc_review_value(core, root, p)
        value["provenance_path"] = receipt["path"]
        value["provenance_hash"] = receipt["hash"]
        value["input_revision"] = p["revision"]
        p = core.apply(
            root,
            {"request_id": "bind",
             "ops": [{"collection": "reviews", "value": value}]},
            p["revision"])
    ev_hash = core.digest((Path(root) / "report.json").read_bytes())
    fp = core.fingerprint(root, p, refs)
    checks = [
        {"check_id": cid, "status": "pass",
         "file_hash": core.digest(BYTES),
         "input_fingerprint": fp,
         "evidence_path": "report.json",
         "evidence_hash": ev_hash}
        for cid in ("structure", "recalculation", "crosscheck", "render")]
    core.adopt_output(
        root, _output_value_for(core, root, p, checks=checks),
        p["revision"], "adopt-x", major_id="specialty_crops")
    return core.load(root)


def _gate_plain_output_project(core, root):
    """Non-economic counterpart: a plain fact (no P3 metadata) plus an
    adopted xlsx output with all four xlsx checks byte-matched — zero
    reviews of any kind."""
    p = _seed_plain(root)
    p = _bind_specialty(root)
    p = core.apply(
        root, {"request_id": "fmt", "ops": [_output_formats_op()]},
        p["revision"])
    _write_xlsx(root, "out.xlsx", BYTES)
    write_text(root, "ev.txt", "EVIDENCE")
    refs = core.sort_target_refs(core._export_refs(p))
    fp = core.fingerprint(root, p, refs)
    checks = [
        {"check_id": cid, "status": "pass",
         "file_hash": core.digest(BYTES),
         "input_fingerprint": fp,
         "evidence_path": "ev.txt",
         "evidence_hash": core.digest(b"EVIDENCE")}
        for cid in ("structure", "recalculation", "crosscheck", "render")]
    core.adopt_output(
        root, _output_value_for(core, root, p, checks=checks),
        p["revision"], "adopt-x", major_id="specialty_crops")
    return core.load(root)


class EconomicOutputFreshnessGateTests(ContractCase):
    """Lane 24 (architect finding C7-FINAL-4): gate() output eligibility
    and _lineage_terminal_verify both require CURRENT economic
    freshness for economic-scoped outputs — a currently-fresh bound
    review over the output's exact target_refs per
    _economic_output_fresh (composing _finance_review_state, not a
    second freshness mechanism).  Lane 19's write-time report binding
    alone no longer suffices at read time; non-economic outputs are
    byte-identical in behavior."""

    def test_economic_output_without_fresh_review_excluded(self):
        """Every legacy conjunct holds (file hash, fingerprint, all four
        xlsx checks pass with byte-matched evidence, Lane-19 report
        binding satisfied at adopt) but no review exists at all ->
        the output is NOT eligible; output_xlsx is blocked.  Previously
        this output would have passed."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=False)
        entries = [
            e for e in core.gate(root, p)
            if e["check_id"] == "output_xlsx"]
        self.assertTrue(entries)
        self.assertTrue(all(e["status"] == "blocked" for e in entries))
        self.assertTrue(all(e["target"] == "project" for e in entries))

    def test_fresh_economic_review_restores_eligibility(self):
        """Positive control: the SAME output fixture plus a genuinely
        fresh bound calculation review over identical target_refs ->
        the output is eligible (output_xlsx pass targeting 'out').
        Proves the new gate composes live review freshness and does
        not over-reject."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=True)
        passes = [
            e for e in core.gate(root, p)
            if e["check_id"] == "output_xlsx" and e["status"] == "pass"]
        self.assertEqual(1, len(passes))
        self.assertEqual("out", passes[0]["target"])

    def test_noneconomic_output_unaffected(self):
        """Scope boundary: an adopted xlsx output over a fact with NO
        P3 metadata is non-economic — the predicate is not applicable,
        so legacy checks alone still make it eligible with zero
        reviews anywhere."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _gate_plain_output_project(core, root)
        passes = [
            e for e in core.gate(root, p)
            if e["check_id"] == "output_xlsx" and e["status"] == "pass"]
        self.assertEqual(1, len(passes))
        self.assertEqual("out", passes[0]["target"])

    def test_stale_economic_terminal_rejected(self):
        """Lineage terminal whose bound economic review goes stale:
        mutating the stored review's status post-hoc flips
        _finance_review_state to 'stale' (review_status conjunct)
        while fingerprint/publication/file-hash all still verify —
        the rejection is attributable specifically to the new economic
        freshness check."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=True)
        terminal = p["outputs"]["out"]
        # sanity: fresh state verifies silently.
        core._lineage_terminal_verify(root, p, terminal)
        p["reviews"]["rev-calc"]["status"] = "fail"
        # isolation proof: fingerprint still matches — every EARLIER
        # conjunct of _lineage_terminal_verify is still satisfied.
        self.assertEqual(
            terminal["input_fingerprint"],
            core.fingerprint(root, p, terminal["target_refs"]))
        with self.assertRaises(ValueError) as cm:
            core._lineage_terminal_verify(root, p, terminal)
        self.assertIn("경제 근거가 최신이 아님", str(cm.exception))

    def test_fresh_and_noneconomic_terminals_verify(self):
        """Regression proof: a fresh economic terminal verifies, and a
        non-economic terminal needs no economic evidence at all —
        both pass unchanged."""
        core = runtime("gg_core")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=True)
        core._lineage_terminal_verify(root, p, p["outputs"]["out"])
        root2 = self.make_project()
        p2 = _gate_plain_output_project(core, root2)
        core._lineage_terminal_verify(root2, p2, p2["outputs"]["out"])

    def _bind_second_fresh_review(self, core, sem, root, p, refs, *,
                                  registry):
        """A second, UNRELATED calculation review over the same
        ``refs``, bound to a different report file (reportB.json) —
        never bound to the output 'out' via its checks[] entries."""
        report_b = _semantics_report(
            sem, core, root, p, refs, registry=registry)
        write_text(root, "reportB.json",
                   json.dumps(report_b, ensure_ascii=False))
        obs_b = {
            "author_session": "s-author",
            "reviewer_session": "s-reviewer",
            "author_id": "author", "reviewer_id": "reviewer",
            "review_kinds": ["calculation"],
            "target_refs": refs,
            "input_fingerprint": core.fingerprint(root, p, refs),
            "report_path": "reportB.json",
            "report_hash": core.digest(
                (Path(root) / "reportB.json").read_bytes()),
            "input_revision": p["revision"],
        }
        receipt_b = core.ingest_review_observation(
            root, obs_b, "observer-2")
        value_b = _calc_review_value(
            core, root, p, rid="rev-b", path="reportB.json")
        value_b["target_refs"] = refs
        value_b["input_fingerprint"] = core.fingerprint(root, p, refs)
        value_b["provenance_path"] = receipt_b["path"]
        value_b["provenance_hash"] = receipt_b["hash"]
        value_b["input_revision"] = p["revision"]
        return core.apply(
            root,
            {"request_id": "bind-b",
             "ops": [{"collection": "reviews", "value": value_b}]},
            p["revision"])

    def test_unrelated_fresh_review_does_not_substitute(self):
        """Lane 26 (C7-FINAL-4-REOPEN) exploit reproduction: output 'out'
        is adopted with evidence bound to report A (report.json) and a
        review bound to A.  The registry then drifts -> A's review goes
        stale — while an UNRELATED review B over the SAME target_refs
        bound to a different, currently-fresh report stays fresh.
        O's OWN bound evidence is stale, so it must be excluded even
        though a fresh review over matching refs exists."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=True)
        refs = p["outputs"]["out"]["target_refs"]
        # Registry drifts (Lane 28: simulated by patching
        # _semantic_context — the deployed file resolves
        # unconditionally, so drift is a patched read, not a late
        # install).  A's report/review declared the pre-drift registry
        # -> stale; B's report is authored against the drifted registry
        # -> fresh.
        with _drifted_context(core):
            drifted = core._semantic_context(root, p)["registry"]
            state, _ = core._finance_review_state(
                root, p, p["reviews"]["rev-calc"],
                context=core._semantic_context(root, p))
            self.assertEqual("stale", state)
            # Unrelated review B: same refs, different report, fresh.
            p2 = self._bind_second_fresh_review(
                core, sem, root, p, refs, registry=drifted)
            state_b, _ = core._finance_review_state(
                root, p2, p2["reviews"]["rev-b"],
                context=core._semantic_context(root, p2))
            self.assertEqual("fresh", state_b)
            entries = [
                e for e in core.gate(root, p2)
                if e["check_id"] == "output_xlsx"]
            self.assertTrue(entries)
            self.assertTrue(
                all(e["status"] == "blocked" for e in entries))
            with self.assertRaises(ValueError) as cm:
                core._lineage_terminal_verify(
                    root, p2, p2["outputs"]["out"])
            self.assertIn("경제 근거가 최신이 아님", str(cm.exception))

    def test_coexisting_unrelated_fresh_review_still_eligible(self):
        """Positive control: with NO drift, O's own bound review is
        fresh and an unrelated second review over the same refs is also
        fresh — O stays eligible via its OWN bound evidence's identity,
        unaffected by the unrelated review's presence."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=True)
        refs = p["outputs"]["out"]["target_refs"]
        p2 = self._bind_second_fresh_review(
            core, sem, root, p, refs, registry=None)
        passes = [
            e for e in core.gate(root, p2)
            if e["check_id"] == "output_xlsx" and e["status"] == "pass"]
        self.assertEqual(1, len(passes))
        self.assertEqual("out", passes[0]["target"])

    def _adopt_second_output(self, core, sem, root, p):
        """Adopt a second same-format output 'out2' over the post-adopt
        refs ({answer,sales,out}), bound to its own report2.json — a
        legal metadata-only superseded_by target/sibling for the public
        apply() link path."""
        refs2 = core.sort_target_refs(core._export_refs(p))
        _write_xlsx(root, "out2.xlsx", BYTES2)
        report2 = _semantics_report(sem, core, root, p, refs2)
        report2["semantics"]["output_file_hash"] = core.digest(BYTES2)
        write_text(root, "report2.json",
                   json.dumps(report2, ensure_ascii=False))
        fp2 = core.fingerprint(root, p, refs2)
        checks2 = [
            {"check_id": cid, "status": "pass",
             "file_hash": core.digest(BYTES2),
             "input_fingerprint": fp2,
             "evidence_path": "report2.json",
             "evidence_hash": core.digest(
                 (Path(root) / "report2.json").read_bytes())}
            for cid in
            ("structure", "recalculation", "crosscheck", "render")]
        core.adopt_output(
            root,
            _output_value_for(
                core, root, p, path="out2.xlsx", data=BYTES2,
                checks=checks2),
            p["revision"], "adopt-y", major_id="specialty_crops")
        return core.load(root)

    def test_apply_rejects_stale_metadata_only_terminal_link(self):
        """Lane 27 (F06-COVERAGE-001): the metadata-only superseded_by
        link driven through the PUBLIC apply() — the same
        _meta_only_link -> _meta_only_link_verify ->
        _lineage_terminal_verify -> _economic_output_fresh composition
        prior lanes only exercised via the private helper.  'out' keeps
        a bound review/report that goes stale when the registry drifts
        in AFTER both outputs are adopted; linking out2 -> 'out' makes
        'out' the lineage terminal, so the link is refused at write
        time with invalid_change/not_committed and canonical bytes stay
        unchanged."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=True)
        p = self._adopt_second_output(core, sem, root, p)
        before = (Path(root) / "project.json").read_bytes()
        link = dict(p["outputs"]["out2"])
        link["superseded_by"] = "out"
        # Registry drifts after adoption (Lane 28: simulated by patching
        # _semantic_context): the terminal's bound report/review
        # declared the pre-drift registry -> stale.
        with _drifted_context(core):
            state, _ = core._finance_review_state(
                root, p, p["reviews"]["rev-calc"],
                context=core._semantic_context(root, p))
            self.assertEqual("stale", state)
            with self.assertRaises(core.OperationError) as cm:
                core.apply(
                    root,
                    {"request_id": "link",
                     "ops": [{"collection": "outputs", "value": link}]},
                    p["revision"])
            self.assertEqual(
                "invalid_change", cm.exception.result["reason"])
            self.assertEqual(
                "not_committed", cm.exception.result["commit_state"])
            # Attribution: only _lineage_terminal_verify's
            # economic-freshness conjunct produces this reason — the
            # rejection comes through the real metadata-only path, not
            # a private helper call.
            self.assertIn("경제 근거가 최신이 아님", str(cm.exception))
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())
        self.assertIsNone(
            core.load(root)["outputs"]["out2"].get("superseded_by"))

    def test_apply_accepts_fresh_metadata_only_terminal_link(self):
        """Positive control for the same public path: identical fixture
        with NO registry drift — the link commits, out2 becomes history
        with 'out' as terminal.  Proves the reject above is attributable
        to economic freshness, not fixture shape."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _gate_econ_output_project(core, root, with_review=True)
        p = self._adopt_second_output(core, sem, root, p)
        link = dict(p["outputs"]["out2"])
        link["superseded_by"] = "out"
        p2 = core.apply(
            root,
            {"request_id": "link",
             "ops": [{"collection": "outputs", "value": link}]},
            p["revision"])
        self.assertEqual(
            "out", p2["outputs"]["out2"]["superseded_by"])
        lineage = core.output_lineage(p2)
        self.assertEqual("history", lineage["out2"]["state"])
        self.assertEqual("out", lineage["out2"]["terminal"])


class StaleRegistryWriteAdmissionTests(ContractCase):
    """Lane 25 (architect finding C7-FINAL-5 / DESIGN.md §7 /
    ACCEPTANCE.md F06): the three economic write boundaries — apply's
    NEW calculation-review registration, ingest_review_observation's
    publication step, and adopt_output's evidence check — reject a
    bound report whose declared registry/refs/coverage no longer
    matches CURRENT state.  Lane 28: the deployed registry resolves
    unconditionally, so "stale" is staged by patching
    ``_semantic_context`` to return a drifted registry (the report was
    valid when authored/observed; current drifted afterwards) — never
    by touching the real deployed file.  Write-time rejection,
    canonical byte-invariant — the report is never admitted and later
    excluded; it is refused up front."""

    def _write_econ_report(
            self, core, sem, root, p, *, registry=None,
            kinds=("calculation",)):
        """Author report.json honestly over the seeded project ``p`` —
        against ``registry`` (None = the live context registry, the
        deployed file since Lane 28).  ``kinds`` sets the report's
        declared review.review_kinds so the bound file honestly
        describes the review kind under test."""
        refs = core.sort_target_refs(core._export_refs(p))
        report = _semantics_report(
            sem, core, root, p, refs, registry=registry)
        report["review"]["review_kinds"] = list(kinds)
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        return refs

    def test_apply_rejects_stale_registry_report(self):
        """apply boundary: the observation was ingested while the
        report was current (declared==current), then the registry
        drifts — the NEW review registration is refused with the
        Lane-25 reason and project.json is byte-unchanged."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        refs = self._write_econ_report(core, sem, root, p)
        obs = _observation(core, root, p)
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        before = (Path(root) / "project.json").read_bytes()
        value = _calc_review_value(core, root, p)
        value["provenance_path"] = receipt["path"]
        value["provenance_hash"] = receipt["hash"]
        value["input_revision"] = p["revision"]
        with _drifted_context(core):
            with self.assertRaises(core.OperationError) as cm:
                core.apply(
                    root,
                    {"request_id": "bind",
                     "ops": [{"collection": "reviews", "value": value}]},
                    p["revision"])
        self.assertEqual(
            "invalid_change", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        # isolation: fingerprint is registry-independent — only the new
        # current-registry/coverage stage can be rejecting here.
        self.assertEqual(
            value["input_fingerprint"],
            core.fingerprint(root, core.load(root), refs))
        self.assertIn("registry", str(cm.exception))
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())
        self.assertNotIn("rev-calc", core.load(root).get("reviews", {}))

    def test_apply_accepts_current_registry_report(self):
        """Positive control: the report's declared
        registry/refs/coverage all match current state, so observation
        + review commit normally."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        self._write_econ_report(core, sem, root, p)
        obs = _observation(core, root, p)
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        value = _calc_review_value(core, root, p)
        value["provenance_path"] = receipt["path"]
        value["provenance_hash"] = receipt["hash"]
        value["input_revision"] = p["revision"]
        p2 = core.apply(
            root,
            {"request_id": "bind",
             "ops": [{"collection": "reviews", "value": value}]},
            p["revision"])
        self.assertIn("rev-calc", p2["reviews"])

    def _apply_kind_review_stale(self, kind):
        """Shared fixture for the Lane-27 apply boundary: an honest
        report authored against the live registry + a schema-2
        observation ingested while declared==current; the caller then
        applies a NEW ``kind``-kind review under a drifted context."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        refs = self._write_econ_report(
            core, sem, root, p, kinds=(kind,))
        obs = _observation(core, root, p, kinds=(kind,))
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        before = (Path(root) / "project.json").read_bytes()
        rid = "rev-" + kind
        value = _calc_review_value(core, root, p, rid=rid)
        value["review_kind"] = kind
        value["provenance_path"] = receipt["path"]
        value["provenance_hash"] = receipt["hash"]
        value["input_revision"] = p["revision"]
        return core, root, p, refs, value, before, rid

    def test_apply_rejects_stale_registry_content_review(self):
        """Lane 27 (C7-L26-001) apply boundary: a NEW content-kind
        review over an economic fact bound to a stale-registry report
        is refused at write time — the admission guard now covers all
        three economic review kinds, not calculation alone."""
        core, root, p, refs, value, before, rid = (
            self._apply_kind_review_stale("content"))
        with _drifted_context(core):
            with self.assertRaises(core.OperationError) as cm:
                core.apply(
                    root,
                    {"request_id": "bind",
                     "ops": [{"collection": "reviews", "value": value}]},
                    p["revision"])
        self.assertEqual(
            "invalid_change", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertEqual(
            value["input_fingerprint"],
            core.fingerprint(root, core.load(root), refs))
        self.assertIn("registry", str(cm.exception))
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())
        self.assertNotIn(rid, core.load(root).get("reviews", {}))

    def test_apply_rejects_stale_registry_logic_review(self):
        """Same boundary for the logic kind — proves the guard widened
        to the full three-kind set, not just 'content' added."""
        core, root, p, refs, value, before, rid = (
            self._apply_kind_review_stale("logic"))
        with _drifted_context(core):
            with self.assertRaises(core.OperationError) as cm:
                core.apply(
                    root,
                    {"request_id": "bind",
                     "ops": [{"collection": "reviews", "value": value}]},
                    p["revision"])
        self.assertEqual(
            "invalid_change", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertIn("registry", str(cm.exception))
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())
        self.assertNotIn(rid, core.load(root).get("reviews", {}))

    def test_apply_accepts_current_registry_content_review(self):
        """Positive control: a NEW content-kind review bound to a
        CURRENT-registry report over the same economic fact commits —
        the widened guard admits honestly-bound content reviews rather
        than rejecting the kind outright."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        self._write_econ_report(
            core, sem, root, p, kinds=("content",))
        obs = _observation(core, root, p, kinds=("content",))
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        value = _calc_review_value(core, root, p, rid="rev-content")
        value["review_kind"] = "content"
        value["provenance_path"] = receipt["path"]
        value["provenance_hash"] = receipt["hash"]
        value["input_revision"] = p["revision"]
        p2 = core.apply(
            root,
            {"request_id": "bind",
             "ops": [{"collection": "reviews", "value": value}]},
            p["revision"])
        self.assertIn("rev-content", p2["reviews"])

    def test_observe_rejects_stale_registry_content_observation(self):
        """Lane 27 (C7-L26-001) ingest boundary: review_kinds
        ('content',) — no 'calculation' — over economic refs is now
        gated; the stale-registry report is refused with
        economic_report_schema_invalid and canonical is untouched."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        self._write_econ_report(
            core, sem, root, p, kinds=("content",))
        before = (Path(root) / "project.json").read_bytes()
        obs = _observation(core, root, p, kinds=("content",))
        with _drifted_context(core):
            with self.assertRaises(core.OperationError) as cm:
                core.ingest_review_observation(root, obs, "observer-1")
        self.assertEqual(
            "economic_report_schema_invalid",
            cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertIn("registry", str(cm.exception))
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())

    def test_observe_rejects_stale_registry_logic_observation(self):
        """Same ingest boundary for review_kinds ('logic',)."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        self._write_econ_report(
            core, sem, root, p, kinds=("logic",))
        before = (Path(root) / "project.json").read_bytes()
        obs = _observation(core, root, p, kinds=("logic",))
        with _drifted_context(core):
            with self.assertRaises(core.OperationError) as cm:
                core.ingest_review_observation(root, obs, "observer-1")
        self.assertEqual(
            "economic_report_schema_invalid",
            cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertIn("registry", str(cm.exception))
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())

    def test_observe_accepts_current_registry_content_observation(self):
        """Positive control: a content-kind observation bound to a
        CURRENT-registry report publishes normally — the widened guard
        admits honest content observations."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        self._write_econ_report(
            core, sem, root, p, kinds=("content",))
        obs = _observation(core, root, p, kinds=("content",))
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        self.assertTrue(receipt["path"])
        self.assertTrue(receipt["hash"])

    def test_observe_rejects_stale_registry_report(self):
        """ingest boundary: a report whose declared registry has drifted
        is refused at the publication step with
        economic_report_schema_invalid; the Lane-25 reason rides in the
        detail and canonical is untouched."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        self._write_econ_report(core, sem, root, p)
        before = (Path(root) / "project.json").read_bytes()
        obs = _observation(core, root, p)
        with _drifted_context(core):
            with self.assertRaises(core.OperationError) as cm:
                core.ingest_review_observation(root, obs, "observer-1")
        self.assertEqual(
            "economic_report_schema_invalid",
            cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertIn("registry", str(cm.exception))
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())

    def test_observe_accepts_current_registry_report(self):
        """Positive control: the report authored against the deployed
        registry publishes normally."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        self._write_econ_report(core, sem, root, p)
        obs = _observation(core, root, p)
        receipt = core.ingest_review_observation(root, obs, "observer-1")
        self.assertTrue(receipt["path"])
        self.assertTrue(receipt["hash"])

    def test_adopt_rejects_stale_registry_report(self):
        """adopt boundary: a byte-matched, schema-valid report bound to
        the output's own bytes (Lane-19 conjunct intact) is still not
        qualifying evidence once the registry has drifted — the
        candidate loop drops it and the whole adopt is refused with
        economic_evidence_missing; canonical untouched."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        self._write_econ_report(core, sem, root, p)
        _write_xlsx(root, "out.xlsx", BYTES)
        report = json.loads(
            (Path(root) / "report.json").read_text(encoding="utf-8"))
        report["semantics"]["output_file_hash"] = core.digest(BYTES)
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        before = (Path(root) / "project.json").read_bytes()
        ov = _output_value_for(
            core, root, p,
            checks=[{"evidence_path": "report.json",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.json").read_bytes())}])
        with _drifted_context(core):
            with self.assertRaises(core.OperationError) as cm:
                core.adopt_output(root, ov, p["revision"], "adopt-stale", major_id="specialty_crops")
        self.assertEqual(
            "economic_evidence_missing", cm.exception.result["reason"])
        self.assertEqual(
            "not_committed", cm.exception.result["commit_state"])
        self.assertEqual(
            before, (Path(root) / "project.json").read_bytes())
        self.assertNotIn("out", core.load(root).get("outputs", {}))

    def test_adopt_accepts_current_registry_report(self):
        """Positive control: report authored against the deployed
        registry with the Lane-19 output_file_hash binding — the
        economic output adopts normally."""
        core = runtime("gg_core")
        sem = runtime("gg_fact_semantics")
        root = self.make_project()
        p = _seed_econ(root)
        p = _bind_specialty(root)
        self._write_econ_report(core, sem, root, p)
        _write_xlsx(root, "out.xlsx", BYTES)
        report = json.loads(
            (Path(root) / "report.json").read_text(encoding="utf-8"))
        report["semantics"]["output_file_hash"] = core.digest(BYTES)
        write_text(root, "report.json",
                   json.dumps(report, ensure_ascii=False))
        ov = _output_value_for(
            core, root, p,
            checks=[{"evidence_path": "report.json",
                     "evidence_hash": core.digest(
                         (Path(root) / "report.json").read_bytes())}])
        result = core.adopt_output(
            root, ov, p["revision"], "adopt-ok", major_id="specialty_crops")
        self.assertEqual("adopted", result["status"])
        self.assertIn("out", core.load(root)["outputs"])


if __name__ == "__main__":
    unittest.main()
