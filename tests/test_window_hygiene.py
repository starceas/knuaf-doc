"""Office preservation boundaries: fake transport, real filesystem/lock.

No Office launch or user files. Script checks cover targeting syntax; native
AppleScript semantics and saved output caches require separate app evidence.
"""
import copy
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from tests._harness import ContractCase, bind_major, runtime

try:
    import fcntl
except ModuleNotFoundError as exc:
    if exc.name != "fcntl":
        raise
    fcntl = None


@unittest.skipUnless(fcntl is not None, "requires native POSIX fcntl locks; Office automation is not exercised")
class OfficeSessionTests(unittest.TestCase):
    def setUp(self):
        self.session = runtime("gg_office_session")
        self.office = runtime("gg_office")
        self.tmp = tempfile.TemporaryDirectory(prefix="office-hygiene-")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.guard = self.root / "guards"
        self.work, self.pdf = self.job("first")

    def job(self, name):
        folder = self.root / ("job-" + name) / "output"
        folder.mkdir(parents=True)
        work = folder / "same-name.docx"
        work.write_bytes(b"synthetic staged file")
        return work, folder / "same-name.pdf"

    def result(self, stdout="GG_OFFICE_SESSION|closed|saved|synthetic", *,
               rc=0, stderr=""):
        return subprocess.CompletedProcess(["synthetic"], rc, stdout, stderr)

    def run_job(self, raw, *, kind="word", prepare=None, work=None, pdf=None):
        script = (self.session.WORD_APPLESCRIPT_TEMPLATE if kind == "word"
                  else self.session.EXCEL_APPLESCRIPT_TEMPLATE).format(timeout=9)
        return self.session.run_session(
            kind, script, [str(work or self.work), str(pdf or self.pdf)], 9,
            run_raw=raw, prepare_word=prepare, guard_dir=self.guard)

    def transport(self, outcome=None):
        return mock.Mock(side_effect=[
            self.result("GG_OFFICE_INVENTORY|0"), outcome or self.result()])

    def assert_no_compensation(self, raw):
        for call in raw.call_args_list:
            script = call.args[0]
            self.assertNotIn("saving no", script)
            self.assertNotIn("quit", script)
        self.assertLessEqual(raw.call_count, 2)

    def test_clean_and_dirty_existing_documents_block_before_launch_or_open(self):
        for dirty in (False, True):
            with self.subTest(dirty=dirty):
                user = {"path": "/user/same-name.docx", "dirty": dirty,
                        "content": "unsaved user work", "open": True}
                before = copy.deepcopy(user)
                raw = mock.Mock(return_value=self.result("GG_OFFICE_INVENTORY|1"))
                prepare = mock.Mock()
                res = self.run_job(raw, prepare=prepare)
                self.assertEqual(res.returncode, 2)
                self.assertEqual(user, before)
                prepare.assert_not_called()
                self.assertEqual(raw.call_count, 1)
                self.assertEqual(res.office_session["cleanup"], "not_attempted")

    def test_unknown_inventory_blocks(self):
        for response in (self.result(""), self.result("bad", rc=1),
                         self.result("GG_OFFICE_INVENTORY|2")):
            with self.subTest(response=response):
                raw = mock.Mock(return_value=response)
                self.assertEqual(self.run_job(raw).returncode, 2)
                self.assertEqual(raw.call_count, 1)

    def test_open_update_save_export_dirty_and_foreign_errors_preserve_job(self):
        cases = [("opening", "open_failed"), ("updating", "field_error"),
                 ("saving", "save_failed"), ("exporting", "pdf_failed"),
                 ("checking_saved", "owned_document_dirty"),
                 ("updating", "foreign_or_missing_document"),
                 ("exporting", "owned_path_not_found")]
        for stage, reason in cases:
            with self.subTest(stage=stage, reason=reason):
                raw = self.transport(self.result(rc=1, stdout="", stderr=
                    "execution error: GG_OFFICE_SESSION|" + stage + "|" + reason))
                res = self.run_job(raw)
                self.assertEqual(res.returncode, 1)
                self.assertEqual(res.office_session["cleanup"], "preserved")
                self.assertTrue(self.work.exists())
                self.assertTrue(self.session.retain_job(self.work.parent.parent))
                self.assert_no_compensation(raw)

    def test_close_error_is_unknown_and_blocks_next_workspace(self):
        raw = self.transport(self.result(rc=1, stdout="", stderr=
            "GG_OFFICE_SESSION|closing|owned_document_still_open"))
        res = self.run_job(raw)
        self.assertEqual(res.office_session["cleanup"], "unknown")
        second_work, second_pdf = self.job("second")
        next_raw = mock.Mock()
        second = self.run_job(next_raw, work=second_work, pdf=second_pdf)
        self.assertEqual(second.returncode, 2)
        self.assertEqual(second.office_session["reason"], "previous_office_session_unknown")
        next_raw.assert_not_called()
        self.assertTrue(self.work.exists())

    def test_timeout_does_not_cleanup_or_allow_next_job(self):
        for when in ("preflight", "engine"):
            with self.subTest(when=when):
                # Separate guard files for independent timeout scenarios.
                self.guard = self.root / ("guards-" + when)
                replies = [TimeoutError("synthetic hung read")] if when == "preflight" else [
                    self.result("GG_OFFICE_INVENTORY|0"), TimeoutError("synthetic hung event")]
                raw = mock.Mock(side_effect=replies)
                with self.assertRaises(TimeoutError) as caught:
                    self.run_job(raw)
                self.assertEqual(caught.exception.office_session["cleanup"], "unknown")
                self.assertTrue(self.work.exists())
                self.assert_no_compensation(raw)
                next_raw = mock.Mock()
                self.assertEqual(self.run_job(next_raw).returncode, 2)
                next_raw.assert_not_called()

    def test_appleevent_timeout_is_unknown_even_when_script_returned(self):
        raw = self.transport(self.result(rc=1, stdout="", stderr=
            "GG_OFFICE_SESSION|opening|AppleEvent timed out (-1712)"))
        self.assertEqual(self.run_job(raw).office_session["cleanup"], "unknown")
        self.assert_no_compensation(raw)

    def test_prepare_failure_records_unknown_and_preserves(self):
        raw = mock.Mock(return_value=self.result("GG_OFFICE_INVENTORY|0"))
        prepare = mock.Mock(side_effect=RuntimeError("launch failure"))
        with self.assertRaisesRegex(RuntimeError, "launch failure"):
            self.run_job(raw, prepare=prepare)
        data = self.session.read_session(self.work.parent.parent)
        self.assertEqual(data["cleanup"], "unknown")
        self.assertEqual(raw.call_count, 1)
        self.assertTrue(self.work.exists())

    def test_unconfirmed_zero_exit_is_not_success(self):
        raw = self.transport(self.result("Word process completed"))
        res = self.run_job(raw)
        self.assertEqual(res.returncode, 2)
        self.assertEqual(res.office_session["cleanup"], "unknown")
        self.assert_no_compensation(raw)

    def test_success_records_document_close_but_no_app_quit(self):
        raw = self.transport()
        res = self.run_job(raw)
        self.assertEqual(res.returncode, 0)
        self.assertEqual(res.office_session["cleanup"], "closed")
        self.assertEqual(res.office_session["save_state"], "saved")
        self.assertEqual(res.office_session["app_quit"], "not_attempted")
        self.assertEqual(self.session.read_session(self.work.parent.parent), res.office_session)
        self.assert_no_compensation(raw)

    def test_other_cli_cannot_enter_same_app(self):
        import fcntl
        self.guard.mkdir()
        lock = self.guard / "word.lock"
        fd = os.open(lock, os.O_RDWR | os.O_CREAT, 0o600)
        self.addCleanup(os.close, fd)
        fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        raw = mock.Mock()
        res = self.run_job(raw)
        self.assertEqual(res.office_session["reason"], "office_session_busy")
        raw.assert_not_called()

    def test_app_locks_are_independent(self):
        import fcntl
        self.guard.mkdir()
        # Excel execution is held, but its lock must not block Word.
        fd = os.open(self.guard / "excel.lock", os.O_RDWR | os.O_CREAT, 0o600)
        self.addCleanup(os.close, fd)
        fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        raw = self.transport()
        self.assertEqual(self.run_job(raw, kind="word").returncode, 0)
        self.assertEqual(raw.call_count, 2)

    def test_inflight_crash_and_malformed_lock_fail_closed(self):
        self.guard.mkdir()
        for content in ('{"state":"inflight"}', '{"state":"unknown"}', '{'):
            with self.subTest(content=content):
                (self.guard / "word.lock").write_text(content, encoding="utf-8")
                raw = mock.Mock()
                self.assertEqual(self.run_job(raw).returncode, 2)
                raw.assert_not_called()

    def test_templates_use_path_for_every_write_and_no_error_cleanup(self):
        for kind, template in [("word", self.session.WORD_APPLESCRIPT_TEMPLATE),
                               ("excel", self.session.EXCEL_APPLESCRIPT_TEMPLATE)]:
            with self.subTest(kind=kind):
                text = template.format(timeout=9)
                if kind == "excel":
                    self.assertIn(self.session.EXCEL_RECALCULATION_HOLD, text)
                    self.assertIn("error ", text)
                    for operation in ("tell application", "open workbook", "calculate",
                                      "save ", "save as ", "close ", "used range"):
                        self.assertNotIn(operation, text)
                    self.assertNotIn("saving no", text)
                    self.assertNotIn("quit", text)
                    self.assertNotIn("active workbook", text)
                    self.assertNotIn("workbook 1", text)
                    self.assertNotIn("repeat 2 times", text)
                    continue
                self.assertNotIn("saving no", text)
                self.assertNotIn("quit", text)
                self.assertNotIn("active document", text)
                self.assertNotIn("active workbook", text)
                self.assertNotIn("document 1", text)
                noun = "document" if kind == "word" else "workbook"
                target = (f"first {noun} whose (full name is jobHFS or full name is jobPOSIX)"
                          if kind == "word" else
                          "item 1 of (get workbooks whose (full name is jobHFS or full name is jobPOSIX))")
                writes = [line.strip() for line in text.splitlines() if line.strip().startswith(
                    ("repaginate ", "update field ", "calculate (", "save ", "save as ", "close "))]
                self.assertTrue(writes)
                self.assertTrue(all(target in line for line in writes), writes)
                self.assertIn("close (" + target + ") saving ask", text)
                self.assertLess(text.index("if not (saved of"), text.index("close ("))
                self.assertNotIn("close ", text.split("on error messageText", 1)[1])

    def test_failure_after_confirmed_close_reports_cleanup_separately(self):
        res = self.run_job(self.transport())
        fail = {"status": "fail", "error": "PDF validation failed"}
        self.office._write_failure_receipt(self.work.parent.parent, fail)
        self.assertEqual(fail["status"], "fail")
        self.assertEqual(fail["office_session"]["cleanup"], "closed")
        self.assertEqual(res.office_session["cleanup"], "closed")

    def test_policy_hold_retains_live_job_but_removes_closed_job(self):
        mc = runtime("gg_major_contract")
        for cleanup, retained in [("preserved", True), ("unknown", True), ("closed", False)]:
            with self.subTest(cleanup=cleanup):
                work, _ = self.job("hold-" + cleanup)
                job = work.parent.parent
                (job / "office_session.json").write_text(json.dumps({
                    "schema": self.session.SCHEMA, "cleanup": cleanup}), encoding="utf-8")
                with mock.patch.object(mc, "reconfirm_output", side_effect=
                                       mc.OutputHeldError("project_revision_changed")):
                    with self.assertRaises(mc.OutputHeldError):
                        self.office._hold_if_moved([object()], object(), job)
                self.assertEqual(work.exists(), retained)

    def test_unreadable_session_never_authorizes_deletion(self):
        job = self.work.parent.parent
        (job / "office_session.json").write_text("{", encoding="utf-8")
        self.assertTrue(self.session.retain_job(job))

    def test_original_or_existing_pdf_cannot_enter_transport(self):
        raw = mock.Mock()
        with self.assertRaises(ValueError):
            self.run_job(raw, work=self.root / "original.docx")
        self.pdf.write_bytes(b"preexisting pdf")
        with self.assertRaises(ValueError):
            self.run_job(raw)
        raw.assert_not_called()


class DarwinExcelHoldTests(ContractCase):
    """No native transport: the accepted P1 disposition holds all Excel jobs."""

    HOLD = "owned_workbook_recalculation_unproven"

    def setUp(self):
        self.office = runtime("gg_office")
        self.session = runtime("gg_office_session")
        self.mc = runtime("gg_major_contract")
        self.root = self.make_project()
        bind_major(self.root)
        self.context = self.mc.output_context(self.root, "specialty_crops")
        self.out = self.root / "new-output"
        self.workspace = self.root / "new-workspace"

    def workbook(self, name, *, sheets=1, formulas=True):
        from openpyxl import Workbook
        wb = Workbook()
        for i in range(sheets):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = "S" + str(i + 1)
            ws["A1"] = "=1+1" if formulas else 2
        path = self.root / (name + ".xlsx")
        wb.save(path)
        return path

    def test_public_legacy_spec_receipt_multisheet_and_no_formula_hold(self):
        for case in ("legacy", "spec", "receipt", "multi-sheet", "no-formula"):
            source = self.workbook(case, sheets=3 if case == "multi-sheet" else 1,
                                   formulas=case != "no-formula")
            kwargs = {}
            if case == "spec":
                spec = self.root / "spec.json"
                spec.write_text(json.dumps({"major_id": "specialty_crops",
                    "profile": "school_17_sheet_v1"}), encoding="utf-8")
                kwargs["spec_file"] = spec
            if case == "receipt":
                digest = self.office.compute_sha256(source.read_bytes())
                receipt = self.root / "fill.json"
                receipt.write_text(json.dumps({
                    "schema": "gg-xlsx-fill-receipt/v1", "written": ["A1"],
                    "template": {"path": str(source), "sha256": digest},
                    "map": {"path": str(source), "sha256": digest},
                    "values": {"path": str(source), "sha256": digest},
                    "output": {"status": "filled", "sha256": digest},
                }), encoding="utf-8")
                kwargs["template_receipt"] = receipt
            before = self._tree_bytes(self.root)
            with self.subTest(case=case), mock.patch.object(sys, "platform", "darwin"), \
                    mock.patch.object(self.office, "_authorize_office",
                        wraps=self.office._authorize_office) as auth, \
                    mock.patch.object(self.office, "_preflight_excel_profile",
                        wraps=self.office._preflight_excel_profile) as profile, \
                    mock.patch.object(self.office, "stage_job",
                        side_effect=AssertionError("STAGE_REACHED")) as stage, \
                    mock.patch.object(self.office, "_run_excel_engine") as engine, \
                    mock.patch.object(self.office, "_publish_outputs") as publish, \
                    mock.patch.object(self.office, "_run_osascript_raw") as raw:
                with self.assertRaises(self.mc.OutputHeldError) as held:
                    self.office.process_excel(source, self.out,
                        workspace=self.workspace, context=self.context, **kwargs)
                self.assertEqual(self.HOLD, held.exception.reason)
                auth.assert_called_once()
                profile.assert_called_once()
                for boundary in (stage, engine, publish, raw):
                    boundary.assert_not_called()
            self.assertEqual(before, self._tree_bytes(self.root))
            self.assertFalse(self.workspace.exists())
            self.assertFalse(self.out.exists())

    def test_low_level_guard_holds_before_inventory_or_session_writes(self):
        folder = self.root / "job-direct" / "output"
        folder.mkdir(parents=True)
        work = folder / "synthetic.xlsx"
        work.write_bytes(b"synthetic staged workbook")
        pdf = folder / "synthetic.pdf"
        guard = self.root / "new-guards"
        raw = mock.Mock(side_effect=lambda script, args, **kw:
            subprocess.CompletedProcess([], 0, "GG_OFFICE_INVENTORY|0" if not args
                else "GG_OFFICE_SESSION|closed|saved|synthetic", ""))
        prepare = mock.Mock()
        def guarded(*args, **kwargs):
            return self.session.run_session(*args, **kwargs, guard_dir=guard)
        script = self.session.EXCEL_APPLESCRIPT_TEMPLATE.format(timeout=9)
        unsafe = self.session.SESSION_MARKER + '\ntell application "Microsoft Excel" to calculate'
        calls = (
            lambda: self.office._run_excel_engine(work, pdf, timeout=9),
            lambda: self.office.run_osascript(unsafe, [str(work), str(pdf)], 9),
            lambda: self.session.run_session("excel", script, [str(work), str(pdf)], 9,
                run_raw=raw, prepare_word=prepare, guard_dir=guard),
        )
        before = self._tree_bytes(self.root)
        for number, call in enumerate(calls):
            with self.subTest(entry=number), mock.patch.object(sys, "platform", "darwin"), \
                    mock.patch.object(self.office, "run_session", side_effect=guarded), \
                    mock.patch.object(self.office, "_run_osascript_raw", raw):
                with self.assertRaises(self.mc.OutputHeldError) as held:
                    call()
                self.assertEqual(self.HOLD, held.exception.reason)
            raw.assert_not_called()
            prepare.assert_not_called()
            self.assertEqual(before, self._tree_bytes(self.root))
            self.assertFalse(guard.exists())
            self.assertFalse(pdf.exists())

    def test_cli_and_mixed_batch_stop_at_excel_without_publishing(self):
        source = self.workbook("batch")
        first, later = self.root / "before.docx", self.root / "after.docx"
        from docx import Document
        for path in (first, later):
            doc = Document()
            doc.add_paragraph("synthetic Word input")
            doc.save(path)
        prior = self.root / "prior-word.pdf"
        def word(path, *args, **kwargs):
            prior.write_bytes(b"synthetic previously published Word PDF")
            return {"status": "converted"}
        with mock.patch.object(sys, "platform", "darwin"), \
                mock.patch.object(self.office, "stage_job",
                    side_effect=AssertionError("STAGE_REACHED")) as stage, \
                mock.patch.object(self.office, "_run_excel_engine") as engine, \
                mock.patch.object(self.office, "_publish_outputs") as publish, \
                mock.patch.object(self.office, "_run_osascript_raw") as raw, \
                mock.patch.object(self.office, "process_word", side_effect=word) as words:
            stdout = io.StringIO()
            with contextlib.redirect_stdout(stdout):
                rc = self.office.main(["excel", str(source), "--out-dir", str(self.out),
                    "--project", str(self.root), "--major", "specialty_crops"])
            self.assertEqual(2, rc)
            result = json.loads(stdout.getvalue())
            self.assertEqual("held", result["status"])
            self.assertEqual(self.HOLD, result["reason"])
            with self.assertRaises(self.mc.OutputHeldError) as held:
                self.office.process_batch([first, source, later], self.out,
                    workspace=self.workspace, context=self.context)
            self.assertEqual(self.HOLD, held.exception.reason)
            words.assert_called_once()
            self.assertEqual(first, words.call_args.args[0])
            self.assertEqual(b"synthetic previously published Word PDF", prior.read_bytes())
            for boundary in (stage, engine, publish, raw):
                boundary.assert_not_called()
        self.assertFalse(self.workspace.exists())
        self.assertFalse(self.out.exists())

    def test_word_and_windows_dispatch_keep_their_contracts(self):
        result = subprocess.CompletedProcess([], 0, "synthetic", "")
        with mock.patch.object(sys, "platform", "darwin"), \
                mock.patch.object(self.office, "run_osascript", return_value=result) as raw:
            self.assertIs(result, self.office._run_word_engine(Path("a.docx"), Path("a.pdf")))
            raw.assert_called_once()
        with mock.patch.object(sys, "platform", "win32"), \
                mock.patch.object(self.office, "run_windows_com", return_value=result) as com:
            self.assertIs(result, self.office._run_excel_engine(Path("a.xlsx"), Path("a.pdf")))
            self.assertEqual("excel", com.call_args.args[0])
            com.assert_called_once()


if __name__ == "__main__":
    unittest.main()
