"""gg_excel_fill: writing student values into the school's own workbook.

Migrated from the GUI contract suite (cb72ae1).  This is where a wrong number
reaches a submitted thesis.  Every refusal below protects either the school
form's own calculations (formulas, merges, cells nobody reviewed) or the
skill's honesty rule that an unanswered question must not become a figure.
"""
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from tests._harness import ContractCase, bind_major, runtime

SHEET = "1. 입력"


class ExcelFillTests(ContractCase):
    def setUp(self):
        self.fill = runtime("gg_excel_fill")
        import openpyxl as _openpyxl

        self.openpyxl = _openpyxl
        tmp = tempfile.TemporaryDirectory(prefix="p1-fill-")
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)
        self.template = self._template()
        self._ctx = None

    def _template(self):
        """Numbers, a label, a formula (C3) and a merged range (B5:D5)."""
        wb = self.openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET
        ws["B1"] = "면적(평):"
        ws["C1"] = 600
        ws["B2"] = "작목:"
        ws["C2"] = "도라지"
        ws["B3"] = "합계"
        ws["C3"] = "=C1*2"
        ws["B5"] = "병합 제목"
        ws.merge_cells("B5:D5")
        path = self.tmp / "학교 양식.xlsx"
        wb.save(path)
        self._keep_merged_non_anchor(path)
        return path

    @staticmethod
    def _keep_merged_non_anchor(path, ref="C5"):
        """Re-add an empty cell node inside a merged range.

        Excel keeps ``<c>`` nodes for merged non-anchor cells (they carry
        styles); openpyxl drops them on merge.  Without this the fixture
        cannot reach the merged-non-anchor refusal at all.
        """
        name = "xl/worksheets/sheet1.xml"
        with zipfile.ZipFile(path) as z:
            items = {n: z.read(n) for n in z.namelist()}
        xml = items[name].decode("utf-8")
        cut = xml.index('<c r="B5"')
        end = xml.index("</c>", cut) + len("</c>")
        items[name] = (xml[:end] + '<c r="%s"/>' % ref
                       + xml[end:]).encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for n, data in items.items():
                z.writestr(n, data)

    def ctx(self):
        """Policy B: the workbook output needs a bound major."""
        if self._ctx is None:
            project = self.make_project()
            bind_major(project)
            self._ctx = runtime("gg_major_contract").output_context(
                project, "specialty_crops")
        return self._ctx

    def entry(self, cell, **over):
        base = {"sheet": SHEET, "cell": cell, "role": "farm.area",
                "period": "2026", "source_note": "인터뷰 Q2", "editable": True}
        base.update(over)
        return base

    def value(self, cell, raw=700, **over):
        base = {"sheet": SHEET, "cell": cell, "value": raw,
                "value_type": "integer", "answer_state": "provided",
                "evidence_ref": {"source_id": "src-answers", "revision": 1,
                                 "locator": "L3", "origin": "factual"}}
        base.update(over)
        return base

    def files(self, entries, values):
        m = self.tmp / "map.json"
        v = self.tmp / "values.json"
        m.write_text(json.dumps({
            "schema": self.fill.MAP_SCHEMA,
            "template": {"sha256": self.fill.sha256(self.template)},
            "entries": entries}, ensure_ascii=False), encoding="utf-8")
        v.write_text(json.dumps({
            "schema": self.fill.VALUES_SCHEMA, "values": values},
            ensure_ascii=False), encoding="utf-8")
        return m, v

    def run_fill(self, entries, values, out_name="채운 사본.xlsx"):
        m, v = self.files(entries, values)
        return self.fill.fill_copy(self.template, m, v, self.tmp / out_name,
                                   context=self.ctx())

    def out(self):
        return self.tmp / "채운 사본.xlsx"

    # --- the happy path, and what must survive it -------------------------

    def test_a_reviewed_cell_is_written_and_the_form_survives(self):
        receipt = self.run_fill([self.entry("C1")], [self.value("C1", 700)])
        ws = self.openpyxl.load_workbook(self.out())[SHEET]
        self.assertEqual(700, ws["C1"].value)
        self.assertEqual("면적(평):", ws["B1"].value)
        self.assertEqual("=C1*2", str(ws["C3"].value).replace(" ", ""))
        self.assertIn("B5:D5", [str(r) for r in ws.merged_cells.ranges])
        self.assertEqual("gg-xlsx-fill-receipt/v1", receipt["schema"])

    def test_the_receipt_binds_every_input_and_the_output(self):
        receipt = self.run_fill([self.entry("C1")], [self.value("C1", 700)])
        self.assertEqual(self.fill.sha256(self.template),
                         receipt["template"]["sha256"])
        for part in ("map", "values", "output"):
            self.assertTrue(receipt[part]["path"])
        self.assertIs(True, receipt["recalcNeeded"])

    def test_a_mapped_cell_with_no_value_is_left_alone(self):
        self.run_fill([self.entry("C1"), self.entry("C2", role="crop")],
                      [self.value("C1", 700)])
        ws = self.openpyxl.load_workbook(self.out())[SHEET]
        self.assertEqual(700, ws["C1"].value)
        self.assertEqual("도라지", ws["C2"].value)

    # --- protecting the school form's own calculations --------------------

    def test_a_formula_cell_cannot_be_written(self):
        with self.assertRaisesRegex(ValueError, "map references formula cell"):
            self.run_fill([self.entry("C3")], [self.value("C3", 1)])

    def test_a_formula_cell_is_refused_even_with_nothing_written_to_it(self):
        with self.assertRaisesRegex(ValueError, "map references formula cell"):
            self.run_fill([self.entry("C1"), self.entry("C3")],
                          [self.value("C1", 700)])

    def test_a_merged_non_anchor_cannot_be_written(self):
        with self.assertRaisesRegex(ValueError,
                                    "map references merged non-anchor"):
            self.run_fill([self.entry("C5")], [self.value("C5", 1)])

    def test_a_cell_outside_the_write_map_is_refused(self):
        with self.assertRaisesRegex(ValueError, "outside write-map"):
            self.run_fill([self.entry("C1")], [self.value("C2", 1)])

    def test_a_sheet_that_does_not_exist_is_refused(self):
        with self.assertRaisesRegex(ValueError, "missing sheet"):
            self.run_fill([self.entry("C1", sheet="없는 시트")], [])

    def test_an_entry_must_be_marked_editable_explicitly(self):
        for editable in (None, False, "true", 1, 0):
            with self.subTest(editable=editable):
                with self.assertRaisesRegex(
                        ValueError, "editable must be explicitly true"):
                    self.run_fill([self.entry("C1", editable=editable)], [])

    def test_an_entry_without_an_editable_key_is_refused(self):
        bare = {k: v for k, v in self.entry("C1").items()
                if k != "editable"}
        with self.assertRaisesRegex(ValueError,
                                    "editable must be explicitly true"):
            self.run_fill([bare], [])

    def test_overlapping_map_entries_are_refused(self):
        with self.assertRaisesRegex(ValueError, "duplicate or overlapping"):
            self.run_fill([self.entry("C1"), self.entry("C1", role="other")],
                          [])

    def test_two_writes_to_one_cell_are_refused(self):
        with self.assertRaisesRegex(ValueError, "duplicate value write"):
            self.run_fill([self.entry("C1")],
                          [self.value("C1", 1), self.value("C1", 2)])

    # --- the map must belong to this exact file ---------------------------

    def test_a_template_changed_since_review_is_refused(self):
        m, v = self.files([self.entry("C1")], [self.value("C1", 700)])
        wb = self.openpyxl.load_workbook(self.template)
        wb[SHEET]["C1"] = 999
        wb.save(self.template)
        with self.assertRaisesRegex(RuntimeError, "template version changed"):
            self.fill.fill_copy(self.template, m, v, self.tmp / "out.xlsx",
                                context=self.ctx())

    def test_an_existing_output_is_never_overwritten(self):
        self.run_fill([self.entry("C1")], [self.value("C1", 700)])
        with self.assertRaises(FileExistsError):
            self.run_fill([self.entry("C1")], [self.value("C1", 800)])

    # --- the honesty rule, enforced at the cell ---------------------------

    def test_an_unanswered_question_cannot_become_a_number(self):
        for state in ("unknown", "not_provided", "explicit_none", "withheld",
                      "not_applicable"):
            with self.subTest(state=state):
                with self.assertRaisesRegex(ValueError,
                                            "unresolved/none answers"):
                    self.run_fill([self.entry("C1")],
                                  [self.value("C1", 0, answer_state=state)])

    def test_an_unanswered_question_may_still_be_left_blank(self):
        self.run_fill([self.entry("C1")], [self.value(
            "C1", None, value_type="blank", answer_state="unknown")])
        ws = self.openpyxl.load_workbook(self.out())[SHEET]
        self.assertIsNone(ws["C1"].value)

    def test_every_written_value_needs_full_evidence(self):
        for missing in ("source_id", "revision", "locator", "origin"):
            with self.subTest(missing=missing):
                ev = {"source_id": "s", "revision": 1, "locator": "L1",
                      "origin": "factual"}
                ev.pop(missing)
                with self.assertRaisesRegex(ValueError, "evidence_ref"):
                    self.run_fill([self.entry("C1")],
                                  [self.value("C1", 700, evidence_ref=ev)])

    def test_an_unknown_evidence_origin_is_refused(self):
        ev = {"source_id": "s", "revision": 1, "locator": "L1",
              "origin": "made_up"}
        with self.assertRaisesRegex(ValueError, "origin must be one of"):
            self.run_fill([self.entry("C1")],
                          [self.value("C1", 700, evidence_ref=ev)])

    def test_a_value_must_match_its_declared_type(self):
        for raw, vtype in (("육백", "integer"), (True, "integer"),
                           (600, "string"), (float("inf"), "number")):
            with self.subTest(raw=raw, vtype=vtype):
                with self.assertRaises(ValueError):
                    self.run_fill([self.entry("C1")],
                                  [self.value("C1", raw, value_type=vtype)])

    # --- the copy is a real workbook, not just a patched zip --------------

    def test_the_output_is_a_readable_xlsx(self):
        self.run_fill([self.entry("C1")], [self.value("C1", 700)])
        with zipfile.ZipFile(self.out()) as z:
            self.assertIsNone(z.testzip())
            self.assertIn("xl/workbook.xml", z.namelist())
        self.assertEqual(
            700, self.openpyxl.load_workbook(self.out())[SHEET]["C1"].value)


if __name__ == "__main__":
    unittest.main()
