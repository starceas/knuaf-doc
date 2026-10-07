"""D8 synthetic counterexamples; no student or professor values are bundled."""
import copy
import json
import tempfile
from pathlib import Path

from tests._harness import ContractCase, bind_major, runtime

S8 = '8. 노무비계획'
S5 = ' 5. 판매계획'
S11 = '11. 생산원가계획'
S9 = '9 .경비계획'


def fixture(path, broken=True):
    from openpyxl import Workbook
    w = Workbook()
    inv = w.active
    inv.title = '3.투자계획'
    inv['C19'] = '입력'
    inv['C24'] = '계'
    dep = w.create_sheet('10. 감가상각비계획 ')
    dep['H12'] = '=100/10'
    loan = w.create_sheet('4. 원리금상환계획')
    loan['C11'] = '=100*0.03'
    s = w.create_sheet(S8)
    s['F6'] = '2026년'
    for row in range(26, 32):
        s[f'X{row}'] = 2018 if row == 26 else f'=X{row-1}+1'
        s[f'Y{row}'] = 100 * 1.05 ** (row-26)
        s[f'Z{row}'] = 80 * 1.04 ** (row-26)
    s['Z28'] = '=Z29'
    for row in range(32, 39):
        s[f'X{row}'] = f'=X{row-1}+1'
    for col, rate in [('Y', 'AC38'), ('Z', 'AD38')]:
        for row in range(32, 39):
            s[f'{col}{row}'] = f'={col}{row-1}*(1+${rate[:2]}$38)'
    s['AC38'] = 0.05
    s['AD38'] = 0.04
    for row, base, rate in [(8, 'Y34', 'AC38'), (11, 'Y34', 'AC38'), (12, 'Z34', 'AD38')]:
        for i, col in enumerate(['F', 'I', 'L', 'O', 'R']):
            prior = ['F', 'I', 'L', 'O', 'R'][i-1]
            s[f'{col}{row}'] = f'={base}' if not i else f'={prior}{row}*(1+${rate[:2]}$38)'
            if broken and row == 11:
                s[f'{col}{row}'] = '=130' if not i else f'={prior}{row}'
            if broken and row == 8:
                s[f'{col}{row}'] = f'={col}9'
            s.cell(row, 7+3*i, 10)
            s.cell(row, 8+3*i, f'={col}{row}*{s.cell(row,7+3*i).coordinate}/1000')
    s.merge_cells('B2:D2')
    s['B2'] = 'Synthetic wage fixture'
    s.print_area = 'B2:T14'
    cost = w.create_sheet(S11)
    for col, src in zip('CDEFG', ['H', 'K', 'N', 'Q', 'T']):
        cost[f'{col}12'] = f"='{S8}'!{src}8"
        cost[f'{col}15'] = f"='{S8}'!{src}11+'{S8}'!{src}12"
        cost[f'{col}11'] = f'={col}12+{col}15'
    sale = w.create_sheet(S5)
    sale['C36'] = '2026년'
    sale['B46'] = '* 판매가격 물가상승률 2.5% 적용'
    for row, price in {38: 'O7', 39: 'O8', 40: 'O9', 41: 'O11', 42: 'O12', 43: 'O13'}.items():
        sale[price] = 20
        for q, r in zip('CEGIK', 'DFHJL'):
            sale[f'{q}{row}'] = 100
            sale[f'{r}{row}'] = f'={q}{row}*${price[0]}${price[1:]}/1000'
    exp = w.create_sheet(S9)
    exp['C18'] = 100
    exp['C38'] = '=C18*1.025'
    w.save(path)
    w.close()


def rate(r, sid, **extra):
    return {'base_year': 2020, 'observation_years': [2020, 2025],
            'observed_values': [100, 100*(1+r)**5], 'application_base_year': 2026,
            'evidence_ref': {'source_id': sid, 'revision': 1,
                             'locator': 'synthetic official-series contract fixture', 'origin': 'factual'}, **extra}


def student_series(sid, observations, *, role='wage', unit='원/일', base_year=None, dimensions=None):
    """Synthetic declared official metadata; never bundled as actual statistics."""
    source = copy.deepcopy(runtime('gg_price_assumptions').load_registry()[
        'official.rda.crop_income.spring_potato.output_price'])
    source.update(id=sid, agency_kind='central_government',
                  statistic_name='합성 회귀 수열', item_name='합성 회귀 항목',
                  allowed_roles=[role], unit=unit, base_year=base_year,
                  observations=observations, unconfirmed_years=[],
                  dimensions=dimensions or {'region': '전국', 'unit': unit},
                  evidence_locator='synthetic test declaration, not actual official values')
    return source


def registered_rate(sid, *, application_base_year=2026):
    source = runtime('gg_price_assumptions').load_registry()[sid]
    return {'base_year': source['base_year'],
            'observation_years': [o['year'] for o in source['observations']],
            'observed_values': [o['value'] for o in source['observations']],
            'application_base_year': application_base_year,
            'evidence_ref': {'source_id': sid, 'revision': 1, 'origin': 'factual',
                             'locator': source['evidence_locator']}}


def assumptions(sales_r=None):
    d = {'sales': {'mode': 'not_applied', 'reason': '확인 불가'}, 'price_sources': []}
    for kind, r in [('wage_male', .06), ('wage_female', .04)]:
        sid = 'official.rda.crop_income.synthetic.' + kind
        d[kind] = rate(r, sid, base_year=None)
        d['price_sources'].append(student_series(sid, [
            {'year': y, 'value': v} for y, v in zip(d[kind]['observation_years'], d[kind]['observed_values'])]))
    if sales_r is not None:
        dims = {'crop': 'synthetic', 'cultivation': 'open_field', 'region': 'national',
                'unit': '원/kg', 'product': 'fresh', 'grade': 'same', 'channel': 'producer'}
        sid = 'official.rda.crop_income.synthetic.output_price'
        d['sales'] = rate(sales_r, sid, base_year=None, dimensions=dims, target_dimensions=dims)
        d['price_sources'].append(student_series(sid, [
            {'year': y, 'value': v} for y, v in zip(d['sales']['observation_years'], d['sales']['observed_values'])],
            role='sales', unit='원/kg', dimensions=dims))
    return d


class WageLinkTests(ContractCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.src = self.root / 'src.xlsx'
        fixture(self.src)
        self.tpl = runtime('gg_excel_template')
        self.patch = runtime('gg_excel_formula_patch')
        self.fill = runtime('gg_excel_fill')
        self.audit = runtime('gg_workbook_audit')
        project = self.make_project()
        bind_major(project)
        self.context = runtime('gg_major_contract').output_context(project, 'specialty_crops')

    def json(self, name, value):
        p = self.root / name
        p.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
        return p

    def apply(self, data, name='patched.xlsx', source=None):
        out = self.root / name
        rec = self.patch.patch_copy(source or self.src, self.json(name+'.map.json', data), out, context=self.context)
        return out, rec

    def test_constant_chain_restored_source_hash_and_cost_links_preserved(self):
        before = self.tpl.sha256(self.src)
        out, rec = self.apply(self.patch.materialize_d8(self.src))
        self.assertEqual(self.tpl.sha256(self.src), before)
        self.assertEqual(len(rec['patched']), 10)
        b = self.audit.Workbook.load(out)
        costs = [b.evaluate(S11, col+'15') for col in 'CDEFG']
        self.assertTrue(all(y > x for x, y in zip(costs, costs[1:])), costs)
        original = self.audit.Workbook.load(self.src)
        for k, f in original.formulas.items():
            if k[0] in [S11, '10. 감가상각비계획 ', '4. 원리금상환계획']:
                self.assertEqual(b.formulas[k].text, f.text)
        self.assertEqual(b.values[S8, 'F6'], original.values[S8, 'F6'])
        import zipfile
        from xml.etree import ElementTree as E
        with zipfile.ZipFile(self.src) as a, zipfile.ZipFile(out) as z:
            for _, target in self.tpl.workbook_sheets(a):
                x, y = E.fromstring(a.read(target)), E.fromstring(z.read(target))
                for tag in ['mergeCells', 'pageMargins', 'pageSetup', 'printOptions']:
                    x1, y1 = x.find(self.tpl.local(tag)), y.find(self.tpl.local(tag))
                    self.assertEqual(None if x1 is None else E.tostring(x1), None if y1 is None else E.tostring(y1))

    def test_normal_x01_and_second_patch_are_byte_identical(self):
        fixture(self.src, broken=False)
        plan = self.patch.materialize_d8(self.src)
        self.assertEqual(plan['patches'], [])
        out, rec = self.apply(plan)
        self.assertEqual(self.src.read_bytes(), out.read_bytes())
        self.assertFalse(rec['recalcNeeded'])
        self.assertEqual(self.patch.materialize_d8(out)['patches'], [])

    def test_disconnected_numeric_daily_wage_supported(self):
        from openpyxl import load_workbook
        w = load_workbook(self.src)
        w[S8]['F11'] = 130
        w.save(self.src)
        w.close()
        plan = self.patch.materialize_d8(self.src)
        self.assertIn('expected_value', next(p for p in plan['patches'] if p['cell'] == 'F11'))
        out, _ = self.apply(plan)
        self.assertEqual(self.audit.Workbook.load(out).formulas[S8, 'F11'].text, 'Y34')

    def test_distinct_rates_no_double_application_and_update_idempotency(self):
        config = assumptions(.02)
        config['general'] = registered_rate('official.kosis.cpi.total')
        out, rec = self.apply(self.patch.materialize_d8(self.src, config))
        b = self.audit.Workbook.load(out)
        self.assertAlmostEqual(b.evaluate(S8, 'I11') / b.evaluate(S8, 'F11'), 1.06)
        self.assertAlmostEqual(b.evaluate(S8, 'I12') / b.evaluate(S8, 'F12'), 1.04)
        self.assertAlmostEqual(b.evaluate(S5, 'L38') / b.evaluate(S5, 'D38'), 1.02**4)
        r = (116.61/100)**(1/5)-1
        self.assertAlmostEqual(b.evaluate(S9, 'C38'), 100*(1+r))
        self.assertIn('2.00000000%', b.values[S5, 'B46'])
        self.assertIn('출처:', b.values[S5, 'B46'])
        self.assertIn('지수 기준연도 해당 없음', b.values[S5, 'B46'])
        self.assertNotIn('None', b.values[S5, 'B46'])
        self.assertEqual(rec['d8']['rates']['wage_male']['daily_base_application_count'], 3)
        self.assertEqual(self.patch.materialize_d8(out, config)['patches'], [])
        config['general'] = registered_rate('official.kosis.farm_purchase.expenses')
        changed, _ = self.apply(self.patch.materialize_d8(out, config), 'updated.xlsx', out)
        self.assertAlmostEqual(self.audit.Workbook.load(changed).evaluate(S9, 'C38'), 100*(130.2/100)**(1/5))

    def test_negative_rate_allowed(self):
        out, _ = self.apply(self.patch.materialize_d8(self.src, assumptions(-.02)))
        b = self.audit.Workbook.load(out)
        self.assertAlmostEqual(b.evaluate(S5, 'L38')/b.evaluate(S5, 'D38'), .98**4)

    def test_sales_not_applied_is_explicit(self):
        out, rec = self.apply(self.patch.materialize_d8(self.src, assumptions()))
        b = self.audit.Workbook.load(out)
        self.assertEqual(b.evaluate(S5, 'D38'), b.evaluate(S5, 'L38'))
        self.assertIn('미반영', rec['d8']['rates']['sales']['reason'])
        self.assertEqual(b.values[S5, 'B46'], '판매가 상승 미반영(확인 불가)')
        c = assumptions(); del c['sales']
        with self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)

    def test_unsourced_invalid_rates_and_dimensions_rejected(self):
        cases = []
        c = assumptions(.02); del c['wage_male']['evidence_ref']; cases.append(c)
        c = assumptions(.02); c['wage_male']['evidence_ref']['source_id'] = 'blog.wage'; cases.append(c)
        c = assumptions(.02); c['wage_male']['evidence_ref']['origin'] = 'assumption'; cases.append(c)
        c = assumptions(.02); c['wage_male']['evidence_ref']['source_id'] = 'official.rda.crop_income.synthetic.output_price'; cases.append(c)
        c = assumptions(.02); c['wage_male']['observed_values'][0] = 0; cases.append(c)
        c = assumptions(.02); c['wage_male']['r'] = .25; cases.append(c)
        c = assumptions(.02); c['sales']['target_dimensions'] = dict(c['sales']['dimensions'], crop='other'); cases.append(c)
        c = assumptions(.02); c['sales']['evidence_ref']['source_id'] = 'official.kosis.farm_sales.total'; cases.append(c)
        c = assumptions(.02); c['wage_female']['application_base_year'] = 2027; cases.append(c)
        for i, c in enumerate(cases):
            with self.subTest(case=i), self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)

    def test_patch_tampering_and_source_drift_fail_closed(self):
        plan = self.patch.materialize_d8(self.src)
        plan['patches'][0]['new_formula'] = '=1'
        with self.assertRaisesRegex(ValueError, 'differs'): self.apply(plan)
        plan = self.patch.materialize_d8(self.src)
        self.src.write_bytes(self.src.read_bytes()+b'changed')
        with self.assertRaisesRegex(RuntimeError, 'source version changed'): self.apply(plan)

    def test_h01_and_x02_not_given_x01_coordinates(self):
        from openpyxl import load_workbook
        w = load_workbook(self.src)
        w['3.투자계획']['C24'] = None
        w['3.투자계획']['C25'] = '계'
        w['10. 감가상각비계획 ']['J7'] = '대식물'
        w.save(self.src); w.close()
        with self.assertRaisesRegex(ValueError, 'layout_variant_unsupported'): self.patch.materialize_d8(self.src)
        with self.assertRaisesRegex(ValueError, 'layout_variant_unsupported'): self.tpl.wage_fill_map(self.src)
        from openpyxl import Workbook
        w = Workbook()
        w.active.title = '4.투자계획'
        w.create_sheet('11. 인건비계획')['F11'] = 130
        w.save(self.src); w.close()
        with self.assertRaisesRegex(ValueError, 'layout_variant_unsupported'): self.patch.materialize_d8(self.src)

    def observations(self):
        sid = 'official.rda.crop_income.synthetic.daily_wage'
        vals = []
        for row, year in zip(range(26, 32), range(2018, 2024)):
            for col, v in [('X', year), ('Y', 100+row), ('Z', 100+row)]:
                vals.append({'sheet': S8, 'cell': f'{col}{row}', 'value': v,
                             'value_type': 'integer', 'answer_state': 'provided',
                             'evidence_ref': {'source_id': sid, 'revision': 1, 'origin': 'factual', 'locator': f'synthetic year {year}, daily wage'}})
        return {'schema': self.fill.VALUES_SCHEMA, 'values': vals,
                'price_sources': [student_series(sid, [{'year': year, 'value': 100+row}
                                    for row, year in zip(range(26, 32), range(2018, 2024))])],
                'wage_observations': {'statistic_id': sid, 'base_year': None, 'unit': '원/일',
                                      'observation_years': list(range(2018, 2024)), 'application_base_year': 2026}}

    def test_formula_write_refusal_before_observation_preparation(self):
        m = self.json('fill.json', self.tpl.wage_fill_map(self.src))
        v = self.json('values.json', self.observations())
        with self.assertRaisesRegex(ValueError, 'formula cell'):
            self.fill.fill_copy(self.src, m, v, self.root/'filled.xlsx', context=self.context)
        self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v3_02_gender_specific_single_source_cannot_fill_both_columns(self):
        index = 0
        for key in ('gender', 'sex', '성별', ' gender ', ' sex ', ' 성별 ', 'Gender', 'SEX'):
            for gender in ('남자', '여자', ['남자'], ['여자'], ['남자', '여자']):
                index += 1
                name = f'gender-{index}.xlsx'
                v = self.observations()
                v['price_sources'][0]['dimensions'][key] = gender
                with self.subTest(key=key, gender=gender), self.assertRaises(ValueError):
                    self.prepared_fill(v, name)
                self.assertFalse((self.root/name).exists())
        self.prepared_fill(self.observations())
        self.assertTrue((self.root/'filled.xlsx').exists())

    def test_v3_02_gender_rejection_preserves_hourly_write_plan(self):
        v = self.hourly_observations()
        v['wage_observations']['unit_conversion'] = self.conversion()
        source = copy.deepcopy(runtime('gg_price_assumptions').load_registry()['official.minimumwage.hourly'])
        source.update(id='synthetic.gender', dimensions={'region': '전국', 'sex': '남자'})
        v['price_sources'] = [source]
        v['wage_observations']['statistic_id'] = source['id']
        for value in v['values']:
            value['evidence_ref']['source_id'] = source['id']
        before = copy.deepcopy(v)
        cells = {(value['sheet'], value['cell']): value for value in v['values']}
        cell_before = copy.deepcopy(cells)
        with self.assertRaisesRegex(ValueError, 'gender'):
            self.fill._wage_observations(dict.fromkeys(cells), cells, v)
        self.assertEqual(cells, cell_before)
        self.assertEqual(v, before)
        source_digest = self.tpl.sha256(self.src)
        with self.assertRaisesRegex(ValueError, 'gender'):
            self.prepared_fill(v)
        self.assertFalse((self.root/'filled.xlsx').exists())
        self.assertEqual(v, before)
        v['price_sources'][0]['dimensions'].pop('sex')
        self.prepared_fill(v)
        self.assertTrue((self.root/'filled.xlsx').exists())
        self.assertEqual(self.tpl.sha256(self.src), source_digest)

    def test_v3_02_rate_gender_must_match_its_column_and_aliases(self):
        for key in ('gender', 'sex', '성별'):
            c = assumptions()
            for kind, gender in (('wage_male', '남자'), ('wage_female', '여자')):
                source = next(s for s in c['price_sources'] if s['id'] == c[kind]['evidence_ref']['source_id'])
                source['dimensions'][key] = gender
            self.patch.materialize_d8(self.src, c)
            for bad in ('opposite', 'conflicting_alias', 'multi_gender'):
                mutated = copy.deepcopy(c)
                dims = mutated['price_sources'][0]['dimensions']
                if bad == 'opposite': dims[key] = '여자'
                elif bad == 'conflicting_alias': dims['sex' if key != 'sex' else 'gender'] = '여자'
                else: dims[key] = ['남자', '여자']
                with self.subTest(key=key, bad=bad), self.assertRaises(ValueError):
                    self.patch.materialize_d8(self.src, mutated)

    def adapter_variants(self):
        scalar = assumptions(.02)
        listed = copy.deepcopy(scalar)
        listed['sales']['target_dimensions'] = copy.deepcopy(listed['sales']['target_dimensions'])
        listed['sales']['target_dimensions']['crop'] = ['synthetic']
        alias = copy.deepcopy(scalar)
        alias['sales']['target_dimensions'] = copy.deepcopy(alias['sales']['target_dimensions'])
        alias['sales']['target_dimensions']['작목'] = alias['sales']['target_dimensions'].pop('crop')
        region = copy.deepcopy(scalar)
        region['wage_male']['target_dimensions'] = {'region': '전국'}
        region_alias = copy.deepcopy(scalar)
        region_alias['wage_male']['target_dimensions'] = {'지역': '전국'}
        return [scalar, listed, alias, region, region_alias]

    def test_v3_03_normal_dimension_variants_materialize_and_cli(self):
        import contextlib
        import io
        for i, c in enumerate(self.adapter_variants()):
            with self.subTest(variant=i):
                before = copy.deepcopy(c)
                plan = self.patch.materialize_d8(self.src, c)
                self.assertEqual(c, before)
                self.assertAlmostEqual(plan['d8']['rates']['sales']['r'], .02)
                self.assertAlmostEqual(plan['d8']['rates']['wage_male']['r'], .06)
                out = self.root / f'cli-{i}.json'
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    code = self.patch.cli(['materialize-d8', '--source', str(self.src),
                                           '--assumptions', str(self.json(f'config-{i}.json', c)), '--out', str(out)])
                self.assertEqual(code, 0)
                self.assertEqual(json.loads(out.read_text()), plan)

    def test_v3_03_normalization_still_rejects_missing_conflicting_or_invalid_crop(self):
        for bad in ({'작목': 'other'}, {'crop': 'synthetic', '작목': 'other'},
                    {'crop': [['synthetic']]}, {'crop': []}, {'crop': None}, {}):
            c = assumptions(.02)
            c['sales']['target_dimensions'] = copy.deepcopy(c['sales']['target_dimensions'])
            dims = c['sales']['target_dimensions']
            dims.pop('crop')
            dims.update(bad)
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                self.patch.materialize_d8(self.src, c)

    def x01_observation_fixture(self):
        # Only the SHA admission is simulated here; the actual registered X01
        # is replayed separately, outside this portable synthetic suite.
        from openpyxl import load_workbook
        w = load_workbook(self.src)
        w[S8]['Z28'] = 82
        w[S8]['Z27'] = '=Z28'
        w.save(self.src)
        w.close()
        return '5d19b6a2e5dcddccb70c68e39abee8f423bba005c008db7b64076326a17b9c80'

    def test_v3_04_exact_x01_observation_formula_is_prepared_without_numeric_writes(self):
        from unittest.mock import patch
        digest = self.x01_observation_fixture()
        real_sha = self.patch.sha256
        with patch.object(self.patch, 'sha256', side_effect=lambda p: digest if p == self.src else real_sha(p)):
            plan = self.patch.materialize_d8(self.src, prepare_observations=True)
        clearing = [p for p in plan['patches'] if p.get('operation') == 'prepare_wage_input']
        self.assertEqual({p['cell'] for p in clearing}, {'X27', 'X28', 'X29', 'X30', 'X31', 'Z27'})
        target = next(p for p in clearing if p['cell'] == 'Z27')
        self.assertEqual(target['expected_formula'], '=Z28')
        self.assertIsNone(target['new_value'])
        self.assertIn('source_sha256', target['evidence_ref']['locator'])
        self.assertFalse(any(p['cell'] == 'Z28' for p in clearing))

    def test_v3_04_x01_exception_rejects_unregistered_sha_formula_and_audit_drift(self):
        from unittest.mock import patch
        digest = self.x01_observation_fixture()
        with self.assertRaisesRegex(ValueError, 'unreviewed observation'):
            self.patch.materialize_d8(self.src, prepare_observations=True)
        real_sha = self.patch.sha256
        audit_path = self.patch.D8_SPEC_PATH.parent.parent / 'formula-audit/x01.json'
        with patch.object(self.patch, 'sha256', side_effect=lambda p: digest if p == self.src else '0'*64 if p == audit_path else real_sha(p)):
            with self.assertRaisesRegex(ValueError, 'audit'):
                self.patch.materialize_d8(self.src, prepare_observations=True)
        real_json = self.patch._json
        for bad in ('source_sha256', 'inventory'):
            audit = real_json(audit_path)
            if bad == 'source_sha256': audit[bad] = '0'*64
            else: audit['inventory'] = [r for r in audit['inventory'] if (r['sheet'], r['cell']) != (S8, 'Z27')]
            with patch.object(self.patch, 'sha256', side_effect=lambda p: digest if p == self.src else real_sha(p)):
                with patch.object(self.patch, '_json', side_effect=lambda p: audit if p == audit_path else real_json(p)):
                    with self.subTest(audit=bad), self.assertRaisesRegex(ValueError, 'audit'):
                        self.patch.materialize_d8(self.src, prepare_observations=True)
        from openpyxl import load_workbook
        w = load_workbook(self.src)
        w[S8]['Z27'] = '=Z29'
        w.save(self.src)
        w.close()
        with patch.object(self.patch, 'sha256', side_effect=lambda p: digest if p == self.src else real_sha(p)):
            with self.assertRaisesRegex(ValueError, 'unreviewed observation'):
                self.patch.materialize_d8(self.src, prepare_observations=True)

    def test_observation_preparation_fill_and_receipt(self):
        before = self.tpl.sha256(self.src)
        out, rec = self.apply(self.patch.materialize_d8(self.src, prepare_observations=True))
        self.assertEqual(sum(p['operation']=='prepare_wage_input' for p in rec['patched']), 6)
        receipt = self.fill.fill_copy(out, self.json('fill.json', self.tpl.wage_fill_map(out)),
                                     self.json('values.json', self.observations()), self.root/'filled.xlsx', context=self.context)
        self.assertEqual(receipt['wage_observations']['observation_range'], [2018, 2023])
        self.assertEqual(receipt['wage_observations']['application_count'], 3)
        self.assertEqual(receipt['wage_observations']['unit'], '원/일')
        self.assertEqual(self.tpl.sha256(self.src), before)
        b = self.audit.Workbook.load(self.root/'filled.xlsx')
        self.assertAlmostEqual(b.evaluate(S8, 'AC38'), (131/126)**(1/5)-1)
        self.assertAlmostEqual(b.evaluate(S8, 'AD38'), (131/126)**(1/5)-1)

    def test_missing_wage_evidence_partial_or_invalid_observations_rejected(self):
        out, _ = self.apply(self.patch.materialize_d8(self.src, prepare_observations=True))
        m = self.json('fill.json', self.tpl.wage_fill_map(out))
        cases = []
        v = self.observations(); del v['values'][1]['evidence_ref']; cases.append(v)
        v = self.observations(); v['values'][1]['evidence_ref']['source_id'] = 'news.wages'; cases.append(v)
        v = self.observations(); v['values'].pop(); cases.append(v)
        v = self.observations(); v['values'][1]['value'] = -10; cases.append(v)
        v = self.observations(); v['values'][0]['value'] = 2020; cases.append(v)
        v = self.observations(); v['wage_observations']['observation_years'][1] = 2018; cases.append(v)
        for v in cases:
            with self.subTest(v=v), self.assertRaises(ValueError):
                self.fill.fill_copy(out, m, self.json('bad.json', v), self.root/'filled.xlsx', context=self.context)
        self.assertFalse((self.root/'filled.xlsx').exists())

    def test_default_map_lists_observations_but_protects_shared_followers(self):
        info = self.tpl.inspect_source(self.src, None, None)
        listed = {e.get('range') for e in info['entries'] if e.get('semanticField') == 'wage_observation'}
        self.assertIn('X26', listed)
        self.assertIn('Y26', listed)
        self.assertNotIn('X27', listed)
        self.assertNotIn('Z28', listed)

    def test_sales_uses_its_own_plan_year_and_preserves_w023(self):
        from openpyxl import load_workbook
        w = load_workbook(self.src)
        w[S8]['F6'] = '2027년'
        w.save(self.src); w.close()
        c = assumptions(.02)
        c['wage_male']['application_base_year'] = 2027
        c['wage_female']['application_base_year'] = 2027
        out, _ = self.apply(self.patch.materialize_d8(self.src, c))
        b = self.audit.Workbook.load(out)
        self.assertEqual(b.values[S8, 'F6'], '2027년')
        self.assertEqual(b.values[S5, 'C36'], '2026년')

    def test_unreviewed_wage_or_observation_formula_refused(self):
        from openpyxl import load_workbook
        w = load_workbook(self.src); w[S8]['I11'] = '=F11*2'; w.save(self.src); w.close()
        with self.assertRaisesRegex(ValueError, 'unreviewed wage formula'): self.patch.materialize_d8(self.src)
        fixture(self.src)
        w = load_workbook(self.src); w[S8]['Z28'] = '=Z29*2'; w.save(self.src); w.close()
        with self.assertRaisesRegex(ValueError, 'unreviewed observation'): self.patch.materialize_d8(self.src, prepare_observations=True)

    def test_generic_map_cannot_clear_formulas_as_observations(self):
        plan = self.patch.materialize_d8(self.src, prepare_observations=True)
        del plan['d8']
        with self.assertRaisesRegex(ValueError, 'preparation'): self.apply(plan)

    def test_empty_observation_years_and_combined_preparation_rates_refused(self):
        with self.assertRaisesRegex(ValueError, 'then apply'):
            self.patch.materialize_d8(self.src, assumptions(), prepare_observations=True)
        out, _ = self.apply(self.patch.materialize_d8(self.src, prepare_observations=True))
        with self.assertRaisesRegex(ValueError, 'end year'):
            self.patch.materialize_d8(out, assumptions())

    def test_shared_year_followers_are_protected_then_explicitly_prepared(self):
        import zipfile
        from xml.etree import ElementTree as E
        with zipfile.ZipFile(self.src) as z:
            parts = {n: z.read(n) for n in z.namelist()}
            target = dict(self.tpl.workbook_sheets(z))[S8]
        root = E.fromstring(parts[target])
        for row in range(27, 32):
            f = root.find(f".//{self.tpl.local('c')}[@r='X{row}']/{self.tpl.local('f')}")
            f.set('t', 'shared'); f.set('si', '0')
            if row == 27: f.set('ref', 'X27:X31')
            else: f.text = None
        parts[target] = E.tostring(root)
        with zipfile.ZipFile(self.src, 'w') as z:
            for n, raw in parts.items(): z.writestr(n, raw)
        info = self.tpl.inspect_source(self.src, None, None)
        self.assertFalse(any(e.get('semanticField') == 'wage_observation' and e.get('range') == 'X28' for e in info['entries']))
        out, rec = self.apply(self.patch.materialize_d8(self.src, prepare_observations=True))
        self.assertEqual(len(rec['expanded_shared_formulas']), 5)
        b = self.audit.Workbook.load(out)
        self.assertNotIn((S8, 'X28'), b.formulas)
        self.assertIn((S8, 'X32'), b.formulas)

    def test_v1_05_registered_sales_hourly_and_index_cannot_fill_daily_wages(self):
        out, _ = self.apply(self.patch.materialize_d8(self.src, prepare_observations=True))
        m = self.json('fill.json', self.tpl.wage_fill_map(out))
        for sid, unit, base in [
            ('official.rda.crop_income.spring_potato.output_price', '원/kg', 2020),
            ('official.minimumwage.hourly', '원/시간', None),
            ('official.moel.establishment_wage.total', '천원/인·월', None),
            ('official.kosis.farm_purchase.labor', '지수(2020=100)', 2020),
        ]:
            v = self.observations()
            v['wage_observations'].update(statistic_id=sid, unit=unit, base_year=base)
            for value in v['values']:
                value['evidence_ref'].update(source_id=sid, locator='봄감자 주산물 단가 원/kg')
                if value['cell'][0] != 'X': value['value'] = 1190
            with self.subTest(sid=sid), self.assertRaises(ValueError):
                self.fill.fill_copy(out, m, self.json('bad.json', v), self.root/'rejected.xlsx', context=self.context)
            self.assertFalse((self.root/'rejected.xlsx').exists())

    def test_v1_06_wrong_index_base_and_unregistered_rate_rejected(self):
        for change in [
            {'base_year': 2015},
            {'base_year': None},
            {'evidence_ref': {'source_id': 'official.rda.crop_income.not_registered.daily_wage',
                              'revision': 1, 'origin': 'factual', 'locator': 'synthetic fixture'}},
        ]:
            c = assumptions()
            c['wage_male'] = registered_rate('official.kosis.farm_purchase.labor')
            c['wage_male'].update(change)
            with self.subTest(change=change), self.assertRaises(ValueError):
                self.patch.materialize_d8(self.src, c)

    def hourly_observations(self):
        sid = 'official.minimumwage.hourly'
        source = runtime('gg_price_assumptions').load_registry()[sid]
        v = self.observations()
        v['price_sources'] = []
        v['wage_observations'].update(statistic_id=sid, unit=source['unit'], base_year=None,
                                     observation_years=[o['year'] for o in source['observations']])
        for row, obs in zip(range(26, 32), source['observations']):
            for value in v['values']:
                if value['cell'][1:] == str(row):
                    value['value'] = obs['year'] if value['cell'][0] == 'X' else obs['value']
                value['evidence_ref']['source_id'] = sid
        return v

    def conversion(self, hours=6.5):
        return {'from': '원/시간', 'to': '원/일', 'hours_per_day': hours,
                'evidence_ref': {'source_id': 'official.rda.synthetic.day_definition',
                                 'revision': 1, 'origin': 'factual',
                                 'locator': 'synthetic explicit labor-day definition; not a default'}}

    def prepared_fill(self, v, name='filled.xlsx'):
        out = self.root/'prepared.xlsx'
        if not out.exists():
            out, _ = self.apply(self.patch.materialize_d8(self.src, prepare_observations=True), 'prepared.xlsx')
        return self.fill.fill_copy(out, self.json('fill.json', self.tpl.wage_fill_map(out)),
                                   self.json('values.json', v), self.root/name, context=self.context)

    def test_v2_02_tampered_hourly_values_and_unobserved_years_rejected(self):
        for case in ('tampered', 'unobserved'):
            v = self.hourly_observations()
            v['wage_observations']['unit_conversion'] = self.conversion()
            years = list(range(2030, 2036)) if case == 'unobserved' else list(range(2020, 2026))
            v['wage_observations'].update(observation_years=years, application_base_year=years[-1]+1)
            for row, year in zip(range(26, 32), years):
                for e in v['values']:
                    if e['cell'][1:] == str(row):
                        e['value'] = year if e['cell'][0] == 'X' else 1
            with self.subTest(case=case), self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
                self.prepared_fill(v, case+'.xlsx')
            self.assertFalse((self.root/(case+'.xlsx')).exists())

    def test_v2_02_each_raw_value_must_match_in_both_gender_columns(self):
        import math
        for row in range(26, 32):
            for col in 'YZ':
                v = self.hourly_observations()
                v['wage_observations']['unit_conversion'] = self.conversion()
                e = next(e for e in v['values'] if e['cell'] == f'{col}{row}')
                # Even a one-float-step difference is not a registered observation.
                e.update(value=math.nextafter(float(e['value']), math.inf), value_type='number')
                with self.subTest(cell=e['cell']), self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
                    self.prepared_fill(v)
                self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v2_02_registered_daily_series_cannot_have_different_female_values(self):
        v = self.observations()
        for e in v['values']:
            if e['cell'][0] == 'Z': e['value'] -= 20
        with self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
            self.prepared_fill(v)
        self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v2_02_bundled_id_cannot_be_shadowed_with_changed_registration(self):
        v = self.hourly_observations()
        v['wage_observations']['unit_conversion'] = self.conversion()
        entry = copy.deepcopy(runtime('gg_price_assumptions').load_registry()[v['wage_observations']['statistic_id']])
        for obs in entry['observations']: obs['value'] = 1
        v['price_sources'] = [entry]
        for e in v['values']:
            if e['cell'][0] in 'YZ': e['value'] = 1
        with self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
            self.prepared_fill(v)
        self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v2_02_conversion_cannot_hide_changed_or_preconverted_raw_values(self):
        for case, scale, hours in [('compensated', 2, 3.25), ('preconverted', 6.5, 6.5)]:
            v = self.hourly_observations()
            v['wage_observations']['unit_conversion'] = self.conversion(hours)
            for e in v['values']:
                if e['cell'][0] in 'YZ': e.update(value=e['value']*scale, value_type='number')
            with self.subTest(case=case), self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
                self.prepared_fill(v)
            self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v2_02_failure_does_not_mutate_or_convert_the_write_plan(self):
        v = self.hourly_observations()
        v['wage_observations']['unit_conversion'] = self.conversion()
        v['values'][-1]['value'] = 1
        cells = {(e['sheet'], e['cell']): e for e in v['values']}
        original = copy.deepcopy(cells)
        with self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
            self.fill._wage_observations(dict.fromkeys(cells), cells, v)
        self.assertEqual(cells, original)
        with self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
            self.prepared_fill(v)
        self.assertFalse((self.root/'filled.xlsx').exists())
        before = self.tpl.sha256(self.root/'prepared.xlsx')
        good = self.hourly_observations()
        good['wage_observations']['unit_conversion'] = self.conversion()
        self.prepared_fill(good)
        self.assertEqual(self.tpl.sha256(self.root/'prepared.xlsx'), before)
        self.assertEqual(self.audit.Workbook.load(self.root/'filled.xlsx').values[S8, 'Z31'], 10030*6.5)

    def test_v2_02_year_value_pairs_cannot_be_shifted_or_reordered(self):
        for case in ('swapped_values', 'missing_year', 'future_years'):
            v = self.hourly_observations()
            v['wage_observations']['unit_conversion'] = self.conversion()
            if case == 'swapped_values':
                for col in 'YZ':
                    a, b = [next(e for e in v['values'] if e['cell'] == f'{col}{row}') for row in (26, 27)]
                    a['value'], b['value'] = b['value'], a['value']
            else:
                years = [2019, 2021, 2022, 2023, 2024, 2025] if case == 'missing_year' else list(range(2030, 2036))
                v['wage_observations'].update(observation_years=years, application_base_year=years[-1]+1)
                for row, year in zip(range(26, 32), years):
                    next(e for e in v['values'] if e['cell'] == f'X{row}')['value'] = year
            with self.subTest(case=case), self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
                self.prepared_fill(v)
            self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v2_02_different_values_require_registration_under_another_id(self):
        for unit in ('원/시간', '원/일'):
            v = self.hourly_observations()
            sid = 'official.synthetic.alternative.' + ('hourly' if unit == '원/시간' else 'daily')
            years = list(range(2030, 2036))
            observations = [{'year': year, 'value': 1+i} for i, year in enumerate(years)]
            v['price_sources'] = [student_series(sid, observations, unit=unit)]
            v['wage_observations'].update(statistic_id=sid, unit=unit,
                                         observation_years=years, application_base_year=2036)
            if unit == '원/시간': v['wage_observations']['unit_conversion'] = self.conversion()
            for row, obs in zip(range(26, 32), observations):
                for e in v['values']:
                    if e['cell'][1:] == str(row):
                        e['value'] = obs['year'] if e['cell'][0] == 'X' else obs['value']
                    e['evidence_ref']['source_id'] = sid
            missing = copy.deepcopy(v); missing['price_sources'] = []
            with self.subTest(unit=unit), self.assertRaisesRegex(ValueError, '알 수 없는 출처 id'):
                self.prepared_fill(missing)
            self.assertFalse((self.root/'filled.xlsx').exists())
            original = copy.deepcopy(v)
            name = sid+'.xlsx'
            receipt = self.prepared_fill(v, name)
            b = self.audit.Workbook.load(self.root/name)
            for row, obs in zip(range(26, 32), observations):
                self.assertEqual(b.values[S8, f'X{row}'], obs['year'])
                for col in 'YZ':
                    self.assertEqual(b.values[S8, f'{col}{row}'], obs['value']*(6.5 if unit == '원/시간' else 1))
            self.assertEqual(receipt['wage_observations']['statistic_id'], sid)
            self.assertEqual(v, original)
            # A declared but unconfirmed year cannot supply an observation.
            v['price_sources'][0]['observations'].pop()
            v['price_sources'][0]['unconfirmed_years'] = [2035]
            with self.subTest(unconfirmed=unit), self.assertRaisesRegex(ValueError, '확정 관측과 불일치'):
                self.prepared_fill(v)
            self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v2_02_all_converted_values_follow_only_registered_raw_times_hours(self):
        for hours in (0.5, 9):  # Explicit synthetic factors, never defaults.
            v = self.hourly_observations()
            v['wage_observations']['unit_conversion'] = self.conversion(hours)
            before = copy.deepcopy(v)
            name = f'hours-{hours}.xlsx'
            receipt = self.prepared_fill(v, name)
            b = self.audit.Workbook.load(self.root/name)
            for row, year, raw in zip(range(26, 32), range(2020, 2026), (8590, 8720, 9160, 9620, 9860, 10030)):
                self.assertEqual(b.values[S8, f'X{row}'], year)
                for col in 'YZ':
                    self.assertEqual(b.values[S8, f'{col}{row}'], raw*hours)
                    written = next(e for e in receipt['written'] if e['cell'] == f'{col}{row}')
                    self.assertEqual(written['newValue'], raw*hours)
            self.assertEqual(v, before)

    def test_v1_05_explicit_hourly_conversion_changes_values_and_preserves_years(self):
        v = self.hourly_observations()
        v['wage_observations']['unit_conversion'] = self.conversion()
        unchanged = copy.deepcopy(v)
        rec = self.prepared_fill(v)
        b = self.audit.Workbook.load(self.root/'filled.xlsx')
        self.assertEqual(b.values[S8, 'Y26'], 8590*6.5)
        self.assertEqual(b.values[S8, 'Z31'], 10030*6.5)
        self.assertEqual(b.values[S8, 'X26'], 2020)
        self.assertEqual(b.values[S8, 'X31'], 2025)
        self.assertEqual(v, unchanged)
        self.assertEqual((self.root/'values.json').read_text(encoding='utf-8'),
                         json.dumps(unchanged, ensure_ascii=False))
        self.assertEqual(rec['wage_observations']['input_unit'], '원/시간')
        self.assertEqual(rec['wage_observations']['unit'], '원/일')
        self.assertEqual(rec['wage_observations']['unit_conversion'], self.conversion())
        self.assertIsNone(rec['wage_observations']['base_year'])
        self.assertEqual(rec['wage_observations']['application_count'], 1)
        written = next(e for e in rec['written'] if e['cell'] == 'Y26')
        self.assertEqual(written['newValue'], 8590*6.5)
        filled = self.root/'filled.xlsx'
        before = self.tpl.sha256(filled)
        self.fill.fill_copy(filled, self.json('again-map.json', self.tpl.wage_fill_map(filled)),
                            self.root/'values.json', self.root/'again.xlsx', context=self.context)
        again = self.audit.Workbook.load(self.root/'again.xlsx')
        self.assertEqual(again.values[S8, 'Y26'], 8590*6.5)
        self.assertEqual(again.values[S8, 'Z31'], 10030*6.5)
        self.assertEqual(self.tpl.sha256(filled), before)

    def test_v1_05_conversion_requires_every_field_and_finite_explicit_hours(self):
        invalid = [None, {}, [], True]
        for field in ['from', 'to', 'hours_per_day', 'evidence_ref']:
            c = self.conversion(); del c[field]; invalid.append(c)
        for hours in [0, -1, True, False, '8', None, 1e308, 10**400, float('nan'), float('inf'), -float('inf')]:
            invalid.append(self.conversion(hours))
        for field, val in [('from', ' 원/시간'), ('from', '원/일'), ('to', '원/시간'),
                           ('to', '원/일 '), ('from', []), ('to', True)]:
            c = self.conversion(); c[field] = val; invalid.append(c)
        for ev in [None, {}, {'source_id': 'official.x', 'revision': True, 'locator': 'x', 'origin': 'factual'},
                   {'source_id': 'official.x', 'revision': 1, 'locator': ' ', 'origin': 'factual'},
                   {'source_id': 'official.x', 'revision': 1, 'locator': 'x', 'origin': 'assumption'},
                   {'source_id': 'official.x', 'revision': 1, 'locator': 'x', 'origin': []}]:
            c = self.conversion(); c['evidence_ref'] = ev; invalid.append(c)
        for i, conversion in enumerate(invalid):
            v = self.hourly_observations()
            v['wage_observations']['unit_conversion'] = conversion
            with self.subTest(i=i), self.assertRaises(ValueError): self.prepared_fill(v)
            self.assertFalse((self.root/'filled.xlsx').exists())

    def test_v1_05_daily_metadata_unit_role_and_value_bypasses_rejected(self):
        cases = []
        for unit in [None, '原/일', ' 원/일', '원/일 ', '원/시간', '원/kg', [], True]:
            v = self.observations(); v['wage_observations']['unit'] = unit; cases.append(v)
        for source_id in ['official.rda.crop_income.not_registered.daily_wage',
                          'OFFICIAL.RDA.CROP_INCOME.SYNTHETIC.DAILY_WAGE',
                          ' official.rda.crop_income.spring_potato.output_price',
                          'official.rda.crop_income.spring_potato.output_price ',
                          'official.rda.crop_income.synthetic/../daily_wage', [], True]:
            v = self.observations(); v['wage_observations']['statistic_id'] = source_id
            for e in v['values']: e['evidence_ref']['source_id'] = source_id
            cases.append(v)
        for value in [0, -1, True, 10**400, float('nan'), float('inf'), -float('inf')]:
            v = self.observations(); v['values'][1].update(value=value, value_type='number'); cases.append(v)
        for role in [[], ['sales'], 'nonsales', None]:
            v = self.observations(); v['price_sources'][0]['allowed_roles'] = role; cases.append(v)
        v = self.observations(); v['wage_observations']['unit_conversion'] = self.conversion(); cases.append(v)
        v = self.observations(); v['values'][1]['unit'] = '원/kg'; cases.append(v)
        v = self.observations(); v['values'][1]['value_type'] = []; cases.append(v)
        v = self.observations(); v['values'][1]['answer_state'] = {}; cases.append(v)
        for i, v in enumerate(cases):
            name = f'rejected-{i}.xlsx'
            with self.subTest(case=i), self.assertRaises(ValueError): self.prepared_fill(v, name)
            self.assertFalse((self.root/name).exists())

    def test_v1_05_lowercase_and_absolute_cell_refs_still_validate_wage_role(self):
        for transform in [str.lower, lambda ref: '$'+ref[0]+'$'+ref[1:]]:
            v = self.observations()
            v['wage_observations'].update(statistic_id='official.rda.crop_income.spring_potato.output_price', unit='원/kg')
            for e in v['values']:
                e['cell'] = transform(e['cell'])
                e['evidence_ref']['source_id'] = v['wage_observations']['statistic_id']
            with self.subTest(transform=transform), self.assertRaises(ValueError): self.prepared_fill(v)

    def test_v1_06_registered_nominal_null_and_index_parity_with_legacy(self):
        pa = runtime('gg_price_assumptions')
        for sid in ['official.kosis.farm_purchase.labor', 'official.minimumwage.hourly']:
            item = registered_rate(sid)
            c = assumptions(); c['wage_male'] = item
            out, _ = self.apply(self.patch.materialize_d8(self.src, c), sid.split('.')[-1]+'.xlsx')
            normalized = pa.normalize_price_assumptions({
                'price_assumptions': {'wage': {'source_id': sid, 'base_year': item['base_year'], 'application_base_year': 2026},
                                     'general': {'status': 'not_applied', 'reason': 'test'},
                                     'sales': {'status': 'not_applied', 'reason': 'test'}}}, list(range(2026, 2031)))['wage']
            b = self.audit.Workbook.load(out)
            self.assertAlmostEqual(b.evaluate(S8, 'AC38'), normalized['rate'])
            self.assertAlmostEqual(b.evaluate(S8, 'I11')/b.evaluate(S8, 'F11'), 1+normalized['rate'])
            self.assertEqual(self.patch.materialize_d8(out, c)['patches'], [])

    def test_v1_06_registered_values_cannot_be_replaced_under_bundled_id(self):
        c = assumptions(); c['wage_male'] = registered_rate('official.kosis.farm_purchase.labor')
        self.patch.materialize_d8(self.src, c)
        c['wage_male']['observed_values'][2] += 1
        with self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)

    def test_v1_06_registration_and_null_basis_are_required_in_both_paths(self):
        c = assumptions(); c.pop('price_sources')
        with self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)
        for i, bad in enumerate([2020, True, '2020', float('nan'), float('inf')]):
            c = assumptions(); c['wage_male']['base_year'] = bad
            with self.subTest(base=bad), self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)
            v = self.observations(); v['wage_observations']['base_year'] = bad
            with self.subTest(fill_base=bad), self.assertRaises(ValueError): self.prepared_fill(v, f'base-{i}.xlsx')
        for key in ['base_year', 'unit']:
            v = self.observations(); v['wage_observations'].pop(key)
            with self.subTest(missing=key), self.assertRaises(ValueError): self.prepared_fill(v)

    def test_v1_06_invalid_registration_containers_and_rate_types_rejected(self):
        for sources in [{}, False, 'series', 1, [True]]:
            c = assumptions(); c['price_sources'] = sources
            with self.subTest(sources=sources), self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)
        for r in [True, '0.06', None, float('nan'), float('inf'), -float('inf'), .9, []]:
            c = assumptions(); c['wage_male']['r'] = r
            with self.subTest(r=r), self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)

    def test_shared_check_rate_is_only_a_check_actual_cagr_is_applied(self):
        c = assumptions()
        c['wage_male']['r'] = .061  # Within FIX-A's documented rounding tolerance.
        out, rec = self.apply(self.patch.materialize_d8(self.src, c))
        b = self.audit.Workbook.load(out)
        self.assertAlmostEqual(rec['d8']['rates']['wage_male']['r'], .06)
        self.assertAlmostEqual(b.evaluate(S8, 'I11')/b.evaluate(S8, 'F11'), 1.06)

    def test_id_whitespace_is_rejected_by_shared_resolution_in_both_paths(self):
        pa = runtime('gg_price_assumptions')
        for i, decorate in enumerate([lambda s: ' '+s, lambda s: s+' ']):
            sid = 'official.kosis.farm_purchase.labor'
            c = assumptions(); c['wage_male'] = registered_rate(sid)
            c['wage_male']['evidence_ref']['source_id'] = decorate(sid)
            with self.assertRaises(ValueError): self.patch.materialize_d8(self.src, c)
            with self.assertRaises(ValueError): pa.resolve_series(decorate(sid))
            with self.assertRaises(ValueError):
                pa.normalize_price_assumptions({'price_assumptions': {
                    'wage': {'source_id': decorate(sid), 'base_year': 2020, 'application_base_year': 2026},
                    'general': {'status': 'not_applied', 'reason': 'test'},
                    'sales': {'status': 'not_applied', 'reason': 'test'}}}, list(range(2026, 2031)))
            v = self.observations()
            canonical = v['wage_observations']['statistic_id']
            v['wage_observations']['statistic_id'] = decorate(canonical)
            for e in v['values']: e['evidence_ref']['source_id'] = decorate(canonical)
            with self.assertRaises(ValueError): self.prepared_fill(v, f'filled-alias-{i}.xlsx')
            self.assertFalse((self.root/f'filled-alias-{i}.xlsx').exists())
