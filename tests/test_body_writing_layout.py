"""Student-visible photo requests and object layout regressions."""
from io import BytesIO
from docx import Document
from tests._harness import ContractCase, bind_major, runtime


class BodyWritingLayoutTests(ContractCase):
    def setUp(self):
        super().setUp()
        self.root = self.make_project()
        bind_major(self.root, 'specialty_crops')
        self.builder = runtime('build_docx')
        self.context = runtime('gg_major_contract').output_context(self.root, 'specialty_crops')

    def convert(self, text):
        doc = self.builder.convert(text, base=self.root, context=self.context)
        data = BytesIO()
        doc.save(data)
        data.seek(0)
        return Document(data)

    def test_photo_guide_is_still_unresolved_and_prints_blue(self):
        guide = '※ 이 자리에 실제 부지 전경 사진 추가 필요'
        issues = runtime('gg_document').check(guide, None)
        self.assertIn('photo_placeholder', {item[0] for item in issues})
        doc = self.convert(guide)
        paragraph = next(p for p in doc.paragraphs if p.text == guide)
        self.assertEqual(str(paragraph.runs[0].font.color.rgb), '0000FF')
        self.assertEqual(paragraph.alignment, 1)
        self.assertEqual(paragraph.paragraph_format.first_line_indent, 0)

    def test_list_and_table_cell_photo_requests_remain_unresolved(self):
        guide = '※ 이 자리에 실제 부지 전경 사진 추가 필요'
        examples = [prefix + guide for prefix in ('', '- ', '* ', '+ ')]
        examples.extend('|대상|사진 요청|\n|---|---|\n|예정지|' + prefix + guide + '|'
                        for prefix in ('', '- ', '* ', '+ '))
        docmod = runtime('gg_document')
        for text in examples:
            with self.subTest(text=text):
                issues = docmod.check(text, None)
                requests = [message for check, message in issues if check == 'photo_placeholder']
                self.assertEqual(len(requests), 1)
                self.assertIn('1건', requests[0])
                self.assertEqual(docmod.photo_requests(docmod.parse(text)), [guide])

    def test_mixed_photo_request_count_keeps_legacy_prefix_without_double_counting(self):
        guide = '※ 이 자리에 실제 부지 전경 사진 추가 필요'
        text = ('- ' + guide + '\n\n|요청|\n|---|\n|' + guide
                + '|\n\n[사진 필요: 실제 시설 전경]\n본문 속 [사진 필요: 실제 작업 장면]')
        issues = runtime('gg_document').check(text, None)
        requests = [message for check, message in issues if check == 'photo_placeholder']
        self.assertEqual(len(requests), 1)
        self.assertIn('4건', requests[0])

    def test_detailed_authoring_marker_does_not_leak_shooting_instructions(self):
        doc = self.convert('[사진 필요: PHOTO-001 | 부지 전경 | 수평 원경 촬영 | 그림 1. 전경 | 80% | 입지 절]')
        printed = '\n'.join(p.text for p in doc.paragraphs)
        self.assertEqual(printed, '※ 이 자리에 부지 전경 사진 추가 필요')
        self.assertNotIn('PHOTO-001', printed)
        self.assertNotIn('수평 원경', printed)

    def test_spacers_stay_outside_table_caption_and_note(self):
        doc = self.convert('앞 본문.\n\n표 1. 계획\n|항목|값|\n|---|---|\n|시험|1|\n주 : 단위 설명.\n\n뒤 본문.')
        children = list(doc.element.body)
        table = doc.tables[0]._tbl
        index = children.index(table)
        self.assertEqual(''.join(children[index - 1].xpath('.//w:t/text()')), '표 1. 계획')
        self.assertEqual(''.join(children[index - 2].xpath('.//w:t/text()')), '')
        self.assertEqual(''.join(children[index + 1].xpath('.//w:t/text()')), '주 : 단위 설명.')
        self.assertEqual(''.join(children[index + 2].xpath('.//w:t/text()')), '')
        self.assertTrue(doc.tables[0].rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next)
        self.assertFalse(children[index + 2].xpath('./w:pPr/w:keepNext')[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') != '0')

    def test_heading_adjacent_object_has_no_extra_spacer(self):
        doc = self.convert('# Ⅰ. 머리말\n표 1. 계획\n|항목|값|\n|---|---|\n|시험|1|\n# Ⅱ. 환경')
        table = doc.tables[0]._tbl
        previous = table.getprevious().getprevious()
        self.assertEqual(''.join(previous.xpath('.//w:t/text()')), 'Ⅰ. 머리말')
        self.assertEqual(''.join(table.getnext().xpath('.//w:t/text()')), 'Ⅱ. 환경')

    def test_photo_guide_caption_and_note_form_one_kept_block(self):
        guide = '※ 이 자리에 실제 부지 전경 사진 추가 필요'
        for request in (guide, '- ' + guide, '[사진 필요: 실제 부지 전경]'):
            with self.subTest(request=request):
                doc = self.convert('앞 본문.\n\n' + request + '\n그림 1. 예정지 전경\n'
                                   '주 : 촬영 안내.\n\n뒤 본문.')
                self.assertEqual([p.text for p in doc.paragraphs],
                                 ['앞 본문.', '', guide, '그림 1. 예정지 전경',
                                  '주 : 촬영 안내.', '', '뒤 본문.'])
                self.assertTrue(doc.paragraphs[2].paragraph_format.keep_with_next)
                self.assertTrue(doc.paragraphs[3].paragraph_format.keep_with_next)
                self.assertFalse(doc.paragraphs[4].paragraph_format.keep_with_next)
                self.assertFalse(doc.paragraphs[1].paragraph_format.keep_with_next)
                self.assertFalse(doc.paragraphs[5].paragraph_format.keep_with_next)

    def test_photo_guide_does_not_chain_to_unrelated_prose(self):
        guide = '※ 이 자리에 실제 부지 전경 사진 추가 필요'
        doc = self.convert(guide + '\n\n뒤 본문.')
        self.assertEqual([p.text for p in doc.paragraphs], [guide, '', '뒤 본문.'])
        self.assertFalse(doc.paragraphs[0].paragraph_format.keep_with_next)
