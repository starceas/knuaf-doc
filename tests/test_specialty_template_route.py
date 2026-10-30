"""Specialty-crop template route: crop-category labels, daily-labor gender
warning and the professor-template print layout of sheet 7.

All workbooks are synthetic stand-ins carrying only the cells a contract
reads; no professor or student file is opened.  Hashes of the registered
template are simulated where the contract binds to them.
"""
import json
import tempfile
from pathlib import Path
from unittest import mock

from tests._harness import ContractCase, bind_major, runtime

S5, S6, S7, S8 = " 5. 판매계획", "6. 생산계획", "7. 영농자재소요계획", "8. 노무비계획"
X01_SHA = "5d19b6a2e5dcddccb70c68e39abee8f423bba005c008db7b64076326a17b9c80"


def labelled_template(path):
    """Stand-in with the X01 crop-bound labels, a formula and numbers."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "3.투자계획"
    wb.active["C24"] = "계"
    wb.create_sheet("10. 감가상각비계획 ")
    sales, prod, mat, labor = (wb.create_sheet(n) for n in (S5, S6, S7, S8))
    sales["B30"] = "미상품과"
    sales["D30"] = "=100%-D25-D29"
    prod["B17"], prod["E25"], prod["E27"] = "묘목수", "묘목수(1년차)", "묘목수(2년차 이후)"
    prod["H8"] = 1000
    mat["C4"], mat["C5"] = "묘목비(원/10a)", 5000
    labor["O24"] = labor["O38"] = "고용(남자, 단가 동일)"
    wb.save(path)


LABOR_LABELS = {"C11": "일용직", "E11": "남", "E12": "여",
                "O24": "고용(남자, 단가 동일)", "R24": "고용(여자)",
                "O38": "고용(남자, 단가 동일)", "R38": "고용(여자)",
                "O25": "일수", "R25": "일수", "O39": "일수", "R39": "일수"}


def labor_sheet(ws, *, male=(), female=(), task_male=(), task_female=()):
    """X01 '8. 노무비계획' labels plus the given hired days."""
    for cell, text in LABOR_LABELS.items():
        ws[cell] = text
    for cells, days in ((("G11", "J11", "M11", "P11", "S11"), male),
                        (("G12", "J12", "M12", "P12", "S12"), female)):
        for cell, value in zip(cells, days):
            ws[cell] = value
    for row, (task, m) in enumerate(task_male, 26):
        ws[f"B{row}"], ws[f"O{row}"] = task, m
    for row, (task, f) in enumerate(task_female, 26):
        ws[f"B{row}"], ws[f"R{row}"] = task, f


def x01_anchors(wb):
    """The two cells detect_layout reads to name a workbook X01."""
    inv = wb["3.투자계획"] if "3.투자계획" in wb.sheetnames else wb.create_sheet("3.투자계획")
    inv["C24"] = "계"
    if "10. 감가상각비계획 " not in wb.sheetnames:
        wb.create_sheet("10. 감가상각비계획 ")


def labor_book(path, *, x01=True, **days):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = S8
    labor_sheet(ws, **days)
    if x01:
        x01_anchors(wb)
    wb.save(path)

class CropLabelTests(ContractCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.tpl = runtime("gg_excel_template")
        self.fill = runtime("gg_excel_fill")
        self.src = self.root / "src.xlsx"
        labelled_template(self.src)
        real = self.tpl.sha256
        patcher = mock.patch.object(
            self.tpl, "sha256",
            side_effect=lambda p: X01_SHA if Path(p) == self.src else real(p))
        patcher.start()
        self.addCleanup(patcher.stop)

    def relabel(self, category, choices=None, name="out.xlsx"):
        project = self.make_project()
        bind_major(project)
        ctx = runtime("gg_major_contract").output_context(project, "specialty_crops")
        fill_map, values = self.tpl.crop_label_map(self.src, category, choices)
        m, v = self.root / (name + ".m.json"), self.root / (name + ".v.json")
        m.write_text(json.dumps(fill_map, ensure_ascii=False), encoding="utf-8")
        v.write_text(json.dumps(values, ensure_ascii=False), encoding="utf-8")
        out = self.root / name
        with mock.patch.object(self.fill, "sha256",
                               side_effect=lambda p: X01_SHA if Path(p) == self.src
                               else self.tpl.sha256.side_effect(p)):
            self.fill.fill_copy(self.src, m, v, out, context=ctx)
        return out

    def test_herbaceous_labels_replace_tree_terms_and_keep_numbers(self):
        out = self.relabel("herbaceous", {f"{S6}!B17": "재배시설"})
        from openpyxl import load_workbook
        wb = load_workbook(out)
        self.assertEqual(wb[S5]["B30"].value, "미상품")
        self.assertEqual(wb[S6]["B17"].value, "재배시설")
        self.assertEqual(wb[S6]["E25"].value, "정식 주수(1년차)")
        self.assertEqual(wb[S7]["C4"].value, "종자·육묘비(원/10a)")
        self.assertEqual(wb[S8]["O24"].value, "고용(남자)")
        self.assertEqual(wb[S5]["D30"].value, "=100%-D25-D29")
        wb.close()
        result = self.tpl.crop_label_check(self.src, out)
        self.assertEqual((result["status"], result["message"]), ("match", self.tpl.NO_DIFFERENCE))
        self.assertEqual(len(result["labelsChanged"]), 7)

    def test_woody_perennial_keeps_tree_terms(self):
        _, values = self.tpl.crop_label_map(self.src, "woody_perennial")
        changed = {(v["sheet"], v["cell"]): v["value"] for v in values["values"]}
        # Only the male-hire header loses its misleading "단가 동일" note.
        self.assertEqual(changed, {(S8, "O24"): "고용(남자)", (S8, "O38"): "고용(남자)"})

    def test_unlisted_label_category_or_moved_label_refused(self):
        with self.assertRaisesRegex(ValueError, "not a listed"):
            self.tpl.crop_label_map(self.src, "herbaceous", {f"{S5}!B30": "불량품"})
        with self.assertRaisesRegex(ValueError, "unknown crop category"):
            self.tpl.crop_label_map(self.src, "orchard")
        with self.assertRaisesRegex(ValueError, "outside the label table"):
            self.tpl.crop_label_map(self.src, "herbaceous", {f"{S5}!B31": "미상품"})
        from openpyxl import load_workbook
        wb = load_workbook(self.src)
        wb[S6]["B17"] = "주수"
        wb.save(self.src)
        wb.close()
        with self.assertRaisesRegex(ValueError, "expected original label"):
            self.tpl.crop_label_map(self.src, "herbaceous")

    def test_unregistered_template_refused(self):
        other = self.root / "other.xlsx"
        labelled_template(other)
        with self.assertRaisesRegex(ValueError, "registered X01/X02"):
            self.tpl.crop_label_map(other, "herbaceous")

    def test_check_fails_on_number_formula_or_unlisted_text_change(self):
        out = self.relabel("herbaceous")
        from openpyxl import load_workbook
        for label, edit in (
                ("number", lambda wb: wb[S6].__setitem__("H8", 999)),
                ("formula", lambda wb: wb[S5].__setitem__("D30", "=1-D25")),
                ("other_text", lambda wb: wb[S5].__setitem__("B31", "합계")),
                ("free_label", lambda wb: wb[S5].__setitem__("B30", "불량품"))):
            with self.subTest(change=label):
                bad = self.root / f"bad-{label}.xlsx"
                wb = load_workbook(out)
                edit(wb)
                wb.save(bad)
                wb.close()
                result = self.tpl.crop_label_check(self.src, bad)
                self.assertEqual((result["status"], result["message"]),
                                 ("warning", self.tpl.ADVISORY_WARNING))
                self.assertTrue(result["violations"])

    def test_label_map_admits_copy_derived_from_registered_template_only(self):
        """R1-01 (reverse order): a prepared or value-filled copy of the
        registered template is admitted by a cell-level formula match with
        the audit; any formula or sheet drift is refused."""
        from openpyxl import load_workbook
        audit_mod = runtime("gg_workbook_audit")
        entries, _ = audit_mod.inventory(audit_mod.Workbook.load(self.src))
        wb = load_workbook(self.src)
        audit = {"schema": "knuaf-workbook-formula-audit/v1", "source_sha256": X01_SHA,
                 "inventory": entries, "sheet_roles": {n: "input" for n in wb.sheetnames}}
        wb.close()
        def copy(name, edit):
            path = self.root / name
            wb = load_workbook(self.src)
            edit(wb)
            wb.save(path)
            wb.close()
            return path
        filled = copy("filled.xlsx", lambda wb: wb[S6].__setitem__("H8", 4321))
        drift = {"formula_changed": copy("f.xlsx", lambda wb: wb[S5].__setitem__("D30", "=1-D25")),
                 "formula_added": copy("a.xlsx", lambda wb: wb[S6].__setitem__("H9", "=H8")),
                 "sheet_added": copy("s.xlsx", lambda wb: wb.create_sheet("추가"))}
        with mock.patch.object(self.tpl, "registered_audit", return_value=audit):
            fill_map, values = self.tpl.crop_label_map(filled, "herbaceous")
            self.assertEqual(fill_map["template"]["sha256"], self.tpl.sha256(filled))
            self.assertIn("X01 derived", values["values"][0]["evidence_ref"]["locator"])
            for label, path in drift.items():
                with self.subTest(drift=label), self.assertRaisesRegex(ValueError, "registered X01/X02"):
                    self.tpl.crop_label_map(path, "herbaceous")

    def test_check_unreadable_after_file_is_read_failure_not_match(self):
        import contextlib, io
        tpl = self.tpl
        bad = self.root / "bad.xlsx"
        bad.write_bytes(b"not a workbook")
        for after in (bad, self.root / "missing.xlsx"):
            with self.subTest(after=after.name):
                result = tpl.crop_label_check(self.src, after)
                self.assertEqual((result["message"], result["status"]), (tpl.READ_FAILURE, "read_failed"))
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    code = tpl.cli(["label-check", "--before", str(self.src), "--after", str(after)])
                # Advisory: exit 0, the read-failure message first.
                self.assertEqual(code, 0)
                printed = json.loads(out.getvalue())
                self.assertEqual(next(iter(printed)), "message")
                self.assertEqual(printed["message"], tpl.READ_FAILURE)
                self.assertIn("판정 근거 아님", printed["meaning"])

    def test_check_fails_on_shared_formula_attribute_or_value_type_change(self):
        """R1-02: a shared follower re-pointed to another anchor (same empty
        formula text, different meaning) and a number stored as text with the
        same digits are violations, not passes."""
        import zipfile
        from openpyxl import load_workbook
        wb = load_workbook(self.src)
        prod = wb[S6]
        prod["H9"], prod["H10"] = "=H8*2", "=H8+1"
        wb.save(self.src)
        wb.close()
        shared = self.root / "shared.xlsx"
        with zipfile.ZipFile(self.src) as z, zipfile.ZipFile(shared, "w") as out:
            target = dict(self.tpl.workbook_sheets(z))[S6]
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == target:
                    text = data.decode()
                    text = text.replace('<f>H8*2</f>', '<f t="shared" ref="H9:I9" si="0">H8*2</f>')
                    text = text.replace('<f>H8+1</f>', '<f t="shared" ref="H10:I10" si="1">H8+1</f>')
                    text = text.replace('<c r="H10"', '<c r="I9"><f t="shared" si="0"/></c><c r="H10"', 1)
                    data = text.encode()
                out.writestr(item, data)
        out = self.relabel_from(shared, "herbaceous")
        self.assertEqual(self.tpl.crop_label_check(shared, out)["status"], "match")
        cases = {"shared_si": ('<c r="I9"><f t="shared" si="0"/>', '<c r="I9"><f t="shared" si="1"/>'),
                 "number_as_text": ('<c r="H8" t="n">', '<c r="H8" t="str">')}
        for label, (old, new) in cases.items():
            with self.subTest(change=label):
                bad = self.root / f"bad-{label}.xlsx"
                with zipfile.ZipFile(out) as z, zipfile.ZipFile(bad, "w") as w:
                    target = dict(self.tpl.workbook_sheets(z))[S6]
                    for item in z.infolist():
                        data = z.read(item.filename)
                        if item.filename == target:
                            self.assertIn(old.encode(), data)
                            data = data.replace(old.encode(), new.encode(), 1)
                        w.writestr(item, data)
                result = self.tpl.crop_label_check(shared, bad)
                self.assertEqual(result["status"], "warning")
                cell = "I9" if label == "shared_si" else "H8"
                self.assertEqual([(v["sheet"], v["cell"]) for v in result["violations"]], [(S6, cell)])

    def relabel_from(self, source, category):
        real = self.tpl.sha256
        with mock.patch.object(self.tpl, "sha256",
                               side_effect=lambda p: X01_SHA if Path(p) == source else real(p)):
            fill_map, values = self.tpl.crop_label_map(source, category)
        m, v = self.root / "rf.m.json", self.root / "rf.v.json"
        m.write_text(json.dumps(fill_map, ensure_ascii=False), encoding="utf-8")
        v.write_text(json.dumps(values, ensure_ascii=False), encoding="utf-8")
        project = self.make_project()
        bind_major(project)
        ctx = runtime("gg_major_contract").output_context(project, "specialty_crops")
        out = self.root / "rf.xlsx"
        with mock.patch.object(self.fill, "sha256",
                               side_effect=lambda p: X01_SHA if Path(p) == source else real(p)):
            self.fill.fill_copy(source, m, v, out, context=ctx)
        return out


class SharedReadCheckTests(ContractCase):
    """One read check runs before every template tool: a file it cannot
    read is blocked in each of them — never a pass, never 'another form'."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.tpl = runtime("gg_excel_template")
        self.patch = runtime("gg_excel_formula_patch")
        self.good = self.root / "good.xlsx"
        labelled_template(self.good)

    def unreadable(self):
        import zipfile
        corrupt = self.root / "corrupt.xlsx"
        corrupt.write_bytes(b"not an xlsx")
        no_workbook = self.root / "no-workbook.xlsx"
        with zipfile.ZipFile(self.good) as z, zipfile.ZipFile(no_workbook, "w") as out:
            for item in z.infolist():
                if item.filename != "xl/workbook.xml":
                    out.writestr(item, z.read(item.filename))
        orphan = self.root / "orphan-shared.xlsx"
        with zipfile.ZipFile(self.good) as z, zipfile.ZipFile(orphan, "w") as out:
            target = dict(self.tpl.workbook_sheets(z))[S5]
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == target:
                    data = data.replace(b"<f>100%-D25-D29</f>", b'<f t="shared" si="7"/>')
                out.writestr(item, data)
        return {"corrupt": corrupt, "missing": self.root / "missing.xlsx",
                "no_workbook_part": no_workbook, "shared_anchor_missing": orphan}

    def test_every_template_tool_blocks_an_unreadable_workbook(self):
        import contextlib, io
        for label, path in self.unreadable().items():
            with self.subTest(file=label):
                with self.assertRaises(self.tpl.TemplateUnreadable):
                    self.tpl.read_template(path)
                with self.assertRaises(self.tpl.TemplateUnreadable):
                    self.tpl.detect_layout(path)
                self.assertEqual([(i["code"], i["message"]) for i in self.tpl.labor_review(path)],
                                 [("LABOR_WORKBOOK_UNREADABLE", self.tpl.READ_FAILURE)])
                self.assertEqual(self.tpl.crop_label_check(self.good, path)["status"], "read_failed")
                self.assertEqual(self.tpl.crop_label_check(path, self.good)["status"], "read_failed")
                for tool in (lambda: self.tpl.crop_label_map(path, "herbaceous"),
                             lambda: self.patch.materialize_d8(path, prepare_observations=True),
                             lambda: self.tpl.wage_fill_map(path)):
                    with self.assertRaises(self.tpl.TemplateUnreadable):
                        tool()
                # Producers make nothing from an unreadable input and say so first.
                for cli, argv in (
                        (self.tpl.cli, ["label-map", "--source", str(path), "--category", "herbaceous",
                                        "--out-map", str(self.root / f"{label}.m"),
                                        "--out-values", str(self.root / f"{label}.v")]),
                        (self.patch.cli, ["materialize-d8", "--source", str(path),
                                          "--out", str(self.root / f"{label}.d8")])):
                    with contextlib.redirect_stdout(io.StringIO()) as out, \
                            contextlib.redirect_stderr(io.StringIO()):
                        self.assertEqual(cli(argv), 2, argv[0])
                    self.assertEqual(json.loads(out.getvalue())["message"], self.tpl.READ_FAILURE)
                    self.assertFalse([p for p in self.root.glob(f"{label}.*") if p.suffix in (".m", ".v", ".d8")])
                # The labor check is advisory: exit 0 with the read-failure message.
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    self.assertEqual(self.tpl.cli(["labor-review", "--workbook", str(path)]), 0)
                self.assertEqual(json.loads(out.getvalue())["message"], self.tpl.READ_FAILURE)

    def test_readable_other_form_is_not_applicable_not_blocked(self):
        from openpyxl import Workbook
        other = self.root / "other.xlsx"
        wb = Workbook()
        wb.active.title = "11. 인건비계획"
        wb.save(other)
        self.assertEqual(self.tpl.read_template(other).layout, "unknown")
        self.assertEqual(self.tpl.labor_review(other), [])
        with self.assertRaisesRegex(ValueError, "layout_variant_unsupported"):
            self.patch.materialize_d8(other, prepare_observations=True)


class PackageAndIdentificationFailureTests(ContractCase):
    """R3: the shared read check opens the file with openpyxl first and
    checks the package; an identification config it cannot read is
    unreadable too.  Each case is blocked in every tool and CLI; a normal
    X01 and a normal other form keep their results."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.tpl = runtime("gg_excel_template")
        self.patch = runtime("gg_excel_formula_patch")
        self.pr = runtime("gg_excel_print")
        self.mc = runtime("gg_major_contract")
        self.x01 = self.root / "x01.xlsx"
        labelled_template(self.x01)
        from openpyxl import Workbook
        self.other = self.root / "other.xlsx"
        wb = Workbook()
        wb.active.title = "11. 인건비계획"
        wb.active["B2"] = "다른 양식"
        wb.save(self.other)

    def rewrite(self, source, name, edit):
        import zipfile
        out = self.root / name
        with zipfile.ZipFile(source) as z, zipfile.ZipFile(out, "w") as w:
            sheets = dict(self.tpl.workbook_sheets(z))
            for item in z.infolist():
                data = edit(item.filename, z.read(item.filename), sheets)
                if data is not None:
                    w.writestr(item, data)
        return out

    def broken_packages(self):
        no_ct = lambda n, d, s: None if n == "[Content_Types].xml" else d
        def sheet_data_removed(n, d, s):
            if n != s[S6]:
                return d
            start, end = d.find(b"<sheetData"), d.find(b"</sheetData>")
            return d[:start] + d[end + len(b"</sheetData>"):]
        return {
            "x01_no_content_types": self.rewrite(self.x01, "a.xlsx", no_ct),
            "other_no_content_types": self.rewrite(self.other, "b.xlsx", no_ct),
            "broken_workbook_part": self.rewrite(
                self.x01, "w.xlsx", lambda n, d, s: b"<broken/>" if n == "xl/workbook.xml" else d),
            "broken_sheet_part": self.rewrite(
                self.x01, "s.xlsx", lambda n, d, s: b"<broken/>" if n == s[S6] else d),
            "sheet_without_sheetData": self.rewrite(self.x01, "d.xlsx", sheet_data_removed),
            "listed_sheet_part_missing": self.rewrite(
                self.x01, "m.xlsx", lambda n, d, s: None if n == s[S6] else d),
            # 추가 2: openpyxl opens these, the package check does not.
            "root_rels_missing": self.rewrite(
                self.x01, "r.xlsx", lambda n, d, s: None if n == "_rels/.rels" else d),
            "root_rels_broken": self.rewrite(
                self.x01, "rb.xlsx", lambda n, d, s: b"<Relationships" if n == "_rels/.rels" else d),
            "other_root_rels_missing": self.rewrite(
                self.other, "or.xlsx", lambda n, d, s: None if n == "_rels/.rels" else d),
            "shared_strings_missing": self.shared_strings_case("ss.xlsx", None),
            "shared_strings_broken": self.shared_strings_case("sb.xlsx", b"<sst"),
            "shared_string_index_out_of_range": self.shared_strings_case("si.xlsx", "index"),
        }

    def shared_strings_case(self, name, replacement):
        """X01 stand-in whose labels are stored as shared strings (t="s"),
        then the table is removed, broken, or a cell points past it."""
        from openpyxl import load_workbook
        base = self.root / "sst-base.xlsx"
        if not base.exists():
            # Excel-style storage: move inline strings into xl/sharedStrings.xml.
            import re
            strings = []
            def to_shared(data):
                def repl(m):
                    strings.append(m.group(2))
                    return m.group(1).replace(b' t="inlineStr"', b' t="s"') + b"<v>%d</v></c>" % (len(strings) - 1)
                return re.sub(rb'(<c [^>]*t="inlineStr"[^>]*>)<is><t>(.*?)</t></is></c>', repl, data)
            src = self.rewrite(self.x01, "sst-tmp.xlsx", lambda n, d, s: to_shared(d)
                               if n.startswith("xl/worksheets/") else d)
            import zipfile
            sst = (b'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="' + self.tpl.NS_MAIN.encode()
                   + b'">' + b"".join(b"<si><t>" + s + b"</t></si>" for s in strings) + b"</sst>")
            with zipfile.ZipFile(src) as z, zipfile.ZipFile(base, "w") as w:
                for item in z.infolist():
                    data = z.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        data = data.replace(b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType='
                                            b'"application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
                    if item.filename == "xl/_rels/workbook.xml.rels":
                        data = data.replace(b"</Relationships>", b'<Relationship Id="rIdSST" Type="http://schemas.'
                                            b'openxmlformats.org/officeDocument/2006/relationships/sharedStrings" '
                                            b'Target="sharedStrings.xml"/></Relationships>')
                    w.writestr(item, data)
                w.writestr("xl/sharedStrings.xml", sst)
            self.assertTrue(strings)
            wb = load_workbook(base)  # the Excel-style control opens and reads the same labels
            self.assertEqual(wb[S5]["B30"].value, "미상품과")
            wb.close()
            self.assertEqual(self.tpl.read_template(base).layout, "x01")
        if replacement == "index":
            return self.rewrite(base, name, lambda n, d, s: d.replace(b't="s"><v>0</v>', b't="s"><v>9999</v>', 1)
                                if n.startswith("xl/worksheets/") else d)
        return self.rewrite(base, name, lambda n, d, s: replacement if n == "xl/sharedStrings.xml" else d)

    def assert_blocked_everywhere(self, path, *, before=None):
        """Checks report the read failure (advisory, exit 0, message first);
        producers make nothing (exit 2, message first, no file)."""
        import contextlib, io
        before = before or self.x01
        tpl, patch = self.tpl, self.patch
        for call in (lambda: tpl.read_template(path), lambda: tpl.detect_layout(path),
                     lambda: tpl.crop_label_map(path, "herbaceous"),
                     lambda: patch.materialize_d8(path, prepare_observations=True),
                     lambda: tpl.wage_fill_map(path),
                     lambda: tpl.inspect_source(path, None, None)):
            with self.assertRaises(tpl.TemplateUnreadable):
                call()
        self.assertEqual([(i["code"], i["status"], i["message"]) for i in tpl.labor_review(path)],
                         [("LABOR_WORKBOOK_UNREADABLE", "pending", tpl.READ_FAILURE)])
        self.assertEqual(tpl.crop_label_check(before, path)["status"], "read_failed")
        self.assertEqual(tpl.crop_label_check(path, before)["status"], "read_failed")
        n = len(list(self.root.iterdir()))
        checks = ((tpl.cli, ["labor-review", "--workbook", str(path)]),
                  (tpl.cli, ["label-check", "--before", str(before), "--after", str(path)]),
                  (tpl.cli, ["label-check", "--before", str(path), "--after", str(before)]))
        producers = ((tpl.cli, ["label-map", "--source", str(path), "--category", "herbaceous",
                                "--out-map", str(self.root / f"lm{n}.json"),
                                "--out-values", str(self.root / f"lv{n}.json")]),
                     (tpl.cli, ["wage-map", "--source", str(path), "--out-map", str(self.root / f"wm{n}.json")]),
                     (tpl.cli, ["inspect", "--source", str(path), "--out-map", str(self.root / f"im{n}.json")]),
                     (patch.cli, ["materialize-d8", "--source", str(path), "--prepare-observations",
                                  "--out", str(self.root / f"d8{n}.json")]))
        for expected, calls in ((0, checks), (2, producers)):
            for cli, argv in calls:
                with self.subTest(cli=argv[0]), contextlib.redirect_stdout(io.StringIO()) as out, \
                        contextlib.redirect_stderr(io.StringIO()):
                    self.assertEqual(cli(argv), expected)
                    printed = json.loads(out.getvalue())
                    self.assertEqual((next(iter(printed)), printed["message"]),
                                     ("message", tpl.READ_FAILURE))
        for out in ("lm", "lv", "wm", "im", "d8"):
            self.assertFalse(list(self.root.glob(f"{out}{n}.json")))
        with self.assertRaises(tpl.TemplateUnreadable):
            self.print_copy(path, f"p{n}")
        self.assertFalse((self.root / f"p{n}.xlsx").exists())
        self.assertFalse((self.root / f"p{n}.json").exists())

    def print_copy(self, src, name):
        project = self.make_project()
        bind_major(project, "specialty_crops")
        ctx = self.mc.output_context(project, "specialty_crops")
        m = self.root / f"{name}.map.json"
        m.write_text(json.dumps({"schema": self.pr.SCHEMA, "source": {"sha256": self.tpl.sha256(src)},
                                 "sheets": [{"sheet": S6, "fit_width": 1, "fit_height": 0,
                                             "reason": "합성"}]}, ensure_ascii=False), encoding="utf-8")
        return self.pr.apply(src, m, self.root / f"{name}.xlsx", self.root / f"{name}.json", context=ctx)

    def test_package_failures_are_blocked_in_every_tool(self):
        import openpyxl
        for label, path in self.broken_packages().items():
            with self.subTest(case=label):
                self.assert_blocked_everywhere(path)
        # Counterexample A: the X01 copy without [Content_Types].xml does not
        # open with openpyxl, and that is the first check's reason.
        with self.assertRaisesRegex(self.tpl.TemplateUnreadable, "does not open"):
            self.tpl.read_template(self.root / "a.xlsx")
        with self.assertRaises(Exception):
            openpyxl.load_workbook(self.root / "a.xlsx")
        # 추가 2: openpyxl alone opens a copy without _rels/.rels; the package
        # half of the first check is what stops it.
        wb = openpyxl.load_workbook(self.root / "r.xlsx")
        wb.close()
        with self.assertRaisesRegex(self.tpl.TemplateUnreadable, r"package part missing: _rels/\.rels"):
            self.tpl.read_template(self.root / "r.xlsx")

    def test_identification_config_failure_is_unreadable_not_another_form(self):
        bad_json = self.root / "anchors.json"
        bad_json.write_text("{not json", encoding="utf-8")
        wrong_schema = self.root / "schema.json"
        wrong_schema.write_text(json.dumps({"schema": "other/v1", "layout_anchors": {}}), encoding="utf-8")
        no_anchors = self.root / "empty.json"
        no_anchors.write_text(json.dumps({"schema": "knuaf-common-workbook-extract/v1",
                                          "layout_anchors": {}}), encoding="utf-8")
        for label, config in (("missing", self.root / "absent.json"), ("unparsable", bad_json),
                              ("wrong_schema", wrong_schema), ("no_anchors", no_anchors)):
            for workbook in (self.x01, self.other):
                with self.subTest(config=label, workbook=workbook.name), \
                        mock.patch.object(self.tpl, "LAYOUT_EXTRACT_PATH", config):
                    self.assert_blocked_everywhere(workbook, before=workbook)

    def test_normal_x01_and_other_form_results_unchanged(self):
        import contextlib, io
        self.assertEqual(self.tpl.detect_layout(self.x01), "x01")
        self.assertEqual(self.tpl.detect_layout(self.other), "unknown")
        labor = self.root / "labor-x01.xlsx"
        labor_book(labor, female=(0, 60, 60, 60, 60))
        for path in (labor, self.other):
            with self.subTest(workbook=path.name):
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    self.assertEqual(self.tpl.cli(["labor-review", "--workbook", str(path)]), 0)
                printed = json.loads(out.getvalue())
                self.assertEqual((printed["message"], printed["status"], printed["issues"]),
                                 (self.tpl.NO_DIFFERENCE, "reviewed", []))
        for path in (self.x01, self.other):
            with self.subTest(workbook=path.name):
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    self.assertEqual(self.tpl.cli(["label-check", "--before", str(path),
                                                   "--after", str(path)]), 0)
                self.assertEqual(json.loads(out.getvalue())["status"], "match")
        with self.assertRaisesRegex(ValueError, "layout_variant_unsupported"):
            self.patch.materialize_d8(self.other, prepare_observations=True)
        self.assertEqual(self.tpl.inspect_source(self.other, None, None, legacy_map=True)["mapStatus"],
                         "manual_map_required")
        self.assertEqual(self.tpl.inspect_source(self.x01, None, None, legacy_map=True)["mapStatus"], "ok")
        receipt = self.print_copy(self.x01, "x01print")
        self.assertEqual(receipt["changes"][0]["sheet"], S6)


class LaborReviewTests(ContractCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.tpl = runtime("gg_excel_template")

    def review(self, **kw):
        path = self.root / "labor.xlsx"
        labor_book(path, **kw)
        return self.tpl.labor_review(path)

    def test_missing_or_corrupt_file_is_read_failure_not_reviewed(self):
        """R1-03 (advisory since R4): a file that cannot be read is not
        'another form' — the read-failure message, not a block."""
        import contextlib, io
        tpl = self.tpl
        corrupt = self.root / "corrupt.xlsx"
        corrupt.write_bytes(b"not an xlsx")
        for path in (corrupt, self.root / "missing.xlsx"):
            with self.subTest(path=path.name):
                issues = tpl.labor_review(path)
                self.assertEqual([(i["code"], i["status"], i["severity"], i["message"]) for i in issues],
                                 [("LABOR_WORKBOOK_UNREADABLE", "pending", "warning", tpl.READ_FAILURE)])
                with contextlib.redirect_stdout(io.StringIO()) as out:
                    code = tpl.cli(["labor-review", "--workbook", str(path)])
                self.assertEqual(code, 0)
                printed = json.loads(out.getvalue())
                self.assertEqual((next(iter(printed)), printed["message"], printed["status"]),
                                 ("message", tpl.READ_FAILURE, "read_failed"))
                self.assertIn("판정 근거 아님", printed["meaning"])
        # A readable non-X01 workbook stays not-applicable.
        other = self.root / "other.xlsx"
        labor_book(other, x01=False, male=(0, 60))
        with contextlib.redirect_stdout(io.StringIO()) as out:
            self.assertEqual(tpl.cli(["labor-review", "--workbook", str(other)]), 0)
        self.assertEqual(json.loads(out.getvalue())["issues"], [])

    def test_male_only_hired_days_warn_with_task_list(self):
        issues = self.review(male=(0, 60, 60, 60, 60), female=(0, 0, 0, 0, 0),
                             task_male=[("수확 및 선별", 60)])
        self.assertEqual([i["code"] for i in issues], ["LABOR_HIRED_MALE_ONLY"])
        self.assertEqual(issues[0]["severity"], "warning")
        self.assertEqual(issues[0]["evidence"]["tasks"][0]["task"], "수확 및 선별")

    def test_task_rows_alone_trigger_warning(self):
        issues = self.review(male=(0,) * 5, female=(0,) * 5,
                             task_male=[("정식", 12)])
        self.assertEqual([i["code"] for i in issues], ["LABOR_HIRED_MALE_ONLY"])

    def test_female_or_split_or_no_hire_is_silent(self):
        for label, kw in (
                ("female", dict(male=(0,) * 5, female=(0, 60, 60, 60, 60),
                                task_female=[("수확", 60)])),
                ("split", dict(male=(10,) * 5, female=(50,) * 5,
                               task_male=[("정식", 10)], task_female=[("수확", 50)])),
                ("none", dict(male=(0,) * 5, female=(0,) * 5))):
            with self.subTest(case=label):
                self.assertEqual(self.review(**kw), [])

    def test_unreadable_labor_sheet_is_read_failure_not_a_pass(self):
        from openpyxl import Workbook, load_workbook
        path = self.root / "labor.xlsx"
        cases = {}
        wb = Workbook(); wb.active.title = "8. 노무계획"; x01_anchors(wb); wb.save(path)
        cases["sheet_missing"] = ("LABOR_SHEET_MISSING", self.tpl.labor_review(path))
        labor_book(path, female=(0, 60))
        wb = load_workbook(path); wb[S8]["R24"] = "고용 인부"; wb.save(path)
        cases["header_changed"] = ("LABOR_LAYOUT_UNSUPPORTED", self.tpl.labor_review(path))
        labor_book(path, male=(0, "60일"))
        cases["text_days"] = ("LABOR_DAYS_UNREADABLE", self.tpl.labor_review(path))
        labor_book(path)
        wb = load_workbook(path); wb[S8]["J12"] = "=60*1"; wb.save(path)  # no cached value
        cases["uncached_formula"] = ("LABOR_DAYS_UNREADABLE", self.tpl.labor_review(path))
        for label, (code, issues) in cases.items():
            with self.subTest(case=label):
                self.assertEqual([i["code"] for i in issues], [code])
                self.assertEqual((issues[0]["status"], issues[0]["severity"], issues[0]["message"]),
                                 ("pending", "warning", self.tpl.READ_FAILURE))
                self.assertTrue(issues[0]["reason"].startswith(self.tpl.READ_FAILURE))
        # Header whitespace differences are the same layout, not unreadable.
        labor_book(path, female=(0, 60))
        wb = load_workbook(path); wb[S8]["R24"] = " 고용 (여자) "; wb.save(path)
        self.assertEqual(self.tpl.labor_review(path), [])

    def test_non_x01_workbook_is_not_applicable(self):
        from openpyxl import Workbook
        path = self.root / "other.xlsx"
        # Another form with male-only days or no labor sheet at all: not in scope.
        labor_book(path, x01=False, male=(0, 60, 60, 60, 60))
        self.assertEqual(self.tpl.labor_review(path), [])
        wb = Workbook(); wb.active.title = "11. 인건비계획"; wb.save(path)
        self.assertEqual(self.tpl.labor_review(path), [])
        # X02 (perennial-asset variant) is a different layout, also not in scope.
        labor_book(path, x01=False, male=(0, 60))
        from openpyxl import load_workbook
        wb = load_workbook(path)
        wb.create_sheet("3.투자계획")["C25"] = "계"
        wb.create_sheet("10. 감가상각비계획 ")["J7"] = "대식물"
        wb.save(path)
        self.assertEqual(self.tpl.detect_layout(path), "x02")
        self.assertEqual(self.tpl.labor_review(path), [])

    def test_review_plan_surfaces_warning_and_read_failure_without_blocking(self):
        from tests.test_skillgaps_wp3_finance import plan_workbook
        from openpyxl import load_workbook
        review = runtime("gg_finance_tables").review_plan
        path = self.root / "plan.xlsx"
        plan_workbook(path)
        # Not an X01 workbook: no labor issue at all.
        self.assertFalse([i for i in review(path)["issues"] if i["code"].startswith("LABOR_")])
        wb = load_workbook(path)
        x01_anchors(wb)
        wb.save(path)
        missing = next(i for i in review(path)["issues"] if i["code"] == "LABOR_SHEET_MISSING")
        # Advisory: a pending warning with the read-failure reason, not a block.
        self.assertEqual((missing["status"], missing["severity"]), ("pending", "warning"))
        self.assertTrue(missing["reason"].startswith(self.tpl.READ_FAILURE))
        wb = load_workbook(path)
        labor_sheet(wb.create_sheet(S8), male=(0, 60, 60, 60, 60))
        wb.save(path)
        wb.close()
        plan = review(path)
        item = next(i for i in plan["issues"] if i["code"] == "LABOR_HIRED_MALE_ONLY")
        self.assertEqual(item["check_id"], "finance_plan_review")
        self.assertEqual(item["severity"], "warning")
        self.assertFalse([i for i in plan["issues"] if i["code"].startswith("LABOR_")
                          and i["status"] != "pending"])


class LaborAdvisoryCheckOutputTests(ContractCase):
    """R5: through review_plan -> core.checks -> ``gg.py check`` the final
    finance_plan_review labor rows still carry "보조 경고, 판정 근거 아님"
    and never block; other finance rows keep their behavior."""

    def check_rows(self, edit):
        import contextlib, io
        from tests.test_skillgaps_integ_plan_review import FinancePlanReviewTests
        case = FinancePlanReviewTests()
        core, root = case._project(mutate=edit)
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            runtime("gg").main(["check", str(root)])
        report = json.loads(out.getvalue())
        labor = [r for r in report if r["check_id"] == "finance_plan_review"
                 and any(str(e.get("code", "")).startswith("LABOR_") for e in r.get("evidence", []))]
        others = [r for r in report if r["check_id"] == "finance_plan_review" and r not in labor]
        return core, labor, others

    def test_final_labor_rows_keep_the_advisory_note_and_do_not_block(self):
        from openpyxl import load_workbook
        def mutate(kind):
            def edit(path):
                wb = load_workbook(path)
                x01_anchors(wb)
                if kind != "LABOR_SHEET_MISSING":
                    male = (0, "60일") if kind == "LABOR_DAYS_UNREADABLE" else (0, 60)
                    labor_sheet(wb.create_sheet(S8), male=male)
                wb.save(path)
                wb.close()
            return edit
        for code in ("LABOR_SHEET_MISSING", "LABOR_HIRED_MALE_ONLY", "LABOR_DAYS_UNREADABLE"):
            with self.subTest(code=code):
                core, labor, others = self.check_rows(mutate(code))
                self.assertEqual([r["evidence"][0]["code"] for r in labor], [code])
                row = labor[0]
                self.assertEqual((row["status"], row["severity"]), ("pending", "warning"))
                self.assertIn("보조 경고, 판정 근거 아님", row["recovery"])
                self.assertFalse(core.blocks_skill_candidate(row))
                if code != "LABOR_HIRED_MALE_ONLY":
                    self.assertTrue(row["reason"].startswith("읽기 실패 — 사람 확인"))
                # Other finance rows keep the old rule: a non-blocked row gets
                # no recovery text (this fixture yields pending warnings).
                self.assertTrue(others)
                for r in others:
                    if r["status"] != "blocked":
                        self.assertNotIn("recovery", r)


class MaterialSheetPrintTests(ContractCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.pr = runtime("gg_excel_print")
        self.tpl = runtime("gg_excel_template")
        self.mc = runtime("gg_major_contract")

    def source(self, *, x02=False, unknown=False):
        from openpyxl import Workbook
        wb = Workbook()
        inv = wb.active
        inv.title = "3.투자계획"
        if not unknown:
            inv["C25" if x02 else "C24"] = "계"
        dep = wb.create_sheet("10. 감가상각비계획 ")
        if x02:
            dep["J7"] = "대식물"
        for name in (S7, "12 손익계획"):
            wb.create_sheet(name)["B2"] = name
        path = self.root / ("x02.xlsx" if x02 else "unknown.xlsx" if unknown else "x01.xlsx")
        wb.save(path)
        return path

    def test_unidentified_layout_material_sheet_is_protected(self):
        src = self.source(unknown=True)
        self.assertEqual(self.tpl.detect_layout(src), "unknown")
        with self.assertRaisesRegex(ValueError, "영농자재소요계획.*X01"):
            self.apply(src, S7, name="unk")
        self.assertFalse((self.root / "unk.xlsx").exists())

    def apply(self, src, sheet, major="specialty_crops", name="out"):
        project = self.make_project()
        bind_major(project, major)
        ctx = self.mc.output_context(project, major)
        m = self.root / f"{name}.map.json"
        m.write_text(json.dumps({"schema": self.pr.SCHEMA,
                                 "source": {"sha256": self.tpl.sha256(src)},
                                 "sheets": [{"sheet": sheet, "fit_width": 1,
                                             "fit_height": 0, "reason": "합성"}]},
                                ensure_ascii=False), encoding="utf-8")
        out = self.root / f"{name}.xlsx"
        return self.pr.apply(src, m, out, self.root / f"{name}.json", context=ctx), out

    def test_x01_material_sheet_layout_cannot_change(self):
        src = self.source()
        with self.assertRaisesRegex(ValueError, "영농자재소요계획.*X01"):
            self.apply(src, S7)
        self.assertFalse((self.root / "out.xlsx").exists())

    def test_other_sheet_and_other_layout_unchanged(self):
        receipt, _ = self.apply(self.source(), "12 손익계획", name="pl")
        self.assertEqual(receipt["changes"][0]["sheet"], "12 손익계획")
        receipt, _ = self.apply(self.source(x02=True), S7, name="x02out")
        self.assertEqual(receipt["changes"][0]["sheet"], S7)
