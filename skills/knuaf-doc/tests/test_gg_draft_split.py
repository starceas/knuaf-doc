#!/usr/bin/env python3
"""gg_draft_split 오프라인 시험 — 합성 원고만 쓴다 (학생 원고 아님).

실행: python3 test_gg_draft_split.py   (스크립트 상위 tests/ 위치에서)
"""

import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location(
    "gg_draft_split", ROOT / "scripts" / "gg_draft_split.py")
split = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(split)

TOC_PATH = ROOT / "tests" / "fixtures" / "sample-toc.json"


def load_toc():
    return split.load_toc(TOC_PATH)


def write_src(d, name, text):
    p = Path(d) / name
    p.write_text(text, encoding="utf-8")
    return p


def run_gg(*args):
    env = dict(os.environ)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    return subprocess.run(
        [sys.executable, "-B", str(ROOT / "scripts" / "gg.py"), *args],
        capture_output=True, text=True, env=env)


NORMAL = """표지 페이지
합성 대학 졸업논문 초안

Ⅰ. 머리말
본 계획서는 합성 작목의 영농창업 계획을 다룬다.

1. 농업환경분석
기후와 토양을 살핀다.
| 항목 | 값 |
| 평균기온 | 12.5℃ |

2. 시장환경분석
수요를 본다.

4. 재배기술
정식 방법을 쓴다.

6. 관련 논문
가나다 등(2020) 합성논문.
홍길동 등(2021) 다른논문.
김철수 등(2022) 세번째.

Ⅴ. 맺음말
맺는다.

Ⅵ. 참고문헌
합성논문 저널 2020.
"""


class Plan(unittest.TestCase):
    def setUp(self):
        self.toc = load_toc()
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.work = Path(self._tmp.name) / "work"
        self.work.mkdir()

    def plan(self, text, name="draft.md"):
        src = write_src(self._tmp.name, name, text)
        return split.build_plan(text.splitlines(), self.toc,
                                src_path=src, fmt="text",
                                workdir=self.work), src

    def test_normal_split(self):
        plan, _ = self.plan(NORMAL)
        proposed = {i["section_id"]: i for i in plan["items"]
                    if i["status"] == "proposed"}
        self.assertEqual(
            set(proposed), {"I", "II-1", "II-2", "III-4", "III-6",
                            "V", "VI"})
        # 목차 앞 본문(표지)은 보류
        reasons = {h["reason"] for h in plan["held"]}
        self.assertIn("unmatched_prefix", reasons)

    def test_duplicate_heading_held(self):
        text = NORMAL + "\n6. 관련 논문\n중복 제목 줄.\n"
        plan, _ = self.plan(text)
        held = [h for h in plan["held"] if h["reason"] ==
                "duplicate_heading"]
        self.assertTrue(held)
        self.assertNotIn("III-6",
                         {i["section_id"] for i in plan["items"]})
        proposed_ranges = [tuple(map(int, i["lines"].split("-")))
                           for i in plan["items"]]
        held_ranges = [tuple(map(int, h["lines"].split("-")))
                       for h in held]
        for start, end in held_ranges:
            self.assertFalse(any(max(start, lo) <= min(end, hi)
                                 for lo, hi in proposed_ranges))

    def test_out_of_order_body_is_not_absorbed_by_previous_section(self):
        text = "Ⅰ. 머리말\n머리말 본문\n6. 관련 논문\n보류 본문\n" \
               "1. 농업환경분석\n역순 본문\nⅤ. 맺음말\n맺음말 본문\n"
        plan, _ = self.plan(text)
        intro = next(i for i in plan["items"] if i["section_id"] == "I")
        ending = next(i for i in plan["items"] if i["section_id"] == "V")
        intro_body = "\n".join(text.splitlines()[intro["line"] - 1:
                                                int(intro["lines"].split("-")[1])])
        ending_body = "\n".join(text.splitlines()[ending["line"] - 1:
                                                 int(ending["lines"].split("-")[1])])
        self.assertNotIn("보류 본문", intro_body)
        self.assertNotIn("역순 본문", intro_body)
        self.assertNotIn("역순 본문", ending_body)
        reasons = {h["reason"] for h in plan["held"]}
        self.assertIn("out_of_order", reasons)

    def test_out_of_order_held(self):
        text = "Ⅰ. 머리말\n앞부분\n6. 관련 논문\n본문\n" \
               "1. 농업환경분석\n뒤늦은 절\n"
        plan, _ = self.plan(text)
        reasons = {h["reason"] for h in plan["held"]}
        self.assertIn("out_of_order", reasons)

    def test_table_middle_boundary_held(self):
        # 비앵커 목차 패턴이 표 셀 안에서도 걸릴 때 — 경계로 쓰지 않고 보류
        toc = load_toc() + [{
            "id": "X-1", "title": "표안후보",
            "patterns": [__import__("re").compile("표셀제목")],
            "order": 7}]
        text = "Ⅰ. 머리말\n서론\n| 항목 | 값 |\n" \
               "| 표셀제목 | 내부셀 |\n| a | b |\nⅤ. 맺음말\n끝\n"
        src = write_src(self._tmp.name, "t.md", text)
        plan = split.build_plan(text.splitlines(), toc, src_path=src,
                                fmt="text", workdir=self.work)
        held = [h for h in plan["held"]
                if h["reason"] == "heading_inside_table"]
        self.assertTrue(held)
        # 표 안 제목은 경계가 되지 않아 표가 잘리지 않는다
        self.assertEqual(
            {i["section_id"] for i in plan["items"]}, {"I", "V"})
        intro = next(i for i in plan["items"] if i["section_id"] == "I")
        self.assertEqual(intro["table_blocks"], 1)

    def test_no_headings(self):
        plan, _ = self.plan("그냥 본문만 있는 문서\n두 번째 줄\n")
        self.assertFalse(plan["items"])
        self.assertTrue(any(h["reason"] == "no_headings"
                            for h in plan["held"]))

    def test_same_line_two_sections_held_independent_of_toc_order(self):
        text = "Ⅰ. 머리말\n합성 본문\n"
        src = write_src(self._tmp.name, "ambiguous.md", text)
        entries = [{"id": "I", "title": "머리말", "patterns": ["머리말"]},
                   {"id": "II", "title": "다음 절", "patterns": ["Ⅰ"]}]
        plans = []
        for ordered in (entries, list(reversed(entries))):
            with self.subTest(order=[entry["id"] for entry in ordered]):
                toc_path = write_src(self._tmp.name, "ambiguous-toc.json",
                                     json.dumps(ordered))
                plan = split.build_plan(
                    text.splitlines(), split.load_toc(toc_path),
                    src_path=src, fmt="text", workdir=self.work)
                plans.append(plan)
                self.assertEqual([], plan["items"])
                self.assertEqual({"I", "II"},
                                 {entry["section_id"] for entry in plan["held"]})
                self.assertTrue(all(entry["reason"] == "ambiguous_heading"
                                    and entry["status"] == "held"
                                    and entry["lines"] == "1-2"
                                    for entry in plan["held"]))
                result = split.apply_plan(self.work, plan, text.splitlines(),
                                          source_path=src)
                self.assertEqual([], result["written"])
                self.assertEqual([], list(self.work.glob("sections/*.md")))
        self.assertEqual(plans[0], plans[1])

    def test_ambiguous_body_is_not_absorbed_by_adjacent_sections(self):
        text = "시작 장\n앞 본문\nⅠ. 머리말\n보류 본문\n끝 장\n끝 본문\n"
        toc_path = write_src(self._tmp.name, "boundary-toc.json", json.dumps([
            {"id": "A", "patterns": ["시작 장"]},
            {"id": "I", "patterns": ["머리말"]},
            {"id": "II", "patterns": ["Ⅰ"]},
            {"id": "V", "patterns": ["끝 장"]}]))
        src = write_src(self._tmp.name, "boundary.md", text)
        plan = split.build_plan(text.splitlines(), split.load_toc(toc_path),
                                src_path=src, fmt="text", workdir=self.work)
        self.assertEqual({"A": "1-2", "V": "5-6"},
                         {entry["section_id"]: entry["lines"]
                          for entry in plan["items"]})
        self.assertEqual({"I": "3-4", "II": "3-4"},
                         {entry["section_id"]: entry["lines"]
                          for entry in plan["held"]})

    def test_existing_file_marked_conflict(self):
        (self.work / "sections").mkdir()
        (self.work / "sections" / "I.md").write_text("기존", encoding="utf-8")
        plan, _ = self.plan(NORMAL)
        intro = next(i for i in plan["items"] if i["section_id"] == "I")
        self.assertTrue(intro["conflict"])


class Apply(unittest.TestCase):
    def setUp(self):
        self.toc = load_toc()
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.work = Path(self._tmp.name) / "work"
        (self.work / "sections").mkdir(parents=True)

    def _make(self, text=NORMAL):
        src = write_src(self._tmp.name, "draft.md", text)
        plan = split.build_plan(text.splitlines(), self.toc,
                                src_path=src, fmt="text",
                                workdir=self.work)
        return src, plan, text.splitlines()

    def test_apply_writes_sections_and_change(self):
        src, plan, lines = self._make()
        result = split.apply_plan(self.work, plan, lines, source_path=src,
                                  change_out="build/change.json")
        self.assertTrue((self.work / "sections" / "I.md").exists())
        self.assertTrue((self.work / "sections" / "III-6.md").exists())
        # 원본 사본 보존
        safe_name = "manuscript-%s.md" % split.hashlib.sha256(
            src.read_bytes()).hexdigest()[:12]
        copied = self.work / "sources" / "imported" / safe_name
        self.assertEqual(copied.read_bytes(), src.read_bytes())
        # 절 파일은 gg:draft 마커 봉투
        text = (self.work / "sections" / "I.md").read_text(
            encoding="utf-8")
        self.assertIn(split.START, text)
        self.assertIn(split.END, text)
        self.assertIn("머리말", text)
        # 변경 JSON: sources + sections + tasks
        change = result["change"]
        cols = [o["collection"] for o in change["ops"]]
        self.assertIn("sources", cols)
        self.assertIn("sections", cols)
        self.assertIn("tasks", cols)
        sec_ops = [o["value"] for o in change["ops"]
                   if o["collection"] == "sections"]
        for rec in sec_ops:
            self.assertIn(rec["status"], {"drafting"})
            self.assertIsInstance(rec["order"], int)
            self.assertTrue(rec["path"].startswith("sections/"))
            self.assertTrue(rec["depends_on"])
        self.assertTrue((self.work / "build" / "change.json").exists())

    def test_apply_never_overwrites(self):
        (self.work / "sections" / "I.md").write_text(
            "기존 절 내용", encoding="utf-8")
        _src, plan, lines = self._make()
        result = split.apply_plan(self.work, plan, lines, source_path=_src)
        self.assertEqual(
            (self.work / "sections" / "I.md").read_text(
                encoding="utf-8"), "기존 절 내용")
        self.assertIn(
            {"section_id": "I", "reason": "conflict_existing_file"},
            result["skipped"])

    def test_plan_and_copy_do_not_expose_source_filename(self):
        src, plan, lines = self._make()
        private_name = "private-identifying-name.md"
        renamed = src.with_name(private_name)
        src.rename(renamed)
        plan = split.build_plan(lines, self.toc, src_path=renamed,
                                fmt="text", workdir=self.work)
        self.assertNotIn(private_name, json.dumps(plan, ensure_ascii=False))
        result = split.apply_plan(self.work, plan, lines, source_path=renamed)
        self.assertNotIn(private_name, json.dumps(result, ensure_ascii=False))

    def test_tampered_target_path_is_rejected_before_writes(self):
        src, plan, lines = self._make()
        plan["items"][0]["target_path"] = "sections/../../outside.md"
        with self.assertRaises(ValueError):
            split.apply_plan(self.work, plan, lines, source_path=src)
        self.assertFalse((self.work.parent / "outside.md").exists())

    def test_existing_ledger_section_is_not_replaced(self):
        src, plan, lines = self._make()
        (self.work / "project.json").write_text(json.dumps({
            "sections": {"I": {"id": "I", "path": "sections/elsewhere.md"}},
            "sources": {},
        }), encoding="utf-8")
        result = split.apply_plan(self.work, plan, lines, source_path=src)
        self.assertIn({"section_id": "I", "reason": "conflict_existing_record"},
                      result["skipped"])
        self.assertNotIn("I", {op["value"]["id"] for op in result["change"]["ops"]
                                if op["collection"] == "sections"})

    def test_apply_skips_held(self):
        _src, plan, lines = self._make(
            NORMAL + "\n6. 관련 논문\n중복 제목 줄.\n")
        result = split.apply_plan(self.work, plan, lines, source_path=_src)
        self.assertFalse((self.work / "sections" / "III-6.md").exists())
        self.assertNotIn("III-6",
                         {w for w in result["written"]})
        intro = (self.work / "sections" / "I.md").read_text(encoding="utf-8")
        self.assertNotIn("가나다 등(2020)", intro)
        self.assertNotIn("중복 제목 줄", intro)

    def test_reapplying_same_source_reuses_its_revision(self):
        init = run_gg("init", str(self.work))
        self.assertEqual(0, init.returncode, init.stdout + init.stderr)
        src, plan, lines = self._make()
        plan1 = dict(plan, items=[i for i in plan["items"]
                                  if i["section_id"] == "I"])
        split.apply_plan(self.work, plan1, lines, source_path=src,
                         change_out="build/first.json")
        first = run_gg("apply", str(self.work), "--change",
                       str(self.work / "build" / "first.json"),
                       "--expected-revision", "0")
        self.assertEqual(0, first.returncode, first.stdout + first.stderr)

        plan2 = dict(plan, items=[i for i in plan["items"]
                                  if i["section_id"] == "V"])
        second_result = split.apply_plan(
            self.work, plan2, lines, source_path=src,
            change_out="build/second.json")
        second_change = second_result["change"]
        self.assertNotIn("sources", [op["collection"]
                                     for op in second_change["ops"]])
        for op in second_change["ops"]:
            if op["collection"] == "sections":
                self.assertEqual(1, op["value"]["depends_on"][0]["revision"])
        second = run_gg("apply", str(self.work), "--change",
                        str(self.work / "build" / "second.json"),
                        "--expected-revision", "1")
        self.assertEqual(0, second.returncode, second.stdout + second.stderr)
        checked = run_gg("check", str(self.work))
        results = json.loads(checked.stdout)
        dependency_failures = [
            row for row in results
            if row.get("check_id") == "dependency_revision"
            and row.get("status") == "fail"
        ]
        self.assertEqual([], dependency_failures)

    def test_plan_roundtrip_via_files(self):
        src, plan, _lines = self._make()
        plan_file = Path(self._tmp.name) / "plan.json"
        plan_file.write_text(json.dumps(plan, ensure_ascii=False),
                             encoding="utf-8")
        # CLI apply 경로: 계획 파일 → 실행 (원본 해시 대조 포함)
        code = split.main(
            [str(self.work), "--apply", "--plan", str(plan_file),
             "--from", str(src),
             "--change-out", "build/c.json"])
        self.assertEqual(code, 0)
        self.assertTrue((self.work / "sections" / "V.md").exists())
        # 원본이 바뀌면 거부
        src.write_text("바뀐 원고", encoding="utf-8")
        code = split.main(
            [str(self.work), "--apply", "--plan", str(plan_file),
             "--from", str(src)])
        self.assertEqual(code, 2)


class Input(unittest.TestCase):
    def test_hwp_blocked(self):
        with tempfile.TemporaryDirectory() as d:
            p = write_src(d, "old.hwp", "bytes")
            lines, _fmt, blocked = split.read_text(p)
        self.assertIsNone(lines)
        self.assertIn("변환", blocked)

    def test_docx(self):
        try:
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("python-docx 없음 — 환경 의존")
        from docx import Document
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "draft.docx"
            doc = Document()
            doc.add_paragraph("Ⅰ. 머리말")
            doc.add_paragraph("서론 본문")
            doc.add_paragraph("1. 농업환경분석")
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "항목"
            t.cell(0, 1).text = "값"
            t.cell(1, 0).text = "면적"
            t.cell(1, 1).text = "1000㎡"
            doc.add_paragraph("Ⅴ. 맺음말")
            doc.add_paragraph("끝")
            doc.save(str(src))
            lines, fmt, blocked = split.read_text(src)
            self.assertEqual(fmt, "docx")
            toc = load_toc()
            work = Path(d) / "work"
            work.mkdir()
            plan = split.build_plan(lines, toc, src_path=src, fmt=fmt,
                                    workdir=work)
            ids = {i["section_id"] for i in plan["items"]}
            self.assertEqual(ids, {"I", "II-1", "V"})
            # 표는 어느 절 안에서도 잘리지 않고 덩어리로 남는다
            intro = next(i for i in plan["items"]
                         if i["section_id"] == "II-1")
            self.assertEqual(intro["table_blocks"], 1)


if __name__ == "__main__":
    unittest.main()
