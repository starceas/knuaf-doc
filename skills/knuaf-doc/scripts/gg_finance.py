"""Single annual crop cash-accounting profile. Values are Decimal strings in won.
No grants, processing, shared facilities, perennial yield curves, or tax model are
silently inferred. Unsupported plans are preserved and explicitly refused.
"""

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import copy
import json
from pathlib import Path
from gg_core import digest

D = Decimal
FIELDS = (
    "quantity",
    "sold",
    "loss",
    "price",
    "variable_cost",
    "fixed_cost",
    "household",
)
COMPOSITE_INPUTS = (
    "site_area",
    "facility_area",
    "net_bed_area",
    "crop_parts",
    "product_forms",
    "cycles",
    "commercialization_rate",
    "loss_rate",
    "commission",
    "transport",
    "labor",
    "repair",
    "household",
    "reinvestment",
)
INPUT_STATUSES = {
    "user_answer",
    "research",
    "explicit_assumption",
    "not_applicable",
    "unresolved",
}


def num(v):
    if not isinstance(v, str):
        raise ValueError("십진 문자열 필요")
    try:
        n = D(v)
    except InvalidOperation as e:
        raise ValueError("유효한 십진 문자열 필요") from e
    if not n.is_finite():
        raise ValueError("유한 수치 필요")
    return n


def _investment(flows, benefits, costs, rate):
    if rate <= -1:
        raise ValueError("할인율 범위 오류")
    npv = sum(v / (1 + rate) ** t for t, v in enumerate(flows))
    pv_b = sum(v / (1 + rate) ** t for t, v in enumerate(benefits))
    pv_c = sum(v / (1 + rate) ** t for t, v in enumerate(costs))
    bc = None if pv_c == 0 else pv_b / pv_c
    nonzero = [x for x in flows if x]
    changes = sum((a > 0) != (b > 0) for a, b in zip(nonzero, nonzero[1:]))
    irr, irr_status = None, (
        "undefined" if changes == 0 else "nonconventional_multiple_or_undefined"
    )
    if changes == 1:

        def value(r):
            return sum(v / (1 + r) ** t for t, v in enumerate(flows))

        lo, hi = D("-.9999"), D("1")
        while value(lo) * value(hi) > 0 and hi < D("1e12"):
            hi *= 2
        if value(lo) * value(hi) <= 0:
            for _ in range(180):
                mid = (lo + hi) / 2
                if value(lo) * value(mid) <= 0:
                    hi = mid
                else:
                    lo = mid
            irr, irr_status = (lo + hi) / 2, "calculated"
    cum, payback = flows[0], D(0) if flows[0] >= 0 else None
    for t, v in enumerate(flows[1:], 1):
        before = cum
        cum += v
        if payback is None and cum >= 0 and v > 0:
            payback = D(t - 1) + (-before) / v
    return {
        "npv": str(npv),
        "bc": str(bc) if bc is not None else None,
        "bc_status": "calculated" if bc is not None else "undefined_zero_cost",
        "irr": str(irr) if irr is not None else None,
        "irr_status": irr_status,
        "payback": str(payback) if payback is not None else None,
    }


def require_classified_workbook_inputs(workbook, *, context=None):
    """Every clear numeric/formula cell must be classified. Coordinates only.

    ``context=None`` (draft calls, e.g. from composite_contract) keeps the
    legacy behavior exactly.  An admission call passes a context dict; the
    workbook must first satisfy the legacy classification contract, then
    every classified cell's ``source_ref`` must appear in
    ``context["resolved_source_refs"]`` (a list)."""
    if not isinstance(workbook, dict):
        raise ValueError("미분류 경제 입력")
    if workbook.get("economic_classification") != "complete":
        raise ValueError("미분류 경제 입력")
    cells = workbook.get("cells")
    if not isinstance(cells, list) or not cells:
        raise ValueError("미분류 경제 입력")
    for cell in cells:
        if not isinstance(cell, dict):
            raise ValueError("미분류 경제 입력")
        if cell.get("map_action") != "clear" or cell.get("kind") not in {"number", "formula"}:
            continue
        record = cell.get("economic")
        if (
            not isinstance(record, dict)
            or record.get("status") not in INPUT_STATUSES
            or not isinstance(record.get("unit"), str)
            or not record["unit"].strip()
            or not isinstance(record.get("period"), str)
            or not record["period"].strip()
            or not isinstance(record.get("source_ref"), str)
            or not record["source_ref"].strip()
            or not isinstance(record.get("meaning_id"), str)
            or not record["meaning_id"].strip()
        ):
            raise ValueError("미분류 경제 입력")
        if record["status"] in {"user_answer", "research", "explicit_assumption"} and record.get("value") is None:
            raise ValueError("미분류 경제 입력")
        if record["status"] == "explicit_assumption" and record.get("assumption_approved") is not True:
            raise ValueError("미분류 경제 입력")
        if record["status"] == "unresolved" and record.get("value") is not None:
            raise ValueError("미분류 경제 입력")
        if record["status"] == "not_applicable" and (
            record.get("value") is not None
            or not isinstance(record.get("reason"), str)
            or not record["reason"].strip()
        ):
            raise ValueError("미분류 경제 입력")
    if context is not None:
        if not context or not isinstance(context, dict):
            raise ValueError("CONTEXT_REQUIRED")
        resolved = context.get("resolved_source_refs")
        if not isinstance(resolved, list):
            raise ValueError("MAP_SLOT_UNRESOLVED")
        for cell in cells:
            if cell.get("map_action") != "clear" or cell.get("kind") not in {"number", "formula"}:
                continue
            if cell["economic"]["source_ref"] not in resolved:
                raise ValueError("MAP_SLOT_UNRESOLVED")
            # P4 hook (stage-g005 D07): a statistical source_ref claims
            # audited provenance — the cell's economic claim (value/unit/
            # locator/pdf binding) must verify against the bound receipt,
            # not merely appear in resolved_source_refs.
            if _p4_is_statistical(cell["economic"]["source_ref"]):
                verdict = _p4_verify_economic(
                    cell["economic"], cell["economic"]["source_ref"],
                    context)
                if verdict is None:
                    raise ValueError("MAP_SLOT_UNRESOLVED")
                if verdict["status"] != "resolved":
                    raise ValueError("STATISTICAL_CLAIM_UNVERIFIED")


def validate_economic_inputs(spec, *, context=None, purpose="draft"):
    """External validation pass over the same spec shape calculate()
    accepts (design-r2 §5/§9, C-lane adapter).

    Never raises for ``purpose="draft"``/``"admission"`` — it returns a
    list of ``check_id="finance_review"`` Issue dicts, and an empty list
    always means "checked, nothing found".  Native-slot/value/declaration
    conflicts are ``status="fail"``; provenance that cannot be connected
    to the supplied context (a declared ``source_ref`` absent from
    ``context["resolved_source_refs"]``, or no usable context at all) is
    ``status="blocked"`` so existing draft calculations are preserved.
    Draft tolerates a well-formed ``"unresolved"`` economic.status;
    admission reports it as a blocking issue and itself requires a real
    context dict.  A workbook cell's slot identity is its
    ``economic.meaning_id``: two clear cells claiming the same slot with
    contradictory ``unit``/``period`` declarations are a fail conflict.
    """

    def issue(code, *, status="fail", fact_id=None, consumer_id=None,
              location=None, reason="", remedy=""):
        return {
            "code": code,
            "check_id": "finance_review",
            "status": status,
            "severity": "error",
            "fact_id": fact_id,
            "consumer_id": consumer_id,
            "location": location or {
                "path": None, "section_id": None, "line": None,
                "column": None, "table": None, "row": None, "cell": None,
            },
            "reason": reason,
            "required_for": ["submission_candidate"],
            "remedy": remedy,
        }

    def loc(*, cell=None, row=None):
        return {
            "path": None, "section_id": None, "line": None,
            "column": None, "table": None, "row": row, "cell": cell,
        }

    if purpose not in {"draft", "admission"}:
        raise ValueError("지원하지 않는 purpose")
    issues = []
    if not isinstance(spec, dict):
        return [issue(
            "ECONOMIC_SPEC_SHAPE",
            reason="경제 입력 spec이 dict가 아니어서 검사 불가",
            remedy="calculate()와 동일한 spec dict로 전달")]
    if purpose == "admission" and (
        not context or not isinstance(context, dict)
    ):
        issues.append(issue(
            "CONTEXT_REQUIRED", status="blocked",
            reason="admission 검증에는 실제 context dict가 필요",
            remedy="resolved_source_refs를 포함한 context를 전달"))
    resolved = (
        context.get("resolved_source_refs")
        if isinstance(context, dict) else None
    )

    def ref_connected(ref):
        return isinstance(resolved, list) and ref in resolved

    periods = spec.get("periods")
    if isinstance(periods, list):
        start = spec.get("start_year")
        for t, row in enumerate(periods, 1):
            if not isinstance(row, dict):
                issues.append(issue(
                    "PERIOD_ROW_SHAPE", location=loc(row=t),
                    reason="%d번째 기간 행이 dict가 아님" % t))
                continue
            vals = {}
            for k in FIELDS:
                try:
                    vals[k] = num(row[k])
                except (KeyError, ValueError):
                    issues.append(issue(
                        "PERIOD_VALUE_INVALID", location=loc(row=t),
                        reason="%d년차 행의 %s 값이 유한 십진 문자열이 아님"
                               % (t, k)))
            if len(vals) == len(FIELDS):
                if any(v < 0 for v in vals.values()):
                    issues.append(issue(
                        "PERIOD_VALUE_NEGATIVE", location=loc(row=t),
                        reason="%d년차 행에 음수 입력" % t))
                elif vals["quantity"] != vals["sold"] + vals["loss"]:
                    issues.append(issue(
                        "PERIOD_VALUE_CONFLICT", location=loc(row=t),
                        reason="%d년차 행 생산량≠판매+손실 — 재고 불일치 선언" % t))
            if (
                type(start) is int
                and row.get("year") != start + t - 1
            ):
                issues.append(issue(
                    "PERIOD_YEAR_CONFLICT", location=loc(row=t),
                    reason="%d번째 행의 선언 연도가 start_year 순서와 불일치" % t))

    inventory = spec.get("input_inventory")
    if isinstance(inventory, dict):
        fields = []
        for field in inventory:
            if isinstance(field, str):
                fields.append(field)
            else:
                issues.append(issue(
                    "ECONOMIC_FIELD_DECLARATION",
                    reason="input_inventory의 키가 문자열이 아님: %r"
                           % (field,)))
        for field in sorted(fields):
            record = inventory[field]
            if (
                not isinstance(record, dict)
                or record.get("status") not in INPUT_STATUSES
                or not isinstance(record.get("source_ref"), str)
                or not record["source_ref"].strip()
            ):
                issues.append(issue(
                    "ECONOMIC_FIELD_DECLARATION", fact_id=field,
                    reason="input_inventory['%s']의 상태/출처 선언 불완전"
                           % field))
                continue
            status = record["status"]
            if (
                status in {"user_answer", "research"}
                and record.get("value") is None
            ):
                issues.append(issue(
                    "ECONOMIC_VALUE_MISSING", fact_id=field,
                    reason="input_inventory['%s']에 값이 없음" % field))
            elif status == "explicit_assumption" and (
                record.get("value") is None
                or record.get("assumption_approved") is not True
            ):
                issues.append(issue(
                    "ECONOMIC_ASSUMPTION_UNAPPROVED", fact_id=field,
                    reason="input_inventory['%s'] 명시 가정이 승인되지 않음"
                           % field))
            elif status == "not_applicable" and (
                record.get("value") is not None
                or not isinstance(record.get("reason"), str)
                or not record["reason"].strip()
            ):
                issues.append(issue(
                    "ECONOMIC_NA_DECLARATION", fact_id=field,
                    reason="input_inventory['%s'] not_applicable 선언 불완전"
                           % field))
            elif status == "unresolved" and purpose == "admission":
                issues.append(issue(
                    "ECONOMIC_STATUS_UNRESOLVED", status="blocked",
                    fact_id=field,
                    reason="input_inventory['%s']이 unresolved — admission은 해소된 입력만 허용"
                           % field))
            if not ref_connected(record["source_ref"]):
                issues.append(issue(
                    "PROVENANCE_UNCONNECTED", status="blocked",
                    fact_id=field,
                    reason="input_inventory['%s']의 source_ref가 context.resolved_source_refs에 없음 — 출처 미연결"
                           % field))
            elif _p4_is_statistical(record["source_ref"]):
                # P4 hook (D07): statistical refs verify the claim itself
                _v = _p4_verify_economic(
                    record, record["source_ref"], context)
                if _v is None:
                    issues.append(issue(
                        "PROVENANCE_UNCONNECTED", status="blocked",
                        fact_id=field,
                        reason="input_inventory['%s']의 통계 출처 해석 "
                               "context(p4_resolver) 없음" % field))
                elif _v["status"] != "resolved":
                    issues.append(issue(
                        "STATISTICAL_CLAIM_UNVERIFIED", status="blocked",
                        fact_id=field,
                        reason="input_inventory['%s']의 통계 주장 불일치: %s"
                               % (field, _v.get("reason"))))

    wb = spec.get("workbook_inputs")
    if wb is not None:
        if (
            not isinstance(wb, dict)
            or wb.get("economic_classification") != "complete"
            or not isinstance(wb.get("cells"), list)
            or not wb["cells"]
        ):
            issues.append(issue(
                "WORKBOOK_UNCLASSIFIED",
                reason="workbook_inputs가 완전 분류된 경제 입력이 아님",
                remedy="require_classified_workbook_inputs의 분류 계약 충족"))
        else:
            slot_decls = {}
            for idx, cell in enumerate(wb["cells"]):
                if isinstance(cell, dict) and isinstance(
                    cell.get("cell"), str
                ):
                    cell_loc = loc(cell=cell["cell"])
                else:
                    cell_loc = loc(cell="cells[%d]" % idx)
                if not isinstance(cell, dict):
                    issues.append(issue(
                        "WORKBOOK_CELL_SHAPE", location=cell_loc,
                        reason="cells[%d]가 dict가 아님" % idx))
                    continue
                if (
                    cell.get("map_action") != "clear"
                    or cell.get("kind") not in {"number", "formula"}
                ):
                    continue
                record = cell.get("economic")
                if not isinstance(record, dict):
                    issues.append(issue(
                        "WORKBOOK_ECONOMIC_MISSING", location=cell_loc,
                        reason="분류 셀에 economic 레코드가 없음"))
                    continue
                for f in ("unit", "period", "meaning_id"):
                    if (
                        not isinstance(record.get(f), str)
                        or not record[f].strip()
                    ):
                        issues.append(issue(
                            "ECONOMIC_FIELD_DECLARATION", location=cell_loc,
                            reason="분류 셀의 %s 선언 불완전" % f))
                ref = record.get("source_ref")
                if not isinstance(ref, str) or not ref.strip():
                    issues.append(issue(
                        "ECONOMIC_FIELD_DECLARATION", location=cell_loc,
                        reason="분류 셀의 source_ref 선언 불완전"))
                elif not ref_connected(ref):
                    issues.append(issue(
                        "PROVENANCE_UNCONNECTED", status="blocked",
                        location=cell_loc,
                        reason="source_ref '%s'가 context.resolved_source_refs에 없음 — 출처 미연결"
                               % ref))
                elif _p4_is_statistical(ref):
                    # P4 hook (D07): statistical refs verify the claim
                    _v = _p4_verify_economic(record, ref, context)
                    if _v is None:
                        issues.append(issue(
                            "PROVENANCE_UNCONNECTED", status="blocked",
                            location=cell_loc,
                            reason="통계 출처 해석 context(p4_resolver) 없음"))
                    elif _v["status"] != "resolved":
                        issues.append(issue(
                            "STATISTICAL_CLAIM_UNVERIFIED", status="blocked",
                            location=cell_loc,
                            reason="통계 주장 불일치: %s"
                                   % _v.get("reason")))
                status = record.get("status")
                if status not in INPUT_STATUSES:
                    issues.append(issue(
                        "ECONOMIC_FIELD_DECLARATION", location=cell_loc,
                        reason="분류 셀의 status 선언이 무효"))
                elif (
                    status in {"user_answer", "research", "explicit_assumption"}
                    and record.get("value") is None
                ):
                    issues.append(issue(
                        "ECONOMIC_VALUE_MISSING", location=cell_loc,
                        reason="분류 셀에 값이 없음"))
                elif (
                    status == "explicit_assumption"
                    and record.get("assumption_approved") is not True
                ):
                    issues.append(issue(
                        "ECONOMIC_ASSUMPTION_UNAPPROVED", location=cell_loc,
                        reason="명시 가정이 승인되지 않음"))
                elif status == "unresolved":
                    if record.get("value") is not None:
                        issues.append(issue(
                            "ECONOMIC_VALUE_CONFLICT", location=cell_loc,
                            reason="unresolved 셀에 잔존 값이 남아 있음"))
                    elif purpose == "admission":
                        issues.append(issue(
                            "ECONOMIC_STATUS_UNRESOLVED", status="blocked",
                            location=cell_loc,
                            reason="unresolved 경제 입력은 admission에서 차단"))
                elif status == "not_applicable" and (
                    record.get("value") is not None
                    or not isinstance(record.get("reason"), str)
                    or not record["reason"].strip()
                ):
                    issues.append(issue(
                        "ECONOMIC_NA_DECLARATION", location=cell_loc,
                        reason="not_applicable 선언 불완전"))
                mid = record.get("meaning_id")
                if isinstance(mid, str) and mid.strip():
                    slot_decls.setdefault(mid, []).append(
                        (record.get("unit"), record.get("period")))
            for mid, decls in slot_decls.items():
                units = {u for u, _ in decls}
                spans = {p for _, p in decls}
                if len(units) > 1 or len(spans) > 1:
                    issues.append(issue(
                        "ECONOMIC_SLOT_CONFLICT",
                        reason="동일 meaning_id '%s' 슬롯에 상충 선언: unit=%s period=%s"
                               % (mid,
                                  sorted(str(u) for u in units),
                                  sorted(str(p) for p in spans))))
    return issues


def composite_contract(spec):
    """Preserve composite contracts; report incompleteness instead of raising.

    Duplicate harvest identities and unclassified workbook inputs remain hard
    errors; incomplete classification only returns an unsupported status with
    the missing inputs.  Complete inputs are reported as calculable, but no
    aggregation formula is ever invented.
    """
    crops = spec.get("crops")
    contracts = spec.get("crop_contracts")
    inventory = spec.get("input_inventory")
    missing = []
    if (
        not isinstance(crops, list)
        or len(crops) < 2
        or any(not isinstance(crop, str) or not crop.strip() for crop in crops)
        or len(set(crops)) != len(crops)
    ):
        missing.append("crops")
    if not isinstance(contracts, list) or not contracts:
        missing.append("crop_contracts")
        contracts = []
    if not isinstance(inventory, dict):
        missing.append("input_inventory")
        inventory = {}
    identities = set()
    covered_crops = set()
    for contract in contracts:
        required = (
            "crop_id",
            "planting_id",
            "part",
            "product_form",
            "cycle",
            "zone",
            "deduplication_note",
        )
        if (
            not isinstance(contract, dict)
            or any(
                not isinstance(contract.get(field), str)
                or not contract[field].strip()
                for field in required
            )
            or (isinstance(crops, list) and contract.get("crop_id") not in crops)
            or (
                contract.get("same_plant_group") is not None
                and (
                    not isinstance(contract["same_plant_group"], str)
                    or not contract["same_plant_group"].strip()
                )
            )
        ):
            missing.append("crop_contracts")
            continue
        identity = tuple(contract[field] for field in required[:-1])
        if identity in identities:
            raise ValueError("중복 수확 계약")
        identities.add(identity)
        covered_crops.add(contract["crop_id"])
    if isinstance(crops, list) and crops and covered_crops != set(crops):
        missing.append("crop_contracts")
    if "workbook_inputs" in spec:
        require_classified_workbook_inputs(spec["workbook_inputs"])
    unresolved = []
    not_applicable = []
    assumptions = []
    for field in COMPOSITE_INPUTS:
        record = inventory.get(field)
        if (
            not isinstance(record, dict)
            or record.get("status") not in INPUT_STATUSES
            or not isinstance(record.get("source_ref"), str)
            or not record["source_ref"].strip()
        ):
            missing.append(field)
            continue
        status = record["status"]
        if status in {"user_answer", "research"} and record.get("value") is None:
            missing.append(field)
        elif status == "explicit_assumption":
            if record.get("value") is None or record.get("assumption_approved") is not True:
                missing.append(field)
            else:
                assumptions.append(field)
        elif status == "not_applicable":
            if (
                record.get("value") is not None
                or not isinstance(record.get("reason"), str)
                or not record["reason"].strip()
            ):
                missing.append(field)
            else:
                not_applicable.append(field)
        elif status == "unresolved":
            # An unresolved answer is still a missing economic input: the
            # contract stays unsupported until the student resolves it.
            # The field stays in unresolved regardless of a stale value so
            # the completeness surface shows what still needs an answer.
            missing.append(field)
            unresolved.append(field)
    base = {
        "reason": "복합 작목 계약은 보존했으나 검증된 배분 산식이 없어 자동 합산하지 않음",
        "profile": spec.get("profile"),
        "input_hash": digest(spec),
        "crops": copy.deepcopy(crops) if isinstance(crops, list) else crops,
        "crop_contracts": copy.deepcopy(contracts),
        "input_inventory": copy.deepcopy(inventory),
        "input_completeness": {
            "required": list(COMPOSITE_INPUTS),
            "unresolved": unresolved,
            "not_applicable": not_applicable,
            "explicit_assumptions": assumptions,
            "calculable": not unresolved and not missing,
        },
    }
    if missing:
        return dict(
            base,
            status="unsupported",
            reason="복합 작목 입력 분류·계약이 불완전해 자동 합산하지 않음",
            missing=list(dict.fromkeys(missing)),
        )
    return dict(base, status="calculable", missing=[])


def _calculate(spec):
    if (
        isinstance(spec.get("crops"), list)
        and len(spec["crops"]) > 1
        and ("crop_contracts" in spec or "input_inventory" in spec)
    ):
        return composite_contract(spec)
    if (
        spec.get("profile") != "single_annual_cash_v1"
        or len(spec.get("crops", [])) != 1
        or any(
            spec.get(k)
            for k in ("shared_facilities", "perennial", "processing", "grants")
        )
    ):
        return {
            "status": "unsupported",
            "reason": "복합·공유·다년생·가공·보조금 자동 산식 미검증. 원계획을 보존하고 자료 정리·서술 지속.",
        }
    if (
        spec.get("accounting_basis") != "cash_pre_tax_no_inventory"
        or not spec.get("source_refs")
        or spec.get("unit") != "원"
        or spec.get("quantity_unit") != "kg"
    ):
        raise ValueError("단위·회계 모델·근거 참조 필요")
    initial = {
        k: num(spec[k])
        for k in (
            "land",
            "facility",
            "equity",
            "loan",
            "loan_rate",
            "discount_rate",
            "salvage",
        )
    }
    if any(v < 0 for v in initial.values()):
        raise ValueError("음수 초기 입력")
    financing = copy.deepcopy(spec.get("financing", {
        "loan_basis": "legacy_unspecified",
        "scenario_loan": spec["loan"],
        "actual_drawdown": None,
    }))
    if "financing" in spec:
        if not isinstance(financing, dict) or financing.get("loan_basis") not in {
            "planned", "actual",
        }:
            raise ValueError("차입 계산 기준은 예정액 또는 실제 실행액으로 명시")
        for key in ("funding_ceiling", "planned_borrowing", "actual_drawdown"):
            if key not in financing:
                raise ValueError("한도·예정액·실행액은 미정이면 null로 명시")
            if financing[key] is not None and num(financing[key]) < 0:
                raise ValueError("음수 차입 입력")
        selected = (
            "planned_borrowing" if financing["loan_basis"] == "planned"
            else "actual_drawdown"
        )
        if financing[selected] is None or num(financing[selected]) != initial["loan"]:
            raise ValueError("명시한 차입 기준과 계산 대출액 불일치")
        financing["scenario_loan"] = spec["loan"]
    years, life, grace, term = spec["years"], spec["life"], spec["grace"], spec["term"]
    if (
        any(type(v) is not int for v in (spec["start_year"], years, life, grace, term))
        or min(years, life, term) <= 0
        or grace < 0
        or years > life
        or len(spec["periods"]) != years
    ):
        raise ValueError("기간 오류 또는 교체투자 필요: 자동 지원 범위 밖")
    if spec.get("repayment") != "equal_principal":
        return {"status": "unsupported", "reason": "원금균등 이외 상환 산식 미검증"}
    if (
        spec.get("investment_basis") != "school_farm_new_business"
        or spec.get("owner_labor_in_costs") is not False
        or years != life
    ):
        return {
            "status": "unsupported",
            "reason": "학교 농가 신규사업 전체 투자·주 투자 내용연수 분석만 지원. 기존 사업 일부 투자/경영주 급여 포함 모델은 별도 검토 필요",
        }
    dep = (initial["facility"] - initial["salvage"]) / life
    if dep < 0:
        raise ValueError("잔존가치가 시설가액 초과")
    cash = initial["equity"] + initial["loan"] - initial["land"] - initial["facility"]
    debt, retained, inventory = initial["loan"], D(0), D(0)
    rows, benefits, costs = [], [D(0)], [initial["land"] + initial["facility"]]
    flows = [-costs[0]]
    household_scope = spec.get("household_scope", "owner_withdrawal")
    if household_scope not in {"owner_withdrawal", "farm_cost"}:
        raise ValueError("가계비 범위는 개인 인출 또는 농장 비용으로 명시")
    for t, raw in enumerate(spec["periods"], 1):
        x = {k: num(raw[k]) for k in FIELDS}
        if any(v < 0 for v in x.values()):
            raise ValueError("음수 생산/비용 입력")
        if raw.get("year") != spec["start_year"] + t - 1:
            raise ValueError("연도 순서 불일치")
        inventory += x["quantity"] - x["sold"] - x["loss"]
        if inventory != 0:
            raise ValueError("지원 프로필은 재고 0: 생산=판매+손실이어야 함")
        revenue = x["sold"] * x["price"]
        operating = x["quantity"] * x["variable_cost"] + x["fixed_cost"]
        withdrawal = x["household"] if household_scope == "owner_withdrawal" else D(0)
        if household_scope == "farm_cost":
            operating += x["household"]
        principal = min(debt, initial["loan"] / term) if t > grace else D(0)
        interest = debt * initial["loan_rate"]
        opening = cash
        profit = revenue - operating - dep - interest
        cash += revenue - operating - principal - interest - withdrawal
        debt -= principal
        retained += profit - withdrawal
        book = initial["facility"] - dep * t
        assets = cash + initial["land"] + book
        capital = initial["equity"] + retained
        if assets.quantize(D("1"), rounding=ROUND_HALF_UP) != (debt + capital).quantize(
            D("1"), rounding=ROUND_HALF_UP
        ):
            raise ValueError("자산=부채+자본 불일치")
        row = {
            "year": raw["year"],
            "revenue": revenue,
            "operating": operating,
            "depreciation": dep,
            "principal": principal,
            "interest": interest,
            "profit": profit,
            "opening_cash": opening,
            "cash": cash,
            "debt": debt,
            "assets": assets,
            "capital": capital,
        }
        rows.append({k: str(v) if isinstance(v, D) else v for k, v in row.items()})
        # School source: remaining book value enters terminal benefits;
        # household costs enter farm investment costs when owner pay is not
        # already included. This is not a market liquidation valuation.
        b = revenue + (initial["land"] + book if t == years else 0)
        c = operating + withdrawal
        benefits.append(b)
        costs.append(c)
        flows.append(b - c)
    inv = _investment(flows, benefits, costs, initial["discount_rate"])
    inv["repayment_horizon"] = grace + term
    inv["repayment_comparison"] = (
        "not_recovered"
        if inv["payback"] is None
        else (
            "longer_than_repayment"
            if D(inv["payback"]) > grace + term
            else "within_repayment"
        )
    )
    return {
        "status": "calculated",
        "profile": spec["profile"],
        "input_hash": digest(spec),
        "income_target": copy.deepcopy(
            spec.get("income_target", {"answer_state": "not_provided"})
        ),
        "financing": financing,
        "household_scope": household_scope,
        "household_scope_explicit": "household_scope" in spec,
        "rounding": {"unit": "원", "mode": "ROUND_HALF_UP", "display_places": 0},
        "rows": rows,
        "investment": inv,
        "risks": (
            ["cash_shortage"]
            if any(D(r["cash"]) < 0 for r in rows)
            or initial["equity"] + initial["loan"]
            < initial["land"] + initial["facility"]
            else []
        ),
        "business_approval": "not_assessed",
        "recalculation": "not_run",
    }


def calculate(spec, *, context=None):
    """Guarded public entry: returning a finance calculation needs an
    explicit major authorized against the canonical binding (policy B)."""
    import gg_major_contract as mc

    mc.authorize_output(mc.OUTPUT_FINANCE_CALCULATION, context, spec=spec)
    return _calculate(spec)


def crosscheck(result, claims):
    issues = []
    for c in claims:
        row = next((r for r in result["rows"] if r["year"] == c["year"]), None)
        if (
            row is None
            or c["unit"] != "원"
            or c["scope"] != "annual"
            or c["metric"] not in row
            or D(c["value"])
            != D(row[c["metric"]]).quantize(D("1"), rounding=ROUND_HALF_UP)
        ):
            issues.append(
                {
                    "location": c["location"],
                    "status": "fail",
                    "reason": "기간·단위·범위·출력 반올림 값 불일치",
                }
            )
    return issues


# Independent investment metrics print at bounded precision for readability
# (FINPRINT01): money rounds to whole won per result["rounding"], ratios and
# payback years to 8 decimals matching verify_recalculated's unit. The exact
# Decimal strings stay next to them in the 정확값 column and in the manifest;
# status words and None are never coerced into numbers.
INVESTMENT_DISPLAY = {
    "npv": (D("1"), "#,##0"),
    "bc": (D(".00000001"), "0.00000000"),
    "irr": (D(".00000001"), "0.00000000"),
    "payback": (D(".00000001"), "0.00000000"),
}


def workbook(spec, path, *, context=None):
    """Guarded workbook build: authorize before any output, reconfirm
    just before the file is saved (policy B)."""
    import gg_major_contract as mc

    if spec.get("profile") == "school_17_sheet_v1":
        from gg_school_excel import school_workbook

        return school_workbook(spec, path, context=context)
    authorization = mc.authorize_output(
        mc.OUTPUT_SCHOOL_WORKBOOK, context, spec=spec)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    result = _calculate(spec)
    if result["status"] != "calculated":
        raise ValueError(result["reason"])
    wb = Workbook()
    a = wb.active
    a.title = "입력"
    a.append(["항목", "값", "단위"])
    fixed = (
        "land",
        "facility",
        "equity",
        "loan",
        "loan_rate",
        "discount_rate",
        "salvage",
        "life",
        "grace",
        "term",
    )
    refs = {}
    for k in fixed:
        a.append(
            [
                k,
                float(spec[k]) if isinstance(spec[k], str) else spec[k],
                (
                    "ratio"
                    if "rate" in k
                    else "년" if k in ("life", "grace", "term") else "원"
                ),
            ]
        )
        refs[k] = "'입력'!$B$" + str(a.max_row)
    a.append(["year"] + list(FIELDS))
    first = a.max_row + 1
    for row in spec["periods"]:
        a.append([row["year"]] + [float(row[k]) for k in FIELDS])
    s = wb.create_sheet("연간계산")
    columns = [
        "year",
        "revenue",
        "operating",
        "depreciation",
        "principal",
        "interest",
        "profit",
        "opening_cash",
        "cash",
        "debt",
        "assets",
        "capital",
    ]
    s.append(columns)
    R = refs
    initial = f"{R['equity']}+{R['loan']}-{R['land']}-{R['facility']}"
    for t in range(1, spec["years"] + 1):
        r = t + 1
        ir = first + t - 1
        opening_debt = R["loan"] if t == 1 else f"J{r-1}"
        household = f"'입력'!H{ir}"
        farm_household = "+" + household if result["household_scope"] == "farm_cost" else ""
        withdrawal = household if result["household_scope"] == "owner_withdrawal" else "0"
        s.append(
            [
                spec["start_year"] + t - 1,
                f"='입력'!C{ir}*'입력'!E{ir}",
                f"='입력'!B{ir}*'입력'!F{ir}+'입력'!G{ir}{farm_household}",
                f"=({R['facility']}-{R['salvage']})/{R['life']}",
                f"=IF({t}>{R['grace']},MIN({opening_debt},{R['loan']}/{R['term']}),0)",
                f"={opening_debt}*{R['loan_rate']}",
                f"=B{r}-C{r}-D{r}-F{r}",
                "=" + initial if t == 1 else f"=I{r-1}",
                f"=H{r}+B{r}-C{r}-E{r}-F{r}-{withdrawal}",
                f"={opening_debt}-E{r}",
                f"=I{r}+{R['land']}+{R['facility']}-D{r}*{t}",
                (
                    f"={R['equity']}+G{r}-{withdrawal}"
                    if t == 1
                    else f"=L{r-1}+G{r}-{withdrawal}"
                ),
            ]
        )
    audit = wb.create_sheet("독립계산_대조전")
    audit.append(columns)
    for row in result["rows"]:
        audit.append([row[k] if k == "year" else float(row[k]) for k in columns])
    inv = wb.create_sheet("투자분석_독립값")
    inv.append(["지표", "표시값", "정확값", "주의"])
    for k, v in result["investment"].items():
        rule = INVESTMENT_DISPLAY.get(k)
        if rule is None or v is None:
            inv.append([k, v, None, "재계산 대조 전"])
            continue
        shown = float(D(v).quantize(rule[0], rounding=ROUND_HALF_UP))
        inv.append([k, shown, v, "재계산 대조 전"])
        inv.cell(inv.max_row, 2).number_format = rule[1]
    calc = wb.create_sheet("투자분석_수식")
    calc.append(
        [
            "연차",
            "순현금흐름",
            "편익",
            "비용",
            "할인편익",
            "할인비용",
            "할인현금흐름",
            "누적현금",
            "회수시점",
        ]
    )
    calc.append(
        [
            0,
            f"=C2-D2",
            0,
            f"={R['land']}+{R['facility']}",
            "=C2",
            "=D2",
            "=B2",
            "=B2",
            '=IF(B2>=0,0,"")',
        ]
    )
    for t in range(1, spec["years"] + 1):
        r = t + 2
        ar = t + 1
        ir = first + t - 1
        withdrawal = (
            f"'입력'!H{ir}" if result["household_scope"] == "owner_withdrawal" else "0"
        )
        terminal = (
            f"+{R['land']}+{R['facility']}-'연간계산'!D{ar}*{t}"
            if t == spec["years"]
            else ""
        )
        calc.append(
            [
                t,
                f"=C{r}-D{r}",
                f"='연간계산'!B{ar}{terminal}",
                f"='연간계산'!C{ar}+{withdrawal}",
                f"=C{r}/(1+{R['discount_rate']})^A{r}",
                f"=D{r}/(1+{R['discount_rate']})^A{r}",
                f"=E{r}-F{r}",
                f"=H{r-1}+B{r}",
                f'=IF(AND(H{r}>=0,H{r-1}<0,B{r}>0),A{r}-1-H{r-1}/B{r},"")',
            ]
        )
    end = calc.max_row
    calc["K1"] = "지표"
    calc["L1"] = "재계산 값"
    metrics = result["investment"]
    for r, (name, formula) in enumerate(
        [
            ("npv", f"=SUM(G2:G{end})"),
            (
                "bc",
                (
                    f"=SUM(E2:E{end})/SUM(F2:F{end})"
                    if metrics["bc"] is not None
                    else metrics["bc_status"]
                ),
            ),
            (
                "irr",
                (
                    # Excel IRR iterates from guess 0.1 and returns #NUM!
                    # when it cannot converge (e.g. a deep negative IRR).
                    # Retry seeded with the independently computed root;
                    # a single sign change gives one real root, so the
                    # seed can only affect convergence, not the answer.
                    f"=IFERROR(IRR(B2:B{end}),IRR(B2:B{end},{metrics['irr']}))"
                    if metrics["irr"] is not None
                    else metrics["irr_status"]
                ),
            ),
            (
                "payback",
                (
                    f"=MIN(I2:I{end})"
                    if metrics["payback"] is not None
                    else "not_recovered"
                ),
            ),
        ],
        2,
    ):
        calc.cell(r, 11, name)
        calc.cell(r, 12, formula)
        calc.cell(r, 12).number_format = "#,##0" if name == "npv" else "0.00000000"
    meta = wb.create_sheet("검증상태")
    meta.append(["상태", "검토 전. 재계산·모든 인쇄 페이지 미검증"])
    meta.append(["input_hash", result["input_hash"]])

    source_refs = ", ".join(str(ref) for ref in spec["source_refs"])
    source_kind = (
        "합성 입력"
        if all(str(ref).startswith("synthetic:") for ref in spec["source_refs"])
        else "등록 입력"
    )
    assumptions = result.get("composite", {}).get("explicit_assumptions", [])
    assumption_text = ", ".join(assumptions) if assumptions else "명시된 입력 외 추가 가정 없음"

    def add_printed_note(sheet, role):
        """Put provenance below the actual printed table, never in a cell note."""
        row = sheet.max_row + 2
        end = get_column_letter(sheet.max_column)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=sheet.max_column)
        cell = sheet.cell(row, 1)
        cell.value = (
            f"출처/입력: {source_kind} ({source_refs}) | "
            f"기간: {spec['start_year']}~{spec['start_year'] + spec['years'] - 1} | "
            f"가정: {assumption_text} | 계산 역할: {role}"
        )
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(size=9, italic=True)
        sheet.row_dimensions[row].height = 42
        # This is intentionally before print_area is set below: the note must
        # be printed under the table rather than hidden in workbook metadata.
        return f"A{row}:{end}{row}"

    add_printed_note(a, "원값과 단위의 입력 표")
    add_printed_note(s, "입력값으로 계산한 연간 재무 결과")
    add_printed_note(audit, "독립 계산 대조 전의 기준 행")
    add_printed_note(inv, "투자지표의 표시값과 정확값 대조")
    add_printed_note(calc, "투자지표 산식과 할인·회수 계산")
    for sheet in wb:
        sheet.freeze_panes = "B2"
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions
        for col in sheet.columns:
            if sheet.max_column > 8:
                # FINPRINT02: bounded width — long English headers
                # (depreciation, opening_cash) must not wrap into
                # neighbouring interest/profit labels.
                header = str(col[0].value or "")
                sheet.column_dimensions[col[0].column_letter].width = min(
                    16, max(11, len(header) + 2)
                )
            else:
                sheet.column_dimensions[col[0].column_letter].width = 24
        for row in sheet:
            for cell in row:
                cell.font = Font(size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if sheet.title in {"연간계산", "독립계산_대조전"} and cell.column > 1:
                    cell.number_format = "#,##0"
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(size=11, bold=True)
    meta.column_dimensions["B"].width = 80
    # 정확값 (column C) keeps the full Decimal string on one line (28 digits).
    inv.column_dimensions["C"].width = 34
    path = Path(path)
    if path.exists():
        raise ValueError("기존 XLSX 덮어쓰기 금지")
    mc.reconfirm_output(authorization, context)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    path.with_suffix(".manifest.json").write_text(
        json.dumps(
            dict(
                result,
                file_hash=digest(path.read_bytes()),
                major_authorization=authorization.to_dict(),
            ),
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return result


def verify_recalculated(spec, path):
    if spec.get("profile") == "school_17_sheet_v1":
        from gg_school_verify import verify_school_recalculated

        return verify_school_recalculated(spec, path)
    from openpyxl import load_workbook

    result = _calculate(spec)
    if result["status"] != "calculated":
        return [{"status": "unsupported", "reason": result["reason"]}]
    wb = load_workbook(path, data_only=True)
    if "연간계산" not in wb:
        return [{"status": "fail", "reason": "연간계산 시트 누락"}]
    s = wb["연간계산"]
    issues = []
    keys = list(result["rows"][0])
    expected_last_row = len(result["rows"]) + 1
    has_provenance_note = (
        s.max_row == expected_last_row + 2
        and isinstance(s.cell(expected_last_row + 2, 1).value, str)
        and s.cell(expected_last_row + 2, 1).value.startswith("출처/입력:")
    )
    if [c.value for c in s[1]] != keys or (
        s.max_row != expected_last_row and not has_provenance_note
    ):
        return [{"cell": "연간계산!A1", "status": "fail", "reason": "연간계산 열/행 구조 변경"}]
    for r, row in enumerate(result["rows"], 2):
        for col, key in enumerate(keys, 1):
            actual = s.cell(r, col).value
            if type(actual) not in (int, float) or not D(str(actual)).is_finite():
                issues.append(
                    {
                        "cell": f"연간계산!{s.cell(r,col).coordinate}",
                        "status": "blocked",
                        "reason": "재계산 캐시 없음/오류",
                    }
                )
            elif D(str(actual)).quantize(D("1"), rounding=ROUND_HALF_UP) != D(
                str(row[key])
            ).quantize(D("1"), rounding=ROUND_HALF_UP):
                issues.append(
                    {
                        "cell": f"연간계산!{s.cell(r,col).coordinate}",
                        "status": "fail",
                        "reason": "독립 계산과 출력 반올림 불일치",
                    }
                )
    if "투자분석_수식" not in wb:
        issues.append({"status": "fail", "reason": "투자분석 수식 시트 누락"})
        return issues
    investment_sheet = wb["투자분석_수식"]
    for r, key in enumerate(("npv", "bc", "irr", "payback"), 2):
        actual = investment_sheet.cell(r, 12).value
        expected = result["investment"][key]
        location = f"투자분석_수식!L{r}"
        if investment_sheet.cell(r, 11).value != key:
            issues.append(
                {"cell": location, "status": "fail", "reason": "투자지표 구조 변경"}
            )
        elif expected is None:
            status = result["investment"].get(key + "_status", "not_recovered")
            if actual != status:
                issues.append(
                    {
                        "cell": location,
                        "status": "fail",
                        "reason": "산정 불가 상태 은폐/변경",
                    }
                )
        elif type(actual) not in (int, float) or not D(str(actual)).is_finite():
            issues.append(
                {
                    "cell": location,
                    "status": "blocked",
                    "reason": "투자지표 재계산 캐시 없음/오류",
                }
            )
        else:
            unit = D("1") if key == "npv" else D(".00000001")
            if D(str(actual)).quantize(unit, rounding=ROUND_HALF_UP) != D(
                expected
            ).quantize(unit, rounding=ROUND_HALF_UP):
                issues.append(
                    {
                        "cell": location,
                        "status": "fail",
                        "reason": "투자지표 독립계산 반올림 불일치",
                    }
                )
    return issues


# ---------------------------------------------------------------------------
# P4 verification adapter (stage-g005 D07/D09): audited statistical
# source_refs resolve through the provenance chain in gg_rda_research.
# Produces the ``resolved_source_refs`` list that ``context`` consumers
# (require_classified_workbook_inputs / validate_economic_inputs) already
# membership-check — statistical refs that cannot be anchored are reported
# as conflicts and never reach the resolved list.  No calculation or
# formula path is touched by this section.


def resolve_economic_source_refs(source_refs, *, resolver=None):
    """P4 adapter: resolve audited statistical source_refs.

    ``resolver`` is an optional ``gg_rda_research.resolver_context()`` dict;
    omitted -> the deployed packs context is built lazily.  Returns
    ``{"resolved_source_refs": [...], "conflicts": [...]}`` for callers to
    splice into the admission ``context`` — a declared statistical ref that
    fails resolution is absent from ``resolved_source_refs``, so the
    existing ``MAP_SLOT_UNRESOLVED``/``ref_connected`` checks fail closed.
    """
    import gg_rda_research

    if not isinstance(source_refs, list):
        raise ValueError("source_refs 목록 필요")
    return gg_rda_research.resolve_statistical_refs(
        source_refs, context=resolver)


def _p4_is_statistical(ref):
    """True when a source_ref claims audited statistical provenance."""
    try:
        import gg_rda_research
    except ImportError:
        return False
    return gg_rda_research.is_statistical_ref(ref)


def _p4_verify_economic(record, ref, context):
    """P4 claim verifier for the finance lane (D07): cross-checks the
    cell/record's economic claim (value/unit/locator/pdf binding) against
    the audited observation its source_ref resolves to.

    ``context["p4_resolver"]`` supplies the verified resolver context —
    a statistical ref with no resolver cannot be verified and returns
    ``None`` (fail closed upstream)."""
    import gg_rda_research

    resolver = (context or {}).get("p4_resolver")
    if not isinstance(resolver, dict):
        return None
    try:
        return gg_rda_research.verify_economic_claim(
            record, ref, context=resolver)
    except Exception as exc:  # verification never crashes admission
        return {"status": "invalid",
                "reason": "통계 주장 검증 실패: %s" % exc}
