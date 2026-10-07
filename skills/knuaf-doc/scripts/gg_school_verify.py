"""School 17-sheet workbook recalculation and accounting verification.

Validates cached recalculation values, formula integrity, 5-year balance equality,
cash rollforwards, repayment nonnegativity, depreciation floors, and sales/profit links.
Reports missing caches as 'blocked', accounting/formula errors as 'fail', and unsupported
plans as 'unsupported'. Narrowly permits #N/A only in 2. 중장기영농목표 C12:G12 when corresponding
year sales is verified zero.
"""

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from gg_finance import num
from gg_price_assumptions import multiplier_d as price_multiplier_d
from gg_school_excel import SCHOOL_SHEETS, check_unsupported, validate

D = Decimal
TOLERANCE = Decimal("0.05")
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!", "#N/A")
GOAL_RATIO_CELLS = {"C12", "D12", "E12", "F12", "G12"}


def calculate_school_expected(spec):
    """Computes independent 5-year accounting and financial schedules in Decimal."""
    years, v, price = validate(spec)
    land = v["land"]
    facility = v["facility"]
    equipment = v["equipment"]
    equity = v["equity"]
    loan = v["loan"]
    loan_rate = v["loan_rate"]
    salvage = v["salvage"]
    materials = v["materials"]
    labor = v["labor"]
    packing = v["packing"]
    transport = v["transport"]
    household = v["household"]
    repair_facility_rate = v["repair_facility_rate"]
    repair_equipment_rate = v["repair_equipment_rate"]
    utility_per_10a = v["utility_per_10a"]
    area_m2 = v["area_m2"]
    production_kg = v["production_kg"]
    purchase_price = v["purchase_price"]
    direct_price = v["direct_price"]
    purchase_share = v["purchase_share"]
    direct_share = v["direct_share"]
    grace = spec["grace"]
    term = spec["term"]
    life = spec["life"]

    pur_hist = spec.get("purchase_price_history")
    if pur_hist and len(pur_hist) == 5:
        pur_avg = sum(num(str(p)) for p in pur_hist) / D(5)
        if abs(purchase_price - pur_avg) > Decimal("0.05"):
            raise ValueError(
                f"수매 평균가격({purchase_price})과 과거 시세 평균({pur_avg}) 불일치"
            )
        effective_purchase_price = pur_avg
    else:
        effective_purchase_price = purchase_price

    dir_hist = spec.get("direct_price_history")
    if dir_hist and len(dir_hist) == 5:
        dir_avg = sum(num(str(p)) for p in dir_hist) / D(5)
        if abs(direct_price - dir_avg) > Decimal("0.05"):
            raise ValueError(
                f"직거래 평균가격({direct_price})과 과거 시세 평균({dir_avg}) 불일치"
            )
        effective_direct_price = dir_avg
    else:
        effective_direct_price = direct_price

    tot_loan = loan
    fac_loan = min(tot_loan, facility)
    rem_loan = tot_loan - fac_loan
    eq_loan = min(rem_loan, equipment)
    rem_loan2 = rem_loan - eq_loan
    land_loan = min(rem_loan2, land)

    fac_self = facility - fac_loan
    eq_self = equipment - eq_loan
    land_self = land - land_loan
    total_capex = facility + equipment + land

    # 세 상승률 분리: 일반물가→자재·수도광열비, 임금→노무비, 판매가→매출.
    # 이중 적용 금지(노무총액에 일반물가 재적용 없음), 감가상각·고정 원리금 제외.
    mult_g = [price_multiplier_d(price["general"], y) for y in years]
    mult_w = [price_multiplier_d(price["wage"], y) for y in years]
    mult_s = [price_multiplier_d(price["sales"], y) for y in years]

    opening_cash = equity
    purchase_rev = (production_kg * purchase_share * effective_purchase_price) / D(1000)
    direct_rev = (production_kg * direct_share * effective_direct_price) / D(1000)
    fac_dep_annual = (facility - salvage) / life if life > 0 else D(0)
    eq_dep_annual = equipment / life if life > 0 else D(0)

    cur_cash = opening_cash
    cur_loan = loan
    cur_equity = equity
    cur_fac_book = facility
    cur_eq_book = equipment

    repair_fac = facility * repair_facility_rate
    repair_eq = equipment * repair_equipment_rate
    utility = utility_per_10a * (area_m2 / D(1000))
    sga = packing + transport

    expected_years = []
    cum_cash = D(0)

    for i, year in enumerate(years):
        t = i + 1
        if t <= grace or cur_loan <= 0 or loan <= 0 or term <= 0:
            principal = D(0)
        else:
            principal = min(cur_loan, loan / term)
        interest = cur_loan * loan_rate
        ending_loan = cur_loan - principal

        fac_dep = fac_dep_annual if t <= life else D(0)
        eq_dep = eq_dep_annual if t <= life else D(0)
        tot_dep = fac_dep + eq_dep

        cur_fac_book -= fac_dep
        cur_eq_book -= eq_dep

        yr_mat = materials * mult_g[i]
        yr_lab = labor * mult_w[i]
        yr_util = utility * mult_g[i]
        yr_rep = repair_fac + repair_eq

        yr_purchase_rev = purchase_rev * mult_s[i]
        yr_direct_rev = direct_rev * mult_s[i]
        rev = yr_purchase_rev + yr_direct_rev

        prod_cash_cost = yr_mat + yr_lab + yr_util + yr_rep
        cogs = prod_cash_cost + tot_dep

        gross_profit = rev - cogs
        operating_profit = gross_profit - sga
        net_profit = operating_profit - interest

        if t == 1:
            cf_in = rev + loan
            cf_out = total_capex + prod_cash_cost + sga + interest + principal
        else:
            cf_in = rev
            cf_out = prod_cash_cost + sga + interest + principal
        net_cf = cf_in - cf_out
        cum_cash += net_cf

        cur_cash += net_cf
        cur_equity += net_profit

        total_assets = cur_cash + land + cur_fac_book + cur_eq_book
        total_liab_eq = ending_loan + cur_equity

        expected_years.append(
            {
                "year": year,
                "revenue": rev,
                "purchase_rev": yr_purchase_rev,
                "direct_rev": yr_direct_rev,
                "materials": yr_mat,
                "labor": yr_lab,
                "utility": yr_util,
                "repairs": yr_rep,
                "prod_cash_cost": prod_cash_cost,
                "tot_dep": tot_dep,
                "fac_dep": fac_dep,
                "eq_dep": eq_dep,
                "cogs": cogs,
                "sga": sga,
                "operating_profit": operating_profit,
                "interest": interest,
                "principal": principal,
                "net_profit": net_profit,
                "household": household,
                "estimated_income": net_profit - household,
                "net_cf": net_cf,
                "cum_cf": cum_cash,
                "cash": cur_cash,
                "fac_book": cur_fac_book,
                "eq_book": cur_eq_book,
                "total_assets": total_assets,
                "ending_loan": ending_loan,
                "equity": cur_equity,
                "total_liab_eq": total_liab_eq,
            }
        )
        cur_loan = ending_loan

    return {
        "years": years,
        "values": v,
        "expected_years": expected_years,
        "effective_purchase_price": effective_purchase_price,
        "effective_direct_price": effective_direct_price,
        "tot_loan": tot_loan,
        "fac_loan": fac_loan,
        "eq_loan": eq_loan,
        "land_loan": land_loan,
        "fac_self": fac_self,
        "eq_self": eq_self,
        "land_self": land_self,
        "capex_self": total_capex - tot_loan,
        "total_capex": total_capex,
    }


def is_valid_unused_grade_cell(formula_val, data_val):
    """True only for exact allowed values in both views (no whitespace normalization)."""
    allowed_values = {None, "", "[해당 없음: 상품만 판매]"}

    if formula_val is not None:
        if formula_val not in allowed_values:
            return False

    if data_val is not None:
        if data_val not in allowed_values:
            return False

    return True


def to_d(val, cell_name=""):
    if val is None:
        raise ValueError(f"{cell_name}: 값 누락")
    if isinstance(val, bool):
        raise ValueError(f"{cell_name}: 잘못된 타입 (bool)")
    if isinstance(val, (int, float)):
        d = D(str(val))
        if not d.is_finite():
            raise ValueError(f"{cell_name}: 비유한 값 ({val})")
        return d
    if isinstance(val, D):
        if not val.is_finite():
            raise ValueError(f"{cell_name}: 비유한 값 ({val})")
        return val
    if isinstance(val, str):
        try:
            d = D(val)
        except InvalidOperation:
            raise ValueError(f"{cell_name}: 수치 아님 ({val})")
        if not d.is_finite():
            raise ValueError(f"{cell_name}: 비유한 값 ({val})")
        return d
    raise ValueError(f"{cell_name}: 잘못된 타입 ({type(val)})")


def verify_school_recalculated(spec, path):
    """Verifies that the recalculated school workbook satisfies all accounting and recalculation rules."""
    unsupported = check_unsupported(spec)
    if unsupported:
        return [unsupported]

    try:
        expected = calculate_school_expected(spec)
    except Exception as e:
        return [{"status": "fail", "reason": f"입력 명세 검증 실패: {e}"}]

    path = Path(path)
    if not path.exists():
        return [{"status": "fail", "reason": f"파일 없음: {path}"}]

    try:
        wb_formula = load_workbook(path, data_only=False)
        wb_data = load_workbook(path, data_only=True)
    except Exception as e:
        return [{"status": "fail", "reason": f"엑셀 파일 로드 실패: {e}"}]

    issues = []

    # 1. 17-sheet name and sequence check
    if tuple(wb_formula.sheetnames) != SCHOOL_SHEETS:
        issues.append(
            {
                "status": "fail",
                "reason": "17시트 이름·순서 불일치",
                "observed": list(wb_formula.sheetnames),
            }
        )
        return issues

    # 2. Hidden state check for model sheet
    if wb_formula["참고1. 모델농장분석"].sheet_state != "hidden":
        issues.append(
            {
                "sheet": "참고1. 모델농장분석",
                "status": "fail",
                "reason": "모델농장분석 숨김 아님",
            }
        )

    # 3. Scan prohibited model sheet formula links across all sheets
    for sheetname in SCHOOL_SHEETS:
        ws_f = wb_formula[sheetname]
        for row in ws_f.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    if (
                        "참고1. 모델농장분석" in cell.value
                        or "모델농장분석" in cell.value
                    ):
                        issues.append(
                            {
                                "cell": f"{sheetname}!{cell.coordinate}",
                                "status": "fail",
                                "reason": "숨김 모델농장분석 시트 수식 참조 금지",
                            }
                        )

    # 4. Check formula caches and error tokens across wb_data
    for sheetname in SCHOOL_SHEETS:
        ws_f = wb_formula[sheetname]
        ws_d = wb_data[sheetname]
        for row in ws_f.iter_rows():
            for cell in row:
                f_val = cell.value
                d_val = ws_d.cell(row=cell.row, column=cell.column).value
                coord = cell.coordinate
                location = f"{sheetname}!{coord}"

                is_formula = isinstance(f_val, str) and f_val.startswith("=")
                if is_formula:
                    if d_val is None or (
                        isinstance(d_val, str) and d_val.startswith("=")
                    ):
                        issues.append(
                            {
                                "cell": location,
                                "status": "blocked",
                                "reason": "재계산 캐시 없음/미실행",
                            }
                        )
                        continue

                if isinstance(d_val, str):
                    for token in ERROR_TOKENS:
                        if token in d_val:
                            if (
                                token == "#N/A"
                                and sheetname == "2. 중장기영농목표"
                                and coord in GOAL_RATIO_CELLS
                            ):
                                col_idx = cell.column - 3
                                yr_rev = expected["expected_years"][col_idx]["revenue"]
                                if yr_rev == 0:
                                    break
                            issues.append(
                                {
                                    "cell": location,
                                    "status": "fail",
                                    "reason": f"수식 오류 발생: {d_val}",
                                }
                            )
                            break

    # 5. Financial invariants & numerical crosschecks
    # Validate all critical numeric caches that are present, even if unrelated caches are missing.
    bs = wb_data["13. 추정대차대조표"] if "13. 추정대차대조표" in wb_data else None
    cf = wb_data["14. 현금흐름계획"] if "14. 현금흐름계획" in wb_data else None
    base = wb_data["1. 기초재무상태조사"] if "1. 기초재무상태조사" in wb_data else None
    repay = wb_data["4. 원리금상환계획"] if "4. 원리금상환계획" in wb_data else None
    dep = wb_data["10. 감가상각비계획 "] if "10. 감가상각비계획 " in wb_data else None
    sales = wb_data[" 5. 판매계획"] if " 5. 판매계획" in wb_data else None
    prod = wb_data["6. 생산계획"] if "6. 생산계획" in wb_data else None
    pl = wb_data["12 손익계획"] if "12 손익계획" in wb_data else None
    cost = wb_data["11. 생산원가계획"] if "11. 생산원가계획" in wb_data else None
    goal = wb_data["2. 중장기영농목표"] if "2. 중장기영농목표" in wb_data else None
    inc = wb_data["15.추정소득분석"] if "15.추정소득분석" in wb_data else None

    # 5.1 Sales & Production Quantities
    if prod and sales:
        try:
            prod_kg = to_d(prod["H9"].value, "6. 생산계획!H9")
            purchase_kg = to_d(sales["E22"].value, " 5. 판매계획!E22")
            direct_kg = to_d(sales["E26"].value, " 5. 판매계획!E26")
            loss_kg = to_d(sales["E30"].value, " 5. 판매계획!E30")
            total_kg = to_d(sales["E31"].value, " 5. 판매계획!E31")

            if abs((purchase_kg + direct_kg + loss_kg) - prod_kg) > TOLERANCE:
                issues.append(
                    {
                        "cell": " 5. 판매계획!E31",
                        "status": "fail",
                        "reason": "채널별 판매량 합계가 생산량과 불일치",
                    }
                )
            if abs(total_kg - prod_kg) > TOLERANCE:
                issues.append(
                    {
                        "cell": " 5. 판매계획!E31",
                        "status": "fail",
                        "reason": "총 판매수량이 생산량과 불일치",
                    }
                )
            # F16: Validate unused grade price cells (rows 8, 9, 11, 12 in E/G/I/K/M and O) are not fabricated
            ws_sales_f = wb_formula[" 5. 판매계획"]
            ws_sales_d = wb_data[" 5. 판매계획"]
            hist_cols = ["E", "G", "I", "K", "M"]
            for r in (8, 9, 11, 12):
                for c_col in hist_cols:
                    f_v = ws_sales_f[f"{c_col}{r}"].value
                    d_v = ws_sales_d[f"{c_col}{r}"].value
                    if not is_valid_unused_grade_cell(f_v, d_v):
                        issues.append(
                            {
                                "cell": f" 5. 판매계획!{c_col}{r}",
                                "status": "fail",
                                "reason": "미사용 등급 시세 임의 수식/수치 조작 금지",
                            }
                        )
                f_o = ws_sales_f[f"O{r}"].value
                for c_col in ("F", "H", "J", "L", "N"):
                    f_sp = ws_sales_f[f"{c_col}{r}"].value
                    d_sp = ws_sales_d[f"{c_col}{r}"].value
                    if f_sp not in (None, "") or d_sp not in (None, ""):
                        issues.append(
                            {
                                "cell": f" 5. 판매계획!{c_col}{r}",
                                "status": "fail",
                                "reason": "미사용 등급 스페이서 열 비어 있어야 함",
                            }
                        )
                d_o = ws_sales_d[f"O{r}"].value
                if not is_valid_unused_grade_cell(f_o, d_o):
                    issues.append(
                        {
                            "cell": f" 5. 판매계획!O{r}",
                            "status": "fail",
                            "reason": "미사용 등급 평균가격 임의 수식/수치 조작 금지",
                        }
                    )

            o7_price = to_d(sales["O7"].value, " 5. 판매계획!O7")
            o10_price = to_d(sales["O10"].value, " 5. 판매계획!O10")
            if abs(o7_price - expected["effective_purchase_price"]) > TOLERANCE:
                issues.append(
                    {
                        "cell": " 5. 판매계획!O7",
                        "status": "fail",
                        "reason": "수매 평균가격 불일치",
                    }
                )
            if abs(o10_price - expected["effective_direct_price"]) > TOLERANCE:
                issues.append(
                    {
                        "cell": " 5. 판매계획!O10",
                        "status": "fail",
                        "reason": "직거래 평균가격 불일치",
                    }
                )
            # Used-grade 5-year history caches must match the spec series
            # (fabricated observations with a tampered mean must not verify).
            hist_cols = ["E", "G", "I", "K", "M"]
            pur_hist = spec.get("purchase_price_history")
            dir_hist = spec.get("direct_price_history")
            if isinstance(pur_hist, list) and len(pur_hist) == 5:
                for i, col in enumerate(hist_cols):
                    c = to_d(
                        sales[f"{col}7"].value, f" 5. 판매계획!{col}7"
                    )
                    if abs(c - D(str(pur_hist[i]))) > TOLERANCE:
                        issues.append(
                            {
                                "cell": f" 5. 판매계획!{col}7",
                                "status": "fail",
                                "reason": "수매 가격 이력 캐시 불일치",
                            }
                        )
            if isinstance(dir_hist, list) and len(dir_hist) == 5:
                for i, col in enumerate(hist_cols):
                    c = to_d(
                        sales[f"{col}10"].value, f" 5. 판매계획!{col}10"
                    )
                    if abs(c - D(str(dir_hist[i]))) > TOLERANCE:
                        issues.append(
                            {
                                "cell": f" 5. 판매계획!{col}10",
                                "status": "fail",
                                "reason": "직거래 가격 이력 캐시 불일치",
                            }
                        )
        except (ValueError, KeyError, TypeError, ArithmeticError) as e:
            issues.append(
                {"status": "fail", "reason": f"독립계산 불가/비수치 캐시: {e}"}
            )
    # 5.2 Depreciation Schedule
    if dep:
        try:
            fac_cost = to_d(dep["E8"].value, "10. 감가상각비계획 !E8")
            eq_cost = to_d(dep["I8"].value, "10. 감가상각비계획 !I8")
            salvage_val = expected["values"]["salvage"]

            cum_fac_dep = D(0)
            cum_eq_dep = D(0)
            for t in range(5):
                r = 12 + t
                exp_yr = expected["expected_years"][t]
                f_dep = to_d(dep.cell(r, 5).value, f"10. 감가상각비계획 !E{r}")
                e_dep = to_d(dep.cell(r, 9).value, f"10. 감가상각비계획 !I{r}")

                if f_dep < D(0) or e_dep < D(0):
                    issues.append(
                        {
                            "cell": f"10. 감가상각비계획 !E{r}",
                            "status": "fail",
                            "reason": "감가상각비 음수 발생",
                        }
                    )
                if abs(f_dep - exp_yr["fac_dep"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"10. 감가상각비계획 !E{r}",
                            "status": "fail",
                            "reason": "시설 감가상각비 독립계산 불일치",
                        }
                    )
                if abs(e_dep - exp_yr["eq_dep"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"10. 감가상각비계획 !I{r}",
                            "status": "fail",
                            "reason": "대농기구 감가상각비 독립계산 불일치",
                        }
                    )

                cum_fac_dep += f_dep
                cum_eq_dep += e_dep

            if cum_fac_dep > (fac_cost - salvage_val) + Decimal("0.0001"):
                issues.append(
                    {
                        "cell": "10. 감가상각비계획 !E12:E16",
                        "status": "fail",
                        "reason": "누적 시설감가상각비가 상각가능액 초과",
                    }
                )
            if cum_eq_dep > eq_cost + Decimal("0.0001"):
                issues.append(
                    {
                        "cell": "10. 감가상각비계획 !I12:I16",
                        "status": "fail",
                        "reason": "누적 농기구감가상각비가 취득가액 초과",
                    }
                )
        except (ValueError, KeyError, TypeError, ArithmeticError) as e:
            issues.append(
                {"status": "fail", "reason": f"독립계산 불가/비수치 캐시: {e}"}
            )

    # 5.3 Repayment Schedule
    if repay:
        try:
            prev_ending_loan = to_d(repay["C8"].value, "4. 원리금상환계획!C8")
            for t in range(5):
                r = 8 + t
                exp_yr = expected["expected_years"][t]
                principal = to_d(repay.cell(r, 4).value, f"4. 원리금상환계획!D{r}")
                interest = to_d(repay.cell(r, 5).value, f"4. 원리금상환계획!E{r}")
                ending_loan = to_d(repay.cell(r, 6).value, f"4. 원리금상환계획!F{r}")

                if principal < D(0):
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!D{r}",
                            "status": "fail",
                            "reason": "원금상환액 음수 발생",
                        }
                    )
                if principal > prev_ending_loan + Decimal("0.0001"):
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!D{r}",
                            "status": "fail",
                            "reason": "원금상환액이 기초 대출잔액 초과",
                        }
                    )
                if interest < D(0):
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!E{r}",
                            "status": "fail",
                            "reason": "이자비용 음수 발생",
                        }
                    )
                if ending_loan < D(0):
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!F{r}",
                            "status": "fail",
                            "reason": "연말융자잔액 음수 발생",
                        }
                    )
                if abs(ending_loan - (prev_ending_loan - principal)) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!F{r}",
                            "status": "fail",
                            "reason": "연말융자잔액 롤포워드 불일치",
                        }
                    )

                if abs(principal - exp_yr["principal"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!D{r}",
                            "status": "fail",
                            "reason": "원금상환액 독립계산 불일치",
                        }
                    )
                if abs(interest - exp_yr["interest"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!E{r}",
                            "status": "fail",
                            "reason": "이자비용 독립계산 불일치",
                        }
                    )
                if abs(ending_loan - exp_yr["ending_loan"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"4. 원리금상환계획!F{r}",
                            "status": "fail",
                            "reason": "연말융자잔액 독립계산 불일치",
                        }
                    )

                prev_ending_loan = ending_loan
        except (ValueError, KeyError, TypeError, ArithmeticError) as e:
            issues.append(
                {"status": "fail", "reason": f"독립계산 불가/비수치 캐시: {e}"}
            )

    # 5.4 Income Statement, Production Cost, Medium-term Goal & Income Analysis
    if pl and sales and cost and goal and inc:
        try:
            for t in range(5):
                col_idx = 3 + t
                col_letter = get_column_letter(col_idx)
                exp_yr = expected["expected_years"][t]

                sales_rev = to_d(
                    sales.cell(50, 4 + t).value,
                    f" 5. 판매계획!{get_column_letter(4+t)}50",
                )
                pl_rev_prod = to_d(
                    pl.cell(7, col_idx).value, f"12 손익계획!{col_letter}7"
                )
                pl_rev_tot = to_d(
                    pl.cell(6, col_idx).value, f"12 손익계획!{col_letter}6"
                )
                cost_cogs = to_d(
                    cost.cell(35, col_idx).value, f"11. 생산원가계획!{col_letter}35"
                )
                pl_cogs = to_d(pl.cell(9, col_idx).value, f"12 손익계획!{col_letter}9")
                pl_gross = to_d(
                    pl.cell(10, col_idx).value, f"12 손익계획!{col_letter}10"
                )
                pl_sga = to_d(pl.cell(11, col_idx).value, f"12 손익계획!{col_letter}11")
                pl_op = to_d(pl.cell(15, col_idx).value, f"12 손익계획!{col_letter}15")
                pl_net = to_d(pl.cell(26, col_idx).value, f"12 손익계획!{col_letter}26")

                # Revenue crosscheck
                if abs(pl_rev_prod - sales_rev) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"12 손익계획!{col_letter}7",
                            "status": "fail",
                            "reason": "손익 생산물매출이 판매계획과 불일치",
                        }
                    )
                if abs(pl_rev_tot - exp_yr["revenue"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"12 손익계획!{col_letter}6",
                            "status": "fail",
                            "reason": "손익 총매출 독립계산 불일치",
                        }
                    )

                # COGS crosscheck
                if abs(pl_cogs - cost_cogs) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"12 손익계획!{col_letter}9",
                            "status": "fail",
                            "reason": "손익 매출원가가 생산원가계획과 불일치",
                        }
                    )
                if abs(pl_cogs - exp_yr["cogs"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"12 손익계획!{col_letter}9",
                            "status": "fail",
                            "reason": "손익 매출원가 독립계산 불일치",
                        }
                    )

                # Gross & Operating & Net profit
                if abs(pl_gross - (pl_rev_tot - pl_cogs)) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"12 손익계획!{col_letter}10",
                            "status": "fail",
                            "reason": "매출총이익 산식 불일치",
                        }
                    )
                if abs(pl_op - (pl_gross - pl_sga)) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"12 손익계획!{col_letter}15",
                            "status": "fail",
                            "reason": "영업이익 산식 불일치",
                        }
                    )
                if abs(pl_net - exp_yr["net_profit"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"12 손익계획!{col_letter}26",
                            "status": "fail",
                            "reason": "당기순이익 독립계산 불일치",
                        }
                    )

                # Goal and Income analysis links
                goal_rev = to_d(
                    goal.cell(10, col_idx).value, f"2. 중장기영농목표!{col_letter}10"
                )
                goal_net = to_d(
                    goal.cell(11, col_idx).value, f"2. 중장기영농목표!{col_letter}11"
                )
                inc_rev = to_d(
                    inc.cell(4, col_idx).value, f"15.추정소득분석!{col_letter}4"
                )
                inc_net = to_d(
                    inc.cell(6, col_idx).value, f"15.추정소득분석!{col_letter}6"
                )
                inc_est = to_d(
                    inc.cell(8, col_idx).value, f"15.추정소득분석!{col_letter}8"
                )
                inc_house = to_d(
                    inc.cell(7, col_idx).value, f"15.추정소득분석!{col_letter}7"
                )

                if abs(goal_rev - pl_rev_tot) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"2. 중장기영농목표!{col_letter}10",
                            "status": "fail",
                            "reason": "중장기목표 매출이 손익계획과 불일치",
                        }
                    )
                if abs(goal_net - pl_net) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"2. 중장기영농목표!{col_letter}11",
                            "status": "fail",
                            "reason": "중장기목표 순이익이 손익계획과 불일치",
                        }
                    )
                if abs(inc_rev - pl_rev_tot) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"15.추정소득분석!{col_letter}4",
                            "status": "fail",
                            "reason": "추정소득분석 매출이 손익계획과 불일치",
                        }
                    )
                if abs(inc_net - pl_net) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"15.추정소득분석!{col_letter}6",
                            "status": "fail",
                            "reason": "추정소득분석 당기순이익이 손익계획과 불일치",
                        }
                    )
                if abs(inc_est - (inc_net - inc_house)) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"15.추정소득분석!{col_letter}8",
                            "status": "fail",
                            "reason": "추정소득 산식(순이익-가계비) 불일치",
                        }
                    )
                if abs(inc_house - exp_yr["household"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"15.추정소득분석!{col_letter}7",
                            "status": "fail",
                            "reason": "추정소득분석 가계비 독립계산 불일치",
                        }
                    )
                if abs(inc_est - exp_yr["estimated_income"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"15.추정소득분석!{col_letter}8",
                            "status": "fail",
                            "reason": "추정소득 독립계산 불일치",
                        }
                    )

                # Goal ratio
                goal_ratio_val = goal.cell(12, col_idx).value
                if exp_yr["revenue"] == 0:
                    if goal_ratio_val != "#N/A":
                        issues.append(
                            {
                                "cell": f"2. 중장기영농목표!{col_letter}12",
                                "status": "fail",
                                "reason": "매출 0일 때 순이익률 #N/A 미표시",
                            }
                        )
                else:
                    goal_ratio = to_d(
                        goal_ratio_val, f"2. 중장기영농목표!{col_letter}12"
                    )
                    expected_ratio = exp_yr["net_profit"] / exp_yr["revenue"]
                    if abs(goal_ratio - expected_ratio) > TOLERANCE:
                        issues.append(
                            {
                                "cell": f"2. 중장기영농목표!{col_letter}12",
                                "status": "fail",
                                "reason": "순이익률 독립계산 불일치",
                            }
                        )
        except (ValueError, KeyError, TypeError, ArithmeticError) as e:
            issues.append(
                {"status": "fail", "reason": f"독립계산 불가/비수치 캐시: {e}"}
            )
    # 5.5 Cash Flow Schedule & Rollforward
    if cf and base:
        try:
            prev_cum_cf = D(0)
            base_cash = to_d(base["C9"].value, "1. 기초재무상태조사!C9")
            prev_bs_cash = base_cash

            for t in range(5):
                col_idx = 4 + t
                col_letter = get_column_letter(col_idx)
                exp_yr = expected["expected_years"][t]

                cf_in = to_d(
                    cf.cell(14, col_idx).value, f"14. 현금흐름계획!{col_letter}14"
                )
                cf_out = to_d(
                    cf.cell(22, col_idx).value, f"14. 현금흐름계획!{col_letter}22"
                )
                cf_net = to_d(
                    cf.cell(23, col_idx).value, f"14. 현금흐름계획!{col_letter}23"
                )
                cf_cum = to_d(
                    cf.cell(24, col_idx).value, f"14. 현금흐름계획!{col_letter}24"
                )

                if abs(cf_net - (cf_in - cf_out)) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"14. 현금흐름계획!{col_letter}23",
                            "status": "fail",
                            "reason": "현금수지균형(유입-유출) 산식 불일치",
                        }
                    )

                if t == 0:
                    if abs(cf_cum - cf_net) > TOLERANCE:
                        issues.append(
                            {
                                "cell": f"14. 현금흐름계획!{col_letter}24",
                                "status": "fail",
                                "reason": "1년차 누적현금수지 불일치",
                            }
                        )
                else:
                    if abs(cf_cum - (prev_cum_cf + cf_net)) > TOLERANCE:
                        issues.append(
                            {
                                "cell": f"14. 현금흐름계획!{col_letter}24",
                                "status": "fail",
                                "reason": "누적현금수지 롤포워드 불일치",
                            }
                        )

                if abs(cf_net - exp_yr["net_cf"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"14. 현금흐름계획!{col_letter}23",
                            "status": "fail",
                            "reason": "당해 현금수지 독립계산 불일치",
                        }
                    )
                if abs(cf_cum - exp_yr["cum_cf"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"14. 현금흐름계획!{col_letter}24",
                            "status": "fail",
                            "reason": "누적 현금수지 독립계산 불일치",
                        }
                    )

                prev_cum_cf = cf_cum
        except (ValueError, KeyError, TypeError, ArithmeticError) as e:
            issues.append(
                {"status": "fail", "reason": f"독립계산 불가/비수치 캐시: {e}"}
            )
    # 5.6 Five-Year Balance Sheet Reconciliation & Invariants
    if bs and base:
        try:
            base_cash = to_d(base["C9"].value, "1. 기초재무상태조사!C9")
            prev_bs_cash = base_cash
            salvage_val = expected["values"]["salvage"]
            for t in range(5):
                col_idx = 3 + t
                col_letter = get_column_letter(col_idx)
                exp_yr = expected["expected_years"][t]

                bs_cash = to_d(
                    bs.cell(6, col_idx).value, f"13. 추정대차대조표!{col_letter}6"
                )
                bs_land = to_d(
                    bs.cell(7, col_idx).value, f"13. 추정대차대조표!{col_letter}7"
                )
                bs_fac = to_d(
                    bs.cell(8, col_idx).value, f"13. 추정대차대조표!{col_letter}8"
                )
                bs_eq = to_d(
                    bs.cell(9, col_idx).value, f"13. 추정대차대조표!{col_letter}9"
                )
                bs_tot_assets = to_d(
                    bs.cell(10, col_idx).value, f"13. 추정대차대조표!{col_letter}10"
                )
                bs_loan = to_d(
                    bs.cell(12, col_idx).value, f"13. 추정대차대조표!{col_letter}12"
                )
                bs_capital = to_d(
                    bs.cell(13, col_idx).value, f"13. 추정대차대조표!{col_letter}13"
                )
                bs_tot_liab_eq = to_d(
                    bs.cell(14, col_idx).value, f"13. 추정대차대조표!{col_letter}14"
                )
                bs_diff = to_d(
                    bs.cell(16, col_idx).value, f"13. 추정대차대조표!{col_letter}16"
                )

                # Component additions
                if (
                    abs(bs_tot_assets - (bs_cash + bs_land + bs_fac + bs_eq))
                    > TOLERANCE
                ):
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}10",
                            "status": "fail",
                            "reason": "자산합계(현금+토지+시설+대농기구) 불일치",
                        }
                    )
                if abs(bs_tot_liab_eq - (bs_loan + bs_capital)) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}14",
                            "status": "fail",
                            "reason": "부채와자본 합계(융자잔액+자본) 불일치",
                        }
                    )

                # Balance Equality: Assets == Liabilities + Equity (exact in thousand won)
                if abs(bs_tot_assets - bs_tot_liab_eq) > Decimal("0.0001") or abs(
                    bs_diff
                ) > Decimal("0.0001"):
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}16",
                            "status": "fail",
                            "reason": f"{exp_yr['year']}년차 대차 불일치 (자산 {bs_tot_assets} != 부채+자본 {bs_tot_liab_eq})",
                        }
                    )

                # Floors
                if bs_fac < salvage_val:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}8",
                            "status": "fail",
                            "reason": "시설 장부가액이 잔존가치 미만",
                        }
                    )
                if bs_eq < D(0):
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}9",
                            "status": "fail",
                            "reason": "대농기구 장부가액 음수 발생",
                        }
                    )
                if bs_loan < D(0):
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}12",
                            "status": "fail",
                            "reason": "대차대조표 융자잔액 음수 발생",
                        }
                    )

                # Cash connection between Cash Flow and Balance Sheet
                if cf:
                    try:
                        cf_net_yr = to_d(
                            cf.cell(23, 4 + t).value,
                            f"14. 현금흐름계획!{get_column_letter(4+t)}23",
                        )
                        if abs(bs_cash - (prev_bs_cash + cf_net_yr)) > TOLERANCE:
                            issues.append(
                                {
                                    "cell": f"13. 추정대차대조표!{col_letter}6",
                                    "status": "fail",
                                    "reason": "대차대조표 현금 잔액과 현금흐름 수지균형 롤포워드 불일치",
                                }
                            )
                    except (ValueError, KeyError, TypeError, ArithmeticError) as e:
                        issues.append(
                            {
                                "cell": f"13. 추정대차대조표!{col_letter}6",
                                "status": "fail",
                                "reason": f"독립계산 불가/비수치 캐시: {e}",
                            }
                        )
                prev_bs_cash = bs_cash

                # Comparison with expected independent values
                if abs(bs_cash - exp_yr["cash"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}6",
                            "status": "fail",
                            "reason": "대차대조표 현금 독립계산 불일치",
                        }
                    )
                if abs(bs_fac - exp_yr["fac_book"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}8",
                            "status": "fail",
                            "reason": "시설 장부가액 독립계산 불일치",
                        }
                    )
                if abs(bs_eq - exp_yr["eq_book"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}9",
                            "status": "fail",
                            "reason": "대농기구 장부가액 독립계산 불일치",
                        }
                    )
                if abs(bs_tot_assets - exp_yr["total_assets"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}10",
                            "status": "fail",
                            "reason": "자산총계 독립계산 불일치",
                        }
                    )
                if abs(bs_loan - exp_yr["ending_loan"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}12",
                            "status": "fail",
                            "reason": "대차대조표 융자잔액 독립계산 불일치",
                        }
                    )
                if abs(bs_capital - exp_yr["equity"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}13",
                            "status": "fail",
                            "reason": "대차대조표 자본총계 독립계산 불일치",
                        }
                    )
                if abs(bs_tot_liab_eq - exp_yr["total_liab_eq"]) > TOLERANCE:
                    issues.append(
                        {
                            "cell": f"13. 추정대차대조표!{col_letter}14",
                            "status": "fail",
                            "reason": "부채와자본 총계 독립계산 불일치",
                        }
                    )
        except (ValueError, KeyError, TypeError, ArithmeticError) as e:
            issues.append(
                {"status": "fail", "reason": f"독립계산 불가/비수치 캐시: {e}"}
            )
    return issues
