#!/usr/bin/env python3
"""Compare explicit narrative channel/year/share/price claims to cached XLSX cells.

A reviewed locator map supplies workbook semantics; no major's cell addresses
or channel aliases are silently imposed on another major.
"""
import argparse
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import json
from pathlib import Path
import re

YEAR = re.compile(r"(?<!\d)(\d{4})년")
START = re.compile(r"시작(?!가)|개시|진출|도입")
NUMBER = r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"
SHARE = re.compile(r"(" + NUMBER + r")\s*%")
PRICE = re.compile(r"(" + NUMBER + r")\s*원\s*/\s*kg", re.I)


def _number(value):
    if isinstance(value, bool) or value is None:
        raise ValueError('missing/non-numeric cached value')
    try:
        result = Decimal(str(value).replace(',', ''))
    except InvalidOperation as error:
        raise ValueError('non-numeric cached value') from error
    if not result.is_finite():
        raise ValueError('non-finite cached value')
    return result


def _cell(workbook, ref):
    if not isinstance(ref, dict) or set(ref) != {'sheet', 'cell'}:
        raise ValueError('locator requires sheet and cell')
    sheet, address = ref['sheet'], ref['cell']
    if not isinstance(sheet, str) or not isinstance(address, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*', address):
        raise ValueError('invalid worksheet/cell locator')
    if sheet not in workbook.sheetnames:
        raise ValueError('worksheet not found: ' + sheet)
    return _number(workbook[sheet][address].value)


def _direct_year_prefix(text, aliases, metric):
    """Accept only a year followed by an optional channel/metric label.

    This is a deliberately narrow input grammar, not prose interpretation.
    Intervening claims or connecting prose cannot prove year attribution.
    """
    channel = '(?:' + '|'.join(re.escape(alias) for alias in aliases) + ')'
    labels = {'start': '', 'share': '(?:판매비율|비율|비중)', 'price': '(?:단가|가격)'}
    label = labels[metric]
    pattern = r'\s*(?:' + channel + r'\s*)?'
    if label:
        pattern += '(?:' + label + r'\s*)?'
    return re.fullmatch(pattern + r'[:：]?\s*', text) is not None


def crosscheck(text, workbook_path, mapping):
    """Return observed comparisons; zero observed claims is not a pass.

    mapping={channels:[{name,aliases:[...],years:{YYYY:{share:{sheet,cell},
    price:{sheet,cell}}}}]}. share is a fraction, price is won/kg.
    """
    from openpyxl import load_workbook

    issues, comparisons = [], []
    channels = mapping.get('channels') if isinstance(mapping, dict) else None
    if not isinstance(channels, list) or not channels:
        raise ValueError('channels must be a non-empty list')
    plans = []
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        names, aliases_seen = set(), set()
        for channel in channels:
            name, aliases, years = channel['name'], channel['aliases'], channel['years']
            if not isinstance(name, str) or not name or name in names:
                raise ValueError('channel names must be unique')
            names.add(name)
            if not isinstance(aliases, list) or not aliases or any(not isinstance(a, str) or not a for a in aliases):
                raise ValueError('channel aliases must be non-empty strings')
            if aliases_seen.intersection(aliases):
                raise ValueError('channel aliases must not overlap')
            aliases_seen.update(aliases)
            if not isinstance(years, dict) or not years:
                raise ValueError('complete planning years required')
            values = {}
            for year, locators in years.items():
                if not re.fullmatch(r'\d{4}', year):
                    raise ValueError('planning year must have four digits')
                share = _cell(workbook, locators['share'])
                if not 0 <= share <= 1:
                    raise ValueError('share must be a fraction between 0 and 1')
                values[int(year)] = {'share': share, 'locators': locators}
                if 'price' in locators:
                    values[int(year)]['price'] = _cell(workbook, locators['price'])
            ordered = sorted(values)
            if ordered != list(range(ordered[0], ordered[-1] + 1)):
                raise ValueError('planning years must be contiguous')
            start = next((year for year in ordered if values[year]['share'] > 0), None)
            plans.append((name, aliases, values, start))
    finally:
        workbook.close()

    for line_number, line in enumerate(text.splitlines(), 1):
        matches = [plan for plan in plans if any(alias in line for alias in plan[1])]
        if not matches:
            continue
        if len(matches) != 1:
            issues.append({'line': line_number, 'status': 'not_checked', 'reason': 'multiple channels in one line; split claims'})
            continue
        name, aliases, values, start = matches[0]
        years = list(YEAR.finditer(line))
        # The year-first extractor cannot attribute preceding metrics. Keep
        # every such candidate visible even when another claim matches.
        if years:
            prefix = line[:years[0].start()]
            for kind, pattern in (('start', START), ('share', SHARE), ('price', PRICE)):
                for metric in pattern.finditer(prefix):
                    issues.append({
                        'line': line_number, 'channel': name, 'metric': kind,
                        'quote': line, 'status': 'not_checked',
                        'reason': 'metric precedes first year; explicit year attribution required',
                        'unparsed': metric[0],
                    })
        for index, match in enumerate(years):
            year = int(match[1])
            following = years[index + 1] if index + 1 < len(years) else None
            segment = line[match.end():following.start() if following else len(line)]
            claims = []
            for kind, pattern in (('start', START), ('share', SHARE), ('price', PRICE)):
                for metric in pattern.finditer(segment):
                    tail = segment[metric.end():].rstrip()
                    closed = not tail or tail.endswith((',', ';', '|', '.', '。'))
                    if following and (
                        not _direct_year_prefix(segment[:metric.start()], aliases, kind) or not closed
                    ):
                        issues.append({
                            'line': line_number, 'channel': name, 'metric': kind,
                            'quote': line, 'status': 'not_checked',
                            'reason': 'metric between explicit years; year attribution is ambiguous',
                            'candidate_years': [year, int(following[1])],
                            'unparsed': metric[0],
                        })
                        continue
                    if kind == 'start':
                        claims.append((kind, Decimal(year), None))
                    else:
                        claims.append((kind, _number(metric[1]), metric[1]))
            for kind, actual, display in claims:
                if kind == 'start':
                    expected = Decimal(start) if start is not None else None
                    locator = {str(y): value['locators']['share'] for y, value in values.items()}
                else:
                    record = values.get(year, {})
                    expected = record.get(kind)
                    locator = record.get('locators', {}).get(kind)
                    if kind == 'share' and expected is not None:
                        expected *= 100
                    if expected is not None:
                        places = len(display.split('.')[1]) if '.' in display else 0
                        expected = expected.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
                result = {'line': line_number, 'channel': name, 'year': year, 'metric': kind,
                          'actual': str(actual), 'expected': str(expected) if expected is not None else None,
                          'locator': locator, 'quote': line}
                result['status'] = 'blocked' if expected is None else ('pass' if actual == expected else 'fail')
                comparisons.append(result)
                if result['status'] != 'pass':
                    issues.append(result)
    status = ('fail' if any(i['status'] == 'fail' for i in issues) else
              'blocked' if any(i['status'] == 'blocked' for i in issues) else
              'not_checked' if not comparisons or issues else 'pass')
    return {'status': status, 'comparisons': comparisons, 'issues': issues}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--body', required=True)
    parser.add_argument('--xlsx', required=True)
    parser.add_argument('--map', required=True, dest='locator_map')
    args = parser.parse_args(argv)
    try:
        result = crosscheck(Path(args.body).read_text(encoding='utf-8'), args.xlsx,
                            json.loads(Path(args.locator_map).read_text(encoding='utf-8')))
    except (OSError, ValueError, KeyError, TypeError) as error:
        result = {'status': 'blocked', 'reason': str(error)}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result['status'] == 'pass' else (1 if result['status'] == 'fail' else 2)


if __name__ == '__main__':
    raise SystemExit(main())
