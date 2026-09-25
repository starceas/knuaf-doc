"""School 17-sheet workbook. Sheet names follow the exemplar formula refs."""

import json
import hashlib
import re
import zipfile
from decimal import Context, Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from gg_core import digest, local
from gg_finance import num

# Irregular names are formula targets. Do not "fix" spaces or periods.
SCHOOL_SHEETS = (
    "목록",
    "1. 기초재무상태조사",
    "2. 중장기영농목표",
    "3.투자계획",
    "4. 원리금상환계획",
    " 5. 판매계획",
    "6. 생산계획",
    "7. 영농자재소요계획",
    "8. 노무비계획",
    "9 .경비계획",
    "10. 감가상각비계획 ",
    "11. 생산원가계획",
    "12 손익계획",
    "13. 추정대차대조표",
    "14. 현금흐름계획",
    "15.추정소득분석",
    "참고1. 모델농장분석",
)
YEAR_LABELS = "가나다라마"
EVIDENCE_KEYS = ("area", "seedlings", "yield", "growth", "commodity")


def q(name):
    return "'" + name.replace("'", "''") + "'"


def plan_years(writing_year):
    if (
        type(writing_year) is not int
        or isinstance(writing_year, bool)
        or writing_year < 1900
        or writing_year > 2100
    ):
        raise ValueError("작성연도 정수 필요")
    start = writing_year + 1
    return [start + i for i in range(5)]


def check_unsupported(spec):
    if spec.get("profile") != "school_17_sheet_v1":
        return None
    if any(spec.get(k) for k in ("processing", "perennial", "shared_facilities")):
        return {
            "status": "unsupported",
            "profile": "school_17_sheet_v1",
            "reason": "가공·다년생·공유시설 자동 산식 미검증. 원계획을 보존합니다.",
        }
    grants = spec.get("grants")
    if grants is not None and grants != 0 and grants != "0" and grants != 0.0:
        return {
            "status": "unsupported",
            "profile": "school_17_sheet_v1",
            "reason": "보조금 자동 산식 미검증. 원계획을 보존합니다.",
        }
    if "crops" in spec and len(spec["crops"]) > 1:
        return {
            "status": "unsupported",
            "profile": "school_17_sheet_v1",
            "reason": "다작목 자동 산식 미검증. 원계획을 보존합니다.",
        }
    if "period_data" in spec or "yearly_data" in spec:
        return {
            "status": "unsupported",
            "profile": "school_17_sheet_v1",
            "reason": "연차별 변동 계획 자동 산식 미검증. 원계획을 보존합니다.",
        }
    return None


def validate(spec, *, resolver=None):
    if spec.get("profile") != "school_17_sheet_v1":
        raise ValueError("school_17_sheet_v1 만 지원")
    if not spec.get("source_refs"):
        raise ValueError("근거 참조 필요")
    if not isinstance(spec.get("source_refs"), list) or not all(
        isinstance(r, str) and r.strip() for r in spec["source_refs"]
    ):
        raise ValueError("근거 참조 목록 필요")
    # P4 reference-fill hook (stage-g005 D09): refs claiming audited
    # statistical provenance (pack@rev#record / audit_key) must resolve
    # through the provenance chain — fabricated refs fail closed.
    stat_result = resolve_statistical_source_refs(
        spec["source_refs"], resolver=resolver)
    if stat_result["conflicts"]:
        raise ValueError(
            "통계 근거 참조 해석 실패: "
            + "; ".join(
                "%s (%s)" % (c["ref"], c["reason"])
                for c in stat_result["conflicts"]
            )
        )
    if spec.get("unit") != "천원" or spec.get("quantity_unit") != "kg":
        raise ValueError("학교 제출 엑셀 단위는 천원, 수량은 kg")

    # F14: strict nonblank string validation for production evidence
    evidence = spec.get("production_evidence")
    if not isinstance(evidence, dict):
        raise ValueError("생산량 산출근거 딕셔너리 필요")
    for k in EVIDENCE_KEYS:
        val = evidence.get(k)
        if not isinstance(val, str) or not val.strip():
            raise ValueError(f"생산량 산출근거 문자열 필요: {k}")

    fields = (
        "area_m2",
        "production_kg",
        "purchase_price",
        "direct_price",
        "purchase_share",
        "direct_share",
        "land",
        "facility",
        "equipment",
        "equity",
        "loan",
        "loan_rate",
        "salvage",
        "materials",
        "labor",
        "packing",
        "transport",
        "household",
        "repair_facility_rate",
        "repair_equipment_rate",
        "utility_per_10a",
        "inflation",
    )
    values = {k: num(spec[k]) for k in fields}

    # F6: Domain constraints
    if values["area_m2"] <= 0:
        raise ValueError("재배면적은 양수 필요")
    if values["production_kg"] < 0:
        raise ValueError("생산량은 0 이상 필요")
    if values["purchase_price"] < 0 or values["direct_price"] < 0:
        raise ValueError("단가는 0 이상 필요")
    if values["purchase_share"] < 0 or values["purchase_share"] > 1:
        raise ValueError("수매 판매비율은 0과 1 사이 필요")
    if values["direct_share"] < 0 or values["direct_share"] > 1:
        raise ValueError("직거래 판매비율은 0과 1 사이 필요")
    if values["purchase_share"] + values["direct_share"] > 1:
        raise ValueError("판매비율 합계 100% 초과")

    for k in (
        "land",
        "facility",
        "equipment",
        "equity",
        "loan",
        "salvage",
        "materials",
        "labor",
        "packing",
        "transport",
        "household",
        "loan_rate",
        "repair_facility_rate",
        "repair_equipment_rate",
        "utility_per_10a",
    ):
        if values[k] < 0:
            raise ValueError(f"{k} 항목은 0 이상 필요")

    if values["inflation"] <= 0:
        raise ValueError("물가상승률은 양수 필요")

    # F4: salvage bound
    if values["salvage"] > values["facility"]:
        raise ValueError("잔존가액은 시설가액 이하 필요")

    # F5: Loan funding allocation bound
    total_capex = values["facility"] + values["equipment"] + values["land"]
    if values["loan"] > total_capex:
        raise ValueError("융자금은 총 투자액(시설+대농기구+토지) 이하 필요")
    if values["equity"] + values["loan"] < total_capex:
        raise ValueError("자부담+융자금은 총 투자액 이상 필요")

    # Price history validation if supplied
    has_pur_hist = (
        "purchase_price_history" in spec and spec["purchase_price_history"] is not None
    )
    has_dir_hist = (
        "direct_price_history" in spec and spec["direct_price_history"] is not None
    )
    if has_pur_hist or has_dir_hist:
        hist_years = spec.get("price_history_years")
        if not isinstance(hist_years, (list, tuple)) or len(hist_years) != 5:
            raise ValueError(
                "가격 이력이 있는 경우 price_history_years는 5개 연도 목록 필요"
            )
        for y in hist_years:
            if type(y) is not int or isinstance(y, bool):
                raise ValueError("price_history_years는 정수 연도 필요")
        for i in range(4):
            if hist_years[i] >= hist_years[i + 1]:
                raise ValueError(
                    "price_history_years는 엄격히 오름차순 정렬된 5개 연도 필요"
                )

    for hist_key in ("purchase_price_history", "direct_price_history"):
        hist = spec.get(hist_key)
        if hist is not None:
            if not isinstance(hist, (list, tuple)) or len(hist) != 5:
                raise ValueError(f"{hist_key}는 5개 연도 관측치 목록 필요")
            for p in hist:
                if num(str(p)) < 0:
                    raise ValueError(f"{hist_key} 가격은 0 이상 필요")
    years = plan_years(spec["writing_year"])
    for key in ("grace", "term", "life"):
        if (
            type(spec.get(key)) is not int
            or isinstance(spec.get(key), bool)
            or spec[key] < 0
        ):
            raise ValueError("거치·상환·내용연수 0 이상의 정수 필요")
    if spec["life"] <= 0:
        raise ValueError("내용연수는 1 이상의 정수 필요")
    if values["loan"] > 0 and spec["term"] <= 0:
        raise ValueError("융자금이 있는 경우 상환기간은 1 이상의 정수 필요")

    return years, values


def school_workbook(spec, path, *, context=None):
    """Guarded school workbook build: authorize before any output,
    reconfirm just before the file is saved (policy B)."""
    import gg_major_contract as mc

    authorization = mc.authorize_output(
        mc.OUTPUT_SCHOOL_WORKBOOK, context, spec=spec)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    unsupported = check_unsupported(spec)
    if unsupported:
        return unsupported

    years, v = validate(spec)
    start = years[0]
    wb = Workbook()
    wb.active.title = SCHOOL_SHEETS[0]
    for name in SCHOOL_SHEETS[1:]:
        wb.create_sheet(name)
    (
        index,
        base,
        goal,
        invest,
        repay,
        sales,
        prod,
        mat,
        labor,
        exp,
        dep,
        cost,
        pl,
        bs,
        cf,
        inc,
        model,
    ) = (wb[name] for name in SCHOOL_SHEETS)
    model.sheet_state = "hidden"

    index["B2"] = "목록"
    index["B4"] = "동반제출 재무계획 시트 목록. 숨김 모델농장분석은 수식 연결하지 않음."
    for i, name in enumerate(SCHOOL_SHEETS, 1):
        index.cell(5 + i, 2, i)
        index.cell(5 + i, 3, name)

    # F1: Pre-investment opening basis (assets/loans are zero before Year 1 capex)
    base["B2"] = "1. 기초재무상태조사"
    base["I2"] = "(단위:천원)"
    base["B4"] = "토지"
    base["C4"] = 0
    base["B5"] = "시설"
    base["C5"] = 0
    base["B6"] = "대농기구"
    base["C6"] = 0
    base["B7"] = "자기자본"
    base["C7"] = float(v["equity"])
    base["B8"] = "융자"
    base["C8"] = 0
    base["B9"] = "기초현금"
    base["C9"] = "=C7+C8-C4-C5-C6"
    base["D4"] = "창업 전 기초재무상태. 신규 투자는 투자계획에 기록."

    prod["B2"] = "6. 생산계획"
    prod["B4"] = "가. 생산량계획"
    prod["B7"] = "10a당 생산량(kg)"
    prod["H8"] = float(v["area_m2"])
    prod["B8"] = "재배면적(㎡)"
    prod["B9"] = "생산량(kg)"
    prod["H9"] = float(v["production_kg"])
    prod["H7"] = "=H9/H8*1000"
    prod["B11"] = "나. 생산량 산출근거"
    evidence = spec["production_evidence"]
    prod["B15"] = "재배면적"
    prod["H15"] = evidence["area"]
    prod["B17"] = "묘목수"
    prod["H17"] = evidence["seedlings"]
    prod["B19"] = "생산량"
    prod["H19"] = evidence["yield"]
    prod["B23"] = "생산증가율"
    prod["H23"] = evidence["growth"]
    prod["B24"] = "상품화율"
    prod["H24"] = evidence["commodity"]
    prod["B25"] = "합계"
    for i, year in enumerate(years):
        prod.cell(25, 5 + i, "=H9")
        prod.cell(4, 5 + i, year if i == 0 else f"={get_column_letter(4 + i)}4+1")
    prod["E4"] = start

    sales_name = q(" 5. 판매계획")
    prod_name = q("6. 생산계획")
    dep_name = q("10. 감가상각비계획 ")
    invest_name = q("3.투자계획")
    repay_name = q("4. 원리금상환계획")
    cost_name = q("11. 생산원가계획")
    pl_name = q("12 손익계획")
    mat_name = q("7. 영농자재소요계획")
    labor_name = q("8. 노무비계획")
    exp_name = q("9 .경비계획")
    cf_name = q("14. 현금흐름계획")
    base_name = q("1. 기초재무상태조사")

    # F8: 5. 판매계획 - 가. 연도별 가격 추이 (5 dated observations & visible citations)
    sales["B2"] = "5. 판매계획"
    sales["B4"] = "가. 연도별 가격 추이"
    sales["I4"] = "(단위 : 원/kg)"
    sales["B6"] = "연도별"
    sales["C6"] = "등급"
    sales["O6"] = "평균가격"

    hist_cols = ["E", "G", "I", "K", "M"]
    hist_years = spec.get("price_history_years")
    if hist_years and len(hist_years) == 5:
        for i, yr in enumerate(hist_years):
            sales[f"{hist_cols[i]}6"] = yr
    sales["B7"] = "평균가격(수매)"
    sales["C7"] = "상 품"
    pur_hist = spec.get("purchase_price_history")
    if pur_hist and len(pur_hist) == 5:
        for i, val in enumerate(pur_hist):
            c_letter = hist_cols[i]
            sales[f"{c_letter}7"] = float(num(str(val)))
        sales["O7"] = "=AVERAGE(E7,G7,I7,K7,M7)"
    else:
        sales["O7"] = float(v["purchase_price"])
    sales["C8"] = "중 품"
    sales["O8"] = "[해당 없음: 상품만 판매]"
    sales["C9"] = "하 품"
    sales["O9"] = "[해당 없음: 상품만 판매]"

    sales["B10"] = "평균가격(직거래)"
    sales["C10"] = "상 품"
    dir_hist = spec.get("direct_price_history")
    if dir_hist and len(dir_hist) == 5:
        for i, val in enumerate(dir_hist):
            c_letter = hist_cols[i]
            sales[f"{c_letter}10"] = float(num(str(val)))
        sales["O10"] = "=AVERAGE(E10,G10,I10,K10,M10)"
    else:
        sales["O10"] = float(v["direct_price"])
    sales["C11"] = "중 품"
    sales["O11"] = "[해당 없음: 상품만 판매]"
    sales["C12"] = "하 품"
    sales["O12"] = "[해당 없음: 상품만 판매]"
    sales["B14"] = f"출처: {', '.join(source_ref_labels(spec['source_refs']))}"

    # F9: 나. 유통채널별 판매계획 (5 year pairs across columns D..M)
    sales["B19"] = "나. 유통채널별 판매계획"
    sales["I19"] = "(단위 : %, kg)"
    sales["B21"] = "유통채널"
    sales["C21"] = "등급"
    sales["B22"] = "수매시장"
    sales["C22"] = "상품"
    sales["C23"] = "중품"
    sales["C24"] = "하품"
    sales["C25"] = "계"
    sales["B26"] = "직거래"
    sales["C26"] = "상품"
    sales["C27"] = "중품"
    sales["C28"] = "하품"
    sales["C29"] = "계"
    sales["B30"] = "미상품과"
    sales["B31"] = "합계"

    for t in range(5):
        col_pct = get_column_letter(4 + 2 * t)
        col_qty = get_column_letter(5 + 2 * t)
        prev_pct = get_column_letter(4 + 2 * (t - 1)) if t > 0 else None

        sales[f"{col_pct}20"] = f"{years[t]}년"
        sales[f"{col_pct}21"] = "판매비율"
        sales[f"{col_qty}21"] = "판매량"

        # 수매
        sales[f"{col_pct}22"] = (
            float(v["purchase_share"]) if t == 0 else f"={prev_pct}22"
        )
        sales[f"{col_qty}22"] = f"={col_pct}22*{prod_name}!$H$9"
        sales[f"{col_pct}23"] = 0
        sales[f"{col_qty}23"] = f"={col_pct}23*{prod_name}!$H$9"
        sales[f"{col_pct}24"] = 0
        sales[f"{col_qty}24"] = f"={col_pct}24*{prod_name}!$H$9"
        sales[f"{col_pct}25"] = f"=SUM({col_pct}22:{col_pct}24)"
        sales[f"{col_qty}25"] = f"=SUM({col_qty}22:{col_qty}24)"

        # 직거래
        sales[f"{col_pct}26"] = float(v["direct_share"]) if t == 0 else f"={prev_pct}26"
        sales[f"{col_qty}26"] = f"={col_pct}26*{prod_name}!$H$9"
        sales[f"{col_pct}27"] = 0
        sales[f"{col_qty}27"] = f"={col_pct}27*{prod_name}!$H$9"
        sales[f"{col_pct}28"] = 0
        sales[f"{col_qty}28"] = f"={col_pct}28*{prod_name}!$H$9"
        sales[f"{col_pct}29"] = f"=SUM({col_pct}26:{col_pct}28)"
        sales[f"{col_qty}29"] = f"=SUM({col_qty}26:{col_qty}28)"

        # 미상품과 & 합계
        sales[f"{col_pct}30"] = f"=1-{col_pct}25-{col_pct}29"
        sales[f"{col_qty}30"] = f"={col_pct}30*{prod_name}!$H$9"
        sales[f"{col_pct}31"] = f"={col_pct}25+{col_pct}29+{col_pct}30"
        sales[f"{col_qty}31"] = f"={col_qty}25+{col_qty}29+{col_qty}30"

        # Apply percent formatting on populated ratio rows 22:31
        for r_idx in range(22, 32):
            sales[f"{col_pct}{r_idx}"].number_format = "0%"

    # F9: 다. 연도별 매출액 추정 (5 years in columns D..H with genuine revenue formulas)
    sales["B38"] = "다. 연도별 매출액 추정"
    sales["C38"] = "판매량(kg)"
    sales["I38"] = "(단위 : 천원, Kg)"
    sales["B39"] = "수매(상품)"
    sales["B40"] = "직거래(상품)"
    sales["B41"] = "미상품과"
    sales["B50"] = "매출합계"
    # Year 1 quantity reference column C
    sales["C39"] = "=E22"
    sales["C40"] = "=E26"
    sales["C41"] = "=E30"
    sales["C50"] = "=C39+C40+C41"

    for t in range(5):
        rev_col = get_column_letter(4 + t)
        qty_col = get_column_letter(5 + 2 * t)
        sales[f"{rev_col}38"] = f"{years[t]}년"
        sales[f"{rev_col}39"] = f"={qty_col}22*$O$7/1000"
        sales[f"{rev_col}40"] = f"={qty_col}26*$O$10/1000"
        sales[f"{rev_col}41"] = 0
        sales[f"{rev_col}50"] = f"={rev_col}39+{rev_col}40+{rev_col}41"

    # F5 & F10: 3.투자계획 - 총괄 및 5개년 연차별 세부투자계획
    tot_loan = v["loan"]
    fac_loan = min(tot_loan, v["facility"])
    rem_loan = tot_loan - fac_loan
    eq_loan = min(rem_loan, v["equipment"])
    rem_loan2 = rem_loan - eq_loan
    land_loan = min(rem_loan2, v["land"])

    invest["B2"] = "3. 투자계획"
    invest["B4"] = "가. 총괄 투자계획"
    invest["I5"] = "(단위:천원)"
    invest["B6"] = "투자연도"
    invest["C6"] = "소요액"
    invest["D7"] = "자부담"
    invest["E7"] = "보조"
    invest["F7"] = "융자"
    invest["B8"] = start
    invest["C8"] = "=SUM(D8:F8)"
    invest["D8"] = "=F24"
    invest["E8"] = 0
    invest["F8"] = "=H24"
    for i in range(1, 5):
        r = 8 + i
        det_r = 24 + 6 * i
        invest.cell(r, 2, f"=B{r-1}+1")
        invest.cell(r, 3, f"=SUM(D{r}:F{r})")
        invest.cell(r, 4, f"=F{det_r}")
        invest.cell(r, 5, f"=G{det_r}")
        invest.cell(r, 6, f"=H{det_r}")
    invest["B13"] = "합    계"
    invest["C13"] = "=SUM(C8:C12)"
    invest["D13"] = "=SUM(D8:D12)"
    invest["E13"] = "=SUM(E8:E12)"
    invest["F13"] = "=SUM(F8:F12)"

    # 나. 연차별 세부투자계획 - 5개년 블록 (행 19~48)
    invest["B15"] = "나. 연차별 세부투자계획"
    invest["B17"] = "연도별"
    invest["C17"] = "투자대상"
    invest["E17"] = "투자금액(천원)"
    invest["C18"] = "투자항목"
    invest["D18"] = "규격"
    invest["E18"] = "계"
    invest["F18"] = "자부담"
    invest["G18"] = "보조"
    invest["H18"] = "융자"

    # Year 1 세부투자 (19..24)
    invest["B19"] = start
    invest["C19"] = "시설"
    invest["D19"] = spec.get("facility_spec") or (
        "[확인 필요]" if v["facility"] > 0 else "-"
    )
    invest["E19"] = "=SUM(F19:H19)"
    invest["F19"] = float(v["facility"] - fac_loan)
    invest["G19"] = 0
    invest["H19"] = float(fac_loan)

    invest["C20"] = "대농기구"
    invest["D20"] = spec.get("equipment_spec") or (
        "[확인 필요]" if v["equipment"] > 0 else "-"
    )
    invest["E20"] = "=SUM(F20:H20)"
    invest["F20"] = float(v["equipment"] - eq_loan)
    invest["G20"] = 0
    invest["H20"] = float(eq_loan)

    invest["C21"] = "토지"
    invest["D21"] = spec.get("land_spec") or ("[확인 필요]" if v["land"] > 0 else "-")
    invest["E21"] = "=SUM(F21:H21)"
    invest["F21"] = float(v["land"] - land_loan)
    invest["G21"] = 0
    invest["H21"] = float(land_loan)

    invest["C22"] = "기타1"
    invest["D22"] = "-"
    invest["E22"] = "=SUM(F22:H22)"
    invest["F22"] = 0
    invest["G22"] = 0
    invest["H22"] = 0

    invest["C23"] = "기타2"
    invest["D23"] = "-"
    invest["E23"] = "=SUM(F23:H23)"
    invest["F23"] = 0
    invest["G23"] = 0
    invest["H23"] = 0

    invest["C24"] = "계"
    invest["E24"] = "=SUM(E19:E23)"
    invest["F24"] = "=SUM(F19:F23)"
    invest["G24"] = "=SUM(G19:G23)"
    invest["H24"] = "=SUM(H19:H23)"

    # Years 2..5 세부투자 블록 (25..48)
    for i in range(1, 5):
        r = 25 + 6 * (i - 1)
        invest.cell(r, 2, f"=B{8 + i}")
        invest.cell(r, 3, "시설")
        invest.cell(r, 4, "-")
        invest.cell(r, 5, f"=SUM(F{r}:H{r})")
        invest.cell(r, 6, 0)
        invest.cell(r, 7, 0)
        invest.cell(r, 8, 0)

        invest.cell(r + 1, 3, "대농기구")
        invest.cell(r + 1, 4, "-")
        invest.cell(r + 1, 5, f"=SUM(F{r+1}:H{r+1})")
        invest.cell(r + 1, 6, 0)
        invest.cell(r + 1, 7, 0)
        invest.cell(r + 1, 8, 0)

        invest.cell(r + 2, 3, "토지")
        invest.cell(r + 2, 4, "-")
        invest.cell(r + 2, 5, f"=SUM(F{r+2}:H{r+2})")
        invest.cell(r + 2, 6, 0)
        invest.cell(r + 2, 7, 0)
        invest.cell(r + 2, 8, 0)

        invest.cell(r + 3, 3, "기타1")
        invest.cell(r + 3, 4, "-")
        invest.cell(r + 3, 5, f"=SUM(F{r+3}:H{r+3})")
        invest.cell(r + 3, 6, 0)
        invest.cell(r + 3, 7, 0)
        invest.cell(r + 3, 8, 0)

        invest.cell(r + 4, 3, "기타2")
        invest.cell(r + 4, 4, "-")
        invest.cell(r + 4, 5, f"=SUM(F{r+4}:H{r+4})")
        invest.cell(r + 4, 6, 0)
        invest.cell(r + 4, 7, 0)
        invest.cell(r + 4, 8, 0)

        invest.cell(r + 5, 3, "계")
        invest.cell(r + 5, 5, f"=SUM(E{r}:E{r+4})")
        invest.cell(r + 5, 6, f"=SUM(F{r}:F{r+4})")
        invest.cell(r + 5, 7, f"=SUM(G{r}:G{r+4})")
        invest.cell(r + 5, 8, f"=SUM(H{r}:H{r+4})")

    # F3 & F5: Bounded principal repayment and authoritative loan linkage
    term_val = spec["term"] if spec.get("term", 0) > 0 else 1
    repay["B2"] = "4. 원리금상환계획"
    repay["B6"] = "연도"
    repay["C6"] = "융자금액"
    repay["D7"] = "원금상환"
    repay["E7"] = "이자"
    repay["F6"] = "연말융자잔액"
    repay["B8"] = start
    repay["C8"] = f"={invest_name}!F8"
    repay["D8"] = (
        f'=IF(B8-$B$8+1<={spec["grace"]},0,IF($C$8=0,0,MIN(C8,$C$8/{term_val})))'
    )
    repay["E8"] = f"=C8*{float(v['loan_rate'])}"
    repay["F8"] = "=C8-D8"
    for i in range(1, 5):
        r = 8 + i
        repay.cell(r, 2, f"=B{r-1}+1")
        repay.cell(r, 3, 0)
        repay.cell(
            r,
            4,
            f'=IF(B{r}-$B$8+1<={spec["grace"]},0,IF($C$8=0,0,MIN(F{r-1},$C$8/{term_val})))',
        )
        repay.cell(r, 5, f"=F{r-1}*{float(v['loan_rate'])}")
        repay.cell(r, 6, f"=F{r-1}-D{r}")

    # F4 & F5 & F10: 10. 감가상각비계획 (annual basis columns and useful-life bounded charges)
    dep["B2"] = "10. 감가상각비계획"
    dep["B4"] = "취득가액"
    dep["D7"] = "영농시설"
    dep["E8"] = f"={invest_name}!E19"
    dep["H7"] = "대농기구"
    dep["I8"] = f"={invest_name}!E20"

    dep["B10"] = "연도별 상각"
    dep["B11"] = "연도"
    dep["C11"] = "영농시설 취득가"
    dep["D11"] = "내용연수"
    dep["E11"] = "영농시설 상각액"
    dep["G11"] = "대농기구 취득가"
    dep["H11"] = "내용연수"
    dep["I11"] = "대농기구 상각액"

    for t in range(5):
        r = 12 + t
        dep.cell(r, 2, start if t == 0 else f"=B{r-1}+1")
        dep.cell(r, 3, "=E8")
        dep.cell(r, 4, spec["life"])
        dep.cell(
            r,
            5,
            f"=IF({t + 1}<={spec['life']},IF({spec['life']}=0,0,(E8-{float(v['salvage'])})/{spec['life']}),0)",
        )
        dep.cell(r, 7, "=I8")
        dep.cell(r, 8, spec["life"])
        dep.cell(
            r,
            9,
            f"=IF({t + 1}<={spec['life']},IF({spec['life']}=0,0,I8/{spec['life']}),0)",
        )

    # 9 .경비계획
    exp["B2"] = "9. 경비계획"
    exp["B5"] = "대농기구 수선비"
    exp["D5"] = float(v["repair_equipment_rate"])
    exp["B6"] = "영농시설 수선비"
    exp["D6"] = float(v["repair_facility_rate"])
    exp["B7"] = "수도광열비 (천원/10a)"
    exp["D7"] = float(v["utility_per_10a"])
    exp["B8"] = "경영면적(m2) > 1년차"
    exp["D8"] = f"={prod_name}!$H$8"
    exp["B9"] = "경영면적(m2) > 2년차 이후"
    exp["D9"] = "=D8"
    exp["B10"] = "임대면적(m2)"
    exp["D10"] = 0
    for t, year in enumerate(years):
        b = 12 + 20 * t
        exp.cell(b, 2, f"{YEAR_LABELS[t]}. {year}년")
        exp.cell(b + 1, 5, "(단위 : 천원)")
        exp.cell(b + 2, 2, "과     목")
        exp.cell(b + 2, 3, "금  액")
        exp.cell(b + 2, 4, "산   출   근   거")
        exp.cell(b + 2, 5, "비고")
        exp.cell(b + 3, 2, "1. 수선비")
        exp.cell(b + 3, 3, f"=SUM(C{b+4}:C{b+5})")
        exp.cell(b + 3, 4, "하위 합계")
        exp.cell(b + 4, 2, "  가. 영농시설")
        if t == 0:
            exp.cell(b + 4, 3, f"={dep_name}!E8*$D$6")
        else:
            exp.cell(b + 4, 3, f"=C{b + 4 - 20}")
        exp.cell(b + 4, 4, "* 구입금액의 시설수선비율 적용")
        exp.cell(b + 5, 2, "  나. 대농기구")
        if t == 0:
            exp.cell(b + 5, 3, f"={dep_name}!I8*$D$5")
        else:
            exp.cell(b + 5, 3, f"=C{b + 5 - 20}")
        exp.cell(b + 5, 4, "* 구입금액의 대농기구수선비율 적용")
        exp.cell(b + 6, 2, "2. 수도 광열비")
        if t == 0:
            exp.cell(b + 6, 3, "=$D$7*$D$8/1000")
        else:
            exp.cell(b + 6, 3, f"=C{b + 6 - 20}*{float(v['inflation'])}")
        exp.cell(b + 6, 4, "수도광열비×면적/1000×물가")
        exp.cell(b + 8, 2, "3. 감가상각비")
        exp.cell(b + 8, 3, f"=SUM(C{b+9}:C{b+11})")
        exp.cell(b + 8, 4, "* 감가상각비계획 참조")
        exp.cell(b + 9, 2, "  가. 영농시설")
        exp.cell(b + 9, 3, f"={dep_name}!E{12 + t}")
        exp.cell(b + 9, 4, "정액상각")
        exp.cell(b + 10, 2, "  나. 대농기구")
        exp.cell(b + 10, 3, f"={dep_name}!I{12 + t}")
        exp.cell(b + 10, 4, "정액상각")
        exp.cell(b + 11, 2, "  다. 생물자산")
        exp.cell(b + 11, 3, 0)
        exp.cell(b + 11, 4, "해당 없으면 0")
        exp.cell(b + 12, 2, "4. 임차료")
        # F11: Rental parent SUM of subitems
        exp.cell(b + 12, 3, f"=SUM(C{b+13}:C{b+14})")
        exp.cell(b + 12, 4, "하위 합계")
        exp.cell(b + 13, 2, "  가. 대농기구")
        exp.cell(b + 13, 3, 0)
        exp.cell(b + 13, 4, "해당 없으면 0")
        exp.cell(b + 14, 2, "  나. 토지")
        exp.cell(b + 14, 3, 0)
        exp.cell(b + 14, 4, "해당 없으면 0")
        exp.cell(b + 15, 2, "5. 위탁영농비")
        exp.cell(b + 15, 3, 0)
        exp.cell(b + 15, 4, "해당 없으면 0")
        exp.cell(b + 16, 2, "6. 소농구비")
        exp.cell(b + 16, 3, 0)
        exp.cell(b + 16, 4, "해당 없으면 0")
        exp.cell(b + 17, 2, "7. 기타 요금")
        exp.cell(b + 17, 3, 0)
        exp.cell(b + 17, 4, "해당 없으면 0")
        exp.cell(b + 18, 2, "합    계")
        exp.cell(
            b + 18,
            3,
            f"=C{b+3}+C{b+6}+C{b+8}+C{b+12}+C{b+15}+C{b+16}+C{b+17}",
        )
        exp.cell(b + 18, 4, "7과목 합계. 해당 없어도 0으로 남김")

    # F10: 7. 영농자재소요계획 & 8. 노무비계획 (basis text & inflation roll-forward)
    mat["B2"] = "7. 영농자재소요계획"
    mat["I2"] = "(단위:천원)"
    labor["B2"] = "8. 노무비계획"
    labor["I2"] = "(단위:천원)"
    for t, year in enumerate(years):
        r = 4 + t
        if t == 0:
            mat.cell(r, 2, year)
            mat.cell(r, 3, float(v["materials"]))
            mat.cell(
                r,
                4,
                spec.get("materials_basis")
                or ("[확인 필요]" if v["materials"] > 0 else "-"),
            )
            labor.cell(r, 2, year)
            labor.cell(r, 3, float(v["labor"]))
            labor.cell(
                r,
                4,
                spec.get("labor_basis") or ("[확인 필요]" if v["labor"] > 0 else "-"),
            )
        else:
            mat.cell(r, 2, f"=B{r-1}+1")
            mat.cell(r, 3, f"=C{r-1}*{float(v['inflation'])}")
            mat.cell(r, 4, "전년 자재비 × 물가상승률")
            labor.cell(r, 2, f"=B{r-1}+1")
            labor.cell(r, 3, f"=C{r-1}*{float(v['inflation'])}")
            labor.cell(r, 4, "전년 노무비 × 물가상승률")
    mat["B10"] = "합계"
    mat["C10"] = "=SUM(C4:C8)"
    labor["B10"] = "합계"
    labor["C10"] = "=SUM(C4:C8)"

    cost["B2"] = "11. 생산원가계획"
    cost["I2"] = "(단위:천원)"
    cost["B4"] = "과목"
    for i, year in enumerate(years):
        cost.cell(4, 3 + i, year if i == 0 else f"={get_column_letter(2 + i)}4+1")
    cost["C4"] = start
    cost["B6"] = "자재비"
    cost["B7"] = "노무비"
    cost["B8"] = "경비"
    cost["B9"] = "감가상각비"
    cost["B31"] = "당기총생산원가"
    cost["B32"] = "감가상각비"
    cost["B33"] = "현금생산원가"
    cost["B35"] = "매출원가"
    for i in range(5):
        col = get_column_letter(3 + i)
        b = 12 + 20 * i
        cost.cell(6, 3 + i, f"={mat_name}!C{4 + i}")
        cost.cell(7, 3 + i, f"={labor_name}!C{4 + i}")
        cost.cell(8, 3 + i, f"={exp_name}!C{b + 18}-{exp_name}!C{b + 8}")
        cost.cell(9, 3 + i, f"={exp_name}!C{b + 8}")
        cost.cell(31, 3 + i, f"={col}6+{col}7+{col}8+{col}9")
        cost.cell(32, 3 + i, f"={col}9")
        cost.cell(33, 3 + i, f"={col}31-{col}32")
        cost.cell(35, 3 + i, f"={col}31")

    pl["B2"] = "12 손익계획"
    pl["B4"] = "과목"
    pl["C4"] = start
    for i in range(1, 5):
        pl.cell(4, 3 + i, f"={get_column_letter(2 + i)}4+1")
        pl.cell(5, 3 + i, "1월1일~12월31일")
    pl["C5"] = "1월1일~12월31일"
    rows = {
        6: ("I. 매출액", "=SUM(C7:C8)"),
        7: ("  1. 생산물 매출", f"={sales_name}!D50"),
        8: ("  2. 부산물 매출", 0),
        9: ("II. 매출원가", f"={cost_name}!C35"),
        10: ("III. 매출총이익", "=C6-C9"),
        11: ("IV. 판매비와 관리비", "=SUM(C12:C14)"),
        12: ("  1. 포장비", float(v["packing"])),
        13: ("  2. 운반비", float(v["transport"])),
        14: ("  3. 기타", 0),
        15: ("V. 영업이익", "=C10-C11"),
        16: ("VI. 영업외 수익", "=SUM(C17:C19)"),
        17: ("  1. 토지/농기계임대료", 0),
        18: ("  2. 보조금", 0),
        19: ("  3. 배당금", 0),
        20: ("VII. 영업외 비용", "=SUM(C21:C22)"),
        21: ("  1. 이자비용", f"={repay_name}!E8"),
        22: ("  2. 기타", 0),
        23: ("VIII. 세금차감전 순익", "=C15+C16-C20"),
        25: ("IX. 세금(법인세 등)", 0),
        26: ("X. 당기순이익", "=C23-C25"),
    }
    for r, (label, formula) in rows.items():
        pl.cell(r, 2, label)
        pl.cell(r, 3, formula)
    for i in range(1, 5):
        col = get_column_letter(3 + i)
        prev = get_column_letter(2 + i)
        pl.cell(7, 3 + i, f"={sales_name}!{get_column_letter(4 + i)}50")
        pl.cell(8, 3 + i, 0)
        pl.cell(6, 3 + i, f"=SUM({col}7:{col}8)")
        pl.cell(9, 3 + i, f"={cost_name}!{col}35")
        pl.cell(10, 3 + i, f"={col}6-{col}9")
        pl.cell(12, 3 + i, float(v["packing"]))
        pl.cell(13, 3 + i, float(v["transport"]))
        pl.cell(14, 3 + i, 0)
        pl.cell(11, 3 + i, f"=SUM({col}12:{col}14)")
        pl.cell(15, 3 + i, f"={col}10-{col}11")
        pl.cell(17, 3 + i, 0)
        pl.cell(18, 3 + i, 0)
        pl.cell(19, 3 + i, 0)
        pl.cell(16, 3 + i, f"=SUM({col}17:{col}19)")
        pl.cell(21, 3 + i, f"={repay_name}!E{8 + i}")
        pl.cell(22, 3 + i, 0)
        pl.cell(20, 3 + i, f"=SUM({col}21:{col}22)")
        pl.cell(23, 3 + i, f"={col}15+{col}16-{col}20")
        pl.cell(25, 3 + i, 0)
        pl.cell(26, 3 + i, f"={col}23-{col}25")

    goal["B2"] = "2. 중장기영농목표"
    goal["B4"] = f"가. 작성기간 : 5개년 계획({years[0]}~{years[4]})"
    goal["B7"] = "나. 매출액 및 순이익 목표"
    goal["B9"] = "구 분"
    goal["C9"] = start
    for i in range(1, 5):
        goal.cell(9, 3 + i, f"={get_column_letter(2 + i)}9+1")
    goal["B10"] = "매출액"
    goal["B11"] = "순이익"
    goal["B12"] = "순이익률(%)"
    for i in range(5):
        col = get_column_letter(3 + i)
        goal.cell(10, 3 + i, f"={pl_name}!{col}6")
        goal.cell(11, 3 + i, f"={pl_name}!{col}26")
        goal.cell(12, 3 + i, f"=IF({col}10=0,NA(),{col}11/{col}10)")
        goal.cell(12, 3 + i).number_format = "0.0%"
    # F1 & F5: Coherent cash flow schedule with land cash purchase
    cf["B2"] = "14. 현금흐름계획"
    cf["B6"] = "조달"
    cf["C7"] = "농산물 판매대금"
    cf["D7"] = f"={pl_name}!C6"
    cf["C8"] = "부산물 매출"
    cf["D8"] = 0
    cf["C9"] = "보조금"
    cf["D9"] = 0
    cf["C10"] = "영농착수금"
    cf["D10"] = 0
    cf["C11"] = "자산매각대금"
    cf["D11"] = 0
    cf["C12"] = "융자금"
    cf["D12"] = f"={invest_name}!F8"
    cf["C13"] = "기타 수입"
    cf["D13"] = 0
    cf["C14"] = "소   계(A)"
    cf["D14"] = "=SUM(D7:D13)"
    cf["C15"] = "시설투자"
    cf["D15"] = f"={invest_name}!E19"
    cf["C16"] = "농기계투자"
    cf["D16"] = f"={invest_name}!E20"
    cf["C17"] = "토지투자"
    cf["D17"] = f"={invest_name}!E21"
    cf["C18"] = "당기총생산 원가"
    cf["D18"] = f"={cost_name}!C33"
    cf["C19"] = "판매 및 일반관리비"
    cf["D19"] = f"={pl_name}!C11"
    cf["C20"] = "영업외 비용"
    cf["D20"] = f"={pl_name}!C20"
    cf["C21"] = "융자원금상환"
    cf["D21"] = f"={repay_name}!D8"
    cf["C22"] = "소   계(B)"
    cf["D22"] = "=SUM(D15:D21)"
    cf["C23"] = "수지균형(C)"
    cf["D23"] = "=D14-D22"
    cf["C24"] = "누적수지균형(D)"
    cf["D24"] = "=D23"
    cf["E24"] = "=D24"
    cf["B25"] = "* 당기 총생산원가 중 감가상각비는 제외. 비현금 비용이기 때문."
    for i in range(1, 5):
        col = get_column_letter(4 + i)
        prev = get_column_letter(3 + i)
        cf.cell(7, 4 + i, f"={pl_name}!{get_column_letter(3 + i)}6")
        cf.cell(8, 4 + i, 0)
        cf.cell(9, 4 + i, 0)
        cf.cell(10, 4 + i, 0)
        cf.cell(11, 4 + i, 0)
        cf.cell(12, 4 + i, 0)
        cf.cell(13, 4 + i, 0)
        cf.cell(14, 4 + i, f"=SUM({col}7:{col}13)")
        cf.cell(15, 4 + i, 0)
        cf.cell(16, 4 + i, 0)
        cf.cell(17, 4 + i, 0)
        cf.cell(18, 4 + i, f"={cost_name}!{get_column_letter(3 + i)}33")
        cf.cell(19, 4 + i, f"={pl_name}!{get_column_letter(3 + i)}11")
        cf.cell(20, 4 + i, f"={pl_name}!{get_column_letter(3 + i)}20")
        cf.cell(21, 4 + i, f"={repay_name}!D{8 + i}")
        cf.cell(22, 4 + i, f"=SUM({col}15:{col}21)")
        cf.cell(23, 4 + i, f"={col}14-{col}22")
        cf.cell(24, 4 + i, f"={prev}24+{col}23")

    # F1 & F2 & F5: Balance Sheet with cumulative depreciation and authoritative references
    bs["B2"] = "13. 추정대차대조표"
    bs["I2"] = "(단위:천원)"
    bs["B4"] = "자산"
    bs["C4"] = start
    for i in range(1, 5):
        bs.cell(4, 3 + i, f"={get_column_letter(2 + i)}4+1")
    bs["B6"] = "현금"
    bs["C6"] = f"={base_name}!C9+{cf_name}!D23"
    bs["B7"] = "토지"
    bs["C7"] = f"={invest_name}!E21"
    bs["B8"] = "시설장부가"
    bs["C8"] = f"={dep_name}!E8-{dep_name}!E12"
    bs["B9"] = "대농기구장부가"
    bs["C9"] = f"={dep_name}!I8-{dep_name}!I12"
    bs["B10"] = "자산합계"
    bs["C10"] = "=C6+C7+C8+C9"
    bs["B12"] = "융자잔액"
    bs["C12"] = f"={repay_name}!F8"
    bs["B13"] = "자본"
    bs["C13"] = f"={base_name}!C7+{pl_name}!C26"
    bs["B14"] = "부채와자본"
    bs["C14"] = "=C12+C13"
    bs["B16"] = "자산=부채+자본"
    bs["C16"] = "=C10-C14"
    for i in range(1, 5):
        col = get_column_letter(3 + i)
        prev = get_column_letter(2 + i)
        bs.cell(6, 3 + i, f"={prev}6+{cf_name}!{get_column_letter(4 + i)}23")
        bs.cell(7, 3 + i, f"={prev}7")
        bs.cell(8, 3 + i, f"={prev}8-{dep_name}!E{12 + i}")
        bs.cell(9, 3 + i, f"={prev}9-{dep_name}!I{12 + i}")
        bs.cell(10, 3 + i, f"={col}6+{col}7+{col}8+{col}9")
        bs.cell(12, 3 + i, f"={repay_name}!F{8 + i}")
        bs.cell(13, 3 + i, f"={prev}13+{pl_name}!{col}26")
        bs.cell(14, 3 + i, f"={col}12+{col}13")
        bs.cell(16, 3 + i, f"={col}10-{col}14")

    inc["B2"] = "15.추정소득분석"
    inc["I2"] = "(단위:천원)"
    inc["B4"] = "매출액"
    inc["B5"] = "생산원가"
    inc["B6"] = "당기순이익"
    inc["B7"] = "가계비"
    inc["B8"] = "추정소득"
    inc["C4"] = f"={pl_name}!C6"
    inc["C5"] = f"={cost_name}!C35"
    inc["C6"] = f"={pl_name}!C26"
    inc["C7"] = float(v["household"])
    inc["C8"] = "=C6-C7"
    for i in range(1, 5):
        col = get_column_letter(3 + i)
        inc.cell(4, 3 + i, f"={pl_name}!{col}6")
        inc.cell(5, 3 + i, f"={cost_name}!{col}35")
        inc.cell(6, 3 + i, f"={pl_name}!{col}26")
        inc.cell(7, 3 + i, float(v["household"]))
        inc.cell(8, 3 + i, f"={col}6-{col}7")

    model["B2"] = "참고1. 모델농장분석"
    model["B4"] = "정성 비교 자료. 다른 시트가 이 시트를 수식으로 참조하지 않음."

    for sheet in wb:
        sheet.freeze_panes = "B2"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 18
        for row in sheet:
            for cell in row:
                cell.font = Font(size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Column widths tailored for content and evidence readability
    exp.column_dimensions["B"].width = 24
    exp.column_dimensions["C"].width = 16
    exp.column_dimensions["D"].width = 38
    exp.column_dimensions["E"].width = 14

    sales.column_dimensions["B"].width = 22
    sales.column_dimensions["C"].width = 12
    for c_letter in ("D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N"):
        sales.column_dimensions[c_letter].width = 13
    sales.column_dimensions["O"].width = 16

    invest.column_dimensions["B"].width = 14
    invest.column_dimensions["C"].width = 16
    invest.column_dimensions["D"].width = 28
    for c_letter in ("E", "F", "G", "H"):
        invest.column_dimensions[c_letter].width = 14

    prod.column_dimensions["B"].width = 22
    prod.column_dimensions["H"].width = 40
    index.column_dimensions["B"].width = 10
    index.column_dimensions["C"].width = 30
    index.column_dimensions["D"].width = 60
    index.merge_cells("B4:D4")

    # Explicit used-range print areas and pagination configurations
    index.print_area = "B2:D22"
    base.print_area = "B2:D10"
    goal.print_area = "B2:G13"

    # 3.투자계획: discontiguous print areas (총괄+1년차, 2~5년차).
    # 연속 인쇄 페이지에서도 열 머리행이 반복되도록 제목 행을 지정한다.
    invest.print_area = "B2:H24,B25:H48"
    invest.print_title_rows = "17:18"

    repay.print_area = "B2:G13"

    # 5. 판매계획: discontiguous print areas (가.가격추이+나.유통채널 [B2:O31] and 다.매출추정 [B38:O50])
    sales.print_area = "B2:O31,B38:O50"

    prod.print_area = "B2:I26"
    mat.print_area = "B2:E10"
    labor.print_area = "B2:E10"

    # 9 .경비계획: discontiguous print areas per summary and annual blocks
    exp.print_area = "B2:E10,B12:E30,B32:E50,B52:E70,B72:E90,B92:E110"

    dep.print_area = "B2:I17"
    cost.print_area = "B2:G36"  # fitToHeight=1 avoids orphan row on page 19
    pl.print_area = "B2:G27"  # fitToHeight=1 avoids orphan row on page 21
    bs.print_area = "B2:G17"
    cf.print_area = "B2:H26"
    inc.print_area = "B2:G9"
    path = Path(path)
    if path.exists():
        raise ValueError("기존 XLSX 덮어쓰기 금지")
    mc.reconfirm_output(authorization, context)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    path.with_suffix(".manifest.json").write_text(
        json.dumps(
            {
                "profile": "school_17_sheet_v1",
                "writing_year": spec["writing_year"],
                "years": years,
                "file_hash": digest(path.read_bytes()),
                "input_hash": digest(spec),
                "recalculation": "not_run",
                "rendering": "not_run",
                "major_authorization": authorization.to_dict(),
                "notice": "학교 17시트 구조·수식 생성. 재계산·인쇄·본문 교차는 별도 검사.",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return {
        "status": "generated",
        "profile": "school_17_sheet_v1",
        "years": years,
        "sheets": list(SCHOOL_SHEETS),
        "recalculation": "not_run",
    }


def inspect_school_workbook(path):
    from openpyxl import load_workbook

    issues = []
    wb = load_workbook(path, data_only=False)
    if tuple(wb.sheetnames) != SCHOOL_SHEETS:
        issues.append(
            {
                "status": "fail",
                "reason": "17시트 이름·순서 불일치",
                "observed": wb.sheetnames,
            }
        )
        return issues
    if wb["참고1. 모델농장분석"].sheet_state != "hidden":
        issues.append({"status": "fail", "reason": "모델농장분석 숨김 아님"})
    exp = wb["9 .경비계획"]
    if exp["D16"].value in (None, "") or exp["D17"].value in (None, ""):
        issues.append({"status": "fail", "reason": "경비 산출근거 열 비어 있음"})
    labels = [exp.cell(12 + i, 2).value for i in (3, 6, 8, 12, 15, 16, 17)]
    if not all(labels):
        issues.append({"status": "fail", "reason": "경비 7과목 라벨 누락"})
    prod = wb["6. 생산계획"]
    for addr in ("H15", "H17", "H19", "H23", "H24"):
        if not prod[addr].value:
            issues.append(
                {"status": "fail", "reason": "생산량 산출근거 문장 누락: " + addr}
            )
    invest = wb["3.투자계획"]
    if not str(invest["B9"].value or "").startswith("="):
        issues.append({"status": "fail", "reason": "투자연도 수식 증가 아님"})
    if invest["D19"].value in (None, "", "입력 규격", "[확인 필요]") and invest[
        "E19"
    ].value not in (0, "0", None, 0.0):
        issues.append({"status": "fail", "reason": "시설 투자 규격 누락"})
    if invest["D20"].value in (None, "", "입력 규격", "[확인 필요]") and invest[
        "E20"
    ].value not in (0, "0", None, 0.0):
        issues.append({"status": "fail", "reason": "대농기구 투자 규격 누락"})
    if invest["D21"].value in (None, "", "입력 규격", "[확인 필요]") and invest[
        "E21"
    ].value not in (0, "0", None, 0.0):
        issues.append({"status": "fail", "reason": "토지 투자 규격 누락"})
    mat = wb["7. 영농자재소요계획"]
    if mat["D4"].value in (None, "", "[확인 필요]") and mat["C4"].value not in (
        0,
        "0",
        None,
        0.0,
    ):
        issues.append({"status": "fail", "reason": "영농자재 산출근거 누락"})
    labor = wb["8. 노무비계획"]
    if labor["D4"].value in (None, "", "[확인 필요]") and labor["C4"].value not in (
        0,
        "0",
        None,
        0.0,
    ):
        issues.append({"status": "fail", "reason": "노무비 산출근거 누락"})
    repay = wb["4. 원리금상환계획"]
    if str(repay["F9"].value or "").replace(" ", "") != "=F8-D9":
        issues.append(
            {"status": "fail", "reason": "상환잔액이 전년 잔액-당해 원금이 아님"}
        )
    sales = wb[" 5. 판매계획"]
    if not str(sales["B14"].value or "").startswith("출처:"):
        issues.append({"status": "fail", "reason": "판매계획 출처 각주 누락"})
    pl = wb["12 손익계획"]
    if " 5. 판매계획" not in str(pl["C7"].value or ""):
        issues.append(
            {"status": "fail", "reason": "손익 매출이 판매계획을 참조하지 않음"}
        )
    goal = wb["2. 중장기영농목표"]
    if "12 손익계획" not in str(goal["C10"].value or ""):
        issues.append(
            {"status": "fail", "reason": "중장기목표가 손익계획을 참조하지 않음"}
        )
    if "NA()" not in str(goal["C12"].value or ""):
        issues.append({"status": "fail", "reason": "매출 0 순이익률 NA 미처리"})
    cf = wb["14. 현금흐름계획"]
    if "감가상각" not in str(cf["B25"].value or ""):
        issues.append({"status": "fail", "reason": "현금흐름 감가상각 제외 각주 없음"})
    return issues


def _period_keys(period):
    """Canonical keys a table header may carry for a fact period."""
    s = str(period or "").strip()
    keys = set()
    m = re.fullmatch(r"(\d{4})\s*년?", s)
    if m:
        keys.add(("year", int(m.group(1))))
    m = re.fullmatch(r"(\d+)\s*년차", s)
    if m:
        keys.add(("nth", int(m.group(1))))
    return keys


def _header_period_keys(cell):
    s = str(cell or "")
    keys = set()
    for m in re.finditer(r"(?<!\d)(\d{4})\s*년", s):
        keys.add(("year", int(m.group(1))))
    for m in re.finditer(r"(?<!\d)(\d+)\s*년차", s):
        keys.add(("nth", int(m.group(1))))
    return keys


def _has_unit(text, unit):
    return (
        re.search(
            r"(?<![가-힣A-Za-z0-9_/])" + re.escape(unit) + r"(?![가-힣A-Za-z0-9_/])",
            text or "",
        )
        is not None
    )


def _has_scope(text, scope):
    return (
        re.search(
            r"(?<![가-힣A-Za-z0-9_])" + re.escape(scope) + r"(?![A-Za-z0-9_])",
            text or "",
        )
        is not None
    )


def _label_cell(cell, label):
    """Match a header cell to the claim label; return (matched, declared unit)."""
    s = str(cell or "").strip()
    if s == label:
        return True, None
    m = re.fullmatch(re.escape(label) + r"\s*[(\[]\s*([^)\]]+?)\s*[)\]]", s)
    if m:
        return True, m.group(1)
    return False, None


_CELL_VALUE = re.compile(
    r"^\s*(?:\(\s*((?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*\)"
    r"|([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?))\s*(.*?)\s*$"
)

# A cell that opens like a parenthesized figure but is malformed must not
# be skipped silently — a wrong paren attempt can never hide behind a
# correct occurrence elsewhere.
_CELL_BAD_PAREN = re.compile(r"^\s*[+-]?\s*\(.*\d")


_TABLE_NOTE = re.compile(r"^(?:주\)|주:|자료:|출처:|단위:)")


def _is_table_note(text):
    return _TABLE_NOTE.match(text or "") is not None


def _body_tables(text):
    """Table nodes plus the context each table actually owns.

    Ownership is settled once for every node before any context is
    collected, in priority order: a ``표`` caption belongs to the table it
    announces; an explicit trailing note (``주)``/``주:``/``자료:``/
    ``출처:``/``단위:``) belongs to the table it directly follows — or the
    block it sits inside — never to a table it merely precedes; and the
    whole contiguous run of unclaimed non-note prose leading into a table
    block is that block's introduction, stopping at the heading that opens
    the block's own section.  Each table then collects only its own nodes
    and unclaimed neighbours inside the existing 3-back/2-forward window;
    a node owned by another table, an unattached note, another object
    block (image / ``그림`` caption), or a heading/new ``표`` caption
    after the table is a boundary — never a scope/unit source.
    """
    from gg_document import parse

    nodes = parse(text)
    n = len(nodes)
    is_table = [nd.get("kind") == "table" for nd in nodes]
    # A "주) X: tail" note is split by the parser into a heading plus a
    # tail paragraph — treat the tail as part of the note when the
    # original text still holds them on a single line.
    raw_lines = {ln.strip() for ln in text.splitlines()}
    note_like = [_is_table_note(nd.get("text")) for nd in nodes]
    for k in range(1, n):
        if (
            nodes[k].get("kind") == "paragraph"
            and nodes[k - 1].get("kind") == "heading"
            and note_like[k - 1]
        ):
            head = nodes[k - 1].get("text") or ""
            tail = nodes[k].get("text") or ""
            pat = re.escape(head) + r":\s+" + re.escape(tail)
            if any(re.fullmatch(pat, ln) for ln in raw_lines):
                note_like[k] = True
    # Other object blocks (images and non-표 captions) are boundaries in
    # both directions — they never serve as a table's introduction.
    is_other = [
        nd.get("kind") == "image"
        or (nd.get("kind") == "caption" and nd.get("label") != "표")
        for nd in nodes
    ]
    owner = [None] * n
    for i, t in enumerate(is_table):
        if not t:
            continue
        k = i - 1
        while k >= 0 and not is_table[k] and note_like[k]:
            k -= 1
        if (
            k >= 0
            and nodes[k].get("kind") == "caption"
            and nodes[k].get("label") == "표"
        ):
            owner[k] = i
            for m in range(k + 1, i):
                owner[m] = i
    for k in range(1, n):
        if is_table[k] or owner[k] is not None or not note_like[k]:
            continue
        if is_table[k - 1]:
            owner[k] = k - 1
        elif owner[k - 1] is not None and note_like[k - 1]:
            owner[k] = owner[k - 1]
    for a in range(n):
        t = a if is_table[a] else owner[a]
        if t is None or not (is_table[a] or nodes[a].get("kind") == "caption"):
            continue
        for k in range(a - 1, -1, -1):
            if (
                is_table[k]
                or owner[k] is not None
                or note_like[k]
                or is_other[k]
            ):
                break
            owner[k] = t
            if nodes[k].get("kind") == "heading":
                break

    tables = []
    for i, node in enumerate(nodes):
        if not is_table[i]:
            continue
        context = []
        for j in range(i - 1, max(i - 4, -1), -1):
            kind_j = nodes[j].get("kind")
            if is_table[j] or is_other[j]:
                break
            if owner[j] is not None:
                if owner[j] != i:
                    break
            elif note_like[j]:
                break
            if kind_j in ("caption", "paragraph", "heading"):
                context.append(nodes[j].get("text") or "")
            if kind_j == "heading" and not note_like[j]:
                break
        for j in range(i + 1, min(i + 3, n)):
            kind_j = nodes[j].get("kind")
            if is_table[j] or is_other[j]:
                break
            if owner[j] is not None:
                if owner[j] != i:
                    break
                context.append(nodes[j].get("text") or "")
                continue
            if (
                kind_j == "heading"
                or note_like[j]
                or (kind_j == "caption" and nodes[j].get("label") == "표")
            ):
                break
            if kind_j in ("caption", "paragraph"):
                context.append(nodes[j].get("text") or "")
        tables.append((node["rows"], "\n".join(context)))
    return tables


_DECL_PLACES = re.compile(
    r"소수점?\s*(?:아래|이하)?\s*(\d+)\s*자리\s*(?:까지|로)"
)
_DECL_INT = re.compile(r"정수\s*(?:로|까지)\s*반올림")
_DECL_BREAK = "(,;:)"
_DECL_SUBJECT = re.compile(r"([가-힣A-Za-z0-9/㎡]+?)\s*(?:은|는|의)\s*$")
_MAX_DECL_PLACES = 10


def _decl_unit(text, unit):
    """Unit mention where a digit may precede (``231,594.3천원``)."""
    return (
        re.search(
            r"(?<![가-힣A-Za-z_/])" + re.escape(unit) + r"(?![가-힣A-Za-z0-9_/])",
            text or "",
        )
        is not None
    )


def _display_places(scope_text, unit):
    """Local display-rounding declaration for ``unit`` inside ``scope_text``.

    Returns ``None`` when no declaration applies — the caller keeps the
    exact Decimal comparison; an ``int`` N when exactly one valid
    ``소수 N자리까지 반올림``/``정수로 반올림`` declaration binds to
    ``unit``; or ``"conflict"`` when an applicable declaration lacks places,
    is out of range, or is contradicted by another declaration — the claim
    then stays a mismatch rather than silently passing on exact comparison.
    """
    found = set()
    broken = False
    for line in str(scope_text or "").splitlines():
        if "반올림" not in line:
            continue
        matches = [
            (m.start(), m.group(1)) for m in _DECL_PLACES.finditer(line)
        ]
        matches += [(m.start(), "0") for m in _DECL_INT.finditer(line)]
        if not matches:
            head = line[: line.index("반올림")]
            w0 = max(head.rfind(ch) for ch in _DECL_BREAK)
            window = head[w0 + 1 :].strip()
            subj = _DECL_SUBJECT.search(window)
            if _decl_unit(window, unit):
                broken = True
            elif subj is not None:
                broken = broken or (subj.group(1) == unit)
            elif not window:
                broken = broken or _decl_unit(head, unit)
            continue
        for start, digits in matches:
            w0 = max(line.rfind(ch, 0, start) for ch in _DECL_BREAK)
            window = line[w0 + 1 : start].strip()
            subj = _DECL_SUBJECT.search(window)
            if _decl_unit(window, unit):
                bound = True
            elif subj is not None:
                bound = subj.group(1) == unit
            elif not window or "반올림" in window:
                bound = _decl_unit(line[:start], unit)
            else:
                bound = False
            if not bound:
                continue
            n = int(digits)
            if n > _MAX_DECL_PLACES:
                broken = True
            else:
                found.add(n)
    if broken or len(found) > 1:
        return "conflict"
    if found:
        return found.pop()
    return None


def _round_half_up(value, places):
    """Quantize a canonical Decimal to ``places`` decimals (ROUND_HALF_UP)."""
    prec = max(28, len(value.as_tuple().digits) + places + 2)
    try:
        return value.quantize(
            Decimal(1).scaleb(-places),
            rounding=ROUND_HALF_UP,
            context=Context(prec=prec),
        )
    except (InvalidOperation, ValueError, OverflowError):
        return None


def _table_claim_verdict(rows, context, label, expected_d, unit, period_keys, scope):
    """(matched, conflicted) for one table against one claim.

    A cell position counts only when the table's adjacent context declares
    the scope, the unit is declared by the item cell or the context, and the
    item row/column meets the period row/column.  Wrong period/scope/unit or
    a same-valued cell under another item can never stand in.
    """
    if not rows or not period_keys or not _has_scope(context, scope):
        return False, False
    decl = _display_places(context, unit)
    ncols = max(len(r) for r in rows)

    def cell(r, c):
        row = rows[r]
        return row[c] if c < len(row) else ""

    matched = conflicted = False
    for r in range(1, len(rows)):
        for c in range(1, ncols):
            for label_cell, period_cell in (
                (cell(r, 0), cell(0, c)),
                (cell(0, c), cell(r, 0)),
            ):
                ok, declared = _label_cell(label_cell, label)
                if not ok or not (_header_period_keys(period_cell) & period_keys):
                    continue
                if declared is not None:
                    if declared != unit:
                        conflicted = True
                        continue
                elif not _has_unit(
                    context + "\n" + str(label_cell) + "\n" + str(period_cell),
                    unit,
                ):
                    continue
                m = _CELL_VALUE.match(str(cell(r, c)))
                if m is None:
                    if _CELL_BAD_PAREN.match(str(cell(r, c))):
                        conflicted = True
                    continue
                rest = m.group(3).strip()
                if rest and rest != unit:
                    conflicted = True
                    continue
                claimed = Decimal(
                    (m.group(1) or m.group(2)).replace(",", "")
                )
                if m.group(1) is not None:
                    claimed = -claimed
                if decl == "conflict":
                    conflicted = True
                    continue
                target = (
                    expected_d
                    if decl is None
                    else _round_half_up(expected_d, decl)
                )
                if target is not None and claimed == target:
                    matched = True
                else:
                    conflicted = True
    return matched, conflicted


def crosscheck_body(text, values):
    """Check explicit scope/period/metric/amount/unit claims, not arbitrary prose.

    ``values`` accepts the legacy ``{label: meta}`` mapping or a list of
    claim records (``fact_id`` identity, ``field_id`` display label, value,
    unit, period, scope) so repeated labels across periods survive.  Table
    claims count only when the connected axes and declared scope/unit admit
    the position; a conflicting sentence or table cell is never covered by
    one correct occurrence.
    """
    issues = []
    if not isinstance(text, str):
        return [{"location": "body", "status": "fail", "reason": "본문 텍스트 없음"}]

    num_pattern = re.compile(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![\d.,])")

    if isinstance(values, dict):
        items = [(None, label, meta) for label, meta in values.items()]
    elif isinstance(values, (list, tuple)):
        items = [
            (
                (v.get("fact_id"), v.get("field_id"), v)
                if isinstance(v, dict)
                else (None, None, v)
            )
            for v in values
        ]
    else:
        items = [(None, None, values)]
    tables = _body_tables(text) if items else []

    for fact_id, label, meta in items:
        issue = {"location": str(label), "status": "fail"}
        if fact_id is not None:
            issue["fact_id"] = fact_id
        if not isinstance(label, str) or not label.strip():
            issue["reason"] = "유효하지 않은 메타데이터/기준값"
            issues.append(issue)
            continue
        issue["location"] = label
        if label not in text:
            issue["reason"] = "본문에 재무 항목 없음"
            issues.append(issue)
            continue

        # Fail closed on missing metadata (non-dict, or missing/blank value/unit/period/scope)
        val_raw = meta.get("value") if isinstance(meta, dict) else None
        unit_req = meta.get("unit") if isinstance(meta, dict) else None
        period_req = meta.get("period") if isinstance(meta, dict) else None
        scope_req = meta.get("scope") if isinstance(meta, dict) else None

        if (
            val_raw is None
            or unit_req is None
            or period_req is None
            or scope_req is None
            or not str(unit_req).strip()
            or not str(period_req).strip()
            or not str(scope_req).strip()
        ):
            issue["reason"] = "유효하지 않은 메타데이터/기준값"
            issues.append(issue)
            continue

        try:
            expected_d = Decimal(str(val_raw).replace(",", ""))
        except InvalidOperation:
            issue["reason"] = f"유효하지 않은 기준값: {val_raw}"
            issues.append(issue)
            continue

        expected_unit = str(unit_req).strip()
        expected_period = str(period_req).strip()
        expected_scope = str(scope_req).strip()

        # Split text into lines/sentences
        sentences = re.split(
            r"(?:\r?\n)+|(?<=[^\d])\.(?:\s+|$)|(?<=\d)\.(?=[^\d\s]|\s+[^\d])|[;•·]+",
            text,
        )

        scope = re.escape(expected_scope)
        period = re.escape(expected_period)
        claim_pattern = re.compile(
            r"(?<!\w)(?:"
            + scope
            + r"\s+"
            + period
            + "|"
            + period
            + r"\s+"
            + scope
            + r")\s+"
            + r"(?P<label>"
            + re.escape(label)
            + r")(?:은|는|이|가)?\s*[:=]?\s*"
        )
        unit_pattern = re.compile(re.escape(expected_unit) + r"(?![A-Za-z0-9/㎡])")
        found_match = False
        conflicting_claim = False
        for sent in sentences:
            if label not in sent:
                continue

            # Split sentence into clauses
            clauses = re.split(r",\s+(?=[^\d])", sent)
            for clause in clauses:
                claim = claim_pattern.search(clause)
                if claim is None:
                    continue
                label_start = claim.start("label")
                valid_occurrence = any(
                    m.start() == label_start for m in re.finditer(re.escape(label), clause)
                )
                if not valid_occurrence:
                    continue
                amount = num_pattern.match(clause, claim.end())
                if amount is None:
                    conflicting_claim = True
                    continue
                unit = unit_pattern.match(clause[amount.end() :].lstrip())
                actual = Decimal(amount.group().replace(",", ""))
                decl = _display_places(clause, expected_unit)
                if decl is None:
                    target = expected_d
                elif decl == "conflict":
                    target = None
                else:
                    target = _round_half_up(expected_d, decl)
                if unit is None or target is None or actual != target:
                    conflicting_claim = True
                else:
                    found_match = True

        for rows, context in tables:
            matched, conflicted = _table_claim_verdict(
                rows,
                context,
                label,
                expected_d,
                expected_unit,
                _period_keys(expected_period),
                expected_scope,
            )
            found_match = found_match or matched
            conflicting_claim = conflicting_claim or conflicted

        if not found_match or conflicting_claim:
            issue["reason"] = "기간·단위·범위·출력 반올림 값 불일치"
            issues.append(issue)

    return issues


# A26 body↔workbook locators: cached cells carrying the authoritative
# first-year value, year axis, and unit banner for each checked claim.
A26_CLAIM_LOCATORS = (
    {
        "label": "매출액",
        "unit": "천원",
        "value_ref": (" 5. 판매계획", "D50"),
        "period_ref": (" 5. 판매계획", "D38"),
        "unit_ref": (" 5. 판매계획", "I38"),
    },
    {
        "label": "당기순이익",
        "unit": "천원",
        "value_ref": ("12 손익계획", "C26"),
        "period_ref": ("12 손익계획", "C4"),
        "unit_ref": ("15.추정소득분석", "I2"),
    },
    {
        "label": "수매 판매량",
        "channel": "수매",
        "unit": "kg",
        "value_ref": (" 5. 판매계획", "E22"),
        "unit_ref": (" 5. 판매계획", "I19"),
    },
    {
        "label": "직거래 판매량",
        "channel": "직거래",
        "unit": "kg",
        "value_ref": (" 5. 판매계획", "E26"),
        "unit_ref": (" 5. 판매계획", "I19"),
    },
)


def _a26_year(raw):
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return int(raw)
    m = re.search(r"(\d{4})", str(raw or ""))
    return int(m.group(1)) if m else None


def crosscheck_body_workbook(text, path, scope):
    """Crosscheck canonical body claims against recalculated workbook cells.

    Each issue names the body claim location and the workbook sheet!cell that
    contradicts it.  A missing or non-numeric cached value is reported as
    ``blocked`` so an unrecalculated workbook cannot masquerade as a mismatch.
    """
    issues = []
    if not isinstance(text, str):
        return [{"location": "body", "status": "fail", "reason": "본문 텍스트 없음"}]
    scope = str(scope or "").strip()
    if not scope:
        return [
            {
                "location": "body",
                "status": "fail",
                "reason": "유효하지 않은 메타데이터/기준값",
            }
        ]
    from openpyxl import load_workbook

    num_pattern = r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![\d.,])"
    scope_p = re.escape(scope)
    sentences = re.split(
        r"(?:\r?\n)+|(?<=[^\d])\.(?:\s+|$)|(?<=\d)\.(?=[^\d\s]|\s+[^\d])|[;•·]+",
        text,
    )
    clauses = []
    for sent in sentences:
        clauses.extend(re.split(r",\s+(?=[^\d])", sent))

    wb = load_workbook(path, read_only=True, data_only=True)
    for loc in A26_CLAIM_LOCATORS:
        label = loc["label"]
        v_sheet, v_cell = loc["value_ref"]
        value_ref = v_sheet + "!" + v_cell
        if v_sheet not in wb.sheetnames:
            issues.append(
                {
                    "location": label,
                    "cell": value_ref,
                    "status": "fail",
                    "reason": "시트 없음: " + v_sheet,
                }
            )
            continue
        cached = wb[v_sheet][v_cell].value
        if isinstance(cached, bool) or not isinstance(cached, (int, float)):
            issues.append(
                {
                    "location": label,
                    "cell": value_ref,
                    "status": "blocked",
                    "reason": "재계산 캐시 없음/미실행",
                }
            )
            continue
        expected_d = Decimal(str(cached))
        expected_unit = loc["unit"]
        u_sheet, u_cell = loc["unit_ref"]
        unit_ref = u_sheet + "!" + u_cell

        expected_year = None
        period_ref = None
        if "period_ref" in loc:
            p_sheet, p_cell = loc["period_ref"]
            expected_year = _a26_year(wb[p_sheet][p_cell].value)
            period_ref = p_sheet + "!" + p_cell

        channel = loc.get("channel")
        if channel:
            claim_re = re.compile(
                r"(?<!\w)"
                + re.escape(channel)
                + r"\s+(?P<amount>"
                + num_pattern
                + r")\s*(?P<unit>[^\s,.;:()×]*)"
            )
        else:
            claim_re = re.compile(
                r"(?<!\w)(?:"
                + scope_p
                + r"\s+(?P<period>\d{4}년|\d+년차)|(?P<period2>\d{4}년|\d+년차)\s+"
                + scope_p
                + r")\s+"
                + re.escape(label)
                + r"(?:은|는|이|가)?\s*[:=]?\s*(?P<amount>"
                + num_pattern
                + r")\s*(?P<unit>[^\s,.;:()×]*)"
            )

        found_match = False
        for clause in clauses:
            decl = _display_places(clause, expected_unit)
            if decl is None:
                effective_d = expected_d
                decl_note = ""
            elif decl == "conflict":
                effective_d = None
                decl_note = " (표시 반올림 선언 모호·상충)"
            else:
                effective_d = _round_half_up(expected_d, decl)
                if effective_d is None:
                    decl_note = " (선언 소수 %d자리 반올림 계산 불가)" % decl
                else:
                    decl_note = (
                        " (소수 %d자리 반올림 표시 선언 기준 %s)"
                        % (decl, effective_d)
                    )
            for m in claim_re.finditer(clause):
                raw_amount = m.group("amount")
                try:
                    claimed_d = Decimal(raw_amount.replace(",", ""))
                except InvalidOperation:
                    claimed_d = None
                raw_unit = m.group("unit") or ""
                claimed_unit = re.sub(
                    r"(?:이다|입니다|이며|이고|에서|으로|였다|된다|임)\.?$",
                    "",
                    raw_unit,
                )
                unit_ok = claimed_unit == expected_unit

                claimed_period = m.groupdict().get("period") or m.groupdict().get(
                    "period2"
                )
                if claimed_period and expected_year is not None:
                    claimed_year = _a26_year(claimed_period)
                    if claimed_period.endswith("년차"):
                        claimed_year = expected_year + int(claimed_period[:-2]) - 1
                    if claimed_year is not None and claimed_year != expected_year:
                        if claimed_d is not None and claimed_d == effective_d:
                            issues.append(
                                {
                                    "location": label,
                                    "cell": period_ref,
                                    "field": "period",
                                    "status": "fail",
                                    "reason": "기간 불일치: 본문 "
                                    + claimed_period
                                    + " ≠ 워크북 "
                                    + str(expected_year)
                                    + "년 ("
                                    + value_ref
                                    + "="
                                    + str(expected_d)
                                    + ")",
                                    "claim": clause.strip(),
                                    "expected": str(expected_year) + "년",
                                    "claimed": claimed_period,
                                }
                            )
                        continue
                bad = False
                if not unit_ok:
                    bad = True
                    issues.append(
                        {
                            "location": label,
                            "cell": unit_ref,
                            "field": "unit",
                            "status": "fail",
                            "reason": "단위 불일치: 본문 "
                            + (claimed_unit or "(없음)")
                            + " ≠ 워크북 "
                            + expected_unit
                            + " ("
                            + unit_ref
                            + ")",
                            "claim": clause.strip(),
                            "expected": expected_unit,
                            "claimed": claimed_unit,
                        }
                    )
                if (
                    claimed_d is None
                    or effective_d is None
                    or claimed_d != effective_d
                ):
                    bad = True
                    issues.append(
                        {
                            "location": label,
                            "cell": value_ref,
                            "field": "value",
                            "status": "fail",
                            "reason": "수치 불일치: 본문 "
                            + raw_amount
                            + " ≠ 워크북 "
                            + str(expected_d)
                            + " ("
                            + value_ref
                            + ")"
                            + decl_note,
                            "claim": clause.strip(),
                            "expected": str(expected_d),
                            "claimed": raw_amount,
                        }
                    )
                if not bad:
                    found_match = True
        if not found_match and not any(
            i["location"] == label and i.get("status") == "fail" for i in issues
        ):
            issues.append(
                {
                    "location": label,
                    "cell": value_ref,
                    "status": "fail",
                    "reason": "본문에 일치/충돌하는 재무 수치 없음",
                }
            )
    return issues


def inspect_outputs(root, p, lineage=None):
    """Inspect current XLSX outputs.

    ``lineage`` is the core-validated ``superseded_by`` verdict map
    (``gg_core.output_lineage``); when omitted it is computed here so
    standalone callers get the same selection.  Only valid chain terminals
    are history-skipped — broken relations stay under inspection.
    """
    if lineage is None:
        from gg_core import output_lineage

        lineage = output_lineage(p)
    issues = []
    for oid, output in p.get("outputs", {}).items():
        if output.get("format") != "xlsx":
            continue
        if (lineage.get(oid) or {}).get("state") == "history":
            continue
        try:
            from openpyxl import load_workbook
        except ImportError:
            issues.append(("xlsx_dependency", "XLSX 검사에 필요한 openpyxl 미설치"))
            continue
        layout_kind = output.get("layout_kind")
        if layout_kind is not None and layout_kind != "provided_template":
            issues.append(("school_excel_17", "알 수 없는 XLSX layout_kind"))
            continue
        # A provided template may share the exemplar's 17 sheet names while
        # having a different cell contract.  It is inspected through its
        # native Excel manifest and generic workbook integrity only.
        if output.get("layout_kind") == "provided_template":
            if output.get("stale") is True:
                issues.append(
                    (
                        "school_excel_17",
                        "provided_template 출력이 입력 변경으로 stale: 재검사 필요: "
                        + str(output.get("id")),
                    )
                )
                continue
            try:
                path = local(root, output["path"])
                if not path.is_file():
                    raise ValueError("제공 템플릿 출력 파일 없음")
                output_hash = hashlib.sha256(path.read_bytes()).hexdigest()
                if output.get("file_hash") and output["file_hash"] != output_hash:
                    raise ValueError("제공 템플릿 출력 file_hash 불일치")

                native = output.get("native_manifest")
                if not isinstance(native, dict):
                    raise ValueError("provided_template native_manifest 필요")
                manifest_path = native.get("path")
                manifest_sha = native.get("sha256")
                if (
                    not isinstance(manifest_path, str)
                    or not manifest_path.strip()
                    or Path(manifest_path).is_absolute()
                    or not isinstance(manifest_sha, str)
                    or not re.fullmatch(r"[0-9a-fA-F]{64}", manifest_sha)
                ):
                    raise ValueError("native_manifest 경로·sha256 형식 오류")
                manifest_file = local(root, manifest_path)
                if not manifest_file.is_file():
                    raise ValueError("native manifest 파일 없음")
                actual_manifest_sha = hashlib.sha256(
                    manifest_file.read_bytes()
                ).hexdigest()
                if actual_manifest_sha.lower() != manifest_sha.lower():
                    raise ValueError("native manifest sha256 불일치")
                manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
                if not isinstance(manifest, dict):
                    raise ValueError("native manifest 객체 필요")
                if manifest.get("status") != "converted":
                    raise ValueError("native manifest status converted 필요")
                if manifest.get("engine") != "Microsoft Excel":
                    raise ValueError("native manifest engine Excel 필요")
                if str(manifest.get("engine_stderr", "")).strip():
                    raise ValueError("native Excel stderr 기록됨")
                if manifest.get("error") or (
                    manifest.get("returncode") not in (None, 0)
                ):
                    raise ValueError("native Excel 실행 오류 기록됨")
                if manifest.get("published_office_sha256") != output_hash:
                    raise ValueError("native manifest published_office_sha256 불일치")
                validation = manifest.get("validation")
                if not isinstance(validation, dict):
                    raise ValueError("native manifest validation 필요")
                if validation.get("school_validation") != "source_template_preservation_only":
                    raise ValueError("provided 템플릿 school_validation 상태 오류")
                if validation.get("structure_issues"):
                    raise ValueError("provided 템플릿 structure_issues 존재")
                if validation.get("school_issues"):
                    raise ValueError("provided 템플릿 school_issues 존재")
                excel = validation.get("excel")
                if (
                    not isinstance(excel, dict)
                    or excel.get("valid") is not True
                    or excel.get("sha256") != output_hash
                ):
                    raise ValueError("native manifest validation.excel valid/sha256 오류")
                # Generic integrity gate: do not apply legacy coordinates or
                # school-specific formulas to a supplied template.
                wb_formula = load_workbook(path, read_only=True, data_only=False)
                names = wb_formula.sheetnames
                if len(names) != 17:
                    raise ValueError("제공 템플릿 17시트 무결성 불일치")
                error_tokens = re.compile(
                    r"#(?:REF!|VALUE!|DIV/0!|N/A|NAME\?|NUM!|NULL!|SPILL!|CALC!|BLOCKED!|CONNECT!|UNKNOWN!)"
                )
                for ws in wb_formula.worksheets:
                    for row in ws.iter_rows():
                        for cell in row:
                            value = cell.value
                            if cell.data_type == "e" or (
                                isinstance(value, str)
                                and cell.data_type == "f"
                                and error_tokens.search(value)
                            ):
                                raise ValueError(
                                    f"native workbook error token: {ws.title}!{cell.coordinate}"
                                )
                wb_values = load_workbook(path, read_only=True, data_only=True)
                for ws in wb_values.worksheets:
                    for row in ws.iter_rows():
                        for cell in row:
                            value = cell.value
                            if cell.data_type == "e" or (
                                isinstance(value, str) and error_tokens.fullmatch(value.strip())
                            ):
                                raise ValueError(
                                    f"native workbook cached error: {ws.title}!{cell.coordinate}"
                                )
            except (
                OSError,
                ValueError,
                KeyError,
                TypeError,
                UnicodeError,
                zipfile.BadZipFile,
                json.JSONDecodeError,
            ) as e:
                issues.append(("school_excel_17", f"provided_template: {e}"))
            continue
        if output.get("stale") is True:
            issues.append(
                (
                    "school_excel_17",
                    "입력 변경으로 stale한 출력: 재검사 필요: "
                    + str(output.get("id") or oid),
                )
            )
            continue
        try:
            path = local(root, output["path"])
            names = load_workbook(path, read_only=True).sheetnames
            if len(names) != 17 and not (set(names) & set(SCHOOL_SHEETS)):
                continue
            for item in inspect_school_workbook(path):
                issues.append(("school_excel_17", item["reason"]))
        except (
            OSError,
            ValueError,
            KeyError,
            TypeError,
            UnicodeError,
            zipfile.BadZipFile,
            json.JSONDecodeError,
        ) as e:
            issues.append(("school_excel_17", str(e)))
    return issues


# ---------------------------------------------------------------------------
# P4 reference-fill section (stage-g005 D07/D09): audited statistical
# source_refs resolve through the provenance chain in gg_rda_research
# before any workbook cell is filled.  No P3 workbook formula or
# calculation path is touched by this section.


def resolve_statistical_source_refs(source_refs, *, resolver=None):
    """P4 adapter: resolve audited statistical refs for reference fill.

    ``resolver`` is an optional ``gg_rda_research.resolver_context()`` dict;
    omitted -> the deployed packs context is built lazily.  Returns
    ``{"resolved_source_refs": [...], "conflicts": [...]}`` — conflicts are
    reported, never silently filled."""
    import gg_rda_research

    return gg_rda_research.resolve_statistical_refs(
        source_refs, context=resolver)


def source_ref_labels(source_refs, *, resolver=None):
    """Display labels for the ``출처:`` fill: statistical refs that
    resolved through the provenance chain carry their physical-line
    binding ``(line N)``; plain refs render verbatim.  Callers must run
    ``validate`` (which fails closed on conflicts) before filling."""
    import gg_rda_research

    labels = []
    for ref in source_refs:
        if gg_rda_research.is_statistical_ref(ref):
            verdict = gg_rda_research.resolve_pack_ref(
                ref, context=resolver) if isinstance(ref, str) else (
                gg_rda_research.audit_source_ref(None, ref,
                                                 context=resolver))
            if verdict["status"] == "resolved":
                line = verdict["audit_key"]["physical_jsonl_line_1based"]
                labels.append("%s (물리행 %d)" % (verdict["source_ref"], line))
                continue
        labels.append(ref if isinstance(ref, str) else str(ref))
    return labels
