"""Offline, selected-aggregate receipt for the pinned 2024 MAFRA insect survey.

Usage: python gg_insect_source_receipt.py --pdf R0.pdf --xlsx R1.xlsx

The CLI has no network or publication path. It reads the two explicit files,
checks their independent byte hashes, then emits only an allowlisted summary.
Provenance that the CLI cannot prove from the bytes (portal dates, URLs) is
reported under "origin_claims" with a per-field basis, never as verified
fact.  The XLSX rights and the R0/R1 farm-count discrepancy stay unresolved.
"""

from __future__ import annotations

import argparse
import contextlib
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
from io import BytesIO, StringIO
import json
import logging
import os
import re
import stat
import sys
import warnings

PDF_SHA256 = "3dd2b75f5588d1b710123f56cc37747521dee786a2c4a13ecf5e6ab6e66541be"
XLSX_SHA256 = "2582fc7703bb3c409225e1a7329a6a591513500bd0a827e3e9f1807b2bcebe8a"
SHEET = "1. 라. 사육곤충, 판매액"
MAX_SOURCE_BYTES = 32 * 1024 * 1024
SCHEMA = "gg-insect-source-receipt/2"

# Provenance bases (DESIGN 14.4/15.2): every origin_claims entry carries the
# evidence level that backs it.  Nothing below is reverified by this CLI.
_SNAPSHOT = "private_portal_snapshot_2026-09-24"
_UNVERIFIED = "research_claim_unverified"
_INFERRED_HOST = "inferred_host"
_AUDIT = "independent_audit_note"

R0_ORIGIN_URL = "https://www.mafra.go.kr/bbs/home/789/589850/download.do"
R1_HOST = "data.mafra.go.kr"
R1_DOWNLOAD_PATH = "/opendata/data/downloadOpenDataWebFile.do"
R1_DATA_ID = "20210727000000001441"
R1_LANDING_URL = ("https://data.mafra.go.kr/opendata/data/"
                  "indexOpenDataDetail.do"
                  "?data_id=20210727000000001441&service_ty=F")
R1_DOWNLOAD_URL = "https://data.mafra.go.kr" + R1_DOWNLOAD_PATH

# R0 "참고 2 연도별 주요 통계" species rows -> R1 "농가수" cells, used only
# for the conflict_locus diagnostic.  Each entry: (anchor regex on the
# species label line, offset from that line to the parenthesized count row,
# R0 printed name, R1 column, R1 printed name).  Only 사슴벌레 prints its
# label on the count row itself in the layout text, so its offset is 0.
_SPECIES_COUNTS = (
    (r"흰\s*점\s*박\s*이", 1, "흰점박이꽃무지", "J", "흰점박이꽃무지"),
    (r"갈\s*색\s*거\s*저\s*리", 1, "갈색거저리", "M", "갈색거저리"),
    (r"귀\s*뚜\s*라\s*미", 1, "귀뚜라미", "V", "귀뚜라미"),
    (r"아\s*메\s*리\s*카", 1, "아메리카동애등에", "S", "동애등에"),
    (r"장\s*수\s*풍\s*뎅\s*이", 1, "장수풍뎅이", "D", "장수풍뎅이"),
    (r"사\s*슴\s*벌\s*레", 0, "사슴벌레", "G", "사슴벌레"),
)
# R1 species columns that together form R0's "기타(반딧불이, 나비 등)" group.
_R1_GROUP_CELLS = {"P": "나비", "Y": "반딧불이", "AB": "누에", "AE": "기타"}
_R1_GROUP_SALES_COLS = ("R", "AA", "AD", "AG")


class ReceiptError(Exception):
    """Closed failure; never include a path, source text, or parser exception."""

    def __init__(self, reason: str, source_id: str):
        super().__init__(reason)
        self.reason = reason
        self.source_id = source_id


def _isolated_parse(func, *args, **kwargs):
    """Run one whole third-party parse phase inside a single quiet box.

    func must carry the phase end to end — open, iterate, extract, close —
    and return only plain values.  stdout/stderr writes and warnings are
    captured and discarded, and logging.disable silences every logger,
    including ones whose handlers were attached before this call, so a
    noisy or failing parser cannot leak source text or local paths
    (DESIGN 14.5).  The previous manager.disable level is always restored.
    """
    previous_disable = logging.root.manager.disable
    try:
        logging.disable(logging.CRITICAL)
        with contextlib.redirect_stdout(StringIO()), \
                contextlib.redirect_stderr(StringIO()), \
                warnings.catch_warnings(record=True):
            return func(*args, **kwargs)
    finally:
        logging.disable(previous_disable)


def _read_pinned(path: str, expected_sha256: str,
                 source_id: str) -> tuple[bytes, str]:
    if not re.fullmatch(r"[0-9a-f]{64}", expected_sha256):
        raise ReceiptError("invalid_expected_hash", source_id)
    try:
        listing = os.lstat(path)
    except OSError:
        raise ReceiptError("source_unreadable", source_id) from None
    if not stat.S_ISREG(listing.st_mode):
        # FIFO, directory, socket, device, and symlink are all refused before
        # any open, so a named pipe can never block the reader.
        raise ReceiptError("not_regular_file", source_id)
    try:
        flags = os.O_RDONLY | os.O_NONBLOCK | getattr(os, "O_CLOEXEC", 0)
        if hasattr(os, "O_NOFOLLOW"):
            flags |= os.O_NOFOLLOW
        fd = os.open(path, flags)
    except OSError:
        raise ReceiptError("source_unreadable", source_id) from None
    try:
        opened = os.fstat(fd)
        if not stat.S_ISREG(opened.st_mode):
            raise ReceiptError("not_regular_file", source_id)
        if (opened.st_dev, opened.st_ino) != (listing.st_dev, listing.st_ino):
            raise ReceiptError("unstable_source", source_id)
        if opened.st_size > MAX_SOURCE_BYTES:
            raise ReceiptError("source_too_large", source_id)
        chunks = []
        size = 0
        while True:
            chunk = os.read(fd, 1 << 20)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_SOURCE_BYTES:
                raise ReceiptError("source_too_large", source_id)
            chunks.append(chunk)
        after = os.fstat(fd)
        if (opened.st_dev, opened.st_ino, opened.st_size,
                opened.st_mtime_ns, opened.st_ctime_ns) != (
                after.st_dev, after.st_ino, after.st_size,
                after.st_mtime_ns, after.st_ctime_ns) or size != after.st_size:
            raise ReceiptError("unstable_source", source_id)
    finally:
        os.close(fd)
    blob = b"".join(chunks)
    digest = hashlib.sha256(blob).hexdigest()
    if digest != expected_sha256:
        raise ReceiptError("hash_mismatch", source_id)
    return blob, digest


def _pdf_pages(blob: bytes) -> tuple[list[str], str]:
    # BytesIO ties parsing to the checked bytes, even if the path changes.
    import pypdf

    reader = pypdf.PdfReader(BytesIO(blob), strict=True)
    try:
        if reader.is_encrypted:
            raise ValueError("encrypted")
        return [page.extract_text(extraction_mode="layout") or ""
                for page in reader.pages], pypdf.__version__
    finally:
        reader.close()


_NUMBER = re.compile(r"(?<![\d,])(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![\d,])")
_PAREN_NUMBER = re.compile(r"\(((?:\d{1,3}(?:,\d{3})+|\d+))\)")


def _one_line(lines: list[str], pattern: str) -> tuple[int, re.Match[str]]:
    matches = [(i, m) for i, line in enumerate(lines)
               if (m := re.search(pattern, line))]
    if len(matches) != 1:
        raise ReceiptError("pdf_table_anchor_missing_or_ambiguous", "R0")
    return matches[0]


def _year_column_starts(lines: list[str]) -> tuple[int, ...]:
    # pypdf layout text preserves horizontal spacing.  Infer the five
    # columns from the table header; a year appearing somewhere on the page
    # is not enough to assign a row token to 2024.
    headers = [line for line in lines if re.search(
        r"(?<!\d)2020\s+2021\s+2022\s+2023\s+2024(?!\d)", line)]
    if len(headers) != 1:
        raise ReceiptError("pdf_year_header_missing_or_ambiguous", "R0")
    header = headers[0]
    starts = []
    for year in range(2020, 2025):
        matches = list(re.finditer(rf"(?<!\d){year}(?!\d)", header))
        if len(matches) != 1:
            raise ReceiptError("pdf_year_header_missing_or_ambiguous", "R0")
        starts.append(matches[0].start())
    gaps = [right - left for left, right in zip(starts, starts[1:])]
    if any(gap < 8 or gap > 20 for gap in gaps):
        raise ReceiptError("pdf_year_header_alignment_invalid", "R0")
    return tuple(starts)


def _aligned_2024_value(line: str, label_end: int,
                        year_starts: tuple[int, ...], *,
                        parenthesized: bool = False) -> Decimal:
    # Classify by absolute layout column, not by ordinal token count.
    edges = [(left + right) / 2 for left, right in
             zip(year_starts, year_starts[1:])]
    left_edge = year_starts[0] - (year_starts[1] - year_starts[0]) / 2
    right_edge = year_starts[4] + (year_starts[4] - year_starts[3]) / 2
    pattern = _PAREN_NUMBER if parenthesized else _NUMBER
    columns: list[list[str]] = [[] for _ in range(5)]
    for match in pattern.finditer(line, label_end):
        start = match.start()
        if start >= right_edge:
            continue  # year-over-year change and percentage columns
        if start < left_edge:
            raise ReceiptError("pdf_year_column_mismatch", "R0")
        column = sum(start >= edge for edge in edges)
        columns[column].append(match.group(1) if parenthesized else match.group())
    if any(len(column) != 1 for column in columns):
        raise ReceiptError("pdf_year_column_mismatch", "R0")
    return Decimal(columns[4][0].replace(",", ""))


def _pdf_locus(lines: list[str],
               year_starts: tuple[int, ...]) -> dict | None:
    """Runtime-extract the farm-count conflict locus from the R0 table.

    Diagnostic-only: any missing or ambiguous anchor degrades the locus to
    not_extracted — it must never fail the pinned receipt or alter the
    farm-count conflict axis (DESIGN 14.4/15.2).
    """
    try:
        species = {}
        for pattern, offset, r0_name, _col, _r1_name in _SPECIES_COUNTS:
            idx, _match = _one_line(lines, pattern)
            species[r0_name] = _aligned_2024_value(
                lines[idx + offset], 0, year_starts, parenthesized=True)
        sales_i, sales_m = _one_line(lines, r"기\s*타")
        group_sales = _aligned_2024_value(
            lines[sales_i], sales_m.end(), year_starts)
        count_i, _count_m = _one_line(lines, r"반\s*딧\s*불\s*이")
        group_count = _aligned_2024_value(
            lines[count_i], 0, year_starts, parenthesized=True)
    except (ReceiptError, IndexError):
        return None
    return {"species_counts": species, "group_sales_2024": group_sales,
            "group_farm_count_2024": group_count}


def _extract_pdf(blob: bytes) -> tuple[dict[str, Decimal], dict | None, str]:
    try:
        pages, version = _isolated_parse(_pdf_pages, blob)
        if len(pages) != 5:
            raise ReceiptError("pdf_page_count_unexpected", "R0")
        p1 = re.sub(r"\s+", "", pages[0])
        p4 = pages[3]
        p4_compact = re.sub(r"\s+", "", p4)
        p5 = re.sub(r"\s+", "", pages[4])
        if ("2024년곤충산업현황실태조사결과" not in p1
                or "곤충생산·가공·유통업을신고한농가및법인" not in p1
                or "꿀벌제외" not in p1
                or "114060" not in p5):
            raise ReceiptError("pdf_population_or_approval_missing", "R0")
        if ("참고2연도별주요통계" not in p4_compact
                or "백만원" not in p4_compact
                or "2024" not in p4_compact
                or not re.search(r"-\s*4\s*-", p4)):
            raise ReceiptError("pdf_table_header_or_page_label_missing", "R0")
        lines = [line for line in p4.splitlines() if line.strip()]
        year_starts = _year_column_starts(lines)
        total_i, total_m = _one_line(lines, r"합\s*계(?=\s)")
        cricket_i, cricket_m = _one_line(lines, r"귀\s*뚜\s*라\s*미")
        fly_i, fly_m = _one_line(lines, r"아메리카")
        if (fly_i + 1 >= len(lines)
                or "동애등에" not in re.sub(r"\s+", "", lines[fly_i + 1])):
            raise ReceiptError("pdf_species_label_incomplete", "R0")
        sales_total = _aligned_2024_value(lines[total_i], total_m.end(), year_starts)
        sales_cricket = _aligned_2024_value(lines[cricket_i], cricket_m.end(), year_starts)
        sales_fly = _aligned_2024_value(lines[fly_i], fly_m.end(), year_starts)
        if total_i + 1 >= len(lines):
            raise ReceiptError("pdf_farm_count_row_missing", "R0")
        count_line = lines[total_i + 1]
        count_label = re.search(r"\(중복포함\)", count_line)
        if count_label is None:
            raise ReceiptError("pdf_farm_count_row_missing", "R0")
        farm_count = _aligned_2024_value(
            count_line, count_label.end(), year_starts, parenthesized=True)
        if any(v < 0 for v in (sales_total, sales_cricket, sales_fly, farm_count)):
            raise ReceiptError("pdf_invalid_value", "R0")
        locus = _pdf_locus(lines, year_starts)
        return {"sales_total": sales_total, "sales_cricket": sales_cricket,
                "sales_fly": sales_fly, "farm_count": farm_count}, locus, version
    except ReceiptError:
        raise
    except Exception:  # parser details may include source text or local paths
        raise ReceiptError("pdf_parse_failed", "R0") from None


def _compact(value: object) -> str:
    return re.sub(r"\s+", "", value) if isinstance(value, str) else ""


def _numeric_cell(sheet: object, address: str) -> Decimal:
    value = sheet[address].value
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ReceiptError("xlsx_required_numeric_cell_missing", "R1")
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        raise ReceiptError("xlsx_invalid_numeric_cell", "R1") from None
    if not number.is_finite() or number < 0:
        raise ReceiptError("xlsx_invalid_numeric_cell", "R1")
    return number


def _optional_numeric(sheet: object, address: str) -> Decimal | None:
    """Locus-cell value; a blank or '-' reads as zero (SUM semantics) and
    any other non-numeric degrades the locus instead of failing the
    receipt."""
    value = sheet[address].value
    if value is None or (isinstance(value, str)
                         and _compact(value) in {"-", "."}):
        return Decimal(0)
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        return None
    if not number.is_finite() or number < 0:
        return None
    return number


def _xlsx_locus(sheet: object) -> dict | None:
    """Runtime-extract R1 cells for the farm-count conflict locus.

    Like _pdf_locus this is diagnostic-only: header or cell surprises
    return None rather than failing the pinned receipt (DESIGN 14.4/15.2).
    """
    species_cols = {col: r1_name for _p, _o, _r0, col, r1_name
                    in _SPECIES_COUNTS}
    for col, name in {**species_cols, **_R1_GROUP_CELLS}.items():
        if (_compact(sheet[col + "4"].value) != name
                or _compact(sheet[col + "5"].value) != "농가수"):
            return None
    for col in _R1_GROUP_SALES_COLS:
        if _compact(sheet[col + "5"].value) != "판매액":
            return None
    species_cells = {}
    for col, name in species_cols.items():
        value = _optional_numeric(sheet, col + "23")
        if value is None:
            return None
        species_cells[name] = {"cell": col + "23", "value": value}
    group_cells = {}
    for col in _R1_GROUP_CELLS:
        value = _optional_numeric(sheet, col + "23")
        if value is None:
            return None
        group_cells[col + "23"] = value
    group_sales = {}
    for col in _R1_GROUP_SALES_COLS:
        value = _optional_numeric(sheet, col + "23")
        if value is None:
            return None
        group_sales[col + "23"] = value
    return {"species_cells": species_cells,
            "group_cells": group_cells,
            "group_farm_count_total": sum(group_cells.values()),
            "group_sales_cells": group_sales,
            "group_sales_total": sum(group_sales.values())}


def _extract_xlsx(blob: bytes) -> tuple[dict[str, Decimal], dict[str, str],
                                      dict | None, str]:
    def _xlsx_values() -> tuple[dict[str, Decimal], dict[str, str], dict | None]:
        # The whole phase — load, lazy cell reads, close — runs inside one
        # isolation so lazy reads cannot leak diagnostics (DESIGN 14.5).
        workbook = openpyxl.load_workbook(BytesIO(blob), read_only=True,
                                          data_only=True, keep_links=False)
        formulas = openpyxl.load_workbook(BytesIO(blob), read_only=True,
                                          data_only=False, keep_links=False)
        try:
            if SHEET not in workbook.sheetnames or SHEET not in formulas.sheetnames:
                raise ReceiptError("xlsx_required_sheet_missing", "R1")
            sheet = workbook[SHEET]
            formula_sheet = formulas[SHEET]
            headers = {address: _compact(sheet[address].value)
                       for address in ("A1", "B3", "C3", "D3", "S4", "U5",
                                       "V4", "X5", "A23")}
            expected = {"A1": "1.곤충사육·가공·유통농가·업체",
                        "B3": "농가수합계", "C3": "판매액",
                        "D3": "연간사육곤충농가수(호),마리수(마리),판매액(백만원)",
                        "S4": "동애등에", "U5": "판매액", "V4": "귀뚜라미",
                        "X5": "판매액", "A23": "계"}
            if headers != expected:
                raise ReceiptError("xlsx_header_or_unit_mismatch", "R1")
            addresses = (("sales_total", "C23"), ("sales_fly", "U23"),
                         ("sales_cricket", "X23"), ("farm_count", "B23"))
            values = {key: _numeric_cell(sheet, address)
                      for key, address in addresses}
            storage = {}
            for key, address in addresses:
                raw_cell = formula_sheet[address]
                if raw_cell.data_type == "f":
                    if not isinstance(raw_cell.value, str) or not raw_cell.value.startswith("=SUM("):
                        raise ReceiptError("xlsx_formula_unexpected", "R1")
                    storage[key] = "cached_formula"
                elif isinstance(raw_cell.value, (int, float)) and not isinstance(raw_cell.value, bool):
                    storage[key] = "literal_numeric"
                else:
                    raise ReceiptError("xlsx_cell_type_unexpected", "R1")
            locus = _xlsx_locus(sheet)
        finally:
            workbook.close()
            formulas.close()
        return values, storage, locus

    try:
        import openpyxl

        values, storage, locus = _isolated_parse(_xlsx_values)
        return values, storage, locus, openpyxl.__version__
    except ReceiptError:
        raise
    except Exception:
        raise ReceiptError("xlsx_parse_failed", "R1") from None


def _decimal(value: Decimal) -> str:
    return format(value, "f")


def _decimal_places(value: Decimal) -> int:
    return max(0, -value.as_tuple().exponent)


def _sales_comparison(metric: str, pdf_value: Decimal,
                      xlsx_value: Decimal) -> dict[str, str]:
    difference = xlsx_value - pdf_value
    if difference == 0:
        relation = "exact_numeric_match"
    elif xlsx_value.quantize(Decimal("1"), rounding=ROUND_HALF_UP) == pdf_value:
        # R0 prints integer KRW-million cells; R1 stores decimals.  Rounding
        # R1 to R0's displayed precision matching is an observed display
        # relation only — not a confirmed official rounding policy or a
        # same-species proof (see relation_note).
        relation = "display_rounding_consistent"
    else:
        relation = "unresolved_numeric_difference"
    row = {"metric": metric, "relation": relation,
           "r1_minus_r0_KRW_million": _decimal(difference)}
    if metric == "sales_fly":
        # R0 prints 아메리카동애등에 while R1 says 동애등에; no alias proof.
        row["caveat"] = "species_label_mapping_unconfirmed"
    return row


def _conflict_locus(pdf_locus: dict | None,
                    xlsx_locus: dict | None) -> dict:
    """Runtime-only description of the reported_farm_count conflict locus.

    Only values extracted and compared during this run may appear here;
    independent-audit observations live separately under audit_notes.
    Extraction failure changes nothing outside this field (DESIGN 15.2).
    """
    if pdf_locus is None or xlsx_locus is None:
        return {"status": "not_extracted", "cause": "unresolved"}
    species = []
    for _pattern, _offset, r0_name, _col, r1_name in _SPECIES_COUNTS:
        r0 = pdf_locus["species_counts"].get(r0_name)
        cell = xlsx_locus["species_cells"].get(r1_name)
        r1 = cell["value"] if isinstance(cell, dict) else None
        if r0 is None or r1 is None:
            return {"status": "not_extracted", "cause": "unresolved"}
        species.append({"species_R0": r0_name, "species_R1": r1_name,
                        "R1_cell": cell["cell"],
                        "R0_2024": _decimal(r0), "R1_2024": _decimal(r1),
                        "relation": "match" if r0 == r1 else "differ"})
    return {
        "status": "extracted",
        "basis": "runtime_extracted",
        "cause": "unresolved",
        "R0": {"physical_page": 4, "row": "기타(반딧불이, 나비 등)",
               "farm_count_2024": _decimal(pdf_locus["group_farm_count_2024"]),
               "sales_2024_KRW_million": _decimal(
                   pdf_locus["group_sales_2024"])},
        "R1": {"farm_count_cells": {k: _decimal(v) for k, v in
                                    xlsx_locus["group_cells"].items()},
               "farm_count_2024_total": _decimal(
                   xlsx_locus["group_farm_count_total"]),
               "sales_cells": {k: _decimal(v) for k, v in
                               xlsx_locus["group_sales_cells"].items()},
               "sales_2024_KRW_million_total": _decimal(
                   xlsx_locus["group_sales_total"])},
        "species_farm_counts": species,
    }


def _origin_claims() -> dict:
    """Per-field provenance claims — explicitly not CLI verification."""
    return {
        "R0": {
            "download_url": {"value": R0_ORIGIN_URL, "basis": _UNVERIFIED},
            "source_publication_date": {
                "value": "2025-09-01", "basis": _UNVERIFIED,
                "partial_evidence": "pdf_metadata_creation_date",
            },
        },
        "R1": {
            "data_id": {"value": R1_DATA_ID, "basis": _SNAPSHOT},
            "download_path": {"value": R1_DOWNLOAD_PATH, "basis": _SNAPSHOT},
            "download_method": {"value": "POST", "basis": _SNAPSHOT},
            "download_parameters": {
                "value": {"file_ty_code": "file", "file_sn": "17"},
                "basis": _SNAPSHOT,
            },
            "portal_dataset_registered_date": {"value": "2021-07-27",
                                               "basis": _SNAPSHOT},
            "portal_last_modified_date": {"value": "2025-09-01",
                                          "basis": _SNAPSHOT},
            "license_scope_field": {"value": "blank", "basis": _SNAPSHOT},
            "file_publication_date": {"value": None, "basis": _SNAPSHOT},
            "host": {"value": R1_HOST, "basis": _INFERRED_HOST},
            "landing_url": {
                "value": R1_LANDING_URL, "status": "composed",
                "parts_basis": {"host": _INFERRED_HOST,
                                "landing_path": _UNVERIFIED,
                                "data_id": _SNAPSHOT,
                                "service_ty=F": _UNVERIFIED},
            },
            "download_url": {
                "value": R1_DOWNLOAD_URL, "status": "composed",
                "parts_basis": {"host": _INFERRED_HOST,
                                "download_path": _SNAPSHOT},
            },
            "r1_portal_identity": "not_rechecked",
        },
        "note": "origin claims are owner-snapshot provenance; "
                "the CLI does not recheck them",
    }


def verify_sources(pdf_path: str, xlsx_path: str) -> dict:
    """Public entry — official pins only, no hash overrides (DESIGN 6.1).

    Both hashes are checked before either parser runs.
    """
    return _verify_with_pins(pdf_path, xlsx_path,
                             expected_pdf_sha256=PDF_SHA256,
                             expected_xlsx_sha256=XLSX_SHA256)


def _verify_with_pins(pdf_path: str, xlsx_path: str, *,
                      expected_pdf_sha256: str,
                      expected_xlsx_sha256: str) -> dict:
    """Isolated-test entry (not a public API): callers may pin different
    hashes, and the receipt then records official_pins:false so the
    evidence path refuses it."""
    pdf_blob, pdf_hash = _read_pinned(pdf_path, expected_pdf_sha256, "R0")
    xlsx_blob, xlsx_hash = _read_pinned(xlsx_path, expected_xlsx_sha256, "R1")
    pdf_values, pdf_locus, pdf_version = _extract_pdf(pdf_blob)
    xlsx_values, xlsx_storage, xlsx_locus, xlsx_version = _extract_xlsx(xlsx_blob)
    official_pins = pdf_hash == PDF_SHA256 and xlsx_hash == XLSX_SHA256

    pdf_cells = {"sales_total": "2024/B:합계/판매액",
                 "sales_cricket": "2024/B:귀뚜라미/판매액",
                 "sales_fly": "2024/B:아메리카동애등에/판매액",
                 "farm_count": "2024/B:합계(중복포함)/농가수"}
    xlsx_cells = {"sales_total": "C23", "sales_cricket": "X23",
                  "sales_fly": "U23", "farm_count": "B23"}
    observations = []
    for key in ("sales_total", "sales_cricket", "sales_fly", "farm_count"):
        unit = "count" if key == "farm_count" else "KRW_million"
        observations.append({
            "metric": "reported_farm_count" if key == "farm_count"
                      else "primary_product_sales",
            "aggregate": key.removeprefix("sales_"), "unit": unit,
            "R0": {"value": _decimal(pdf_values[key]), "physical_page": 4,
                   "display_decimal_places": _decimal_places(pdf_values[key]),
                   "printed_page_label": "4", "table": "참고 2 연도별 주요 통계",
                   "row_column": pdf_cells[key]},
            "R1": {"value": _decimal(xlsx_values[key]), "sheet": SHEET,
                   "display_decimal_places": _decimal_places(xlsx_values[key]),
                   "cell": xlsx_cells[key], "storage": xlsx_storage[key]},
        })
    comparisons = [_sales_comparison(key, pdf_values[key], xlsx_values[key])
                   for key in ("sales_total", "sales_cricket", "sales_fly")]
    count_conflict = pdf_values["farm_count"] != xlsx_values["farm_count"]
    comparisons.append({"metric": "reported_farm_count",
                        "relation": "source_conflict" if count_conflict
                        else "exact_numeric_match",
                        "r1_minus_r0_count": _decimal(
                            xlsx_values["farm_count"] - pdf_values["farm_count"])})
    conflict = ({"metric": "reported_farm_count",
                 "R0": _decimal(pdf_values["farm_count"]),
                 "R1": _decimal(xlsx_values["farm_count"]),
                 "status": "source_conflict", "disposition": "quarantined",
                 "resolution": "unresolved_definition_timing_or_erratum"}
                if count_conflict else None)
    parsed_selection = {"R0": {k: _decimal(v) for k, v in pdf_values.items()},
                        "R1": {k: _decimal(v) for k, v in xlsx_values.items()}}
    selection_sha256 = hashlib.sha256(json.dumps(
        parsed_selection, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")).hexdigest()
    return {
        "schema": SCHEMA,
        "status": "partial_verification_source_conflict" if count_conflict
                  else "selected_aggregates_parsed",
        "scope": "selected_2024_aggregate_values_only",
        "extracted_at_utc": datetime.now(timezone.utc).isoformat(),
        "selection_sha256": selection_sha256,
        "official_pins": official_pins,
        "origin_claims": _origin_claims() if official_pins else None,
        "conflict_locus": _conflict_locus(pdf_locus, xlsx_locus),
        "audit_notes": [
            {"basis": _AUDIT,
             "note": "R1 전체 셀에 2,393 값이 없음 (독립 감사 2026-09-25)"},
            {"basis": _AUDIT,
             "note": "R0 기타 묶음과 R1 누에 열(AB23) 사이 1곳 차이라는 "
                     "산술 추정 — 원문 근거 없음, 미확정, 어느 값도 "
                     "승격하지 않음"},
        ],
        "relation_note": "display_rounding_consistent·exact_numeric_match는 "
                         "표시 자릿수 반올림·값 일치 관측일 뿐, 공식 반올림 "
                         "정책이나 동일 종 확인이 아님",
        "sources": {
            "R0": {"kind": "government_pdf", "sha256": pdf_hash,
                   "bytes": len(pdf_blob), "parser": "pypdf",
                   "parser_version": pdf_version, "ocr": False,
                   "physical_pages": 5, "population_locator": "physical_page_1",
                   "approval_locator": "physical_page_5",
                   "table_locator": "physical_page_4_printed_4",
                   "visual_crosscheck": "not_performed_by_cli"},
            "R1": {"kind": "official_xlsx", "sha256": xlsx_hash,
                   "bytes": len(xlsx_blob), "parser": "openpyxl",
                   "parser_version": xlsx_version, "read_only": True,
                   "data_only": True,
                   "formula_recalculation": "not_performed",
                   "sheet": SHEET,
                   "headers": {"A1": "1. 곤충 사육·가공·유통 농가·업체",
                               "B3": "농가수 합계", "C3": "판매액",
                               "D3": "연간 사육곤충 농가수(호), 마리수(마리), 판매액(백만원)",
                               "S4": "동애등에", "U5": "판매액",
                               "V4": "귀뚜라미", "X5": "판매액", "A23": "계"},
                   "unit_locator": "D3", "population_locator": "A1"},
        },
        "survey": {"year": 2024, "period": "2024-01-01/2024-12-31",
                   "population": "곤충생산·가공·유통업을 신고한 농가 및 법인 "
                                 "(R0 physical_page_1)",
                   "population_note": "꿀벌 제외 표기는 R0 1쪽 판매액 순위 "
                                      "문장에 붙어 있음(모집단 정의 아님)",
                   "use_scope": "industry_context_only"},
        "observations": observations,
        "comparisons": comparisons,
        "conflicts": [conflict] if conflict else [],
        "rights": {"R0": "attribution_indicated_in_w21; "
                        "redistribution_not_assessed_by_cli",
                   "R1": "unconfirmed; source_and_extracted_rows_"
                         "not_cleared_for_publication"},
        "publication": "not_promoted_to_catalog_or_manifest",
        "claim_limits": ["no_promotional_claims", "no_farm_price_claims",
                         "no_yield_claims", "no_farm_revenue_projection"],
    }


class _QuietParser(argparse.ArgumentParser):
    """Argument errors become a fixed rejection on stdout — the offending
    argv text (which may be a private path) is never echoed (DESIGN 14.5)."""

    def error(self, message):
        sys.stdout.write(json.dumps(
            {"schema": SCHEMA, "status": "rejected",
             "reason": "invalid_arguments"}, ensure_ascii=False) + "\n")
        self.exit(2)


def main(argv: list[str] | None = None) -> int:
    parser = _QuietParser(description=__doc__)
    parser.add_argument("--pdf", required=True, help="local pinned R0 PDF")
    parser.add_argument("--xlsx", required=True, help="local pinned R1 XLSX")
    args = parser.parse_args(argv)
    try:
        receipt = verify_sources(args.pdf, args.xlsx)
    except ReceiptError as exc:
        print(json.dumps({"schema": SCHEMA, "status": "rejected",
                          "reason": exc.reason, "source_id": exc.source_id}))
        return 2
    print(json.dumps(receipt, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
