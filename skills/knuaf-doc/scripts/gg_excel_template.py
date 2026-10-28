#!/usr/bin/env python3
"""Source-bound reusable XLSX template inspection and blank-copy utility.

The utility edits worksheet XML in a zip-preserving pass. This keeps styles,
drawings, merges, links, print settings and formula expressions intact while
removing only cells listed in an explicit, SHA-bound input map. Formula cached
values are invalidated in the output so a blank template never presents old
calculated results as current financial results.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import re
import sys
import zipfile
import unicodedata
from xml.dom import minidom
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
ET.register_namespace("", NS_MAIN)

CELL_RE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")
RANGE_RE = re.compile(r"^([A-Z]+)([1-9][0-9]*)(?::([A-Z]+)([1-9][0-9]*))?$")

LAYOUT_EXTRACT_PATH = (
    Path(__file__).resolve().parents[1] / "references"
    / "common-workbooks" / "kang-finance-workbook.json"
)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def resolve_source(path: Path) -> Path:
    """Resolve a user-supplied NFC/NFD path without guessing another workbook."""
    if path.exists():
        return path
    parent = path.parent
    wanted = unicodedata.normalize("NFD", path.name)
    matches = [p for p in parent.glob("*.xlsx") if unicodedata.normalize("NFD", p.name) == wanted]
    if len(matches) == 1:
        return matches[0]
    raise FileNotFoundError(f"source workbook not found (NFC/NFD exact match failed): {path}")


def col_num(col: str) -> int:
    n = 0
    for c in col:
        n = n * 26 + ord(c) - 64
    return n


def col_name(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


MAX_EXPAND_REF_CELLS = 16384
MAX_MERGED_TOTAL_CELLS = 262144
EXCEL_MAX_COL = 16384  # XFD
EXCEL_MAX_ROW = 1048576


def expand_ref(ref: str) -> list[str]:
    m = RANGE_RE.match(ref.upper().replace("$", ""))
    if not m:
        raise ValueError(f"invalid cell/range reference: {ref}")
    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), m.group(4)
    c2, r2 = c2 or c1, int(r2 or r1)
    lo_c, hi_c = min(col_num(c1), col_num(c2)), max(col_num(c1), col_num(c2))
    lo_r, hi_r = min(r1, r2), max(r1, r2)
    if hi_c > EXCEL_MAX_COL or hi_r > EXCEL_MAX_ROW:
        raise ValueError(f"cell/range reference outside Excel bounds: {ref}")
    n = (hi_r - lo_r + 1) * (hi_c - lo_c + 1)
    if n > MAX_EXPAND_REF_CELLS:
        raise ValueError(f"cell range too large to expand: {ref} ({n} cells)")
    return [f"{col_name(c)}{r}" for r in range(lo_r, hi_r + 1)
            for c in range(lo_c, hi_c + 1)]


def local(tag: str) -> str:
    return f"{{{NS_MAIN}}}{tag}"


def is_reference_text(value: str) -> bool:
    """Detect narrative/reference text that is not a fixed worksheet label."""
    if value.startswith("*"):
        return True
    markers = ("강원도", "원주시", "농촌진흥청", "임산물", "소득조사", "유통정보",
               "물가상승률", "택배", "일당 기준", "판매량", "판매가격", "3년생", "2,000",
               "60,000", "경제성분석", "기준으로 작성", "`", "주소")
    return len(value) > 24 and any(m in value for m in markers)


def text_value(cell: ET.Element, shared: list[str]) -> str | int | float | None:
    t = cell.attrib.get("t")
    if t == "inlineStr":
        return "".join(cell.itertext())
    v = cell.find(local("v"))
    if v is None:
        return None
    raw = v.text or ""
    if t == "s":
        try:
            return shared[int(raw)]
        except (ValueError, IndexError):
            return raw
    if t == "b":
        return raw == "1"
    try:
        return float(raw) if any(x in raw for x in ".eE") else int(raw)
    except ValueError:
        return raw


def load_shared(z: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(z.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    return ["".join(si.itertext()) for si in root.findall(local("si"))]


def workbook_sheets(z: zipfile.ZipFile) -> list[tuple[str, str]]:
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rid_to_target = {r.attrib["Id"]: r.attrib["Target"] for r in rels}
    out = []
    for s in wb.findall(f"{{{NS_MAIN}}}sheets/{{{NS_MAIN}}}sheet"):
        rid = s.attrib.get(f"{{{NS_REL}}}id")
        target = rid_to_target.get(rid, "")
        if target.startswith("/"):
            target = target[1:]
        elif not target.startswith("xl/"):
            target = "xl/" + target
        out.append((s.attrib["name"], target))
    return out


def merged_cells(root: ET.Element) -> tuple[set[str], set[str]]:
    anchors, nonanchors = set(), set()
    mc = root.find(local("mergeCells"))
    if mc is None:
        return anchors, nonanchors
    for m in mc.findall(local("mergeCell")):
        ref = m.attrib.get("ref", "")
        cells = expand_ref(ref)
        if cells:
            anchors.add(cells[0])
            nonanchors.update(cells[1:])
        if len(anchors) + len(nonanchors) > MAX_MERGED_TOTAL_CELLS:
            raise ValueError("merged cell ranges exceed safety bound")
    return anchors, nonanchors


def _validated_anchors(anchors) -> dict:
    """The variant -> anchor-list map, or ``TemplateUnreadable``: an empty
    or malformed anchor list can never classify a workbook."""
    if (not isinstance(anchors, dict) or not anchors
            or any(not isinstance(variant_anchors, list) or not variant_anchors
                   or not all(_anchor_well_formed(a) for a in variant_anchors)
                   for variant_anchors in anchors.values())):
        raise TemplateUnreadable("layout identification config unreadable: "
                                 "anchor list missing or malformed")
    return anchors


def _load_layout_anchors() -> dict:
    """Shipped layout anchors keyed by variant ("x01"/"x02").  A missing,
    unparsable or wrong-schema extract raises ``TemplateUnreadable``: a
    workbook that cannot be identified is never "another form"."""
    try:
        doc = json.loads(LAYOUT_EXTRACT_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise TemplateUnreadable(f"layout identification config unreadable: {exc}") from exc
    if not isinstance(doc, dict) or doc.get("schema") != "knuaf-common-workbook-extract/v1":
        raise TemplateUnreadable("layout identification config unreadable: schema mismatch")
    return _validated_anchors(doc.get("layout_anchors"))


def _anchor_well_formed(anchor) -> bool:
    """An anchor is {sheet: str, cell: CELL_RE, expect: exactly one of
    equals|empty|present}.  An empty or malformed anchor list never
    classifies a variant."""
    if not isinstance(anchor, dict):
        return False
    if not isinstance(anchor.get("sheet"), str):
        return False
    if not (isinstance(anchor.get("cell"), str)
            and CELL_RE.match(anchor["cell"])):
        return False
    expect = anchor.get("expect")
    if not isinstance(expect, dict):
        return False
    if sorted(expect) not in (["empty"], ["equals"], ["present"]):
        return False
    if "empty" in expect and expect["empty"] is not True:
        return False
    if "present" in expect and expect["present"] is not True:
        return False
    return True


def _anchor_holds(expect, cell_state) -> bool:
    if not isinstance(expect, dict) or not isinstance(cell_state, dict):
        return False
    value = cell_state.get("value")
    # A formula whose cached value was invalidated (or never stored)
    # is still content: "present" holds and "empty" does not.
    has_formula = cell_state.get("formula") is True
    if "equals" in expect:
        return value == expect["equals"]
    if "empty" in expect:
        return not has_formula and value in (None, "")
    if "present" in expect:
        return has_formula or value not in (None, "")
    return False


class TemplateUnreadable(ValueError):
    """The workbook could not be opened or parsed.  Callers report this as
    blocked — never as a pass and never as "not this template"."""


class TemplateRead:
    """One parsed workbook: sheet order, every stored cell, layout.

    ``cells[(sheet, ref)]`` holds ``kind`` (``f`` formula, ``n`` number,
    ``s`` string, ``b``/``e``/``d``), the normalized ``value``, and for a
    formula its stored ``formula`` text, ``fattrs`` (every ``<f>``
    attribute) and ``resolved`` text (shared followers translated from
    their anchor).  ``display`` is the cell as ``text_value`` reads it,
    cached formula results included."""
    __slots__ = ("path", "sheets", "cells", "layout")

    def __init__(self, path, sheets, cells, layout):
        self.path, self.sheets, self.cells, self.layout = path, sheets, cells, layout


def _cell_record(cell: ET.Element, shared: list[str]) -> dict | None:
    f = cell.find(local("f"))
    display = text_value(cell, shared)
    if f is not None:
        return {"kind": "f", "value": None, "formula": f.text or "",
                "fattrs": tuple(sorted(f.attrib.items())), "display": display}
    t = cell.attrib.get("t", "n")
    v = cell.find(local("v"))
    if t == "inlineStr":
        kind, value = "s", "".join(cell.itertext())
    elif v is None or v.text is None:
        return None
    elif t == "s":
        try:
            kind, value = "s", shared[int(v.text)]
        except (ValueError, IndexError):
            raise TemplateUnreadable(f"shared string index unreadable: {cell.attrib.get('r')}")
    elif t == "str":
        kind, value = "s", v.text
    elif t == "n":
        try:
            kind, value = "n", repr(float(v.text))
        except ValueError:
            raise TemplateUnreadable(f"number cell unreadable: {cell.attrib.get('r')}")
    elif t in ("b", "e", "d"):
        kind, value = t, v.text
    else:
        raise TemplateUnreadable(f"unknown cell type {t!r}: {cell.attrib.get('r')}")
    if kind == "s" and value == "":
        return None
    return {"kind": kind, "value": value, "display": display}


def _layout_of(cells: dict) -> str:
    # An unreadable or malformed identification config raises: it is never
    # "another form".
    anchors = _validated_anchors(_load_layout_anchors())
    for variant, variant_anchors in anchors.items():
        states = []
        for a in variant_anchors:
            rec = cells.get((a["sheet"], a["cell"]))
            states.append({} if rec is None else
                          {"value": rec["display"], "formula": rec["kind"] == "f"})
        if all(_anchor_holds(a["expect"], s) for a, s in zip(variant_anchors, states)):
            return variant
    return "unknown"


def _opened_sheet_names(path: Path) -> list[str]:
    """First check of the shared read: open the file with openpyxl, the
    library the later checks read it with.  Any failure is unreadable."""
    import warnings
    import openpyxl
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(path)
    except Exception as exc:  # openpyxl raises many parser types
        raise TemplateUnreadable(f"workbook does not open: {path.name}: "
                                 f"{type(exc).__name__}: {exc}") from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


_SHEET_ROOTS = {local("worksheet"), local("chartsheet")}
_OFFICE_DOCUMENT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
_SHARED_STRINGS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"
_REQUIRED_PARTS = (("[Content_Types].xml", "{http://schemas.openxmlformats.org/package/2006/content-types}Types"),
                   ("_rels/.rels", f"{{{NS_PKG}}}Relationships"),
                   ("xl/workbook.xml", local("workbook")),
                   ("xl/_rels/workbook.xml.rels", f"{{{NS_PKG}}}Relationships"))


def _part_root(z: zipfile.ZipFile, name: str, tag: str) -> ET.Element:
    if name not in z.namelist():
        raise TemplateUnreadable(f"package part missing: {name}")
    try:
        root = ET.fromstring(z.read(name))
    except ET.ParseError as exc:
        raise TemplateUnreadable(f"package part is not XML: {name}: {exc}") from exc
    if root.tag != tag:
        raise TemplateUnreadable(f"package part has the wrong root: {name}")
    return root


def _package_shared_strings(z: zipfile.ZipFile) -> list[str]:
    """Package structure half of the first check (openpyxl opens files
    whose ``_rels/.rels`` is missing or broken): the four required parts
    exist and parse with their roots, the package relationship points at
    ``xl/workbook.xml``, and a listed shared-string part exists and
    parses.  Returns the shared strings."""
    roots = {name: _part_root(z, name, tag) for name, tag in _REQUIRED_PARTS}
    office = [r for r in roots["_rels/.rels"] if r.attrib.get("Type") == _OFFICE_DOCUMENT]
    if len(office) != 1 or office[0].attrib.get("Target", "").lstrip("/") != "xl/workbook.xml":
        raise TemplateUnreadable("package relationship does not point at xl/workbook.xml")
    listed = [r for r in roots["xl/_rels/workbook.xml.rels"] if r.attrib.get("Type") == _SHARED_STRINGS]
    if listed or "xl/sharedStrings.xml" in z.namelist():
        root = _part_root(z, "xl/sharedStrings.xml", local("sst"))
        return ["".join(si.itertext()) for si in root.findall(local("si"))]
    return []


def read_template(source: Path) -> TemplateRead:
    """The shared read check every template tool runs first.

    1. The file opens with openpyxl, and the package structure holds:
       ``[Content_Types].xml``, ``_rels/.rels`` (pointing at the workbook),
       ``xl/workbook.xml`` and ``xl/_rels/workbook.xml.rels`` exist and
       parse with their roots; a listed ``xl/sharedStrings.xml`` exists and
       parses; at least one sheet; every listed sheet part exists with a
       worksheet (with ``sheetData``) or chartsheet root; the sheet list
       equals the one openpyxl opened.
    2. Every cell parses (``t="s"`` indices inside the shared-string
       table) and every shared formula resolves.
    3. The layout is identified with the shipped anchors.
    Any failure — including an unreadable identification config — raises
    ``TemplateUnreadable``.  ``layout == "unknown"`` therefore always
    means "read and identified as another form"."""
    from gg_workbook_audit import translate
    try:
        path = resolve_source(Path(source))
    except FileNotFoundError as exc:
        raise TemplateUnreadable(str(exc)) from exc
    opened = _opened_sheet_names(path)
    try:
        with zipfile.ZipFile(path) as z:
            shared = _package_shared_strings(z)
            sheets = workbook_sheets(z)
            if not sheets:
                raise TemplateUnreadable("workbook lists no sheets")
            if [name for name, _ in sheets] != opened:
                raise TemplateUnreadable("sheet list differs from the opened workbook")
            for name, target in sheets:
                if target not in z.namelist():
                    raise TemplateUnreadable(f"sheet part missing: {name!r} ({target})")
            cells = {}
            for name, target in sheets:
                root = ET.fromstring(z.read(target))
                if root.tag not in _SHEET_ROOTS or (
                        root.tag == local("worksheet") and root.find(local("sheetData")) is None):
                    raise TemplateUnreadable(f"sheet part is not a worksheet: {name!r} ({target})")
                anchors, followers = {}, []
                for c in root.iter(local("c")):
                    ref = c.attrib.get("r")
                    if not ref or not CELL_RE.match(ref):
                        raise TemplateUnreadable(f"cell without a valid address in {name!r}")
                    rec = _cell_record(c, shared)
                    if rec is None:
                        continue
                    if rec["kind"] == "f":
                        attrs = dict(rec["fattrs"])
                        if attrs.get("t") == "shared" and not rec["formula"]:
                            followers.append((ref, rec, attrs.get("si")))
                        else:
                            rec["resolved"] = rec["formula"]
                            if attrs.get("t") == "shared":
                                anchors[attrs.get("si")] = (ref, rec["formula"])
                    cells[name, ref] = rec
                for ref, rec, si in followers:
                    if si not in anchors:
                        raise TemplateUnreadable(f"shared formula anchor missing: {name}!{ref}")
                    anchor, text = anchors[si]
                    rec["resolved"] = translate(text, anchor, ref)
    except TemplateUnreadable:
        raise
    except (OSError, KeyError, ValueError, IndexError, ET.ParseError, zipfile.BadZipFile) as exc:
        raise TemplateUnreadable(f"workbook unreadable: {Path(source).name}: {exc}") from exc
    return TemplateRead(path, [name for name, _ in sheets], cells, _layout_of(cells))


def detect_layout(source: Path) -> str:
    """Classify the source as ``x01`` | ``x02`` | ``unknown`` by checking
    the shipped C1 ``layout_anchors`` against the source bytes.

    A variant is named only when every one of its anchors holds.  The
    column-J predicate is exclusive, so both variants can never hold
    together.  An empty or malformed anchor list never classifies.
    ``unknown`` means the workbook was read and is another form; a file
    or identification config that cannot be read raises
    ``TemplateUnreadable`` (callers block or protect, never pass)."""
    return read_template(source).layout


def _check_map_layout(map_data: dict, source: Path) -> None:
    """Refuse a map whose declared layout cannot cover this source.

    Detection re-reads the source bytes on every call — a map status is
    never trusted.  Missing ``layout`` (a legacy map) is treated as
    ``default_x01``; ``reviewed_custom`` is accepted only when its
    variant equals the detected one.  Every other origin or status is
    ``map_layout_invalid``."""
    detected = detect_layout(source)
    layout = map_data.get("layout")
    if layout is None:
        layout = {"mapOrigin": "default_x01"}
    if not isinstance(layout, dict):
        raise ValueError("map_layout_invalid")
    origin = layout.get("mapOrigin")
    status = map_data.get("mapStatus", "ok")
    if origin == "default_x01":
        if detected != "x01":
            raise ValueError("layout_variant_unsupported")
        if status != "ok":
            raise ValueError("map_layout_invalid")
        return
    if origin == "reviewed_custom":
        if layout.get("variant") != detected or status != "ok":
            raise ValueError("map_layout_invalid")
        return
    raise ValueError("map_layout_invalid")


def inspect_source(source: Path, out_map: Path | None, report: Path | None) -> dict:
    source = resolve_source(source).resolve()
    # Shared read check first: an unreadable file stops here (blocked).
    detected = read_template(source).layout
    if out_map and out_map.exists():
        raise FileExistsError(f"refusing to overwrite map: {out_map}")
    if report and report.exists():
        raise FileExistsError(f"refusing to overwrite report: {report}")
    digest = sha256(source)
    with zipfile.ZipFile(source) as z:
        shared = load_shared(z)
        sheets = workbook_sheets(z)
        inventory = []
        sheet_summaries = []
        for idx, (name, target) in enumerate(sheets):
            root = ET.fromstring(z.read(target))
            cells = []
            formulas = 0
            for c in root.findall(f".//{local('c')}"):
                ref = c.attrib.get("r")
                if not ref:
                    continue
                f = c.find(local("f"))
                if f is not None:
                    formulas += 1
                cells.append({"cell": ref, "value": text_value(c, shared),
                              "formula": f.text if f is not None else None,
                              "hasFormula": f is not None,
                              "style": c.attrib.get("s")})
            sheet_summaries.append({"name": name, "index": idx, "xml": target,
                                    "cellCount": len(cells), "formulaCount": formulas})
            inventory.extend({"sheet": name, **c} for c in cells)
    entries = default_entries(inventory)
    # ``detected`` comes from the shared read check above.
    result = {
        "schema": "gg-xlsx-template-map/v1",
        "source": {"path": str(source), "sha256": digest, "size": source.stat().st_size},
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "sheets": sheet_summaries,
        "entries": entries,
        "inventory": inventory,
        "layout": {"variant": detected, "mapOrigin": "default_x01"},
        "mapStatus": ("ok" if detected == "x01"
                     else "manual_map_required"),
        "mapReason": (None if detected == "x01"
                      else "layout_variant_unsupported"),
        "policy": {
            "clearOnlyListed": True,
            "protectFormulas": True,
            "protectMergedNonAnchors": True,
            "invalidateFormulaCaches": True,
            "sourceVersionFailClosed": True,
            "noOverwrite": True,
            "legacyErrorsPreserved": True,
            "unclassifiedTextDisclosed": True,
            "publicationReady": False,
            "referenceSheetPreserved": "참고1. 모델농장분석",
        },
    }
    if out_map:
        out_map.parent.mkdir(parents=True, exist_ok=True)
        out_map.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    if report:
        report.parent.mkdir(parents=True, exist_ok=True)
        report.write_text(json.dumps({k: result[k] for k in ("schema", "source", "sheets", "entries", "policy")},
                                     ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def e(sheet: str, refs: str, semantic: str, reason: str, action: str = "clear") -> dict:
    return {"sheet": sheet, "range": refs, "semanticField": semantic, "reason": reason,
            "action": action, "explicit": True}


def wage_fill_map(source: Path) -> dict:
    """Value-only map; formula observations must first be explicitly prepared.

    The fill engine requires a complete, sourced six-year observation block.
    This map never grants permission to overwrite a formula.
    """
    if read_template(source).layout != "x01":
        raise ValueError("layout_variant_unsupported: wage observations require x01")
    return {"schema": "gg-xlsx-fill-map/v1",
            "template": {"sha256": sha256(source)},
            "entries": [{"sheet": "8. 노무비계획", "range": "X26:Z31",
                         "role": "wage.observation", "period": "observation",
                         "source_note": "raw years and both wage columns must match registered observations and target gender before explicit hourly conversion; gender-specific single sources cannot fill both columns; no default hours",
                         "editable": True}]}


def template_price_sources(price_sources: list) -> dict:
    """Index the template's explicit list using the shared entry validator."""
    from gg_price_assumptions import check_source_entry
    if not isinstance(price_sources, list):
        raise ValueError("price_sources must be a list")
    sources = {}
    for entry in price_sources:
        sid = check_source_entry(entry, origin="price_sources")
        if sid in sources:
            raise ValueError(f"duplicate price_sources id: {sid}")
        sources[sid] = entry
    return sources


def wage_observation_source(metadata: dict, price_sources: list) -> dict:
    """Resolve observation provenance through the shared price contract."""
    from gg_price_assumptions import (
        load_registry, resolve_series, check_role, check_unit, check_base_year,
    )
    if "base_year" not in metadata:
        raise ValueError("wage_observations.base_year required; nominal series use null")
    source = resolve_series(metadata.get("statistic_id"), registry=load_registry(),
                            student_sources=template_price_sources(price_sources))
    check_role(source, "wage", origin="wage_observations")
    check_unit(source, ("원/일", "원/시간"), origin="wage_observations")
    check_base_year(metadata["base_year"], source, origin="wage_observations")
    return source


def check_wage_gender(source: dict, gender: str, *, origin: str) -> None:
    """A gender-specific series can supply only its matching school column.

    The single-source observation format fills both columns, so such a
    series is rejected there. Separate male/female rate inputs can use it.
    """
    from gg_price_assumptions import canonical_dimensions
    dimensions = canonical_dimensions(source.get("dimensions"))
    for key, value in dimensions.items():
        if key.casefold() in {"gender", "sex", "성별"} and value != gender:
            raise ValueError(f"{origin}: source gender {key}={value!r} "
                             f"does not match target gender {gender!r}; "
                             "separate gender sources required")


def default_entries(inventory: list[dict]) -> list[dict]:
    """Return an explicit cell map.  Rectangles that contain labels are avoided."""
    out = [
        e("1. 기초재무상태조사", "B9:F9", "existing_land_asset", "sample author's land facts"),
        e("1. 기초재무상태조사", "H9", "existing_land_asset_note", "sample author's land note"),
        e("1. 기초재무상태조사", "C39", "cash_account_name", "sample cash account"),
        e("1. 기초재무상태조사", "E39", "cash_and_equivalents", "sample cash balance"),
        e("3.투자계획", "B8:C8", "investment_summary_inputs", "sample first-year investment summary"),
        e("3.투자계획", "E8", "investment_summary_subsidy", "sample first-year subsidy input"),
        e("3.투자계획", "B19:D23", "annual_investment_plan", "sample investment detail facts"),
        e("3.투자계획", "E20:E23", "annual_investment_amounts", "sample investment amounts"),
        e("3.투자계획", "F19:H23", "annual_investment_funding", "sample funding sources"),
        e("3.투자계획", "I20:I22", "annual_investment_notes", "sample investment notes"),
        e("3.투자계획", "G28:H28", "annual_investment_funding_later", "sample later-year funding"),
        e("3.투자계획", "G25:H27", "annual_investment_funding_later", "sample later-year funding"),
        e("4. 원리금상환계획", "B38:C38", "loan_terms", "sample loan product and timing"),
        e("4. 원리금상환계획", "E38:F38", "loan_terms", "sample loan condition and rate"),
        e(" 5. 판매계획", "D22:L24", "sales_channel_mix", "sample procurement channel percentages"),
        e(" 5. 판매계획", "D26:L28", "sales_channel_mix", "sample direct channel percentages"),
        e(" 5. 판매계획", "D30", "sales_channel_mix_residual", "sample residual channel percentage"),
        e("6. 생산계획", "H8:T8", "production_area", "sample cultivation area inputs"),
        e("6. 생산계획", "H9:T9", "production_output", "sample production output inputs"),
        e("6. 생산계획", "H15:H17", "production_basis_notes", "sample production basis narrative"),
        e("6. 생산계획", "H19:H24", "production_basis_notes", "sample production basis narrative"),
        e("6. 생산계획", "H25:H28", "production_basis_notes", "sample production basis narrative"),
        e("6. 생산계획", "E38", "legacy_note_artifact", "stray sample note artifact"),
        e("7. 영농자재소요계획", "D15:E15", "annual_material_area_basis", "sample first-year area basis"),
        e("7. 영농자재소요계획", "C19", "annual_material_other_cost", "sample first-year other material cost"),
        e("7. 영농자재소요계획", "G19", "annual_material_other_amount", "sample first-year other amount"),
        e("7. 영농자재소요계획", "D25:E25", "annual_material_area_basis", "sample second-year area basis"),
        e("7. 영농자재소요계획", "C29", "annual_material_other_cost", "sample second-year other material cost"),
        e("7. 영농자재소요계획", "G29", "annual_material_other_amount", "sample second-year other amount"),
        e("7. 영농자재소요계획", "D35:E35", "annual_material_area_basis", "sample third-year area basis"),
        e("7. 영농자재소요계획", "C39", "annual_material_other_cost", "sample third-year other material cost"),
        e("7. 영농자재소요계획", "G39", "annual_material_other_amount", "sample third-year other amount"),
        e("7. 영농자재소요계획", "D45:E45", "annual_material_area_basis", "sample fourth-year area basis"),
        e("7. 영농자재소요계획", "C49", "annual_material_other_cost", "sample fourth-year other material cost"),
        e("7. 영농자재소요계획", "G49", "annual_material_other_amount", "sample fourth-year other amount"),
        e("7. 영농자재소요계획", "D55:E55", "annual_material_area_basis", "sample fifth-year area basis"),
        e("7. 영농자재소요계획", "C59", "annual_material_other_cost", "sample fifth-year other material cost"),
        e("7. 영농자재소요계획", "G59", "annual_material_other_amount", "sample fifth-year other amount"),
        e("8. 노무비계획", "G8:J8", "self_labor_days", "sample self-labor workdays"),
        e("8. 노무비계획", "F11:G12", "hired_labor_rates_days", "sample hired-labor rates and days"),
        e("8. 노무비계획", "J11:S12", "hired_labor_days", "sample hired-labor workdays"),
        e("8. 노무비계획", "I26:K30", "labor_detail_self_days", "sample first-year labor detail days"),
        e("8. 노무비계획", "O26:S30", "labor_detail_hired_days", "sample first-year hired detail days"),
        e("8. 노무비계획", "I37:K41", "labor_detail_self_days", "sample later-year labor detail days"),
        e("8. 노무비계획", "O37:S41", "labor_detail_hired_days", "sample later-year hired detail days"),
        e("8. 노무비계획", "B17:B20", "labor_basis_notes", "sample wage and workday assumptions"),
        e("9 .경비계획", "D5:D10", "expense_assumptions", "sample expense rates and areas"),
        e("9 .경비계획", "C23", "biological_asset_expense", "sample zero biological-asset expense"),
        e("9 .경비계획", "C43", "biological_asset_expense", "sample zero biological-asset expense"),
        e("9 .경비계획", "C49", "other_expense", "sample zero other expense"),
        e("9 .경비계획", "C63", "biological_asset_expense", "sample zero biological-asset expense"),
        e("9 .경비계획", "C69", "other_expense", "sample zero other expense"),
        e("9 .경비계획", "C83", "biological_asset_expense", "sample zero biological-asset expense"),
        e("9 .경비계획", "C89", "other_expense", "sample zero other expense"),
        e("9 .경비계획", "C103", "biological_asset_expense", "sample zero biological-asset expense"),
        e("9 .경비계획", "C109", "other_expense", "sample zero other expense"),
        e("10. 감가상각비계획 ", "C9:D10", "depreciation_assumptions", "sample facility assumptions"),
        e("10. 감가상각비계획 ", "F9:H10", "depreciation_assumptions", "sample machinery assumptions"),
        e("10. 감가상각비계획 ", "C11:H11", "depreciation_years_used", "sample asset years used"),
        e("11. 생산원가계획", "C24:G24", "biological_asset_cost", "sample zero biological-asset cost"),
        e("11. 생산원가계획", "C28:G28", "contract_farming_cost", "sample zero contract-farming cost"),
        e("12 손익계획", "E30:E33", "packaging_cost_inputs", "sample packaging inputs"),
        e("12 손익계획", "C35:G35", "packaging_production_quantity", "sample packaging quantity inputs"),
        e("12 손익계획", "B38", "packaging_note", "sample packaging calculation narrative"),
        e("14. 현금흐름계획", "E15:E16", "later_year_investment_cashflow", "sample later-year investment cashflow"),
        e("14. 현금흐름계획", "D17", "perennial_investment_cashflow", "sample perennial investment cashflow"),
        # This sheet is an example-only reference and is deliberately retained.
        e("참고1. 모델농장분석", "A1:E19", "reference_model_farm", "example-only reference sheet; retain verbatim", "preserve"),
    ]

    # Hard-coded legacy outputs and public benchmark tables are retained and
    # explicitly classified so they cannot be mistaken for missing input cells.
    preserve = [
        ("2. 중장기영농목표", "C12", "zero-denominator guard; derived output stored as a legacy constant"),
        ("3.투자계획", "F9:F12", "summary financing zeros; legacy derived constants"),
        ("3.투자계획", "B12:F12", "year-5 summary row; legacy structure/constants"),
        ("4. 원리금상환계획", "B10:F32", "repayment schedule outputs; legacy constants without formulas"),
        ("4. 원리금상환계획", "B44:F66", "loan detail schedule outputs; legacy constants without formulas"),
        (" 5. 판매계획", "M6:M11", "public benchmark years/prices; example reference data"),
        (" 5. 판매계획", "E7:M13", "public benchmark price history; example reference data"),
        (" 5. 판매계획", "D25:M25", "channel totals; legacy constants"),
        (" 5. 판매계획", "D29:M29", "channel totals; legacy constants"),
        (" 5. 판매계획", "D31:M31", "channel totals; legacy constants"),
        (" 5. 판매계획", "C43:K43", "sales output row; formula/cache legacy constants"),
        ("8. 노무비계획", "W7:Y19", "public wage benchmark/projection table; legacy constants"),
        ("8. 노무비계획", "U28:U41", "labor detail totals; legacy constants"),
        ("8. 노무비계획", "G13", "labor subtotal zero; legacy constant"),
        ("8. 노무비계획", "P13", "labor subtotal zero; legacy constant"),
        ("10. 감가상각비계획 ", "H20:H22", "later-year asset assumptions copied from year one; legacy constants"),
        ("10. 감가상각비계획 ", "G23:H23", "later-year usage years; derived constants"),
        ("10. 감가상각비계획 ", "H32:H34", "later-year asset assumptions copied from year one; legacy constants"),
        ("10. 감가상각비계획 ", "G35:H35", "later-year usage years; derived constants"),
        ("10. 감가상각비계획 ", "H44:H46", "later-year asset assumptions copied from year one; legacy constants"),
        ("10. 감가상각비계획 ", "G47:H47", "later-year usage years; derived constants"),
        ("10. 감가상각비계획 ", "H56:H58", "later-year asset assumptions copied from year one; legacy constants"),
        ("10. 감가상각비계획 ", "G59:H59", "later-year usage years; derived constants"),
        ("11. 생산원가계획", "E5:G5", "production cost outputs; legacy constants"),
        ("12 손익계획", "C11:G13", "selling/admin outputs; legacy constants"),
        ("13. 추정대차대조표", "D15:H15", "balance-sheet outputs; legacy constants"),
        ("13. 추정대차대조표", "D26:H32", "balance-sheet outputs; legacy constants"),
        ("15.추정소득분석", "G26:I28", "income analysis outputs; legacy constants"),
        ("15.추정소득분석", "G31:I32", "income analysis outputs; legacy constants"),
        ("15.추정소득분석", "E28", "zero-denominator guard; derived output stored as a legacy constant"),
        ("15.추정소득분석", "E32", "zero-denominator guard; derived output stored as a legacy constant"),
        (" 5. 판매계획", "B14:B15", "public benchmark source notes; fixed reference text"),
        (" 5. 판매계획", "B32", "public distribution source note; fixed reference text"),
        (" 5. 판매계획", "B45:B46", "pricing and channel method notes; fixed reference text"),
        ("7. 영농자재소요계획", "G5", "public cost-survey source note; fixed reference text"),
        ("7. 영농자재소요계획", "C6:C10", "materials methodology/source notes; fixed reference text"),
        ("9 .경비계획", "D16:D17", "expense calculation basis notes; fixed reference text"),
        ("9 .경비계획", "D20", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D36:D37", "expense calculation basis notes; fixed reference text"),
        ("9 .경비계획", "D40", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D45", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D56:D57", "expense calculation basis notes; fixed reference text"),
        ("9 .경비계획", "D60", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D65", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D76:D77", "expense calculation basis notes; fixed reference text"),
        ("9 .경비계획", "D80", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D85", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D96:D97", "expense calculation basis notes; fixed reference text"),
        ("9 .경비계획", "D100", "expense calculation basis note; fixed reference text"),
        ("9 .경비계획", "D105", "expense calculation basis note; fixed reference text"),
        ("12 손익계획", "B39", "shipping benchmark source note; fixed reference text"),
        ("14. 현금흐름계획", "B25", "cash-flow method note; fixed reference text"),
    ]
    out.extend(e(sheet, refs, "legacy_constant", reason, "preserve") for sheet, refs, reason in preserve)
    for cell in inventory:
        if (cell.get("sheet") == "8. 노무비계획"
                and re.fullmatch(r"[XYZ](?:2[6-9]|3[01])", cell.get("cell", ""))
                and not cell.get("hasFormula", cell.get("formula") is not None)):
            out.append(e(cell["sheet"], cell["cell"], "wage_observation",
                         "replace example observation with sourced official daily wage/year"))
    return out


def map_entries_for_source(map_data: dict, source: Path) -> tuple[list[dict], list[dict]]:
    got = sha256(source)
    expected = map_data.get("source", {}).get("sha256")
    if expected != got:
        raise RuntimeError(f"source version changed: map sha256={expected}, current sha256={got}")
    entries = map_data.get("entries")
    if not isinstance(entries, list) or not entries:
        raise ValueError("map must contain non-empty entries")
    return entries, map_data.get("inventory", [])


_MC_QNAME_ATTRIBUTES = {
    "Ignorable", "MustUnderstand", "ProcessContent", "PreserveAttributes",
    "PreserveElements", "Requires",
}


def _validate_mc_qnames(document: minidom.Document, source_name: str) -> None:
    """Fail closed when MC QName-valued attributes contain undeclared prefixes.

    ElementTree serialization can silently change namespace prefixes while leaving
    lexical values such as ``mc:Ignorable=\"x15\"`` untouched.  A resulting
    workbook then fails to open.  Walk the original DOM with an in-scope prefix
    map so nested namespace rebindings are handled correctly before any mutation.
    """
    def walk(node: minidom.Node, nsmap: dict[str, str]) -> None:
        if node.nodeType != node.ELEMENT_NODE:
            return
        current = dict(nsmap)
        for i in range(node.attributes.length):
            attr = node.attributes.item(i)
            if attr.name == "xmlns":
                current[""] = attr.value
            elif attr.name.startswith("xmlns:"):
                current[attr.name.split(":", 1)[1]] = attr.value
        for i in range(node.attributes.length):
            attr = node.attributes.item(i)
            is_mc_attr = attr.namespaceURI == NS_MC and attr.localName in _MC_QNAME_ATTRIBUTES
            # Choice/@Requires is intentionally unqualified by OOXML, while
            # the other MC controls carry the mc: prefix.
            is_choice_requires = (
                attr.localName == "Requires" and attr.namespaceURI in (None, "")
                and node.namespaceURI == NS_MC
            )
            if not (is_mc_attr or is_choice_requires):
                continue
            for token in attr.value.split():
                if attr.localName in {"ProcessContent", "PreserveAttributes", "PreserveElements"}:
                    # These controls contain QName lists; unprefixed local
                    # names are valid and need no prefix declaration.
                    if ":" not in token:
                        continue
                    prefix = token.split(":", 1)[0]
                else:
                    # Ignorable/MustUnderstand/Requires contain namespace
                    # prefix lists, including bare prefixes such as x15/hs.
                    prefix = token.split(":", 1)[0] if ":" in token else token
                if prefix and prefix not in current:
                    raise ValueError(
                        f"undeclared namespace prefix {prefix!r} in mc:{attr.localName} "
                        f"of {source_name}"
                    )
        for child in node.childNodes:
            if child.nodeType == child.ELEMENT_NODE:
                walk(child, current)

    walk(document.documentElement, {})


def _dom_elements(document: minidom.Document, namespace: str, local_name: str) -> list[minidom.Element]:
    """Return namespace-aware DOM elements without relying on lexical prefixes."""
    return list(document.getElementsByTagNameNS(namespace, local_name))


def _dom_attr(element: minidom.Element, name: str) -> str | None:
    return element.getAttribute(name) if element.hasAttribute(name) else None


def _dom_clear_cell_value(cell: minidom.Element) -> None:
    for child in list(cell.childNodes):
        if child.nodeType == child.ELEMENT_NODE and (
            child.namespaceURI == NS_MAIN and child.localName in {"v", "is"}
        ):
            cell.removeChild(child)
    if cell.hasAttribute("t"):
        cell.removeAttribute("t")


def _dom_invalidate_formula_cache(cell: minidom.Element) -> bool:
    formula = next((child for child in cell.childNodes
                    if child.nodeType == child.ELEMENT_NODE
                    and child.namespaceURI == NS_MAIN and child.localName == "f"), None)
    if formula is None:
        return False
    value = next((child for child in cell.childNodes
                  if child.nodeType == child.ELEMENT_NODE
                  and child.namespaceURI == NS_MAIN and child.localName == "v"), None)
    if value is not None:
        for child in list(value.childNodes):
            value.removeChild(child)
    return True


def _serialize_dom(xml_bytes: bytes, source_name: str) -> minidom.Document:
    """Parse and validate an XML part while retaining original namespace lexemes."""
    document = minidom.parseString(xml_bytes)
    _validate_mc_qnames(document, source_name)
    return document


def blank_copy(source: Path, map_path: Path, out: Path, *, context=None) -> dict:
    import gg_major_contract as mc

    authorization = mc.authorize_output(mc.OUTPUT_SCHOOL_WORKBOOK, context)
    mc.require_finance_profile(authorization, "school_17_sheet_v1")
    if out.exists():
        raise FileExistsError(f"refusing to overwrite output: {out}")
    map_data = json.loads(map_path.read_text(encoding="utf-8"))
    entries, _ = map_entries_for_source(map_data, source)
    _check_map_layout(map_data, source)
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source, "r") as zin:
        sheets = dict(workbook_sheets(zin))
        shared = load_shared(zin)
        by_sheet: dict[str, set[str]] = {}
        preserved_by_sheet: dict[str, dict[str, str]] = {}
        action_by_cell: dict[tuple[str, str], str] = {}
        for ent in entries:
            if "action" not in ent:
                raise ValueError("map entry missing explicit action")
            action = ent["action"]
            if action not in {"clear", "preserve"}:
                raise ValueError(f"invalid map action: {action!r}")
            if ent.get("sheet") not in sheets:
                raise ValueError(f"map references missing sheet: {ent.get('sheet')}")
            refs = ent.get("range") or ent.get("cell")
            if not refs:
                raise ValueError("map entry missing range/cell")
            expanded = expand_ref(refs)
            for cell in expanded:
                key = (ent["sheet"], cell)
                prior = action_by_cell.get(key)
                if prior and prior != action:
                    raise ValueError(f"conflicting map actions for {ent['sheet']}!{cell}")
                action_by_cell[key] = action
            if action == "clear":
                by_sheet.setdefault(ent["sheet"], set()).update(expanded)
            else:
                reason = ent.get("reason", "mapped preserve cell")
                preserved_by_sheet.setdefault(ent["sheet"], {}).update({r: reason for r in expanded})
        modified: dict[str, bytes] = {}
        receipt = {"schema": "gg-xlsx-template-receipt/v1", "source": {"path": str(source.resolve()),
                    "sha256": sha256(source), "size": source.stat().st_size},
                   "output": {"path": str(out.resolve())}, "cleared": [], "preserved": [],
                   "formulaCellsProtected": 0, "mergedCellsProtected": 0,
                   "formulaCachesInvalidated": 0, "ambiguousCount": 0, "errors": []}
        for name, target in sheets.items():
            sheet_bytes = zin.read(target)
            root = ET.fromstring(sheet_bytes)
            # Mutate the original namespace-aware DOM. ElementTree's serializer
            # can rewrite prefixes and drop declarations referenced by MC QName
            # values such as mc:Ignorable="x15".
            sheet_dom = _serialize_dom(sheet_bytes, target)
            dom_cells = {
                _dom_attr(cell, "r"): cell
                for cell in _dom_elements(sheet_dom, NS_MAIN, "c")
                if _dom_attr(cell, "r")
            }
            anchors, nonanchors = merged_cells(root)
            wanted = by_sheet.get(name, set())
            seen_wanted: set[str] = set()
            for c in root.findall(f".//{local('c')}"):
                ref = c.attrib.get("r")
                if not ref:
                    continue
                dom_cell = dom_cells.get(ref)
                if dom_cell is None:
                    raise ValueError(f"worksheet cell missing from namespace-preserving DOM: {target}!{ref}")
                if ref in wanted:
                    seen_wanted.add(ref)
                formula = c.find(local("f"))
                if formula is not None:
                    # Invalidate every formula cache in a blank template; keep f text.
                    v = c.find(local("v"))
                    if v is not None and v.text is not None:
                        _dom_invalidate_formula_cache(dom_cell)
                        receipt["formulaCachesInvalidated"] += 1
                    if ref in wanted:
                        receipt["formulaCellsProtected"] += 1
                        receipt["preserved"].append({"sheet": name, "cell": ref, "reason": "formula"})
                    continue
                if ref not in wanted:
                    if ref in preserved_by_sheet.get(name, {}):
                        receipt["preserved"].append({"sheet": name, "cell": ref,
                                                     "reason": preserved_by_sheet[name][ref]})
                    continue
                if ref in nonanchors:
                    receipt["mergedCellsProtected"] += 1
                    receipt["preserved"].append({"sheet": name, "cell": ref, "reason": "merged_nonanchor"})
                    continue
                # Explicitly mapped data cell: remove value while retaining style and cell node.
                old = text_value(c, shared)
                _dom_clear_cell_value(dom_cell)
                receipt["cleared"].append({"sheet": name, "cell": ref, "oldValue": old})
            missing_wanted = wanted - seen_wanted
            if missing_wanted:
                # A map entry names a cell that has no <c> element in this
                # source worksheet at all, so the loop above never inspected
                # it. Silently accepting this risks masking a map generated
                # against a different/stale copy of the source. Fail closed
                # instead of treating "never seen" as "nothing to clear".
                raise ValueError(
                    "map references cell missing from source worksheet: "
                    f"{name}!{sorted(missing_wanted)[0]}"
                )
            modified[target] = sheet_dom.toxml(encoding="utf-8")
        targeted = {(r["sheet"], r["cell"]) for r in receipt["cleared"]}
        preserved = {(r["sheet"], r["cell"]) for r in receipt["preserved"]}
        # Hard-coded numerics left outside the explicit map are intentionally reported
        # as review items; they may be valid benchmarks, controls, or legacy sample data.
        unclassified = []
        unclassified_text = []
        for name, target in sheets.items():
            root = ET.fromstring(zin.read(target))
            for c in root.findall(f".//{local('c')}"):
                ref = c.attrib.get("r")
                if not ref or c.find(local("f")) is not None or (name, ref) in targeted or (name, ref) in preserved:
                    continue
                val = text_value(c, shared)
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    unclassified.append({"sheet": name, "cell": ref, "value": val})
                elif isinstance(val, str) and val and is_reference_text(val):
                    # Narrative/reference text outside the explicit map is
                    # retained but disclosed. Ordinary worksheet labels are
                    # treated as fixed structure and are not ambiguity items.
                    unclassified_text.append({"sheet": name, "cell": ref, "value": val})
        receipt["unclassifiedNumericCells"] = len(unclassified)
        receipt["unclassifiedNumericSample"] = unclassified[:40]
        receipt["unclassifiedTextCells"] = len(unclassified_text)
        receipt["unclassifiedTextSample"] = unclassified_text[:40]
        receipt["ambiguousCount"] = len(unclassified)
        receipt["classification"] = {
            "clearMapped": len(receipt["cleared"]),
            "preserveMapped": len(receipt["preserved"]),
            "unclassifiedNumeric": len(unclassified),
            "unclassifiedText": len(unclassified_text),
            "unclassifiedValuesRemain": bool(unclassified or unclassified_text),
        }
        # Force Excel recalculation on open without touching formulas.
        try:
            wb_dom = _serialize_dom(zin.read("xl/workbook.xml"), "xl/workbook.xml")
            calc_nodes = _dom_elements(wb_dom, NS_MAIN, "calcPr")
            calc = calc_nodes[0] if calc_nodes else None
            if calc is None:
                calc = wb_dom.createElementNS(NS_MAIN, "calcPr")
                wb_dom.documentElement.appendChild(calc)
            for key, value in (("calcMode", "auto"), ("fullCalcOnLoad", "1"), ("forceFullCalc", "1")):
                calc.setAttribute(key, value)
            modified["xl/workbook.xml"] = wb_dom.toxml(encoding="utf-8")
        except KeyError:
            receipt["errors"].append("workbook.xml missing; calculation mode not updated")
        mc.reconfirm_output(authorization, context)
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = modified.get(info.filename, zin.read(info.filename))
                zout.writestr(info, data)
    receipt["majorAuthorization"] = authorization.to_dict()
    receipt["output"]["sha256"] = sha256(out)
    receipt["output"]["size"] = out.stat().st_size
    # Any retained text outside the map (labels, notes, or legacy sample prose)
    # keeps the artifact review-only; the receipt must not call it publication-ready.
    retained_legacy = any("legacy" in reason.lower() or "reference" in reason.lower()
                          for reasons in preserved_by_sheet.values() for reason in reasons.values())
    receipt["output"]["status"] = "partial" if (receipt["ambiguousCount"] or receipt.get("unclassifiedTextCells") or retained_legacy) else "blank_template"
    receipt["output"]["calculationStatus"] = "recalculate_on_open; cached_formula_values_invalidated"
    return receipt


COMMON_WORKBOOKS_DIR = Path(__file__).resolve().parents[1] / "references" / "common-workbooks"
CROP_LABELS_PATH = COMMON_WORKBOOKS_DIR / "x01-crop-labels.json"
CROP_LABEL_SCHEMA = "knuaf-template-crop-labels/v2"
FORMULA_AUDIT_SCHEMA = "knuaf-workbook-formula-audit/v1"
# The X01/X02 wage observation block: observation preparation removes its
# formulas and the wage fill writes sourced numbers there.
OBSERVATION_BLOCK = ("8. 노무비계획", "X26:Z31")


def _crop_label_table() -> dict:
    table = _json_file(CROP_LABELS_PATH)
    if table.get("schema") != CROP_LABEL_SCHEMA:
        raise ValueError("crop label table schema mismatch")
    return table


def _json_file(path: Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def registered_audit(binding: dict) -> dict:
    """The pinned formula audit of a registered template: the audit file
    must hash to ``audit_sha256`` and describe the template ``sha256``."""
    path = COMMON_WORKBOOKS_DIR / binding["audit_file"]
    if sha256(path) != binding["audit_sha256"]:
        raise ValueError(f"registered template audit hash differs: {binding['audit_file']}")
    audit = _json_file(path)
    if audit.get("schema") != FORMULA_AUDIT_SCHEMA or audit.get("source_sha256") != binding["sha256"]:
        raise ValueError(f"registered template audit source differs: {binding['audit_file']}")
    return audit


def derivation_mismatches(read: TemplateRead, audit: dict, *, open_range=OBSERVATION_BLOCK) -> list[str]:
    """Cell-by-cell proof that ``read`` is the audited template after only
    the defined copy transforms.  Every audited formula must be present
    with the same resolved text (a shared formula expanded to plain
    formulas is the same formula), no other cell may hold a formula, and
    the sheet set must match.  Only the observation block (``open_range``)
    may have lost its formulas — prepared, or filled with numbers.  Value
    and label cells are the fill path and are not compared here."""
    out = []
    if set(read.sheets) != set(audit.get("sheet_roles", {})):
        out.append("sheet set differs: " + ", ".join(sorted(set(read.sheets) ^ set(audit.get("sheet_roles", {})))))
    expected = {(row["sheet"], row["cell"]): row["formula"] for row in audit["inventory"]}
    actual = {key: rec["resolved"] for key, rec in read.cells.items() if rec["kind"] == "f"}
    sheet, refs = open_range
    open_cells = {(sheet, ref) for ref in expand_ref(refs)}
    for key in sorted(set(expected) | set(actual)):
        want, got = expected.get(key), actual.get(key)
        if want == got or (key in open_cells and want is not None and got is None):
            continue
        out.append(f"{key[0]}!{key[1]}: registered formula {want!r}, found {got!r}")
    return out


def registered_template(read: TemplateRead, bindings: dict) -> tuple[str, str]:
    """``(ref, "registered"|"derived")`` for a workbook that is a registered
    template or a copy of one after only the defined transforms; raises
    ValueError naming the first differing cells otherwise."""
    digest = sha256(read.path)
    for ref, binding in bindings.items():
        if digest == binding["sha256"]:
            return ref, "registered"
    problems = []
    for ref, binding in bindings.items():
        if read.layout != ref.lower():
            problems.append(f"{ref}: layout {read.layout}")
            continue
        mismatches = derivation_mismatches(read, registered_audit(binding))
        if not mismatches:
            return ref, "derived"
        problems.append(f"{ref}: {len(mismatches)} cells differ ({'; '.join(mismatches[:3])})")
    raise ValueError("not a registered template or a copy derived from one by the defined "
                     "transforms: " + " | ".join(problems))


def crop_label_map(source: Path, category: str, choices: dict | None = None) -> tuple[dict, dict]:
    """Fill map + values that rename only the template's crop-bound labels.

    ``category`` is one of the table's crop categories (asked from the
    student when the major/crop does not settle it).  ``choices`` picks a
    listed alternative per ``"<sheet>!<cell>"``; the first listed label is
    the default.  The source is the registered X01/X02 or a copy of it
    whose every formula still matches the registered audit
    (``registered_template``); each original label must still be there —
    a template whose label moved needs the table re-aligned first."""
    table = _crop_label_table()
    if category not in table["categories"]:
        raise ValueError(f"unknown crop category: {category!r}; "
                         f"choose one of {sorted(table['categories'])}")
    read = read_template(source)
    try:
        ref, admitted = registered_template(read, table["templates"])
    except ValueError as exc:
        raise ValueError("crop labels are aligned to registered X01/X02 only; "
                         "re-align x01-crop-labels.json for a new template year — " + str(exc)) from exc
    digest = sha256(read.path)
    choices = dict(choices or {})
    entries, values = [], []
    for row in table["cells"]:
        key = f"{row['sheet']}!{row['cell']}"
        options = row["labels"][category]
        label = choices.pop(key, options[0])
        if label not in options:
            raise ValueError(f"{key}: {label!r} is not a listed {category} label {options}")
        stored = read.cells.get((row["sheet"], row["cell"]))
        if stored is None or stored["kind"] != "s" or stored["value"] != row["original"]:
            raise ValueError(f"{key}: expected original label {row['original']!r}")
        if label == row["original"]:
            continue
        entries.append({"sheet": row["sheet"], "cell": row["cell"], "role": "label.crop_category",
                        "period": None, "editable": True,
                        "source_note": f"작목 분류 {category} 칸 이름; 원래 글자 {row['original']!r}"})
        values.append({"sheet": row["sheet"], "cell": row["cell"], "value": label,
                       "value_type": "string", "answer_state": "provided",
                       "evidence_ref": {"source_id": "template.crop_labels", "revision": 1,
                                        "locator": f"x01-crop-labels.json {key} {category}; "
                                                   f"{ref} {admitted}",
                                        "origin": "assumption"}})
    if choices:
        raise ValueError("choices name cells outside the label table: " + ", ".join(sorted(choices)))
    fill_map = {"schema": "gg-xlsx-fill-map/v1", "template": {"sha256": digest}, "entries": entries}
    fill_values = {"schema": "gg-xlsx-fill-values/v1", "values": values}
    return fill_map, fill_values


def _compared(rec: dict | None):
    """What label-check compares: formulas by stored text, every ``<f>``
    attribute (shared ``t``/``si``/``ref``) and resolved text; other cells
    by kind and value — a number stored as text is a different cell."""
    if rec is None:
        return None
    if rec["kind"] == "f":
        return ("f", rec["formula"], rec["fattrs"], rec["resolved"])
    return (rec["kind"], rec["value"])


def crop_label_check(before: Path, after: Path) -> dict:
    """Prove a relabel changed only listed label cells: every formula
    (text, attributes, resolved meaning), every value and its type, and
    the sheet list stay identical.  A workbook that cannot be read is
    ``blocked``, never a pass."""
    table = _crop_label_table()
    allowed = {(r["sheet"], r["cell"]): r for r in table["cells"]}
    try:
        a, b = read_template(before), read_template(after)
    except TemplateUnreadable as exc:
        return {"status": "blocked", "reason": str(exc), "labelsChanged": [], "violations": []}
    changed, violations = [], []
    if a.sheets != b.sheets:
        violations.append({"sheets": {"old": a.sheets, "new": b.sheets}})
    for key in sorted(set(a.cells) | set(b.cells)):
        old, new = _compared(a.cells.get(key)), _compared(b.cells.get(key))
        if old == new:
            continue
        row = allowed.get(key)
        ok = (row is not None and old == ("s", row["original"]) and new is not None
              and new[0] == "s" and any(new[1] in v for v in row["labels"].values()))
        (changed if ok else violations).append(
            {"sheet": key[0], "cell": key[1], "old": old, "new": new})
    return {"status": "pass" if not violations else "fail",
            "labelsChanged": changed, "violations": violations,
            "formulasCompared": sum(1 for v in a.cells.values() if v["kind"] == "f"),
            "numbersCompared": sum(1 for v in a.cells.values() if v["kind"] == "n")}


LABOR_SHEET = "8. 노무비계획"
# X01 '8. 노무비계획': 고용 노동비 일용직 남(11행)·여(12행) 연도별 작업 일수 칸과
# 작업별 산출근거(1년차 26~33행, 2년차 이후 40~47행)의 남자 O·여자 R 일수 칸.
LABOR_SUMMARY_DAYS = {"남자": ("G11", "J11", "M11", "P11", "S11"),
                      "여자": ("G12", "J12", "M12", "P12", "S12")}
LABOR_TASK_BLOCKS = ((range(26, 34), "1년차"), (range(40, 48), "2년차 이후"))
# 위 좌표가 X01 배치인지 확인하는 칸 이름(앞뒤 공백·칸 안 공백 무시).
LABOR_LAYOUT_ANCHORS = (("C11", "일용직"), ("E11", "남"), ("E12", "여"),
                        ("R24", "고용(여자)"), ("R38", "고용(여자)"),
                        ("O25", "일수"), ("R25", "일수"), ("O39", "일수"), ("R39", "일수"))
LABOR_MALE_HEADERS = ("고용(남자,단가동일)", "고용(남자)")


def _labor_blocked(code, reason, evidence):
    # Not a pass: blocked status, warning severity like the check it replaces.
    return {"code": code, "status": "blocked", "severity": "warning", "reason": reason,
            "remedy": "교수님 양식(X01) '8. 노무비계획' 배치인지 확인하고, 엑셀에서 재계산·저장해 "
                      "수식 캐시를 만든 뒤 다시 검사한다",
            "evidence": evidence}


def labor_review(path) -> list[dict]:
    """Daily-labor gender review on a professor-template (X01) workbook.

    The shared read check (``read_template``) runs first: a file that is
    missing or cannot be read is blocked (``LABOR_WORKBOOK_UNREADABLE``).
    A readable workbook not identified as X01 (another major's form, X02,
    unknown) is not in scope and yields no issue.  On an X01 workbook,
    hired days only in the male column are a warning
    (``LABOR_HIRED_MALE_ONLY``).  A missing sheet (``LABOR_SHEET_MISSING``),
    a different labor layout (``LABOR_LAYOUT_UNSUPPORTED``) or a days cell
    that cannot be read as a number — text, or a formula without a cached
    value (``LABOR_DAYS_UNREADABLE``) — is blocked, never a silent pass.
    Only a truly empty, formula-free cell counts as 0 days."""
    path = Path(path)
    try:
        layout = read_template(path).layout
    except TemplateUnreadable as exc:
        return [_labor_blocked("LABOR_WORKBOOK_UNREADABLE",
                               "통합문서를 열거나 읽지 못해 일용직 성별을 확인하지 못함",
                               {"workbook": path.name, "error": str(exc)})]
    if layout != "x01":
        return []
    import openpyxl
    try:
        workbook = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:  # openpyxl raises many parser types
        return [_labor_blocked("LABOR_WORKBOOK_UNREADABLE",
                               "통합문서를 열거나 읽지 못해 일용직 성별을 확인하지 못함",
                               {"workbook": path.name, "error": f"{type(exc).__name__}: {exc}"})]
    try:
        return _labor_review_x01(workbook, formulas)
    finally:
        workbook.close()
        formulas.close()


def _labor_review_x01(workbook, formulas) -> list[dict]:
    if LABOR_SHEET not in workbook.sheetnames:
        return [_labor_blocked("LABOR_SHEET_MISSING",
                               "'8. 노무비계획' 시트를 찾지 못해 일용직 성별을 확인하지 못함",
                               {"sheets": list(workbook.sheetnames)})]
    ws, wf = workbook[LABOR_SHEET], formulas[LABOR_SHEET]
    norm = lambda v: re.sub(r"\s+", "", v) if isinstance(v, str) else v
    mismatched = [{"cell": ref, "expected": text, "found": ws[ref].value}
                  for ref, text in LABOR_LAYOUT_ANCHORS if norm(ws[ref].value) != text]
    mismatched += [{"cell": ref, "expected": list(LABOR_MALE_HEADERS), "found": ws[ref].value}
                   for ref in ("O24", "O38") if norm(ws[ref].value) not in LABOR_MALE_HEADERS]
    if mismatched:
        return [_labor_blocked("LABOR_LAYOUT_UNSUPPORTED",
                               "'8. 노무비계획' 칸 배치가 교수님 양식과 달라 남녀 고용 일수를 읽지 못함",
                               {"mismatched": mismatched})]
    unreadable = []

    def days(ref):
        value = ws[ref].value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return value
        source = wf[ref].value
        if value is None and not (isinstance(source, str) and source.startswith("=")):
            return 0
        if isinstance(value, str) and not value.strip() and source == value:
            return 0
        unreadable.append({"cell": ref, "value": value, "formula": source
                           if isinstance(source, str) and source.startswith("=") else None})
        return 0

    male = sum(days(r) for r in LABOR_SUMMARY_DAYS["남자"])
    female = sum(days(r) for r in LABOR_SUMMARY_DAYS["여자"])
    tasks = []
    for rows, period in LABOR_TASK_BLOCKS:
        for row in rows:
            m, f = days(f"O{row}"), days(f"R{row}")
            if m or f:
                tasks.append({"period": period, "row": row, "task": ws[f"B{row}"].value,
                              "male_days": m, "female_days": f})
    if unreadable:
        return [_labor_blocked("LABOR_DAYS_UNREADABLE",
                               "고용 일수 칸을 숫자로 읽지 못함(글자이거나 수식 캐시 없음): "
                               + ", ".join(u["cell"] for u in unreadable),
                               {"cells": unreadable})]
    task_male = sum(t["male_days"] for t in tasks)
    task_female = sum(t["female_days"] for t in tasks)
    if (male > 0 and female == 0) or (task_male > 0 and task_female == 0):
        return [{
            "code": "LABOR_HIRED_MALE_ONLY", "status": "pending", "severity": "warning",
            "reason": "일용직 고용 일수가 남자 칸에만 있음(남 %s일·여 %s일, 작업별 남 %s일·여 %s일)"
                      % (male, female, task_male, task_female),
            "remedy": "작업별(수확·선별·포장 등)로 남녀 누구를 고용하는지 학생에게 확인해 칸을 나누고, "
                      "근거 자료에 성별 구분이 없으면 작업 관행으로 정한 이유를 각주에 적는다",
            "evidence": {"summary": {"male_days": male, "female_days": female}, "tasks": tasks},
        }]
    return []


def cli(argv: list[str]) -> int:
    p = argparse.ArgumentParser(description="inspect and create a source-bound reusable XLSX template")
    sp = p.add_subparsers(dest="cmd", required=True)
    i = sp.add_parser("inspect")
    i.add_argument("--source", required=True, type=Path)
    i.add_argument("--out-map", required=True, type=Path)
    i.add_argument("--report", type=Path)
    w = sp.add_parser("wage-map")
    w.add_argument("--source", required=True, type=Path)
    w.add_argument("--out-map", required=True, type=Path)
    lm = sp.add_parser("label-map", help="작목 분류별 칸 이름 채우기 맵·값 생성")
    lm.add_argument("--source", required=True, type=Path)
    lm.add_argument("--category", required=True)
    lm.add_argument("--choice", action="append", default=[], metavar="SHEET!CELL=LABEL")
    lm.add_argument("--out-map", required=True, type=Path)
    lm.add_argument("--out-values", required=True, type=Path)
    lc = sp.add_parser("label-check", help="칸 이름 변경 뒤 수식·숫자 불변 확인")
    lc.add_argument("--before", required=True, type=Path)
    lc.add_argument("--after", required=True, type=Path)
    lr = sp.add_parser("labor-review", help="일용직 성별 고용 일수 경고")
    lr.add_argument("--workbook", required=True, type=Path)
    c = sp.add_parser("clear")
    c.add_argument("--source", required=True, type=Path)
    c.add_argument("--map", required=True, type=Path)
    c.add_argument("--out", required=True, type=Path)
    c.add_argument("--receipt", type=Path)
    c.add_argument("--project", default=None, type=Path,
                   help="프로젝트 정본 폴더")
    c.add_argument("--major", default=None, help="명시 전공 ID")
    args = p.parse_args(argv)
    try:
        if args.cmd == "label-map":
            choices = {}
            for item in args.choice:
                key, sep, label = item.partition("=")
                if not sep:
                    raise ValueError(f"--choice needs SHEET!CELL=LABEL: {item!r}")
                choices[key] = label
            fill_map, fill_values = crop_label_map(args.source, args.category, choices)
            for path, data in ((args.out_map, fill_map), (args.out_values, fill_values)):
                path.parent.mkdir(parents=True, exist_ok=True)
                with path.open("x", encoding="utf-8") as stream:
                    json.dump(data, stream, ensure_ascii=False, indent=2)
            print(json.dumps({"status": "mapped", "labels": len(fill_values["values"]),
                              "map": str(args.out_map), "values": str(args.out_values)},
                             ensure_ascii=False))
            return 0
        if args.cmd == "label-check":
            result = crop_label_check(args.before, args.after)
            print(json.dumps(result, ensure_ascii=False))
            return {"pass": 0, "fail": 1}.get(result["status"], 2)
        if args.cmd == "labor-review":
            issues = labor_review(args.workbook)
            blocked = any(i["status"] == "blocked" for i in issues)
            print(json.dumps({"status": "blocked" if blocked else "reviewed", "issues": issues},
                             ensure_ascii=False))
            return 2 if blocked else 0
        if args.cmd == "wage-map":
            result = wage_fill_map(resolve_source(args.source))
            args.out_map.parent.mkdir(parents=True, exist_ok=True)
            with args.out_map.open("x", encoding="utf-8") as stream:
                json.dump(result, stream, ensure_ascii=False, indent=2)
            print(json.dumps({"status": "mapped", "map": str(args.out_map)}, ensure_ascii=False))
        elif args.cmd == "inspect":
            r = inspect_source(args.source, args.out_map, args.report)
            print(json.dumps({"status": "inspected", "source": r["source"], "entries": len(r["entries"]),
                              "sheets": len(r["sheets"]), "map": str(args.out_map),
                              "layout": r["layout"]["variant"],
                              "mapStatus": r["mapStatus"]}, ensure_ascii=False))
        else:
            if args.out.exists():
                raise FileExistsError(f"refusing to overwrite output: {args.out}")
            if args.receipt and args.receipt.exists():
                raise FileExistsError(f"refusing to overwrite receipt: {args.receipt}")
            # Stage both files first.  Publication uses rollback on ordinary
            # exceptions; the receipt is staged before the workbook is made
            # visible, but two separate renames are not crash-atomic.
            staged_out = args.out.with_name(f".{args.out.name}.tmp-{os.getpid()}")
            staged_receipt = (args.receipt.with_name(f".{args.receipt.name}.tmp-{os.getpid()}")
                              if args.receipt else None)
            if staged_out.exists() or (staged_receipt and staged_receipt.exists()):
                raise FileExistsError("staging path already exists")
            published_out = False
            try:
                import gg_major_contract as mc
                context = (mc.output_context(args.project, args.major)
                           if args.project is not None else None)
                r = blank_copy(resolve_source(args.source).resolve(), args.map.resolve(), staged_out.resolve(),
                               context=context)
                r["output"]["path"] = str(args.out.resolve())
                if staged_receipt:
                    staged_receipt.parent.mkdir(parents=True, exist_ok=True)
                    staged_receipt.write_text(json.dumps(r, ensure_ascii=False, indent=2), encoding="utf-8")
                # Reconfirm after the receipt is staged, right before the
                # first final replace (policy B).
                mc.reconfirm_output(
                    mc.OutputAuthorization(**r["majorAuthorization"]), context)
                os.replace(staged_out, args.out)
                published_out = True
                if staged_receipt:
                    os.replace(staged_receipt, args.receipt)
            except Exception:
                for pending in (staged_out, staged_receipt):
                    if pending and pending.exists():
                        pending.unlink()
                if published_out and args.out.exists():
                    args.out.unlink()
                raise
            print(json.dumps({"status": r["output"]["status"], "out": r["output"],
                              "cleared": len(r["cleared"]), "formulaProtected": r["formulaCellsProtected"],
                              "formulaCachesInvalidated": r["formulaCachesInvalidated"]}, ensure_ascii=False))
        return 0
    except Exception as exc:
        import gg_major_contract as mc
        if isinstance(exc, mc.OutputHeldError):
            print(json.dumps({"status": "held", "reason": exc.reason,
                              "detail": str(exc),
                              "guidance": exc.detail.get("guidance")},
                             ensure_ascii=False))
            return 2
        if isinstance(exc, (OSError, ValueError, RuntimeError,
                            zipfile.BadZipFile)):
            print(f"BLOCK: {exc}", file=sys.stderr)
            return 2
        raise


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):
        getattr(_stream, "reconfigure", lambda **_: None)(encoding="utf-8")
    raise SystemExit(cli(sys.argv[1:]))
