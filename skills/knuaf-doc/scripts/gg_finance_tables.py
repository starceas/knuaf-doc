"""특용작물전공 Ⅳ장 재무표 전개기 (D9).

채워진 학과 통합문서(17시트)의 캐시된 셀 값만 읽어 Ⅳ장 표 묶음을
Markdown으로 만든다. 표 지도는
``references/specialty-crops/finance-table-map.json`` — 라벨·앵커로 블록을
찾으므로 학과 양식 버전 사이의 행 좌표 차이를 견딘다.

규칙:
- 특용작물전공 전용. ``major_id``가 ``specialty_crops``가 아니면 거부한다.
- 셀이 비어 있거나 ``#`` 오류이거나 수식 캐시가 없으면 값을 만들지 않고
  ``확인 불가``로 표시하며 경고 목록에 남긴다.
- 표가 통째로 비면 만들지 않는다(빈 표를 수치로 채우지 않음).
- 공식 통계 각주(물가·임금·판매가 상승률)는 입력 JSON으로 받는다. 없으면
  각주를 만들지 않고 ``확인 불가`` 경고만 남긴다.
- 상승률 각주는 항상 참고값이다. r이 워크북 계산에 실제로 쓰였는지는
  이 도구가 증명하지 않으며, 출력 문구도 통합문서에서 따로 확인해야
  한다고 밝힌다(Q-DESIGN-1 처분 A).
- 연도별 비용표는 워크북의 계획연도 집합·개수와 일치해야 발행한다.
- 입력 통합문서는 읽기만 하며 절대 수정하지 않는다(사본 사용 전제).
"""

import argparse
import json
import math
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

import gg_major_contract as mc
import gg_price_assumptions as pa

MAJOR_ID = "specialty_crops"
MISSING = "확인 불가"
FINANCE_PROFILE = "school_17_sheet_v1"

_DEFAULT_MAP = (
    Path(__file__).resolve().parent.parent
    / "references" / "specialty-crops" / "finance-table-map.json"
)

_WS_RUN = re.compile(r"\s{2,}")
_SHEET_KEY = re.compile(r"[\s.]+")


def _disp(text):
    """표시용 텍스트: 2자 이상 공백런·줄바꿈은 제거, 단일 공백은 유지."""
    if text is None:
        return ""
    s = _WS_RUN.sub("", str(text).strip())
    return s.replace("\n", " ").strip()


def _norm(text):
    """앵커 비교용 정규화: 모든 공백 제거."""
    if text is None:
        return ""
    return re.sub(r"\s+", "", str(text))


def _cell_safe(text):
    """마크다운 표 셀 안전화."""
    return str(text).replace("|", "/").replace("\n", " ").strip()


def _col(letter):
    n = 0
    for ch in letter.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _fmt_number(v):
    """소수는 2자리까지, 정수는 천 단위 구분."""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        if v == int(v):
            return f"{int(v):,}"
        s = f"{v:,.2f}".rstrip("0").rstrip(".")
        return s
    return _disp(v)


def _fmt_int_comma(v):
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return f"{round(v):,}"
    return _disp(v)


def _num1(v):
    """소수 1자리 반올림; 정수면 소수점을 뗀다."""
    n = round(v, 1)
    return str(int(n)) if n == int(n) else str(n)


def _fmt_ratio(v):
    """셀이 비율값(0.05=5%)인 행 → 백분율 숫자로 변환해 표시."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _num1(v * 100)
    return _disp(v)


def _fmt_percent_value(v):
    """셀이 이미 백분율값(5=5%)인 행 → 숫자 그대로 표시."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _num1(v)
    return _disp(v)


def _fmt_ratio_cell(v):
    """비율값 저장 셀(연이자율 등) → ×100 뒤 % 기호."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _num1(v * 100) + "%"
    return _disp(v)


def _fmt_percent_value_cell(v):
    """백분율값 저장 셀 → 숫자 그대로 + % 기호."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _num1(v) + "%"
    return _disp(v)


def _fmt_text(v):
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return _disp(v)


_FORMATS = {
    "number": _fmt_number,
    "int_comma": _fmt_int_comma,
    "ratio": _fmt_ratio,
    "percent_value": _fmt_percent_value,
    "ratio_cell": _fmt_ratio_cell,
    "percent_value_cell": _fmt_percent_value_cell,
    "text": _fmt_text,
}


_XLS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_XLS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CELL_REF = re.compile(r"^([A-Z]+)(\d+)$")
_RANGE_REF = re.compile(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$")


def _sheet_files(path):
    """시트 이름 → xl/worksheets/*.xml 경로."""
    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        target_by_rid = {
            rel.get("Id"): rel.get("Target") for rel in rels
        }
        out = {}
        for sh in wb.iter(_XLS_MAIN + "sheet"):
            rid = sh.get(_XLS_REL + "id")
            target = target_by_rid.get(rid)
            if not target:
                continue
            target = target.lstrip("/")
            if not target.startswith("xl/"):
                target = "xl/" + target
            out[sh.get("name")] = target
        return out


def _sheet_meta(path, sheet_file):
    """시트 XML에서 수식 셀 참조(공유 수식 따르는 셀 포함)와 병합 범위를 읽는다."""
    formula_refs = set()
    merges = []
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read(sheet_file))
    for mc in root.iter(_XLS_MAIN + "mergeCell"):
        ref = mc.get("ref")
        if ref:
            merges.append(ref)
    for cell in root.iter(_XLS_MAIN + "c"):
        ref = cell.get("r")
        if ref is None:
            continue
        if cell.find(_XLS_MAIN + "f") is not None:
            formula_refs.add(ref)
    return formula_refs, merges


def _ref_to_rc(ref):
    m = _CELL_REF.match(ref)
    if not m:
        return None
    return int(m.group(2)), _col(m.group(1))


def _range_cells(ref):
    m = _RANGE_REF.match(ref)
    if not m:
        return []
    r1, r2 = int(m.group(2)), int(m.group(4))
    c1, c2 = _col(m.group(1)), _col(m.group(3))
    return [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


class _SheetView:
    """시트 한 장: 캐시 값 + 수식 존재 여부 + 병합 채우기.

    read_only 모드로 값만 읽고(큰 파일·손상 스타일 내성), 수식·병합은
    시트 XML에서 직접 읽는다."""

    def __init__(self, ws_values, formula_refs, merges):
        self.cells = {}
        self.max_row = 0
        self.max_col = 0
        for row in ws_values.iter_rows():
            for cell in row:
                if cell.value is not None:
                    self.cells[(cell.row, cell.column)] = cell.value
                    if cell.row > self.max_row:
                        self.max_row = cell.row
                    if cell.column > getattr(self, "max_col", 0):
                        self.max_col = cell.column
        self.formulas = {
            rc for rc in (_ref_to_rc(f) for f in formula_refs) if rc
        }
        self.covered = {}
        for ref in merges:
            cells = _range_cells(ref)
            if not cells:
                continue
            top = cells[0]
            for rc in cells[1:]:
                self.covered[rc] = top

    def _coord(self, r, c):
        return self.covered.get((r, c), (r, c))

    def value(self, r, c):
        """캐시 값. 세로 병합(같은 열의 top-left)만 채운다 — 가로 병합은
        라벨 텍스트를 값 열로 복제하므로 데이터 판독에 쓰지 않는다."""
        if r < 1 or c < 1:
            return None
        top = self.covered.get((r, c))
        if top is None:
            return self.cells.get((r, c))
        if top[1] != c:
            return None
        return self.cells.get(top)

    def value_wide(self, r, c):
        """병합 방향 무관 채움 — 머리글(그룹 라벨) 조합 전용."""
        if r < 1 or c < 1:
            return None
        return self.cells.get(self._coord(r, c))

    def raw(self, r, c):
        if r < 1 or c < 1:
            return None
        return self.cells.get((r, c))

    def status(self, r, c):
        """('ok', value) | ('empty'|'error'|'uncached_formula', None)."""
        v = self.value(r, c)
        if isinstance(v, str) and v.startswith("#"):
            return ("error", None)
        if isinstance(v, float) and not math.isfinite(v):
            return ("error", None)
        if v is None:
            if self._coord(r, c) in self.formulas or (r, c) in self.formulas:
                return ("uncached_formula", None)
            return ("empty", None)
        return ("ok", v)


def _find_sheet(workbook, name):
    """공백·마침표를 지운 이름으로 시트를 찾는다."""
    want = _SHEET_KEY.sub("", name)
    for title in workbook.sheetnames:
        if _SHEET_KEY.sub("", title) == want:
            return workbook[title]
    return None


def _find_anchor(view, extract, start_row=1):
    col = _col(extract.get("anchor_col", "B"))
    regex = extract.get("anchor_regex")
    regex = re.compile(regex) if regex else None
    anchor = _norm(extract.get("anchor"))
    search_after = extract.get("search_after")
    row = start_row
    if search_after:
        mark = _norm(search_after)
        while row <= view.max_row:
            if mark in _norm(view.value(row, col)):
                break
            row += 1
        if row > view.max_row:
            return None
    for r in range(row, view.max_row + 1):
        text = _norm(view.value(r, col))
        if not text:
            continue
        if regex is not None and regex.match(text):
            return r
        if regex is None and anchor and anchor in text:
            return r
    return None


def _header_label_parts(view, col_letter, header_rows):
    parts = []
    for r in header_rows:
        t = _disp(view.value_wide(r, _col(col_letter)))
        if t and (not parts or parts[-1] != t):
            parts.append(t)
    return parts


def _resolve_headers(view, cols, header_rows, header_map):
    """열 머리글: header_map 우선, 없으면 마지막 머리글(중복이면 상위+하위)."""
    header_map = header_map or {}
    parts_by_col = {c: _header_label_parts(view, c, header_rows) for c in cols}
    leafs = {c: (p[-1] if p else "") for c, p in parts_by_col.items()}
    counts = {}
    for c in cols:
        counts[leafs[c]] = counts.get(leafs[c], 0) + 1
    out = []
    for col_letter in cols:
        if col_letter in header_map:
            out.append(header_map[col_letter])
            continue
        leaf = leafs[col_letter]
        if not leaf:
            out.append("열" + col_letter)
            continue
        if counts[leaf] > 1:
            parent_parts = parts_by_col[col_letter][:-1]
            parent = parent_parts[-1] if parent_parts else ""
            out.append(parent + " " + leaf if parent
                       else leaf + "(" + col_letter + ")")
        else:
            out.append(leaf)
    return out


def _render_cell(view, r, col_idx, col_letter, role, fmt, warnings,
                 table_id, sheet_title, label):
    st, v = view.status(r, col_idx)
    addr = "'%s'!%s%d" % (sheet_title, col_letter, r)
    if st != "ok":
        if st == "empty" and role != "value":
            return ""
        warnings.append({
            "table": table_id,
            "cell": addr,
            "row_label": label,
            "reason": st,
        })
        return MISSING
    if role == "structural":
        return _disp(v)
    return _FORMATS.get(fmt, _fmt_number)(v)


def _row_label(view, r, cols):
    for c in cols:
        t = _disp(view.value(r, c))
        if t:
            return t
    return ""


def _extract_block_rows(view, extract, anchor_row, table_id, sheet_title,
                        warnings, notes, col_idxs, extra_stop_re=None):
    """앵커 아래 데이터 행을 모은다. (rows, last_scanned_row) 반환."""
    anchor_col = _col(extract.get("anchor_col", "B"))
    data_start = anchor_row + extract["data_offset"]
    stop_before = [_norm(s) for s in extract.get("stop_before", [])]
    stop_re = extract.get("stop_before_regex")
    stop_re = re.compile(stop_re) if stop_re else None
    stop_after = [_norm(s) for s in extract.get("stop_after", [])]
    skip_re = extract.get("skip_row_regex")
    skip_re = re.compile(skip_re) if skip_re else None
    keep_rows = {_norm(s) for s in extract.get("keep_rows", [])}
    value_cols = set(extract.get("value_cols", []))
    note_cols = set(extract.get("note_cols", []))
    col_letters = extract["cols"]
    col_fmt = extract.get("col_formats", {})
    # 행별 단위 계약: 키(정규화 라벨 부분문자열) → 저장 단위
    # (percent_value=셀이 이미 백분율값, ratio=셀이 비율). 긴 키 우선.
    row_units = extract.get("row_units", {})
    unit_keys = sorted(row_units, key=len, reverse=True)
    # 라벨이 있는데 값 열이 전부 빈 행: 지도에 등록된 미보유/미사용
    # 행만 생략하고 나머지는 '확인 불가'로 보존한다.
    optional_blank = [_norm(s) for s in extract.get("optional_when_blank", [])]

    rows = []
    r = data_start
    r_stop = view.max_row
    seen_note = False
    while r <= view.max_row:
        probe = _norm(view.value(r, anchor_col))
        if any(m and m in probe for m in stop_before):
            break
        if stop_re and stop_re.match(probe):
            break
        if extra_stop_re and extra_stop_re.match(probe):
            break
        hits_after = any(m and m in probe for m in stop_after)

        statuses = {}
        for letter in col_letters:
            ci = col_idxs[letter]
            statuses[letter] = view.status(r, ci)
        label = _row_label(view, r, [col_idxs[l] for l in col_letters])
        first_text = ""
        is_skip = False
        for letter in col_letters:
            st, v = statuses[letter]
            if st == "ok" and _disp(v):
                first_text = _disp(v)
                if skip_re and skip_re.match(_norm(v)):
                    is_skip = True
                break
        if is_skip:
            if hits_after:
                r_stop = r
                break
            r += 1
            continue
        if first_text.startswith("*"):
            notes.append(first_text)
            seen_note = True
        else:
            first_label = _norm(view.value(r, _col(col_letters[0])))
            structural = first_label in keep_rows
            has_value = any(
                statuses[l][0] in ("ok", "error", "uncached_formula")
                and (statuses[l][0] != "ok" or _disp(statuses[l][1]) != "")
                for l in value_cols
            )
            if structural:
                rendered = {}
                for letter in col_letters:
                    st, v = statuses[letter]
                    rendered[letter] = _disp(v) if st == "ok" else ""
                rows.append((r, rendered))
            elif has_value:
                norm_label = _norm(label)
                unit = next(
                    (row_units[k] for k in unit_keys
                     if k and k in norm_label), None)
                rendered = {}
                for letter in col_letters:
                    role = (
                        "value" if letter in value_cols
                        else ("note" if letter in note_cols else "label")
                    )
                    fmt = col_fmt.get(letter, "number")
                    if unit and letter in value_cols:
                        fmt = unit
                    rendered[letter] = _render_cell(
                        view, r, col_idxs[letter], letter, role, fmt,
                        warnings, table_id, sheet_title, label)
                rows.append((r, rendered))
            elif label:
                # 라벨 열의 각 셀 텍스트로 선택적 생략을 판정한다 — 병합된
                # 상위 라벨("자가 노동비")만 보면 "가족" 같은 행을 못 잡는다.
                label_cells = [
                    _norm(view.value(r, col_idxs[l]))
                    for l in col_letters if l not in value_cols]
                if any(k and any(k in t for t in label_cells)
                       for k in optional_blank):
                    pass  # 지도가 명시한 미보유/미사용 행만 생략
                elif seen_note:
                    notes.append("* " + label)  # '*' 주석의 이어지는 줄
                else:
                    rendered = {}
                    for letter in col_letters:
                        role = (
                            "value" if letter in value_cols
                            else ("note" if letter in note_cols
                                  else "label"))
                        rendered[letter] = _render_cell(
                            view, r, col_idxs[letter], letter, role,
                            col_fmt.get(letter, "number"), warnings,
                            table_id, sheet_title, label)
                    rows.append((r, rendered))
        if hits_after:
            r_stop = r
            break
        r += 1
    else:
        r_stop = view.max_row
    return rows, r_stop


def _trailing_notes(view, r, cols, notes):
    """데이터 종료 직후 '*' 주석 줄을 모은다."""
    col_idxs = [_col(c) for c in cols]
    while r <= view.max_row:
        texts = []
        for c in col_idxs:
            v = view.value(r, c)
            if v is not None and _disp(v):
                texts.append(_disp(v))
        if not texts:
            r += 1
            continue
        if texts[0].startswith("*"):
            notes.append(texts[0])
            r += 1
            continue
        break


def _blank_repeat(rows, cols, label_cols):
    """병합 채움으로 반복된 라벨은 첫 행만 표시한다."""
    prev = {}
    out = []
    for r, rendered in rows:
        new = dict(rendered)
        for letter in label_cols:
            if letter in prev and prev[letter] == new.get(letter, ""):
                new[letter] = ""
            else:
                prev[letter] = new.get(letter, "")
        out.append((r, new))
    return out


def _emit_table(view, spec, extract, anchor_row, table_id, sheet_title,
                warnings, notes, extra_stop_re=None):
    """블록 하나 → (headers, rows, notes) 또는 None(생략)."""
    col_letters = extract["cols"]
    col_idxs = {l: _col(l) for l in col_letters}
    header_rows = [anchor_row + o for o in extract["header_offsets"]]
    headers = _resolve_headers(
        view, col_letters, header_rows, extract.get("header_map"))
    rows, r_stop = _extract_block_rows(
        view, extract, anchor_row, table_id, sheet_title, warnings, notes,
        col_idxs, extra_stop_re=extra_stop_re)
    if not rows:
        return None
    if extract.get("trailing_notes"):
        _trailing_notes(view, r_stop + 1, col_letters, notes)
    label_cols = [l for l in col_letters
                  if l not in set(extract.get("value_cols", []))
                  and l not in set(extract.get("note_cols", []))]
    if extract.get("blank_repeat_labels"):
        rows = _blank_repeat(rows, col_letters, label_cols)
    return headers, rows, col_letters


def _emit_year_grid(view, spec, extract, sheet_title, warnings):
    """연도 열 그룹(연도별 일당/일수/금액)에서 부속 열 하나를 뽑는다."""
    anchor_row = _find_anchor(view, extract)
    if anchor_row is None:
        return None, "anchor_missing"
    year_row = anchor_row + extract["year_row_offset"]
    sub_row = anchor_row + extract["sub_row_offset"]
    pick = extract["pick"]
    want_sub = _norm(pick["sub"])

    # 연도 그룹: 연도 행에서 raw(병합 전) 값이 연도 패턴인 열만 그룹 시작.
    groups = []
    c = 1
    while c <= view.max_col:
        v = view.raw(year_row, c)
        if isinstance(v, str) and re.match(r"^\d{4}년", _disp(v)):
            groups.append((c, _disp(v)))
        elif isinstance(v, (int, float)) and 1900 <= v <= 2100:
            groups.append((c, "%d년" % int(v)))
        c += 1
    if not groups:
        return None, "year_headers_missing"
    picked = []
    for i, (start, year) in enumerate(groups):
        end = groups[i + 1][0] if i + 1 < len(groups) else start + 3
        sub_col = None
        for c in range(start, end):
            if _norm(view.value(sub_row, c)) == want_sub:
                sub_col = c
                break
        if sub_col is None:
            warnings.append({
                "table": spec.get("id"), "cell": "%s열(%s)" % (year, pick["sub"]),
                "row_label": "", "reason": "sub_column_missing"})
            continue
        picked.append((sub_col, "%s %s" % (year, pick["sub"])))

    label_letters = extract["label_cols"]
    col_letters = label_letters + [str(i) for i in range(len(picked))]
    stop_re = extract.get("stop_before_regex")
    stop_re = re.compile(stop_re) if stop_re else None
    stop_before = [_norm(s) for s in extract.get("stop_before", [])]
    anchor_col = _col(extract.get("anchor_col", "B"))
    fmt = pick.get("format", "number")

    optional_blank = [_norm(s) for s in extract.get("optional_when_blank", [])]
    data_start = anchor_row + extract["data_offset"]
    notes = []
    rows = []
    r = data_start
    seen_note = False
    while r <= view.max_row:
        probe = _norm(view.value(r, anchor_col))
        if stop_re and stop_re.match(probe):
            break
        if any(m and m in probe for m in stop_before):
            break
        statuses = {}
        for letter in label_letters:
            statuses[letter] = view.status(r, _col(letter))
        for i, (pc, _) in enumerate(picked):
            statuses[str(i)] = view.status(r, pc)
        # 주석 행
        first_text = ""
        for letter in label_letters + [str(i) for i in range(len(picked))]:
            st, v = statuses[letter]
            if st == "ok" and _disp(v):
                first_text = _disp(v)
                break
        if first_text.startswith("*"):
            notes.append(first_text)
            seen_note = True
            r += 1
            continue
        label = _row_label(view, r, [_col(l) for l in label_letters])
        has_value = any(
            statuses[str(i)][0] in ("ok", "error", "uncached_formula")
            and (statuses[str(i)][0] != "ok" or _disp(statuses[str(i)][1]) != "")
            for i in range(len(picked))
        )
        if has_value:
            rendered = {}
            for letter in label_letters:
                st, v = statuses[letter]
                rendered[letter] = _disp(v) if st == "ok" else ""
            for i, (pc, _) in enumerate(picked):
                st, v = statuses[str(i)]
                if st != "ok":
                    warnings.append({
                        "table": spec.get("id"),
                        "cell": "'%s'!%s%d" % (sheet_title,
                                             _col_letter(pc), r),
                        "row_label": label, "reason": st})
                    rendered[str(i)] = MISSING
                else:
                    rendered[str(i)] = _FORMATS.get(fmt, _fmt_number)(v)
            rows.append((r, rendered))
        elif label:
            label_cells = [
                _norm(view.value(r, _col(l))) for l in label_letters]
            if any(k and any(k in t for t in label_cells)
                   for k in optional_blank):
                pass  # 지도가 명시한 미보유/미사용 행만 생략
            elif seen_note:
                notes.append("* " + label)  # '*' 주석의 이어지는 줄
            else:
                # 라벨은 있는데 값이 전부 빈 행은 '확인 불가'로 보존
                rendered = {}
                for letter in label_letters:
                    st, v = statuses[letter]
                    rendered[letter] = _disp(v) if st == "ok" else ""
                for i, (pc, _) in enumerate(picked):
                    st, _ = statuses[str(i)]
                    warnings.append({
                        "table": spec.get("id"),
                        "cell": "'%s'!%s%d" % (sheet_title,
                                             _col_letter(pc), r),
                        "row_label": label, "reason": st})
                    rendered[str(i)] = MISSING
                rows.append((r, rendered))
        r += 1
    if not rows:
        return None, "empty"
    headers = [_disp(view.value(sub_row, _col(l))) or l for l in label_letters]
    headers = [extract.get("label_header_map", {}).get(l, headers[i])
               for i, l in enumerate(label_letters)]
    headers += [h for _, h in picked]
    if extract.get("blank_repeat_labels", True):
        rows = _blank_repeat(rows, col_letters, label_letters)
    return (headers, rows, col_letters), notes


def _col_letter(idx):
    s = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        s = chr(65 + rem) + s
    return s


# 각주 구조화 입력의 허용 필드 — 이 밖의 키(자유 문구 note·label·
# source_name·locator 등)는 정책 우회 통로가 될 수 있어 거부한다.
# 계약: {source_id, base_year, observation?{start_year,end_year},
#        rate?(검산용), target_dimensions?, unit?(등록 단위 대조),
#        application_base_year?(구 계약 필드 — 받아도 출력·검증에
#        쓰지 않는다. 각주는 r이 실제로 쓰였는지 주장하지 않으므로)}
#   또는 {status:"not_applied", reason}.
_FOOTNOTE_KEYS = frozenset({
    "status", "reason", "source_id", "base_year", "observation",
    "rate", "application_base_year", "target_dimensions", "unit",
})


def _same_rate(a, b):
    """D8 수치 직렬화 회귀가 공유하는 유한 수치 비교(허용오차 1e-12)."""
    return (type(a) in (int, float) and type(b) in (int, float)
            and math.isfinite(a) and math.isfinite(b)
            and math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12))


def _resolve_footnote(kind, cfg, info, registry):
    """구조화 각주 입력 → '* ...' 줄. 실패 시 (None, 사유).

    입력은 가격 가정 계약과 같다: source_id·base_year·observation·
    검산용 rate. 출처 해석·역할·기준연도·단위·차원·관측·CAGR 검증은
    모두 gg_price_assumptions의 공개 함수로 하고, 출력하는 r은
    항상 관측 수열의 끝점 CAGR 계산값이다. 출력 각주는 참고값이다 —
    r이 워크북 계산에 쓰였는지는 이 도구가 증명하지 않는다."""
    label = cfg.get("label") or kind
    role = cfg.get("role")
    origin = "footnote.%s" % kind
    if not isinstance(info, dict):
        return None, "구조화 객체 입력만 허용(자유 문자열 각주 제거)"
    extra = sorted(set(info) - _FOOTNOTE_KEYS)
    if extra:
        return None, "허용되지 않은 각주 필드: " + ", ".join(extra)
    status = info.get("status", "applied")
    if status == "not_applied":
        reason = info.get("reason")
        if not isinstance(reason, str) or not reason.strip():
            return None, "미적용 선택은 사유(reason) 필요"
        return "* %s 미적용(명시 선택): %s" % (label, reason.strip()), None
    if status != "applied":
        return None, "status는 applied|not_applied"
    sid = info.get("source_id")
    if not isinstance(sid, str) or not sid.strip():
        return None, "출처 id(source_id) 필요"
    try:
        # id는 정규화하지 않는다 — 대소문자·공백 변형은 미등록 id로 거부.
        src = pa.resolve_series(sid, registry)
        pa.check_role(src, role, origin=origin)
        if "base_year" not in info:
            raise ValueError(
                "%s: 기준연도(base_year) 필드 필요(명목 수열은 null)" % origin)
        base_year = pa.check_base_year(info["base_year"], src, origin=origin)
        if "unit" in info:
            pa.check_unit(src, info["unit"], origin=origin)
        if cfg.get("expected_unit"):
            pa.check_unit(src, cfg["expected_unit"], origin=origin)
        pa.check_dimensions(src, info.get("target_dimensions"),
                            required=("crop",) if role == "sales" else (), origin=origin)
        res = pa.resolve_series_rate(
            src, info.get("observation"), info.get("rate"), origin=origin)
    except ValueError as e:
        return None, str(e)
    r = res["rate"]
    obs = res["observation"]
    interval = res["interval_years"]
    basis = ("지수 기준연도 %d" % base_year) if base_year is not None \
        else "지수 기준연도 해당 없음"
    parts = ["출처: %s %s %s [%s]" % (src.get("agency") or "",
             src.get("statistic_name") or "", src.get("item_name") or "", sid),
             basis, "관측기간 %d~%d년" % (obs["start_year"], obs["end_year"]),
             "산식 r=(끝값/시작값)^(1/%d)-1" % interval]
    line = ("* %s r=%.4f%%는 참고값이다 (%s). 이 값이 위 표의 계산에 "
            "쓰였는지는 통합문서에서 따로 확인해야 한다."
            % (label, r * 100, ", ".join(parts)))
    return line, None


def _footnote_lines(table, footnotes, warnings, registry, table_map):
    """각주 입력 → '* ...' 줄 목록. 입력 없으면 확인 불가 경고만.

    각주는 구조화 객체(가격 가정 형태)만 받는다. 자유 문자열과
    임의 문구 필드는 거부하고, r은 레지스트리 관측 수열로 다시 계산해
    출력한다. 출력 각주는 항상 참고값이다."""
    kinds = table.get("footnotes", [])
    if not kinds:
        return []
    footnotes = footnotes or {}
    kind_cfg = table_map.get("footnote_kinds") or {}
    lines = []
    for kind in kinds:
        label = (kind_cfg.get(kind) or {}).get("label") or kind
        info = footnotes.get(kind)
        if info is None:
            warnings.append({
                "table": table.get("id"),
                "cell": None,
                "row_label": "",
                "reason": "footnote_missing",
                "detail": "%s 각주 입력 없음(%s)" % (kind, MISSING),
            })
            continue
        if registry is None:
            line, err = None, "가격 레지스트리를 읽을 수 없음"
        else:
            line, err = _resolve_footnote(
                kind, kind_cfg.get(kind) or {}, info, registry)
        if line is None:
            warnings.append({
                "table": table.get("id"),
                "cell": None,
                "row_label": "",
                "reason": "footnote_invalid",
                "detail": "%s 각주 거부: %s" % (kind, err),
            })
            lines.append("* %s %s" % (label, MISSING))
            continue
        lines.append(line)
    return lines


def _eun_neun(text):
    """조사 은/는 선택: 받침 있으면 '은', 없거나 비한글 끝이면 '는'."""
    s = str(text).strip()
    if s and "가" <= s[-1] <= "힣" and (ord(s[-1]) - 0xAC00) % 28:
        return "은"
    return "는"


def _render_markdown_table(headers, rows, col_letters):
    out = [
        "| " + " | ".join(_cell_safe(h) for h in headers) + " |",
        "|" + "---|" * len(headers),
    ]
    for _, cells in rows:
        out.append("| " + " | ".join(
            _cell_safe(cells.get(l, "")) for l in col_letters) + " |")
    return out


def load_table_map(map_path=None):
    path = Path(map_path) if map_path else _DEFAULT_MAP
    return json.loads(path.read_text(encoding="utf-8"))


def _check_layout(table_map, view_for):
    """학과 템플릿 지문 검사 → (layout_id, 실패 signature 목록).

    지도에 ``layout``이 없으면 (None, []). signature 하나라도
    시트 부재 또는 pattern 미일치면 미지원 레이아웃이다."""
    layout = table_map.get("layout")
    if not layout:
        return None, []
    failed = []
    for sig in layout.get("signatures", []):
        view, _title = view_for(sig["sheet"])
        col = _col(sig.get("col", "B"))
        max_row = int(sig.get("max_row", 15))
        pattern = re.compile(sig["pattern"])
        ok = False
        if view is not None:
            for r in range(1, min(max_row, view.max_row) + 1):
                v = view.value(r, col)
                if v is not None and pattern.search(_norm(v)):
                    ok = True
                    break
        if not ok:
            failed.append({
                "sheet": sig["sheet"], "pattern": sig["pattern"],
                "desc": sig.get("desc"),
                "detail": "시트 '%s' %s열 %d행 이내에 '%s' 없음"
                          % (sig["sheet"], sig.get("col", "B"),
                             max_row, sig.get("desc") or sig["pattern"])})
    return layout.get("id"), failed


def _plan_years(table_map, view_for):
    """Use workbook labels; never infer missing years from the first anchor.

    W-023 leaves some school sheets one year apart. Annual cost groups use
    the materials sheet's confirmed labels; the production header independently
    confirms their count, without changing either sheet's years.
    """
    cfg = table_map.get("plan_years")
    if not cfg:
        return None, None  # Custom maps without annual-plan declarations.

    def headers(spec):
        view, _ = view_for(spec["sheet"])
        if view is None:
            raise ValueError("계획연도 시트 없음")
        years = []
        for ref in spec["header_cells"]:
            m = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", ref)
            value = view.value(int(m[2]), _col(m[1]))
            label = re.fullmatch(r"(\d{4})년?", _norm(value))
            if not label:
                raise ValueError("계획연도 머리글 확인 불가: %s" % ref)
            years.append(int(label[1]))
        if len(set(years)) != len(years) or years != sorted(years):
            raise ValueError("계획연도 머리글 중복·역순")
        return years

    try:
        if "header_cells" in cfg:
            return headers(cfg), None
        count_years = headers(cfg["count_from"])
        view, _ = view_for(cfg["sheet"])
        if view is None:
            raise ValueError("계획연도 시트 없음")
        pattern = re.compile(cfg["anchor_regex"])
        years = [int(re.search(r"(\d{4})년", _norm(view.value(r, _col(cfg["anchor_col"]))))[1])
                 for r in range(1, view.max_row + 1)
                 if pattern.match(_norm(view.value(r, _col(cfg["anchor_col"]))))]
        if (len(years) != len(count_years) or len(set(years)) != len(years)
                or years != sorted(years)):
            raise ValueError("자재 계획연도 개수·중복·순서 불일치")
        return years, None
    except (ValueError, KeyError, TypeError) as e:
        return [], str(e)


def generate(workbook_path, *, major_id, footnotes=None, map_path=None,
             registry=None):
    """통합문서 경로 → {'markdown', 'emitted', 'skipped', 'warnings',
    'missing_required', 'layout'}.

    ``major_id``는 명시 필수이며 ``specialty_crops``가 아니면 거부한다.
    통합문서는 읽기만 한다(쓰기·재계산 없음). 지도가 ``layout`` 지문을
    선언하면 학과 제공 템플릿 레이아웃과 다를 때
    ``status='unsupported_layout'``으로 조기 반환한다. ``registry``는
    각주 검증용 가격 수열 레지스트리 주입 지점(생략 시 번들 JSON).
    상승률 각주는 항상 참고값으로 출력한다."""
    if major_id != MAJOR_ID:
        raise ValueError("unsupported_major: %r" % (major_id,))
    table_map = load_table_map(map_path)
    path = Path(workbook_path)
    wb_values = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet_files = _sheet_files(path)

    sections = [map_section for map_section in table_map["sections"]]
    md = [table_map.get("chapter", "Ⅳ. 재무계획"), ""]
    emitted = []
    skipped = []
    warnings = []
    table_no = 0

    views = {}

    def view_for(sheet_name):
        if sheet_name not in views:
            ws_v = _find_sheet(wb_values, sheet_name)
            view = None
            if ws_v is not None:
                sheet_file = sheet_files.get(ws_v.title)
                formulas, merges = (
                    _sheet_meta(path, sheet_file) if sheet_file
                    else (set(), []))
                view = _SheetView(ws_v, formulas, merges)
            views[sheet_name] = (
                view, ws_v.title if ws_v is not None else sheet_name)
        return views[sheet_name]

    layout_id, failed_sigs = _check_layout(table_map, view_for)
    if failed_sigs:
        wb_values.close()
        return {
            "status": "unsupported_layout",
            "layout": layout_id,
            "failed_signatures": failed_sigs,
            "markdown": None,
            "emitted": [],
            "skipped": [],
            "missing_required": [],
            "warnings": [{
                "table": None, "cell": None, "row_label": "",
                "reason": "unsupported_layout",
                "detail": "학과 제공 통합문서 레이아웃 지문 불일치",
            }],
        }

    registry_err = None
    if footnotes is not None and registry is None:
        try:
            registry = pa.load_registry()
        except (OSError, ValueError) as e:
            registry_err = str(e)
    if registry_err:
        warnings.append({
            "table": None, "cell": None, "row_label": "",
            "reason": "registry_unavailable",
            "detail": registry_err})

    plan_years, plan_error = _plan_years(table_map, view_for)
    invalid_year_groups = []

    def emit_table_block(tid, headers, rows, col_letters, title, unit,
                         sheet_title, notes, foot_lines):
        nonlocal table_no
        table_no += 1
        caption = "표 %d. %s(단위: %s)" % (table_no, title, unit)
        # 본문 참조 문장(object_reference) → 캡션 → 표 → 출처 → 주석 순.
        # 출처는 반드시 표 바로 뒤에 둔다: 주석이 사이에 끼면 인접 표가
        # 출처 줄을 가로채 gg_document.check의 object_credit이 실패한다.
        md.append(table_map["lead_template"].format(
            title=title, josa=_eun_neun(title), no=table_no))
        md.append(caption)
        md.extend(_render_markdown_table(headers, rows, col_letters))
        md.append(
            table_map["source_credit_template"].format(sheet=sheet_title))
        # 원문 상승률 주석은 출처·관측 검증 없이 복사하지 않는다.
        # 해당 상승률은 구조화 입력을 검증한 참고값 각주로만 안내한다.
        rate_note_pattern = table_map.get("workbook_rate_note_pattern")
        md.extend(note for note in notes if not rate_note_pattern
                  or not re.search(rate_note_pattern, _norm(note)))
        md.extend(foot_lines)
        md.append("")
        emitted.append({"table": tid, "title": title,
                        "rows": len(rows), "number": table_no})

    for section in sections:
        md.append(section["heading"])
        md.append("")
        current_group = None
        for table in section["tables"]:
            group = table.get("group")
            if group and group != current_group:
                md.append(group)
                md.append("")
                current_group = group
            view, sheet_title = view_for(table["sheet"])
            if view is None:
                skipped.append({"table": table["id"],
                                "reason": "sheet_missing",
                                "sheet": table["sheet"]})
                continue
            extract = table["extract"]
            foot_lines = _footnote_lines(table, footnotes, warnings,
                                         registry, table_map)
            kind = extract["kind"]
            if kind == "year_blocks":
                regex = re.compile(extract["anchor_regex"])
                anchor_col = _col(extract.get("anchor_col", "B"))
                anchor_rows = [
                    r for r in range(1, view.max_row + 1)
                    if regex.match(_norm(view.value(r, anchor_col)))
                ]
                if not anchor_rows:
                    skipped.append({"table": table["id"],
                                    "reason": "anchor_missing",
                                    "sheet": sheet_title})
                    continue
                anchor_years = [int(re.search(r"(\d{4})년", _norm(view.value(r, anchor_col)))[1])
                                for r in anchor_rows]
                if plan_years is not None and (plan_error or anchor_years != plan_years
                        or len(set(anchor_years)) != len(anchor_years)):
                    invalid_year_groups.append({"table": table["id"],
                        "sheet": sheet_title, "reason": "year_blocks_incomplete",
                        "expected_years": plan_years, "actual_years": anchor_years,
                        "detail": plan_error or "계획연도 블록 누락·중복·순서 불일치"})
                    continue
                year_re = re.compile(r"(\d{4})년")
                emitted_sub = 0
                for ar in anchor_rows:
                    m = year_re.search(_norm(view.value(ar, anchor_col)))
                    year = m.group(0) if m else "연도미상"
                    sub_id = "%s@%s" % (table["id"], year)
                    notes = []
                    result = _emit_table(
                        view, table, extract, ar, sub_id, sheet_title,
                        warnings, notes, extra_stop_re=regex)
                    if result is None:
                        skipped.append({"table": sub_id,
                                        "reason": "empty",
                                        "sheet": sheet_title})
                        continue
                    headers, rows, col_letters = result
                    title = table["title_template"].format(year=year)
                    emit_table_block(sub_id, headers, rows, col_letters,
                                     title, table["unit"], sheet_title,
                                     notes, foot_lines)
                    emitted_sub += 1
                if not emitted_sub:
                    skipped.append({"table": table["id"],
                                    "reason": "empty",
                                    "sheet": sheet_title})
                if plan_years is not None and emitted_sub != len(plan_years):
                    invalid_year_groups.append({"table": table["id"],
                        "sheet": sheet_title, "reason": "year_blocks_incomplete",
                        "expected_years": plan_years, "actual_count": emitted_sub,
                        "detail": "계획연도 블록 일부가 비어 있음"})
            elif kind == "year_grid":
                result, extra = _emit_year_grid(
                    view, table, extract, sheet_title, warnings)
                if result is None:
                    skipped.append({"table": table["id"], "reason": extra,
                                    "sheet": sheet_title})
                    continue
                headers, rows, col_letters = result
                emit_table_block(table["id"], headers, rows, col_letters,
                                 table["title"], table["unit"], sheet_title,
                                 extra, foot_lines)
            else:
                anchor_row = _find_anchor(view, extract)
                if anchor_row is None:
                    skipped.append({"table": table["id"],
                                    "reason": "anchor_missing",
                                    "sheet": sheet_title})
                    continue
                notes = []
                result = _emit_table(
                    view, table, extract, anchor_row, table["id"],
                    sheet_title, warnings, notes)
                if result is None:
                    skipped.append({"table": table["id"], "reason": "empty",
                                    "sheet": sheet_title})
                    continue
                headers, rows, col_letters = result
                emit_table_block(table["id"], headers, rows, col_letters,
                                 table["title"], table["unit"], sheet_title,
                                 notes, foot_lines)

    emitted_ids = {e["table"] for e in emitted}
    skipped_reason = {}
    for s in skipped:
        skipped_reason.setdefault(s["table"], s["reason"])
    missing_required = list(invalid_year_groups)
    for section in sections:
        for table in section["tables"]:
            if table.get("optional"):
                continue
            tid = table["id"]
            if tid in emitted_ids or any(
                    e.startswith(tid + "@") for e in emitted_ids):
                continue
            missing_required.append({
                "table": tid,
                "reason": skipped_reason.get(tid, "no_output"),
                "sheet": table["sheet"]})
    wb_values.close()
    return {
        "status": "generated",
        "layout": layout_id,
        "markdown": "\n".join(md).rstrip() + "\n",
        "emitted": emitted,
        "skipped": skipped,
        "missing_required": missing_required,
        "warnings": warnings,
    }


def run(base, xlsx, *, major_id=None, input_path=None, out="build/finance-tables.md"):
    """전공 가드 + 재무 프로필 게이트를 거쳐 표 묶음 Markdown을 쓴다.

    ``(value, exit_code)`` 반환 — ``gg.py finance-tables``과 이 모듈의
    ``main()``이 공유한다. ``base``는 작업 폴더(project.json이 있는
    루트), ``xlsx``는 재계산 완료 통합문서의 명시 경로. 출력 보류는
    ``OutputHeldError``를 올리지 않고 ``{"status": "held"}`` 값으로
    돌려준다."""
    try:
        ctx = mc.output_context(base, major_id)
        auth = mc.authorize_output(mc.OUTPUT_SCHOOL_PAPER, ctx)
        mc.require_finance_profile(auth, FINANCE_PROFILE)
    except mc.OutputHeldError as e:
        return {
            "status": "held",
            "reason": e.reason,
            "detail": str(e),
            "guidance": e.detail.get("guidance"),
        }, 2
    base_path = Path(base).resolve()
    out_path = (base_path / out).resolve()
    if not out_path.is_relative_to(base_path):
        raise ValueError("작업 폴더 밖 경로는 허용하지 않음: " + str(out))
    if out_path.exists():
        raise ValueError("기존 산출물을 덮어쓰지 않음: 새 경로 지정")
    footnotes = None
    if input_path:
        src = (base_path / input_path).resolve()
        if not src.is_relative_to(base_path):
            raise ValueError("작업 폴더 밖 경로는 허용하지 않음: "
                             + str(input_path))
        footnotes = json.loads(src.read_text(encoding="utf-8"))
        if isinstance(footnotes, dict) and "footnotes" in footnotes:
            footnotes = footnotes["footnotes"]
    result = generate(xlsx, major_id=auth.major_id, footnotes=footnotes)
    if result.get("status") == "unsupported_layout":
        failed = result.get("failed_signatures") or []
        return {
            "status": "held",
            "reason": "unsupported_workbook_layout",
            "detail": "학과 제공 통합문서 레이아웃만 지원한다 — "
                      "지문 불일치로 전개 보류: "
                      + "; ".join(
                          "시트 '%s'에 %s" % (f["sheet"], f.get("detail", ""))
                          for f in failed),
            "layout": result.get("layout"),
            "failed_signatures": failed,
        }, 2
    if result["missing_required"]:
        return {
            "status": "failed",
            "reason": "required_tables_missing",
            "detail": "필수 표·앵커 누락으로 발행하지 않음",
            "missing": result["missing_required"],
            "emitted": result["emitted"],
            "skipped": result["skipped"],
            "warnings": result["warnings"],
        }, 2
    import gg_document as gd
    issues = gd.check(
        result["markdown"], str(base_path), major_id=auth.major_id)
    errors = [list(i) for i in issues
              if i[0] not in gd.INFO_CHECKS | gd.WARNING_CHECKS]
    if errors:
        return {
            "status": "failed",
            "reason": "policy_gate_failed",
            "detail": "문서 정책 게이트(금칙어·출처 등) 실패로 발행하지 않음",
            "issues": errors,
            "emitted": result["emitted"],
            "skipped": result["skipped"],
            "warnings": result["warnings"],
        }, 2
    try:
        mc.reconfirm_output(auth, ctx)
    except mc.OutputHeldError as e:
        return {
            "status": "held",
            "reason": e.reason,
            "detail": str(e),
            "guidance": e.detail.get("guidance"),
        }, 2
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(result["markdown"], encoding="utf-8")
    return {
        "path": str(out_path),
        "status": "generated",
        "major_id": auth.to_dict()["major_id"],
        "tables_emitted": len(result["emitted"]),
        "emitted": result["emitted"],
        "skipped": result["skipped"],
        "warnings": result["warnings"],
    }, 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("base")
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--input")
    ap.add_argument("--out", default="build/finance-tables.md")
    ap.add_argument("--major")
    a = ap.parse_args(argv)
    try:
        value, code = run(a.base, a.xlsx, major_id=a.major,
                          input_path=a.input, out=a.out)
    except (ValueError, OSError, KeyError, TypeError) as e:
        print(json.dumps({"status": "blocked", "reason": str(e)},
                         ensure_ascii=False))
        return 2
    print(json.dumps(value, ensure_ascii=False, indent=2))
    return code


if __name__ == "__main__":
    sys.exit(main())
